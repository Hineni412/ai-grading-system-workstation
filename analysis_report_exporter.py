"""学生个人考试分析报告导出 + 班级分析页面数据装配。

产物形态与数据口径见已确认实施方案（.zcode/plans/考试分析报告_实施方案.md）：
- personal_analysis_html：每名正常参考学生一份自包含 HTML，打包 zip，附未生成清单；
- 班级分析（教师版）的内嵌页面与独立 HTML 共用分班数据装配、提示词与
  AI 叙述缓存，页面状态与生成 job 见 backend/class_analysis.py。

与计划的三处已确认偏差：
1. 解析失败不重发模型请求，由 LLMClient.json_from_text 本地确定性修复，
   仍失败则该报告降级为「无 AI 叙述版」（数据段照常渲染）；
2. 费用估算只给 token 粗估，不给金额；
3. AI 叙述按 (session_id, score_revision, rendition_version, 报告键) 缓存到
   受控目录，重新生成命中缓存不再调用模型；个人报告额外兼容读取
   backend.report_exports.LEGACY_PERSONAL_NARRATIVE_VERSIONS 中的旧版叙述。
"""

from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import math
import re
import statistics
import sqlite3
import zipfile
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from functools import lru_cache
from dataclasses import dataclass, field, replace
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from PIL import Image

from analysis_report_prompts import (
    CLASS_ALIAS_PREFIX,
    PERSONAL_STUDENT_ALIAS,
    PERSONAL_SYSTEM_PROMPT,
    personal_output_token_limit,
)
from answer_key_utils import answer_forms_map
from answer_region_geometry import (
    answer_regions_with_template_source_sizes,
    scaled_region_bbox,
)
from api_profiles import get_api_profile_store, resolve_profile_for_task
from backend.analytics.service import load_session_score_type_maps
from backend.llm.policy import policy_overrides_from_profile
from backend.llm.trace import safe_endpoint_host
from backend.repositories.access import (
    GradingRepositoryAccess,
    as_grading_repositories,
)
from backend.repositories.compat import open_grading_repositories
from export_names import safe_filename_fragment, session_export_path_name
from grading_completeness import resolve_grading_completeness
from llm_client import LLMSettings, normalize_openai_base_url
from path_manager import resolve_stored_file_path
from question_id_contract import (
    QuestionIdContractError,
    canonicalize_question_document,
    question_id_coordinates,
    resolve_known_question_id,
)
from backend.review.service import REVIEW_CONFIRMED_REASON
from report import _natural_question_order

PERSONAL_ANALYSIS_REPORT_TYPE = "personal_analysis_html"
ANALYSIS_REPORT_TYPES = (PERSONAL_ANALYSIS_REPORT_TYPE,)

# 方案 §7/§8 的固定标注文案。
AI_DISCLAIMER = "AI 分析 · 仅供参考"
AI_FAILED_NOTE = "AI 分析生成失败，可重新生成"
SMALL_SAMPLE_LIMIT = 10

# 题目截图同时用于图文分析与报告展示，保留手写公式的清晰度。
_SHOT_PADDING = 12
_SHOT_MAX_WIDTH = 1600
_SHOT_JPEG_QUALITY = 85
_SHOT_BUDGET_BYTES = 12 * 1024 * 1024

_QUESTION_TYPE_LABELS = {
    "choice": "选择",
    "single_choice": "选择",
    "multiple_choice": "选择",
    "fill_blank": "填空",
    "fill_in_blank": "填空",
    "calculation": "计算",
    "proof": "证明",
    "answer": "解答",
    "subjective": "解答",
    "comprehensive": "解答",
    "constructed_response": "解答",
    "judgement": "判断",
    "true_false": "判断",
    "direct_answer": "作答",
}

# 客观题批改记录里的内部调试码 → 家长可读文本；
# 口径与前端 frontend/src/utils/grading-reasons.ts 保持一致。
_OBJECTIVE_ANSWER_PREFIX = "objective_answer="
_REASON_CODE_LABELS = {
    "objective_api_disabled": "客观题识别未启用，已转教师复核",
    "objective_api_not_configured": "客观题识别模型配置不完整，已转教师复核",
    "objective_paper_model_failed": "客观题识别请求失败，已转教师复核",
    "objective_paper_region_failed": "无法读取选填题作答区域",
    "objective_region_not_found": "未找到此题的有效作答区域",
    "missing_question_result": "AI 未返回此题的识别结果",
    "missing_answer_field": "AI 未返回此题的作答内容",
    "duplicate_question_result": "AI 返回了重复的识别结果",
    "paper_key_mismatch": "识别结果与当前答卷不一致",
    "low_confidence": "作答辨识度较低，需要教师复核",
    "needs_review": "AI 建议教师复核",
    "objective_needs_review": "客观题识别结果需要教师复核",
    "objective_score_uncertain": "答案识别存在不确定性",
    "prompt_injection_or_score_bait": "作答区出现与答题无关的批改指令",
    "discarded_answer_only": "只识别到已经涂抹或作废的答案",
    "assignment_changed": "批改期间答卷匹配发生变化",
    "missing_detail_question_ids": "AI 未完整返回所有小问的评分结果",
    "duplicate_detail_question_id": "AI 重复返回了同一小问",
    "unexpected_detail_question_id": "AI 返回了不属于本题的小问",
    "no_numeric_value": "未能可靠识别填写的数值，需要教师确认",
    "multiple options selected": "识别到选择了多个选项",
    # 阅卷模型自由发挥的英文 error_summary 代码（历史数据实测值，见错因体系改造方案）。
    "blank_or_no_valid_work": "未识别到有效作答内容",
    "no_valid_answer": "未识别到有效作答内容",
    "insufficient_work_shown": "作答过程不完整",
    "incomplete_process": "作答过程不完整",
    "missing_steps": "作答缺少关键步骤",
    "unclear_or_missing_process": "作答过程不完整",
    "formatting_issue": "作答书写不规范",
    "not_simplified": "结果未化到最简",
    "review_required": "需要教师复核确认",
    "calculation_error": "计算有误",
    "sign_error": "符号处理有误",
    "incorrect_reasoning": "推理过程有误",
    "incorrect_reason": "所给理由有误",
    "concept_misunderstanding": "概念理解有误",
    "conceptual_error": "概念理解有误",
}

_CJK_RE = re.compile(r"[\u3400-\u9fff]")


def _has_cjk(text: str) -> bool:
    return bool(_CJK_RE.search(text))

# 分数段按满分等比缩放（满分 100 时即方案 §3 的固定分段）。
_BAND_CUTOFFS = (85, 70, 60, 40, 0)
_PASS_CUTOFF = 60


def resolve_content_generation_settings() -> LLMSettings | None:
    """解析设置页「内容生成」任务当前绑定的模型；api_key/model 缺失即未配置。"""
    profile = resolve_profile_for_task(get_api_profile_store(), "content_generation")
    api_key = str(profile.get("api_key") or "").strip()
    config_api_key = str(profile.get("config_api_key") or api_key).strip()
    config_model = str(profile.get("config_model") or "").strip()
    if not api_key or not config_api_key or not config_model:
        return None
    grading_model = str(profile.get("grading_model") or config_model)
    return LLMSettings(
        api_key=api_key,
        base_url=normalize_openai_base_url(
            str(profile.get("base_url") or "https://api.openai.com/v1")
        ),
        ocr_model=str(profile.get("ocr_model") or grading_model),
        grading_model=grading_model,
        config_model=config_model,
        config_api_key=config_api_key,
        config_base_url=normalize_openai_base_url(
            str(
                profile.get("config_base_url")
                or profile.get("base_url")
                or "https://api.openai.com/v1"
            )
        ),
        policy_profile=policy_overrides_from_profile(profile),
    )


def content_generation_public_info() -> tuple[str | None, str | None]:
    """返回可对外展示的（服务名, 模型名），不含密钥与完整路径。"""
    profile = resolve_profile_for_task(get_api_profile_store(), "content_generation")
    model = str(profile.get("config_model") or "").strip() or None
    profile_name = str(profile.get("name") or "").strip()
    host = safe_endpoint_host(
        profile.get("config_base_url") or profile.get("base_url")
    )
    service = " @ ".join(part for part in (profile_name, host) if part) or None
    return service, model


class AnalysisNarrativeCache:
    """AI 叙述 JSON 的本地缓存：命中即不再调用模型（重新生成幂等、不重复扣费）。"""

    def __init__(self, cache_dir: Path) -> None:
        self.cache_dir = Path(cache_dir)

    @staticmethod
    def cache_key(
        *,
        session_id: int,
        score_revision: str,
        rendition_version: str,
        report_key: str,
    ) -> str:
        material = json.dumps(
            {
                "session_id": int(session_id),
                "score_revision": str(score_revision),
                "rendition_version": str(rendition_version),
                "report_key": str(report_key),
            },
            ensure_ascii=False,
            sort_keys=True,
        )
        return hashlib.sha256(material.encode("utf-8")).hexdigest()

    def _path(self, key: str) -> Path:
        return self.cache_dir / f"{key}.json"

    def load(self, key: str) -> dict[str, Any] | None:
        try:
            payload = json.loads(self._path(key).read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            return None
        narrative = payload.get("narrative") if isinstance(payload, dict) else None
        return narrative if isinstance(narrative, dict) else None

    def store(self, key: str, narrative: dict[str, Any]) -> None:
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        target = self._path(key)
        temp = target.with_name(f".{target.stem}.tmp")
        try:
            temp.write_text(
                json.dumps(
                    {"version": 1, "narrative": narrative},
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            temp.replace(target)
        except OSError:
            temp.unlink(missing_ok=True)


def estimate_prompt_tokens(text: str) -> int:
    """token 粗估：中文场景约 1.5 字符/token，仅用于生成前的量级提示。"""
    return math.ceil(len(text) / 1.5)


@dataclass
class _QuestionInfo:
    question_id: str
    question_type: str
    max_score: float
    stem_summary: str
    canonical_answer: str
    attempts: int = 0
    score_sum: float = 0.0
    question_text: str = ""
    question_markup: str = ""
    reference_analysis: str = ""
    reference_images: list[tuple[str, bytes]] = field(default_factory=list)

    @property
    def class_rate(self) -> float | None:
        full = self.max_score * self.attempts
        if full <= 0:
            return None
        return self.score_sum / full

    @property
    def class_avg(self) -> float | None:
        if self.attempts <= 0:
            return None
        return self.score_sum / self.attempts


@dataclass
class _StudentQuestionRecord:
    question_id: str
    score: float
    max_score: float
    deduction_reason: str
    error_category: str
    error_summary: str
    secondary_errors: list[str] = field(default_factory=list)
    student_answer: str = ""
    evidence_steps: list[str] = field(default_factory=list)
    missing_steps: list[str] = field(default_factory=list)
    teacher_confirmed: bool = False
    teacher_comment: str = ""

    @property
    def lost(self) -> bool:
        return self.max_score > 0 and self.score < self.max_score - 1e-6

    @property
    def lost_points(self) -> float:
        return max(0.0, self.max_score - self.score) if self.lost else 0.0


@dataclass
class _StudentReportData:
    result_id: int
    student_id: int
    student_code: str
    student_name: str
    class_name: str
    student_score: float
    needs_review: bool
    graded_at: str
    rank: int = 0
    records: list[_StudentQuestionRecord] = field(default_factory=list)
    material_notes: list[str] = field(default_factory=list)
    knowledge_mastery: dict[str, dict[str, Any]] = field(default_factory=dict)

    @property
    def lost_points_total(self) -> float:
        return sum(record.lost_points for record in self.records)


@dataclass
class _SessionAnalysisData:
    session_id: int
    session_name: str
    subject: str
    full_score: float
    graded_at: str
    present: int
    roster_absent: int
    small_sample: bool
    questions: list[_QuestionInfo]
    students: list[_StudentReportData]
    skipped: list[dict[str, str]]
    stats: dict[str, Any]
    knowledge_backfill: dict[str, list[dict[str, str]]]
    attendance_by_class: dict[str, dict[str, int]] = field(default_factory=dict)
    class_name: str | None = None
    rubric: dict[str, Any] = field(default_factory=dict)
    question_assessments: dict[str, dict[str, Any]] = field(default_factory=dict)
    knowledge_structure: dict[str, Any] = field(default_factory=dict)


def assemble_session_analysis(
    db: GradingRepositoryAccess | Any,
    session_id: int,
    *,
    data_root: Path | None = None,
    page_only: bool = False,
    include_answer_evidence: bool = False,
    include_knowledge: bool = False,
) -> _SessionAnalysisData:
    """轻量页面按需读取标签或作答文字，不重新识别原卷。"""
    repositories = as_grading_repositories(db)
    resolved_data_root = data_root or _infer_data_root(repositories.db_path)
    snapshot = repositories.reports.get_session_report_snapshot(
        int(session_id),
        question_bank_path=None if page_only and not include_knowledge else _question_bank_db_path(repositories.db_path),
    )
    session_row = repositories.sessions.get_grading_session(int(session_id)) or {}
    rubric = _load_rubric(session_row, resolved_data_root)
    score_map, type_map = load_session_score_type_maps(
        session_row or None,
        data_root=resolved_data_root,
    )
    stem_map = _stem_summary_map(rubric)
    answer_map = _canonical_answer_map(session_row, resolved_data_root)

    attendance_by_identity = {
        (
            str(row.get("student_code") or ""),
            str(row.get("class_name") or "未分班"),
        ): str(row.get("attendance_status") or "").strip()
        for row in snapshot.attendance
    }
    present_count = sum(
        1
        for row in snapshot.attendance
        if str(row.get("attendance_status") or "").strip() == "present"
    )
    roster_absent = sum(
        1
        for row in snapshot.attendance
        if str(row.get("attendance_status") or "").strip() in {"absent", "scan_issue"}
    )

    details_by_result: dict[int, list[dict[str, Any]]] = {}
    for detail in snapshot.details:
        details_by_result.setdefault(int(detail.get("result_id") or 0), []).append(
            detail
        )

    students: list[_StudentReportData] = []
    skipped: list[dict[str, str]] = []
    covered_identities: set[tuple[str, str]] = set()
    for result in snapshot.results:
        identity = (
            str(result.get("student_code") or ""),
            str(result.get("class_name") or "未分班"),
        )
        covered_identities.add(identity)
        label = {
            "class_name": identity[1],
            "student_code": identity[0],
            "student_name": str(result.get("student_name") or ""),
        }
        attendance_status = attendance_by_identity.get(identity)
        if attendance_status in {"absent", "scan_issue"}:
            skipped.append(
                {**label, "reason": "缺考" if attendance_status == "absent" else "扫描异常"}
            )
            continue
        result_id = int(result.get("result_id") or 0)
        completeness = resolve_grading_completeness(
            result.get("raw_json"),
            rubric=rubric if isinstance(rubric, dict) else None,
            details=details_by_result.get(result_id, []),
        )
        status = (
            str(completeness.get("status") or "").strip()
            if isinstance(completeness, dict)
            else ""
        )
        if status != "complete":
            missing = (
                completeness.get("missing_question_ids")
                if isinstance(completeness, dict)
                else []
            )
            missing_text = (
                "、".join(str(item) for item in missing if str(item).strip())
                if isinstance(missing, list)
                else ""
            )
            skipped.append(
                {
                    **label,
                    "reason": (
                        f"批改不完整（缺失题目：{missing_text}）"
                        if missing_text
                        else "批改结果无效或不完整"
                    ),
                }
            )
            continue
        students.append(
            _StudentReportData(
                result_id=result_id,
                student_id=int(result.get("student_id") or 0),
                student_code=identity[0],
                student_name=str(result.get("student_name") or ""),
                class_name=identity[1],
                student_score=float(result.get("student_score") or 0),
                needs_review=bool(result.get("needs_human_review")),
                graded_at=str(result.get("graded_at") or "")[:10],
                records=_merge_student_records(
                    repositories,
                    result_id,
                    details_by_result.get(result_id, []),
                    score_map,
                    student_answers=_student_answer_map(result.get("raw_json")) if not page_only or include_answer_evidence else None,
                    grading_evidence=_grading_detail_map(result.get("raw_json")) if not page_only or include_answer_evidence else None,
                    include_secondary_errors=not page_only,
                ),
            )
        )
    locks_by_student_question = {
        (
            int(lock["student_id"]),
            resolve_known_question_id(str(lock["question_id"]), score_map)
            or str(lock["question_id"]),
        ): lock
        for lock in snapshot.locks
    }
    for student in students:
        for record in student.records:
            lock = locks_by_student_question.get((student.student_id, record.question_id))
            if lock is not None:
                record.teacher_confirmed = True
                record.teacher_comment = _sanitize_grading_text(lock.get("deduction_reason"))
        # Match the review queue's per-item rules; the original whole-paper AI
        # flag remains historical after a teacher confirms a flagged answer.
        from backend.review.service import _is_substantive_review_reason
        from grading_completeness import is_objective_detail
        pending_review = False
        for detail in details_by_result.get(student.result_id, []):
            qid = resolve_known_question_id(str(detail.get("question_id") or ""), score_map) or str(detail.get("question_id") or "")
            lock = locks_by_student_question.get((student.student_id, qid))
            if lock is not None:
                pending_review |= abs(float(lock.get("max_score") or 0) - float(score_map.get(qid) or 0)) > 1e-6
            else:
                pending_review |= _is_substantive_review_reason(
                    str(detail.get("deduction_reason") or ""),
                    str(detail.get("error_category") or ""),
                    detail.get("confidence_score"),
                    objective=is_objective_detail(detail),
                )
        student.needs_review = pending_review
    # 名册中缺考且无结果的学生也要进入未生成清单。
    attendance_labels = {"absent": "缺考", "scan_issue": "扫描异常"}
    for row in snapshot.attendance:
        status = str(row.get("attendance_status") or "").strip()
        identity = (
            str(row.get("student_code") or ""),
            str(row.get("class_name") or "未分班"),
        )
        if status in attendance_labels and identity not in covered_identities:
            skipped.append(
                {
                    "class_name": identity[1],
                    "student_code": identity[0],
                    "student_name": str(row.get("student_name") or ""),
                    "reason": attendance_labels[status],
                }
            )

    # 名次：按总分降序，同分同名次（method="min"）。
    ordered = sorted(
        students,
        key=lambda item: (-item.student_score, item.student_code, item.student_name),
    )
    previous_score: float | None = None
    previous_rank = 0
    for index, student in enumerate(ordered, start=1):
        if previous_score is None or student.student_score < previous_score - 1e-9:
            previous_rank = index
            previous_score = student.student_score
        student.rank = previous_rank
    students = ordered

    questions: list[_QuestionInfo] = []
    for qid in _natural_question_order([str(qid) for qid in score_map]):
        questions.append(
            _QuestionInfo(
                question_id=qid,
                question_type=str(type_map.get(qid) or ""),
                max_score=float(score_map.get(qid) or 0),
                stem_summary=stem_map.get(qid) or "",
                canonical_answer=answer_map.get(qid) or "",
            )
        )
    info_by_qid = {info.question_id: info for info in questions}
    for student in students:
        for record in student.records:
            info = info_by_qid.get(record.question_id)
            if info is None:
                continue
            info.attempts += 1
            info.score_sum += min(record.score, record.max_score or record.score)

    # 大题满分与实际评分的小问同时存在时，只展示真正参与评分的行。
    scored_parents = {
        _parent_question_id(info.question_id)
        for info in questions
        if info.attempts and _parent_question_id(info.question_id) != info.question_id
    }
    questions = [
        info for info in questions
        if info.attempts or info.question_id not in scored_parents
    ]
    attendance_by_class: dict[str, dict[str, int]] = {}
    for row in snapshot.attendance:
        counts = attendance_by_class.setdefault(
            str(row.get("class_name") or "未分班"), {"present": 0, "absent": 0},
        )
        status = str(row.get("attendance_status") or "").strip()
        counts["present"] += int(status == "present")
        counts["absent"] += int(status in {"absent", "scan_issue"})

    scores = [student.student_score for student in students]
    stats = _score_distribution(scores, _full_score(rubric, score_map, snapshot))
    session_name = (
        str(session_row.get("session_name") or "").strip() or f"考试批改_{session_id}"
    )
    present = present_count if snapshot.attendance else len(students)
    return _SessionAnalysisData(
        session_id=int(session_id),
        session_name=session_name,
        subject=str(rubric.get("subject") or "数学").strip() or "数学",
        full_score=_full_score(rubric, score_map, snapshot),
        graded_at=max((student.graded_at for student in students), default=""),
        present=present,
        roster_absent=roster_absent,
        small_sample=present < SMALL_SAMPLE_LIMIT,
        questions=questions,
        students=students,
        skipped=skipped,
        stats=stats,
        knowledge_backfill=snapshot.knowledge_backfill,
        question_assessments=snapshot.question_assessments,
        attendance_by_class=attendance_by_class,
        rubric=rubric,
    )


def split_session_analysis_by_class(
    data: _SessionAnalysisData,
) -> dict[str, _SessionAnalysisData]:
    """个人报告、班级页面与班级导出共用的本班统计和同分名次。"""
    names = {student.class_name for student in data.students}
    names.update(item["class_name"] for item in data.skipped)
    names.update(data.attendance_by_class)
    groups: dict[str, _SessionAnalysisData] = {}
    for name in sorted(names, key=lambda value: [
        (0, int(part)) if part.isdigit() else (1, part)
        for part in re.split(r"(\d+)", value)
    ]):
        students = sorted(
            (replace(student, material_notes=list(student.material_notes))
             for student in data.students if student.class_name == name),
            key=lambda student: (-student.student_score, student.student_code, student.student_name),
        )
        previous_score = None
        rank = 0
        for index, student in enumerate(students, 1):
            if previous_score is None or student.student_score < previous_score - 1e-9:
                rank, previous_score = index, student.student_score
            student.rank = rank
        questions = []
        for question in data.questions:
            records = [record for student in students for record in student.records
                       if record.question_id == question.question_id]
            questions.append(replace(
                question, attempts=len(records),
                score_sum=sum(min(record.score, record.max_score or record.score) for record in records),
            ))
        skipped = [item for item in data.skipped if item["class_name"] == name]
        counts = data.attendance_by_class.get(name)
        present = counts["present"] if counts else len(students)
        absent = counts["absent"] if counts else sum(item["reason"] in {"缺考", "扫描异常"} for item in skipped)
        groups[name] = replace(
            data, class_name=name, students=students, questions=questions, skipped=skipped,
            present=present, roster_absent=absent, small_sample=present < SMALL_SAMPLE_LIMIT,
            stats=_score_distribution([student.student_score for student in students], data.full_score),
            graded_at=max((student.graded_at for student in students), default=""),
            attendance_by_class={name: counts} if counts else {},
        )
    return groups


def _infer_data_root(db_path: Path) -> Path | None:
    path = Path(db_path)
    return path.parent.parent if path.parent.name == "databases" else None


def _question_bank_db_path(db_path: Path) -> Path | None:
    path = Path(db_path)
    data_root = path.parent.parent if path.parent.name == "databases" else path.parent
    candidate = data_root / "databases" / "question_bank.db"
    return candidate if candidate.is_file() else None


def _load_rubric(session_row: dict[str, Any], data_root: Path | None) -> dict[str, Any]:
    rubric_path = resolve_stored_file_path(
        session_row.get("rubric_path"),
        data_root=data_root,
    )
    try:
        rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
        if not isinstance(rubric, dict):
            return {}
        return canonicalize_question_document(rubric)
    except (
        OSError,
        UnicodeDecodeError,
        json.JSONDecodeError,
        QuestionIdContractError,
    ):
        return {}


def _stem_summary_map(rubric: dict[str, Any]) -> dict[str, str]:
    stems: dict[str, str] = {}
    questions = rubric.get("questions") if isinstance(rubric, dict) else []
    if not isinstance(questions, list):
        return stems
    for question in questions:
        if not isinstance(question, dict):
            continue
        question_id = str(question.get("question_id") or "").strip()
        stem = str(question.get("stem_summary") or "").strip()
        if question_id and stem:
            stems[question_id] = stem
        parts = question.get("parts")
        if not isinstance(parts, list):
            continue
        for part in parts:
            if not isinstance(part, dict):
                continue
            part_id = str(part.get("part_id") or "").strip()
            part_stem = str(part.get("stem_summary") or "").strip()
            # 小问缺题干摘要时回退大问摘要，仍没有则渲染层用题号替代（方案 §9）。
            if part_id and (part_stem or stem):
                stems[part_id] = part_stem or stem
    return stems


# 标准答案展示收敛参数：判分用的等价变体常含标点/单位差异的重复全文，
# 报告展示前做规范化去重、限量、限长。
_ANSWER_FORM_MAX_LEN = 120
_ANSWER_DISPLAY_MAX_FORMS = 3
_ANSWER_DISPLAY_MAX_CHARS = 400

_ANSWER_NORMALIZE_PAIRS = (
    ("：", ":"),
    ("，", ","),
    ("。", "."),
    ("（", "("),
    ("）", ")"),
    ("＝", "="),
    ("；", ";"),
    ("°", "度"),
)


def _normalize_answer_form(text: str) -> str:
    normalized = re.sub(r"\s+", "", str(text))
    for full, half in _ANSWER_NORMALIZE_PAIRS:
        normalized = normalized.replace(full, half)
    return normalized.lower()


def _display_answer_text(forms: list[str]) -> str:
    """把判分用的全部等价变体收敛为家长可读的标准答案文本。"""
    seen: set[str] = set()
    unique: list[str] = []
    for form in forms:
        text = str(form or "").strip()
        key = _normalize_answer_form(text)
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(text)
    if not unique:
        return ""
    # 长文本（证明过程等）只保留代表形式；短答案最多并列 3 个。
    if any(len(text) > _ANSWER_FORM_MAX_LEN for text in unique):
        unique = unique[:1]
    else:
        unique = unique[:_ANSWER_DISPLAY_MAX_FORMS]
    display = " 或 ".join(unique)
    if len(display) > _ANSWER_DISPLAY_MAX_CHARS:
        display = display[:_ANSWER_DISPLAY_MAX_CHARS].rstrip() + " ……"
    return display


def _canonical_answer_map(
    session_row: dict[str, Any],
    data_root: Path | None,
) -> dict[str, str]:
    answer_path = resolve_stored_file_path(
        session_row.get("answer_key_path"),
        data_root=data_root,
    )
    try:
        answer_key = json.loads(answer_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return {
        qid: display
        for qid, forms in answer_forms_map(answer_key).items()
        if (display := _display_answer_text(forms))
    }


def _full_score(
    rubric: dict[str, Any],
    score_map: dict[str, float],
    snapshot: Any,
) -> float:
    try:
        total = float(rubric.get("total_score") or 0)
    except (TypeError, ValueError):
        total = 0.0
    if total > 0:
        return total
    # 只加总大题满分，避免小问与大问重复计数。
    parent_total = 0.0
    for qid, max_score in score_map.items():
        coordinates = question_id_coordinates(qid)
        if coordinates is not None and coordinates[1] is not None:
            continue
        parent_total += float(max_score or 0)
    if parent_total > 0:
        return parent_total
    return max(
        (float(result.get("total_score") or 0) for result in snapshot.results),
        default=0.0,
    )


def _score_distribution(scores: list[float], full_score: float) -> dict[str, Any]:
    if not scores:
        return {
            "avg": None,
            "median": None,
            "max": None,
            "min": None,
            "pass_rate": None,
            "excellent_count": 0,
            "bands": [],
        }
    scale = full_score / 100 if full_score > 0 else 1.0
    bands: list[dict[str, Any]] = []
    for index, cutoff in enumerate(_BAND_CUTOFFS):
        lower = cutoff * scale
        upper = full_score if index == 0 else _BAND_CUTOFFS[index - 1] * scale
        members = [
            score for score in scores
            if lower <= score and (score <= upper if index == 0 else score < upper)
        ]
        bands.append(
            {
                "label": f"{_fmt_num(lower)} – {'不足' if index else ''}{_fmt_num(upper)} 分",
                "count": len(members),
                "ratio": len(members) / len(scores),
            }
        )
    passing = [score for score in scores if score >= _PASS_CUTOFF * scale]
    return {
        "avg": round(sum(scores) / len(scores), 2),
        "median": round(statistics.median(scores), 2),
        "max": max(scores),
        "min": min(scores),
        "pass_rate": round(len(passing) / len(scores), 4),
        "excellent_count": len(
            [score for score in scores if score >= _BAND_CUTOFFS[0] * scale]
        ),
        "bands": bands,
    }


def _sanitize_grading_text(value: object) -> str:
    """批改记录文本转家长可读口径：内部调试码翻译，无中文的文本丢弃。

    阅卷模型偶发用英文写扣分原因；报告面向家长，没有中文信息的文本
    不如不显示（其余中文字段仍可说明情况）。
    """
    text = str(value or "").strip()
    if not text:
        return ""
    lowered = text.lower()
    if lowered.startswith(_OBJECTIVE_ANSWER_PREFIX):
        answer = text[len(_OBJECTIVE_ANSWER_PREFIX):].strip()
        return f"作答识别为「{answer}」，与参考答案不符" if answer else ""
    direct = _REASON_CODE_LABELS.get(lowered)
    if direct is not None:
        return direct
    for code, label in _REASON_CODE_LABELS.items():
        if code in lowered:
            return label
    return text if _has_cjk(text) else ""


def _format_secondary_error(item: object) -> str:
    """次要错因可能是 dict（category/summary/evidence），格式化为可读文本。"""
    if isinstance(item, dict):
        summary = _sanitize_grading_text(item.get("summary"))
        category = _sanitize_grading_text(item.get("category"))
        evidence = str(item.get("evidence") or "").strip()
        text = summary or category
        if evidence and evidence not in text:
            text = f"{text}（{evidence}）" if text else evidence
        return text
    return _sanitize_grading_text(item)


# 学生作答识别文本在 raw_json grading_details 里的兼容键（与 ai_grader 一致）。
_STUDENT_ANSWER_KEYS = ("observed_answer", "student_answer", "answer_observed", "raw_answer", "recognized_answer")
_STUDENT_ANSWER_MAX_LEN = 2000


def _grading_detail_map(raw_json: Any) -> dict[str, dict[str, Any]]:
    """兼容整卷与混合批改保留的作答证据，不根据分数推断空白。"""
    parsed = raw_json
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except json.JSONDecodeError:
            return {}
    if not isinstance(parsed, dict):
        return {}
    result: dict[str, dict[str, Any]] = {}
    metadata = parsed.get("detail_metadata")
    if isinstance(metadata, dict):
        for qid, item in metadata.items():
            if isinstance(item, dict):
                result[str(qid)] = dict(item)
    details = parsed.get("grading_details")
    for item in details if isinstance(details, list) else []:
        if isinstance(item, dict) and (qid := str(item.get("question_id") or "").strip()):
            result.setdefault(qid, {}).update(item)
    return result


def _student_answer_map(raw_json: Any) -> dict[str, str]:
    """从批改结果提取作答文字；图文报告同时使用原卷核对识别结果。"""
    answers: dict[str, str] = {}
    for qid, item in _grading_detail_map(raw_json).items():
        text = ""
        for key in _STUDENT_ANSWER_KEYS:
            value = item.get(key)
            if value:
                text = str(value).strip()
                break
        if text:
            # 兼容旧批改记录把同一段 observed_answer / evidence_steps 拼接两次。
            for evidence in _evidence_texts(item.get("evidence_steps")):
                doubled = f"{evidence} {evidence}"
                if text == doubled:
                    text = evidence
                    break
            answers[qid] = text[:_STUDENT_ANSWER_MAX_LEN]
    return answers


def _merge_student_records(
    repositories: GradingRepositoryAccess,
    result_id: int,
    details: list[dict[str, Any]],
    score_map: dict[str, float],
    student_answers: dict[str, str] | None = None,
    grading_evidence: dict[str, dict[str, Any]] | None = None,
    include_secondary_errors: bool = True,
) -> list[_StudentQuestionRecord]:
    """把同一学生同一题的多条明细合并为一条（兼容小问拆行存储）。"""
    secondary_map: dict[str, list[str]] = {}
    try:
        for row in repositories.results.get_result_details(result_id) if include_secondary_errors else []:
            qid = str(row.get("question_id") or "").strip()
            canonical = resolve_known_question_id(qid, score_map) or qid
            errors = row.get("secondary_errors")
            if canonical and isinstance(errors, list):
                merged = secondary_map.setdefault(canonical, [])
                for item in errors:
                    text = _format_secondary_error(item)
                    if text and text not in merged:
                        merged.append(text)
    except Exception:
        secondary_map = {}

    merged: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for detail in details:
        raw_qid = str(detail.get("question_id") or "").strip()
        if not raw_qid:
            continue
        qid = resolve_known_question_id(raw_qid, score_map) or raw_qid
        if qid not in merged:
            merged[qid] = {
                "score": 0.0,
                "deduction_reasons": [],
                "error_categories": [],
                "error_summaries": [],
                "objective_observed": "",
            }
            order.append(qid)
        bucket = merged[qid]
        bucket["score"] += float(detail.get("score_awarded") or 0)
        for field_name, bucket_key in (
            ("deduction_reason", "deduction_reasons"),
            ("error_category", "error_categories"),
            ("error_summary", "error_summaries"),
        ):
            raw_text = str(detail.get(field_name) or "").strip()
            # 客观题的作答识别结果只存在于调试串里，先取出再清洗。
            if (
                field_name == "deduction_reason"
                and raw_text.lower().startswith(_OBJECTIVE_ANSWER_PREFIX)
                and not bucket["objective_observed"]
            ):
                bucket["objective_observed"] = raw_text.split("=", 1)[1].strip()
            text = _sanitize_grading_text(raw_text)
            if text and text not in bucket[bucket_key]:
                bucket[bucket_key].append(text)

    answers = {
        resolve_known_question_id(qid, score_map) or qid: text
        for qid, text in (student_answers or {}).items()
    }
    evidence_by_qid = {
        resolve_known_question_id(qid, score_map) or qid: item
        for qid, item in (grading_evidence or {}).items()
    }
    records: list[_StudentQuestionRecord] = []
    for qid in order:
        bucket = merged[qid]
        max_score = float(score_map.get(qid) or 0)
        score = min(bucket["score"], max_score) if max_score > 0 else bucket["score"]
        evidence = evidence_by_qid.get(qid, {})
        records.append(
            _StudentQuestionRecord(
                question_id=qid,
                score=score,
                max_score=max_score,
                deduction_reason="；".join(bucket["deduction_reasons"]),
                error_category="；".join(bucket["error_categories"]),
                error_summary="；".join(bucket["error_summaries"]),
                secondary_errors=secondary_map.get(qid, []),
                student_answer=(
                    answers.get(qid) or bucket["objective_observed"] or ""
                ),
                evidence_steps=_evidence_texts(evidence.get("evidence_steps")),
                missing_steps=_evidence_texts(evidence.get("missing_steps")),
            )
        )
    order_index = {
        qid: index for index, qid in enumerate(_natural_question_order(order))
    }
    records.sort(
        key=lambda record: order_index.get(record.question_id, len(order_index))
    )
    return records


def _docx_question_figures(source: Any) -> dict[str, list[bytes]]:
    """兼容原题重复出现在解析页的 Word：只取题干之后、解析开始之前的图。"""
    if getattr(source, "suffix", None) != ".docx":
        return {}
    from docx import Document
    from backend.document_parsing.question_blocks import (
        _extract_question_marker_number, _looks_like_answer_section_heading,
    )
    document = Document(io.BytesIO(source.private_source_bytes))
    stems = {
        str(item.get("question_id")): re.sub(r"\s+", "", str(item.get("question_text") or item.get("text") or ""))[:18]
        for item in source.private_blocks if isinstance(item, dict)
    }
    images: dict[str, list[bytes]] = {}
    current = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        number = _extract_question_marker_number(text)
        if number is not None:
            qid = f"Q{number}"
            stem = stems.get(qid)
            # 已从前面的题干取得图时，不再重复补入解析页中的题干副本。
            current = qid if qid not in images and stem and stem in re.sub(r"\s+", "", text) else None
        elif _looks_like_answer_section_heading(text) or re.match(
            r"^(?:【|\[)?(?:分析|解答|解析|答案|点评|解：|证明：)", text,
        ):
            current = None
        if current is None:
            continue
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = document.part.related_parts.get(relationship)
            if part is None:
                continue
            blob = part.blob
            try:
                with Image.open(io.BytesIO(blob)) as figure:
                    figure.verify()
            except (OSError, ValueError):
                continue
            bucket = images.setdefault(current, [])
            if blob not in bucket:
                bucket.append(blob)
    return images


def _enrich_personal_questions(
    repositories: GradingRepositoryAccess,
    data: _SessionAnalysisData,
    data_root: Path | None,
    *,
    include_images: bool = True,
) -> None:
    """读取本场评分依据和已绑定原题；考试错因整理仅复用文字投影。"""
    from backend.config_workspace.sources import ConfigSourceService
    from backend.document_parsing.question_blocks import rich_text_for_model

    session = repositories.sessions.get_grading_session(data.session_id) or {}
    root = data_root or _infer_data_root(repositories.db_path)
    rubric = _load_rubric(session, root)
    try:
        answer_key = json.loads(resolve_stored_file_path(
            session.get("answer_key_path"), data_root=root,
        ).read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        answer_key = {}

    def question_map(document: Any) -> dict[str, dict[str, Any]]:
        result: dict[str, dict[str, Any]] = {}
        for item in document.get("questions", []) if isinstance(document, dict) else []:
            if not isinstance(item, dict):
                continue
            qid = str(item.get("question_id") or "")
            result[qid] = item
            for part in item.get("parts", []) if isinstance(item.get("parts"), list) else []:
                if isinstance(part, dict):
                    raw_id = str(part.get("part_id") or part.get("question_id") or "")
                    canonical = resolve_known_question_id(raw_id, known_ids) or raw_id
                    result[canonical] = {**item, **part} if canonical == qid else part
        return result

    known_ids = {info.question_id: info.max_score for info in data.questions}
    rubric_items = question_map(rubric)
    answer_items = question_map(answer_key)
    source = None
    # Active uploads may be uncommitted candidates.  Only use the source whose
    # existing publication identity matches this exam's bound original paper.
    bound_source = str(session.get("source_paper_sha256") or "").strip()
    if bound_source and root is not None:
        try:
            candidate = ConfigSourceService(root / "config" / "uploaded").load_active_record(
                session_id=data.session_id,
            )
            if candidate.sha256 == bound_source:
                source = candidate
        except (OSError, ValueError, RuntimeError):
            pass
    source_items = {
        str(item.get("question_id") or ""): item
        for item in (source.private_blocks if source is not None else [])
        if isinstance(item, dict)
    }
    source_figures = _docx_question_figures(source) if source is not None and include_images else {}
    # Use the native formula alongside the exact importer's HTML projection.
    # This preserves root indices and mixed fractions that plain text loses.
    source_formulas = _source_formula_markup(source) if source is not None else {}
    formula_pattern = re.compile('|'.join(re.escape(key) for key in sorted(source_formulas, key=len, reverse=True))) if source_formulas else None

    def texts(item: dict[str, Any], fields: tuple[str, ...]) -> str:
        return "\n".join(dict.fromkeys(
            text for key in fields
            if (text := rich_text_for_model(item.get(key) or ""))
        ))

    for info in data.questions:
        parent = _parent_question_id(info.question_id)
        original = source_items.get(parent, {})
        rubric_item = rubric_items.get(info.question_id, rubric_items.get(parent, {}))
        answer_item = answer_items.get(info.question_id, answer_items.get(parent, {}))
        # These fields are alternative representations of the same stem.
        stem = next((value for key in ("question_html", "question_text", "text")
                     if (value := texts(original, (key,)))), "")
        if not stem:
            stem = texts(rubric_items.get(parent, {}), ("question_text", "text", "stem", "question_html"))
        options = rubric_items.get(parent, {}).get("options")
        if options:
            option_text = json.dumps(options, ensure_ascii=False) if isinstance(options, (dict, list)) else str(options)
            stem = "\n".join(filter(None, (stem, rich_text_for_model(option_text))))
        part_text = texts(rubric_item, ("question_text", "text", "stem_summary"))
        if parent != info.question_id and part_text and part_text not in stem:
            stem = "\n".join(filter(None, (stem, f"本小问：{part_text}")))
        info.question_text = stem
        info.question_markup = str(original.get("question_html") or "")
        if formula_pattern and info.question_markup and 'data-latex=' not in info.question_markup:
            info.question_markup = formula_pattern.sub(lambda match: source_formulas[match.group(0)], info.question_markup)
        info.reference_analysis = texts(answer_item, ("analysis", "full_answer", "explanation"))
        if not info.reference_analysis:
            info.reference_analysis = texts(original, ("analysis_html", "analysis", "answer_html", "answer_text"))
        if source is not None and include_images:
            assets = source.private_question_images.get(parent, {})
            for role in ("question", "answer"):
                encoded = assets.get(role)
                values = [encoded] if isinstance(encoded, str) else encoded if isinstance(encoded, list) else []
                for value in values:
                    try:
                        blob = base64.b64decode(value, validate=True)
                        if blob:
                            info.reference_images.append((role, blob))
                    except (ValueError, TypeError):
                        continue
            if not any(role == "question" for role, _blob in info.reference_images):
                info.reference_images.extend(("question", blob) for blob in source_figures.get(parent, []))


def _source_formula_markup(source: Any) -> dict[str, str]:
    if getattr(source, 'suffix', None) != '.docx':
        return {}
    from docx import Document
    from question_bank.importers.docx_importer import _math_text
    from question_bank.services.inline_math import omml_parts
    try:
        document = Document(io.BytesIO(source.private_source_bytes))
        variants: dict[str, set[tuple[str, str]]] = {}
        for node in document.element.xpath('.//m:oMath'):
            formula = omml_parts(node)
            original = _math_text(node)
            if formula and original and (len(original) > 2 or any(c in original for c in '√∛')):
                variants.setdefault(original, set()).add(formula)
        return {original: _report_math_span(*next(iter(values)))
                for original, values in variants.items() if len(values) == 1}
    except (OSError, ValueError, zipfile.BadZipFile):
        return {}


def _enrich_personal_knowledge(
    repositories: GradingRepositoryAccess, data: _SessionAnalysisData,
    data_root: Path | None, *, student_ids: set[int] | None = None,
) -> None:
    """Freeze the existing semester diagnosis once for the whole export batch."""
    from integration.diagnosis_profile_service import DiagnosisProfileService
    from question_bank.current_knowledge import CurrentKnowledgeResolver, CurrentKnowledgeUnavailable
    from question_bank.solution_evidence.part_assessments import reading
    root = data_root or _infer_data_root(repositories.db_path)
    path = root / 'databases' / 'question_bank.db'
    snapshot = {'as_of': datetime.now().strftime('%Y-%m-%d %H:%M'), 'catalog': [], 'associations': [],
                'note': '当前掌握度暂不可用；保留本卷考查范围与得分，不以得分率代替掌握度。'}
    data.knowledge_structure = snapshot
    selected = [student for student in data.students if student_ids is None or student.student_id in student_ids]
    for student in selected:
        student.knowledge_mastery = {}
    if not path.is_file() or not selected:
        return
    session = repositories.sessions.get_grading_session(data.session_id) or {}
    volume = str(session.get('curriculum_volume_id') or '').strip()
    if not volume:
        snapshot['note'] = '这场考试尚未关联教学学期，暂不展示当前掌握度；下方保留本卷考查范围与得分。'
        return
    try:
        resolver = CurrentKnowledgeResolver.from_active_database(path)
        # Exact current identities only; never join two nodes by a short label.
        for entries in data.knowledge_backfill.values():
            for entry in entries:
                identities = resolver.resolve(entry.get('stable_key') or entry.get('path') or '')
                if len(identities) == 1:
                    entry['stable_key'] = identities[0].stable_key
        with reading(path) as connection:
            service = DiagnosisProfileService(repositories.db_path, path, grading_db=repositories,
                                              question_bank_connection=connection, data_root=root)
            profile = service.build_tag_profiles(
                scope={'mode': 'selected', 'student_ids': [str(student.student_id) for student in selected], 'use_historical_fallback': False},
                exam_scope={'mode': 'semester', 'curriculum_volume_id': volume},
            )
        snapshot.update(catalog=profile.get('knowledge_catalog') or [],
                        associations=profile.get('knowledge_associations') or [])
        profiles = {str(item['student_id']): item for item in profile.get('students', [])}
        for student in selected:
            student.knowledge_mastery = {
                str(item['knowledge_key']): item
                for item in profiles.get(str(student.student_id), {}).get('weak_points', [])
                # A failed mastery calculation can leave exam-only diagnostic
                # rates. Those are not current mastery and must not be relabeled.
                if 'effective_weight' in item and 'direct_evidence_count' in item
            }
        if snapshot['catalog']:
            snapshot['note'] = '仅展开本卷涉及的知识与技能。当前掌握度综合本学期考试、训练、题目难度与时间；本次考试得分单独列出。'
    except (CurrentKnowledgeUnavailable, OSError, sqlite3.Error, ValueError, TypeError):
        # A missing graph must never prevent delivery of the scored report.
        return


def _evidence_texts(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if isinstance(item, str) and item.strip()]


def _parent_question_id(question_id: str) -> str:
    coordinates = question_id_coordinates(question_id)
    if coordinates is None:
        return str(question_id or "").strip()
    return f"Q{coordinates[0]}"


def _fmt_num(value: object) -> str:
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return str(value or "")
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


def _question_type_label(question_type: str) -> str:
    text = str(question_type or "").strip()
    if not text:
        return ""
    label = _QUESTION_TYPE_LABELS.get(text)
    if label is not None:
        return label
    # 未收录的英文类型码不原样渲染给家长；已是中文的兼容输入照常显示。
    return text if _has_cjk(text) else ""


def _question_display_label(question_id: str) -> str:
    """Q9 → 第9题；Q10(P1) → 第10(1)题。"""
    coordinates = question_id_coordinates(question_id)
    if coordinates is None:
        return str(question_id or "")
    parent, part = coordinates
    if part is None:
        return f"第{parent}题"
    return f"第{parent}({part})题"


def _esc(value: object) -> str:
    return html.escape(str(value or ""), quote=True)


# ---------------------------------------------------------------------------
# 模型 payload（方案 §6.2 / §6.4；学生一律用代号，不含姓名与磁盘路径）
# ---------------------------------------------------------------------------


def build_personal_payload(
    data: _SessionAnalysisData,
    student: _StudentReportData,
) -> dict[str, Any]:
    info_by_qid = {info.question_id: info for info in data.questions}
    questions: list[dict[str, Any]] = []
    for record in student.records:
        info = info_by_qid.get(record.question_id)
        class_rate = info.class_rate if info is not None else None
        item: dict[str, Any] = {
            "question_id": record.question_id,
            "display_label": _question_display_label(record.question_id),
            "type": record_type_label(record, info),
            "max_score": record.max_score,
            "score": record.score,
            "class_rate": round(class_rate, 4) if class_rate is not None else None,
            "lost": record.lost,
            "direct_knowledge": data.knowledge_backfill.get(record.question_id, data.knowledge_backfill.get(_parent_question_id(record.question_id), [])),
            "part_assessment": data.question_assessments.get(record.question_id),
        }
        if record.lost:
            item["stem_summary"] = (
                info.stem_summary if info is not None and info.stem_summary else record.question_id
            )
            item["canonical_answer"] = (
                info.canonical_answer if info is not None else ""
            )
            item["student_answer"] = record.student_answer or None
            item["question_text"] = info.question_text if info is not None else ""
            item["reference_analysis"] = info.reference_analysis if info is not None else ""
            item["grading_record"] = {
                "deduction_reason": record.deduction_reason or None,
                "error_category": record.error_category or None,
                "error_summary": record.error_summary or None,
                "secondary_errors": record.secondary_errors,
                "evidence_steps": record.evidence_steps,
                "missing_steps": record.missing_steps,
                "teacher_confirmed": record.teacher_confirmed,
                "teacher_comment": record.teacher_comment or None,
            }
        questions.append(item)
    return {
        "exam": {
            "title": data.session_name,
            "class_name": data.class_name,
            "subject": data.subject,
            "full_score": data.full_score,
            "graded_at": student.graded_at,
        },
        "student": {
            "alias": PERSONAL_STUDENT_ALIAS,
            "total_score": student.student_score,
            "class_stats": {
                "n": data.present,
                "avg": data.stats["avg"],
                "median": data.stats["median"],
                "max": data.stats["max"],
                "min": data.stats["min"],
                "rank": student.rank,
            },
            "small_sample": data.small_sample,
        },
        "questions": questions,
        "knowledge_statistics_note": "知识点按可匹配的小问直接考查范围统计；小问难度是公式难度，得分率是本次考试表现，均不能直接当作当前掌握度。综合小问的失分不能推断为其中每个知识点或步骤都不会。",
    }


def record_type_label(
    record: _StudentQuestionRecord,
    info: _QuestionInfo | None,
) -> str:
    return str(info.question_type if info is not None else "") or ""


def build_class_payload(data: _SessionAnalysisData) -> dict[str, Any]:
    aliases = {
        student.student_id: f"{CLASS_ALIAS_PREFIX}{index}"
        for index, student in enumerate(data.students, start=1)
    }
    bands = {
        band["label"].replace(" ", ""): band["count"] for band in data.stats["bands"]
    }
    students_payload: list[dict[str, Any]] = []
    for student in data.students:
        students_payload.append(
            {
                "alias": aliases[student.student_id],
                "total": student.student_score,
                "needs_review": student.needs_review,
                "lost": [
                    {
                        "question_id": record.question_id,
                        "lost_points": round(record.lost_points, 2),
                        "record": _record_brief_text(record),
                    }
                    for record in student.records
                    if record.lost
                ],
            }
        )
    records_by_question: dict[str, list[dict[str, Any]]] = {}
    for student in data.students:
        for record in student.records:
            if not record.lost:
                continue
            records_by_question.setdefault(record.question_id, []).append(
                {
                    "alias": aliases[student.student_id],
                    "score": record.score,
                    "deduction_reason": record.deduction_reason or None,
                    "error_category": record.error_category or None,
                }
            )
    questions_payload = [
        {
            "question_id": info.question_id,
            "max_score": info.max_score,
            "class_rate": (
                round(info.class_rate, 4) if info.class_rate is not None else None
            ),
            "stem_summary": info.stem_summary or info.question_id,
            "canonical_answer": info.canonical_answer,
            "records": records_by_question.get(info.question_id, []),
        }
        for info in data.questions
    ]
    return {
        "exam": {
            "title": data.session_name,
            "subject": data.subject,
            "full_score": data.full_score,
            "class_name": data.class_name,
            "present": data.present,
            "roster_absent": data.roster_absent,
        },
        "score_distribution": {
            "avg": data.stats["avg"],
            "median": data.stats["median"],
            "max": data.stats["max"],
            "min": data.stats["min"],
            "pass_rate": data.stats["pass_rate"],
            "bands": bands,
        },
        "students": students_payload,
        "questions": questions_payload,
    }


def _record_brief_text(record: _StudentQuestionRecord) -> str:
    parts = [
        text
        for text in (
            record.deduction_reason,
            record.error_category,
            record.error_summary,
        )
        if text
    ]
    return "；".join(parts) or "未作答或无批改记录"


# ---------------------------------------------------------------------------
# 班级分析内嵌页面的出参装配（教师本人页面，学生用真实姓名，不脱敏）
# ---------------------------------------------------------------------------


def build_class_page_data(data: _SessionAnalysisData, *, compact: bool = False) -> dict[str, Any]:
    """把班级版装配数据序列化为页面 JSON。

    出参形状与前端解码器 frontend/src/api/class-analysis.ts 冻结对齐：
    - present / roster_absent（姓名数组）在 data 顶层；
    - score_distribution.bands 为 {分数段: 人数} 字典；
    - students[].lost[] 为 {question_id, lost_points, record}；
    - questions[].class_rate 恒为数字（无作答记录时 0）。
    - compact 模式只返回逐题失分名单与归并错因，省略逐人清单和重复错因正文。
    """
    bands = {
        str(band["label"]): int(band["count"]) for band in data.stats["bands"]
    }
    absent_names = [
        name
        for item in data.skipped
        if str(item.get("reason") or "") in {"缺考", "扫描异常"}
        and (name := str(item.get("student_name") or "").strip())
    ]
    records_by_question: dict[str, list[tuple[_StudentReportData, _StudentQuestionRecord]]] = {}
    for student in data.students:
        for record in student.records:
            if record.lost:
                records_by_question.setdefault(record.question_id, []).append((student, record))
    questions: list[dict[str, Any]] = []
    for info in data.questions:
        records: list[dict[str, Any]] = []
        causes: dict[str, dict[str, Any]] = {}
        for student, record in records_by_question.get(info.question_id, []):
            if compact:
                records.append({
                    "student_id": student.student_id,
                    "student_code": student.student_code,
                    "student_name": student.student_name,
                    "class_name": student.class_name,
                    "score": record.score,
                })
                reason = record.error_summary or record.deduction_reason or record.error_category or "未记录具体错因"
                # 仅归并已有记录中的相同错因，不把不同表述推断为同一知识错误。
                seen = set()
                for text in re.split(r"[；;\n]+", reason):
                    text = text.strip().rstrip("。.")
                    if not text:
                        continue
                    key = re.sub(r"\s+", "", text).replace("，", ",").replace("：", ":")
                    if key in seen:
                        continue
                    seen.add(key)
                    bucket = causes.setdefault(key, {"reason": text, "count": 0})
                    bucket["count"] += 1
                if not seen:
                    bucket = causes.setdefault("未记录具体错因", {"reason": "未记录具体错因", "count": 0})
                    bucket["count"] += 1
            else:
                records.append({
                    "student_id": student.student_id,
                    "student_code": student.student_code,
                    "student_name": student.student_name,
                    "score": record.score,
                    "max_score": record.max_score,
                    "lost_points": round(record.lost_points, 2),
                    "deduction_reason": record.deduction_reason or None,
                    "error_category": record.error_category or None,
                    "error_summary": record.error_summary or None,
                })
        questions.append(
            {
                "question_id": info.question_id,
                "question_type": info.question_type,
                "max_score": info.max_score,
                "class_rate": (
                    round(info.class_rate, 4) if info.class_rate is not None else 0
                ),
                "class_avg": (
                    round(info.class_avg, 2) if info.class_avg is not None else None
                ),
                "stem_summary": info.stem_summary,
                "canonical_answer": info.canonical_answer,
                "records": records,
                **({"causes": sorted(causes.values(), key=lambda item: -item["count"])} if compact else {}),
            }
        )
    students: list[dict[str, Any]] = []
    for student in ([] if compact else data.students):
        lost_items: list[dict[str, Any]] = []
        for record in student.records:
            if not record.lost:
                continue
            record_payload = {
                "deduction_reason": record.deduction_reason or None,
                "error_category": record.error_category or None,
                "error_summary": record.error_summary or None,
            }
            if not any(record_payload.values()):
                record_payload = None
            lost_items.append(
                {
                    "question_id": record.question_id,
                    "lost_points": round(record.lost_points, 2),
                    "record": record_payload,
                }
            )
        students.append(
            {
                "student_id": student.student_id,
                "student_code": student.student_code,
                "student_name": student.student_name,
                "class_name": student.class_name,
                "total_score": student.student_score,
                "rank": student.rank,
                "needs_review": student.needs_review,
                "lost_points_total": round(student.lost_points_total, 2),
                "lost": lost_items,
            }
        )
    stats = data.stats
    return {
        "exam": {
            "session_id": data.session_id,
            "title": data.session_name,
            "subject": data.subject,
            "full_score": data.full_score,
            "graded_at": data.graded_at,
        },
        "present": data.present,
        "roster_absent": absent_names,
        "score_distribution": {
            "avg": stats["avg"],
            "median": stats["median"],
            "max": stats["max"],
            "min": stats["min"],
            "pass_rate": stats["pass_rate"],
            "bands": bands,
        },
        "questions": questions,
        "students": students,
        "skipped": data.skipped,
    }


def class_narrative_with_student_names(
    narrative: dict[str, Any],
    students: list[_StudentReportData],
) -> dict[str, Any]:
    """保留结构中的 alias 键，把正文里的学生代号映射回本班姓名。"""
    alias_to_student = {
        f"{CLASS_ALIAS_PREFIX}{index}": student
        for index, student in enumerate(students, start=1)
    }
    # 中文紧邻 S1 时也应识别代号；Unicode 的 \b 会把中文和数字都视为单词字符。
    alias_pattern = re.compile(r"(?<![A-Za-z0-9_])S\d+(?![A-Za-z0-9_])")

    def map_text(value: Any) -> Any:
        if isinstance(value, str):
            return alias_pattern.sub(
                lambda match: alias_to_student[match[0]].student_name
                if match[0] in alias_to_student else match[0], value,
            )
        if isinstance(value, list):
            return [map_text(item) for item in value]
        if isinstance(value, dict):
            return {key: item if key == "alias" else map_text(item) for key, item in value.items()}
        return value

    mapped = map_text(narrative)
    notes: list[dict[str, Any]] = []
    for item in _narrative_items(mapped, "student_notes"):
        if not isinstance(item, dict):
            continue
        entry = dict(item)
        student = alias_to_student.get(_narrative_text(item.get("alias")))
        if student is not None:
            entry["student_name"] = student.student_name
            entry["student_code"] = student.student_code
        notes.append(entry)
    mapped["student_notes"] = notes
    return mapped


def build_report_prompt(
    system_prompt: str,
    payload: dict[str, Any],
) -> str:
    return (
        system_prompt
        + "\n\n输入 JSON：\n"
        + json.dumps(payload, ensure_ascii=False)
    )


# ---------------------------------------------------------------------------
# 错题截图（方案 §4）
# ---------------------------------------------------------------------------


def _find_question_regions(
    regions: list[dict[str, Any]],
    parent_question_id: str,
) -> list[dict[str, Any]]:
    """保留父题的全部作答区域；未标整题区域时使用其全部小问区域。"""
    parent_regions: list[dict[str, Any]] = []
    child_regions: list[dict[str, Any]] = []
    for region in regions:
        region_id = str(region.get("mapped_question_id") or region.get("detected_question_id") or "").strip()
        if _parent_question_id(region_id) != parent_question_id:
            continue
        coordinates = question_id_coordinates(region_id)
        if coordinates is None or coordinates[1] is None:
            parent_regions.append(region)
        else:
            child_regions.append(region)
    return parent_regions or child_regions


def _region_page(region: dict[str, Any]) -> str:
    return (
        "back"
        if str(region.get("page") or "front").strip().lower() == "back"
        else "front"
    )


def _crop_region_data_uri(
    image_path: Path,
    region: dict[str, Any] | None,
    *,
    max_width: int = _SHOT_MAX_WIDTH,
) -> str | None:
    try:
        with Image.open(image_path) as image:
            image.load()
            bbox = (
                scaled_region_bbox(region, image.width, image.height, padding=_SHOT_PADDING)
                if region is not None else (0, 0, image.width, image.height)
            )
            cropped = image.crop(bbox)
            if cropped.width > max_width:
                ratio = max_width / cropped.width
                cropped = cropped.resize(
                    (max_width, max(1, round(cropped.height * ratio))),
                    Image.Resampling.LANCZOS,
                )
            buffer = io.BytesIO()
            cropped.convert("RGB").save(
                buffer,
                format="JPEG",
                quality=_SHOT_JPEG_QUALITY,
            )
    except Exception:
        return None
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    return f"data:image/jpeg;base64,{encoded}"


def _student_paper_context(
    repositories: GradingRepositoryAccess,
    data: _SessionAnalysisData,
    student: _StudentReportData,
) -> dict[str, Any] | None:
    if student.result_id > 0:
        context = repositories.results.get_result_context(student.result_id)
        if (
            isinstance(context, dict)
            and int(context.get("session_id") or 0) == data.session_id
            and int(context.get("student_id") or 0) == student.student_id
        ):
            return context
        return None
    # Manual scores have no session_results row.  The current exam-paper
    # assignment is authoritative; do not guess among duplicate assignments.
    candidates = [
        row for row in repositories.papers.get_session_paper_identities(data.session_id)
        if int(row.get("student_id") or 0) == student.student_id
        and str(row.get("match_status") or "") == "matched"
    ]
    return candidates[0] if len(candidates) == 1 else None


def capture_lost_question_shots(
    repositories: GradingRepositoryAccess,
    data: _SessionAnalysisData,
    student: _StudentReportData,
    *,
    regions: list[dict[str, Any]] | None,
    data_root: Path | None,
    paper_context: dict[str, Any] | None = None,
) -> dict[str, dict[str, str]]:
    """为丢分大题生成全部截图；同题多图以序号区分，缺图不阻断报告。"""
    if not regions:
        return {}
    lost_records = [record for record in student.records if record.lost]
    if not lost_records:
        return {}

    if paper_context is None:
        paper_context = _student_paper_context(repositories, data, student)

    # 区域粒度是整道大题：同一大题多小问丢分只嵌一次，按丢分从大到小排序。
    lost_by_parent: dict[str, float] = {}
    for record in lost_records:
        parent = _parent_question_id(record.question_id)
        lost_by_parent[parent] = lost_by_parent.get(parent, 0.0) + record.lost_points
    ordered_parents = sorted(
        lost_by_parent,
        key=lambda parent: (-lost_by_parent[parent], parent),
    )

    shots: dict[str, dict[str, str]] = {}
    used_bytes = 0
    for parent in ordered_parents:
        matching = _find_question_regions(regions, parent)
        for index, region in enumerate(matching):
            page = _region_page(region)
            raw_path = str((paper_context or {}).get(f"{page}_image") or "")
            if not raw_path:
                continue
            image_path = resolve_stored_file_path(raw_path, data_root=data_root)
            data_uri = _crop_region_data_uri(image_path, region)
            if data_uri is None:
                continue
            encoded_size = len(data_uri) * 3 // 4
            if used_bytes + encoded_size > _SHOT_BUDGET_BYTES:
                continue
            used_bytes += encoded_size
            region_id = str(region.get("mapped_question_id") or region.get("detected_question_id") or parent)
            shots[parent if index == 0 else f"{parent}:{index}"] = {
                "data_uri": data_uri,
                "parent_question_id": parent,
                "region_question_id": region_id,
                "caption": f"{_question_display_label(region_id)}作答区截图（原卷截图）",
            }
    return shots


def _personal_image_inputs(
    data: _SessionAnalysisData,
    student: _StudentReportData,
    shots: dict[str, dict[str, str]],
    paper_context: dict[str, Any] | None,
    data_root: Path | None,
) -> tuple[list[bytes], list[dict[str, Any]]]:
    """有序组织原卷、题图和作答概览；序号与同一次图文请求严格对应。"""
    images: list[bytes] = []
    image_map: list[dict[str, Any]] = []
    total_bytes = 0
    student.material_notes = []

    def append(blob: bytes, *, kind: str, label: str, question_ids: list[str]) -> None:
        nonlocal total_bytes
        for index, existing in enumerate(images):
            if blob == existing and image_map[index]["kind"] == kind:
                image_map[index]["question_ids"] = list(dict.fromkeys(
                    [*image_map[index]["question_ids"], *question_ids]
                ))
                return
        if len(images) >= 48 or total_bytes + len(blob) > _SHOT_BUDGET_BYTES:
            student.material_notes.append(f"{label}未纳入图片分析，不能推断图中细节。")
            return
        images.append(blob)
        total_bytes += len(blob)
        image_map.append({
            "image_number": len(images), "kind": kind,
            "label": label, "question_ids": question_ids,
        })

    lost = [record for record in student.records if record.lost]
    for key, shot in shots.items():
        parent = shot.get("parent_question_id") or key
        region_id = shot.get("region_question_id") or parent
        region_coordinates = question_id_coordinates(region_id)
        question_ids = [
            record.question_id for record in student.records
            if record.question_id == region_id or (
                region_coordinates is not None and (
                    question_id_coordinates(record.question_id) == region_coordinates
                    or (region_coordinates[1] is None and _parent_question_id(record.question_id) == parent)
                )
            )
        ]
        append(
            base64.b64decode(shot["data_uri"].split(",", 1)[1]),
            kind="student_work", label=f"{_question_display_label(region_id)}学生作答原图",
            question_ids=question_ids,
        )
    info_by_qid = {info.question_id: info for info in data.questions}
    for record in lost:
        info = info_by_qid.get(record.question_id)
        if info is None:
            continue
        for role, blob in info.reference_images:
            append(
                blob, kind=f"reference_{role}",
                label=f"{_question_display_label(_parent_question_id(record.question_id))}{'原题及题图' if role == 'question' else '参考答案解析'}",
                question_ids=[record.question_id],
            )
    # The overview supplies distribution context, not a record of working time.
    for page, label in (("front", "正面"), ("back", "反面")):
        raw_path = str((paper_context or {}).get(f"{page}_image") or "")
        if not raw_path:
            continue
        uri = _crop_region_data_uri(
            resolve_stored_file_path(raw_path, data_root=data_root), None,
            max_width=1200,
        )
        if uri is not None:
            append(
                base64.b64decode(uri.split(",", 1)[1]), kind="student_overview",
                label=f"学生答卷{label}概览，仅用于观察全卷作答分布",
                question_ids=[record.question_id for record in student.records],
            )
    if not any(item["kind"].startswith("student_") for item in image_map):
        student.material_notes.append("未取得该生可读取的答卷图片；缺少作答文字不代表空白，不能推断未作答原因。")
    return images, image_map


def load_session_regions(
    repositories: GradingRepositoryAccess,
    session_id: int,
    *,
    data_root: Path | None,
) -> list[dict[str, Any]]:
    try:
        return answer_regions_with_template_source_sizes(
            repositories,
            int(session_id),
            data_root=data_root,
        )
    except Exception:
        return []


# ---------------------------------------------------------------------------
# HTML 渲染（版面结构与视觉按已确认原型，自包含、截图 base64 内嵌）
# ---------------------------------------------------------------------------

# 家长版个人报告版式（v3 一页式：成绩卡 → 历次成绩 → 答题一览 →
# 重点跟进 → 考查点 → 失分题详解附录），与已验收原型一致。
_PERSONAL_CSS = """
:root{--ink:#1f2d3d;--muted:#6b7a8c;--line:#e3e8ef;--bg:#f4f6f9;--card:#fff;
--good:#2f855a;--good-soft:#e8f5ee;--warn:#b7791f;--warn-soft:#fdf3e1;
--bad:#c0392b;--bad-soft:#fbeaea;--gray:#8a97a8;--gray-soft:#eef1f5;--accent:#2b6cb0}
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:"PingFang SC","Microsoft YaHei UI","Microsoft YaHei",sans-serif;
color:var(--ink);background:var(--bg);line-height:1.7;font-size:15px;
font-variant-numeric:tabular-nums}
.page{max-width:760px;margin:0 auto;padding:18px 14px 40px}
.card{background:var(--card);border-radius:12px;box-shadow:0 1px 2px rgba(31,45,61,.06);
padding:20px 22px;margin-bottom:16px}
h2{font-size:17px;margin-bottom:12px}
h2 .aitag{font-size:11px;font-weight:400;color:var(--warn);background:var(--warn-soft);
border:1px solid #ecd9b0;border-radius:999px;padding:1px 9px;margin-left:8px;vertical-align:2px}
.note{font-size:12px;color:var(--muted)}
.meta{font-size:12.5px;color:var(--muted)}
.review-banner{background:var(--warn-soft);color:var(--warn);border-radius:8px;
padding:8px 12px;font-size:13px;margin-bottom:12px}
.name{font-size:24px;font-weight:700;margin-top:6px}
.score-line{display:flex;align-items:baseline;gap:4px;margin:8px 0 14px}
.score-line .big{font-size:46px;font-weight:800;line-height:1}
.score-line .of{font-size:16px;color:var(--muted)}
.stat3{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin-bottom:14px}
.stat3 .cell{border:1px solid var(--line);border-radius:10px;padding:9px 10px;text-align:center}
.stat3 .v{font-size:16px;font-weight:700;white-space:nowrap}
.stat3 .l{font-size:11.5px;color:var(--muted);white-space:nowrap}
.bands{margin:6px 0 12px}
.band{display:grid;grid-template-columns:96px 1fr 44px;align-items:center;gap:8px;
font-size:12px;color:var(--muted);padding:2px 0}
.band .track{height:12px;background:var(--gray-soft);border-radius:6px;overflow:hidden}
.band .track i{display:block;height:100%;background:#b9c6d4;border-radius:6px}
.band.me{color:var(--ink);font-weight:600}
.band.me .track i{background:var(--accent)}
.band .cnt{text-align:right}
.me-tag{display:inline-block;font-size:11px;color:var(--accent);margin-left:6px}
.concl{font-size:14.5px;margin-top:6px}
.strength{color:var(--good);font-size:13.5px;margin-top:6px}
/* B 历次成绩 */
.hchart{position:relative;height:140px;margin:4px 0 2px}
.hchart svg{position:absolute;inset:0;width:100%;height:100%}
.hdot{position:absolute;width:9px;height:9px;border-radius:50%;background:var(--accent);
border:2px solid #fff;transform:translate(-50%,-50%);box-shadow:0 0 0 1px var(--line)}
.hdot.cur{width:11px;height:11px}
.hval{position:absolute;transform:translate(-50%,-135%);font-size:12px;color:var(--ink);white-space:nowrap}
.hval.cur{font-weight:700}
.hcols{display:flex;margin-top:6px}
.hcol{flex:1;min-width:0;text-align:center;font-size:12px;color:var(--muted);line-height:1.5}
.hcol .hn{white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.hcol.cur{color:var(--ink);font-weight:700}
/* C 答题一览 */
.qgrid{display:grid;grid-template-columns:repeat(6,1fr);gap:8px}
@media(min-width:700px){.qgrid{grid-template-columns:repeat(10,1fr)}}
.qcell{border:1px solid var(--line);border-radius:8px;background:#fff;padding:6px 2px 4px;
text-align:center;cursor:pointer;position:relative;font-family:inherit}
.qcell .qn{display:block;font-size:13px;font-weight:700}
.qcell .qs{display:block;font-size:11px;color:var(--muted)}
.qcell .sym{position:absolute;left:4px;top:3px;font-size:11px;line-height:1}
.qcell .hard{position:absolute;right:3px;top:2px;font-size:9px;color:var(--warn);line-height:1}
.qcell.full{border-color:#bfe3d2;background:var(--good-soft)} .qcell.full .sym{color:var(--good)}
.qcell.part{border-color:#f0d9ab;background:var(--warn-soft)} .qcell.part .sym{color:var(--warn)}
.qcell.zero{border-color:#eec6c0;background:var(--bad-soft)} .qcell.zero .sym{color:var(--bad)}
.qcell.blank{border-style:dashed;background:var(--gray-soft)} .qcell.blank .sym{color:var(--gray)}
.qcell[aria-expanded="true"]{outline:2px solid var(--accent);outline-offset:1px}
.legend{font-size:12px;color:var(--muted);margin-top:10px}
#qpanel{margin-top:10px}
.qdetail-box{border:1px solid var(--line);border-radius:10px;padding:14px 16px;background:#fcfdfe}
.qdetail-box .qd-head{font-size:14px;font-weight:700}
.qdetail-box .qd-points{font-size:13px;color:var(--muted);margin-top:4px}
.qdetail-box .qd-points .pt-tag{display:inline-block;background:var(--gray-soft);border-radius:6px;
padding:0 8px;margin:2px 4px 2px 0;font-size:12px;color:var(--ink)}
.qd-fb{font-size:13.5px;margin-top:8px}
.errline{font-size:12.5px;color:var(--warn);margin-top:5px}
.qmore-btn{margin-top:10px;border:1px solid var(--accent);color:var(--accent);background:#fff;
border-radius:8px;padding:5px 14px;font-size:13px;cursor:pointer;font-family:inherit}
.qmore{margin-top:10px;border-top:1px dashed var(--line);padding-top:10px}
/* 共用题干/截图/分析 */
.stem{white-space:normal;overflow-wrap:anywhere;font-size:13.5px;color:var(--ink)}
.stem table{border-collapse:collapse;width:100%;margin:8px 0}
.stem td,.stem th{border:1px solid #cbd5e1;padding:5px 8px;text-align:center}
.reference-figures{display:flex;flex-wrap:wrap;gap:12px;margin:10px 0}
.reference-figures img{max-width:100%;max-height:180px;object-fit:contain}
.shot{border:1px solid var(--line);border-radius:8px;padding:8px;margin:10px 0 2px;background:#fcfcfd}
.shot img{max-width:100%;display:block;margin:0 auto;border-radius:4px}
.shot .cap{font-size:11.5px;color:var(--muted);margin-top:6px;text-align:center}
.kv{display:grid;grid-template-columns:88px 1fr;gap:4px 10px;font-size:13px}
.kv dt{color:var(--muted)} .kv dd{min-width:0;white-space:pre-wrap;overflow-wrap:anywhere}
.kv dd b.ans{color:var(--good)}
.compact-analysis{font-size:13.5px;line-height:1.75}
.compact-analysis p{margin:3px 0}
.compact-analysis b{color:#334155}
.feedback,.solution,.revision{margin:10px 0}
.solution{border-left:3px solid #b8cde9;padding:2px 12px}
.solution ol{margin:4px 0;padding-left:22px}
.review-note{background:#fff7e8;border-left:3px solid #e9972d;padding:8px 12px;margin:10px 0}
.review-note b{color:#a45b00}
.compact-analysis details{margin:8px 0;color:#475569}
.compact-analysis summary{cursor:pointer;color:var(--accent)}
.aidraft{margin-top:10px;background:var(--warn-soft);border:1px dashed #f5d08c;border-radius:8px;
padding:10px 12px;font-size:13px}
.aidraft .cap{font-size:11.5px;color:var(--warn);font-weight:700;margin-bottom:4px}
.aidraft p+p{margin-top:6px}
.qm{display:inline-block;max-width:100%;vertical-align:baseline;padding:3px 2px 6px}
.qm-display{display:block;text-align:center;margin:8px 0}
.qm-error{color:#9a3412}
.katex{font-size:1.04em}
.stem,.kv dd,.compact-analysis p{overflow-x:auto;overflow-y:hidden;overflow-wrap:anywhere}
/* D 跟进卡片 */
.dcard{border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin-bottom:12px}
.dcard .dhead{display:flex;align-items:center;gap:9px}
.dcard .no{display:inline-flex;align-items:center;justify-content:center;width:22px;height:22px;
border-radius:50%;background:var(--accent);color:#fff;font-size:12px;font-weight:700;flex:none}
.dcard .ttl{font-size:14.5px;font-weight:700}
.dcard .meta2{font-size:12.5px;color:var(--muted);margin-top:5px}
.dcard .qtag{display:inline-block;background:var(--gray-soft);border-radius:6px;padding:0 7px;
margin-left:5px;font-size:11.5px;color:var(--ink)}
.dcard .help{font-size:13.5px;margin-top:7px}
.recur{display:inline-block;font-size:11px;color:var(--warn);background:var(--warn-soft);
border:1px solid #ecd9b0;border-radius:999px;padding:0 8px;margin-left:8px;vertical-align:1px}
.dcard details{margin-top:9px;border-top:1px dashed var(--line);padding-top:8px}
.dcard summary{cursor:pointer;font-size:12.5px;color:var(--accent)}
.dgreen{background:var(--good-soft);border:1px solid #bfe3d2;border-radius:10px;
padding:14px 16px;color:var(--good);font-size:14px}
/* E 考查点 */
.pt-counts{font-size:13.5px;margin-bottom:10px}
.pt-counts b{margin-right:14px}
.pt-counts .g{color:var(--good)} .pt-counts .m{color:var(--warn)} .pt-counts .b{color:var(--bad)}
.pt-row{display:flex;align-items:flex-start;gap:8px;padding:7px 0;border-top:1px solid var(--line);
font-size:13.5px;flex-wrap:wrap}
.pt-row .nm{min-width:0}
.pt-row .dots{display:inline-flex;gap:7px;margin-left:auto;flex-wrap:wrap}
.dot{display:inline-flex;flex-direction:column;align-items:center;width:30px}
.dot i{width:18px;height:18px;border-radius:50%;display:flex;align-items:center;justify-content:center;
font-size:11px;font-style:normal;line-height:1}
.dot s{text-decoration:none;font-size:9px;color:var(--muted);margin-top:1px}
.dot i.full{background:var(--good-soft);color:var(--good);border:1px solid #bfe3d2}
.dot i.part{background:var(--warn-soft);color:var(--warn);border:1px solid #f0d9ab}
.dot i.zero{background:var(--bad-soft);color:var(--bad);border:1px solid #eec6c0}
.tagcloud{margin-top:8px}
.tagcloud .sec{font-size:12px;color:var(--muted);margin:6px 0 3px}
.tagcloud .tg{display:inline-block;background:var(--good-soft);color:var(--good);border-radius:6px;
padding:1px 9px;margin:2px 5px 2px 0;font-size:12.5px}
/* F 附录 */
details.appendix summary{cursor:pointer;font-size:13px;color:var(--accent)}
.qcard{border:1px solid var(--line);border-left:4px solid var(--bad);border-radius:10px;
padding:14px 16px;margin:12px 0}
.qcard .head{display:flex;justify-content:space-between;align-items:baseline;flex-wrap:wrap;
gap:6px;margin-bottom:8px}
.qcard .head b{font-size:14.5px}
.qcard .score{font-size:13px;color:var(--bad);font-weight:700}
.qpart+.qpart{border-top:1px solid var(--line);margin-top:14px;padding-top:14px}
.qpart h3{display:flex;justify-content:space-between;gap:12px;font-size:13px;margin-bottom:8px}
.qpart h3 span{color:var(--muted);font-weight:400}
/* footer */
footer{font-size:12px;color:var(--muted);padding:6px 4px;line-height:1.8}
@media(max-width:560px){
 body{font-size:14px}
 .page{padding:10px 8px 28px}
 .card{padding:16px 14px}
 .qgrid{grid-template-columns:repeat(6,1fr)}
 .kv{grid-template-columns:70px minmax(0,1fr)}
}
@page{size:A4;margin:13mm 14mm}
@media print{
 body{background:#fff}
 .card{box-shadow:none;border:1px solid var(--line);border-radius:0}
 details:not([open])>*:not(summary){display:block}
 #qpanel{display:none!important}
 .qcell{outline:none!important}
 .shot img{max-height:95mm;width:auto;object-fit:contain}
 .qpart,.shot,.dcard{break-inside:avoid}
 p,dd,li{orphans:3;widows:3}
}
"""

# C 节方格 → 原位详情面板（<template> 片段克隆，不滚动、不跳转）。
_PERSONAL_GRID_JS = """
(function(){
  var panel = document.getElementById('qpanel');
  var openQ = null;
  function renderMath(root){
    root.querySelectorAll('.qm[data-latex]').forEach(function(el){
      try{ katex.render(el.dataset.latex, el, {throwOnError:true, output:'htmlAndMathml', displayMode:el.dataset.display === 'true'}); }
      catch(e){ el.classList.add('qm-error'); el.title='公式格式需核对，已保留原文'; }
    });
  }
  if (!panel) return;
  document.querySelectorAll('.qcell').forEach(function(btn){
    btn.addEventListener('click', function(){
      var q = btn.getAttribute('data-q');
      if(openQ === q){
        panel.hidden = true; panel.innerHTML=''; openQ = null;
        btn.setAttribute('aria-expanded','false'); return;
      }
      var tpl = document.querySelector('template[data-q="' + q + '"]');
      if(!tpl) return;
      panel.innerHTML=''; panel.appendChild(tpl.content.cloneNode(true));
      panel.hidden=false; openQ=q; renderMath(panel);
      document.querySelectorAll('.qcell[aria-expanded="true"]').forEach(function(b){b.setAttribute('aria-expanded','false');});
      btn.setAttribute('aria-expanded','true');
    });
  });
  panel.addEventListener('click', function(e){
    var t = e.target.closest('.qmore-btn');
    if(!t) return;
    var m = panel.querySelector('.qmore');
    if(!m) return;
    m.hidden = !m.hidden;
    t.textContent = m.hidden ? '看原卷和解法' : '收起原卷和解法';
  });
})();
"""

_PERSONAL_KATEX_JS = """
document.querySelectorAll('.qm[data-latex]').forEach(el => {
  try { katex.render(el.dataset.latex, el, {throwOnError:true, output:'htmlAndMathml', displayMode:el.dataset.display === 'true'}); }
  catch (_) { el.classList.add('qm-error'); el.title = '公式格式需核对，已保留原文'; }
});
"""


def _narrative_items(narrative: dict[str, Any] | None, key: str) -> list[Any]:
    if not isinstance(narrative, dict):
        return []
    value = narrative.get(key)
    return value if isinstance(value, list) else []


def _narrative_text(value: object) -> str:
    return str(value or "").strip()


def _knowledge_rows(
    backfill: dict[str, list[dict[str, str]]],
    records: list[_StudentQuestionRecord],
) -> list[dict[str, Any]]:
    """按题库知识点标签聚合得分率；无标签数据时返回空，渲染层省略该板块。"""
    buckets: dict[str, dict[str, Any]] = {}
    for record in records:
        if record.max_score <= 0:
            continue
        entries = backfill.get(record.question_id)
        if entries is None:
            entries = backfill.get(_parent_question_id(record.question_id))
        for entry in entries or []:
            path = str(entry.get("path") or "").strip()
            if not path:
                continue
            bucket = buckets.setdefault(
                path,
                {
                    "label": str(entry.get("label") or "").strip() or "未命名知识点",
                    "score": 0.0,
                    "full": 0.0,
                },
            )
            bucket["score"] += min(record.score, record.max_score)
            bucket["full"] += record.max_score
    rows = []
    for bucket in buckets.values():
        full = float(bucket["full"])
        rows.append(
            {
                "label": bucket["label"],
                "rate": (float(bucket["score"]) / full) if full > 0 else None,
            }
        )
    rows.sort(key=lambda row: (row["rate"] is None, row["rate"] or 0))
    return rows


def _personal_knowledge_view(data: _SessionAnalysisData, student: _StudentReportData) -> dict[str, Any]:
    snapshot = data.knowledge_structure
    catalog = {str(item['knowledge_key']): item for item in snapshot.get('catalog', [])}
    nodes: dict[str, dict[str, Any]] = {}
    for record in student.records:
        entries = data.knowledge_backfill.get(record.question_id)
        if entries is None:
            entries = data.knowledge_backfill.get(_parent_question_id(record.question_id), [])
        for entry in entries:
            key = str(entry.get('stable_key') or entry.get('path') or '')
            if not key or record.max_score <= 0:
                continue
            item = catalog.get(key, {})
            path = str(item.get('knowledge_point') or entry.get('path') or entry.get('label') or '')
            segments = [part.strip() for part in re.split('[|｜]', path) if part.strip()]
            label = str(entry.get('label') or (segments[-1] if segments else key))
            kind = item.get('node_kind') or ('skill' if key.startswith('sk_') or label.startswith('技能') else 'topic')
            if kind not in {'topic', 'skill'}:
                kind = 'topic'
            chapter, section, cursor, seen = '', '', item, set()
            while cursor and str(cursor.get('knowledge_key')) not in seen:
                seen.add(str(cursor.get('knowledge_key')))
                if cursor.get('node_kind') == 'chapter':
                    chapter = str(cursor['knowledge_point']).split('｜')[-1].split('|')[-1]
                elif cursor.get('node_kind') == 'section':
                    section = str(cursor['knowledge_point']).split('｜')[-1].split('|')[-1]
                cursor = catalog.get(str(cursor.get('parent_knowledge_key') or ''), {})
            chapter = chapter or ('｜'.join(segments[:-2]) if len(segments) >= 3 else '本卷知识与技能')
            section = section or (segments[-2] if len(segments) >= 3 else '')
            mastery = student.knowledge_mastery.get(key, {})
            value = mastery.get('mastery')
            if value is not None and (not isinstance(value, (float, int)) or not math.isfinite(value)):
                value = None
            node = nodes.setdefault(key, {
                'key': key, 'label': re.sub(r'^技能[·・：:]\s*', '', label), 'kind': kind,
                'chapter': chapter, 'section': section, 'mastery': value,
                'evidence_count': int(mastery.get('evidence_count') or 0),
                'questions': [], 'score': 0.0, 'full': 0.0,
                'step_questions': sorted({str(ref.get('question_id')) for ref in mastery.get('source_question_refs', [])
                                          if int(ref.get('session_id') or 0) == data.session_id
                                          and (ref.get('assessment') or {}).get('granularity') == 'step'}),
            })
            if record.question_id not in {question['id'] for question in node['questions']}:
                node['score'] += min(record.score, record.max_score)
                node['full'] += record.max_score
                node['questions'].append({'id': record.question_id, 'label': _question_display_label(record.question_id),
                                          'score': record.score, 'full': record.max_score, 'lost': record.lost})
    edges = [dict(edge) for edge in snapshot.get('associations', [])
             if edge.get('topic_key') in nodes and edge.get('skill_key') in nodes]
    return {'nodes': list(nodes.values()), 'edges': edges, 'as_of': snapshot.get('as_of') or datetime.now().strftime('%Y-%m-%d %H:%M'),
            'note': snapshot.get('note') or '当前掌握度暂不可用；以下保留本卷考查范围与得分，不以得分率代替掌握度。'}


def _class_knowledge_view(data: _SessionAnalysisData) -> dict[str, Any]:
    """Aggregate students' independent mastery; missing evidence is not a zero."""
    nodes: dict[str, dict[str, Any]] = {}
    for student in data.students:
        for item in _personal_knowledge_view(data, student)['nodes']:
            node = nodes.setdefault(item['key'], {
                **item, 'score': 0.0, 'full': 0.0, 'questions': {}, 'step_questions': set(),
            })
            node['score'] += item['score']
            node['full'] += item['full']
            node['step_questions'].update(item['step_questions'])
            for question in item['questions']:
                bucket = node['questions'].setdefault(question['id'], {
                    'id': question['id'], 'label': question['label'],
                    'score': 0.0, 'full': 0.0, 'count': 0,
                })
                bucket['score'] += question['score']
                bucket['full'] += question['full']
                bucket['count'] += 1
    for node in nodes.values():
        values = [student.knowledge_mastery.get(node['key'], {}).get('mastery') for student in data.students]
        known = [float(value) for value in values
                 if isinstance(value, (int, float)) and math.isfinite(value) and 0 <= value <= 1]
        node['mastery'] = statistics.mean(known) if known else None
        node['coverage'] = len(known)
        node['student_count'] = len(data.students)
        node['distribution'] = {
            'low': sum(value < .6 for value in known),
            'mid': sum(.6 <= value < .75 for value in known),
            'good': sum(value >= .75 for value in known),
            'missing': len(data.students) - len(known),
        }
        node['step_questions'] = sorted(node['step_questions'])
        node.pop('evidence_count', None)
        node['questions'] = [{**question,
                              'score': question['score'] / question['count'],
                              'full': question['full'] / question['count']}
                             for question in node['questions'].values()]
    snapshot = data.knowledge_structure
    return {'mode': 'class', 'student_count': len(data.students), 'nodes': list(nodes.values()),
            'edges': [dict(edge) for edge in snapshot.get('associations', [])
                      if edge.get('topic_key') in nodes and edge.get('skill_key') in nodes],
            'as_of': snapshot.get('as_of') or datetime.now().strftime('%Y-%m-%d %H:%M'),
            'note': snapshot.get('note') or '当前掌握度暂不可用；保留本卷考查范围与得分，不以得分率代替掌握度。'}


def _knowledge_view_html(view: dict[str, Any]) -> str:
    if not view['nodes']:
        return ''
    class_mode = view.get('mode') == 'class'
    chapters: dict[str, list[dict[str, Any]]] = {}
    for node in view['nodes']:
        chapters.setdefault(node['chapter'], []).append(node)
    boards = []
    counts = {'good': 0, 'mid': 0, 'low': 0, 'missing': 0}

    def node_html(node):
        value = node['mastery']
        band = 'missing' if value is None else 'low' if value < .6 else 'mid' if value < .75 else 'good'
        counts[band] += 1
        percent = '证据不足' if value is None else f'{_fmt_num(math.floor(value * 1000) / 10)}%'
        status = {'missing': '证据不足', 'low': '待补强', 'mid': '需巩固', 'good': '较稳定'}[band]
        detail = f'本卷 {_fmt_num(node["score"])} / {_fmt_num(node["full"])} 分'
        distribution = ''
        if class_mode:
            detail = f'有证据 {node["coverage"]} / {node["student_count"]} 人'
            distribution = '<span class="kn-distribution">' + ''.join(
                f'<span class="kn-{key}"><i class="kn-dot"></i>{label} {node["distribution"][key]}</span>'
                for key, label in [('low', '补强'), ('mid', '巩固'), ('good', '稳定'), ('missing', '无证据')]
            ) + '</span>'
        value_label = '有证据学生平均掌握度' if class_mode else f'当前掌握度：{status}'
        return (f'<button type="button" class="kn-node kn-{band}" data-key="{_esc(node["key"])}" aria-pressed="false">'
                f'<span class="kn-node-top"><span class="kn-name">{_esc(node["label"])}</span>'
                f'<strong class="kn-value" title="{value_label}">{percent}</strong></span>'
                f'<span class="kn-node-meta">{_esc(node["section"])}<span>{detail}</span></span>'
                f'{distribution}</button>')

    for chapter, nodes in chapters.items():
        topics, skills = [n for n in nodes if n['kind'] == 'topic'], [n for n in nodes if n['kind'] == 'skill']
        columns = []
        for title, group, kind in [('知识点', topics, 'topics'), ('技能点', skills, 'skills')]:
            columns.append(f'<div class="kn-column kn-{kind}"><h4>{title}<span>{len(group)} 项</span></h4>'
                           + (''.join(node_html(node) for node in group) or '<p class="kn-empty">本卷暂无直接考查记录</p>') + '</div>')
        boards.append(f'<div class="kn-chapter"><h3>{_esc(chapter)}</h3><div class="kn-board">'
                      '<svg class="kn-lines" aria-hidden="true"></svg>' + ''.join(columns) + '</div></div>')
    data_json = json.dumps(view, ensure_ascii=False, separators=(',', ':')).replace('<', r'\u003c').replace('>', r'\u003e').replace('&', r'\u0026')
    title = '班级知识与技能掌握图' if class_mode else '知识与技能掌握图'
    scope_note = (f'统计本班 {view["student_count"]} 名已有成绩的学生；缺考及未形成有效成绩者不计入。'
                  '节点百分比仅对有证据学生取平均；各档显示人数，无证据单列，不参与平均。') if class_mode else ''
    footnote = ('节点颜色按有证据学生的平均掌握度划分；备课时同时查看分布，避免平均值掩盖差异。'
                '“本卷得分率”为相关小问整体得分率，不是技能独立得分。') if class_mode else (
                '节点百分比为当前掌握度；“本卷”分数为相关小问的整体得分，不能直接归因到其中每一步。')
    return (f'<section class="knowledge-map" id="knowledge-map"><div class="kn-heading"><h2>{title}</h2>'
            f'<span class="kn-date">截至 {_esc(view["as_of"])}</span></div>'
            f'<p class="kn-intro">{_esc(view["note"])}</p>'
            + (f'<p class="kn-scope">{scope_note}</p>' if scope_note else '')
            + '<p class="kn-thresholds">待补强 &lt;60% · 需巩固 60%–不足75% · 较稳定 ≥75%</p>'
            + '<div class="kn-legend">'
            + ('<span>按节点平均值：</span>' if class_mode else '')
            + ''.join(f'<span><i class="kn-dot kn-{band}"></i>{label} <b>{counts[band]} 项</b></span>'
                      for band, label in [('low', '待补强'), ('mid', '需巩固'), ('good', '较稳定'), ('missing', '证据不足')])
            + '</div><p class="kn-help">点选知识点或技能点，查看关联与本卷表现。实线：有同小问依据；虚线：仅同题出现。关联不表示掌握度相同。</p>'
            + ''.join(boards)
            + '<div class="kn-inspector" aria-live="polite"><p>点选上方节点，查看相关题目与证据。</p></div>'
            f'<p class="kn-footnote">{footnote}未达满分的教师复核步骤按未达成计入证据。</p>'
            f'<script type="application/json" class="kn-data">{data_json}</script></section>')


def _ai_block(text: str, *, failed: bool) -> str:
    note = AI_FAILED_NOTE if failed else text
    return (
        '<div class="aidraft">'
        f'<div class="cap">{_esc(AI_DISCLAIMER)}</div>{_report_paragraphs(note)}</div>'
    )


def _report_paragraphs(value: object) -> str:
    text = _report_display_text(value)
    # Split prose after recognizing TeX so multiline display math stays whole.
    return ''.join(f'<p>{line}</p>' for line in _report_inline_math(text.strip()).split('<br>') if line.strip())


def _report_inline_math(text: str) -> str:
    """Keep authored TeX intact; convert only explicit linear math notation."""
    explicit = re.compile(r'(?<!\\)(\$\$[\s\S]+?\$\$|\\\[[\s\S]+?\\\]|\\\([\s\S]+?\\\)|\$(?:\\.|[^$])+?\$)')
    math_chars = r'A-Za-z0-9∠△°√∛±²³=＝＋+−－÷×·*/:：.()（）^_≤≥≠⊥∥?\-'
    tokens = re.compile(r'[A-Za-z0-9∠△°√∛±(（][' + math_chars + r']*(?:[ \t]+[' + math_chars + r']+)*')

    def plain(value):
        return _esc(value).replace('\n', '<br>')

    def linear(match):
        value = match.group(0)
        # A Chinese explanatory parenthesis immediately after an expression
        # belongs to the prose (e.g. √18（字迹模糊）), not its radicand.
        suffix = '（' if value.endswith('（') else ''
        if suffix:
            value = value[:-1]
        if not re.search(r'[√∛∠△°=＝÷×⊥∥²³≤≥≠±]', value):
            return plain(value + suffix)
        tex = _report_linear_tex(value)
        return (_report_math_span(tex, value) if tex else plain(value)) + plain(suffix)

    result, end = [], 0
    for match in explicit.finditer(text):
        # Escape prose before substituting formulas, never regex over HTML.
        preceding, cursor = text[end:match.start()], 0
        for token in tokens.finditer(preceding):
            result.extend((plain(preceding[cursor:token.start()]), linear(token)))
            cursor = token.end()
        result.append(plain(preceding[cursor:]))
        authored = match.group(0)
        width = 1 if authored.startswith('$') and not authored.startswith('$$') else 2
        result.append(_report_math_span(authored[width:-width], authored,
                                       display=authored.startswith(('$$', r'\['))))
        end = match.end()
    remaining, cursor = text[end:], 0
    for token in tokens.finditer(remaining):
        result.extend((plain(remaining[cursor:token.start()]), linear(token)))
        cursor = token.end()
    result.append(plain(remaining[cursor:]))
    return ''.join(result)


def _report_math_span(tex: str, fallback: str, *, display: bool = False) -> str:
    return (f'<span class="qm{" qm-display" if display else ""}" data-latex="{_esc(tex)}"'
            f'{" data-display=\"true\"" if display else ""}>{_esc(fallback)}</span>')


def _report_linear_tex(value: str) -> str | None:
    """Parse grouping/root boundaries without calculating or simplifying answers.

    Parenthesized denominators are explicit; ambiguous a/bc remains a/bc.
    A bare digit before √ is a coefficient. Root indices come from native
    Word math, authored TeX, or the unambiguous ∛ character.
    """
    source = value.translate(str.maketrans({'＝':'=', '＋':'+', '－':'-', '−':'-', '（':'(', '）':')', '：':':'}))
    symbols = {'∠':r'\angle ', '△':r'\triangle ', '°':r'^{\circ}',
               '÷':r'\div ', '×':r'\times ', '·':r'\cdot ', '⊥':r'\perp ', '∥':r'\parallel ',
               '²':'^{2}', '³':'^{3}', '±':r'\pm ', '≤':r'\leq ', '≥':r'\geq ', '≠':r'\neq '}

    def atom(pos, depth):
        if depth > 24 or pos >= len(source):
            raise ValueError('incomplete formula')
        start, char = pos, source[pos]
        if char == '(':
            inner, pos = expression(pos + 1, depth + 1, True)
            if pos >= len(source) or source[pos] != ')':
                raise ValueError('unbalanced formula')
            return '(' + inner + ')', inner, pos + 1
        if char in '√∛':
            wrapped, bare, pos = atom(pos + 1, depth + 1)
            tex = (r'\sqrt[3]{' if char == '∛' else r'\sqrt{') + bare + '}'
            return tex, tex, pos
        if char.isdigit():
            match = re.match(r'\d+(?:\.\d+)?', source[pos:])
            pos += len(match.group(0))
        elif char.isalpha():
            match = re.match(r'[A-Za-z]+', source[pos:])
            if not match:
                raise ValueError('unknown atom')
            pos += len(match.group(0))
        else:
            raise ValueError('unknown atom')
        text = source[start:pos]
        if pos < len(source) and source[pos] == '°':
            text += symbols['°']
            pos += 1
        return text, text, pos

    def expression(pos, depth, nested=False):
        parts = []
        while pos < len(source):
            char = source[pos]
            if char == ')':
                if nested:
                    break
                raise ValueError('unbalanced formula')
            # ``str.isalnum()`` also matches superscript characters such as
            # ²/³, but those are operators handled by ``symbols`` below and
            # cannot be consumed by the ASCII digit/letter regexes in atom().
            if ('0' <= char <= '9' or 'A' <= char <= 'Z' or
                    'a' <= char <= 'z' or char in '(√∛'):
                wrapped, bare, pos = atom(pos, depth)
                if pos < len(source) and source[pos] == '/':
                    denominator_start = pos + 1
                    try:
                        dw, db, after = atom(denominator_start, depth)
                    except ValueError:
                        parts.append(wrapped)
                        continue
                    explicit_denominator = source[denominator_start] == '(' or source[denominator_start] in '√∛'
                    simple_denominator = len(db) == 1 or db.replace('.', '').isdigit()
                    adjacent = after < len(source) and (source[after].isalnum() or source[after] in '(√∛')
                    if explicit_denominator or (simple_denominator and not adjacent):
                        wrapped = r'\frac{' + bare + '}{' + db + '}'
                        pos = after
                parts.append(wrapped)
            else:
                parts.append(symbols.get(char, char))
                pos += 1
        return ''.join(parts), pos
    try:
        return expression(0, 0)[0]
    except (ValueError, RecursionError):
        return None


def _report_display_text(value: object) -> str:
    return re.sub(
        r"\bQ\d+(?:\s*\(P?\d+\))?",
        lambda match: _question_display_label(match.group(0)),
        str(value or ""), flags=re.IGNORECASE,
    )


def _question_analysis_html(item: dict[str, Any] | None) -> str:
    if not item:
        return _ai_block("", failed=True)
    feedback = _narrative_text(item.get("feedback")) or "\n".join(dict.fromkeys(
        text for key in ("observation", "possible_cause", "analysis")
        if (text := _narrative_text(item.get(key)))
    ))
    review = _narrative_text(item.get("review_note"))
    body = f'<div class="feedback"><b>本题反馈</b>{_report_paragraphs(feedback)}</div>' if feedback else ""
    if review:
        body += f'<div class="review-note"><b>报告分析提示 · 建议核对</b>{_report_paragraphs(review)}</div>'
    steps = _narrative_items(item, "solution_steps")
    if steps:
        label = "参考解法" if item.get("solution_source") == "reference" else "AI 参考解法"
        body += f'<div class="solution"><b>{label}</b><ol>' + "".join(
            f'<li>{_report_paragraphs(step)}</li>' for step in steps if isinstance(step, str) and step.strip()
        ) + '</ol></div>'
        detail = _narrative_text(item.get("full_solution"))
        if detail:
            body += '<details class="solution-detail"><summary>查看完整解法</summary>' + _report_paragraphs(detail) + '</details>'
    # Compatibility for previously generated narratives; keep supporting evidence
    # available without restoring five repetitive rows to the main report.
    if "feedback" not in item:
        details = "\n".join(_narrative_text(item.get(key)) for key in ("evidence", "verification"))
        if details.strip():
            body += '<details><summary>查看分析依据</summary>' + _report_paragraphs(details) + '</details>'
    if not body:
        return _ai_block("", failed=True)
    return (
        '<div class="compact-analysis">' + body + '</div>'
    )


def _report_stem_html(info: _QuestionInfo | None, fallback: str) -> str:
    """Render source table cells as cells, instead of model-input pipe text."""
    if info is None or not info.question_markup:
        return _report_paragraphs(fallback.replace("[图片]", ""))
    from html.parser import HTMLParser
    from backend.document_parsing.question_blocks import _INLINE_IMAGE_MARKER

    class StemMarkup(HTMLParser):
        tags = {"p", "div", "br", "table", "thead", "tbody", "tr", "td", "th", "sup", "sub", "b", "strong", "em", "span"}

        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts = []
            self.math_stack = []

        def handle_starttag(self, tag, attrs):
            if tag in self.tags:
                spans = "".join(f' {key}="{value}"' for key, value in attrs
                                if tag in {"td", "th"} and key in {"colspan", "rowspan"}
                                and value and value.isdigit())
                latex = dict(attrs).get("data-latex")
                if tag == "span" and latex:
                    spans += f' class="qm" data-latex="{_esc(latex)}"'
                if tag != 'br':
                    self.math_stack.append(bool(latex))
                self.parts.append(f"<{tag}{spans}>")

        def handle_endtag(self, tag):
            if tag in self.tags and tag != "br":
                if self.math_stack:
                    self.math_stack.pop()
                self.parts.append(f"</{tag}>")

        def handle_data(self, value):
            if value.strip():
                value = value.replace("[图片]", "")
                self.parts.append(_esc(value) if any(self.math_stack) else _report_inline_math(value))

    parser = StemMarkup()
    parser.feed(_INLINE_IMAGE_MARKER.sub("", info.question_markup))
    return "".join(parser.parts)


@lru_cache(maxsize=1)
def _report_knowledge_assets() -> tuple[str, str]:
    root = Path(__file__).resolve().parent / 'backend' / 'report_assets'
    return ((root / 'personal_knowledge.css').read_text(encoding='utf-8'),
            (root / 'personal_knowledge.js').read_text(encoding='utf-8'))


@lru_cache(maxsize=1)
def _report_math_assets() -> str:
    """Embed the same KaTeX engine as the question bank, including offline fonts."""
    root = Path(__file__).resolve().parent / "backend" / "report_assets" / "katex"
    css = (root / "katex.min.css").read_text(encoding="utf-8")
    def font_source(match):
        font = root / match.group(1)
        return 'src:url(data:font/woff2;base64,' + base64.b64encode(font.read_bytes()).decode() + ') format("woff2")'
    css = re.sub(r'src:url\((fonts/[^)]+\.woff2)\)[^;}]*(?=[;}])', font_source, css)
    js = (root / "katex.min.js").read_text(encoding="utf-8")
    return '<style>' + css + '</style><script>' + js.replace('</script', '<\\/script') + '</script>'


def _question_cell_label(question_id: str) -> str:
    """Q14(P8) → 14(8)；Q7 → 7。"""
    coordinates = question_id_coordinates(question_id)
    if coordinates is None:
        return str(question_id)
    parent, part = coordinates
    return f"{parent}({part})" if part is not None else str(parent)


def _report_class_label(class_name: str) -> str:
    name = str(class_name or "").strip() or "未分班"
    return name if "班" in name else f"{name}班"


def _personal_cell_status(record: _StudentQuestionRecord) -> str:
    if record.score == 0 and (
        record.error_category == "未作答"
        or "未作答" in (record.deduction_reason or "")
        or (record.error_summary or "") in {"blank", "blank_or_no_valid_work"}
    ):
        return "blank"
    if record.max_score > 0 and record.score >= record.max_score:
        return "full"
    if record.score == 0:
        return "zero"
    return "part"


def _lost_group_label(
    parent: str,
    student: _StudentReportData,
    info_by_qid: dict[str, _QuestionInfo],
) -> str:
    """「第14题计算」形式：大题显示名 + 题型标签。"""
    info = info_by_qid.get(parent)
    if info is None:
        info = next(
            (
                info_by_qid.get(r.question_id)
                for r in student.records
                if _parent_question_id(r.question_id) == parent
            ),
            None,
        )
    type_label = _question_type_label(info.question_type if info is not None else "")
    return _question_display_label(parent) + type_label


def _score_card_conclusion(
    student: _StudentReportData,
    info_by_qid: dict[str, _QuestionInfo],
) -> str:
    scored = [r for r in student.records if r.max_score > 0]
    n = len(scored)
    full = sum(1 for r in scored if r.score >= r.max_score)
    lost = [r for r in scored if r.lost]
    if not lost:
        return f"全部 {n} 个小题都拿到满分。"
    by_parent: dict[str, float] = {}
    for r in lost:
        parent = _parent_question_id(r.question_id)
        by_parent[parent] = by_parent.get(parent, 0.0) + r.lost_points
    ordered = sorted(
        by_parent.items(),
        key=lambda kv: (-kv[1], question_id_coordinates(kv[0]) or (10**9, 0)),
    )
    total_lost = sum(by_parent.values())
    if len(ordered) == 1:
        return (
            f"{n} 个小题中 {full} 个满分，丢的 {_fmt_num(total_lost)} 分都在"
            f"{_lost_group_label(ordered[0][0], student, info_by_qid)}。"
        )
    (g1, x), (g2, y) = ordered[0], ordered[1]
    return (
        f"{n} 个小题中 {full} 个满分，丢分最多的是"
        f"{_lost_group_label(g1, student, info_by_qid)}（{_fmt_num(x)}分）和"
        f"{_lost_group_label(g2, student, info_by_qid)}（{_fmt_num(y)}分）。"
    )


def _narrative_analysis_index(
    narrative: dict[str, Any] | None,
    info_by_qid: dict[str, _QuestionInfo],
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for item in _narrative_items(narrative, "question_analyses"):
        if not isinstance(item, dict):
            continue
        raw = str(item.get("question_id") or "").strip()
        result[resolve_known_question_id(raw, info_by_qid) or raw] = item
    return result


def _history_chart_html(entries: list[dict[str, Any]], current_sid: int) -> str:
    """历次成绩：SVG 只画折线（不缩放描边），圆点/分数/场次名用 HTML 等宽列。"""
    n = len(entries)

    def x_of(i: int) -> float:
        return (i + 0.5) / n * 100

    def y_of(ratio: float) -> float:
        return 16 + (1 - min(1.0, max(0.0, ratio))) * 76  # 顶部留分数值空间，纵轴固定 0..满分

    avg_pts: list[tuple[float, float]] = []
    me_segs: list[list[tuple[float, float]]] = []
    cur_seg: list[tuple[float, float]] = []
    dots: list[tuple[int, float, float, float]] = []
    for i, e in enumerate(entries):
        full = e.get("full") or 100
        avg = e.get("avg")
        if avg is not None:
            avg_pts.append((x_of(i), y_of(float(avg) / full)))
        if e.get("score") is not None:
            cur_seg.append((x_of(i), y_of(float(e["score"]) / full)))
            dots.append((i, x_of(i), y_of(float(e["score"]) / full), float(e["score"])))
        else:
            if len(cur_seg) > 1:
                me_segs.append(cur_seg)
            cur_seg = []
    if len(cur_seg) > 1:
        me_segs.append(cur_seg)

    svg = ['<svg viewBox="0 0 100 100" preserveAspectRatio="none" aria-hidden="true">']
    if len(avg_pts) > 1:
        pts = " ".join(f"{x:.2f},{y:.2f}" for x, y in avg_pts)
        svg.append(
            f'<polyline points="{pts}" fill="none" stroke="#8a97a8" stroke-width="1.5" '
            f'stroke-dasharray="4 3" vector-effect="non-scaling-stroke"/>'
        )
    for seg in me_segs:
        pts = " ".join(f"{x:.2f},{y:.2f}" for x, y in seg)
        svg.append(
            f'<polyline points="{pts}" fill="none" stroke="#2b6cb0" stroke-width="2" '
            f'vector-effect="non-scaling-stroke"/>'
        )
    svg.append("</svg>")

    overlays = []
    for i, x, y, score in dots:
        cur = entries[i]["session_id"] == current_sid
        overlays.append(
            f'<span class="hdot{" cur" if cur else ""}" style="left:{x:.2f}%;top:{y:.2f}%"></span>'
            f'<span class="hval{" cur" if cur else ""}" style="left:{x:.2f}%;top:{y:.2f}%">{_esc(_fmt_num(score))}</span>'
        )

    cols = []
    for e in entries:
        name = str(e.get("name") or "")
        short = name if len(name) <= 6 else name[:6] + "…"
        rank = e.get("rank")
        rank_text = f"第{rank}名" if e.get("score") is not None and rank else (
            "有成绩" if e.get("score") is not None else "未参加"
        )
        cur = e["session_id"] == current_sid
        cols.append(
            f'<div class="hcol{" cur" if cur else ""}">'
            f'<div class="hn">{_esc(short)}</div><div>{_esc(rank_text)}</div></div>'
        )
    return (
        '<div class="hchart">'
        + "".join(svg)
        + "".join(overlays)
        + '</div><div class="hcols">'
        + "".join(cols)
        + "</div>"
    )


def _point_question_result(
    student: _StudentReportData,
    key: str,
    question: dict[str, Any],
    info_by_qid: dict[str, _QuestionInfo],
    session_id: int,
) -> str:
    """考查点单题结果：有本场判定点观测取 achieved 均值，否则用得分率。"""
    mastery = student.knowledge_mastery.get(key) or {}
    for ref in mastery.get("source_question_refs") or []:
        if int(ref.get("session_id") or 0) != session_id:
            continue
        raw = str(ref.get("question_id") or "")
        if (resolve_known_question_id(raw, info_by_qid) or raw) != question["id"]:
            continue
        obs = [
            o
            for o in ((ref.get("assessment") or {}).get("point_observations") or [])
            if isinstance(o, dict) and o.get("stable_key") == key
        ]
        if obs:
            mean = sum(float(o.get("achieved") or 0) for o in obs) / len(obs)
            return "full" if mean == 1 else "zero" if mean == 0 else "part"
    if question["full"] > 0 and question["score"] >= question["full"]:
        return "full"
    if question["score"] == 0:
        return "zero"
    return "part"


def _point_recurrence_sessions(
    student: _StudentReportData, key: str, session_id: int
) -> list[str]:
    """「以前也有失分」徽章：返回命中过该点失分的历史场次名（去重）。"""
    names: list[str] = []
    mastery = student.knowledge_mastery.get(key) or {}
    for ref in mastery.get("source_question_refs") or []:
        if int(ref.get("session_id") or 0) == session_id:
            continue
        obs = [
            o
            for o in ((ref.get("assessment") or {}).get("point_observations") or [])
            if isinstance(o, dict) and o.get("stable_key") == key
        ]
        if obs:
            hit = any(float(o.get("achieved") or 0) < 1 for o in obs)
        else:
            hit = float(ref.get("score_awarded") or 0) < float(ref.get("full_score") or 0)
        if hit:
            name = str(ref.get("session_name") or "").strip()
            if name and name not in names:
                names.append(name)
    return names


def _personal_exam_points(
    view: dict[str, Any],
    student: _StudentReportData,
    info_by_qid: dict[str, _QuestionInfo],
    session_id: int,
) -> list[dict[str, Any]]:
    """本次考查点集合：全部 skill 节点 + 有未覆盖小题的 topic 节点。"""
    nodes = view.get("nodes") or []
    skill_covered = {
        q["id"] for n in nodes if n.get("kind") == "skill" for q in n.get("questions") or []
    }
    points = []
    for node in nodes:
        kind = node.get("kind")
        if kind == "skill":
            questions = list(node.get("questions") or [])
        elif kind == "topic":
            questions = [
                q for q in node.get("questions") or [] if q["id"] not in skill_covered
            ]
            if not questions:
                continue
        else:
            continue
        results = [
            (q, _point_question_result(student, node["key"], q, info_by_qid, session_id))
            for q in questions
        ]
        syms = [r for _q, r in results]
        if all(s == "full" for s in syms):
            status = "good"
        elif all(s == "zero" for s in syms):
            status = "bad"
        else:
            status = "mid"
        points.append(
            {
                "key": node["key"],
                "label": node["label"],
                "section": node.get("section") or "",
                "questions": questions,
                "results": results,
                "status": status,
                "recur": _point_recurrence_sessions(student, node["key"], session_id),
            }
        )
    return points


def _points_for_question(points: list[dict[str, Any]], qid: str) -> list[str]:
    labels: list[str] = []
    for p in points:
        if any(q["id"] == qid for q in p["questions"]) and p["label"] not in labels:
            labels.append(p["label"])
    return labels


_CARD_QSEG_RE = re.compile(r"第([0-9０-９、，,和及与\s()（）\-—~～至]+?)题")
_CARD_QGROUP_RE = re.compile(
    r"(\d+)((?:\s*[（(]\s*\d+\s*[）)](?:\s*[-—~～至]\s*[（(]?\s*\d*\s*[）)]?)?)*)"
)
_FULLWIDTH_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")
_CARD_TYPE_WORDS = {
    "选择题": {"choice", "single_choice", "multiple_choice"},
    "填空题": {"fill_blank", "fill_in_blank"},
}


def _part_numbers(group_text: str) -> list[int]:
    """(2)(6)(7)(8) → [2,6,7,8]；(1)-(4)/(1)至(4) → [1,2,3,4]。"""
    nums: list[int] = []
    prev = None
    pos = 0
    for match in re.finditer(r"\d+", group_text):
        sep = group_text[pos : match.start()]
        num = int(match.group(0))
        if prev is not None and re.search(r"[-—~～至]", sep):
            nums.extend(range(prev + 1, num + 1))
        else:
            nums.append(num)
        prev = num
        pos = match.end()
    return nums


def _follow_up_question_ids(
    problem: dict[str, Any],
    suggestion: dict[str, Any] | None,
    student: _StudentReportData,
    info_by_qid: dict[str, _QuestionInfo],
) -> list[str]:
    """重点跟进卡片的关联题号：优先叙述里的 question_ids，否则按正文文字解析。

    两种来源都只保留该生有失分记录的题号；全部为空时调用方隐藏元信息行。
    """
    lost_ids = {r.question_id for r in student.records if r.lost}
    out: list[str] = []

    def add(qid: str) -> None:
        if qid in lost_ids and qid not in out:
            out.append(qid)

    raw_ids = problem.get("question_ids")
    if isinstance(raw_ids, (list, tuple)):
        for raw in raw_ids:
            text = str(raw).strip()
            if not text:
                continue
            add(resolve_known_question_id(text, info_by_qid) or text)
    if out:
        return out

    # 兼容旧版叙述：正文里的「第…题」片段 + 「选择题/填空题」类别。
    text = (
        str(problem.get("title") or "")
        + str(problem.get("detail") or "")
        + str((suggestion or {}).get("detail") or "")
    ).translate(_FULLWIDTH_DIGITS)

    def add_parent(num: str) -> None:
        parent = resolve_known_question_id(f"Q{num}", info_by_qid) or f"Q{num}"
        for r in student.records:
            if _parent_question_id(r.question_id) == parent:
                add(r.question_id)

    for seg in _CARD_QSEG_RE.finditer(text):
        for grp in _CARD_QGROUP_RE.finditer(seg.group(1)):
            parent_num, part_group = grp.group(1), grp.group(2)
            parts = _part_numbers(part_group) if part_group.strip() else []
            if parts:
                for p in parts:
                    raw = f"Q{parent_num}(P{p})"
                    add(resolve_known_question_id(raw, info_by_qid) or raw)
            else:
                add_parent(parent_num)
    remainder = _CARD_QSEG_RE.sub(" ", text)
    for word, types in _CARD_TYPE_WORDS.items():
        if word in remainder:
            for r in student.records:
                if not r.lost:
                    continue
                info = info_by_qid.get(r.question_id)
                if info is not None and info.question_type in types:
                    add(r.question_id)
    return out


def _shots_for_parent(shots: dict[str, dict[str, str]], parent: str) -> list[dict[str, str]]:
    return [
        shot
        for key, shot in shots.items()
        if (shot.get("parent_question_id") or key) == parent
    ]


def _shots_html(shots: dict[str, dict[str, str]], parent: str, title: str) -> str:
    return "".join(
        f'<div class="shot"><img src="{shot["data_uri"]}" alt="{_esc(title)}作答截图">'
        f'<div class="cap">{_esc(shot.get("caption") or "学生作答（原卷截图）")}</div></div>'
        for shot in _shots_for_parent(shots, parent)
    )


# 教师确认流程写入 teacher_comment 的占位文案：不算真实批语，不向家长展示。
_TEACHER_COMMENT_PLACEHOLDERS = frozenset(
    {REVIEW_CONFIRMED_REASON, "教师已确认", "教师已确认最终分", "已复核"}
)


def _real_teacher_comment(record: _StudentQuestionRecord) -> str:
    """真实老师批语；占位确认文案视为无批语。"""
    comment = (record.teacher_comment or "").strip()
    return "" if comment in _TEACHER_COMMENT_PLACEHOLDERS else comment


def _grading_fallback_html(
    record: _StudentQuestionRecord,
    info: _QuestionInfo | None,
    heading: str = "",
    show_answer: bool = True,
    show_comment: bool = True,
) -> str:
    """无 AI 叙述项时的降级块：真实老师批语 + 参考答案。

    不展示 deduction_reason/error_summary 等原始批改记录与确认占位文案；
    两者都没有时不输出块。
    """
    comment = _real_teacher_comment(record) if show_comment else ""
    answer = (info.canonical_answer if info is not None else "") or ""
    if not comment and not (show_answer and answer):
        return ""
    head = f"<b>{_esc(heading)}</b>" if heading else ""
    body = ""
    if comment:
        body += f"<div>老师批语：{_report_inline_math(comment)}</div>"
    if show_answer and answer:
        body += f'<div>参考答案：<b class="ans">{_report_inline_math(answer)}</b></div>'
    return f'<div class="qd-fb">{head}{body}</div>'


def _analysis_or_fallback_html(
    record: _StudentQuestionRecord,
    info: _QuestionInfo | None,
    analysis: dict[str, Any] | None,
    *,
    show_answer: bool = True,
    show_comment: bool = True,
) -> str:
    """AI 叙述项为空或没有可展示内容时退化为老师批语/参考答案，不显示生成失败提示。"""
    if analysis:
        rendered = _question_analysis_html(analysis)
        if AI_FAILED_NOTE not in rendered:
            return rendered
    return _grading_fallback_html(
        record, info, show_answer=show_answer, show_comment=show_comment
    )


def _stem_block_html(info: _QuestionInfo | None, fallback: str, title: str) -> str:
    html_text = f'<div class="stem">{_report_stem_html(info, fallback)}</div>'
    if info is not None:
        figures = "".join(
            f'<img src="data:image/png;base64,{base64.b64encode(blob).decode()}" alt="{_esc(title)}题图">'
            for role, blob in info.reference_images
            if role == "question"
        )
        if figures:
            html_text += f'<div class="reference-figures">{figures}</div>'
    return html_text


def _question_detail_template(
    record: _StudentQuestionRecord,
    info: _QuestionInfo | None,
    points_labels: list[str],
    analysis: dict[str, Any] | None,
    shots: dict[str, dict[str, str]],
    errors: list[dict[str, Any]] | None = None,
) -> str:
    """答题一览方格对应的隐藏详情片段（<template>，点击方格原位展开）。"""
    parent = _parent_question_id(record.question_id)
    title = _question_display_label(record.question_id)
    type_label = _question_type_label(info.question_type if info is not None else "")
    avg = info.class_avg if info is not None else None
    head_parts = [title]
    if type_label:
        head_parts.append(type_label)
    head_parts.append(f"得 {_fmt_num(record.score)}/{_fmt_num(record.max_score)}")
    if avg is not None:
        head_parts.append(f"全班平均 {_fmt_num(round(avg, 1))}")
    head = " · ".join(head_parts)
    body = [f'<div class="qdetail-box"><div class="qd-head">{_esc(head)}</div>']
    if record.lost and errors:
        labels = "；".join(dict.fromkeys(
            " · ".join(
                part for part in (str(row.get("category") or ""), str(row.get("pattern") or ""))
                if part
            )
            for row in errors
        ))
        if labels:
            body.append(
                f'<div class="errline">错误归类（AI 辅助）：{_esc(labels)}</div>'
            )
    if points_labels:
        tags = "".join(f'<span class="pt-tag">{_esc(t)}</span>' for t in points_labels)
        body.append(f'<div class="qd-points">考这些：{tags}</div>')
    if record.lost:
        feedback = _narrative_text((analysis or {}).get("feedback"))
        if analysis and feedback:
            body.append(f'<div class="qd-fb">{_report_paragraphs(feedback)}</div>')
        more = ['<div class="qmore" hidden>']
        fallback = (
            (info.question_text or info.stem_summary) if info is not None else ""
        ) or title
        parent_title = _question_display_label(parent)
        more.append(_stem_block_html(info, fallback, parent_title))
        more.append(_shots_html(shots, parent, parent_title))
        kv = []
        if record.student_answer:
            kv.append(f"<dt>学生作答</dt><dd>{_report_inline_math(record.student_answer)}</dd>")
        elif _personal_cell_status(record) == "blank":
            kv.append("<dt>学生作答</dt><dd>未作答</dd>")
        if kv:
            more.append(f'<dl class="kv">{"".join(kv)}</dl>')
        more.append(_analysis_or_fallback_html(record, info, analysis))
        more.append("</div>")
        body.append('<button type="button" class="qmore-btn">看原卷和解法</button>')
        body.append("".join(more))
    body.append("</div>")
    return "".join(body)


def _lost_appendix_html(
    student: _StudentReportData,
    info_by_qid: dict[str, _QuestionInfo],
    analysis_by_qid: dict[str, dict[str, Any]],
    shots: dict[str, dict[str, str]],
    error_map: dict[str, list[dict[str, Any]]] | None = None,
) -> str:
    """失分题详解附录：全部失分小问按大题分组，默认收起、打印展开。"""
    lost = [r for r in student.records if r.lost]
    if not lost:
        return ""
    groups: dict[str, list[_StudentQuestionRecord]] = {}
    for r in lost:
        groups.setdefault(_parent_question_id(r.question_id), []).append(r)
    ordered = sorted(
        groups.items(), key=lambda kv: question_id_coordinates(kv[0]) or (10**9, 0)
    )
    cards = []
    for parent, recs in ordered:
        info = info_by_qid.get(parent) or next(
            (info_by_qid.get(r.question_id) for r in recs), None
        )
        parent_records = [
            r for r in student.records if _parent_question_id(r.question_id) == parent
        ]
        gscore = sum(r.score for r in parent_records)
        gmax = sum(r.max_score for r in parent_records)
        title = _question_display_label(parent)
        type_label = _question_type_label(info.question_type if info is not None else "")
        fallback = ((info.question_text or info.stem_summary) if info is not None else "") or title
        parts = [
            f'<div class="qcard"><div class="head"><b>{_esc(title)}'
            f'{(" · " + _esc(type_label)) if type_label else ""}</b>'
            f'<span class="score">得 {_fmt_num(gscore)} 分 / 满分 {_fmt_num(gmax)} 分</span></div>',
            _stem_block_html(info, fallback, title),
            _shots_html(shots, parent, title),
        ]
        for record in sorted(recs, key=lambda r: question_id_coordinates(r.question_id) or (10**9, 0)):
            part_info = info_by_qid.get(record.question_id)
            analysis = analysis_by_qid.get(record.question_id)
            if (
                analysis
                and analysis.get("solution_source") == "reference"
                and not (part_info and part_info.reference_analysis)
            ):
                analysis = {**analysis, "solution_source": "ai"}
            kv = []
            if record.student_answer:
                kv.append(f"<dt>学生作答</dt><dd>{_report_inline_math(record.student_answer)}</dd>")
            elif _personal_cell_status(record) == "blank":
                kv.append("<dt>学生作答</dt><dd>未作答</dd>")
            answer = part_info.canonical_answer if part_info is not None else ""
            if answer and not (analysis and analysis.get("solution_steps")):
                kv.append(f'<dt>标准答案</dt><dd><b class="ans">{_report_inline_math(answer)}</b></dd>')
            teacher_comment = _real_teacher_comment(record)
            if teacher_comment:
                kv.append(f"<dt>老师批语</dt><dd>{_report_inline_math(teacher_comment)}</dd>")
            error_rows = (error_map or {}).get(record.question_id) or []
            error_text = "；".join(dict.fromkeys(
                " · ".join(
                    part for part in (str(row.get("category") or ""), str(row.get("pattern") or ""))
                    if part
                )
                for row in error_rows
            ))
            if error_text:
                kv.append(f"<dt>错误归类</dt><dd>{_esc(error_text)}（AI 辅助）</dd>")
            heading = (
                f'<h3>{_esc(_question_display_label(record.question_id))}'
                f'<span>得 {_fmt_num(record.score)} / {_fmt_num(record.max_score)} 分</span></h3>'
                if len(parent_records) > 1
                else ""
            )
            # kv 中已有标准答案与老师批语行，降级块不重复输出
            analysis_html = _analysis_or_fallback_html(
                record, part_info, analysis, show_answer=False, show_comment=False
            )
            parts.append(
                f'<div class="qpart">{heading}'
                + (f'<dl class="kv">{"".join(kv)}</dl>' if kv else "")
                + analysis_html
                + "</div>"
            )
        parts.append("</div>")
        cards.append("".join(parts))
    return (
        f'<details class="appendix"><summary>全部失分题详解（{len(lost)} 道小题）</summary>'
        + "".join(cards)
        + "</details>"
    )


def _load_student_histories(
    repositories: GradingRepositoryAccess,
    data: _SessionAnalysisData,
    data_root: Path | None,
) -> dict[int, list[dict[str, Any]]]:
    """同教学学期内、不晚于本场的历次成绩，每次导出只读取一轮。

    任何历史读取异常都按无历史处理，不阻断报告导出。
    """
    try:
        session_row = (
            repositories.sessions.get_grading_session(data.session_id) or {}
        )
        volume = str(session_row.get("curriculum_volume_id") or "").strip()
        if not volume:
            return {}
        candidate_ids = [
            int(row["id"])
            for row in repositories.sessions.list_grading_sessions()
            if str(row.get("curriculum_volume_id") or "").strip() == volume
        ]
        assembled: list[tuple[int, _SessionAnalysisData, dict[int, tuple]]] = []
        for sid in candidate_ids:
            try:
                session_data = assemble_session_analysis(
                    repositories, sid, data_root=data_root, page_only=True
                )
            except Exception:
                continue
            if (
                sid != data.session_id
                and data.graded_at
                and session_data.graded_at
                and session_data.graded_at > data.graded_at
            ):
                continue
            by_student = {
                s.student_id: (group, s)
                for group in split_session_analysis_by_class(session_data).values()
                for s in group.students
            }
            assembled.append((sid, session_data, by_student))
        assembled.sort(key=lambda item: (item[1].graded_at or "\uffff", item[0]))
        histories: dict[int, list[dict[str, Any]]] = {
            s.student_id: [] for s in data.students
        }
        for sid, session_data, by_student in assembled:
            for student in data.students:
                entry: dict[str, Any] = {
                    "session_id": sid,
                    "name": session_data.session_name,
                    "full": session_data.full_score,
                    "avg": session_data.stats.get("avg"),
                    "score": None,
                    "rank": None,
                    "present": None,
                }
                hit = by_student.get(student.student_id)
                if hit is not None:
                    group, matched = hit
                    entry.update(
                        score=matched.student_score,
                        rank=matched.rank,
                        present=group.present,
                        avg=group.stats.get("avg"),
                    )
                histories[student.student_id].append(entry)
        return histories
    except Exception:
        return {}


def _render_personal_html(
    data: _SessionAnalysisData,
    student: _StudentReportData,
    narrative: dict[str, Any] | None,
    shots: dict[str, dict[str, str]],
    history: list[dict[str, Any]] | None = None,
    error_map: dict[str, list[dict[str, Any]]] | None = None,
    error_history: dict[str, dict[str, list[str]]] | None = None,
) -> str:
    info_by_qid = {info.question_id: info for info in data.questions}
    analysis_by_qid = _narrative_analysis_index(narrative, info_by_qid)
    history = [e for e in (history or []) if isinstance(e, dict)]
    error_map = error_map or {}
    error_history = error_history or {}

    def question_errors(question_id: str) -> list[dict[str, Any]]:
        return error_map.get(question_id) or []

    def error_label(question_id: str) -> str:
        """本题错误归类文案：大类 · 典型错法；多条用「；」连接。"""
        parts = []
        for row in question_errors(question_id):
            text = " · ".join(
                part for part in (str(row.get("category") or ""), str(row.get("pattern") or "")) if part
            )
            if text and text not in parts:
                parts.append(text)
        return "；".join(parts)

    def error_recurrence(question_ids: list[str]) -> list[str]:
        """卡片关联题在以往场次出现过同大类/同错法的场次名。"""
        names: list[str] = []
        for qid in question_ids:
            for row in question_errors(qid):
                for key, value in (("categories", row.get("category")), ("patterns", row.get("pattern"))):
                    for name in (error_history.get(key) or {}).get(str(value or ""), []):
                        if name and name not in names:
                            names.append(name)
        return names
    points = _personal_exam_points(
        _personal_knowledge_view(data, student), student, info_by_qid, data.session_id
    )

    # ---- A 成绩卡 ----
    meta = (
        f"{_esc(data.session_name)} · {_esc(data.subject)} · "
        f"{_esc(_report_class_label(student.class_name))} · {_esc(student.graded_at)}"
    )
    prev_v, prev_l = "首次记录", "上次"
    prior = [
        e
        for e in history
        if e["session_id"] != data.session_id and e.get("score") is not None
    ]
    if prior:
        prev = prior[-1]
        prev_v = f"{_fmt_num(prev['score'])}分"
        if prev.get("rank") and prev.get("present"):
            prev_l = f"上次 第{prev['rank']}/{prev['present']}名"
    avg = data.stats.get("avg")
    avg_text = f"{float(avg):.1f}" if avg is not None else "—"
    bands = list(data.stats.get("bands") or [])
    scale = (data.full_score / 100) if data.full_score > 0 else 1.0
    my_band = None
    for i, cutoff in enumerate(_BAND_CUTOFFS):
        lower = cutoff * scale
        upper = data.full_score if i == 0 else _BAND_CUTOFFS[i - 1] * scale
        if lower <= student.student_score <= upper if i == 0 else lower <= student.student_score < upper:
            my_band = len(bands) - 1 - i
            break
    max_count = max((b["count"] for b in bands), default=1) or 1
    band_rows = []
    for disp_i, band in enumerate(reversed(bands)):  # 低分段在上
        width = band["count"] / max_count * 100
        me = disp_i == my_band
        band_rows.append(
            f'<div class="band{" me" if me else ""}"><span>{_esc(band["label"].replace(" ", ""))}'
            f'{"<span class=me-tag>孩子在这里</span>" if me else ""}</span>'
            f'<div class="track"><i style="width:{width:.0f}%"></i></div>'
            f'<span class="cnt">{band["count"]}人</span></div>'
        )
    review = (
        '<div class="review-banner">有题目等待老师复核，分数可能微调</div>'
        if student.needs_review
        else ""
    )
    strength = ""
    strengths = [
        _narrative_text(i)
        for i in _narrative_items(narrative, "strengths")
        if _narrative_text(i)
    ]
    if strengths:
        strength = f'<div class="strength">✓ 值得肯定：{_report_inline_math(strengths[0])}</div>'
    small_sample_note = (
        '<div class="note" style="margin-top:6px">本场参考人数较少，班级对比仅供参考。</div>'
        if data.small_sample
        else ""
    )
    card_a = f"""
<div class="card">
  {review}
  <div class="meta">{meta}</div>
  <div class="name">{_esc(student.student_name)}</div>
  <div class="score-line"><span class="big">{_fmt_num(student.student_score)}</span><span class="of">/ {_fmt_num(data.full_score)}</span></div>
  <div class="stat3">
    <div class="cell"><div class="v">{student.rank}/{data.present}</div><div class="l">班级名次</div></div>
    <div class="cell"><div class="v">{avg_text}</div><div class="l">班级平均</div></div>
    <div class="cell"><div class="v">{_esc(prev_v)}</div><div class="l">{_esc(prev_l)}</div></div>
  </div>
  <div class="bands">{''.join(band_rows)}</div>
  {small_sample_note}
  <div class="concl">{_esc(_score_card_conclusion(student, info_by_qid))}</div>
  {strength}
</div>"""

    # ---- B 历次成绩 ----
    card_b = ""
    scored_history = [e for e in history if e.get("score") is not None]
    if len(scored_history) >= 2:
        card_b = f"""
<div class="card">
  <h2>历次成绩</h2>
  <div class="hist">{_history_chart_html(history, data.session_id)}</div>
  <div class="note">实线为孩子的分数，虚线为班级平均。每次试卷难度不同，名次比分数更可比。</div>
</div>"""

    # ---- C 答题一览 ----
    scored_records = [r for r in student.records if r.max_score > 0]
    counts = {"full": 0, "part": 0, "zero": 0, "blank": 0}
    cells, templates = [], []
    for record in scored_records:
        info = info_by_qid.get(record.question_id)
        status = _personal_cell_status(record)
        counts[status] = counts.get(status, 0) + 1
        sym = {"full": "✓", "part": "△", "zero": "✗", "blank": "—"}[status]
        hard = info is not None and info.class_rate is not None and info.class_rate < 0.5
        cells.append(
            f'<button type="button" class="qcell {status}" data-q="{_esc(record.question_id)}" '
            f'aria-expanded="false"><span class="sym">{sym}</span>'
            f'{"<span class=hard>▲</span>" if hard else ""}'
            f'<span class="qn">{_esc(_question_cell_label(record.question_id))}</span>'
            f'<span class="qs">{_fmt_num(record.score)}/{_fmt_num(record.max_score)}</span></button>'
        )
        analysis = analysis_by_qid.get(record.question_id)
        labels = _points_for_question(points, record.question_id)
        templates.append(
            f'<template data-q="{_esc(record.question_id)}">'
            + _question_detail_template(
                record, info, labels, analysis, shots, errors=question_errors(record.question_id),
            )
            + "</template>"
        )
    card_c = f"""
<div class="card">
  <h2>本卷答题一览</h2>
  <div class="qgrid">{''.join(cells)}</div>
  <div class="legend">✓ 满分 {counts['full']}　△ 部分得分 {counts['part']}　✗ 零分 {counts['zero']}　— 未作答 {counts['blank']}　▲ 全班平均得分不到一半</div>
  {''.join(f'<div class="note">{_esc(note)}</div>' for note in student.material_notes)}
  <div id="qpanel" hidden></div>
  {''.join(templates)}
</div>"""

    # ---- D 重点跟进 ----
    points_by_qid: dict[str, list[dict[str, Any]]] = {}
    for p in points:
        for q in p["questions"]:
            points_by_qid.setdefault(q["id"], []).append(p)
    if narrative is not None:
        problems = [
            i for i in _narrative_items(narrative, "problems") if isinstance(i, dict)
        ][:3]
        suggestions = [
            i for i in _narrative_items(narrative, "suggestions") if isinstance(i, dict)
        ]
        dcards = []
        for idx, problem in enumerate(problems):
            suggestion = suggestions[idx] if idx < len(suggestions) else None
            qids = _follow_up_question_ids(problem, suggestion, student, info_by_qid)
            lost_sum = sum(
                r.lost_points for r in student.records if r.question_id in qids
            )
            if qids and lost_sum > 0:
                def qtag(question_id: str) -> str:
                    label = _question_display_label(question_id)
                    categories = [
                        str(row["category"]) for row in question_errors(question_id)
                        if row.get("category")
                    ]
                    if categories:
                        label += f"·{'+'.join(dict.fromkeys(categories))}"
                    return f'<span class="qtag">{_esc(label)}</span>'

                tags = "".join(
                    qtag(q)
                    for q in sorted(
                        qids, key=lambda q: question_id_coordinates(q) or (10**9, 0)
                    )
                )
                meta_line = f'<div class="meta2">丢 {_fmt_num(lost_sum)} 分{tags}</div>'
            else:
                meta_line = ""
            recur_names: list[str] = []
            error_recur = error_recurrence(qids)
            for q in qids:
                for p in points_by_qid.get(q, []):
                    for nm in p["recur"]:
                        if nm not in recur_names:
                            recur_names.append(nm)
            # 有本场错因归类时优先按同类错误/同错法提示复发；否则回退到考查点复发。
            badge = (
                f'<span class="recur" title="{_esc("、".join(error_recur))}">同类错误以前出现过</span>'
                if error_recur
                else (
                    f'<span class="recur" title="{_esc("、".join(recur_names))}">以前也有失分</span>'
                    if recur_names
                    else ""
                )
            )
            why = [f'<div class="qd-fb">{_report_paragraphs(_narrative_text(problem.get("detail")))}</div>']
            shown_parents = set()
            for q in qids:
                item = analysis_by_qid.get(q)
                fb = _narrative_text((item or {}).get("feedback"))
                if fb:
                    why.append(
                        f'<div class="qd-fb"><b>{_esc(_question_display_label(q))}</b>'
                        f'{_report_paragraphs(fb)}</div>'
                    )
                elif item is None:
                    rec = next((r for r in student.records if r.question_id == q), None)
                    if rec is not None:
                        fallback = _grading_fallback_html(
                            rec, info_by_qid.get(q), heading=_question_display_label(q)
                        )
                        if fallback:
                            why.append(fallback)
                parent = _parent_question_id(q)
                if parent not in shown_parents:
                    shown_parents.add(parent)
                    first = next(iter(_shots_for_parent(shots, parent)), None)
                    if first:
                        why.append(
                            f'<div class="shot"><img src="{first["data_uri"]}" alt="作答截图">'
                            f'<div class="cap">{_esc(first.get("caption") or "学生作答（原卷截图）")}</div></div>'
                        )
            sugg_line = ""
            if suggestion and _narrative_text(suggestion.get("detail")):
                sugg_line = (
                    f'<div class="help">在家这样帮：'
                    f'{_report_inline_math(_report_display_text(_narrative_text(suggestion.get("detail"))))}</div>'
                )
            dcards.append(
                f'<div class="dcard"><div class="dhead"><span class="no">{idx + 1}</span>'
                f'<span class="ttl">{_esc(_report_display_text(_narrative_text(problem.get("title"))))}</span>{badge}</div>'
                f'{meta_line}'
                f'{sugg_line}'
                f'<details><summary>为什么这样判断</summary>{"".join(why)}</details></div>'
            )
        if not problems and not [r for r in student.records if r.lost]:
            dcards.append('<div class="dgreen">这次没有失分，保持现在的做题习惯。</div>')
        card_d = f"""
<div class="card">
  <h2>这次重点跟进<span class="aitag">{_esc(AI_DISCLAIMER)}</span></h2>
  {''.join(dcards)}
</div>"""
    else:
        lost = [r for r in student.records if r.lost]
        if not lost:
            inner = '<div class="dgreen">这次没有失分，保持现在的做题习惯。</div>'
        else:
            by_parent: dict[str, list[_StudentQuestionRecord]] = {}
            for r in lost:
                by_parent.setdefault(_parent_question_id(r.question_id), []).append(r)
            ordered = sorted(
                by_parent.items(),
                key=lambda kv: (
                    -sum(r.lost_points for r in kv[1]),
                    question_id_coordinates(kv[0]) or (10**9, 0),
                ),
            )
            cards = []
            for parent, recs in ordered[:3]:
                comments = []
                for r in recs:
                    text = _real_teacher_comment(r)
                    if text and text not in comments:
                        comments.append(text)
                lost_sum = sum(r.lost_points for r in recs)
                error_labels = [
                    error_label(r.question_id) for r in recs if error_label(r.question_id)
                ]
                recur_names = error_recurrence([r.question_id for r in recs])
                badge = (
                    f'<span class="recur" title="{_esc("、".join(recur_names))}">同类错误以前出现过</span>'
                    if recur_names
                    else ""
                )
                cards.append(
                    f'<div class="dcard"><div class="dhead"><span class="ttl">'
                    f'{_esc(_lost_group_label(parent, student, info_by_qid))}</span>{badge}</div>'
                    f'<div class="meta2">丢 {_fmt_num(lost_sum)} 分</div>'
                    + (
                        f'<div class="errline">错误归类（AI 辅助）：{_esc("；".join(dict.fromkeys(error_labels)))}</div>'
                        if error_labels
                        else ""
                    )
                    + (
                        f'<div class="help">老师批语：{_esc("；".join(comments[:2]))}</div>'
                        if comments
                        else ""
                    )
                    + "</div>"
                )
            inner = "".join(cards)
        card_d = f"""
<div class="card">
  <h2>这次重点跟进</h2>
  {inner}
</div>"""

    # ---- E 考查点 ----
    n_good = sum(1 for p in points if p["status"] == "good")
    n_mid = sum(1 for p in points if p["status"] == "mid")
    n_bad = sum(1 for p in points if p["status"] == "bad")
    rows_html = []
    for p in [p for p in points if p["status"] == "bad"] + [
        p for p in points if p["status"] == "mid"
    ]:
        sym_map = {"full": "✓", "part": "△", "zero": "✗"}
        dots = "".join(
            f'<span class="dot"><i class="{r}">{sym_map[r]}</i>'
            f'<s>{_esc(_question_cell_label(q["id"]))}</s></span>'
            for q, r in p["results"]
        )
        badge = (
            f'<span class="recur" title="{_esc("、".join(p["recur"]))}">以前也有失分</span>'
            if p["recur"]
            else ""
        )
        rows_html.append(
            f'<div class="pt-row"><span class="nm">{_esc(p["label"])}{badge}</span>'
            f'<span class="dots">{dots}</span></div>'
        )
    good_sections: dict[str, list[str]] = {}
    for p in points:
        if p["status"] == "good":
            good_sections.setdefault(p["section"] or "其他", []).append(p["label"])
    cloud = "".join(
        f'<div class="sec">{_esc(sec)}</div><div>'
        + "".join(f'<span class="tg">{_esc(label)}</span>' for label in labels)
        + "</div>"
        for sec, labels in good_sections.items()
    )
    card_e = ""
    if points:
        card_e = f"""
<div class="card">
  <h2>本次考查点</h2>
  <div class="pt-counts"><b class="g">做得好 {n_good}</b><b class="m">部分做到 {n_mid}</b><b class="b">需要补 {n_bad}</b></div>
  {''.join(rows_html)}
  {f'<div class="tagcloud">{cloud}</div>' if cloud else ''}
  <div class="note" style="margin-top:8px">只根据这次考试判断；同一道题可能同时考几个点。</div>
</div>"""

    # ---- F 附录 ----
    appendix = _lost_appendix_html(
        student, info_by_qid, analysis_by_qid, shots, error_map=error_map,
    )
    card_f = f'<div class="card">{appendix}</div>' if appendix else ""

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{_esc(student.student_name)} · 家长报告 · {_esc(data.session_name)}</title>
<style>{_PERSONAL_CSS}</style>
{_report_math_assets()}
</head>
<body>
<div class="page">
{card_a}
{card_b}
{card_c}
{card_d}
{card_e}
{card_f}
<footer>
<p>本报告根据阅卷记录自动生成；标“{_esc(AI_DISCLAIMER)}”的内容由人工智能辅助生成，仅供参考，如有疑问请联系任课老师。</p>
<p>班级对比只用匿名统计。报告生成时间：{generated_at}。</p>
</footer>
</div>
<script>{_PERSONAL_KATEX_JS}{_PERSONAL_GRID_JS}</script>
</body>
</html>
"""


# ---------------------------------------------------------------------------
# 生成器与 preflight
# ---------------------------------------------------------------------------


def _render_class_html(
    data: _SessionAnalysisData,
    narrative: dict[str, Any] | None,
    cause_counts: dict[str, list[tuple[str, int]]] | None = None,
) -> str:
    page = build_class_page_data(data)
    class_name = data.class_name or "未分班"
    class_label = class_name if class_name.endswith("班") else f"{class_name}班"
    failed = narrative is None
    narrative = class_narrative_with_student_names(narrative, data.students) if narrative is not None else {}
    analyzed = len(data.students)
    fmt = lambda value: _fmt_num(value) if value is not None else '—'
    math_text = lambda value: _report_inline_math(str(value))
    aliases = {f'S{i}': student['student_name'] for i, student in enumerate(page['students'], 1)}
    stats = page['score_distribution']
    metrics = [('参考人数', page['present']), ('平均分', fmt(stats['avg'])),
               ('中位数', fmt(stats['median'])), ('最高 / 最低', fmt(stats['max'])+' / '+fmt(stats['min'])),
               ('及格率', f"{stats['pass_rate']*100:.1f}%")]
    metric_html = ''.join(f'<div><small>{_esc(label)}</small><strong>{_esc(str(value))}</strong></div>'
                          for label, value in metrics)
    findings = ''.join(f'<article><h3>{math_text(item["title"])}</h3>{_report_paragraphs(item["detail"])}</article>'
                       for item in narrative.get('key_findings', []))
    if failed:
        findings = f'<p>{_esc(AI_FAILED_NOTE)}</p>'
    issues = ''.join(f'<article><h3>{math_text(item["title"])}</h3>{_report_paragraphs(item["evidence"])}'
                     f'<div class="action">{_report_paragraphs(item["teaching_action"])}</div></article>'
                     for item in narrative.get('common_issues', []))
    bands = ''.join(f'<div class="band"><span>{_esc(label)}</span><div class="track">'
                    f'<i style="width:{count/max(1,analyzed)*100:.2f}%"></i></div><b>{count}人</b></div>'
                    for label, count in stats['bands'].items())
    def _cause_text(question_id: str) -> str:
        counts = (cause_counts or {}).get(str(question_id)) or []
        return '；'.join(f'{category} {count}人' for category, count in counts[:3]) or '—'

    questions = ''.join(
        f'<tr class="qrow" id="score-{_esc(q["question_id"])}"><td>{_question_display_label(q["question_id"])}</td>'
        f'<td>{fmt(q["max_score"])}</td><td>{fmt(q["class_avg"])}</td>'
        f'<td><div class="rate"><i style="width:{q["class_rate"]*100:.2f}%"></i>'
        f'<span>{q["class_rate"]*100:.1f}%</span></div></td><td>{len(q["records"])}人</td>'
        f'<td>{_esc(_cause_text(q["question_id"]))}</td>'
        + (f'<td class="answer"><details><summary>查看答案</summary><div>{math_text(q["canonical_answer"])}</div></details></td></tr>'
           if q['canonical_answer'] else '<td class="answer">—</td></tr>')
        for q in page['questions'] if q['class_avg'] is not None)
    roster = ''.join(f'<tr><td>{student["rank"]}</td><td>{_esc(student["student_name"])}</td>'
                     f'<td>{fmt(student["total_score"])}</td>'
                     f'<td>{_esc("、".join(_question_display_label(r["question_id"]) for r in student["lost"]) or "无")}</td></tr>'
                     for student in page['students'])
    notes = ''.join(f'<article><h3>{_esc(aliases.get(item["alias"], item["alias"]))}</h3>'
                    f'{_report_paragraphs(item["note"])}<div class="action">{_report_paragraphs(item["suggestion"])}</div></article>'
                    for item in narrative.get('student_notes', []))
    knowledge_html = _knowledge_view_html(_class_knowledge_view(data))
    if not knowledge_html:
        knowledge_html = ('<section id="knowledge-map"><h2>班级知识与技能掌握图</h2>'
                          '<p class="muted">本卷尚无可可靠匹配的知识与技能标签，暂不展示掌握图；逐题成绩仍可查看。</p></section>')
    knowledge_css, knowledge_js = _report_knowledge_assets()
    styles = """
    *{box-sizing:border-box}body{margin:0;background:#f2f5f9;color:#24354b;font:15px/1.8 "Microsoft YaHei",sans-serif}
    main{max-width:1080px;margin:auto;padding:40px 28px}header{border-bottom:3px solid #527eb0;padding-bottom:22px}
    h1{font-size:28px;line-height:1.45;margin:8px 0}h2{font-size:20px;margin:0 0 18px}h3{font-size:16px;margin:0 0 6px}
    p{margin:5px 0 12px}small,.muted{color:#63758a}section{background:white;border:1px solid #e1e7ef;border-radius:14px;padding:26px;margin:24px 0}
    .metrics{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;margin-top:24px}.metrics div{background:#fff;border-radius:10px;padding:14px}
    .metrics small{display:block}.metrics strong{display:block;font-size:24px}
    .band{display:grid;grid-template-columns:135px 1fr 50px;align-items:center;gap:12px;margin:10px 0}
    .track,.rate{height:20px;border-radius:5px;background:#edf2f8;overflow:hidden}.track i,.rate i{display:block;height:100%;background:#6f94c5}
    .rate{position:relative;min-width:100px;height:26px}.rate span{position:absolute;inset:0;text-align:center;color:#173153;font-size:13px;line-height:26px}.rate i{background:#c0d3ea}
    .table-wrap{overflow-x:auto}table{width:100%;border-collapse:collapse;font-size:14px}th,td{border-bottom:1px solid #e4eaf0;text-align:left;padding:10px}
    th{background:#f5f8fb;color:#52677f}th:not(:last-child),td:not(:last-child){white-space:nowrap}.answer{min-width:120px;max-width:290px;overflow-wrap:anywhere}
    article{border-left:3px solid #8aabc9;padding:0 0 0 16px;margin:20px 0}.action{color:#395e86;background:#f4f8fc;padding:10px 14px;border-radius:7px}.action p:last-child{margin-bottom:0}
    summary{cursor:pointer;font-size:19px;font-weight:bold}.answer summary{font-size:12px;font-weight:400;color:#30567f}.answer details>div{padding-top:8px}.roster td:last-child{font-size:12px;color:#617188}
    .qm{display:inline-block;max-width:100%;vertical-align:baseline;padding:3px 2px 6px}.qm-display{display:block;text-align:center;margin:8px 0}.qm-error{color:#9a3412}
    article p,article h3,.answer{overflow-x:auto;overflow-y:hidden;overflow-wrap:anywhere}.katex{font-size:1.05em}footer{color:#69788a;font-size:12px}
    @media(max-width:600px){main{padding:22px 14px}section{padding:18px 14px}.metrics{grid-template-columns:repeat(2,1fr)}h1{font-size:23px}.band{grid-template-columns:116px 1fr 36px;font-size:12px;gap:5px}th,td{padding:8px 6px;font-size:12px}}
    @media print{body{background:white}main{padding:0}article,.kn-node{break-inside:avoid}h2,h3{break-after:avoid}details:not([open])>*:not(summary){display:block}footer{margin-top:20px}}
    """
    math_js = """
    document.querySelectorAll('.qm[data-latex]').forEach(el => {
      try { katex.render(el.dataset.latex, el, {throwOnError:true, output:'htmlAndMathml', displayMode:el.dataset.display === 'true'}); }
      catch (_) { el.classList.add('qm-error'); el.title = '公式格式需核对，已保留原文'; }
    });
    """
    return (
        '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        f'<title>{_esc(class_label)} · 班级报告</title>{_report_math_assets()}<style>{styles}{knowledge_css}</style></head><body><main>'
        f'<header><small>讲评课参考 · 教师版</small><h1>{math_text(page["exam"]["title"])}<br>{_esc(class_label)} · 班级报告</h1>'
        f'<p class="muted">按本班当前成绩统计，教师复核分优先 · 满分{fmt(page["exam"]["full_score"])}分</p></header>'
        f'<div class="metrics">{metric_html}</div><section><h2>本次最值得关注的结果</h2>{findings}</section>'
        f'<section><h2>分数分布</h2>{bands}<p class="muted">各分数段互不重叠，参与统计{analyzed}人。</p></section>'
        '<section><h2>逐题得分</h2><div class="table-wrap"><table><thead><tr><th>题目</th><th>满分</th><th>均分</th>'
        f'<th>得分率</th><th>未得满分</th><th>主要错误</th><th>参考答案</th></tr></thead><tbody>{questions}</tbody></table></div></section>'
        f'{knowledge_html}<section><h2>下一节讲评课</h2>{issues}<h3>分层安排</h3>{_report_paragraphs(narrative.get("grouping_advice", ""))}</section>'
        f'<section><h2>个别跟进</h2>{notes}</section><section><details><summary>全班成绩与失分题目（{analyzed}人）</summary>'
        '<div class="table-wrap"><table class="roster"><thead><tr><th>名次</th><th>姓名</th><th>成绩</th><th>失分题目</th></tr></thead>'
        f'<tbody>{roster}</tbody></table></div></details></section>'
        '<footer>AI 分析 · 仅供参考。统计使用本班当前成绩，教师复核分优先；错因依据现有作答证据与批改记录，不据分数推断学生态度或作答时间。</footer>'
        f'</main><script>{math_js}{knowledge_js}</script></body></html>'
    )


class AnalysisReportGenerator:
    """契约与 ReportGenerator 一致：(repositories, output_dir) 构造，导出返回 Path。"""

    def __init__(
        self,
        db: GradingRepositoryAccess | Path,
        output_dir: Path,
        *,
        llm_client_factory: Callable[[], Any] | None = None,
        narrative_cache_dir: Path | None = None,
        data_root: Path | None = None,
        reports_dir: Path | None = None,
    ) -> None:
        self.repositories = (
            open_grading_repositories(Path(db))
            if isinstance(db, Path)
            else as_grading_repositories(db)
        )
        self.db_path = self.repositories.db_path
        self.output_dir = Path(output_dir)
        self.llm_client_factory = llm_client_factory
        self.cache = (
            AnalysisNarrativeCache(narrative_cache_dir)
            if narrative_cache_dir is not None
            else None
        )
        self.data_root = data_root or _infer_data_root(self.db_path)
        # 受控 reports 目录：读取 .class_analysis 里的学生错因记录；缺省时不展示归类。
        self.reports_dir = Path(reports_dir) if reports_dir is not None else None
        self._llm_client: Any = None
        self._llm_client_resolved = False

    def export_session(
        self,
        session_id: int,
        report_type: str,
        *,
        score_revision: str = "",
        student_ids: set[int] | None = None,
    ) -> Path:
        if report_type not in ANALYSIS_REPORT_TYPES:
            raise ValueError(f"不支持的分析报告类型: {report_type}")
        revision = str(score_revision or "").strip() or _compute_score_revision(
            self.repositories,
            int(session_id),
        )
        self.output_dir.mkdir(parents=True, exist_ok=True)
        data = assemble_session_analysis(
            self.repositories,
            int(session_id),
            data_root=self.data_root,
        )
        return self._export_personal(data, revision, student_ids=student_ids)

    def export_classes(
        self,
        session_id: int,
        *,
        class_names: set[str] | None = None,
        score_revision: str = "",
    ) -> list[Path]:
        """以班级页面相同的数据、提示词和缓存规则批量导出自包含 HTML。"""
        from backend.class_analysis import (
            _class_narrative, question_category_counts, session_error_records,
        )
        from analysis_report_prompts import CLASS_SYSTEM_PROMPT

        data = assemble_session_analysis(self.repositories, session_id, data_root=self.data_root)
        revision = score_revision or _compute_score_revision(self.repositories, session_id)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        cache = self.cache or AnalysisNarrativeCache(self.output_dir / ".analysis_narrative_cache")
        _enrich_personal_knowledge(self.repositories, data, self.data_root)
        # 错因记录按整场装配（指纹口径），再按各班学生过滤统计；状态缺失时为空。
        error_records = (
            session_error_records(
                self.repositories, session_id, self.reports_dir, data_root=self.data_root,
            )
            if self.reports_dir is not None else {}
        )
        client = self._client()
        files = []
        for name, group in split_session_analysis_by_class(data).items():
            if not group.students or (class_names is not None and name not in class_names):
                continue
            narrative = _class_narrative(
                client=client, cache=cache, session_id=session_id, revision=revision,
                class_name=name, prompt=build_report_prompt(CLASS_SYSTEM_PROMPT, build_class_payload(group)),
            )
            cause_counts = question_category_counts(
                error_records, student_ids=[s.student_id for s in group.students])
            target = self.output_dir / f"{safe_filename_fragment(name, '未分班')}_班级报告.html"
            target.write_text(
                _render_class_html(group, narrative, cause_counts=cause_counts),
                encoding="utf-8")
            files.append(target)
        return files

    def _client(self) -> Any:
        # 未配置内容生成模型时返回 None，全部报告降级为无 AI 叙述版，
        # 不静默改用阅卷模型（方案 §5）。
        if not self._llm_client_resolved:
            self._llm_client_resolved = True
            self._llm_client = (
                self.llm_client_factory()
                if self.llm_client_factory is not None
                else None
            )
        return self._llm_client

    def _cached_narrative(
        self,
        *,
        session_id: int,
        revision: str,
        report_type: str,
        report_key: str,
    ) -> dict[str, Any] | None:
        """当前叙述版本优先；个人报告再按旧版叙述版本依次兼容读取。

        旧版本命中直接返回，不调用模型、不写回新 key。
        """
        if self.cache is None:
            return None
        from backend.report_exports import (
            LEGACY_PERSONAL_NARRATIVE_VERSIONS,
            report_narrative_version,
        )

        versions = [report_narrative_version(report_type)]
        if report_type == PERSONAL_ANALYSIS_REPORT_TYPE:
            versions.extend(LEGACY_PERSONAL_NARRATIVE_VERSIONS)
        for rendition_version in versions:
            cached = self.cache.load(
                AnalysisNarrativeCache.cache_key(
                    session_id=session_id,
                    score_revision=revision,
                    rendition_version=rendition_version,
                    report_key=report_key,
                )
            )
            if cached is not None:
                return cached
        return None

    def _narrative(
        self,
        *,
        session_id: int,
        revision: str,
        report_type: str,
        report_key: str,
        prompt: str,
        max_tokens: int,
        image_blobs: list[bytes] | None = None,
    ) -> dict[str, Any] | None:
        from backend.report_exports import report_narrative_version

        key = AnalysisNarrativeCache.cache_key(
            session_id=session_id,
            score_revision=revision,
            rendition_version=report_narrative_version(report_type),
            report_key=report_key,
        )
        cached = self._cached_narrative(
            session_id=session_id,
            revision=revision,
            report_type=report_type,
            report_key=report_key,
        )
        if cached is not None:
            return cached
        client = self._client()
        if client is None:
            return None
        try:
            options = {"temperature": 0.3, "max_tokens": max_tokens}
            if image_blobs:
                narrative = client.json_from_images_once(
                    prompt, image_blobs,
                    use_config_client=True,
                    extra_kwargs=options,
                )
            else:
                narrative = client.json_from_text(prompt, extra_kwargs=options)
        except Exception:
            # 模型超时/解析失败：本地修复仍失败则降级，不暗中重发（已确认偏差）。
            return None
        if not isinstance(narrative, dict):
            return None
        if self.cache is not None:
            self.cache.store(key, narrative)
        return narrative

    def _export_personal(
        self,
        data: _SessionAnalysisData,
        revision: str,
        *,
        student_ids: set[int] | None = None,
    ) -> Path:
        if not data.students:
            raise ValueError("该场次没有可生成个人报告的学生。")
        _enrich_personal_questions(self.repositories, data, self.data_root)
        _enrich_personal_knowledge(self.repositories, data, self.data_root, student_ids=student_ids)
        regions = load_session_regions(
            self.repositories,
            data.session_id,
            data_root=self.data_root,
        )
        staging_subdir = self.output_dir / "personal_reports"
        staging_subdir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        used_names: set[str] = set()
        report_files: list[Path] = []
        scoped_students = [
            (group, student)
            for group in split_session_analysis_by_class(data).values()
            for student in group.students
            if student_ids is None or student.student_id in student_ids
        ]
        if not scoped_students:
            raise ValueError("所选学生没有可生成个人报告的成绩。")
        session_data = data
        # Resolve once on the caller thread. Only model/cache work runs in workers;
        # repository access, image preparation and rendering stay on this thread.
        all_cached = self.cache is not None and all(
            self._cached_narrative(
                session_id=data.session_id,
                revision=revision,
                report_type=PERSONAL_ANALYSIS_REPORT_TYPE,
                report_key=f"personal:{student.student_id}",
            )
            is not None
            for data, student in scoped_students
        )
        client = None if all_cached else self._client()
        # 同教学学期历次成绩每次导出只读取一轮，渲染时按学生取用。
        histories = _load_student_histories(
            self.repositories, session_data, self.data_root
        )
        # 错因整理产物：本场学生×题记录 + 历次同类/同错法场次索引。
        # 未整理或输入已过期的题不返回记录，报告相应位置不显示错误类型。
        error_state: dict[str, Any] = {}
        error_sources: list[dict[str, Any]] = []
        history_error_index: dict[int, Any] = {}
        error_session_names: dict[int, str] = {}
        if self.reports_dir is not None:
            from backend.class_analysis import (
                ClassAnalysisStateStore,
                build_cause_inputs,
                collect_student_error_index,
                student_error_map,
            )

            store = ClassAnalysisStateStore(self.reports_dir)
            error_state = store.load(data.session_id) or {}
            error_sources = build_cause_inputs(data)
            history_ids = {
                int(entry["session_id"])
                for entries in histories.values()
                for entry in entries
                if entry.get("session_id") != data.session_id
            }
            error_session_names = {
                int(entry["session_id"]): str(entry.get("name") or "")
                for entries in histories.values()
                for entry in entries
                if entry.get("session_id") != data.session_id
            }
            history_error_index = collect_student_error_index(
                store, sorted(history_ids), db=self.repositories, data_root=self.data_root,
            )
            error_maps = {
                student.student_id: student_error_map(error_state, student, error_sources)
                for _group, student in scoped_students
            }
        else:
            error_maps = {}
        execution = getattr(getattr(client, "config_gateway", None), "execution_snapshot", None)
        parallel_limit = min(3, max(1, int(getattr(execution, "max_in_flight", 3))))

        def prepare(executor, data, student):
            paper_context = _student_paper_context(self.repositories, data, student)
            shots = capture_lost_question_shots(
                self.repositories, data, student, regions=regions,
                data_root=self.data_root, paper_context=paper_context,
            )
            images, image_map = _personal_image_inputs(
                data, student, shots, paper_context, self.data_root,
            )
            payload = build_personal_payload(data, student)
            payload["image_map"] = image_map
            payload["material_notes"] = student.material_notes
            future = executor.submit(
                self._narrative,
                session_id=data.session_id,
                revision=revision,
                report_type=PERSONAL_ANALYSIS_REPORT_TYPE,
                report_key=f"personal:{student.student_id}",
                prompt=build_report_prompt(PERSONAL_SYSTEM_PROMPT, payload),
                max_tokens=personal_output_token_limit(sum(record.lost for record in student.records)),
                image_blobs=images,
            )
            filename = (
                f"{safe_filename_fragment(student.student_code, '未知学号')}"
                f"_{safe_filename_fragment(student.student_name, '未知姓名')}"
                "_个人报告.html"
            )
            if filename in used_names:
                filename = filename.replace(".html", f"_{student.student_id}.html")
            used_names.add(filename)
            report_path = staging_subdir / filename
            report_files.append(report_path)
            return future, (data, student, shots, report_path)

        students = iter(scoped_students)
        with ThreadPoolExecutor(
            max_workers=parallel_limit, thread_name_prefix="personal-report"
        ) as executor:
            pending = {}
            while True:
                # Bound prepared images as well as in-flight model requests.
                while len(pending) < parallel_limit:
                    entry = next(students, None)
                    if entry is None:
                        break
                    future, context = prepare(executor, *entry)
                    pending[future] = context
                if not pending:
                    break
                completed, _ = wait(pending, return_when=FIRST_COMPLETED)
                for future in completed:
                    data, student, shots, report_path = pending.pop(future)
                    history_index = history_error_index.get(student.student_id)
                    error_history = (
                        {
                            key: {
                                label: sorted(
                                    {error_session_names.get(sid, "") for sid in sids} - {""}
                                )
                                for label, sids in bucket.items()
                            }
                            for key, bucket in history_index.items()
                        }
                        if history_index
                        else None
                    )
                    html_text = _render_personal_html(
                        data,
                        student,
                        future.result(),
                        shots,
                        history=histories.get(student.student_id, []),
                        error_map=error_maps.get(student.student_id),
                        error_history=error_history,
                    )
                    report_path.write_text(html_text, encoding="utf-8")

        data = session_data

        checklist_lines = [
            "未生成个人报告的学生清单",
            f"场次：{data.session_name}",
            "",
        ]
        if data.skipped:
            checklist_lines.extend(
                f"{item['class_name']} {item['student_code']} {item['student_name']}：{item['reason']}"
                for item in data.skipped
            )
        else:
            checklist_lines.append("无")
        checklist_path = staging_subdir / "未生成清单.txt"
        checklist_path.write_text("\n".join(checklist_lines) + "\n", encoding="utf-8")

        zip_path = self.output_dir / session_export_path_name(
            data.session_name,
            "个人分析报告",
            "zip",
            timestamp,
        )
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in [*report_files, checklist_path]:
                archive.write(path, arcname=path.name)
        return zip_path


def _compute_score_revision(
    repositories: GradingRepositoryAccess,
    session_id: int,
) -> str:
    from backend.report_exports import score_revision

    return score_revision(repositories, session_id)


def build_analysis_preflight(
    db: GradingRepositoryAccess | Any,
    session_id: int,
    report_type: str,
    *,
    score_revision: str,
    cache_dir: Path,
    reports_dir: Path | None = None,
) -> dict[str, Any]:
    """生成前的费用与调用预估：错因整理与报告叙述分开计数；只给 token 粗估。"""
    if report_type not in ANALYSIS_REPORT_TYPES:
        raise ValueError(f"不支持的分析报告类型: {report_type}")
    from backend.report_exports import (
        LEGACY_PERSONAL_NARRATIVE_VERSIONS,
        report_narrative_version,
    )

    repositories = as_grading_repositories(db)
    data = assemble_session_analysis(repositories, int(session_id))
    _enrich_personal_questions(repositories, data, None)
    cache = AnalysisNarrativeCache(cache_dir)
    renditions = [report_narrative_version(report_type)]
    if report_type == PERSONAL_ANALYSIS_REPORT_TYPE:
        renditions.extend(LEGACY_PERSONAL_NARRATIVE_VERSIONS)
    entries = [
        (
            f"personal:{student.student_id}",
            build_report_prompt(
                PERSONAL_SYSTEM_PROMPT,
                build_personal_payload(data, student),
            ),
            personal_output_token_limit(sum(record.lost for record in student.records)),
        )
        for data in split_session_analysis_by_class(data).values()
        for student in data.students
    ]

    cache_hits = 0
    estimated_tokens = 0
    for report_key, prompt, max_tokens in entries:
        hit = any(
            cache.load(
                AnalysisNarrativeCache.cache_key(
                    session_id=int(session_id),
                    score_revision=score_revision,
                    rendition_version=rendition,
                    report_key=report_key,
                )
            )
            is not None
            for rendition in renditions
        )
        if hit:
            cache_hits += 1
            continue
        # 文本粗估 = 输入 token（字符数/1.5）+ 输出上限；图片计费由模型决定。
        estimated_tokens += estimate_prompt_tokens(prompt) + max_tokens

    # 错因整理是报告导出的前置阶段：与实际任务同口径估算——已整理且输入未变
    # 的题与整理失败且输入未变的题都不重复调用。已关联题库的选择题走选项诊断
    # （可复用，未复用时也只计 1 次）；填空题错误答案库全覆盖时零调用。
    cause_call_count = 0
    cause_total_questions = 0
    cause_estimated_tokens = 0
    if report_type == PERSONAL_ANALYSIS_REPORT_TYPE:
        from backend.class_analysis import (
            CAUSE_ANALYSIS_PROMPT,
            CAUSE_ANALYSIS_VERSION,
            ClassAnalysisStateStore,
            assemble_cause_data,
            build_cause_inputs,
            known_cause_patterns,
            plan_cause_question,
            _cause_input_fingerprint,
            cause_input_matches,
        )
        from backend.error_patterns import (
            OPTION_ANALYSIS_PROMPT,
            bank_confirmed_triggers,
            build_option_analysis_input,
            session_bank_context,
        )

        store = ClassAnalysisStateStore(Path(reports_dir) if reports_dir is not None else cache_dir.parent)
        stored = ((store.load(int(session_id)) or {}).get("cause_analysis") or {}).get("questions") or {}
        cause_data = assemble_cause_data(repositories, int(session_id))
        sources = build_cause_inputs(
            cause_data,
            known_patterns=known_cause_patterns(
                store, _question_bank_db_path(repositories.db_path), int(session_id),
            ),
        )
        cause_total_questions = len(sources)
        qb_path = _question_bank_db_path(repositories.db_path)
        option_scope = True
        bank_context = session_bank_context(qb_path, int(session_id))
        confirmed_by_bank = bank_confirmed_triggers(
            qb_path,
            sorted({bid for ctx in bank_context.values() for bid in ctx["bank_ids"]}),
        )
        qtypes = {info.question_id: str(info.question_type or "") for info in cause_data.questions}
        for source in sources:
            fingerprint = _cause_input_fingerprint(source)
            saved = stored.get(source["question_id"]) or {}
            if (saved.get("version") == CAUSE_ANALYSIS_VERSION
                    and cause_input_matches(saved, source) and saved.get("result")):
                continue
            if saved.get("failed") and saved.get("failed_input_fingerprint") == fingerprint:
                continue
            plan = plan_cause_question(
                store, int(session_id), source,
                qtype=qtypes.get(source["question_id"], ""),
                option_scope=option_scope,
                ctx=bank_context.get(_parent_question_id(source["question_id"])) or {},
                confirmed_by_bank=confirmed_by_bank,
                question_bank_path=qb_path,
                retry_failed=False,
            )
            if not plan["needs_call"]:
                continue
            cause_call_count += 1
            if plan["path"] == "option":
                option = plan["option"]
                cause_estimated_tokens += (
                    estimate_prompt_tokens(
                        OPTION_ANALYSIS_PROMPT + "\n" + json.dumps(
                            build_option_analysis_input(
                                option["text"], option["correct"],
                                str(source.get("reference_analysis") or "")),
                            ensure_ascii=False)
                    ) + 8000
                )
            else:
                cause_estimated_tokens += (
                    estimate_prompt_tokens(
                        CAUSE_ANALYSIS_PROMPT + "\n" + json.dumps(source, ensure_ascii=False)
                    ) + 12000
                )

    configured = resolve_content_generation_settings() is not None
    service_name, model_name = content_generation_public_info()
    return {
        "report_type": report_type,
        "configured": configured,
        "service_name": service_name if configured else None,
        "model_name": model_name if configured else None,
        "call_count": len(entries) - cache_hits,
        "estimated_total_tokens": estimated_tokens,
        "cache_hits": cache_hits,
        "cause_call_count": cause_call_count,
        "cause_total_questions": cause_total_questions,
        "cause_estimated_tokens": cause_estimated_tokens,
    }
