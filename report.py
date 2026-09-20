from __future__ import annotations

import json
import math
import re
import statistics
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from backend.repositories.access import (
    GradingRepositoryAccess,
    as_grading_repositories,
)
from backend.repositories.compat import open_grading_repositories
from export_names import session_export_path_name
from grading_completeness import resolve_grading_completeness
from path_manager import resolve_stored_file_path
from question_id_contract import resolve_known_question_id


class ReportGenerator:
    def __init__(
        self,
        db_path: GradingRepositoryAccess | Path,
        reports_dir: Path,
    ) -> None:
        self.repositories = (
            open_grading_repositories(Path(db_path))
            if isinstance(db_path, Path)
            else as_grading_repositories(db_path)
        )
        self.db_path = self.repositories.db_path
        self.reports_dir = Path(reports_dir)

    def export_session(
        self,
        session_id: int,
        *,
        score_excel_options: dict[str, object] | None = None,
    ) -> Path:
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        snapshot = self.repositories.reports.get_session_report_snapshot(
            int(session_id),
            question_bank_path=self._question_bank_db_path(),
        )
        session_name = (
            str(snapshot.session.get("session_name") or "")
            if snapshot.session is not None
            else ""
        ) or f"考试批改_{session_id}"
        output_path = self.reports_dir / session_export_path_name(
            session_name,
            "成绩报表",
            "xlsx",
            timestamp,
        )
        df_results = pd.DataFrame(
            snapshot.results,
            columns=[
                "result_id",
                "student_id",
                "student_code",
                "student_name",
                "class_name",
                "total_score",
                "student_score",
                "needs_human_review",
                "graded_at",
                "raw_json",
            ],
        )
        df_attendance = pd.DataFrame(
            snapshot.attendance,
            columns=[
                "student_code",
                "student_name",
                "class_name",
                "attendance_status",
                "source_reason",
                "created_at",
            ],
        )
        df_details = pd.DataFrame(
            snapshot.details,
            columns=[
                "result_id",
                "student_id",
                "student_code",
                "student_name",
                "class_name",
                "question_id",
                "score_awarded",
                "ai_score_awarded",
                "deduction_reason",
                "knowledge_id",
                "knowledge_ids",
                "error_category",
                "error_summary",
            ],
        )

        if df_results.empty:
            raise ValueError("该会话暂无批改结果，无法导出报表。")

        rubric = self._load_rubric_from_session(snapshot.session)
        score_map, type_map = self._question_maps_from_rubric(rubric)
        knowledge_label_map = self._knowledge_label_map_from_rubric(rubric)
        answer_map = self._load_session_answer_map(snapshot.session)
        pending_review_counts = _pending_review_counts(df_results)

        result_status = self._result_statuses(
            df_results,
            df_details,
            rubric,
            df_attendance,
        )
        df_results = df_results.copy()
        df_results["report_status"] = df_results["result_id"].map(
            lambda result_id: result_status.get(int(result_id), ("无效", ""))[0]
        )
        df_results["missing_questions"] = df_results["result_id"].map(
            lambda result_id: result_status.get(int(result_id), ("无效", ""))[1]
        )
        eligible_result_ids = {
            int(row["result_id"])
            for row in df_results.to_dict(orient="records")
            if row.get("report_status") == "完整"
        }
        eligible_results = df_results[
            df_results["result_id"].isin(eligible_result_ids)
        ].copy()
        eligible_details = df_details[
            df_details["result_id"].isin(eligible_result_ids)
        ].copy()

        options = _normalized_excel_options(score_excel_options)
        hidden_student_ids = self._hidden_student_ids(
            eligible_results,
            options,
        )
        score_detail = self._build_compact_session_report(
            df_results,
            df_details,
            score_map,
            rubric=rubric,
            eligible_result_ids=eligible_result_ids,
        )
        question_print = self._build_question_print_analysis(
            eligible_details,
            score_map,
            type_map,
            hidden_student_ids,
            answer_map,
        )
        error_detail = self._build_error_detail_sheet(
            eligible_details,
            score_map,
        )
        ai_teacher_compare = self._build_ai_teacher_comparison_sheet(
            df_details,
            snapshot.locks,
            score_map,
        )
        knowledge_summary = self._build_session_knowledge_summary_by_class(
            eligible_details,
            score_map,
            knowledge_label_map,
            snapshot.knowledge_backfill,
        )
        exceptions = self._build_exception_sheet(
            df_results,
            df_attendance,
        )
        expected_count = (
            len(df_attendance) if not df_attendance.empty else len(df_results)
        )
        attendance_stats = {
            "expected": expected_count,
            "not_counted": max(0, expected_count - len(eligible_results)),
        }
        knowledge_note = (
            "标注“未命名知识点”的题目在题库中未能可靠匹配知识点标签，暂不能细分。"
            if not knowledge_summary.empty
            and (knowledge_summary["知识点"] == "未命名知识点").any()
            else None
        )

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_overview_sheet(
            workbook.create_sheet("考试总览"),
            session_name,
            eligible_results,
            options,
            hidden_student_ids,
            attendance_stats,
        )
        self._write_class_score_sheet(
            workbook.create_sheet("班级成绩总表"),
            session_name,
            eligible_results,
            pending_review_counts,
        )
        self._write_question_print_sheet(
            workbook.create_sheet("小题分析打印"),
            session_name,
            question_print,
            hidden_student_ids,
            options,
        )
        self._write_dataframe_sheet(
            workbook.create_sheet("成绩与小题明细"),
            score_detail,
            freeze_cell="D2",
            landscape=True,
        )
        self._write_dataframe_sheet(
            workbook.create_sheet("AI与人工分对比"),
            ai_teacher_compare,
            freeze_cell="E3",
            landscape=True,
            top_notes=[
                f"{session_name} · AI与人工分对比",
                "本页仅列出教师人工打分或确认过的题目行；未人工干预的题目不再逐行列出。",
            ],
        )
        self._write_dataframe_sheet(
            workbook.create_sheet("错因明细"),
            error_detail,
            freeze_cell="D2",
            landscape=True,
        )
        self._write_dataframe_sheet(
            workbook.create_sheet("知识点分析"),
            knowledge_summary,
            freeze_cell="C3" if knowledge_note else "C2",
            landscape=True,
            top_notes=(
                [f"{session_name} · 知识点分析", knowledge_note]
                if knowledge_note
                else None
            ),
        )
        self._write_dataframe_sheet(
            workbook.create_sheet("缺考与异常"),
            exceptions,
            freeze_cell="A2",
            landscape=False,
        )
        for sheet in workbook.worksheets:
            self._apply_used_range_borders(sheet)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output_path)

        return output_path

    def _result_statuses(
        self,
        df_results: pd.DataFrame,
        df_details: pd.DataFrame,
        rubric: dict[str, object],
        df_attendance: pd.DataFrame,
    ) -> dict[int, tuple[str, str]]:
        detail_rows_by_result: dict[int, list[dict[str, object]]] = {}
        for detail in df_details.to_dict(orient="records"):
            detail_rows_by_result.setdefault(
                int(detail.get("result_id") or 0),
                [],
            ).append(detail)
        attendance_by_identity = {
            (
                str(row.get("student_code") or ""),
                str(row.get("class_name") or "未分班"),
            ): str(row.get("attendance_status") or "").strip()
            for row in df_attendance.to_dict(orient="records")
        }
        statuses: dict[int, tuple[str, str]] = {}
        for result in df_results.to_dict(orient="records"):
            result_id = int(result.get("result_id") or 0)
            identity = (
                str(result.get("student_code") or ""),
                str(result.get("class_name") or "未分班"),
            )
            attendance_status = attendance_by_identity.get(identity)
            if attendance_status in {"absent", "scan_issue"}:
                statuses[result_id] = ("无效", "")
                continue
            completeness = resolve_grading_completeness(
                result.get("raw_json"),
                rubric=rubric,
                details=detail_rows_by_result.get(result_id, []),
            )
            statuses[result_id] = self._completeness_fields(
                {"grading_completeness": completeness}
            )
        return statuses

    def _hidden_student_ids(
        self,
        eligible_results: pd.DataFrame,
        options: dict[str, object],
    ) -> set[int]:
        hidden = {
            int(student_id)
            for student_id in options["manual_hidden_student_ids"]
        }
        bottom_n = int(options["hide_bottom_n"])
        if not options["hide_bottom_enabled"] or bottom_n <= 0:
            return hidden
        for _class_name, group in eligible_results.groupby(
            eligible_results["class_name"].fillna("未分班"),
            dropna=False,
        ):
            # Preserve the old small-class safeguard: a rule for eight names
            # must not silently hide every name in a class of eight or fewer.
            if len(group) <= bottom_n:
                continue
            ordered = group.assign(
                _student_code=group["student_code"].fillna("").astype(str),
                _student_name=group["student_name"].fillna("").astype(str),
            ).sort_values(
                by=[
                    "student_score",
                    "_student_code",
                    "_student_name",
                    "result_id",
                ],
                ascending=[True, True, True, True],
                kind="stable",
            )
            hidden.update(
                int(value)
                for value in ordered.head(bottom_n)["student_id"].tolist()
            )
        return hidden

    def _build_question_print_analysis(
        self,
        df_details: pd.DataFrame,
        score_map: dict[str, float],
        type_map: dict[str, str],
        hidden_student_ids: set[int],
        answer_map: dict[str, str] | None = None,
    ) -> pd.DataFrame:
        columns = [
            "班级",
            "题号",
            "题型",
            "满分",
            "统计人数",
            "全班得分率",
            "平均得分",
            "失分人数",
            "显示姓名数",
            "隐藏姓名数",
            "失分同学",
            "主要错因",
        ]
        if df_details.empty:
            return pd.DataFrame(columns=columns)
        records = _normalize_question_detail_records(
            df_details.to_dict(orient="records"),
            score_map,
        )
        if not records:
            return pd.DataFrame(columns=columns)
        work_df = pd.DataFrame(records)
        answer_map = answer_map or {}
        objective_types = {
            "choice",
            "single_choice",
            "multiple_choice",
            "fill_blank",
            "objective",
            "judgement",
        }
        rows: list[dict[str, object]] = []
        class_names = sorted(
            {
                str(value or "未分班")
                for value in work_df["class_name"].tolist()
            },
            key=_class_sort_key,
        )
        for class_name in class_names:
            class_df = work_df[
                work_df["class_name"].fillna("未分班").astype(str)
                == class_name
            ]
            qids = _natural_question_order(
                class_df["question_id"].dropna().astype(str).unique().tolist()
            )
            for qid in qids:
                qdf = class_df[
                    class_df["question_id"].astype(str) == qid
                ].copy()
                full_score = float(score_map.get(qid) or 0)
                if full_score <= 0:
                    full_score = max(
                        float(qdf["score_awarded"].max() or 0),
                        0.0,
                    )
                lost = [
                    item
                    for item in qdf.to_dict(orient="records")
                    if full_score > 0
                    and float(item.get("score_awarded") or 0)
                    < full_score - 1e-6
                ]
                displayed = [
                    _student_label(item)
                    for item in lost
                    if int(item.get("student_id") or 0)
                    not in hidden_student_ids
                ]
                hidden_count = len(lost) - len(displayed)
                score_sum = float(qdf["score_awarded"].sum())
                full_sum = full_score * len(qdf)
                rows.append(
                    {
                        "班级": class_name,
                        "题号": qid,
                        "题型": str(type_map.get(qid) or ""),
                        "满分": full_score,
                        "统计人数": len(qdf),
                        "全班得分率": (
                            score_sum / full_sum if full_sum > 0 else None
                        ),
                        "平均得分": (
                            score_sum / len(qdf) if len(qdf) else None
                        ),
                        "失分人数": len(lost),
                        "显示姓名数": len(displayed),
                        "隐藏姓名数": hidden_count,
                        "失分同学": (
                            "、".join(displayed)
                            if displayed
                            else ("姓名已隐藏" if lost else "无")
                        ),
                        "主要错因": _summarize_error_categories(
                            qdf.to_dict(orient="records"),
                            full_score,
                            canonical_answer=answer_map.get(qid, ""),
                            is_objective=str(type_map.get(qid) or "")
                            in objective_types,
                        ),
                    }
                )
        return (
            pd.DataFrame(rows, columns=columns)
            .sort_values(
                by=["全班得分率"],
                ascending=[True],
                na_position="last",
                kind="stable",
            )
            .reset_index(drop=True)
        )

    def _build_error_detail_sheet(
        self,
        df_details: pd.DataFrame,
        score_map: dict[str, float],
    ) -> pd.DataFrame:
        columns = [
            "班级",
            "学号",
            "学生姓名",
            "题号",
            "得分",
            "满分",
            "扣分",
            "扣分原因",
            "错误类别",
            "错误摘要",
        ]
        if df_details.empty:
            return pd.DataFrame(columns=columns)
        records = _normalize_question_detail_records(
            df_details.to_dict(orient="records"),
            score_map,
        )
        rows: list[dict[str, object]] = []
        for item in records:
            qid = str(item.get("question_id") or "")
            full_score = float(score_map.get(qid) or 0)
            awarded = float(item.get("score_awarded") or 0)
            if full_score <= 0 or awarded >= full_score - 1e-6:
                continue
            deduction_reason = _translate_deduction_reason(
                item.get("deduction_reason"),
                fallback="AI 未提供明确扣分依据，建议教师复核",
            )
            rows.append(
                {
                    "班级": item.get("class_name") or "未分班",
                    "学号": item.get("student_code"),
                    "学生姓名": item.get("student_name"),
                    "题号": qid,
                    "得分": awarded,
                    "满分": full_score,
                    "扣分": full_score - awarded,
                    "扣分原因": deduction_reason,
                    "错误类别": (
                        _public_grading_reason(item.get("error_category"))
                        or "—"
                    ),
                    "错误摘要": (
                        _public_grading_reason(item.get("error_summary"))
                        or "—"
                    ),
                }
            )
        qid_order = {
            qid: index
            for index, qid in enumerate(
                _natural_question_order(
                    [str(row["题号"]) for row in rows]
                )
            )
        }
        rows.sort(
            key=lambda row: (
                _class_sort_key(row["班级"]),
                str(row["学号"] or ""),
                qid_order.get(str(row["题号"]), len(qid_order)),
            )
        )
        return pd.DataFrame(rows, columns=columns)

    def _build_ai_teacher_comparison_sheet(
        self,
        df_details: pd.DataFrame,
        locks: list[dict[str, Any]],
        score_map: dict[str, float],
    ) -> pd.DataFrame:
        columns = [
            "班级",
            "学号",
            "学生姓名",
            "题号",
            "满分",
            "AI 得分",
            "人工得分",
            "最终得分",
        ]
        if df_details.empty:
            return pd.DataFrame(columns=columns)
        # A student can hold locks from several scan batches; the row query
        # orders by lock id, so the last write wins and keeps the latest lock.
        lock_by_key: dict[tuple[int, str], dict[str, Any]] = {}
        for lock in locks:
            key = (
                int(lock.get("student_id") or 0),
                str(lock.get("question_id") or ""),
            )
            lock_by_key[key] = lock
        rows: list[dict[str, object]] = []
        for item in df_details.to_dict(orient="records"):
            qid = str(item.get("question_id") or "")
            lock = lock_by_key.get(
                (int(item.get("student_id") or 0), qid)
            )
            # Only rows a teacher actually scored belong in this sheet;
            # unreviewed AI rows add no comparison information.
            if lock is None:
                continue
            full_score = float(score_map.get(qid) or 0)
            if full_score <= 0:
                full_score = float(lock.get("max_score") or 0)
            rows.append(
                {
                    "班级": item.get("class_name") or "未分班",
                    "学号": item.get("student_code"),
                    "学生姓名": item.get("student_name"),
                    "题号": qid,
                    "满分": full_score,
                    "AI 得分": _number_or_none(item.get("ai_score_awarded")),
                    "人工得分": float(lock["score_awarded"]),
                    "最终得分": _number_or_none(item.get("score_awarded")),
                }
            )
        qid_order = {
            qid: index
            for index, qid in enumerate(
                _natural_question_order(
                    [str(row["题号"]) for row in rows]
                )
            )
        }
        rows.sort(
            key=lambda row: (
                _class_sort_key(row["班级"]),
                str(row["学号"] or ""),
                str(row["学生姓名"] or ""),
                qid_order.get(str(row["题号"]), len(qid_order)),
            )
        )
        return pd.DataFrame(rows, columns=columns)

    def _build_exception_sheet(
        self,
        df_results: pd.DataFrame,
        df_attendance: pd.DataFrame,
    ) -> pd.DataFrame:
        columns = ["班级", "学号", "学生姓名", "状态", "说明"]
        rows: list[dict[str, object]] = []
        seen: set[tuple[str, str, str]] = set()
        for result in df_results.to_dict(orient="records"):
            if result.get("report_status") == "完整":
                continue
            key = (
                str(result.get("class_name") or "未分班"),
                str(result.get("student_code") or ""),
                str(result.get("student_name") or ""),
            )
            seen.add(key)
            missing = str(result.get("missing_questions") or "")
            rows.append(
                {
                    "班级": key[0],
                    "学号": key[1],
                    "学生姓名": key[2],
                    "状态": "批改不完整",
                    "说明": f"缺失题目：{missing}" if missing else "结果无效或不完整",
                }
            )
        attendance_labels = {
            "absent": "缺考",
            "scan_issue": "扫描异常",
            "present": "正常参考",
        }
        for item in df_attendance.to_dict(orient="records"):
            status = str(item.get("attendance_status") or "")
            if status == "present":
                continue
            key = (
                str(item.get("class_name") or "未分班"),
                str(item.get("student_code") or ""),
                str(item.get("student_name") or ""),
            )
            if key in seen:
                continue
            rows.append(
                {
                    "班级": key[0],
                    "学号": key[1],
                    "学生姓名": key[2],
                    "状态": attendance_labels.get(status, status or "异常"),
                    "说明": item.get("source_reason") or "",
                }
            )
        rows.sort(
            key=lambda row: (
                _class_sort_key(row["班级"]),
                str(row["学号"] or ""),
            )
        )
        return pd.DataFrame(rows, columns=columns)

    def _write_overview_sheet(
        self,
        sheet: Worksheet,
        session_name: str,
        eligible_results: pd.DataFrame,
        options: dict[str, object],
        hidden_student_ids: set[int],
        attendance_stats: dict[str, int] | None = None,
    ) -> None:
        sheet.merge_cells("A1:H1")
        sheet["A1"] = f"{session_name} · 考试总览"
        sheet.merge_cells("A2:H2")
        sheet["A2"] = "统计仅包含正常参考且完成全部题目批改的学生。姓名精简不会改变任何统计数字。"
        scores = [
            float(value)
            for value in eligible_results["student_score"].dropna().tolist()
        ]
        classes = sorted(
            {
                str(value or "未分班")
                for value in eligible_results["class_name"].tolist()
            },
            key=_class_sort_key,
        )
        sheet["A3"] = "统计人数"
        sheet["B3"] = len(eligible_results)
        sheet["C3"] = "班级数"
        sheet["D3"] = len(classes)
        sheet["E3"] = "整体均分"
        sheet["F3"] = round(sum(scores) / len(scores), 2) if scores else None
        sheet["G3"] = "最高 / 最低"
        sheet["H3"] = (
            f"{_format_score(max(scores))} / {_format_score(min(scores))}"
            if scores
            else "-"
        )
        if attendance_stats and attendance_stats.get("expected"):
            sheet.merge_cells("A4:H4")
            sheet["A4"] = (
                f"应考 {attendance_stats['expected']} 人"
                f" · 计入统计 {len(eligible_results)} 人"
                f" · 缺考或未计入 {attendance_stats.get('not_counted', 0)} 人"
                "（名单见“缺考与异常”页）"
            )
        sheet.merge_cells("A5:H5")
        sheet["A5"] = _excel_option_note(
            options,
            hidden_count=len(hidden_student_ids),
        )
        headers = ["班级", "统计人数", "均分", "中位数", "最高分", "最低分"]
        for column, header in enumerate(headers, start=1):
            sheet.cell(7, column, header)
        row = 8
        for class_name in classes:
            class_scores = [
                float(value)
                for value in eligible_results[
                    eligible_results["class_name"].fillna("未分班").astype(str)
                    == class_name
                ]["student_score"].dropna().tolist()
            ]
            values: list[object] = [
                class_name,
                len(class_scores),
                round(sum(class_scores) / len(class_scores), 2)
                if class_scores
                else None,
                round(statistics.median(class_scores), 2)
                if class_scores
                else None,
                max(class_scores) if class_scores else None,
                min(class_scores) if class_scores else None,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row, column, value)
            row += 1
        self._style_title_sheet(sheet, header_rows={7})
        sheet.freeze_panes = "A8"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 18
        for column in "BCDEFGH":
            sheet.column_dimensions[column].width = 14
        sheet.print_area = f"A1:H{max(8, row - 1)}"
        self._set_print_layout(sheet, landscape=True, repeat_rows="1:7")

    def _write_class_score_sheet(
        self,
        sheet: Worksheet,
        session_name: str,
        eligible_results: pd.DataFrame,
        pending_review_counts: dict[int, int] | None = None,
    ) -> None:
        sheet.merge_cells("A1:H1")
        sheet["A1"] = f"{session_name} · 班级成绩总表"
        current_row = 2
        class_names = sorted(
            {
                str(value or "未分班")
                for value in eligible_results["class_name"].tolist()
            },
            key=_class_sort_key,
        )
        pending_review_counts = pending_review_counts or {}
        for class_index, class_name in enumerate(class_names):
            class_df = eligible_results[
                eligible_results["class_name"].fillna("未分班").astype(str)
                == class_name
            ].copy()
            class_df.sort_values(
                by=[
                    "student_score",
                    "student_code",
                    "student_name",
                    "student_id",
                ],
                ascending=[False, True, True, True],
                kind="stable",
                inplace=True,
            )
            scores = [
                float(value)
                for value in class_df["student_score"].dropna().tolist()
            ]
            sheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=8,
            )
            sheet.cell(
                current_row,
                1,
                f"{class_name} · 正常参考且批改完整 {len(class_df)} 人",
            )
            current_row += 1
            summary = [
                "班级均分",
                round(sum(scores) / len(scores), 2) if scores else None,
                "中位数",
                round(statistics.median(scores), 2) if scores else None,
                "最高分",
                max(scores) if scores else None,
                "最低分",
                min(scores) if scores else None,
            ]
            for column, value in enumerate(summary, start=1):
                sheet.cell(current_row, column, value)
            current_row += 1
            headers = [
                "班级排名",
                "学号",
                "学生姓名",
                "学生得分",
                "试卷总分",
                "得分率",
                "待复核题数",
                "批改时间",
            ]
            header_row = current_row
            for column, header in enumerate(headers, start=1):
                sheet.cell(header_row, column, header)
            current_row += 1
            for rank, result in enumerate(
                class_df.to_dict(orient="records"),
                start=1,
            ):
                total_score = float(result.get("total_score") or 0)
                student_score = float(result.get("student_score") or 0)
                pending_count = pending_review_counts.get(
                    int(result.get("result_id") or 0), 0
                )
                if pending_count > 0:
                    review_cell: object = pending_count
                else:
                    review_cell = "是" if result.get("needs_human_review") else ""
                values = [
                    rank,
                    result.get("student_code"),
                    result.get("student_name"),
                    student_score,
                    total_score,
                    student_score / total_score if total_score > 0 else None,
                    review_cell,
                    result.get("graded_at"),
                ]
                for column, value in enumerate(values, start=1):
                    sheet.cell(current_row, column, value)
                sheet.cell(current_row, 6).number_format = "0.0%"
                current_row += 1
            self._style_table_header(sheet, header_row, 8)
            if class_index < len(class_names) - 1:
                sheet.row_breaks.append(Break(id=current_row - 1))
                current_row += 1
        self._style_title_sheet(sheet)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A4"
        widths = [10, 14, 14, 12, 12, 12, 12, 20]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.print_area = f"A1:H{max(3, current_row - 1)}"
        self._set_print_layout(sheet, landscape=True, repeat_rows="1:1")

    def _write_question_print_sheet(
        self,
        sheet: Worksheet,
        session_name: str,
        question_print: pd.DataFrame,
        hidden_student_ids: set[int],
        options: dict[str, object],
    ) -> None:
        sheet.merge_cells("A1:L1")
        sheet["A1"] = f"{session_name} · 小题分析打印"
        sheet.merge_cells("A2:L2")
        sheet["A2"] = "得分率、均分和失分人数始终按全部完整成绩计算；隐藏规则仅缩短“失分同学”姓名。"
        sheet.merge_cells("A3:L3")
        sheet["A3"] = _excel_option_note(
            options,
            hidden_count=len(hidden_student_ids),
        )
        headers = list(question_print.columns)
        for column, header in enumerate(headers, start=1):
            sheet.cell(5, column, header)
        for row_index, values in enumerate(
            dataframe_to_rows(question_print, index=False, header=False),
            start=6,
        ):
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value)
            sheet.cell(row_index, 6).number_format = "0.0%"
            sheet.cell(row_index, 7).number_format = "0.00"
        self._style_title_sheet(sheet, header_rows={5})
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A6"
        widths = [12, 10, 12, 9, 10, 12, 11, 10, 11, 11, 34, 28]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in range(6, sheet.max_row + 1):
            sheet.cell(row, 11).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            sheet.cell(row, 12).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            sheet.row_dimensions[row].height = 34
        sheet.auto_filter.ref = f"A5:L{max(5, sheet.max_row)}"
        sheet.print_area = f"A1:L{max(5, sheet.max_row)}"
        self._set_print_layout(sheet, landscape=True, repeat_rows="1:5")

    def _write_dataframe_sheet(
        self,
        sheet: Worksheet,
        frame: pd.DataFrame,
        *,
        freeze_cell: str,
        landscape: bool,
        top_notes: list[str] | None = None,
    ) -> None:
        top_notes = [text for text in (top_notes or []) if text]
        header_row = len(top_notes) + 1
        last_column = get_column_letter(max(1, len(frame.columns) or 1))
        for index, text in enumerate(top_notes, start=1):
            sheet.merge_cells(f"A{index}:{last_column}{index}")
            sheet.cell(index, 1, text)
            sheet.cell(index, 1).alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )
        for row_offset, values in enumerate(
            dataframe_to_rows(frame, index=False, header=True)
        ):
            for column, value in enumerate(values, start=1):
                sheet.cell(header_row + row_offset, column, value)
        self._style_title_sheet(sheet, header_rows={header_row})
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = freeze_cell
        if sheet.max_column > 0:
            sheet.auto_filter.ref = (
                f"A{header_row}:{get_column_letter(sheet.max_column)}"
                f"{max(header_row, sheet.max_row)}"
            )
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            values = [
                str(sheet.cell(row, column).value or "")
                for row in range(
                    header_row, min(sheet.max_row, header_row + 79) + 1
                )
            ]
            longest = max(
                (_excel_display_width(value) for value in values),
                default=8,
            )
            sheet.column_dimensions[letter].width = min(
                38,
                max(10, longest * 1.15 + 2),
            )
            header = str(sheet.cell(header_row, column).value or "").strip()
            if header == "得分率":
                for row_index in range(header_row + 1, sheet.max_row + 1):
                    sheet.cell(row_index, column).number_format = "0.0%"
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=header_row),
            start=header_row,
        ):
            required_lines = 1
            for cell in row:
                header = str(
                    sheet.cell(header_row, cell.column).value or ""
                ).strip()
                column_width = float(
                    sheet.column_dimensions[
                        get_column_letter(cell.column)
                    ].width
                    or 10
                )
                text_width = _excel_display_width(
                    str(cell.value or "")
                )
                wrap_text = (
                    isinstance(cell.value, str)
                    and text_width > max(8, column_width - 2)
                )
                cell.alignment = Alignment(
                    horizontal=(
                        "left"
                        if header
                        in {
                            "学生姓名",
                            "缺失题目",
                            "扣分原因",
                            "错误类别",
                            "错误摘要",
                            "知识点",
                            "涉及题目",
                            "说明",
                        }
                        or header.endswith("扣分原因")
                        else "center"
                    ),
                    vertical="top",
                    wrap_text=wrap_text,
                )
                if row_index > header_row:
                    cell.border = Border(
                        bottom=Side(
                            style="thin",
                            color="E7EEF4",
                        )
                    )
                if wrap_text:
                    required_lines = max(
                        required_lines,
                        math.ceil(
                            text_width / max(8, column_width - 2)
                        ),
                    )
            if row_index > header_row and required_lines > 1:
                sheet.row_dimensions[row_index].height = min(
                    72,
                    17 * required_lines,
                )
        self._set_print_layout(
            sheet,
            landscape=landscape,
            repeat_rows=f"1:{header_row}",
        )

    def _style_title_sheet(
        self,
        sheet: Worksheet,
        *,
        header_rows: set[int] | None = None,
    ) -> None:
        if sheet["A1"].value and sheet.merged_cells.ranges:
            sheet["A1"].font = Font(
                name="Microsoft YaHei",
                size=18,
                bold=True,
                color="FFFFFF",
            )
            sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
            sheet["A1"].alignment = Alignment(
                horizontal="left",
                vertical="center",
            )
            sheet.row_dimensions[1].height = 30
        for merged_range in sheet.merged_cells.ranges:
            if merged_range.min_row in {2, 3, 5}:
                cell = sheet.cell(
                    merged_range.min_row,
                    merged_range.min_col,
                )
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                )
        for header_row in header_rows or set():
            self._style_table_header(
                sheet,
                header_row,
                sheet.max_column,
            )
        for row in sheet.iter_rows():
            for cell in row:
                if cell.font.name is None:
                    cell.font = Font(name="Microsoft YaHei", size=10)

    def _style_table_header(
        self,
        sheet: Worksheet,
        row: int,
        max_column: int,
    ) -> None:
        for column in range(1, max_column + 1):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(
                name="Microsoft YaHei",
                size=10,
                bold=True,
                color="17365D",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(
                bottom=Side(style="thin", color="9FBAD0"),
            )
        sheet.row_dimensions[row].height = 26

    @staticmethod
    def _apply_used_range_borders(sheet: Worksheet) -> None:
        grid_side = Side(style="thin", color="AAB7C4")
        grid_border = Border(
            left=grid_side,
            right=grid_side,
            top=grid_side,
            bottom=grid_side,
        )
        for row in sheet.iter_rows(
            min_row=1,
            max_row=sheet.max_row,
            min_col=1,
            max_col=sheet.max_column,
        ):
            for cell in row:
                cell.border = grid_border

    @staticmethod
    def _set_print_layout(
        sheet: Worksheet,
        *,
        landscape: bool,
        repeat_rows: str,
    ) -> None:
        sheet.page_setup.orientation = (
            sheet.ORIENTATION_LANDSCAPE
            if landscape
            else sheet.ORIENTATION_PORTRAIT
        )
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = repeat_rows
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4

    def _build_compact_session_report(
        self,
        df_results: pd.DataFrame,
        df_details: pd.DataFrame,
        score_map: dict[str, float],
        *,
        rubric: dict | None = None,
        eligible_result_ids: set[int] | None = None,
    ) -> pd.DataFrame:
        question_ids = _natural_question_order(df_details["question_id"].dropna().astype(str).unique().tolist())
        rows_by_result: dict[int, dict[str, object]] = {}
        detail_rows_by_result: dict[int, list[dict[str, object]]] = {}
        for detail in df_details.to_dict(orient="records"):
            result_id = int(detail.get("result_id") or 0)
            detail_rows_by_result.setdefault(result_id, []).append(detail)
        # Rank only complete results. Summary statistics live once in the
        # dedicated print sheet instead of being repeated on every student row.
        class_stats = {}
        for class_name, group in df_results.groupby(df_results["class_name"].fillna("未分班")):
            if eligible_result_ids is not None:
                group = group[
                    group["result_id"].isin(eligible_result_ids)
                ].copy()
            if group.empty:
                continue
            group = group.sort_values(by="student_score", ascending=False)
            group["班级排名"] = group["student_score"].rank(method="min", ascending=False).astype(int)
            for _, r in group.iterrows():
                class_stats[int(r["result_id"])] = {
                    "班级排名": r["班级排名"],
                }

        for result in df_results.to_dict(orient="records"):
            rid = int(result["result_id"])
            stats = class_stats.get(rid, {"班级排名": "-"})
            completeness = resolve_grading_completeness(
                result.get("raw_json"),
                rubric=rubric if isinstance(rubric, dict) else None,
                details=detail_rows_by_result.get(rid, []),
            )
            completeness_fields = self._completeness_fields(
                {"grading_completeness": completeness} if isinstance(completeness, dict) else result.get("raw_json")
            )
            
            rows_by_result[rid] = {
                "班级": result.get("class_name") or "未分班",
                "班级排名": stats["班级排名"],
                "学生姓名": result.get("student_name"),
                "学号": result.get("student_code"),
                "总分": result.get("student_score"),
                "批改完整性": completeness_fields[0],
                "缺失题目": completeness_fields[1],
            }

        for detail in df_details.to_dict(orient="records"):
            result_id = int(detail.get("result_id") or 0)
            qid = str(detail.get("question_id") or "").strip()
            row = rows_by_result.get(result_id)
            if not qid or row is None:
                continue
            awarded = float(detail.get("score_awarded") or 0)
            full = score_map.get(qid)
            row[f"{qid}得分"] = f"{_format_score(awarded)}/{_format_score(full)}" if full is not None else _format_score(awarded)
            reason = _translate_deduction_reason(detail.get("deduction_reason"))
            if full is not None and awarded >= float(full) - 1e-6:
                reason = ""
            row[f"{qid}扣分原因"] = reason

        columns = ["班级", "班级排名", "学生姓名", "学号", "总分", "批改完整性", "缺失题目"]
        for qid in question_ids:
            columns.extend([f"{qid}得分", f"{qid}扣分原因"])
            
        df_out = pd.DataFrame(rows_by_result.values()).reindex(columns=columns)

        # Sort by class (numeric-aware: 9 before 10), then total score.
        df_out["_class_sort"] = df_out["班级"].map(_class_sort_key)
        df_out.sort_values(
            by=["_class_sort", "总分"],
            ascending=[True, False],
            inplace=True,
        )
        df_out.drop(columns=["_class_sort"], inplace=True)
        return df_out

    def _load_session_rubric(self, session_id: int) -> dict:
        rubric_path = self.repositories.reports.get_session_rubric_path(
            int(session_id)
        )
        return self._load_rubric_from_session(
            {"rubric_path": rubric_path} if rubric_path else None
        )

    def _load_rubric_from_session(
        self,
        session: dict | None,
    ) -> dict:
        if session is None:
            return {}
        rubric_path = self._resolve_stored_file_path(session.get("rubric_path"))
        if not rubric_path.exists():
            return {}
        try:
            rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
        except Exception:
            return {}
        return rubric if isinstance(rubric, dict) else {}

    def _load_session_question_maps(self, session_id: int) -> tuple[dict[str, float], dict[str, str]]:
        return self._question_maps_from_rubric(
            self._load_session_rubric(session_id)
        )

    def _load_session_answer_map(self, session: dict | None) -> dict[str, str]:
        """question_id -> canonical answer from the session's answer key."""
        if not isinstance(session, dict):
            return {}
        path = self._resolve_stored_file_path(session.get("answer_key_path"))
        if not path.exists():
            return {}
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
        answers: dict[str, str] = {}
        questions = data.get("questions") if isinstance(data, dict) else []
        for question in questions if isinstance(questions, list) else []:
            if not isinstance(question, dict):
                continue
            qid = str(question.get("question_id") or "").strip()
            canonical = str(question.get("canonical_answer") or "").strip()
            if qid and canonical:
                answers[qid] = canonical
            parts = question.get("parts")
            for part in parts if isinstance(parts, list) else []:
                if not isinstance(part, dict):
                    continue
                pid = str(part.get("part_id") or "").strip()
                part_answer = str(
                    part.get("canonical_answer") or part.get("answer") or ""
                ).strip()
                if pid and (part_answer or canonical):
                    answers[pid] = part_answer or canonical
        return answers

    def _question_maps_from_rubric(
        self,
        rubric: dict,
    ) -> tuple[dict[str, float], dict[str, str]]:
        if not rubric:
            return {}, {}
        score_map: dict[str, float] = {}
        type_map: dict[str, str] = {}
        questions = rubric.get("questions") if isinstance(rubric, dict) else []
        if not isinstance(questions, list):
            return score_map, type_map
        for question in questions:
            if not isinstance(question, dict):
                continue
            qid = str(question.get("question_id") or "").strip()
            qtype = str(question.get("question_type") or "").strip()
            if qid:
                score_map[qid] = float(question.get("max_score") or 0)
                type_map[qid] = qtype
            parts = question.get("parts")
            if isinstance(parts, list):
                for part in parts:
                    if not isinstance(part, dict):
                        continue
                    pid = str(part.get("part_id") or "").strip()
                    if pid:
                        score_map[pid] = float(part.get("part_score") or 0)
                        type_map[pid] = qtype
        return score_map, type_map

    def _knowledge_label_map_from_rubric(
        self,
        rubric: dict,
    ) -> dict[str, str]:
        if not rubric:
            return {}

        labels: dict[str, str] = {}
        questions = rubric.get("questions") if isinstance(rubric, dict) else []
        if not isinstance(questions, list):
            return labels
        for question in questions:
            if not isinstance(question, dict):
                continue
            for point in question.get("knowledge_points", []) or []:
                if not isinstance(point, dict):
                    continue
                kid = str(point.get("knowledge_id") or point.get("id") or "").strip()
                label = _clean_knowledge_label(kid, point.get("knowledge_name") or point.get("name") or "")
                if kid and label:
                    labels[kid] = label
            primary_id = str(question.get("knowledge_id") or "").strip()
            primary_label = _clean_knowledge_label(
                primary_id,
                question.get("knowledge_name") or question.get("stem_summary") or "",
            )
            if primary_id and primary_label:
                labels.setdefault(primary_id, primary_label)
        return labels

    def _completeness_fields(self, raw_json: object) -> tuple[str, str]:
        completeness = resolve_grading_completeness(raw_json)
        if not isinstance(completeness, dict):
            return "无效", ""
        status = str(completeness.get("status") or "").strip()
        status_label = {
            "complete": "完整",
            "incomplete": "不完整",
            "invalid": "无效",
        }.get(status, "无效")
        missing_ids = completeness.get("missing_question_ids") if isinstance(completeness, dict) else []
        missing_list = [str(item).strip() for item in missing_ids if str(item).strip()] if isinstance(missing_ids, list) else []
        return status_label, "、".join(missing_list)

    def _resolve_stored_file_path(self, path_value: object) -> Path:
        data_root = self.db_path.parent.parent if self.db_path.parent.name == "databases" else None
        return resolve_stored_file_path(path_value, data_root=data_root)

    def _question_bank_db_path(self) -> Path | None:
        data_root = (
            self.db_path.parent.parent
            if self.db_path.parent.name == "databases"
            else self.db_path.parent
        )
        candidate = data_root / "databases" / "question_bank.db"
        return candidate if candidate.is_file() else None

    def _build_session_knowledge_summary_by_class(
        self,
        df_details: pd.DataFrame,
        score_map: dict[str, float],
        knowledge_label_map: dict[str, str],
        knowledge_backfill: dict[str, list[dict[str, str]]] | None = None,
    ) -> pd.DataFrame:
        columns = ["班级", "知识点", "涉及题目", "累计得分", "累计满分", "得分率", "失分人数"]
        if df_details.empty:
            return pd.DataFrame(columns=columns)

        detail_records = _normalize_question_detail_records(df_details.to_dict(orient="records"), score_map)
        if not detail_records:
            return pd.DataFrame(columns=columns)

        buckets: dict[tuple[str, str], dict[str, object]] = {}
        lost_students: dict[tuple[str, str], set[str]] = {}
        for item in detail_records:
            qid = str(item.get("question_id") or "").strip()
            full_score = float(score_map.get(qid) or 0)
            if full_score <= 0:
                continue
            awarded = max(0.0, min(float(item.get("score_awarded") or 0), full_score))
            class_name = str(item.get("class_name") or "未分班").strip() or "未分班"
            student_name = _student_label(item)
            for bucket_key, display_label in _knowledge_bucket_labels(
                item,
                knowledge_label_map,
                knowledge_backfill,
            ):
                key = (class_name, bucket_key)
                bucket = buckets.setdefault(
                    key,
                    {
                        "label": display_label,
                        "score_sum": 0.0,
                        "full_sum": 0.0,
                        "questions": set(),
                    },
                )
                bucket["score_sum"] = float(bucket["score_sum"]) + awarded
                bucket["full_sum"] = float(bucket["full_sum"]) + full_score
                questions = bucket["questions"]
                if isinstance(questions, set):
                    questions.add(qid)
                if awarded < full_score - 1e-6:
                    lost_students.setdefault(key, set()).add(student_name)

        rows: list[dict[str, object]] = []
        for (class_name, _bucket_key), bucket in buckets.items():
            score_sum = float(bucket["score_sum"])
            full_sum = float(bucket["full_sum"])
            questions = bucket["questions"] if isinstance(bucket["questions"], set) else set()
            rows.append(
                {
                    "班级": class_name,
                    "知识点": str(bucket.get("label") or "未命名知识点"),
                    "涉及题目": "、".join(_natural_question_order([str(q) for q in questions])),
                    "累计得分": round(score_sum, 2),
                    "累计满分": round(full_sum, 2),
                    "得分率": round(score_sum / full_sum, 4) if full_sum > 0 else 0,
                    "失分人数": len(lost_students.get((class_name, _bucket_key), set())),
                }
            )
        rows.sort(
            key=lambda row: (
                _class_sort_key(row["班级"]),
                float(row["得分率"] or 0),
                str(row["知识点"] or ""),
            )
        )
        return pd.DataFrame(rows, columns=columns)

def _normalized_excel_options(
    raw_options: dict[str, object] | None,
) -> dict[str, object]:
    options = dict(raw_options or {})
    enabled = bool(options.get("hide_bottom_enabled", True))
    try:
        bottom_n = max(
            0,
            min(100, int(options.get("hide_bottom_n", 8))),
        )
    except (TypeError, ValueError):
        bottom_n = 8
    if not enabled:
        bottom_n = 0
    raw_ids = options.get("manual_hidden_student_ids")
    manual_ids = (
        sorted(
            {
                int(student_id)
                for student_id in raw_ids
                if isinstance(student_id, int)
                and not isinstance(student_id, bool)
                and student_id > 0
            }
        )
        if isinstance(raw_ids, list)
        else []
    )
    return {
        "hide_bottom_enabled": enabled,
        "hide_bottom_n": bottom_n,
        "manual_hidden_student_ids": manual_ids,
    }


def _excel_option_note(
    options: dict[str, object],
    *,
    hidden_count: int,
) -> str:
    parts: list[str] = []
    bottom_n = int(options.get("hide_bottom_n") or 0)
    if options.get("hide_bottom_enabled") and bottom_n > 0:
        parts.append(f"每班隐藏总分最后 {bottom_n} 名")
    manual_count = len(
        options.get("manual_hidden_student_ids")
        if isinstance(options.get("manual_hidden_student_ids"), list)
        else []
    )
    if manual_count:
        parts.append(f"另手动选择 {manual_count} 人")
    if not parts:
        return "错题姓名不精简；所有失分学生姓名均显示。"
    return (
        "错题姓名精简："
        + "，".join(parts)
        + f"；实际命中 {hidden_count} 人。"
    )


def _pending_review_counts(df_results: pd.DataFrame) -> dict[int, int]:
    """result_id -> count of questions still flagged need_review."""
    counts: dict[int, int] = {}
    for row in df_results.to_dict(orient="records"):
        raw = row.get("raw_json")
        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except (json.JSONDecodeError, TypeError):
                raw = {}
        metadata = raw.get("detail_metadata") if isinstance(raw, dict) else None
        pending = 0
        if isinstance(metadata, dict):
            for item in metadata.values():
                if isinstance(item, dict) and (
                    item.get("need_review") or item.get("needs_human_review")
                ):
                    pending += 1
        counts[int(row.get("result_id") or 0)] = pending
    return counts


def _natural_question_order(question_ids: list[str]) -> list[str]:
    import re

    def key(value: str) -> tuple[int, str]:
        match = re.search(r"\d+", value)
        number = int(match.group(0)) if match else 10**9
        return number, value

    return sorted(question_ids, key=key)


def _normalize_question_detail_records(
    records: list[dict[str, object]],
    score_map: dict[str, float],
) -> list[dict[str, object]]:
    merged: dict[tuple[str, str], dict[str, object]] = {}
    order: list[tuple[str, str]] = []
    for index, record in enumerate(records):
        raw_qid = str(record.get("question_id") or "").strip()
        if not raw_qid:
            continue
        qid = _canonical_question_id_for_score(raw_qid, score_map)
        result_key = str(
            record.get("result_id")
            or record.get("student_code")
            or record.get("student_name")
            or index
        )
        key = (result_key, qid)
        if key not in merged:
            item = dict(record)
            item["question_id"] = qid
            item["score_awarded"] = 0.0
            item["_deduction_reasons"] = []
            item["_error_categories"] = []
            item["_error_summaries"] = []
            item["_knowledge_ids"] = []
            merged[key] = item
            order.append(key)

        item = merged[key]
        item["score_awarded"] = float(item.get("score_awarded") or 0) + float(record.get("score_awarded") or 0)
        for field, private_field in (
            ("deduction_reason", "_deduction_reasons"),
            ("error_category", "_error_categories"),
            ("error_summary", "_error_summaries"),
        ):
            value = _clean_grading_text(record.get(field))
            if value:
                item[private_field].append(value)
        item["_knowledge_ids"].extend(kid for kid in _knowledge_ids_from_detail(record) if kid != "UNKNOWN")

    normalized: list[dict[str, object]] = []
    for key in order:
        item = merged[key]
        qid = str(item.get("question_id") or "")
        full_score = float(score_map.get(qid) or 0)
        if full_score > 0 and float(item.get("score_awarded") or 0) > full_score:
            item["score_awarded"] = full_score

        for private_field, public_field in (
            ("_deduction_reasons", "deduction_reason"),
            ("_error_categories", "error_category"),
            ("_error_summaries", "error_summary"),
        ):
            values = _unique_texts(item.pop(private_field, []))
            if values:
                item[public_field] = "；".join(values)

        knowledge_ids = _unique_texts(item.pop("_knowledge_ids", []))
        if knowledge_ids:
            item["knowledge_id"] = knowledge_ids[0]
            item["knowledge_ids"] = json.dumps(knowledge_ids, ensure_ascii=False)
        normalized.append(item)
    return normalized


def _canonical_question_id_for_score(question_id: str, score_map: dict[str, float]) -> str:
    qid = question_id.strip()
    return resolve_known_question_id(qid, score_map) or qid


def _excel_display_width(value: str) -> int:
    return sum(
        2
        if unicodedata.east_asian_width(character) in {"W", "F"}
        else 1
        for character in value
    )


def _unique_texts(values: list[object]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _format_score(value) -> str:
    if value is None:
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


_OBJECTIVE_ANSWER_PREFIX = "objective_answer="


def _clean_grading_text(value: object) -> str:
    """Return stripped text, treating pandas NaN/None-ish markers as empty."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null", "nat"}:
        return ""
    return text


def _public_grading_reason(value: object, fallback: str = "") -> str:
    text = _clean_grading_text(value)
    if not text:
        return fallback
    labels = {
        "objective_api_disabled": "客观题识别未启用，已转教师复核",
        "objective_api_not_configured": "客观题识别模型配置不完整，已转教师复核",
        "objective_paper_model_failed": "客观题识别请求失败，已转教师复核",
        "objective_paper_region_failed": "无法读取选填题作答区域",
        "objective_region_not_found": "未找到此题的有效作答区域",
        "missing_question_result": "AI 未返回此题的识别结果",
        "duplicate_question_result": "AI 返回了重复的识别结果",
        "paper_key_mismatch": "识别结果与当前答卷不一致",
        "low_confidence": "作答辨识度较低，需要教师复核",
        "needs_review": "AI 建议教师复核",
        "objective_needs_review": "客观题识别结果需要教师复核",
        "objective_score_uncertain": "答案识别存在不确定性",
        "teacher_score_locked": "教师已确认最终分",
        "no_numeric_value": "无法确定数值，需复核",
        "equivalence_uncertain": "等价关系待判定，需复核",
    }
    if text.lower() in labels:
        return labels[text.lower()]
    if re.search(r"[\u3400-\u9fff]", text):
        return text
    return fallback or "自动处理未完成，请教师复核"


def _translate_deduction_reason(value: object, fallback: str = "") -> str:
    """Translate a stored deduction_reason for teacher-facing sheets."""
    text = _clean_grading_text(value)
    if not text:
        return fallback
    if text.startswith(_OBJECTIVE_ANSWER_PREFIX):
        answer = text[len(_OBJECTIVE_ANSWER_PREFIX):].strip()
        return f"识别作答：{answer}" if answer else "识别作答缺失"
    return _public_grading_reason(text, fallback=fallback)


def _class_sort_key(name: object) -> tuple[int, int, str]:
    """Order class labels numerically (9 before 10), text labels last."""
    text = str(name or "").strip()
    digits = re.sub(r"\D", "", text)
    if digits:
        return (0, int(digits), text)
    return (1, 0, text)


def _student_label(item: dict) -> str:
    name = str(item.get("student_name") or "").strip()
    code = str(item.get("student_code") or "").strip()
    return name or code or "未知学生"


def _number_or_none(value: object) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(number):
        return None
    return number


def _knowledge_ids_from_detail(item: dict) -> list[str]:
    raw = item.get("knowledge_ids")
    values: list[str] = []
    if isinstance(raw, str) and raw.strip():
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            parsed = raw
        if isinstance(parsed, list):
            values.extend(str(value).strip() for value in parsed if str(value).strip())
        else:
            values.extend(_split_knowledge_text(str(parsed)))
    elif isinstance(raw, list):
        values.extend(str(value).strip() for value in raw if str(value).strip())
    fallback = str(item.get("knowledge_id") or "").strip()
    if fallback:
        values.append(fallback)

    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result or ["UNKNOWN"]


def _split_knowledge_text(value: str) -> list[str]:
    return [
        part.strip()
        for part in value.replace("；", ",").replace(";", ",").replace("|", ",").replace("、", ",").split(",")
        if part.strip()
    ]


def _clean_knowledge_label(knowledge_id: str, label: object) -> str:
    import re

    text = str(label or "").strip()
    kid = str(knowledge_id or "").strip()
    if kid and text.startswith(kid):
        text = text[len(kid):].strip()
        text = re.sub(r"^[\s:：\-|·、]+", "", text).strip()
    if not text or text == kid or re.fullmatch(r"[A-Za-z]+\d*_\d+", text):
        return "未命名知识点"
    return text


_PART_SUFFIX_RE = re.compile(r"[\(（]\s*P?\s*\d+\s*[\)）]\s*$", re.IGNORECASE)


def _parent_question_id(question_id: str) -> str:
    """Strip a trailing part suffix, e.g. ``Q10(P1)`` -> ``Q10``."""
    return _PART_SUFFIX_RE.sub("", str(question_id or "").strip())


def _knowledge_bucket_labels(
    item: dict,
    knowledge_label_map: dict[str, str],
    knowledge_backfill: dict[str, list[dict[str, str]]] | None = None,
) -> list[tuple[str, str]]:
    """Return (bucket_key, display_label) pairs for one detail record.

    Stored knowledge ids that resolve through the rubric label map keep the
    historical behavior (bucket and display on the same label).  When nothing
    resolves — new sessions only persist placeholder ids — fall back to the
    read-only question-bank backfill, bucketing by the full hierarchical tag
    path (so same-named leaves in different chapters stay separate) while
    displaying the leaf label.  Part-level ids such as ``Q10(P1)`` fall back
    to their parent question ``Q10`` for the backfill lookup.
    """
    labels: list[str] = []
    for kid in _knowledge_ids_from_detail(item):
        label = _clean_knowledge_label(kid, knowledge_label_map.get(kid, ""))
        if label and label != "未命名知识点" and label not in labels:
            labels.append(label)
    backfill_map = knowledge_backfill or {}
    qid = str(item.get("question_id") or "").strip()
    refined_entries = backfill_map.get(qid, backfill_map.get(_parent_question_id(qid)))
    refined = refined_entries == [] or any("stable_key" in entry for entry in refined_entries or [])
    if labels and not refined:
        return [(label, label) for label in labels]

    entries = backfill_map.get(qid)
    if entries is None and qid:
        parent = _parent_question_id(qid)
        if parent and parent != qid:
            entries = backfill_map.get(parent)
    pairs: list[tuple[str, str]] = []
    seen: set[str] = set()
    for entry in entries or []:
        path = str(entry.get("path") or "").strip()
        if not path or path in seen:
            continue
        seen.add(path)
        label = str(entry.get("label") or "").strip() or "未命名知识点"
        pairs.append((path, label))
    return pairs or [("未命名知识点", "未命名知识点")]


def _loss_entry_label(item: dict) -> str:
    """Short per-item label used by the 主要错因 aggregation.

    Prefer the human-readable summary; fall back to the recognized answer
    stored in ``deduction_reason`` (``objective_answer=X``), then to the
    stored reason/category text.
    """
    summary = _clean_grading_text(item.get("error_summary"))
    if summary:
        return _short_loss_label(_public_grading_reason(summary))
    reason = _clean_grading_text(item.get("deduction_reason"))
    if reason.startswith(_OBJECTIVE_ANSWER_PREFIX):
        answer = reason[len(_OBJECTIVE_ANSWER_PREFIX):].strip()
        return answer or "未识别作答"
    if reason:
        return _short_loss_label(_public_grading_reason(reason))
    category = _clean_grading_text(item.get("error_category"))
    if category:
        return _short_loss_label(_public_grading_reason(category))
    return "原因未记录"


def _short_loss_label(text: str) -> str:
    """Collapse boilerplate reasons into short countable labels."""
    normalized = str(text or "").strip()
    if not normalized:
        return ""
    if "未见有效作答" in normalized or "未作答" in normalized or "未完成作答" in normalized:
        return "未作答"
    if "作废" in normalized:
        return "答案作废"
    if "人工复核已确认" in normalized or "教师已确认" in normalized:
        return "教师已确认"
    return normalized


def _summarize_error_categories(
    items: list[dict],
    full_score: float,
    *,
    canonical_answer: str = "",
    is_objective: bool = False,
) -> str:
    counts: dict[str, int] = {}
    for item in items:
        awarded = float(item.get("score_awarded") or 0)
        if full_score > 0 and awarded >= full_score - 1e-6:
            continue
        label = _loss_entry_label(item)
        if not label:
            continue
        counts[label] = counts.get(label, 0) + 1
    text = "、".join(
        f"{key}×{value}"
        for key, value in sorted(
            counts.items(), key=lambda kv: (-kv[1], kv[0])
        )[:5]
    )
    if is_objective and canonical_answer and text:
        text = f"{text}（正确：{canonical_answer}）"
    return text
