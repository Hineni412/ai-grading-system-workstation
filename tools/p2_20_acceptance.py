from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import hmac
import json
import math
import os
import re
import secrets
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from collections.abc import Awaitable, Callable
from pathlib import Path
from typing import Any
from urllib.error import URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path[:1]:
    sys.path.insert(0, str(REPO_ROOT))

from tools import p1_29_acceptance as workspace_tools
from question_id_contract import QuestionIdCatalog, QuestionIdContractError


PACKAGE = "P2-20"
METADATA_FILENAME = "p2-20-acceptance.json"
SERVER_STATE_FILENAME = "server-state.json"
SERVER_CLAIM_FILENAME = "server-start.claim"
SERVER_STOP_FILENAME = "server-stop.request"
_CONFIG_KEY_RE = re.compile(
    r"^[ \t]*(?P<key>DATA_DIR|LOGS_DIR|[\"']DATA_DIR[\"']|[\"']LOGS_DIR[\"'])[ \t]*:"
)
_SAFE_FINISH_REASONS = frozenset(
    {"stop", "length", "content_filter", "tool_calls", "function_call"}
)
_CAPTURE_NEXT_MODEL_REQUEST_ENV = "AI_GRADING_P2_20_CAPTURE_NEXT_MODEL_REQUEST"
_MANUAL_MODEL_ROUTES = {
    "ohmycdn": "https://apic1.ohmycdn.com/v1",
    "ohmygpt": "https://api.ohmygpt.com/v1",
}
_DEFAULT_MANUAL_MODEL_ROUTE = "ohmygpt"
_MANUAL_MODEL_ACTIONS = frozenset({"q11", "q12", "score"})
_MANUAL_MODEL_TOKEN_RE = re.compile(r"^[0-9a-f]{32}$")
_BATCH_QUESTION_IDS_RE = re.compile(
    r"BATCH_QUESTION_IDS_JSON=(\[[^\r\n]+\])"
)
_SCORE_QUESTION_IDS_RE = re.compile(
    r"SCORE_QUESTION_IDS_JSON=(\[[^\r\n]+\])"
)
_MANUAL_OUTCOME_KEYS = frozenset(
    {"success", "upstream_http_error", "transport_error", "unknown"}
)


class AcceptanceError(RuntimeError):
    """Raised when the P2-20 acceptance workspace is unsafe."""


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


def validate_workspace(
    workspace: Path | str,
    *,
    require_empty: bool = True,
) -> Path:
    candidate = Path(workspace).expanduser().resolve()
    temp_root = Path(tempfile.gettempdir()).resolve()
    if candidate == temp_root or not _is_relative_to(candidate, temp_root):
        raise AcceptanceError(
            "workspace must be strictly below the system temporary directory"
        )
    if any(part.casefold() == "user_data" for part in candidate.parts):
        raise AcceptanceError("workspace must not contain a user_data path segment")
    if candidate.exists() and not candidate.is_dir():
        raise AcceptanceError("workspace must be an empty directory")
    if require_empty and candidate.exists() and next(candidate.iterdir(), None) is not None:
        raise AcceptanceError("workspace must be empty")
    return candidate


def load_metadata(workspace: Path | str) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata_path = target / METADATA_FILENAME
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("prepared workspace metadata is unavailable") from exc
    common_keys = {
        "package",
        "source_sha",
        "state",
        "authorized_session_id",
    }
    if not isinstance(metadata, dict) or not common_keys.issubset(metadata):
        raise AcceptanceError("prepared workspace metadata is invalid")
    state = metadata["state"]
    if metadata["package"] != PACKAGE or state not in {
        "code_prepared",
        "inputs_prepared",
        "runtime_ready",
    }:
        raise AcceptanceError("prepared workspace metadata is invalid")
    if state == "code_prepared":
        if set(metadata) != common_keys or metadata["authorized_session_id"] is not None:
            raise AcceptanceError("prepared workspace metadata is invalid")
    else:
        input_keys = common_keys | {
            "paper_count",
            "source_fingerprint",
            "input_manifest",
        }
        if state == "inputs_prepared":
            expected_keys = input_keys
        else:
            expected_keys = input_keys | {
                "profile_fingerprint",
                "model_request_budget",
            }
            if metadata.get("model_request_budget") is None:
                expected_keys |= {
                    "model_request_count_baseline",
                    "direct_diagnostic_requests",
                }
        if set(metadata) != expected_keys:
            raise AcceptanceError("prepared workspace metadata is invalid")
        if (
            not isinstance(metadata["authorized_session_id"], int)
            or not isinstance(metadata["paper_count"], int)
            or not isinstance(metadata["source_fingerprint"], str)
            or len(metadata["source_fingerprint"]) != 64
            or metadata["input_manifest"] != "acceptance_inputs/manifest.json"
        ):
            raise AcceptanceError("prepared workspace metadata is invalid")
        if state == "runtime_ready" and (
            not isinstance(metadata["profile_fingerprint"], str)
            or len(metadata["profile_fingerprint"]) != 64
            or (
                metadata["model_request_budget"] is not None
                and (
                    isinstance(metadata["model_request_budget"], bool)
                    or not isinstance(metadata["model_request_budget"], int)
                )
            )
        ):
            raise AcceptanceError("prepared workspace metadata is invalid")
        if state == "runtime_ready" and metadata["model_request_budget"] is None:
            if (
                isinstance(metadata["model_request_count_baseline"], bool)
                or not isinstance(metadata["model_request_count_baseline"], int)
                or metadata["model_request_count_baseline"] < 0
                or isinstance(metadata["direct_diagnostic_requests"], bool)
                or not isinstance(metadata["direct_diagnostic_requests"], int)
                or metadata["direct_diagnostic_requests"] < 0
            ):
                raise AcceptanceError("prepared workspace metadata is invalid")
    source_sha = metadata["source_sha"]
    if not isinstance(source_sha, str) or len(source_sha) != 40:
        raise AcceptanceError("prepared workspace metadata is invalid")
    return metadata


def prepare_code_workspace(
    source_ref: str,
    workspace: Path | str,
    *,
    repo_root: Path,
) -> dict[str, object]:
    target = validate_workspace(workspace)
    try:
        source_sha = workspace_tools.resolve_source_sha(
            source_ref,
            repo_root=Path(repo_root).resolve(),
        )
    except workspace_tools.AcceptanceError as exc:
        raise AcceptanceError(str(exc)) from exc
    target.mkdir(parents=True, exist_ok=False)
    try:
        workspace_tools._archive_commit(
            source_sha,
            repo_root=Path(repo_root).resolve(),
            workspace=target,
        )
    except workspace_tools.AcceptanceError as exc:
        raise AcceptanceError(str(exc)) from exc
    metadata: dict[str, object] = {
        "package": PACKAGE,
        "source_sha": source_sha,
        "state": "code_prepared",
        "authorized_session_id": None,
    }
    temporary_path = target / f"{METADATA_FILENAME}.tmp"
    temporary_path.write_text(
        json.dumps(metadata, ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    temporary_path.replace(target / METADATA_FILENAME)
    return load_metadata(target)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _database_generation(database: Path) -> dict[str, str | None]:
    return {
        suffix or "main": (
            _sha256_file(Path(f"{database}{suffix}"))
            if Path(f"{database}{suffix}").is_file()
            else None
        )
        for suffix in ("", "-wal", "-shm")
    }


def _resolve_source_file(raw_path: object, source_data_root: Path) -> Path:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise AcceptanceError("authorized exam contains a missing source file")
    candidate = Path(raw_path)
    candidates = [candidate]
    if not candidate.is_absolute():
        candidates.append(source_data_root / candidate)
    parts = list(candidate.parts)
    user_data_indexes = [
        index for index, part in enumerate(parts) if part.casefold() == "user_data"
    ]
    if user_data_indexes:
        relative_parts = parts[user_data_indexes[-1] + 1 :]
        candidates.append(source_data_root.joinpath(*relative_parts))
    resolved_root = source_data_root.resolve()
    for possible in candidates:
        try:
            resolved = possible.resolve(strict=True)
        except OSError:
            continue
        if (
            resolved.is_file()
            and _is_relative_to(resolved, resolved_root)
            and not resolved.is_symlink()
        ):
            return resolved
    raise AcceptanceError("authorized exam source file is missing or outside the data root")


def _combined_fingerprint(
    database_generation: dict[str, str | None],
    file_hashes: dict[str, str],
) -> str:
    canonical = json.dumps(
        {"database": database_generation, "files": file_hashes},
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest()


def _write_metadata(workspace: Path, metadata: dict[str, object]) -> None:
    temporary_path = workspace / f"{METADATA_FILENAME}.tmp"
    temporary_path.write_text(
        json.dumps(metadata, ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    temporary_path.replace(workspace / METADATA_FILENAME)


def prepare_authorized_inputs(
    workspace: Path | str,
    *,
    source_data_root: Path | str,
    session_id: int,
    expected_session_name: str,
    paper_limit: int,
) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "code_prepared":
        raise AcceptanceError("authorized inputs have already been prepared")
    if (
        isinstance(session_id, bool)
        or not isinstance(session_id, int)
        or session_id < 1
        or not isinstance(expected_session_name, str)
        or not expected_session_name.strip()
    ):
        raise AcceptanceError("authorized exam identity is invalid")
    if (
        isinstance(paper_limit, bool)
        or not isinstance(paper_limit, int)
        or paper_limit < 1
        or paper_limit > 3
    ):
        raise AcceptanceError("paper limit must be between 1 and 3")

    data_root = Path(source_data_root).expanduser().resolve()
    if not data_root.is_dir() or _is_relative_to(data_root, target):
        raise AcceptanceError("source data root is missing or unsafe")
    database = data_root / "databases" / "grading_system.db"
    if not database.is_file():
        raise AcceptanceError("source grading database is missing")

    database_before = _database_generation(database)
    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        session = connection.execute(
            """
            SELECT id, session_name, source_paper_path
            FROM grading_sessions
            WHERE id = ? AND COALESCE(is_deleted, 0) = 0
            """,
            (session_id,),
        ).fetchone()
        if session is None or session["session_name"] != expected_session_name:
            raise AcceptanceError("authorized exam does not match the requested session")
        confirmed_template = connection.execute(
            """
            SELECT 1
            FROM session_templates
            WHERE session_id = ? AND is_confirmed = 1
            LIMIT 1
            """,
            (session_id,),
        ).fetchone()
        if confirmed_template is None:
            raise AcceptanceError("authorized exam has no confirmed template")
        papers = connection.execute(
            """
            SELECT DISTINCT
                p.id AS paper_id,
                p.front_image,
                p.back_image,
                s.student_code,
                s.name,
                s.class_name
            FROM exam_papers AS p
            JOIN students AS s ON s.id = p.student_id
            JOIN session_results AS r
              ON r.session_id = p.session_id
             AND r.paper_id = p.id
             AND r.student_id = s.id
            WHERE p.session_id = ?
              AND p.match_status = 'matched'
              AND p.processing_status = 'graded'
              AND TRIM(COALESCE(s.student_code, '')) <> ''
              AND TRIM(COALESCE(s.name, '')) <> ''
            ORDER BY p.id
            LIMIT ?
            """,
            (session_id, paper_limit),
        ).fetchall()
    except sqlite3.Error as exc:
        raise AcceptanceError("authorized exam database could not be read safely") from exc
    finally:
        connection.close()
    if len(papers) != paper_limit:
        raise AcceptanceError("authorized exam has too few eligible answer sheets")

    source_paper = _resolve_source_file(session["source_paper_path"], data_root)
    template_candidates = sorted(
        (data_root / "templates" / f"session_{session_id}").glob("*.pdf")
    )
    if len(template_candidates) != 1:
        raise AcceptanceError("authorized exam must have exactly one template PDF")
    template_pdf = _resolve_source_file(str(template_candidates[0]), data_root)

    selected_files: list[tuple[str, Path]] = [
        (f"source-paper{source_paper.suffix.lower()}", source_paper),
        ("template.pdf", template_pdf),
    ]
    for index, paper in enumerate(papers, start=1):
        front = _resolve_source_file(paper["front_image"], data_root)
        back = _resolve_source_file(paper["back_image"], data_root)
        selected_files.extend(
            [
                (f"answer-sheets/answer-{index:02}-front{front.suffix.lower()}", front),
                (f"answer-sheets/answer-{index:02}-back{back.suffix.lower()}", back),
            ]
        )
    hashes_before = {
        logical_name: _sha256_file(source_path)
        for logical_name, source_path in selected_files
    }

    staging = target / "acceptance_inputs-preparing"
    published = target / "acceptance_inputs"
    backup = target / "authorized-inputs-backup.zip"
    if staging.exists() or published.exists() or backup.exists():
        raise AcceptanceError("acceptance input destination is not empty")
    staging.mkdir()
    for logical_name, source_path in selected_files:
        destination = staging.joinpath(*Path(logical_name).parts)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source_path, destination)
    roster_path = staging / "roster.csv"
    with roster_path.open("w", encoding="utf-8-sig", newline="") as roster_file:
        writer = csv.writer(roster_file, lineterminator="\n")
        writer.writerow(("student_code", "name", "class_name"))
        for paper in papers:
            writer.writerow(
                (
                    paper["student_code"],
                    paper["name"],
                    paper["class_name"] or "",
                )
            )
    manifest = {
        "package": PACKAGE,
        "authorized_session_id": session_id,
        "paper_count": paper_limit,
        "files": {
            logical_name: hashes_before[logical_name]
            for logical_name, _source_path in selected_files
        },
    }
    (staging / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=True, sort_keys=True, separators=(",", ":")),
        encoding="utf-8",
    )

    hashes_after = {
        logical_name: _sha256_file(source_path)
        for logical_name, source_path in selected_files
    }
    database_after = _database_generation(database)
    if hashes_after != hashes_before or database_after != database_before:
        raise AcceptanceError("authorized source changed while the copy was prepared")

    with zipfile.ZipFile(
        backup,
        mode="x",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for path in sorted(staging.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(staging).as_posix())
    staging.replace(published)
    prepared_metadata: dict[str, object] = {
        "package": PACKAGE,
        "source_sha": metadata["source_sha"],
        "state": "inputs_prepared",
        "authorized_session_id": session_id,
        "paper_count": paper_limit,
        "source_fingerprint": _combined_fingerprint(
            database_before,
            hashes_before,
        ),
        "input_manifest": "acceptance_inputs/manifest.json",
    }
    _write_metadata(target, prepared_metadata)
    return load_metadata(target)


def _current_authorized_source_fingerprint(
    source_data_root: Path | str,
    *,
    session_id: int,
    paper_limit: int,
) -> str:
    data_root = Path(source_data_root).expanduser().resolve()
    database = data_root / "databases" / "grading_system.db"
    if not data_root.is_dir() or not database.is_file():
        raise AcceptanceError("authorized source data is unavailable")

    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        session = connection.execute(
            """
            SELECT source_paper_path
            FROM grading_sessions
            WHERE id = ? AND COALESCE(is_deleted, 0) = 0
            """,
            (session_id,),
        ).fetchone()
        papers = connection.execute(
            """
            SELECT DISTINCT p.front_image, p.back_image
            FROM exam_papers AS p
            JOIN students AS s ON s.id = p.student_id
            JOIN session_results AS r
              ON r.session_id = p.session_id
             AND r.paper_id = p.id
             AND r.student_id = s.id
            WHERE p.session_id = ?
              AND p.match_status = 'matched'
              AND p.processing_status = 'graded'
              AND TRIM(COALESCE(s.student_code, '')) <> ''
              AND TRIM(COALESCE(s.name, '')) <> ''
            ORDER BY p.id
            LIMIT ?
            """,
            (session_id, paper_limit),
        ).fetchall()
    except sqlite3.Error as exc:
        raise AcceptanceError("authorized source could not be verified safely") from exc
    finally:
        connection.close()
    if session is None or len(papers) != paper_limit:
        raise AcceptanceError("authorized source no longer matches the prepared scope")

    source_paper = _resolve_source_file(session["source_paper_path"], data_root)
    template_candidates = sorted(
        (data_root / "templates" / f"session_{session_id}").glob("*.pdf")
    )
    if len(template_candidates) != 1:
        raise AcceptanceError("authorized source no longer matches the prepared scope")
    selected_files: list[tuple[str, Path]] = [
        (f"source-paper{source_paper.suffix.lower()}", source_paper),
        (
            "template.pdf",
            _resolve_source_file(str(template_candidates[0]), data_root),
        ),
    ]
    for index, paper in enumerate(papers, start=1):
        front = _resolve_source_file(paper["front_image"], data_root)
        back = _resolve_source_file(paper["back_image"], data_root)
        selected_files.extend(
            [
                (f"answer-sheets/answer-{index:02}-front{front.suffix.lower()}", front),
                (f"answer-sheets/answer-{index:02}-back{back.suffix.lower()}", back),
            ]
        )
    return _combined_fingerprint(
        _database_generation(database),
        {
            logical_name: _sha256_file(source_path)
            for logical_name, source_path in selected_files
        },
    )


def _acceptance_data_file(
    raw_path: object,
    *,
    data_root: Path,
    label: str,
    suffix: str | None = None,
) -> Path:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise AcceptanceError(f"{label} is missing")
    path = Path(raw_path).expanduser()
    if not path.is_absolute():
        path = data_root / path
    resolved = path.resolve()
    if not _is_relative_to(resolved, data_root) or not resolved.is_file():
        raise AcceptanceError(f"{label} is outside the acceptance workspace or missing")
    if suffix is not None and resolved.suffix.casefold() != suffix.casefold():
        raise AcceptanceError(f"{label} has an unexpected file type")
    return resolved


def _load_json_object(path: Path, *, label: str) -> dict[str, object]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError(f"{label} is not valid JSON") from exc
    if not isinstance(value, dict):
        raise AcceptanceError(f"{label} is not a JSON object")
    return value


def _config_summary(
    session: sqlite3.Row,
    *,
    data_root: Path,
) -> tuple[list[str], set[str], QuestionIdCatalog, str]:
    rubric_path = _acceptance_data_file(
        session["rubric_path"],
        data_root=data_root,
        label="rubric",
        suffix=".json",
    )
    answer_path = _acceptance_data_file(
        session["answer_key_path"],
        data_root=data_root,
        label="answer key",
        suffix=".json",
    )
    rubric = _load_json_object(rubric_path, label="rubric")
    answer = _load_json_object(answer_path, label="answer key")
    rubric_questions = rubric.get("questions")
    answer_questions = answer.get("questions")
    if not isinstance(rubric_questions, list) or not isinstance(
        answer_questions,
        list,
    ):
        raise AcceptanceError("configuration questions are missing")

    def question_ids(
        questions: list[object],
        *,
        label: str,
    ) -> list[str]:
        values: list[str] = []
        for question in questions:
            if not isinstance(question, dict):
                raise AcceptanceError(f"{label} contains an invalid question")
            question_id = str(question.get("question_id") or "").strip()
            if not question_id or question_id in values:
                raise AcceptanceError(f"{label} question identities are invalid")
            values.append(question_id)
        if not values:
            raise AcceptanceError(f"{label} questions are missing")
        return values

    rubric_ids = question_ids(rubric_questions, label="rubric")
    answer_ids = question_ids(answer_questions, label="answer key")
    if set(rubric_ids) != set(answer_ids):
        raise AcceptanceError("configuration question identities conflict")

    try:
        question_catalog = QuestionIdCatalog.from_document(rubric)
    except QuestionIdContractError as exc:
        raise AcceptanceError(
            "rubric question identities are invalid"
        ) from exc
    if (
        not question_catalog.parent_ids
        or set(question_catalog.parent_ids) != set(rubric_ids)
        or not question_catalog.detail_ids
    ):
        raise AcceptanceError("rubric question identities are invalid")
    scoring_ids = (
        set(question_catalog.parent_ids)
        | set(question_catalog.detail_ids)
        | set(question_catalog.aliases)
    )
    config_hash = _combined_fingerprint(
        None,
        {
            "rubric.json": _sha256_file(rubric_path),
            "answer.json": _sha256_file(answer_path),
        },
    )
    return rubric_ids, scoring_ids, question_catalog, config_hash


def _template_summary(
    database: Path,
    *,
    session_id: int,
    data_root: Path,
    question_catalog: QuestionIdCatalog,
) -> tuple[int, str]:
    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        templates = connection.execute(
            """
            SELECT front_template_path, back_template_path, ai_analysis_path,
                   template_config_path, regions_path, is_confirmed,
                   regions_snapshot_pending, regions_snapshot_token
            FROM session_templates
            WHERE session_id = ?
            """,
            (session_id,),
        ).fetchall()
        regions = connection.execute(
            """
            SELECT mapped_question_id, is_confirmed, mapping_status,
                   multi_region_confirmed
            FROM answer_regions
            WHERE session_id = ?
            ORDER BY id
            """,
            (session_id,),
        ).fetchall()
    except sqlite3.Error as exc:
        raise AcceptanceError(
            "acceptance template could not be verified safely"
        ) from exc
    finally:
        connection.close()
    if len(templates) != 1 or not bool(templates[0]["is_confirmed"]):
        raise AcceptanceError("acceptance template is missing or unconfirmed")
    template = templates[0]
    if bool(template["regions_snapshot_pending"]) or template[
        "regions_snapshot_token"
    ]:
        raise AcceptanceError("region snapshot is incomplete")

    template_files: dict[str, str] = {}
    for field, label in (
        ("front_template_path", "front template"),
        ("back_template_path", "back template"),
        ("ai_analysis_path", "template analysis"),
        ("template_config_path", "template configuration"),
        ("regions_path", "region snapshot"),
    ):
        path = _acceptance_data_file(
            template[field],
            data_root=data_root,
            label=label,
        )
        template_files[label] = _sha256_file(path)

    if not regions:
        raise AcceptanceError("confirmed answer regions are missing")
    covered_detail_ids: set[str] = set()
    mapped_regions: dict[str, list[bool]] = {}
    for region in regions:
        mapped_question_id = str(region["mapped_question_id"] or "").strip()
        if (
            not mapped_question_id
            or not bool(region["is_confirmed"])
            or str(region["mapping_status"] or "") == "unbound"
            ):
            raise AcceptanceError("answer regions contain an unconfirmed mapping")
        if mapped_question_id == "__student_name__":
            resolved_question_id = mapped_question_id
            expanded = ()
        else:
            resolved_question_id = question_catalog.resolve(mapped_question_id)
            expanded = (
                question_catalog.expand(resolved_question_id)
                if resolved_question_id is not None
                else ()
            )
        if resolved_question_id is None or (
            resolved_question_id != "__student_name__" and not expanded
        ):
            raise AcceptanceError("answer regions conflict with configuration")
        existing_mappings = mapped_regions.setdefault(
            resolved_question_id,
            [],
        )
        existing_mappings.append(bool(region["multi_region_confirmed"]))
        for detail_id in expanded:
            if detail_id in covered_detail_ids and not existing_mappings[:-1]:
                raise AcceptanceError("answer regions conflict with configuration")
            covered_detail_ids.add(detail_id)
    if any(
        len(confirmations) > 1 and not all(confirmations)
        for confirmations in mapped_regions.values()
    ):
        raise AcceptanceError("duplicate answer regions are not confirmed")
    if covered_detail_ids != set(question_catalog.detail_ids):
        raise AcceptanceError("answer regions are missing configured questions")
    return len(regions), _combined_fingerprint(None, template_files)


def _grading_summary(
    database: Path,
    *,
    session_id: int,
    data_root: Path,
    expected_paper_count: int,
    scoring_ids: set[str],
) -> None:
    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        papers = connection.execute(
            """
            SELECT p.id, p.student_id, p.front_image, p.back_image,
                   p.match_status, p.processing_status,
                   s.student_code, s.name
            FROM exam_papers AS p
            LEFT JOIN students AS s ON s.id = p.student_id
            WHERE p.session_id = ?
            ORDER BY p.id
            """,
            (session_id,),
        ).fetchall()
        results = connection.execute(
            """
            SELECT id, student_id, paper_id, total_score, student_score,
                   raw_json
            FROM session_results
            WHERE session_id = ?
            ORDER BY id
            """,
            (session_id,),
        ).fetchall()
        details = connection.execute(
            """
            SELECT d.result_id, d.question_id, d.score_awarded
            FROM session_details AS d
            JOIN session_results AS r ON r.id = d.result_id
            WHERE r.session_id = ?
            ORDER BY d.result_id, d.id
            """,
            (session_id,),
        ).fetchall()
    except sqlite3.Error as exc:
        raise AcceptanceError(
            "acceptance grading could not be verified safely"
        ) from exc
    finally:
        connection.close()

    if len(papers) != expected_paper_count or len(results) != expected_paper_count:
        raise AcceptanceError("grading is incomplete or conflicted")
    paper_students: dict[int, int] = {}
    for paper in papers:
        student_id = paper["student_id"]
        if (
            student_id is None
            or str(paper["match_status"]) != "matched"
            or str(paper["processing_status"]) != "graded"
            or not str(paper["student_code"] or "").strip()
            or not str(paper["name"] or "").strip()
        ):
            raise AcceptanceError("grading is incomplete or conflicted")
        paper_id = int(paper["id"])
        paper_students[paper_id] = int(student_id)
        _acceptance_data_file(
            paper["front_image"],
            data_root=data_root,
            label="graded paper front image",
        )
        _acceptance_data_file(
            paper["back_image"],
            data_root=data_root,
            label="graded paper back image",
        )

    details_by_result: dict[int, list[sqlite3.Row]] = {}
    for detail in details:
        details_by_result.setdefault(int(detail["result_id"]), []).append(detail)
    seen_papers: set[int] = set()
    seen_students: set[int] = set()
    for result in results:
        result_id = int(result["id"])
        paper_id = int(result["paper_id"])
        student_id = int(result["student_id"])
        if (
            paper_id in seen_papers
            or student_id in seen_students
            or paper_students.get(paper_id) != student_id
        ):
            raise AcceptanceError("grading is incomplete or conflicted")
        seen_papers.add(paper_id)
        seen_students.add(student_id)
        result_details = details_by_result.get(result_id, [])
        if not result_details:
            raise AcceptanceError("grading is incomplete or conflicted")
        seen_question_ids: set[str] = set()
        awarded_total = 0.0
        for detail in result_details:
            question_id = str(detail["question_id"] or "").strip()
            try:
                score = float(detail["score_awarded"])
            except (TypeError, ValueError) as exc:
                raise AcceptanceError(
                    "grading is incomplete or conflicted"
                ) from exc
            if (
                not question_id
                or question_id not in scoring_ids
                or question_id in seen_question_ids
                or not math.isfinite(score)
                or score < 0
            ):
                raise AcceptanceError("grading is incomplete or conflicted")
            seen_question_ids.add(question_id)
            awarded_total += score
        try:
            student_score = float(result["student_score"])
            total_score = float(result["total_score"])
        except (TypeError, ValueError) as exc:
            raise AcceptanceError("grading is incomplete or conflicted") from exc
        if (
            not math.isfinite(student_score)
            or not math.isfinite(total_score)
            or total_score <= 0
            or student_score < 0
            or student_score > total_score + 1e-6
            or abs(awarded_total - student_score) > 1e-6
        ):
            raise AcceptanceError("grading is incomplete or conflicted")
        try:
            raw_json = json.loads(str(result["raw_json"]))
        except json.JSONDecodeError as exc:
            raise AcceptanceError("grading is incomplete or conflicted") from exc
        completeness = (
            raw_json.get("grading_completeness")
            if isinstance(raw_json, dict)
            else None
        )
        if (
            not isinstance(completeness, dict)
            or completeness.get("status") != "complete"
            or any(
                completeness.get(key)
                for key in (
                    "missing_question_ids",
                    "duplicate_question_ids",
                    "unexpected_question_ids",
                )
            )
        ):
            raise AcceptanceError("grading is incomplete or conflicted")


def _pending_teacher_review_count(
    database: Path,
    *,
    session_id: int,
) -> int:
    try:
        from backend.review.service import ReviewApplicationService
        from db_manager import DBManager

        db = DBManager(database)
        session = db.get_grading_session(session_id)
        if not session:
            raise AcceptanceError("acceptance session is missing or deleted")
        questions = ReviewApplicationService(db).list_questions(
            session_id,
            session,
        )
    except AcceptanceError:
        raise
    except Exception as exc:
        raise AcceptanceError(
            "teacher review could not be verified safely"
        ) from exc
    return sum(int(question.needs_review_count) for question in questions)


def _report_summary(
    database: Path,
    *,
    session_id: int,
    data_root: Path,
    report_file: Path | str | None,
) -> tuple[int, str]:
    report_path = _acceptance_data_file(
        str(report_file) if report_file is not None else None,
        data_root=data_root,
        label="exported report",
        suffix=".xlsx",
    )
    reports_root = (data_root / "reports").resolve()
    if not _is_relative_to(report_path, reports_root):
        raise AcceptanceError("exported report is outside the reports directory")

    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        expected_rows = connection.execute(
            """
            SELECT s.student_code, r.student_score
            FROM session_results AS r
            JOIN students AS s ON s.id = r.student_id
            WHERE r.session_id = ?
            ORDER BY r.id
            """,
            (session_id,),
        ).fetchall()
    except sqlite3.Error as exc:
        raise AcceptanceError(
            "reviewed results could not be verified safely"
        ) from exc
    finally:
        connection.close()
    expected = {
        str(row["student_code"]): float(row["student_score"])
        for row in expected_rows
    }

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            report_path,
            read_only=True,
            data_only=True,
        )
        try:
            if "成绩与小题明细" not in workbook.sheetnames:
                raise AcceptanceError("exported report is missing the score sheet")
            rows = workbook["成绩与小题明细"].iter_rows(values_only=True)
            header_row = next(rows)
            headers = {
                str(value).strip(): index
                for index, value in enumerate(header_row)
                if value is not None and str(value).strip()
            }
            if "学号" not in headers or "总分" not in headers:
                raise AcceptanceError(
                    "exported report is missing required score columns"
                )
            actual: dict[str, float] = {}
            for row in rows:
                raw_code = row[headers["学号"]]
                raw_score = row[headers["总分"]]
                if raw_code is None and raw_score is None:
                    continue
                student_code = str(raw_code or "").strip()
                try:
                    score = float(raw_score)
                except (TypeError, ValueError) as exc:
                    raise AcceptanceError(
                        "report conflicts with reviewed results"
                    ) from exc
                if (
                    not student_code
                    or student_code in actual
                    or not math.isfinite(score)
                ):
                    raise AcceptanceError(
                        "report conflicts with reviewed results"
                    )
                actual[student_code] = score
        finally:
            workbook.close()
    except AcceptanceError:
        raise
    except (OSError, KeyError, StopIteration, ValueError) as exc:
        raise AcceptanceError(
            "exported report could not be verified safely"
        ) from exc
    if set(actual) != set(expected) or any(
        abs(actual[student_code] - score) > 1e-6
        for student_code, score in expected.items()
    ):
        raise AcceptanceError("report conflicts with reviewed results")
    return len(actual), _sha256_file(report_path)


def build_consistency_report(
    workspace: Path | str,
    *,
    stage: str,
    session_id: int,
    source_data_root: Path | str,
    report_file: Path | str | None = None,
) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "runtime_ready":
        raise AcceptanceError("acceptance runtime is not ready")
    clean_stage = str(stage or "").strip()
    stages = ("session", "config", "template", "grading", "review", "report")
    if clean_stage not in stages:
        raise AcceptanceError("unsupported consistency stage")
    if isinstance(session_id, bool) or not isinstance(session_id, int) or session_id < 1:
        raise AcceptanceError("acceptance session identity is invalid")

    authorized_session_id = int(metadata["authorized_session_id"])
    paper_limit = int(metadata["paper_count"])
    source_fingerprint = _current_authorized_source_fingerprint(
        source_data_root,
        session_id=authorized_session_id,
        paper_limit=paper_limit,
    )
    if source_fingerprint != metadata["source_fingerprint"]:
        raise AcceptanceError("authorized source changed after preparation")

    database = target / "acceptance_data" / "databases" / "grading_system.db"
    if not database.is_file():
        raise AcceptanceError("acceptance database is unavailable")
    connection = sqlite3.connect(
        f"file:{database.as_posix()}?mode=ro",
        uri=True,
    )
    connection.row_factory = sqlite3.Row
    try:
        session = connection.execute(
            """
            SELECT status, COALESCE(is_deleted, 0) AS is_deleted,
                   rubric_path, answer_key_path
            FROM grading_sessions
            WHERE id = ?
            """,
            (session_id,),
        ).fetchone()
        papers = int(
            connection.execute(
                "SELECT COUNT(*) FROM exam_papers WHERE session_id = ?",
                (session_id,),
            ).fetchone()[0]
        )
        results = int(
            connection.execute(
                "SELECT COUNT(*) FROM session_results WHERE session_id = ?",
                (session_id,),
            ).fetchone()[0]
        )
        details = int(
            connection.execute(
                """
                SELECT COUNT(*)
                FROM session_details AS d
                JOIN session_results AS r ON r.id = d.result_id
                WHERE r.session_id = ?
                """,
                (session_id,),
            ).fetchone()[0]
        )
    except sqlite3.Error as exc:
        raise AcceptanceError("acceptance session could not be verified safely") from exc
    finally:
        connection.close()
    if session is None or bool(session["is_deleted"]):
        raise AcceptanceError("acceptance session is missing or deleted")

    data_root = (target / "acceptance_data").resolve()
    report: dict[str, object] = {
        "package": PACKAGE,
        "stage": clean_stage,
        "ok": True,
        "source_unchanged": True,
        "session_id": session_id,
        "session_status": str(session["status"]),
        "counts": {
            "papers": papers,
            "results": results,
            "details": details,
        },
    }
    if stages.index(clean_stage) >= stages.index("config"):
        question_ids, scoring_ids, question_catalog, config_hash = _config_summary(
            session,
            data_root=data_root,
        )
        report.update(
            {
                "question_count": len(question_ids),
                "config_sha256": config_hash,
            }
        )
    if stages.index(clean_stage) >= stages.index("template"):
        region_count, template_hash = _template_summary(
            database,
            session_id=session_id,
            data_root=data_root,
            question_catalog=question_catalog,
        )
        report.update(
            {
                "region_count": region_count,
                "template_confirmed": True,
                "template_sha256": template_hash,
            }
        )
    if stages.index(clean_stage) >= stages.index("grading"):
        _grading_summary(
            database,
            session_id=session_id,
            data_root=data_root,
            expected_paper_count=paper_limit,
            scoring_ids=scoring_ids,
        )
        report["grading_complete"] = True
    if stages.index(clean_stage) >= stages.index("review"):
        pending_review_count = _pending_teacher_review_count(
            database,
            session_id=session_id,
        )
        if pending_review_count:
            raise AcceptanceError("teacher review is incomplete")
        report.update(
            {
                "teacher_review_complete": True,
                "pending_review_count": 0,
            }
        )
    if clean_stage == "report":
        report_row_count, report_hash = _report_summary(
            database,
            session_id=session_id,
            data_root=data_root,
            report_file=report_file,
        )
        report.update(
            {
                "report_matches_results": True,
                "report_row_count": report_row_count,
                "report_sha256": report_hash,
            }
        )
    evidence_dir = target / "acceptance_evidence"
    evidence_dir.mkdir(exist_ok=True)
    evidence_path = evidence_dir / f"consistency-{clean_stage}.json"
    temporary_path = evidence_path.with_suffix(".json.tmp")
    temporary_path.write_text(
        json.dumps(report, ensure_ascii=True, sort_keys=True, separators=(",", ":")),
        encoding="utf-8",
    )
    temporary_path.replace(evidence_path)
    return report


def _require_nonempty_string(profile: dict[str, object], key: str) -> str:
    value = profile.get(key)
    if not isinstance(value, str) or not value.strip():
        raise AcceptanceError("active API profile is not fully configured")
    return value.strip()


def _validate_upstream_url(value: str) -> str:
    parsed = urlparse(value)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise AcceptanceError("active API profile is not fully configured")
    return value.rstrip("/")


def _validate_proxy_base_url(value: str) -> str:
    parsed = urlparse(value)
    if (
        parsed.scheme != "http"
        or parsed.hostname != "127.0.0.1"
        or parsed.port is None
        or not parsed.path.rstrip("/").endswith("/acceptance-llm/v1")
    ):
        raise AcceptanceError("acceptance proxy must use a dedicated loopback URL")
    return value.rstrip("/")


def prepare_runtime_profile(
    workspace: Path | str,
    *,
    source_profile_path: Path | str,
    proxy_base_url: str,
    max_forwarded_requests: int | None,
    initial_forwarded_requests: int = 0,
    direct_diagnostic_requests: int = 0,
) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "inputs_prepared":
        raise AcceptanceError("authorized inputs must be prepared before runtime config")
    manual_unbounded = max_forwarded_requests is None
    if not manual_unbounded and (
        isinstance(max_forwarded_requests, bool)
        or not isinstance(max_forwarded_requests, int)
        or max_forwarded_requests < 1
        or max_forwarded_requests > 9
    ):
        raise AcceptanceError("model request budget must be between 1 and 9")
    if (
        isinstance(initial_forwarded_requests, bool)
        or not isinstance(initial_forwarded_requests, int)
        or initial_forwarded_requests < 0
        or isinstance(direct_diagnostic_requests, bool)
        or not isinstance(direct_diagnostic_requests, int)
        or direct_diagnostic_requests < 0
        or (
            not manual_unbounded
            and (initial_forwarded_requests or direct_diagnostic_requests)
        )
    ):
        raise AcceptanceError("initial model request counts are invalid")
    proxy_url = _validate_proxy_base_url(proxy_base_url)
    source_path = Path(source_profile_path).expanduser().resolve()
    if not source_path.is_file() or _is_relative_to(source_path, target):
        raise AcceptanceError("machine-local API profile is missing or unsafe")
    profile_fingerprint = _sha256_file(source_path)
    try:
        profiles = json.loads(source_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("machine-local API profile cannot be read") from exc
    if not isinstance(profiles, list) or not profiles or not isinstance(profiles[-1], dict):
        raise AcceptanceError("machine-local API profile is invalid")
    active: dict[str, object] = profiles[-1]
    provider = _require_nonempty_string(active, "provider")
    config_provider = str(active.get("config_provider") or provider).strip()
    if provider != "custom-openai-compatible" or config_provider != provider:
        raise AcceptanceError("P2-20 requires an OpenAI-compatible active profile")
    grading_base_url = _validate_upstream_url(
        _require_nonempty_string(active, "base_url")
    )
    grading_api_key = _require_nonempty_string(active, "api_key")
    grading_model = _require_nonempty_string(active, "grading_model")
    config_base_url = _validate_upstream_url(
        str(active.get("config_base_url") or grading_base_url)
    )
    config_api_key = str(active.get("config_api_key") or grading_api_key).strip()
    config_model = _require_nonempty_string(active, "config_model")
    if not config_api_key:
        raise AcceptanceError("active API profile is not fully configured")

    config_dir = target / "acceptance_config"
    if config_dir.exists():
        raise AcceptanceError("acceptance runtime config already exists")
    config_dir.mkdir()
    runtime_profile = {
        "name": "P2-20 Acceptance",
        "provider": provider,
        "base_url": proxy_url,
        "api_key": "acceptance-proxy",
        "grading_model": grading_model,
        "ocr_model": "p2-20-local-ocr",
        "config_provider": config_provider,
        "config_base_url": proxy_url,
        "config_api_key": "acceptance-proxy",
        "config_model": config_model,
        "objective_enabled": False,
        "grading_max_workers": 1,
        "precheck_max_workers": 1,
        "hybrid_inflight_workers": 1,
        "llm_grading_max_retries": 0,
        "llm_config_generation_max_retries": 0,
        "llm_recognition_max_retries": 0,
    }
    upstream = {
        "grading": {
            "base_url": grading_base_url,
            "api_key": grading_api_key,
            "model": grading_model,
        },
        "config": {
            "base_url": config_base_url,
            "api_key": config_api_key,
            "model": config_model,
        },
        "max_forwarded_requests": max_forwarded_requests,
    }
    if manual_unbounded:
        upstream.update(
            {
                "initial_forwarded_requests": initial_forwarded_requests,
                "direct_diagnostic_requests": direct_diagnostic_requests,
                "manual_request_control": {
                    "default_route_id": _DEFAULT_MANUAL_MODEL_ROUTE,
                    "routes": dict(_MANUAL_MODEL_ROUTES),
                },
            }
        )
    runtime_profile_path = config_dir / "api_profiles.json"
    upstream_path = config_dir / "upstream.json"
    runtime_profile_path.write_text(
        json.dumps([runtime_profile], ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    upstream_path.write_text(
        json.dumps(upstream, ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    runtime_profile_path.chmod(0o600)
    upstream_path.chmod(0o600)
    if _sha256_file(source_path) != profile_fingerprint:
        raise AcceptanceError("machine-local API profile changed while it was copied")
    runtime_metadata: dict[str, object] = {
        **metadata,
        "state": "runtime_ready",
        "profile_fingerprint": profile_fingerprint,
        "model_request_budget": max_forwarded_requests,
    }
    if manual_unbounded:
        runtime_metadata.update(
            {
                "model_request_count_baseline": initial_forwarded_requests,
                "direct_diagnostic_requests": direct_diagnostic_requests,
            }
        )
    _write_metadata(target, runtime_metadata)
    return load_metadata(target)


ModelForwarder = Callable[
    [dict[str, str], dict[str, object]],
    Awaitable[tuple[int, dict[str, str], bytes]],
]


def _load_upstream_config(path: Path) -> dict[str, object]:
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("acceptance upstream config is unavailable") from exc
    if not isinstance(config, dict):
        raise AcceptanceError("acceptance upstream config is invalid")
    maximum = config.get("max_forwarded_requests")
    manual_unbounded = maximum is None
    if manual_unbounded:
        initial = config.get("initial_forwarded_requests")
        direct = config.get("direct_diagnostic_requests")
        control = config.get("manual_request_control")
        if (
            isinstance(initial, bool)
            or not isinstance(initial, int)
            or initial < 0
            or isinstance(direct, bool)
            or not isinstance(direct, int)
            or direct < 0
            or not isinstance(control, dict)
            or set(control) != {"default_route_id", "routes"}
            or control.get("default_route_id") != _DEFAULT_MANUAL_MODEL_ROUTE
            or control.get("routes") != _MANUAL_MODEL_ROUTES
        ):
            raise AcceptanceError("acceptance upstream config is invalid")
    elif (
        isinstance(maximum, bool)
        or not isinstance(maximum, int)
        or not 1 <= maximum <= 9
    ):
        raise AcceptanceError("acceptance upstream config is invalid")
    for kind in ("grading", "config"):
        upstream = config.get(kind)
        if not isinstance(upstream, dict):
            raise AcceptanceError("acceptance upstream config is invalid")
        for key in ("base_url", "api_key", "model"):
            if not isinstance(upstream.get(key), str) or not upstream[key].strip():
                raise AcceptanceError("acceptance upstream config is invalid")
        _validate_upstream_url(str(upstream["base_url"]))
    return config


def _load_budget_state(state_path: Path, maximum: int) -> dict[str, int]:
    if not state_path.exists():
        return {
            "forwarded_requests": 0,
            "local_ocr_requests": 0,
            "max_forwarded_requests": maximum,
        }
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("model budget state is unavailable") from exc
    expected = {
        "forwarded_requests",
        "local_ocr_requests",
        "max_forwarded_requests",
    }
    if (
        not isinstance(state, dict)
        or set(state) != expected
        or any(
            isinstance(state[key], bool) or not isinstance(state[key], int)
            for key in expected
        )
        or state["forwarded_requests"] < 0
        or state["local_ocr_requests"] < 0
        or state["max_forwarded_requests"] != maximum
        or state["forwarded_requests"] > maximum
    ):
        raise AcceptanceError("model budget state is invalid")
    return state


def _write_budget_state(state_path: Path, state: dict[str, int]) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = state_path.with_suffix(f"{state_path.suffix}.tmp")
    temporary.write_text(
        json.dumps(state, ensure_ascii=True, separators=(",", ":")),
        encoding="utf-8",
    )
    temporary.replace(state_path)


def _write_manual_model_state(
    state_path: Path,
    state: dict[str, object],
) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = state_path.with_suffix(f"{state_path.suffix}.tmp")
    try:
        with temporary.open("w", encoding="utf-8") as destination:
            json.dump(
                state,
                destination,
                ensure_ascii=True,
                sort_keys=True,
                separators=(",", ":"),
            )
            destination.flush()
            os.fsync(destination.fileno())
        temporary.chmod(0o600)
        temporary.replace(state_path)
        state_path.chmod(0o600)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def _manual_prompt_texts(payload: dict[str, object]) -> list[str]:
    messages = payload.get("messages")
    if not isinstance(messages, list):
        return []
    texts: list[str] = []
    for message in messages:
        if not isinstance(message, dict):
            continue
        content = message.get("content")
        if isinstance(content, str):
            texts.append(content)
            continue
        if not isinstance(content, list):
            continue
        for item in content:
            if (
                isinstance(item, dict)
                and item.get("type") == "text"
                and isinstance(item.get("text"), str)
            ):
                texts.append(str(item["text"]))
    return texts


def _manual_model_action(payload: dict[str, object]) -> str | None:
    combined = "\n".join(_manual_prompt_texts(payload))
    batch_matches = _BATCH_QUESTION_IDS_RE.findall(combined)
    score_matches = _SCORE_QUESTION_IDS_RE.findall(combined)
    if len(batch_matches) + len(score_matches) != 1:
        return None
    raw_ids = batch_matches[0] if batch_matches else score_matches[0]
    try:
        question_ids = json.loads(raw_ids)
    except json.JSONDecodeError:
        return None
    if not isinstance(question_ids, list) or not all(
        isinstance(item, str) for item in question_ids
    ):
        return None
    if batch_matches:
        if question_ids == ["Q11"]:
            return "q11"
        if question_ids == ["Q12"]:
            return "q12"
        return None
    return (
        "score"
        if question_ids == [f"Q{number}" for number in range(1, 13)]
        else None
    )


def _manual_model_state_template(
    config: dict[str, object],
) -> dict[str, object]:
    initial = int(config["initial_forwarded_requests"])
    direct = int(config["direct_diagnostic_requests"])
    control = dict(config["manual_request_control"])
    return {
        "schema_version": 1,
        "mode": "manual_unbounded",
        "forwarded_requests": initial,
        "direct_diagnostic_requests": direct,
        "local_ocr_requests": 0,
        "selected_route_id": str(control["default_route_id"]),
        "unclassified_requests": initial,
        "outcome_counts": {
            "success": 0,
            "upstream_http_error": 0,
            "transport_error": 0,
            "unknown": 0,
        },
        "active_permit": None,
        "in_flight_request": None,
        "last_request": None,
        "recent_action": None,
    }


def _validate_manual_model_state(
    state: object,
    config: dict[str, object],
) -> dict[str, object]:
    expected_keys = {
        "schema_version",
        "mode",
        "forwarded_requests",
        "direct_diagnostic_requests",
        "local_ocr_requests",
        "selected_route_id",
        "unclassified_requests",
        "outcome_counts",
        "active_permit",
        "in_flight_request",
        "last_request",
        "recent_action",
    }
    control = dict(config["manual_request_control"])
    routes = dict(control["routes"])
    if not isinstance(state, dict) or set(state) != expected_keys:
        raise AcceptanceError("manual model request state is invalid")
    integer_keys = (
        "forwarded_requests",
        "direct_diagnostic_requests",
        "local_ocr_requests",
        "unclassified_requests",
    )
    if (
        state.get("schema_version") != 1
        or state.get("mode") != "manual_unbounded"
        or any(
            isinstance(state.get(key), bool)
            or not isinstance(state.get(key), int)
            or int(state[key]) < 0
            for key in integer_keys
        )
        or state.get("direct_diagnostic_requests")
        != config["direct_diagnostic_requests"]
        or state.get("unclassified_requests")
        != config["initial_forwarded_requests"]
        or state.get("selected_route_id") not in routes
    ):
        raise AcceptanceError("manual model request state is invalid")
    outcome_counts = state.get("outcome_counts")
    if (
        not isinstance(outcome_counts, dict)
        or set(outcome_counts) != _MANUAL_OUTCOME_KEYS
        or any(
            isinstance(value, bool) or not isinstance(value, int) or value < 0
            for value in outcome_counts.values()
        )
    ):
        raise AcceptanceError("manual model request state is invalid")
    for optional_mapping in (
        "active_permit",
        "in_flight_request",
        "last_request",
        "recent_action",
    ):
        if state[optional_mapping] is not None and not isinstance(
            state[optional_mapping], dict
        ):
            raise AcceptanceError("manual model request state is invalid")
    active_permit = state["active_permit"]
    if isinstance(active_permit, dict) and (
        set(active_permit)
        != {
            "action",
            "route_id",
            "token",
            "source_job_id",
            "created_at_unix_ms",
        }
        or active_permit.get("action") not in _MANUAL_MODEL_ACTIONS
        or active_permit.get("route_id") not in routes
        or not isinstance(active_permit.get("token"), str)
        or not _MANUAL_MODEL_TOKEN_RE.fullmatch(active_permit["token"])
        or isinstance(active_permit.get("source_job_id"), bool)
        or not isinstance(active_permit.get("source_job_id"), int)
        or int(active_permit["source_job_id"]) <= 0
        or isinstance(active_permit.get("created_at_unix_ms"), bool)
        or not isinstance(active_permit.get("created_at_unix_ms"), int)
        or int(active_permit["created_at_unix_ms"]) <= 0
    ):
        raise AcceptanceError("manual model request state is invalid")
    in_flight_request = state["in_flight_request"]
    if isinstance(in_flight_request, dict) and (
        set(in_flight_request)
        != {
            "number",
            "action",
            "route_id",
            "started_at_unix_ms",
        }
        or isinstance(in_flight_request.get("number"), bool)
        or not isinstance(in_flight_request.get("number"), int)
        or int(in_flight_request["number"]) != int(state["forwarded_requests"])
        or in_flight_request.get("action") not in _MANUAL_MODEL_ACTIONS
        or in_flight_request.get("route_id") not in routes
        or isinstance(in_flight_request.get("started_at_unix_ms"), bool)
        or not isinstance(in_flight_request.get("started_at_unix_ms"), int)
        or int(in_flight_request["started_at_unix_ms"]) <= 0
    ):
        raise AcceptanceError("manual model request state is invalid")
    last_request = state["last_request"]
    if isinstance(last_request, dict) and (
        set(last_request)
        != {
            "number",
            "action",
            "route_id",
            "outcome",
            "upstream_status_code",
            "transport_error_type",
            "completed_at_unix_ms",
        }
        or isinstance(last_request.get("number"), bool)
        or not isinstance(last_request.get("number"), int)
        or not 1 <= int(last_request["number"]) <= int(state["forwarded_requests"])
        or last_request.get("action") not in _MANUAL_MODEL_ACTIONS
        or last_request.get("route_id") not in routes
        or last_request.get("outcome") not in _MANUAL_OUTCOME_KEYS
        or (
            last_request.get("upstream_status_code") is not None
            and (
                isinstance(last_request.get("upstream_status_code"), bool)
                or not isinstance(last_request.get("upstream_status_code"), int)
                or not 100 <= int(last_request["upstream_status_code"]) <= 599
            )
        )
        or not isinstance(last_request.get("transport_error_type"), str)
        or not re.fullmatch(
            r"[A-Za-z_][A-Za-z0-9_.]{0,99}|",
            str(last_request["transport_error_type"]),
        )
        or isinstance(last_request.get("completed_at_unix_ms"), bool)
        or not isinstance(last_request.get("completed_at_unix_ms"), int)
        or int(last_request["completed_at_unix_ms"]) <= 0
    ):
        raise AcceptanceError("manual model request state is invalid")
    recent_action = state["recent_action"]
    if isinstance(recent_action, dict) and (
        set(recent_action)
        != {"action", "route_id", "token", "source_job_id", "job_id"}
        or recent_action.get("action") not in _MANUAL_MODEL_ACTIONS
        or recent_action.get("route_id") not in routes
        or not isinstance(recent_action.get("token"), str)
        or not _MANUAL_MODEL_TOKEN_RE.fullmatch(recent_action["token"])
        or any(
            isinstance(recent_action.get(key), bool)
            or not isinstance(recent_action.get(key), int)
            or int(recent_action[key]) <= 0
            for key in ("source_job_id", "job_id")
        )
    ):
        raise AcceptanceError("manual model request state is invalid")
    in_flight = 1 if state["in_flight_request"] is not None else 0
    classified = sum(int(value) for value in outcome_counts.values())
    if int(state["forwarded_requests"]) != (
        int(state["unclassified_requests"]) + classified + in_flight
    ):
        raise AcceptanceError("manual model request state is invalid")
    return copy.deepcopy(state)


class AcceptanceModelControl:
    def __init__(
        self,
        config: dict[str, object],
        state_path: Path,
    ) -> None:
        self._config = copy.deepcopy(config)
        self._state_path = state_path
        self._lock = threading.Lock()
        control = dict(config["manual_request_control"])
        self._routes = dict(control["routes"])
        if state_path.exists():
            try:
                raw_state = json.loads(state_path.read_text(encoding="utf-8"))
            except (OSError, UnicodeError, json.JSONDecodeError) as exc:
                raise AcceptanceError(
                    "manual model request state is unavailable"
                ) from exc
            self._state = _validate_manual_model_state(raw_state, config)
        else:
            self._state = _manual_model_state_template(config)
        recovered = False
        in_flight = self._state.get("in_flight_request")
        if isinstance(in_flight, dict):
            outcomes = dict(self._state["outcome_counts"])
            outcomes["unknown"] = int(outcomes["unknown"]) + 1
            self._state["outcome_counts"] = outcomes
            self._state["last_request"] = {
                "number": int(in_flight.get("number") or 0),
                "action": str(in_flight.get("action") or ""),
                "route_id": str(in_flight.get("route_id") or ""),
                "outcome": "unknown",
                "upstream_status_code": None,
                "transport_error_type": "",
                "completed_at_unix_ms": int(round(time.time() * 1000.0)),
            }
            self._state["in_flight_request"] = None
            recovered = True
        if self._state.get("active_permit") is not None:
            self._state["active_permit"] = None
            recovered = True
        _validate_manual_model_state(self._state, config)
        if recovered or not state_path.exists():
            _write_manual_model_state(state_path, self._state)

    def _persist(self) -> None:
        _validate_manual_model_state(self._state, self._config)
        _write_manual_model_state(self._state_path, self._state)

    def snapshot(self) -> dict[str, object]:
        with self._lock:
            return self.snapshot_unlocked()

    def issue_permit(
        self,
        *,
        action: str,
        route_id: str,
        token: str,
        source_job_id: int,
    ) -> dict[str, object]:
        clean_action = str(action).strip().lower()
        clean_route = str(route_id).strip()
        clean_token = str(token).strip().lower()
        clean_source_job_id = int(source_job_id)
        if (
            clean_action not in _MANUAL_MODEL_ACTIONS
            or clean_route not in self._routes
            or not _MANUAL_MODEL_TOKEN_RE.fullmatch(clean_token)
            or clean_source_job_id <= 0
        ):
            raise AcceptanceError("manual model request permit is invalid")
        with self._lock:
            existing = self._state.get("active_permit")
            permit = {
                "action": clean_action,
                "route_id": clean_route,
                "token": clean_token,
                "source_job_id": clean_source_job_id,
                "created_at_unix_ms": int(round(time.time() * 1000.0)),
            }
            if isinstance(existing, dict):
                comparable_keys = {
                    "action",
                    "route_id",
                    "token",
                    "source_job_id",
                }
                if all(existing.get(key) == permit[key] for key in comparable_keys):
                    return self.snapshot_unlocked()
                raise AcceptanceError("another manual model request is pending")
            if self._state.get("in_flight_request") is not None:
                raise AcceptanceError("a manual model request is already running")
            self._state["active_permit"] = permit
            self._state["selected_route_id"] = clean_route
            self._persist()
            return self.snapshot_unlocked()

    def revoke_pending_permit(self, token: str) -> bool:
        clean_token = str(token).strip().lower()
        with self._lock:
            permit = self._state.get("active_permit")
            if not isinstance(permit, dict) or permit.get("token") != clean_token:
                return False
            self._state["active_permit"] = None
            self._persist()
            return True

    def replay_action_job_id(
        self,
        *,
        action: str,
        route_id: str,
        token: str,
        source_job_id: int,
    ) -> int | None:
        clean = {
            "action": str(action).strip().lower(),
            "route_id": str(route_id).strip(),
            "token": str(token).strip().lower(),
            "source_job_id": int(source_job_id),
        }
        with self._lock:
            recent = self._state.get("recent_action")
            if not isinstance(recent, dict) or recent.get("token") != clean["token"]:
                return None
            if any(recent.get(key) != value for key, value in clean.items()):
                raise AcceptanceError(
                    "manual action token was already used for another request"
                )
            job_id = int(recent.get("job_id") or 0)
            if job_id <= 0:
                raise AcceptanceError("manual action history is invalid")
            return job_id

    def record_action_job(
        self,
        *,
        action: str,
        route_id: str,
        token: str,
        source_job_id: int,
        job_id: int,
    ) -> None:
        record = {
            "action": str(action).strip().lower(),
            "route_id": str(route_id).strip(),
            "token": str(token).strip().lower(),
            "source_job_id": int(source_job_id),
            "job_id": int(job_id),
        }
        if (
            record["action"] not in _MANUAL_MODEL_ACTIONS
            or record["route_id"] not in self._routes
            or not _MANUAL_MODEL_TOKEN_RE.fullmatch(str(record["token"]))
            or int(record["source_job_id"]) <= 0
            or int(record["job_id"]) <= 0
        ):
            raise AcceptanceError("manual action history is invalid")
        with self._lock:
            existing = self._state.get("recent_action")
            if isinstance(existing, dict) and existing.get("token") == record["token"]:
                if existing != record:
                    raise AcceptanceError(
                        "manual action token was already used for another request"
                    )
                return
            self._state["recent_action"] = record
            self._persist()

    def snapshot_unlocked(self) -> dict[str, object]:
        active = self._state.get("active_permit")
        recent = self._state.get("recent_action")
        return {
            "mode": "manual_unbounded",
            "forwarded_requests": int(self._state["forwarded_requests"]),
            "direct_diagnostic_requests": int(
                self._state["direct_diagnostic_requests"]
            ),
            "total_external_requests": int(self._state["forwarded_requests"])
            + int(self._state["direct_diagnostic_requests"]),
            "local_ocr_requests": int(self._state["local_ocr_requests"]),
            "max_forwarded_requests": None,
            "selected_route_id": str(self._state["selected_route_id"]),
            "routes": [
                {"id": route_id, "base_url": base_url}
                for route_id, base_url in sorted(self._routes.items())
            ],
            "outcome_counts": copy.deepcopy(self._state["outcome_counts"]),
            "unclassified_requests": int(
                self._state["unclassified_requests"]
            ),
            "active_permit": (
                {
                    "action": str(active.get("action") or ""),
                    "route_id": str(active.get("route_id") or ""),
                    "source_job_id": int(active.get("source_job_id") or 0),
                }
                if isinstance(active, dict)
                else None
            ),
            "in_flight_request": (
                copy.deepcopy(self._state["in_flight_request"])
                if isinstance(self._state.get("in_flight_request"), dict)
                else None
            ),
            "last_request": copy.deepcopy(self._state["last_request"]),
            "recent_action": (
                {
                    "action": str(recent.get("action") or ""),
                    "route_id": str(recent.get("route_id") or ""),
                    "source_job_id": int(recent.get("source_job_id") or 0),
                    "job_id": int(recent.get("job_id") or 0),
                }
                if isinstance(recent, dict)
                else None
            ),
        }

    def increment_local_ocr(self) -> None:
        with self._lock:
            self._state["local_ocr_requests"] = (
                int(self._state["local_ocr_requests"]) + 1
            )
            self._persist()

    def authorize(
        self,
        payload: dict[str, object],
        upstream: dict[str, str],
    ) -> tuple[dict[str, str], int]:
        action = _manual_model_action(payload)
        with self._lock:
            permit = self._state.get("active_permit")
            if (
                action is None
                or not isinstance(permit, dict)
                or permit.get("action") != action
            ):
                raise AcceptanceError(
                    "manual model request permission is required"
                )
            route_id = str(permit["route_id"])
            request_number = int(self._state["forwarded_requests"]) + 1
            self._state["forwarded_requests"] = request_number
            self._state["active_permit"] = None
            self._state["in_flight_request"] = {
                "number": request_number,
                "action": action,
                "route_id": route_id,
                "started_at_unix_ms": int(round(time.time() * 1000.0)),
            }
            self._persist()
            routed_upstream = dict(upstream)
            routed_upstream["base_url"] = self._routes[route_id]
            routed_upstream["_trust_env"] = "false"
            routed_upstream["_route_id"] = route_id
            return routed_upstream, request_number

    def finish(
        self,
        request_number: int,
        *,
        outcome: str,
        status_code: int | None,
        transport_error_type: str = "",
    ) -> None:
        if outcome not in _MANUAL_OUTCOME_KEYS - {"unknown"}:
            raise AcceptanceError("manual model request outcome is invalid")
        with self._lock:
            in_flight = self._state.get("in_flight_request")
            if (
                not isinstance(in_flight, dict)
                or int(in_flight.get("number") or 0) != int(request_number)
            ):
                raise AcceptanceError("manual model request state is invalid")
            outcomes = dict(self._state["outcome_counts"])
            outcomes[outcome] = int(outcomes[outcome]) + 1
            self._state["outcome_counts"] = outcomes
            self._state["last_request"] = {
                "number": int(request_number),
                "action": str(in_flight.get("action") or ""),
                "route_id": str(in_flight.get("route_id") or ""),
                "outcome": outcome,
                "upstream_status_code": status_code,
                "transport_error_type": str(transport_error_type or ""),
                "completed_at_unix_ms": int(round(time.time() * 1000.0)),
            }
            self._state["in_flight_request"] = None
            self._persist()


def mount_acceptance_model_control(
    app: Any,
    *,
    controller: AcceptanceModelControl,
    control_token: str | None,
) -> None:
    from fastapi import Header, HTTPException

    def authorize_control(value: str | None) -> None:
        if (
            control_token is None
            or value is None
            or not hmac.compare_digest(value, control_token)
        ):
            raise HTTPException(status_code=404)

    @app.get("/acceptance-control/model-state", include_in_schema=False)
    def acceptance_model_state(
        x_acceptance_control_token: str | None = Header(default=None),
    ):
        authorize_control(x_acceptance_control_token)
        return controller.snapshot()

    @app.post("/acceptance-control/model-permit", include_in_schema=False)
    def acceptance_model_permit(
        payload: dict[str, Any],
        x_acceptance_control_token: str | None = Header(default=None),
    ):
        authorize_control(x_acceptance_control_token)
        try:
            return controller.issue_permit(
                action=str(payload.get("action") or ""),
                route_id=str(payload.get("route_id") or ""),
                token=str(payload.get("token") or ""),
                source_job_id=int(payload.get("source_job_id") or 0),
            )
        except (AcceptanceError, TypeError, ValueError) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc


def _manual_continuation_job_state(
    manager: Any,
    session_id: int,
) -> dict[str, object]:
    active_jobs, _total = manager.list(
        session_id=int(session_id),
        job_types=("config_generation",),
        statuses=("queued", "running", "paused"),
        limit=1,
    )
    if active_jobs:
        job = active_jobs[0]
        return {
            "available_action": None,
            "source_job_id": int(job.id),
            "job": {
                "id": int(job.id),
                "status": str(job.status),
                "progress": float(job.progress),
                "stage": str(job.stage or ""),
                "detail": str(job.detail or ""),
            },
            "q11_status": "waiting",
            "q12_status": "waiting",
            "score_status": "waiting",
            "message": "当前请求仍在运行，请等待结果。",
        }
    jobs, _total = manager.list(
        session_id=int(session_id),
        job_types=("config_generation",),
        limit=1,
    )
    if not jobs:
        return {
            "available_action": None,
            "source_job_id": None,
            "job": None,
            "q11_status": "blocked",
            "q12_status": "blocked",
            "score_status": "blocked",
            "message": "没有找到可续接的配置生成任务。",
        }
    job = jobs[0]
    result = job.result if isinstance(job.result, dict) else {}
    failed_ids = {
        str(item)
        for item in result.get("failed_question_ids") or []
        if str(item).strip()
    }
    score_pending = bool(result.get("score_allocation_pending"))
    outcome = str(result.get("outcome") or "")
    if "Q11" in failed_ids:
        available_action = "q11"
        message = "Q11 尚未成功，可以再次人工请求。"
    elif "Q12" in failed_ids:
        available_action = "q12"
        message = "Q11 已保存，Q12 可以再次人工请求。"
    elif not failed_ids and score_pending:
        available_action = "score"
        message = "Q1～Q12 已保存，可以人工启动 AI 统一配分。"
    elif outcome == "complete" and job.status == "succeeded":
        available_action = None
        message = "Q11、Q12 和 AI 统一配分均已完成。"
    else:
        available_action = None
        message = "当前任务状态不能安全续接，请重新连接 Codex 核对。"
    q11_status = "failed" if "Q11" in failed_ids else "complete"
    q12_status = (
        "blocked"
        if "Q11" in failed_ids
        else ("failed" if "Q12" in failed_ids else "complete")
    )
    score_status = (
        "ready"
        if available_action == "score"
        else ("complete" if outcome == "complete" and job.status == "succeeded" else "blocked")
    )
    return {
        "available_action": available_action,
        "source_job_id": int(job.id),
        "job": {
            "id": int(job.id),
            "status": str(job.status),
            "progress": float(job.progress),
            "stage": str(job.stage or ""),
            "detail": str(job.detail or ""),
        },
        "q11_status": q11_status,
        "q12_status": q12_status,
        "score_status": score_status,
        "message": message,
    }


def _manual_continuation_html(control_token: str) -> str:
    template = """<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>P2-20 手动续接</title>
  <style>
    :root { color-scheme: light; font-family: "Microsoft YaHei", system-ui, sans-serif; }
    body { margin: 0; background: #f4f7fb; color: #1f2937; }
    main { width: min(920px, calc(100% - 32px)); margin: 32px auto; }
    .card { background: white; border: 1px solid #dbe3ef; border-radius: 14px;
      box-shadow: 0 8px 24px rgba(36, 56, 85, .08); padding: 24px; margin-bottom: 18px; }
    h1 { margin: 0 0 8px; font-size: 26px; }
    h2 { margin: 0 0 14px; font-size: 18px; }
    p { line-height: 1.65; }
    .notice { background: #fff8e6; border-left: 4px solid #d99800; padding: 12px 14px;
      border-radius: 8px; }
    .grid { display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 12px; }
    .stage { border: 1px solid #dbe3ef; border-radius: 10px; padding: 16px; }
    button, select { width: 100%; min-height: 44px; border-radius: 8px; font-size: 16px; }
    button { border: 0; background: #1769e0; color: white; font-weight: 700; cursor: pointer; }
    button:disabled { background: #aeb9c8; cursor: not-allowed; }
    select { border: 1px solid #aeb9c8; background: white; padding: 0 10px; }
    label { display: block; font-weight: 700; margin-bottom: 8px; }
    .counts { display: flex; gap: 18px; flex-wrap: wrap; font-variant-numeric: tabular-nums; }
    .status { min-height: 24px; font-weight: 700; }
    .error { color: #a61b1b; white-space: pre-wrap; }
    .ok { color: #176b3a; }
    .muted { color: #5d6b7e; font-size: 14px; }
    .confirm { display: flex; gap: 10px; align-items: flex-start; margin-top: 14px; }
    .confirm input { margin-top: 5px; }
    @media (max-width: 720px) { .grid { grid-template-columns: 1fr; } }
  </style>
</head>
<body>
<main>
  <section class="card">
    <h1>P2-20 手动续接</h1>
    <p>本页只操作隔离验收副本。模型请求<strong>不设总次数上限</strong>，但每次都必须由你点击确认；<strong>程序不会自动重试</strong>。</p>
    <div class="notice">请先完全退出代理软件和 TUN 模式，再勾选确认。每次断连也会计作一次，因为站点后台可能已经处理。</div>
  </section>
  <section class="card">
    <h2>下一次请求线路</h2>
    <label for="route">Base URL</label>
    <select id="route">
      <option value="ohmygpt">https://api.ohmygpt.com/v1</option>
      <option value="ohmycdn">https://apic1.ohmycdn.com/v1</option>
    </select>
    <label class="confirm"><input id="proxy-off" type="checkbox"><span>我已完全退出代理；我知道每次点击都可能产生一次模型费用。</span></label>
  </section>
  <section class="card">
    <h2>累计计数</h2>
    <div class="counts">
      <span>项目内：<strong id="project-count">—</strong></span>
      <span>独立对照：<strong id="direct-count">—</strong></span>
      <span>合计：<strong id="total-count">—</strong></span>
    </div>
    <p id="running" class="muted">正在读取本机状态……</p>
    <p><strong>最近一次：</strong><span id="last-result">暂无新请求结果。</span></p>
  </section>
  <section class="card">
    <h2>按顺序继续</h2>
    <div class="grid">
      <div class="stage"><p id="q11-status" class="status">等待状态</p><button id="q11" disabled>仅请求 Q11</button></div>
      <div class="stage"><p id="q12-status" class="status">等待状态</p><button id="q12" disabled>仅请求 Q12</button></div>
      <div class="stage"><p id="score-status" class="status">等待状态</p><button id="score" disabled>进行 AI 统一配分</button></div>
    </div>
    <p id="message" class="status"></p>
    <p id="error" class="error"></p>
  </section>
</main>
<script>
const controlToken = __CONTROL_TOKEN__;
const headers = {"Content-Type": "application/json", "X-Acceptance-Control-Token": controlToken};
const route = document.getElementById("route");
const proxyOff = document.getElementById("proxy-off");
const buttons = {q11: document.getElementById("q11"), q12: document.getElementById("q12"), score: document.getElementById("score")};
let currentState = null;
let selectedInitialized = false;
let busy = false;
function makeToken() {
  const bytes = new Uint8Array(16); crypto.getRandomValues(bytes);
  return Array.from(bytes, value => value.toString(16).padStart(2, "0")).join("");
}
function stageText(value) {
  return {complete: "已完成", failed: "上次失败，可重试", ready: "可以开始", blocked: "尚未解锁", waiting: "正在等待"}[value] || "等待状态";
}
function lastRequestText(last) {
  if (!last) return "暂无新请求结果。";
  const outcome = {
    success: "成功",
    upstream_http_error: `上游返回 HTTP ${last.upstream_status_code || "错误"}`,
    transport_error: `连接中断（${last.transport_error_type || "未知类型"}）`,
    unknown: "结果未知（程序中断前已计数）"
  }[last.outcome] || "状态未知";
  return `项目内第 ${last.number} 次，线路 ${last.route_id}，${outcome}`;
}
function render(state) {
  currentState = state;
  document.getElementById("project-count").textContent = state.forwarded_requests;
  document.getElementById("direct-count").textContent = state.direct_diagnostic_requests;
  document.getElementById("total-count").textContent = state.total_external_requests;
  document.getElementById("last-result").textContent = lastRequestText(state.last_request);
  if (!selectedInitialized) { route.value = state.selected_route_id; selectedInitialized = true; }
  document.getElementById("q11-status").textContent = stageText(state.q11_status);
  document.getElementById("q12-status").textContent = stageText(state.q12_status);
  document.getElementById("score-status").textContent = stageText(state.score_status);
  document.getElementById("message").textContent = state.message || "";
  const inflight = state.in_flight_request;
  const elapsed = inflight
    ? Math.max(0, Math.floor((Date.now() - Number(inflight.started_at_unix_ms)) / 1000))
    : 0;
  document.getElementById("running").textContent = inflight
    ? `项目内第 ${inflight.number} 次正在等待，已等待 ${elapsed} 秒，线路：${inflight.route_id}`
    : (state.job && ["queued", "running", "paused"].includes(state.job.status)
      ? `任务正在处理：${state.job.stage || "等待模型请求"}`
      : "当前没有正在发送的模型请求。");
  const enabled = !busy && proxyOff.checked && state.available_action;
  Object.entries(buttons).forEach(([action, button]) => {
    button.disabled = !(enabled && state.available_action === action);
  });
  route.disabled = busy || Boolean(state.job && ["queued", "running", "paused"].includes(state.job.status));
}
async function refresh() {
  try {
    const response = await fetch("/acceptance-control/manual-state", {headers});
    if (!response.ok) throw new Error(`读取状态失败（${response.status}）`);
    render(await response.json());
    document.getElementById("error").textContent = "";
  } catch (error) {
    document.getElementById("error").textContent = String(error.message || error);
  }
}
async function runAction(action) {
  if (!currentState || currentState.available_action !== action || busy || !proxyOff.checked) return;
  const nextNumber = Number(currentState.forwarded_requests) + 1;
  const selectedText = route.options[route.selectedIndex].text;
  if (!window.confirm(`即将进行项目内第 ${nextNumber} 次模型调用。\\n线路：${selectedText}\\n确认继续吗？`)) return;
  busy = true; render(currentState);
  try {
    const response = await fetch("/acceptance-control/manual-action", {
      method: "POST", headers,
      body: JSON.stringify({action, route_id: route.value, source_job_id: currentState.source_job_id, token: makeToken()})
    });
    if (!response.ok) {
      const body = await response.json().catch(() => ({}));
      throw new Error(body.detail || `提交失败（${response.status}）`);
    }
    await refresh();
  } catch (error) {
    document.getElementById("error").textContent = String(error.message || error);
  } finally {
    busy = false; if (currentState) render(currentState);
  }
}
Object.entries(buttons).forEach(([action, button]) => button.addEventListener("click", () => runAction(action)));
proxyOff.addEventListener("change", () => currentState && render(currentState));
refresh(); setInterval(refresh, 2000);
</script>
</body>
</html>"""
    return template.replace(
        "__CONTROL_TOKEN__",
        json.dumps(control_token, ensure_ascii=True),
    )


def mount_manual_continuation(
    app: Any,
    *,
    controller: AcceptanceModelControl,
    control_token: str | None,
    session_id: int,
    submit_action: Callable[..., object] | None = None,
) -> None:
    from fastapi import Header, HTTPException
    from fastapi.responses import HTMLResponse

    clean_session_id = int(session_id)
    action_response_cache: dict[str, object] = {}

    def authorize_control(value: str | None) -> None:
        if (
            control_token is None
            or value is None
            or not hmac.compare_digest(value, control_token)
        ):
            raise HTTPException(status_code=404)

    @app.get("/manual", include_in_schema=False, response_class=HTMLResponse)
    def manual_continuation_page():
        if control_token is None:
            raise HTTPException(status_code=404)
        return HTMLResponse(
            _manual_continuation_html(control_token),
            headers={
                "Cache-Control": "no-store",
                "Content-Security-Policy": (
                    "default-src 'self'; "
                    "script-src 'unsafe-inline'; "
                    "style-src 'unsafe-inline'; "
                    "connect-src 'self'; "
                    "img-src 'none'; "
                    "frame-ancestors 'none'; "
                    "base-uri 'none'; "
                    "form-action 'none'"
                ),
                "X-Content-Type-Options": "nosniff",
                "X-Frame-Options": "DENY",
            },
        )

    @app.get("/acceptance-control/manual-state", include_in_schema=False)
    def manual_continuation_state(
        x_acceptance_control_token: str | None = Header(default=None),
    ):
        authorize_control(x_acceptance_control_token)
        manager = getattr(app.state, "job_manager", None)
        if manager is None:
            raise HTTPException(status_code=503, detail="任务服务尚未就绪")
        return {
            **controller.snapshot(),
            **_manual_continuation_job_state(manager, clean_session_id),
        }

    @app.post("/acceptance-control/manual-action", include_in_schema=False)
    def manual_continuation_action(
        payload: dict[str, Any],
        x_acceptance_control_token: str | None = Header(default=None),
    ):
        authorize_control(x_acceptance_control_token)
        manager = getattr(app.state, "job_manager", None)
        if manager is None:
            raise HTTPException(status_code=503, detail="任务服务尚未就绪")
        try:
            action = str(payload.get("action") or "").strip().lower()
            route_id = str(payload.get("route_id") or "").strip()
            token = str(payload.get("token") or "").strip().lower()
            source_job_id = int(payload.get("source_job_id") or 0)
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=409, detail="人工请求参数无效") from exc
        try:
            replay_job_id = controller.replay_action_job_id(
                action=action,
                route_id=route_id,
                token=token,
                source_job_id=source_job_id,
            )
        except AcceptanceError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        if replay_job_id is not None:
            cached = action_response_cache.get(token)
            if cached is not None:
                return cached
            existing = manager.get(replay_job_id)
            if existing is None:
                raise HTTPException(status_code=409, detail="原任务记录已不可用")
            from backend.api.routers.jobs import _job_response

            return _job_response(existing)

        job_state = _manual_continuation_job_state(manager, clean_session_id)
        if (
            job_state.get("available_action") != action
            or job_state.get("source_job_id") != source_job_id
        ):
            raise HTTPException(
                status_code=409,
                detail="当前阶段不能执行这个人工请求，请刷新后重试。",
            )
        try:
            controller.issue_permit(
                action=action,
                route_id=route_id,
                token=token,
                source_job_id=source_job_id,
            )
            if submit_action is not None:
                response = submit_action(
                    manager=manager,
                    session_id=clean_session_id,
                    source_job_id=source_job_id,
                    action=action,
                    route_id=route_id,
                    token=token,
                )
            else:
                from backend.api.routers.config import (
                    retry_session_config_generation,
                )
                from backend.api.schemas.config import (
                    ConfigGenerationRetryRequest,
                )
                from db_manager import DBManager

                path_manager = getattr(app.state, "path_manager", None)
                if path_manager is None:
                    raise AcceptanceError("验收数据服务尚未就绪")
                retry_question_ids = (
                    ["Q11"]
                    if action == "q11"
                    else (["Q12"] if action == "q12" else None)
                )
                response = retry_session_config_generation(
                    clean_session_id,
                    ConfigGenerationRetryRequest(
                        source_job_id=source_job_id,
                        retry_question_ids=retry_question_ids,
                        client_request_token=token,
                    ),
                    DBManager(path_manager.db_path),
                    manager,
                )
            job_id = int(
                response.get("id")
                if isinstance(response, dict)
                else getattr(response, "id", 0)
            )
            controller.record_action_job(
                action=action,
                route_id=route_id,
                token=token,
                source_job_id=source_job_id,
                job_id=job_id,
            )
            action_response_cache[token] = response
            return response
        except BaseException:
            controller.revoke_pending_permit(token)
            raise


def _canonical_json_bytes(value: object) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _nonnegative_int(value: object) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        return 0
    return max(0, value)


def _model_response_diagnostic_fields(content: bytes) -> dict[str, object]:
    fields: dict[str, object] = {
        "response_bytes": len(content),
        "response_sha256": hashlib.sha256(content).hexdigest() if content else "",
        "finish_reason": "",
        "prompt_tokens": 0,
        "completion_tokens": 0,
        "total_tokens": 0,
    }
    if not content:
        return fields
    try:
        payload = json.loads(content)
    except (UnicodeError, json.JSONDecodeError):
        return fields
    if not isinstance(payload, dict):
        return fields
    choices = payload.get("choices")
    if isinstance(choices, list) and choices and isinstance(choices[0], dict):
        finish_reason = str(choices[0].get("finish_reason") or "")
        fields["finish_reason"] = (
            finish_reason
            if finish_reason in _SAFE_FINISH_REASONS
            else ("other" if finish_reason else "")
        )
    usage = payload.get("usage")
    if isinstance(usage, dict):
        fields["prompt_tokens"] = _nonnegative_int(usage.get("prompt_tokens"))
        fields["completion_tokens"] = _nonnegative_int(
            usage.get("completion_tokens")
        )
        fields["total_tokens"] = _nonnegative_int(usage.get("total_tokens"))
    return fields


def _prepare_model_proxy_diagnostic_log(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8"):
        pass
    path.chmod(0o600)


def _prepare_model_request_capture(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(f"{path.suffix}.tmp")
    if path.exists() or temporary.exists():
        raise AcceptanceError("model request capture destination is not empty")
    probe = path.with_suffix(f"{path.suffix}.probe")
    probe_created = False
    try:
        with probe.open("xb") as destination:
            probe_created = True
            destination.write(b"capture-ready")
            destination.flush()
            os.fsync(destination.fileno())
        probe.chmod(0o600)
    finally:
        if probe_created:
            probe.unlink(missing_ok=True)


def _write_model_request_capture(path: Path, payload: dict[str, object]) -> None:
    temporary = path.with_suffix(f"{path.suffix}.tmp")
    if path.exists() or temporary.exists():
        raise AcceptanceError("model request capture has already been written")
    payload_bytes = _canonical_json_bytes(payload)
    try:
        with temporary.open("xb") as destination:
            destination.write(payload_bytes)
            destination.flush()
            os.fsync(destination.fileno())
        temporary.chmod(0o600)
        temporary.replace(path)
        path.chmod(0o600)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def _append_model_proxy_diagnostic(
    path: Path | None,
    lock: threading.Lock,
    event: dict[str, object],
) -> None:
    if path is None:
        return
    serialized = json.dumps(
        event,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    )
    with lock:
        existing = path.read_bytes()
        temporary = path.with_suffix(f"{path.suffix}.tmp")
        try:
            with temporary.open("wb") as destination:
                destination.write(existing)
                destination.write(serialized.encode("utf-8"))
                destination.write(b"\n")
                destination.flush()
                os.fsync(destination.fileno())
            temporary.chmod(0o600)
            temporary.replace(path)
            path.chmod(0o600)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise


def _try_append_model_proxy_diagnostic(
    path: Path | None,
    lock: threading.Lock,
    event: dict[str, object],
) -> None:
    try:
        _append_model_proxy_diagnostic(path, lock, event)
    except OSError as exc:
        print(
            json.dumps(
                {
                    "schema_version": 1,
                    "event": "model_proxy_diagnostic_write_failed",
                    "error_type": type(exc).__name__,
                },
                ensure_ascii=True,
                sort_keys=True,
                separators=(",", ":"),
            ),
            file=sys.stderr,
            flush=True,
        )


def _model_proxy_diagnostic_event(
    *,
    payload: dict[str, object],
    request_number: int,
    maximum: int | None,
    started: float,
    outcome: str,
    status_code: int | None,
    transport_error_type: str = "",
    response_content: bytes = b"",
    route_id: str = "",
    upstream_host: str = "",
) -> dict[str, object]:
    payload_bytes = _canonical_json_bytes(payload)
    event = {
        "schema_version": 1,
        "event": "model_proxy_forward",
        "recorded_at_unix_ms": int(round(time.time() * 1000.0)),
        "outcome": outcome,
        "upstream_status_code": status_code,
        "transport_error_type": transport_error_type,
        "elapsed_ms": max(0, int(round((time.monotonic() - started) * 1000.0))),
        "model": str(payload.get("model") or ""),
        "request_number": request_number,
        "max_forwarded_requests": maximum,
        "payload_bytes": len(payload_bytes),
        "payload_sha256": hashlib.sha256(payload_bytes).hexdigest(),
        **_model_response_diagnostic_fields(response_content),
    }
    if route_id:
        event["route_id"] = str(route_id)
        event["upstream_host"] = str(upstream_host)
    return event


async def _default_model_forwarder(
    upstream: dict[str, str],
    payload: dict[str, object],
) -> tuple[int, dict[str, str], bytes]:
    import httpx

    url = f"{upstream['base_url'].rstrip('/')}/chat/completions"
    payload_bytes = _canonical_json_bytes(payload)
    try:
        async with httpx.AsyncClient(
            timeout=600.0,
            trust_env=upstream.get("_trust_env") != "false",
            follow_redirects=False,
        ) as client:
            response = await client.post(
                url,
                headers={
                    "Authorization": f"Bearer {upstream['api_key']}",
                    "Content-Type": "application/json",
                },
                content=payload_bytes,
            )
    except httpx.HTTPError as exc:
        raise AcceptanceError("approved model request failed at the provider") from exc
    return (
        response.status_code,
        {"content-type": response.headers.get("content-type", "application/json")},
        response.content,
    )


def create_model_budget_proxy(
    upstream_config_path: Path | str,
    *,
    state_path: Path | str,
    diagnostic_log_path: Path | str | None = None,
    request_capture_path: Path | str | None = None,
    forwarder: ModelForwarder | None = None,
):
    from fastapi import FastAPI, Header, HTTPException
    from fastapi.responses import Response

    upstream_path = Path(upstream_config_path).resolve()
    budget_path = Path(state_path).resolve()
    config = _load_upstream_config(upstream_path)
    raw_maximum = config["max_forwarded_requests"]
    maximum = int(raw_maximum) if isinstance(raw_maximum, int) else None
    if maximum is None and request_capture_path is not None:
        raise AcceptanceError(
            "manual mode does not allow model request capture"
        )
    manual_control = (
        AcceptanceModelControl(config, budget_path)
        if maximum is None
        else None
    )
    state = (
        _load_budget_state(budget_path, maximum)
        if maximum is not None
        else None
    )
    if state is not None:
        _write_budget_state(budget_path, state)
    state_lock = threading.Lock()
    diagnostic_lock = threading.Lock()
    diagnostic_path = (
        Path(diagnostic_log_path).resolve()
        if diagnostic_log_path is not None
        else None
    )
    if diagnostic_path is not None:
        _prepare_model_proxy_diagnostic_log(diagnostic_path)
    capture_path = (
        Path(request_capture_path).resolve()
        if request_capture_path is not None
        else None
    )
    capture_pending = capture_path is not None
    if capture_path is not None:
        _prepare_model_request_capture(capture_path)
    dispatch = forwarder or _default_model_forwarder
    app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)
    app.state.model_control = manual_control

    @app.post("/v1/chat/completions")
    async def chat_completions(
        payload: dict[str, Any],
        authorization: str | None = Header(default=None),
    ):
        if authorization != "Bearer acceptance-proxy":
            raise HTTPException(status_code=401, detail="acceptance proxy key required")
        model = payload.get("model")
        if model == "p2-20-local-ocr":
            if manual_control is not None:
                manual_control.increment_local_ocr()
            else:
                assert state is not None
                with state_lock:
                    state["local_ocr_requests"] += 1
                    _write_budget_state(budget_path, state)
            return {
                "id": "p2-20-local-ocr",
                "object": "chat.completion",
                "created": int(time.time()),
                "model": "p2-20-local-ocr",
                "choices": [
                    {
                        "index": 0,
                        "message": {
                            "role": "assistant",
                            "content": "NOT_FOUND",
                        },
                        "finish_reason": "stop",
                    }
                ],
            }
        candidates = [
            dict(config[kind])
            for kind in ("config", "grading")
            if isinstance(config[kind], dict) and config[kind].get("model") == model
        ]
        if not candidates:
            raise HTTPException(status_code=403, detail="model is outside P2-20 scope")
        nonlocal capture_pending
        selected_upstream = candidates[0]
        if manual_control is not None:
            try:
                selected_upstream, request_number = manual_control.authorize(
                    payload,
                    selected_upstream,
                )
            except AcceptanceError as exc:
                raise HTTPException(status_code=428, detail=str(exc)) from exc
            if capture_pending:
                assert capture_path is not None
                _write_model_request_capture(capture_path, payload)
                capture_pending = False
        else:
            assert state is not None
            assert maximum is not None
            with state_lock:
                if state["forwarded_requests"] >= maximum:
                    raise HTTPException(
                        status_code=429,
                        detail="P2-20 model request budget exhausted",
                    )
                if capture_pending:
                    assert capture_path is not None
                    _write_model_request_capture(capture_path, payload)
                    capture_pending = False
                state["forwarded_requests"] += 1
                request_number = state["forwarded_requests"]
                _write_budget_state(budget_path, state)
        started = time.monotonic()
        route_id = str(selected_upstream.get("_route_id") or "")
        upstream_host = (
            urlparse(str(selected_upstream.get("base_url") or "")).hostname
            or ""
        )
        try:
            status_code, headers, content = await dispatch(
                selected_upstream,
                payload,
            )
        except AcceptanceError as exc:
            cause = exc.__cause__
            if manual_control is not None:
                try:
                    manual_control.finish(
                        request_number,
                        outcome="transport_error",
                        status_code=None,
                        transport_error_type=(
                            type(cause).__name__
                            if cause
                            else type(exc).__name__
                        ),
                    )
                except OSError:
                    pass
            _try_append_model_proxy_diagnostic(
                diagnostic_path,
                diagnostic_lock,
                _model_proxy_diagnostic_event(
                    payload=payload,
                    request_number=request_number,
                    maximum=maximum,
                    started=started,
                    outcome="transport_error",
                    status_code=None,
                    transport_error_type=type(cause).__name__ if cause else type(exc).__name__,
                    route_id=route_id,
                    upstream_host=upstream_host,
                ),
            )
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        outcome = (
            "success"
            if 200 <= status_code < 300
            else "upstream_http_error"
        )
        if manual_control is not None:
            try:
                manual_control.finish(
                    request_number,
                    outcome=outcome,
                    status_code=status_code,
                )
            except OSError:
                pass
        _try_append_model_proxy_diagnostic(
            diagnostic_path,
            diagnostic_lock,
            _model_proxy_diagnostic_event(
                payload=payload,
                request_number=request_number,
                maximum=maximum,
                started=started,
                outcome=outcome,
                status_code=status_code,
                response_content=content,
                route_id=route_id,
                upstream_host=upstream_host,
            ),
        )
        return Response(
            content=content,
            status_code=status_code,
            media_type=headers.get("content-type", "application/json"),
        )

    return app


def _replace_config_value(line: str, value: str) -> str:
    if line.endswith("\r\n"):
        ending = "\r\n"
    elif line.endswith("\n"):
        ending = "\n"
    else:
        ending = ""
    return f"{value}{ending}"


def configure_staged_runtime(workspace: Path | str) -> None:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "runtime_ready":
        raise AcceptanceError("acceptance runtime profile has not been prepared")
    config_path = target / "config" / "app_config.yaml"
    try:
        lines = config_path.read_text(encoding="utf-8").splitlines(keepends=True)
    except (OSError, UnicodeError) as exc:
        raise AcceptanceError("staged app config is unavailable") from exc
    positions: dict[str, list[int]] = {"DATA_DIR": [], "LOGS_DIR": []}
    for index, line in enumerate(lines):
        match = _CONFIG_KEY_RE.match(line)
        if match:
            positions[match.group("key").strip("\"'")].append(index)
    if any(len(indexes) != 1 for indexes in positions.values()):
        raise AcceptanceError(
            "staged config must contain exactly one DATA_DIR and LOGS_DIR key"
        )
    lines[positions["DATA_DIR"][0]] = _replace_config_value(
        lines[positions["DATA_DIR"][0]],
        "DATA_DIR: acceptance_data",
    )
    lines[positions["LOGS_DIR"][0]] = _replace_config_value(
        lines[positions["LOGS_DIR"][0]],
        "LOGS_DIR: acceptance_logs",
    )
    temporary = config_path.with_suffix(".yaml.tmp")
    temporary.write_text("".join(lines), encoding="utf-8", newline="")
    temporary.replace(config_path)
    for relative in ("acceptance_data", "acceptance_logs", "acceptance_ops"):
        directory = target / relative
        directory.mkdir(exist_ok=False)


def _sanitize_build_log(text: str, *, workspace: Path, repo_root: Path) -> str:
    sanitized = text
    replacements = (
        (workspace, "<workspace>"),
        (repo_root, "<source-repo>"),
    )
    for path, placeholder in replacements:
        native = str(path.resolve())
        for variant in {native, native.replace("\\", "/"), native.replace("/", "\\")}:
            sanitized = re.sub(
                re.escape(variant),
                lambda _match, value=placeholder: value,
                sanitized,
                flags=re.IGNORECASE,
            )
    return sanitized


def build_and_copy_frontend(
    workspace: Path | str,
    *,
    repo_root: Path | str,
    expected_source_sha: str,
) -> None:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["source_sha"] != expected_source_sha:
        raise AcceptanceError("frontend build source does not match prepared code")
    source_root = Path(repo_root).resolve()
    resolved = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=source_root,
        capture_output=True,
        text=True,
        check=False,
    )
    if resolved.returncode != 0 or resolved.stdout.strip() != expected_source_sha:
        raise AcceptanceError("frontend build requires the exact source commit")
    status = subprocess.run(
        ["git", "status", "--porcelain=v1", "--untracked-files=no"],
        cwd=source_root,
        capture_output=True,
        text=True,
        check=False,
    )
    if status.returncode != 0 or status.stdout.strip():
        raise AcceptanceError("frontend build source must have no tracked changes")
    npm = shutil.which("npm.cmd") or shutil.which("npm")
    if not npm:
        raise AcceptanceError("npm is required to build the acceptance frontend")
    completed = subprocess.run(
        [npm, "run", "build"],
        cwd=source_root / "frontend",
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    log_path = target / "acceptance_logs" / "frontend-build.log"
    log_path.write_text(
        _sanitize_build_log(
            (completed.stdout or "") + (completed.stderr or ""),
            workspace=target,
            repo_root=source_root,
        ),
        encoding="utf-8",
    )
    if completed.returncode != 0:
        raise AcceptanceError("frontend build failed; see the sanitized workspace log")
    source_dist = source_root / "frontend" / "dist"
    if not (source_dist / "index.html").is_file() or not (
        source_dist / "assets"
    ).is_dir():
        raise AcceptanceError("frontend build did not produce a complete dist")
    destination = target / "frontend" / "dist"
    if destination.exists():
        raise AcceptanceError("acceptance frontend dist already exists")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(source_dist, destination)


def _validate_staged_config(workspace: Path) -> None:
    config_path = workspace / "config" / "app_config.yaml"
    try:
        lines = config_path.read_text(encoding="utf-8").splitlines()
    except (OSError, UnicodeError) as exc:
        raise AcceptanceError("prepared app config is unavailable") from exc
    values: dict[str, list[str]] = {"DATA_DIR": [], "LOGS_DIR": []}
    for line in lines:
        match = _CONFIG_KEY_RE.match(line)
        if match:
            key = match.group("key").strip("\"'")
            values[key].append(line.split(":", 1)[1].strip())
    if values != {
        "DATA_DIR": ["acceptance_data"],
        "LOGS_DIR": ["acceptance_logs"],
    }:
        raise AcceptanceError("prepared app config is not acceptance-safe")


def mount_acceptance_control_health(
    app: Any,
    *,
    control_token: str | None,
    source_sha: str,
) -> None:
    from fastapi import Header, HTTPException

    @app.get("/acceptance-control/health", include_in_schema=False)
    def acceptance_control_health(
        x_acceptance_control_token: str | None = Header(default=None),
    ):
        if (
            control_token is None
            or x_acceptance_control_token is None
            or not hmac.compare_digest(
                x_acceptance_control_token,
                control_token,
            )
        ):
            raise HTTPException(status_code=404)
        return {
            "package": PACKAGE,
            "source_sha": source_sha,
        }


def create_acceptance_app(
    workspace: Path | str,
    control_token: str | None = None,
):
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "runtime_ready":
        raise AcceptanceError("acceptance runtime is not ready")
    _validate_staged_config(target)
    frontend_dist = target / "frontend" / "dist"
    if not (frontend_dist / "index.html").is_file() or not (
        frontend_dist / "assets"
    ).is_dir():
        raise AcceptanceError("acceptance frontend dist is missing")
    launcher_root = Path(__file__).resolve().parents[1]
    if launcher_root != target:
        raise AcceptanceError("acceptance server must run from the staged workspace")

    os.environ["AI_GRADING_DATA_DIR"] = str(target / "acceptance_data")
    os.environ["AI_GRADING_API_PROFILES_PATH"] = str(
        target / "acceptance_config" / "api_profiles.json"
    )
    os.environ["AI_GRADING_OPS_STATE_DIR"] = str(target / "acceptance_ops")
    for name in (
        "LLM_API_KEY",
        "LLM_CONFIG_API_KEY",
        "LLM_OBJECTIVE_API_KEY",
        "OPENAI_API_KEY",
        "QUESTION_BANK_TAGGING_API_KEY",
        "QUESTION_BANK_TAGGING_REVIEW_API_KEY",
    ):
        os.environ[name] = ""

    from path_manager import PathManager

    paths = PathManager()
    if (
        paths.project_root.resolve() != target
        or paths.data_root.resolve() != (target / "acceptance_data").resolve()
        or paths.api_profiles_path.resolve()
        != (target / "acceptance_config" / "api_profiles.json").resolve()
        or paths.ops_state_dir.resolve() != (target / "acceptance_ops").resolve()
    ):
        raise AcceptanceError("acceptance runtime resolved outside the workspace")
    paths.ensure_directories()

    from db_manager import DBManager
    from question_bank.database.schema import initialize_database

    DBManager(paths.db_path).initialize()
    initialize_database(paths.qb_db_path)

    from backend.api.app import create_app
    from fastapi import HTTPException
    from fastapi.responses import FileResponse
    from fastapi.staticfiles import StaticFiles

    app = create_app(path_manager=paths)
    request_capture_path = (
        target / "acceptance_evidence" / "model-request-capture.json"
        if os.getenv(_CAPTURE_NEXT_MODEL_REQUEST_ENV) == "1"
        else None
    )
    proxy = create_model_budget_proxy(
        target / "acceptance_config" / "upstream.json",
        state_path=target / "acceptance_config" / "budget-state.json",
        diagnostic_log_path=target / "acceptance_logs" / "model-proxy.jsonl",
        request_capture_path=request_capture_path,
    )
    app.mount("/acceptance-llm", proxy, name="p2-20-model-budget")
    app.mount(
        "/assets",
        StaticFiles(directory=frontend_dist / "assets"),
        name="p2-20-assets",
    )

    mount_acceptance_control_health(
        app,
        control_token=control_token,
        source_sha=str(metadata["source_sha"]),
    )
    model_control = getattr(proxy.state, "model_control", None)
    if isinstance(model_control, AcceptanceModelControl):
        mount_manual_continuation(
            app,
            controller=model_control,
            control_token=control_token,
            session_id=int(metadata["authorized_session_id"]),
        )

    @app.get("/{frontend_path:path}", include_in_schema=False)
    def serve_frontend(frontend_path: str):
        if frontend_path.startswith(
            ("api/", "acceptance-llm/", "acceptance-control/")
        ):
            raise HTTPException(status_code=404)
        return FileResponse(frontend_dist / "index.html")

    return app


def _server_state_path(workspace: Path) -> Path:
    return workspace / "acceptance_config" / SERVER_STATE_FILENAME


def _server_claim_path(workspace: Path) -> Path:
    return workspace / "acceptance_config" / SERVER_CLAIM_FILENAME


def _server_stop_path(workspace: Path) -> Path:
    return workspace / "acceptance_config" / SERVER_STOP_FILENAME


def _validate_control_token(value: object) -> str:
    if not isinstance(value, str) or not 8 <= len(value) <= 256:
        raise AcceptanceError("acceptance server control token is invalid")
    return value


def _load_server_claim(workspace: Path) -> dict[str, object] | None:
    claim_path = _server_claim_path(workspace)
    if not claim_path.exists():
        return None
    try:
        claim = json.loads(claim_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("acceptance server start claim is unavailable") from exc
    expected = {
        "package",
        "source_sha",
        "port",
        "control_token",
        "owner_pid",
    }
    if (
        not isinstance(claim, dict)
        or set(claim) != expected
        or claim["package"] != PACKAGE
        or not isinstance(claim["source_sha"], str)
        or len(claim["source_sha"]) != 40
        or isinstance(claim["port"], bool)
        or not isinstance(claim["port"], int)
        or not 1024 <= claim["port"] <= 65535
        or isinstance(claim["owner_pid"], bool)
        or not isinstance(claim["owner_pid"], int)
        or claim["owner_pid"] <= 0
    ):
        raise AcceptanceError("acceptance server start claim is invalid")
    _validate_control_token(claim["control_token"])
    return claim


def _claim_server_start(
    workspace: Path,
    *,
    source_sha: str,
    port: int,
    control_token: str,
) -> None:
    _validate_control_token(control_token)
    claim = {
        "package": PACKAGE,
        "source_sha": source_sha,
        "port": port,
        "control_token": control_token,
        "owner_pid": os.getpid(),
    }
    claim_path = _server_claim_path(workspace)
    try:
        descriptor = os.open(
            claim_path,
            os.O_WRONLY | os.O_CREAT | os.O_EXCL,
            0o600,
        )
    except FileExistsError as exc:
        raise AcceptanceError(
            "acceptance server start is already claimed; stop it before restarting"
        ) from exc
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            json.dump(claim, stream, ensure_ascii=True, separators=(",", ":"))
    except BaseException:
        claim_path.unlink(missing_ok=True)
        raise


def _remove_server_claim_for_token(workspace: Path, control_token: str) -> None:
    try:
        claim = _load_server_claim(workspace)
    except AcceptanceError:
        return
    if claim is not None and hmac.compare_digest(
        str(claim["control_token"]),
        control_token,
    ):
        _server_claim_path(workspace).unlink(missing_ok=True)


def _load_server_state(workspace: Path) -> dict[str, object] | None:
    state_path = _server_state_path(workspace)
    if not state_path.exists():
        return None
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("acceptance server state is unavailable") from exc
    expected = {
        "package",
        "source_sha",
        "state",
        "pid",
        "port",
        "control_token",
    }
    if (
        not isinstance(state, dict)
        or set(state) != expected
        or state["package"] != PACKAGE
        or state["state"] not in {"starting", "running"}
        or isinstance(state["pid"], bool)
        or not isinstance(state["pid"], int)
        or state["pid"] <= 0
        or isinstance(state["port"], bool)
        or not isinstance(state["port"], int)
        or not 1024 <= state["port"] <= 65535
        or not isinstance(state["source_sha"], str)
        or len(state["source_sha"]) != 40
    ):
        raise AcceptanceError("acceptance server state is invalid")
    _validate_control_token(state["control_token"])
    return state


def _write_server_state(workspace: Path, state: dict[str, object]) -> None:
    state_path = _server_state_path(workspace)
    temporary = state_path.with_name(
        f".{state_path.name}.{os.getpid()}.{secrets.token_hex(8)}.tmp"
    )
    try:
        temporary.write_text(
            json.dumps(state, ensure_ascii=True, separators=(",", ":")),
            encoding="utf-8",
        )
        temporary.replace(state_path)
    finally:
        temporary.unlink(missing_ok=True)


def _remove_server_state_for_pid(workspace: Path, pid: int) -> None:
    try:
        state = _load_server_state(workspace)
    except AcceptanceError:
        return
    if state is not None and state["pid"] == pid:
        _server_state_path(workspace).unlink(missing_ok=True)


def _server_is_healthy(
    port: int,
    *,
    control_token: str,
    source_sha: str,
) -> bool:
    try:
        request = Request(
            f"http://127.0.0.1:{port}/acceptance-control/health",
            headers={"X-Acceptance-Control-Token": control_token},
        )
        with urlopen(request, timeout=0.5) as response:
            if response.status != 200:
                return False
            payload = json.loads(response.read().decode("utf-8"))
        return payload == {
            "package": PACKAGE,
            "source_sha": source_sha,
        }
    except (OSError, URLError, UnicodeError, ValueError, json.JSONDecodeError):
        return False


def _process_is_running(pid: int) -> bool:
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        process_query_limited_information = 0x1000
        still_active = 259
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.OpenProcess.argtypes = [
            wintypes.DWORD,
            wintypes.BOOL,
            wintypes.DWORD,
        ]
        kernel32.OpenProcess.restype = wintypes.HANDLE
        kernel32.GetExitCodeProcess.argtypes = [
            wintypes.HANDLE,
            ctypes.POINTER(wintypes.DWORD),
        ]
        kernel32.GetExitCodeProcess.restype = wintypes.BOOL
        kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel32.CloseHandle.restype = wintypes.BOOL
        handle = kernel32.OpenProcess(
            process_query_limited_information,
            False,
            pid,
        )
        if not handle:
            return False
        try:
            exit_code = wintypes.DWORD()
            if not kernel32.GetExitCodeProcess(
                handle,
                ctypes.byref(exit_code),
            ):
                return False
            return exit_code.value == still_active
        finally:
            kernel32.CloseHandle(handle)
    try:
        os.kill(pid, 0)
    except (OSError, ProcessLookupError):
        return False
    return True


def _ensure_loopback_port_available(port: int) -> None:
    probe = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        if sys.platform == "win32":
            probe.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        probe.bind(("127.0.0.1", port))
    except OSError as exc:
        raise AcceptanceError("acceptance server port is already in use") from exc
    finally:
        probe.close()


def _validate_prepared_proxy_port(workspace: Path, port: int) -> None:
    try:
        profiles = json.loads(
            (
                workspace / "acceptance_config" / "api_profiles.json"
            ).read_text(encoding="utf-8")
        )
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AcceptanceError("prepared API profile is unavailable") from exc
    expected_proxy = f"http://127.0.0.1:{port}/acceptance-llm/v1"
    if (
        not isinstance(profiles, list)
        or len(profiles) != 1
        or profiles[0].get("base_url") != expected_proxy
    ):
        raise AcceptanceError("serve port does not match the prepared model proxy")


def start_acceptance_server(
    workspace: Path | str,
    port: int,
    *,
    ready_timeout: float = 30.0,
) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    if metadata["state"] != "runtime_ready":
        raise AcceptanceError("acceptance runtime is not ready")
    validated_port = _validate_port(port)
    _validate_prepared_proxy_port(target, validated_port)
    existing = _load_server_state(target)
    claim = _load_server_claim(target)
    if existing is not None:
        ownership_matches = (
            claim is not None
            and existing["source_sha"] == metadata["source_sha"]
            and claim["source_sha"] == existing["source_sha"]
            and existing["port"] == validated_port
            and claim["port"] == existing["port"]
            and hmac.compare_digest(
                str(claim["control_token"]),
                str(existing["control_token"]),
            )
            and _process_is_running(int(existing["pid"]))
        )
        if ownership_matches and existing["state"] == "starting":
            raise AcceptanceError(
                "acceptance server start is still in progress"
            )
        if (
            ownership_matches
            and existing["state"] == "running"
            and _server_is_healthy(
                validated_port,
                control_token=str(existing["control_token"]),
                source_sha=str(existing["source_sha"]),
            )
        ):
            result = safe_status(target)
            result.update({"server_state": "running", "port": validated_port})
            return result
        raise AcceptanceError(
            "acceptance server state is stale; stop it before restarting"
        )
    if claim is not None:
        raise AcceptanceError(
            "acceptance server start is already claimed; stop it before restarting"
        )
    control_token = secrets.token_urlsafe(32)
    _claim_server_start(
        target,
        source_sha=str(metadata["source_sha"]),
        port=validated_port,
        control_token=control_token,
    )
    process: subprocess.Popen[bytes] | None = None
    try:
        _ensure_loopback_port_available(validated_port)
        _server_stop_path(target).unlink(missing_ok=True)
        stdout_path = target / "acceptance_logs" / "server-stdout.log"
        stderr_path = target / "acceptance_logs" / "server-stderr.log"
        creationflags = 0
        if sys.platform == "win32":
            creationflags = (
                getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
                | getattr(subprocess, "CREATE_NO_WINDOW", 0)
                | getattr(subprocess, "DETACHED_PROCESS", 0)
            )
        command = [
            sys.executable,
            str(target / "tools" / "p2_20_acceptance.py"),
            "serve",
            "--workspace",
            str(target),
            "--port",
            str(validated_port),
            "--control-token",
            control_token,
        ]
        with stdout_path.open("ab") as stdout, stderr_path.open("ab") as stderr:
            process = subprocess.Popen(
                command,
                cwd=target,
                stdin=subprocess.DEVNULL,
                stdout=stdout,
                stderr=stderr,
                close_fds=True,
                creationflags=creationflags,
            )
    except BaseException:
        _remove_server_claim_for_token(target, control_token)
        raise
    assert process is not None
    startup_failure = "acceptance server did not become healthy in time"
    deadline = time.monotonic() + max(1.0, float(ready_timeout))
    while time.monotonic() < deadline:
        if process.poll() is not None:
            _remove_server_state_for_pid(target, process.pid)
            _remove_server_claim_for_token(target, control_token)
            raise AcceptanceError(
                "acceptance server exited during startup; see workspace logs"
            )
        if _server_is_healthy(
            validated_port,
            control_token=control_token,
            source_sha=str(metadata["source_sha"]),
        ):
            registered = _load_server_state(target)
            if (
                registered is None
                or registered["pid"] != process.pid
                or not hmac.compare_digest(
                    str(registered["control_token"]),
                    control_token,
                )
            ):
                break
            registered["state"] = "running"
            try:
                _write_server_state(target, registered)
            except OSError:
                startup_failure = "acceptance server running state could not be saved"
                break
            result = safe_status(target)
            if result["server_state"] == "running":
                result.update({"port": validated_port})
                return result
            startup_failure = "acceptance server lost health during startup"
            break
        time.sleep(0.2)
    _server_stop_path(target).touch(exist_ok=True)
    try:
        process.wait(timeout=5)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)
    _remove_server_state_for_pid(target, process.pid)
    _remove_server_claim_for_token(target, control_token)
    raise AcceptanceError(startup_failure)


def stop_acceptance_server(
    workspace: Path | str,
    *,
    stop_timeout: float = 20.0,
) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    state = _load_server_state(target)
    if state is None:
        claim = _load_server_claim(target)
        if claim is not None:
            if _process_is_running(int(claim["owner_pid"])):
                raise AcceptanceError(
                    "acceptance server start is still in progress; retry stop shortly"
                )
            _remove_server_claim_for_token(
                target,
                str(claim["control_token"]),
            )
        result = safe_status(target)
        result.update({"server_state": "stopped"})
        return result
    if state["source_sha"] != metadata["source_sha"]:
        raise AcceptanceError("acceptance server source does not match the workspace")
    claim = _load_server_claim(target)
    if (
        claim is None
        or claim["source_sha"] != state["source_sha"]
        or claim["port"] != state["port"]
        or not hmac.compare_digest(
            str(claim["control_token"]),
            str(state["control_token"]),
        )
    ):
        raise AcceptanceError("acceptance server ownership is inconsistent")
    port = int(state["port"])
    _server_stop_path(target).touch(exist_ok=True)
    deadline = time.monotonic() + max(1.0, float(stop_timeout))
    while time.monotonic() < deadline:
        if not _process_is_running(int(state["pid"])):
            _server_state_path(target).unlink(missing_ok=True)
            _server_stop_path(target).unlink(missing_ok=True)
            _remove_server_claim_for_token(
                target,
                str(state["control_token"]),
            )
            result = safe_status(target)
            result.update({"server_state": "stopped", "port": port})
            return result
        time.sleep(0.2)
    raise AcceptanceError("acceptance server did not stop in time")


def run_acceptance_server(
    workspace: Path | str,
    port: int,
    control_token: str,
) -> None:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    validated_port = _validate_port(port)
    _validate_control_token(control_token)
    claim = _load_server_claim(target)
    if (
        claim is None
        or claim["source_sha"] != metadata["source_sha"]
        or claim["port"] != validated_port
        or not hmac.compare_digest(
            str(claim["control_token"]),
            control_token,
        )
    ):
        raise AcceptanceError("acceptance server start claim does not match")
    state = {
        "package": PACKAGE,
        "source_sha": metadata["source_sha"],
        "state": "starting",
        "pid": os.getpid(),
        "port": validated_port,
        "control_token": control_token,
    }
    try:
        _write_server_state(target, state)
        app = create_acceptance_app(target, control_token)
        import uvicorn

        server = uvicorn.Server(
            uvicorn.Config(app, host="127.0.0.1", port=validated_port)
        )
        stop_path = _server_stop_path(target)

        def watch_for_stop() -> None:
            while not server.should_exit:
                if stop_path.exists():
                    server.should_exit = True
                    return
                time.sleep(0.2)

        watcher = threading.Thread(target=watch_for_stop, daemon=True)
        watcher.start()
        server.run()
    finally:
        _server_stop_path(target).unlink(missing_ok=True)
        _remove_server_state_for_pid(target, os.getpid())
        _remove_server_claim_for_token(target, control_token)


def safe_status(workspace: Path | str) -> dict[str, object]:
    target = validate_workspace(workspace, require_empty=False)
    metadata = load_metadata(target)
    result = {
        key: metadata[key]
        for key in (
            "package",
            "state",
            "source_sha",
            "authorized_session_id",
            "paper_count",
            "model_request_budget",
        )
        if key in metadata
    }
    budget_path = target / "acceptance_config" / "budget-state.json"
    if "model_request_budget" in metadata:
        maximum = metadata["model_request_budget"]
        if maximum is None:
            config = _load_upstream_config(
                target / "acceptance_config" / "upstream.json"
            )
            if budget_path.exists():
                try:
                    raw_state = json.loads(
                        budget_path.read_text(encoding="utf-8")
                    )
                except (OSError, UnicodeError, json.JSONDecodeError) as exc:
                    raise AcceptanceError(
                        "manual model request state is unavailable"
                    ) from exc
                manual_state = _validate_manual_model_state(
                    raw_state,
                    config,
                )
            else:
                manual_state = _manual_model_state_template(config)
            result.update(
                {
                    "forwarded_requests": int(
                        manual_state["forwarded_requests"]
                    ),
                    "direct_diagnostic_requests": int(
                        manual_state["direct_diagnostic_requests"]
                    ),
                    "total_external_requests": int(
                        manual_state["forwarded_requests"]
                    )
                    + int(manual_state["direct_diagnostic_requests"]),
                    "local_ocr_requests": int(
                        manual_state["local_ocr_requests"]
                    ),
                }
            )
        else:
            state = _load_budget_state(
                budget_path,
                int(maximum),
            )
            result.update(
                {
                    "forwarded_requests": state["forwarded_requests"],
                    "local_ocr_requests": state["local_ocr_requests"],
                }
            )
    server = _load_server_state(target)
    if server is None:
        result["server_state"] = "stopped"
    else:
        port = int(server["port"])
        process_running = _process_is_running(int(server["pid"]))
        if not process_running:
            visible_state = "unhealthy"
        elif server["state"] == "starting":
            visible_state = "starting"
        elif _server_is_healthy(
            port,
            control_token=str(server["control_token"]),
            source_sha=str(server["source_sha"]),
        ):
            visible_state = "running"
        else:
            visible_state = "unhealthy"
        result.update(
            {
                "server_state": visible_state,
                "port": port,
            }
        )
        if metadata.get("model_request_budget") is None:
            result["manual_url"] = f"http://127.0.0.1:{port}/manual"
    return result


def _validate_port(port: int) -> int:
    if isinstance(port, bool) or not isinstance(port, int) or not 1024 <= port <= 65535:
        raise AcceptanceError("port must be between 1024 and 65535")
    return port


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="P2-20 controlled acceptance launcher")
    subparsers = parser.add_subparsers(dest="command", required=True)
    prepare = subparsers.add_parser("prepare")
    prepare.add_argument("--source-ref", required=True)
    prepare.add_argument("--workspace", type=Path, required=True)
    prepare.add_argument("--source-data-root", type=Path, required=True)
    prepare.add_argument("--session-id", type=int, required=True)
    prepare.add_argument("--session-name", required=True)
    prepare.add_argument("--profile-path", type=Path, required=True)
    prepare.add_argument("--port", type=int, required=True)
    request_mode = prepare.add_mutually_exclusive_group()
    request_mode.add_argument(
        "--model-request-budget",
        type=int,
        choices=range(1, 10),
    )
    request_mode.add_argument(
        "--unlimited-model-requests",
        action="store_true",
    )
    prepare.add_argument(
        "--initial-model-request-count",
        type=int,
        default=0,
    )
    prepare.add_argument(
        "--direct-diagnostic-request-count",
        type=int,
        default=0,
    )
    serve = subparsers.add_parser("serve")
    serve.add_argument("--workspace", type=Path, required=True)
    serve.add_argument("--port", type=int, required=True)
    serve.add_argument("--control-token", required=True)
    start = subparsers.add_parser("start")
    start.add_argument("--workspace", type=Path, required=True)
    start.add_argument("--port", type=int, required=True)
    stop = subparsers.add_parser("stop")
    stop.add_argument("--workspace", type=Path, required=True)
    status = subparsers.add_parser("status")
    status.add_argument("--workspace", type=Path, required=True)
    report = subparsers.add_parser("report")
    report.add_argument("--workspace", type=Path, required=True)
    report.add_argument(
        "--stage",
        choices=("session", "config", "template", "grading", "review", "report"),
        required=True,
    )
    report.add_argument("--session-id", type=int, required=True)
    report.add_argument("--source-data-root", type=Path, required=True)
    report.add_argument("--report-file", type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        if args.command == "prepare":
            port = _validate_port(args.port)
            code_metadata = prepare_code_workspace(
                args.source_ref,
                args.workspace,
                repo_root=REPO_ROOT,
            )
            prepare_authorized_inputs(
                args.workspace,
                source_data_root=args.source_data_root,
                session_id=args.session_id,
                expected_session_name=args.session_name,
                paper_limit=3,
            )
            prepare_runtime_profile(
                args.workspace,
                source_profile_path=args.profile_path,
                proxy_base_url=(
                    f"http://127.0.0.1:{port}/acceptance-llm/v1"
                ),
                max_forwarded_requests=(
                    None
                    if args.unlimited_model_requests
                    else (args.model_request_budget or 4)
                ),
                initial_forwarded_requests=args.initial_model_request_count,
                direct_diagnostic_requests=args.direct_diagnostic_request_count,
            )
            configure_staged_runtime(args.workspace)
            build_and_copy_frontend(
                args.workspace,
                repo_root=REPO_ROOT,
                expected_source_sha=str(code_metadata["source_sha"]),
            )
            result = safe_status(args.workspace)
        elif args.command == "status":
            result = safe_status(args.workspace)
        elif args.command == "start":
            result = start_acceptance_server(args.workspace, args.port)
        elif args.command == "stop":
            result = stop_acceptance_server(args.workspace)
        elif args.command == "report":
            result = build_consistency_report(
                args.workspace,
                stage=args.stage,
                session_id=args.session_id,
                source_data_root=args.source_data_root,
                report_file=args.report_file,
            )
        elif args.command == "serve":
            port = _validate_port(args.port)
            target = validate_workspace(args.workspace, require_empty=False)
            _validate_prepared_proxy_port(target, port)
            run_acceptance_server(target, port, args.control_token)
            return 0
        else:
            raise AcceptanceError("unsupported acceptance command")
        print(json.dumps(result, ensure_ascii=True, separators=(",", ":")))
        return 0
    except KeyboardInterrupt:
        return 130
    except (AcceptanceError, OSError, KeyError, ValueError) as exc:
        print(
            json.dumps(
                {"ok": False, "error": str(exc)},
                ensure_ascii=True,
                separators=(",", ":"),
            ),
            file=sys.stderr,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
