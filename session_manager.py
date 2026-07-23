from __future__ import annotations

import ast
import copy
import io
import json
import os
import re
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator, Mapping, Sequence
from xml.etree import ElementTree

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from answer_normalizer import complete_answer_set_values
from equivalence_engine import merge_equivalent_forms
from llm_client import LLMClient
from question_bank.services.ai_tagging_service import KNOWLEDGE_POINT_OPTIONS
from question_bank.models.skill_catalog import (
    SkillResolutionRequest,
    SkillRole,
)
from score_policy import (
    enforce_integer_scores_by_type,
    MAX_QUESTION_SCORE,
    OBJECTIVE_TYPES,
    _normalize_type,
)


DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS = 600.0
DEFAULT_CONFIG_GENERATION_RETRY_DELAYS = (2.0, 6.0)
DEFAULT_CONFIG_GENERATION_BATCH_SIZE = 3


class _QuestionGenerationRequestError(RuntimeError):
    def __init__(self, question_id: str, attempts: int, category: str, original: Exception) -> None:
        super().__init__(str(original))
        self.question_id = question_id
        self.attempts = attempts
        self.category = category
        self.original = original


def _config_generation_extra_kwargs() -> dict[str, float]:
    raw_timeout = os.getenv("AI_GRADING_CONFIG_TIMEOUT_SECONDS")
    try:
        timeout = float(raw_timeout) if raw_timeout else DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS
    except (TypeError, ValueError):
        timeout = DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS
    return {"timeout": max(120.0, timeout)}


def save_uploaded_json(upload_dir: Path, file_name: str, file_bytes: bytes) -> Path:
    upload_dir.mkdir(parents=True, exist_ok=True)
    output_path = upload_dir / file_name

    parsed = json.load(io.BytesIO(file_bytes))
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)

    return output_path


def extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    seen_lines: set[str] = set()

    for block in _iter_doc_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text and text not in seen_lines:
                lines.append(text)
                seen_lines.add(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    cell_text = "\n".join(
                        para.text.strip()
                        for para in cell.paragraphs
                        if para.text.strip()
                    )
                    if cell_text:
                        cells.append(cell_text)
                row_text = " | ".join(cells)
                if row_text and row_text not in seen_lines:
                    lines.append(row_text)
                    seen_lines.add(row_text)

    # Some exam files put questions in Word text boxes/shapes. python-docx
    # does not expose those as normal paragraphs, so read raw XML text too.
    for text in _extract_docx_xml_text(file_bytes):
        if text and text not in seen_lines:
            lines.append(text)
            seen_lines.add(text)

    if not lines:
        raise ValueError("Word 鏂囨。鏈В鏋愬埌鏈夋晥鏂囨湰")

    return "\n".join(lines)


def _iter_doc_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _extract_docx_xml_text(file_bytes: bytes) -> list[str]:
    """Extract text from XML streams: main doc, headers/footers, footnotes, and text boxes."""
    xml_names = [
        "word/document.xml",
        "word/footnotes.xml",
        "word/endnotes.xml",
    ]
    result: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            xml_names.extend(
                name
                for name in archive.namelist()
                if re.match(r"word/(header|footer)\d+\.xml$", name)
            )
            for name in xml_names:
                if name not in archive.namelist():
                    continue
                raw_xml = archive.read(name)
                root = ElementTree.fromstring(raw_xml)
                # ① Extract OMML math formula as readable placeholder before iterating paragraphs
                # OMML uses namespace http://schemas.openxmlformats.org/officeDocument/2006/math
                for omath in root.iter():
                    tag = omath.tag
                    if not (tag.endswith("}oMath") or tag.endswith("}oMathPara")):
                        continue
                    # Collect all literal text runs inside the formula
                    math_chunks: list[str] = []
                    for node in omath.iter():
                        if node.tag.endswith("}t") and node.text:
                            math_chunks.append(node.text)
                        elif node.tag.endswith("}r") and node.text:
                            math_chunks.append(node.text)
                    formula_text = "".join(math_chunks).strip()
                    if formula_text:
                        placeholder = f"[公式: {formula_text}]"
                    else:
                        placeholder = "[数学公式]"
                    if placeholder not in result:
                        result.append(placeholder)

                # ② Extract paragraphs from main body AND from text boxes (txbxContent)
                for paragraph in root.iter():
                    if not paragraph.tag.endswith("}p"):
                        continue
                    chunks: list[str] = []
                    for node in paragraph.iter():
                        if node.tag.endswith("}t") and node.text:
                            chunks.append(node.text)
                        elif node.tag.endswith("}tab"):
                            chunks.append("\t")
                        elif node.tag.endswith("}br"):
                            chunks.append("\n")
                    text = re.sub(r"[ \t]+", " ", "".join(chunks)).strip()
                    if text and text not in result:
                        result.append(text)
    except Exception:
        return []
    return result


def _aligned_whole_generation_rules() -> str:
    return (
        "\n\n与分题生成模式一致的硬性规则：\n"
        "1. 每个 parts 项必须输出 response_mode：exact_objective、short_answer_points、"
        "process_required 或 visual_construction。\n"
        "2. 选择题和普通填空题只按最终答案判分，不得要求推理或计算过程。\n"
        "3. 要求列出全部可能答案的填空题必须使用 match_mode=complete_set，并输出 required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；少写、错写、多写均不得分。\n"
        "4. 必须输出具体知识点 knowledge_name、knowledge_id、knowledge_points，不得用题干或"
        "“几何综合/代数综合/综合应用”充当知识点。\n"
        "5. 主观题评分点必须写出可核验的必要条件、式子或结论，不得只写通用描述。\n"
        "6. 作图题应输出 visual_requirements 和必要踩分点，不得把答案图臆造为唯一文字答案。\n"
        "7. 总分严格为100；单题不超过18分；相同类型客观题必须同分，其他题型不要求同分。\n"
    )


def _build_whole_text_generation_prompt(doc_text: str) -> str:
    return (
        "这是 Word 整卷单次请求。你必须在本次响应中一次完成所有题目的解析、评分点生成与赋分；"
        "不要建议后续补充请求。\n"
        + _build_generation_prompt("", include_source_text=False)
        + _aligned_whole_generation_rules()
        + f"\nWord 原文（唯一文本来源）：\n{doc_text}"
    )


def _build_whole_image_generation_prompt() -> str:
    return (
        "这是 PDF 整卷视觉单次请求。后续附带的图片按顺序对应整份试卷及答案/解析的各页原图。"
        "图片是唯一权威内容来源，不得参考、猜测或恢复任何 PDF 抽取文字。"
        "你必须在本次响应中一次完成所有题目的解析、评分点生成与赋分；不要建议后续补充请求。\n\n"
        + _build_generation_prompt("", include_source_text=False)
        + _aligned_whole_generation_rules()
    )


def _finalize_whole_generation_payload(
    payload: dict[str, Any],
    generation_mode: str,
) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("AI 整卷生成未返回有效 JSON 对象。")
    normalize_generated_config_schema(payload)
    force_payload_total_score(payload, target_total=100.0)
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["generation_mode"] = generation_mode
    meta["score_allocation_mode"] = "single_request_local_normalization"
    meta["single_request"] = True
    refresh_generated_config_quality_warnings(payload)
    try:
        validate_generated_config(payload)
    except Exception:
        _dump_failed_generated_payload(payload)
        raise
    return payload


def generate_grading_config_from_docx_text(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    if report:
        report(0.18, "Word 整卷单次请求", "直接把整份 Word 文本发送给模型，一次完成解析与赋分。")
    prompt = _build_whole_text_generation_prompt(doc_text)
    payload = llm_client.json_from_text_once(
        prompt,
        model=model_name,
        extra_kwargs=_config_generation_extra_kwargs(),
    )
    return _finalize_whole_generation_payload(payload, "whole_word_text_single_request")


def generate_grading_config_from_docx_text_legacy(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    if report:
        report(0.18, "旧版整卷生成", "直接把整份 Word 文本发送给模型生成评分标准。")
    prompt = _build_generation_prompt(doc_text)
    payload = llm_client.json_from_text(prompt, model=model_name, extra_kwargs=_config_generation_extra_kwargs())
    validate_generated_config(payload)
    if _needs_objective_repair(payload, doc_text):
        if report:
            report(0.70, "旧版整卷修复", "检测到客观题可能缺失，正在请求模型修复。")
        payload = llm_client.json_from_text(
            _build_objective_repair_prompt(doc_text, payload),
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
        )
        validate_generated_config(payload)
    return payload


def _config_generation_retry_delays() -> tuple[float, ...]:
    raw = str(os.getenv("AI_GRADING_CONFIG_RETRY_DELAYS") or "").strip()
    if not raw:
        return DEFAULT_CONFIG_GENERATION_RETRY_DELAYS
    try:
        values = tuple(max(0.0, float(item.strip())) for item in raw.split(",") if item.strip())
    except ValueError:
        return DEFAULT_CONFIG_GENERATION_RETRY_DELAYS
    return values or DEFAULT_CONFIG_GENERATION_RETRY_DELAYS


def _is_transient_config_generation_error(exc: Exception) -> bool:
    status_code = getattr(exc, "status_code", None)
    if status_code == 429 or (isinstance(status_code, int) and 500 <= status_code <= 599):
        return True
    class_name = exc.__class__.__name__.lower()
    if any(token in class_name for token in ("apiconnection", "apitimeout", "ratelimit", "timeout", "connection")):
        return True
    message = str(exc).lower()
    return any(
        token in message
        for token in (
            "connection reset",
            "connection aborted",
            "connection error",
            "temporarily unavailable",
            "temporary upstream",
            "timed out",
            "timeout",
            "rate limit",
            "too many requests",
            "service unavailable",
            "bad gateway",
            "gateway timeout",
        )
    )


def _call_question_generation_with_retry(
    question_id: str,
    request: Callable[[], dict[str, Any]],
) -> tuple[dict[str, Any], int]:
    retry_delays = _config_generation_retry_delays()
    attempts = 0
    while True:
        attempts += 1
        try:
            return request(), attempts
        except Exception as exc:
            transient = _is_transient_config_generation_error(exc)
            retry_index = attempts - 1
            if not transient or retry_index >= len(retry_delays):
                category = "transient_network" if transient else "non_retryable"
                raise _QuestionGenerationRequestError(question_id, attempts, category, exc) from exc
            time.sleep(retry_delays[retry_index])


def _word_block_image_blobs(block: dict[str, Any], *, limit: int = 8) -> list[bytes]:
    raw_paths: list[str] = []
    image_paths = block.get("image_paths")
    if isinstance(image_paths, list):
        raw_paths.extend(str(path).strip() for path in image_paths if str(path).strip())
    for key in ("question_html", "answer_html", "analysis_html"):
        value = str(block.get(key) or "")
        raw_paths.extend(_image_paths_from_rich_text(value))

    blobs: list[bytes] = []
    seen: set[str] = set()
    for raw_path in raw_paths:
        candidates = [Path(raw_path)]
        if not Path(raw_path).is_absolute():
            candidates.append(Path.cwd() / raw_path)
        path = next((candidate for candidate in candidates if candidate.is_file()), None)
        if path is None:
            continue
        resolved = str(path.resolve())
        if resolved in seen:
            continue
        seen.add(resolved)
        try:
            blobs.append(path.read_bytes())
        except OSError:
            continue
        if len(blobs) >= limit:
            break
    return blobs


def _generate_question_block_results(
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any] | None], list[dict[str, Any]], dict[str, int], int, bool]:
    max_workers = _bounded_int(
        os.getenv("AI_GRADING_CONFIG_WORKERS") or os.getenv("AI_GRADING_MAX_WORKERS"),
        8,
        1,
        1000,
    )
    worker_count = min(max_workers, len(question_blocks))
    results: list[dict[str, Any] | None] = [None] * len(question_blocks)
    failures: list[dict[str, Any]] = []
    attempt_counts: dict[str, int] = {}
    image_semantic_mode = bool(q_images) or any(
        str(block.get("semantic_source") or "").strip() == "images"
        for block in question_blocks
    )

    if report:
        qids = ", ".join(str(block.get("question_id") or "") for block in question_blocks if block.get("question_id"))
        report(0.18, "Split paper", f"Parsed {len(question_blocks)} questions; workers={worker_count}; qids: {qids}")

    import base64

    with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="config-q") as executor:
        future_map = {}
        for index, block in enumerate(question_blocks):
            qid = str(block.get("question_id") or f"Q{index + 1}")
            image_data = q_images.get(qid) if isinstance(q_images, dict) else None
            has_question_images = bool(isinstance(image_data, dict) and image_data.get("question"))
            image_semantic_source = str(block.get("semantic_source") or "").strip() == "images"
            word_image_blobs = (
                _word_block_image_blobs(block)
                if not has_question_images and not image_semantic_source
                else []
            )
            prompt = (
                _build_single_question_image_generation_prompt(
                    block,
                    has_answer_image=bool(isinstance(image_data, dict) and image_data.get("answer")),
                )
                if has_question_images or image_semantic_source
                else _build_single_question_generation_prompt(block, doc_text)
            )
            if has_question_images and isinstance(image_data, dict):
                img_list = [base64.b64decode(image_data["question"], validate=True)]
                if image_data.get("answer"):
                    img_list.append(base64.b64decode(image_data["answer"], validate=True))

                def request(prompt: str = prompt, img_list: list[bytes] = img_list) -> dict[str, Any]:
                    return llm_client.json_from_images(prompt, img_list, model=model_name)
            elif word_image_blobs:
                def request(prompt: str = prompt, img_list: list[bytes] = word_image_blobs) -> dict[str, Any]:
                    return llm_client.json_from_images(prompt, img_list, model=model_name)
            else:
                def request(prompt: str = prompt) -> dict[str, Any]:
                    return llm_client.json_from_text(prompt, model=model_name)

            future_map[executor.submit(_call_question_generation_with_retry, qid, request)] = index

        if report:
            qids = ", ".join(str(block.get("question_id") or "") for block in question_blocks if block.get("question_id"))
            report(0.20, "Submit per-question requests", f"Submitted {len(future_map)} question requests; waiting for model: {qids}")
        completed = 0
        total = len(question_blocks)
        for future in as_completed(future_map):
            index = future_map[future]
            qid = str(question_blocks[index].get("question_id") or f"Q{index + 1}")
            try:
                results[index], attempt_counts[qid] = future.result()
            except _QuestionGenerationRequestError as exc:
                attempt_counts[qid] = exc.attempts
                failures.append(
                    {
                        "question_id": qid,
                        "attempts": exc.attempts,
                        "category": exc.category,
                        "error": str(exc.original),
                    }
                )
            except Exception as exc:
                attempt_counts[qid] = 1
                failures.append(
                    {
                        "question_id": qid,
                        "attempts": 1,
                        "category": "unexpected",
                        "error": str(exc),
                    }
                )
            completed += 1
            if report:
                progress = 0.18 + 0.7 * (completed / total)
                report(progress, "Parse question", f"Question {index + 1} parsed (progress: {completed}/{total})")

    return results, failures, attempt_counts, worker_count, image_semantic_mode


def _append_unmergeable_question_failures(
    merged: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    attempt_counts: dict[str, int],
) -> list[dict[str, Any]]:
    questions = merged.get("rubric", {}).get("questions", [])
    merged_qids = {
        str(question.get("question_id") or "").strip()
        for question in questions
        if isinstance(question, dict)
    } if isinstance(questions, list) else set()
    failed_qids = {str(item.get("question_id") or "").strip() for item in failures}
    for block in question_blocks:
        qid = str(block.get("question_id") or "").strip()
        if not qid or qid in merged_qids or qid in failed_qids:
            continue
        failures.append(
            {
                "question_id": qid,
                "attempts": int(attempt_counts.get(qid, 1)),
                "category": "unmergeable_schema",
                "error": "AI returned a result that could not be merged into the rubric schema",
            }
        )
    return failures


def _generate_grading_config_by_question_blocks(
    question_blocks: list[dict[str, str]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    _validate_image_semantic_inputs(question_blocks, q_images)
    results, failures, attempt_counts, worker_count, image_semantic_mode = _generate_question_block_results(
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )

    if report:
        report(0.90, "Assemble rubric", "All questions parsed; assembling final scoring rubric...")

    merged = _merge_single_question_payloads(results, question_blocks)
    failures = _append_unmergeable_question_failures(merged, question_blocks, failures, attempt_counts)
    _attach_parallel_generation_meta(merged, question_blocks, failures, worker_count, attempt_counts)
    _ensure_question_blocks_covered(merged, question_blocks)
    _apply_local_question_facts(merged, question_blocks)
    normalize_generated_config_schema(merged)

    if failures:
        meta = merged.setdefault("meta", {})
        meta["score_allocation_mode"] = "pending_failed_questions"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = True
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    # ① Phase 2: dedicated score-allocation AI call
    # All per-question prompts set score placeholders (=1); this step assigns real scores.
    return retry_grading_config_score_allocation(
        merged,
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        include_document_text=not image_semantic_mode,
    )


def _validate_image_semantic_inputs(
    question_blocks: list[dict[str, Any]],
    q_images: dict[str, Any] | None,
) -> None:
    image_qids = [
        str(block.get("question_id") or "").strip()
        for block in question_blocks
        if isinstance(block, dict) and str(block.get("semantic_source") or "").strip() == "images"
    ]
    if not image_qids:
        return
    if not isinstance(q_images, dict) or not q_images:
        raise ValueError("PDF 图片裁题状态已丢失，请重新点击“① 拆分试卷”后再生成评分标准。")

    import base64

    invalid_qids: list[str] = []
    for qid in image_qids:
        image_data = q_images.get(qid)
        question_image = image_data.get("question") if isinstance(image_data, dict) else None
        if not isinstance(question_image, str) or not question_image.strip():
            invalid_qids.append(qid)
            continue
        try:
            base64.b64decode(question_image, validate=True)
            answer_image = image_data.get("answer")
            if answer_image:
                base64.b64decode(str(answer_image), validate=True)
        except (ValueError, TypeError):
            invalid_qids.append(qid)
    if invalid_qids:
        raise ValueError(
            "PDF 图片裁题数据无效，未退回纯文本模式。请重新拆题："
            + ", ".join(invalid_qids)
        )


def preview_question_blocks_from_docx_text(doc_text: str) -> list[dict[str, Any]]:
    """本地拆题预览，不调用任何 AI。

    供 UI 在发送 AI 之前展示拆题结果，让人工确认每道题拆分是否正确、答案是否靠谱、
    题型判断是否准确（可在预览中删题或切换题型）。直接复用本地拆题逻辑。
    """
    return _extract_local_question_blocks(doc_text)


_INLINE_IMAGE_MARKER = re.compile(r"\[\[IMAGE:(?P<path>.+?)\]\]")
_INLINE_MAIN_QUESTION_MARKER = re.compile(
    r"(?P<prefix>[。！？!?．.][ \t\r\n]*)"
    r"(?P<number>\d{1,2})[ \t]*[.．、](?![ \t]*\d)[ \t]*"
)


class _ControlledDocxWriteError(RuntimeError):
    pass


def preview_question_blocks_from_docx_bytes(
    file_bytes: bytes,
    *,
    fallback_doc_text: str = "",
    temporary_root: str | Path | None = None,
    asset_root: str | Path | None = None,
    register_created_file: Callable[[Path], None] | None = None,
    write_created_file: Callable[[Path, bytes], None] | None = None,
) -> list[dict[str, Any]]:
    """富文本拆题预览（复用题库 import_docx），保留公式 HTML 与图片，不调用 AI。

    解析失败时回退到纯文本拆题（preview_question_blocks_from_docx_text）。
    """
    from backend.document_parsing import parse_docx_question_blocks

    return parse_docx_question_blocks(
        file_bytes,
        fallback_doc_text=fallback_doc_text,
        temporary_root=(
            temporary_root
            if temporary_root is not None
            else _resolve_upload_config_dir()
        ),
        asset_root=asset_root,
        register_created_file=register_created_file,
        write_created_file=write_created_file,
    )


def _extract_rich_question_blocks(
    file_bytes: bytes,
    *,
    temporary_root: str | Path | None = None,
    asset_root: str | Path | None = None,
    register_created_file: Callable[[Path], None] | None = None,
    write_created_file: Callable[[Path, bytes], None] | None = None,
) -> list[dict[str, Any]] | None:
    """用题库 import_docx + map_rich_content_by_number 提取每题富文本块。"""
    from question_bank.importers.docx_importer import import_docx
    from question_bank.importers.batch_importer import map_rich_content_by_number

    tmp_dir = (
        Path(temporary_root)
        if temporary_root is not None
        else Path(_resolve_upload_config_dir())
    )
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    tmp_path = tmp_dir / f"_rich_split_{ts}.docx"

    def controlled_writer(path: Path, content: bytes) -> None:
        assert write_created_file is not None
        try:
            write_created_file(path, content)
        except Exception as error:
            raise _ControlledDocxWriteError() from error

    extracted = import_docx(
        io.BytesIO(file_bytes),
        source_name=tmp_path,
        asset_root=asset_root,
        asset_root_is_output_dir=write_created_file is not None,
        register_created_file=register_created_file,
        write_created_file=(controlled_writer if write_created_file is not None else None),
    )
    content = map_rich_content_by_number(
        _split_inline_main_question_paragraphs(extracted.rich_paragraphs),
        source_file=str(tmp_path),
    )

    question_map = content.get("question") if isinstance(content, dict) else {}
    answer_map = content.get("answer") if isinstance(content, dict) else {}
    if not isinstance(question_map, dict) or not question_map:
        return None

    # 题号按数值升序
    def _num_key(k: str) -> int:
        try:
            return int(str(k).strip())
        except (TypeError, ValueError):
            return 999

    choice_answers = _extract_choice_answer_sequence(
        _rich_blocks_plain_text(_flatten_answer_blocks(answer_map))
    )

    blocks: list[dict[str, Any]] = []
    for num_key in sorted(question_map.keys(), key=_num_key):
        number = _num_key(num_key)
        if not (1 <= number <= 99):
            continue
        q_blocks = question_map.get(num_key) or []
        a_blocks = answer_map.get(num_key) if isinstance(answer_map, dict) else None
        block = _parse_rich_question_blocks(number, q_blocks, a_blocks, choice_answers)
        if block:
            blocks.append(block)
    return blocks or None


def _split_inline_main_question_paragraphs(
    rich_paragraphs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Split only consecutive main-question markers embedded in DOCX paragraphs."""
    normalized: list[dict[str, Any]] = []
    current_number: int | None = None
    in_answer_section = False
    plain_paragraphs = [
        _strip_inline_html(
            _INLINE_IMAGE_MARKER.sub("", str(paragraph.get("text") or ""))
        ).strip()
        for paragraph in rich_paragraphs
    ]

    for index, paragraph in enumerate(rich_paragraphs):
        text = str(paragraph.get("text") or "").strip()
        plain_text = plain_paragraphs[index]
        if _looks_like_answer_section_heading(plain_text):
            in_answer_section = True
        if in_answer_section or not text:
            normalized.append(paragraph)
            continue

        leading_number = _extract_question_marker_number(plain_text)
        if leading_number is not None:
            current_number = leading_number
        if current_number is None:
            normalized.append(paragraph)
            continue

        text_without_images = _INLINE_IMAGE_MARKER.sub("", text).strip()
        segments, final_number = _split_consecutive_inline_main_questions(
            text_without_images,
            current_number=current_number,
            following_number=_next_leading_main_question_number(
                plain_paragraphs,
                after_index=index,
            ),
        )
        if len(segments) == 1:
            normalized.append(paragraph)
            continue

        segments = _attach_inline_images_to_source_segments(
            paragraph,
            segments=segments,
            current_number=current_number,
            final_number=final_number,
        )
        for segment in segments:
            normalized.append({**paragraph, "text": segment})
        current_number = final_number

    return normalized


def _split_consecutive_inline_main_questions(
    text: str,
    *,
    current_number: int,
    following_number: int | None,
) -> tuple[list[str], int]:
    split_offsets: list[int] = []
    expected_number = current_number + 1
    for marker in _INLINE_MAIN_QUESTION_MARKER.finditer(text):
        marker_number = int(marker.group("number"))
        if marker_number != expected_number:
            continue
        split_offsets.append(marker.start("number"))
        expected_number += 1
    if not split_offsets or expected_number != following_number:
        return [text], current_number

    boundaries = [0, *split_offsets, len(text)]
    segments = [
        text[start:end].strip()
        for start, end in zip(boundaries, boundaries[1:])
        if text[start:end].strip()
    ]
    return segments, expected_number - 1


def _attach_inline_images_to_source_segments(
    paragraph: dict[str, Any],
    *,
    segments: list[str],
    current_number: int,
    final_number: int,
) -> list[str]:
    fallback_paths = _image_paths_from_rich_text(
        str(paragraph.get("text") or "")
    )

    def attach_to_last(paths: list[str]) -> list[str]:
        assigned = list(segments)
        for path in paths:
            assigned[-1] = f"{assigned[-1].rstrip()}\n[[IMAGE:{path}]]"
        return assigned

    image_relationships = paragraph.get("image_relationships")
    raw_xml = str(paragraph.get("xml") or "")
    if not isinstance(image_relationships, dict) or not image_relationships or not raw_xml:
        return attach_to_last(fallback_paths)

    try:
        root = ElementTree.fromstring(raw_xml)
    except ElementTree.ParseError:
        return attach_to_last(fallback_paths)

    text_parts: list[str] = []
    positioned_images: list[tuple[int, str]] = []
    text_length = 0
    for element in root.iter():
        local_name = str(element.tag).split("}")[-1]
        if local_name == "t":
            value = str(element.text or "")
            text_parts.append(value)
            text_length += len(value)
            continue
        if local_name != "blip":
            continue
        relationship_id = next(
            (
                str(value)
                for key, value in element.attrib.items()
                if str(key).split("}")[-1] == "embed"
            ),
            "",
        )
        image_path = str(image_relationships.get(relationship_id) or "").strip()
        if image_path:
            positioned_images.append((text_length, image_path))

    if not positioned_images:
        return attach_to_last(fallback_paths)

    source_text = "".join(text_parts)
    expected_number = current_number + 1
    split_offsets: list[int] = []
    for marker in _INLINE_MAIN_QUESTION_MARKER.finditer(source_text):
        marker_number = int(marker.group("number"))
        if marker_number != expected_number:
            continue
        split_offsets.append(marker.start("number"))
        expected_number += 1
        if marker_number == final_number:
            break
    if len(split_offsets) != len(segments) - 1:
        return attach_to_last(fallback_paths)

    assigned = list(segments)
    positioned_paths: set[str] = set()
    for image_offset, image_path in positioned_images:
        segment_index = sum(image_offset >= offset for offset in split_offsets)
        assigned[segment_index] = (
            f"{assigned[segment_index].rstrip()}\n[[IMAGE:{image_path}]]"
        )
        positioned_paths.add(image_path)
    for image_path in fallback_paths:
        if image_path not in positioned_paths:
            assigned[-1] = f"{assigned[-1].rstrip()}\n[[IMAGE:{image_path}]]"
    return assigned


def _next_leading_main_question_number(
    texts: list[str],
    *,
    after_index: int,
) -> int | None:
    for text in texts[after_index + 1 :]:
        if _looks_like_answer_section_heading(text):
            return None
        number = _extract_question_marker_number(text)
        if number is not None:
            return number
    return None


def _flatten_answer_blocks(answer_map: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    if isinstance(answer_map, dict):
        for _k, blocks in answer_map.items():
            if isinstance(blocks, list):
                out.extend(b for b in blocks if isinstance(b, dict))
    return out


def _rich_block_text(block: Any) -> str:
    if isinstance(block, dict):
        return str(block.get("text") or "")
    return str(block or "")


def _rich_blocks_plain_text(blocks: list[Any]) -> str:
    parts = [_rich_block_text(b) for b in blocks]
    return _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", "\n".join(parts)))


def _strip_inline_html(value: str) -> str:
    """去掉公式 HTML 标签，得到可用于答案/题型推断的纯文本。"""
    text = re.sub(r"</?(?:sub|sup|u|table|tbody|tr|td|th|br)\b[^>]*>", "", str(value or ""), flags=re.IGNORECASE)
    return text


def _image_paths_from_rich_text(value: str) -> list[str]:
    seen: list[str] = []
    for m in _INLINE_IMAGE_MARKER.finditer(str(value or "")):
        p = m.group("path").strip()
        if p and p not in seen:
            seen.append(p)
    return seen


def _parse_answer_section_blocks(answer_blocks: list[Any]) -> tuple[str, str]:
    """从卷末"参考答案"块中提取（最终答案, 解答过程）。

    答案块通常重列了题干/选项，需跳过；只保留 【分析】/【解答】 后、【点评】 前的解答内容，
    并优先抽取"故选X""故答案为：X"作为最终答案。
    """
    in_solution = False
    solution_lines: list[str] = []
    final_answer = ""
    for blk in answer_blocks:
        text = _rich_block_text(blk)
        for raw in (text.splitlines() if "\n" in text else [text]):
            line = str(raw or "").strip()
            if not line:
                continue
            if "【点评】" in line or "【点睛】" in line:
                in_solution = False
                continue
            if "【解答】" in line or "【分析】" in line or "【解析】" in line:
                in_solution = True
                marker = next(m for m in ("【解答】", "【分析】", "【解析】") if m in line)
                after = line.split(marker, 1)[1].strip()
                if after:
                    solution_lines.append(after)
                continue
            if line.startswith("【") and "】" in line:
                # 其他标记（如【答案】）后内容并入解答
                in_solution = True
                after = line.split("】", 1)[1].strip()
                if after:
                    solution_lines.append(after)
                continue
            if in_solution:
                solution_lines.append(line)

    solution_text = "\n".join(solution_lines).strip()
    # 抽取最终答案：故选X / 故答案为：X
    plain = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", solution_text))
    m = re.search(r"故\s*选\s*[:：]?\s*([A-Da-d]+)", plain)
    if m:
        final_answer = m.group(1).upper()
    else:
        m = re.search(r"故\s*答\s*案\s*为\s*[:：]?\s*([^。\n．]+)", plain)
        if m:
            final_answer = m.group(1).strip().rstrip("．.")
    return final_answer, solution_text


def _parse_rich_question_blocks(
    number: int,
    question_blocks: list[Any],
    answer_blocks: list[Any] | None,
    choice_answers: dict[str, str],
) -> dict[str, Any] | None:
    """把单题的富文本块（题干区，含内联【答案】【解析】）拆成 题干/答案/解析 富文本。"""
    # 题干区逐行送入分桶（沿用纯文本分桶逻辑），保留 HTML 公式与 [[IMAGE:]] 标记
    lines: list[str] = []
    for blk in question_blocks:
        text = _rich_block_text(blk)
        lines.extend(text.splitlines() if "\n" in text else [text])

    bucket = "stem"
    stem_lines: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []
    for raw in lines:
        line = str(raw or "").strip()
        if not line:
            continue
        if "【答案】" in line:
            bucket = "answer"
            after = line.split("【答案】", 1)[1].strip()
            if after:
                answer_lines.append(after)
            continue
        if ("【解析】" in line) or ("【点睛】" in line):
            bucket = "analysis"
            marker = "【解析】" if "【解析】" in line else "【点睛】"
            after = line.split(marker, 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if line.startswith("【") and "】" in line:
            bucket = "analysis"
            after = line.split("】", 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if bucket == "stem":
            stem_lines.append(line)
        elif bucket == "answer":
            answer_lines.append(line)
        else:
            analysis_lines.append(line)

    # 若题干区没有内联答案，但存在独立答案块（卷末"参考答案"区），从中解析。
    # 答案块结构通常为：[重列题干/选项] 【分析】… 【解答】…故选C/故答案为：X 【点评】…
    # 需跳过重列的题干/选项，只取标记后的解答与最终答案。
    if answer_blocks and not answer_lines:
        ans_final, ana_text = _parse_answer_section_blocks(answer_blocks)
        if ans_final:
            answer_lines.append(ans_final)
        if ana_text and not analysis_lines:
            analysis_lines.append(ana_text)

    question_html = _strip_leading_question_number(number, "\n".join(stem_lines).strip())
    answer_html = "\n".join(answer_lines).strip()
    analysis_html = "\n".join(analysis_lines).strip()

    image_paths = _image_paths_from_rich_text(question_html)

    question_text = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", question_html)).strip()
    answer_text = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", answer_html)).strip()

    if not question_text and not answer_text and not analysis_html and not image_paths:
        return None

    num_str = str(number)
    qtype = _infer_local_question_type(question_text, answer_text, num_str)
    canonical = _extract_canonical_answer_for_local_question(
        number=num_str,
        qtype=qtype,
        question_text=question_text,
        answer_text=answer_text,
        choice_answers=choice_answers,
    )
    accepted = _local_accepted_forms(canonical, qtype)
    has_answer = bool(answer_text or canonical)
    return {
        "question_id": f"Q{number}",
        "text": question_text,
        "question_text": question_text,
        "question_html": question_html,
        "answer_text": answer_text,
        "answer_html": answer_html,
        "analysis": _strip_inline_html(analysis_html).strip(),
        "analysis_html": analysis_html,
        "image_paths": image_paths,
        "question_type": qtype,
        "canonical_answer": canonical,
        "accepted_forms": accepted,
        "local_answer_trusted": has_answer,
        "needs_review": (not bool(canonical)) if qtype in {"choice", "fill_blank"} else False,
    }


def _resolve_upload_config_dir() -> str:
    try:
        from path_manager import PathManager  # type: ignore

        return str(PathManager().upload_config_dir)
    except Exception:
        return str(Path.cwd() / "user_data" / "config" / "uploaded")


def generate_grading_config_from_confirmed_blocks(
    confirmed_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    """用人工确认后的题块直接进入单题并发生成 + 统一赋分管线，跳过重新拆题。

    confirmed_blocks 应来自 preview_question_blocks_from_docx_text 的输出，
    经人工删除/题型修正后传入。
    """
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in confirmed_blocks
    ]
    return _generate_grading_config_by_question_blocks(
        locked_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )


class _BatchSchemaMismatch(ValueError):
    pass


def generate_grading_config_in_batches(
    confirmed_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    batch_size: int = DEFAULT_CONFIG_GENERATION_BATCH_SIZE,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Generate one rubric in small sequential batches with no hidden model calls."""
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in confirmed_blocks
    ]
    if not locked_blocks:
        raise ValueError("至少需要一道已确认题目。")
    clean_size = max(1, min(20, int(batch_size)))
    return _run_config_generation_batches(
        existing_payload=None,
        question_blocks=locked_blocks,
        doc_text=doc_text,
        llm_client=llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        batch_size=clean_size,
        retry_question_ids=None,
        checkpoint=checkpoint,
    )


def retry_failed_grading_config_batches(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    retry_question_ids: Sequence[str] | None = None,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    meta = existing_payload.get("meta") if isinstance(existing_payload, dict) else None
    failed_batches = meta.get("failed_batches") if isinstance(meta, dict) else None
    if not isinstance(failed_batches, list) or not failed_batches:
        return _score_completed_batch_draft_once(
            existing_payload,
            question_blocks,
            doc_text,
            llm_client,
            model_name=model_name,
            report=report,
            q_images=q_images,
            checkpoint=checkpoint,
        )

    failed_by_id: dict[str, list[str]] = {}
    for item in failed_batches:
        if not isinstance(item, dict):
            continue
        batch_id = str(item.get("batch_id") or "").strip()
        raw_ids = item.get("question_ids")
        ids = [str(qid).strip() for qid in raw_ids or [] if str(qid).strip()]
        if batch_id and ids:
            failed_by_id[batch_id] = ids
    if not failed_by_id:
        raise ValueError("失败批次记录已损坏，请重新开始分批生成。")

    if retry_question_ids is None:
        selected_question_ids = {
            qid for ids in failed_by_id.values() for qid in ids
        }
    else:
        selected_ids = list(
            dict.fromkeys(str(qid).strip() for qid in retry_question_ids if str(qid).strip())
        )
        selected_set = set(selected_ids)
        selected_batch_ids = [
            batch_id
            for batch_id, ids in failed_by_id.items()
            if set(ids).issubset(selected_set)
        ]
        expected_set = {
            qid for batch_id in selected_batch_ids for qid in failed_by_id[batch_id]
        }
        if not selected_set or selected_set != expected_set:
            raise ValueError("只能按完整失败批次重试，不能只选择批次中的部分题目。")
        selected_question_ids = expected_set

    existing_size = meta.get("batch_size") if isinstance(meta, dict) else None
    try:
        batch_size = int(existing_size)
    except (TypeError, ValueError):
        batch_size = DEFAULT_CONFIG_GENERATION_BATCH_SIZE
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in question_blocks
    ]
    return _run_config_generation_batches(
        existing_payload=existing_payload,
        question_blocks=locked_blocks,
        doc_text=doc_text,
        llm_client=llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        batch_size=batch_size,
        retry_question_ids=selected_question_ids,
        checkpoint=checkpoint,
    )


def failed_grading_config_batches(payload: dict[str, Any]) -> list[dict[str, Any]]:
    meta = payload.get("meta") if isinstance(payload, dict) else None
    value = meta.get("failed_batches") if isinstance(meta, dict) else None
    return copy.deepcopy(value) if isinstance(value, list) else []


def _run_config_generation_batches(
    *,
    existing_payload: dict[str, Any] | None,
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None,
    report: Any,
    q_images: dict[str, Any] | None,
    batch_size: int,
    retry_question_ids: set[str] | None,
    checkpoint: Callable[[dict[str, Any]], None] | None,
) -> dict[str, Any]:
    _validate_unique_batch_question_ids(question_blocks)
    _validate_image_semantic_inputs(question_blocks, q_images)
    batch_groups: list[list[dict[str, Any]]] = []
    objective_group: list[dict[str, Any]] = []
    for block in question_blocks:
        question_type = str(block.get("question_type") or "").strip().lower()
        if question_type in {"choice", "fill_blank"}:
            objective_group.append(block)
            if len(objective_group) >= batch_size:
                batch_groups.append(objective_group)
                objective_group = []
            continue
        if objective_group:
            batch_groups.append(objective_group)
            objective_group = []
        batch_groups.append([block])
    if objective_group:
        batch_groups.append(objective_group)

    batches = [
        {
            "batch_id": f"B{index + 1:03d}",
            "question_ids": [
                str(block.get("question_id") or "").strip() for block in group
            ],
            "blocks": group,
        }
        for index, group in enumerate(batch_groups)
    ]
    merged = copy.deepcopy(existing_payload) if existing_payload is not None else {
        "rubric": {"exam_title": "generated", "total_score": 100, "questions": []},
        "answer_key": {"questions": []},
        "meta": {"warnings": []},
    }
    states = _existing_batch_states(merged)
    if retry_question_ids is None:
        targets = list(batches)
    else:
        targets = [
            batch
            for batch in batches
            if set(batch["question_ids"]).issubset(retry_question_ids)
        ]
        targeted_question_ids = {
            qid for batch in targets for qid in batch["question_ids"]
        }
        if targeted_question_ids != retry_question_ids:
            raise ValueError("旧草稿的失败题目不能安全映射到新的分批规则。")
    total_targets = len(targets)
    for completed, batch in enumerate(targets, start=1):
        batch_id = str(batch["batch_id"])
        question_ids = list(batch["question_ids"])
        if report:
            progress = 0.10 + 0.75 * ((completed - 1) / max(1, total_targets))
            report(
                progress,
                "分批生成评分标准",
                f"正在生成批次 {batch_id}（{', '.join(question_ids)}）",
            )
        try:
            prompt, image_blobs = _build_config_batch_request(
                list(batch["blocks"]), doc_text, q_images
            )
            if image_blobs:
                batch_payload = llm_client.json_from_images_once(
                    prompt,
                    image_blobs,
                    model=model_name,
                    extra_kwargs=_config_generation_extra_kwargs(),
                    use_config_client=True,
                )
            else:
                batch_payload = llm_client.json_from_text_once(
                    prompt,
                    model=model_name,
                    extra_kwargs=_config_generation_extra_kwargs(),
                )
            _validate_exact_batch_payload(
                batch_payload,
                question_ids,
                require_canonical_answer=False,
            )
            normalize_generated_config_schema(batch_payload)
            _validate_exact_batch_payload(batch_payload, question_ids)
            _replace_retry_question_payloads(
                merged,
                batch_payload,
                set(question_ids),
                question_blocks,
            )
            batch_meta = batch_payload.get("meta") if isinstance(batch_payload, dict) else None
            repair = batch_meta.get("local_json_repair") if isinstance(batch_meta, dict) else None
            state: dict[str, Any] = {
                "batch_id": batch_id,
                "question_ids": question_ids,
                "status": "succeeded",
            }
            if isinstance(repair, dict):
                state["local_json_repair"] = {
                    "repaired": bool(repair.get("repaired")),
                    "operations": [str(item) for item in repair.get("operations") or []],
                    "response_chars": int(repair.get("response_chars") or 0),
                    "response_sha256": str(repair.get("response_sha256") or ""),
                }
            states[batch_id] = state
        except Exception as exc:  # each explicit submission makes exactly one request
            if isinstance(exc, _BatchSchemaMismatch):
                category = "schema_mismatch"
            elif "响应字符数" in str(exc) and "响应摘要" in str(exc):
                category = "invalid_json"
            elif _is_transient_config_generation_error(exc):
                category = "transient_network"
            else:
                category = "model_request"
            states[batch_id] = {
                "batch_id": batch_id,
                "question_ids": question_ids,
                "status": "failed",
                "category": category,
                "error": _safe_batch_failure_message(exc),
            }
        _attach_batch_generation_meta(merged, batches, states, batch_size)
        if checkpoint:
            checkpoint(copy.deepcopy(merged))

    _apply_local_question_facts(merged, question_blocks)
    normalize_generated_config_schema(merged)
    _attach_batch_generation_meta(merged, batches, states, batch_size)
    failed = failed_grading_config_batches(merged)
    meta = merged.setdefault("meta", {})
    if failed:
        meta["score_allocation_mode"] = "pending_failed_batches"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = False
        meta["score_allocation_failed"] = False
        meta.pop("score_allocation_error", None)
        meta.pop("score_allocation_failure_category", None)
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    merged = _score_completed_batch_draft_once(
        merged,
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        checkpoint=checkpoint,
    )
    if report:
        report(0.92, "分批生成完成", f"{len(batches)} 个批次和整卷 AI 统一配分均已完成。")
    return merged


def _score_completed_batch_draft_once(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Allocate scores once after every generation batch has succeeded."""
    if failed_grading_config_batches(existing_payload):
        raise ValueError("仍有失败批次，不能进行整卷 AI 统一配分。")
    payload = copy.deepcopy(existing_payload)
    meta = payload.setdefault("meta", {})
    if (
        isinstance(meta, dict)
        and bool(meta.get("score_allocation_ai_success"))
        and not bool(meta.get("score_allocation_pending"))
    ):
        return payload
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["score_allocation_mode"] = "dedicated_ai_scoring"
    meta["score_allocation_ai_success"] = False
    meta["score_allocation_pending"] = True
    meta["score_allocation_failed"] = False
    meta.pop("score_allocation_error", None)
    meta.pop("score_allocation_failure_category", None)
    if checkpoint:
        checkpoint(copy.deepcopy(payload))

    structure_summary = _score_allocation_structure_summary(payload)
    question_ids = [str(item.get("question_id") or "") for item in structure_summary]
    prompt = _build_score_allocation_prompt(
        structure_summary,
        doc_text,
        include_document_text=not (
            bool(q_images)
            or any(
                str(block.get("semantic_source") or "").strip() == "images"
                for block in question_blocks
            )
        ),
    )
    prompt = (
        f"SCORE_QUESTION_IDS_JSON={json.dumps(question_ids, ensure_ascii=False)}\n"
        + prompt
    )
    if report:
        report(
            0.88,
            "AI 统一配分",
            f"正在根据 {len(question_ids)} 道题的完整评分步骤统一配置 100 分。",
        )
    score_repair: dict[str, Any] | None = None
    try:
        score_data = llm_client.json_from_text_once(
            prompt,
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
        )
        score_meta = score_data.get("meta") if isinstance(score_data, dict) else None
        raw_score_repair = (
            score_meta.get("local_json_repair")
            if isinstance(score_meta, dict)
            else None
        )
        if isinstance(raw_score_repair, dict):
            score_repair = {
                "repaired": bool(raw_score_repair.get("repaired")),
                "operations": [
                    str(item) for item in raw_score_repair.get("operations") or []
                ],
                "response_chars": int(raw_score_repair.get("response_chars") or 0),
                "response_sha256": str(
                    raw_score_repair.get("response_sha256") or ""
                ),
            }
        _validate_exact_score_allocation_payload(score_data, structure_summary)
        _apply_score_allocation(payload, score_data)
    except Exception as exc:
        meta["score_allocation_failed"] = True
        meta["score_allocation_failure_category"] = (
            "transient_network"
            if _is_transient_config_generation_error(exc)
            else "model_request"
        )
        meta["score_allocation_error"] = _safe_score_allocation_failure_message(exc)
        if report:
            report(
                0.91,
                "AI 统一配分失败",
                "评分标准批次已保存在本机；没有自动重试，也没有使用本地分值替代。",
            )
        if checkpoint:
            checkpoint(copy.deepcopy(payload))
        return payload
    payload = finalize_completed_grading_config_draft(
        payload,
        question_blocks,
        q_images=q_images,
    )
    meta = payload.setdefault("meta", {})
    meta["score_allocation_mode"] = "dedicated_ai_scoring"
    meta["score_allocation_ai_success"] = True
    meta["score_allocation_pending"] = False
    meta["score_allocation_failed"] = False
    meta.pop("score_allocation_error", None)
    meta.pop("score_allocation_failure_category", None)
    if score_repair is not None:
        meta["score_allocation_local_json_repair"] = score_repair
    else:
        meta.pop("score_allocation_local_json_repair", None)
    if checkpoint:
        checkpoint(copy.deepcopy(payload))
    return payload


def _safe_score_allocation_failure_message(exc: Exception) -> str:
    status_code = getattr(exc, "status_code", None)
    if isinstance(status_code, int):
        return f"AI 统一配分失败（HTTP {status_code}），未自动重试。"
    return "AI 统一配分或本地校验失败，未自动重试。"


def _validate_exact_score_allocation_payload(
    score_data: dict[str, Any],
    structure_summary: list[dict[str, Any]],
) -> None:
    if not isinstance(score_data, dict):
        raise ValueError("AI 统一配分结果顶层不是 JSON 对象。")
    raw_scores = score_data.get("question_scores")
    if not isinstance(raw_scores, list):
        raise ValueError("AI 统一配分缺少 question_scores。")
    expected_ids = [str(item.get("question_id") or "") for item in structure_summary]
    actual_ids = [
        str(item.get("question_id") or "") if isinstance(item, dict) else ""
        for item in raw_scores
    ]
    if actual_ids != expected_ids or len(set(actual_ids)) != len(actual_ids):
        raise ValueError("AI 统一配分的题号缺失、重复、越界或顺序不一致。")

    total_score = 0
    objective_scores: dict[str, int] = {}
    for expected, actual in zip(structure_summary, raw_scores):
        if not isinstance(actual, dict):
            raise ValueError("AI 统一配分的题目结构无效。")
        question_score = _strict_positive_score(actual.get("max_score"))
        if question_score > MAX_QUESTION_SCORE:
            raise ValueError("AI 统一配分存在超过单题上限的分值。")
        total_score += question_score
        question_type = str(expected.get("question_type") or "").strip()
        if question_type in {"choice", "fill_blank", "judgement", "true_false"}:
            previous = objective_scores.setdefault(question_type, question_score)
            if previous != question_score:
                raise ValueError("AI 统一配分未保持同类型客观题同分。")

        expected_parts = expected.get("parts")
        actual_parts = actual.get("parts")
        if not isinstance(expected_parts, list) or not isinstance(actual_parts, list):
            raise ValueError("AI 统一配分的分问结构无效。")
        expected_part_ids = [
            str(item.get("part_id") or "") if isinstance(item, dict) else ""
            for item in expected_parts
        ]
        actual_part_ids = [
            str(item.get("part_id") or "") if isinstance(item, dict) else ""
            for item in actual_parts
        ]
        if actual_part_ids != expected_part_ids:
            raise ValueError("AI 统一配分改变了分问结构。")

        part_total = 0
        for expected_part, actual_part in zip(expected_parts, actual_parts):
            if not isinstance(expected_part, dict) or not isinstance(actual_part, dict):
                raise ValueError("AI 统一配分的分问结构无效。")
            part_score = _strict_positive_score(actual_part.get("part_score"))
            part_total += part_score
            expected_steps = expected_part.get("steps")
            actual_steps = actual_part.get("steps")
            if not isinstance(expected_steps, list) or not isinstance(actual_steps, list):
                raise ValueError("AI 统一配分的步骤结构无效。")
            expected_step_ids = [
                str(item.get("step_id") or "") if isinstance(item, dict) else ""
                for item in expected_steps
            ]
            actual_step_ids = [
                str(item.get("step_id") or "") if isinstance(item, dict) else ""
                for item in actual_steps
            ]
            if actual_step_ids != expected_step_ids:
                raise ValueError("AI 统一配分改变了评分步骤结构。")
            step_total = sum(
                _strict_positive_score(item.get("step_score"))
                for item in actual_steps
                if isinstance(item, dict)
            )
            if step_total != part_score:
                raise ValueError("AI 统一配分的步骤分之和不等于分问分。")
        if part_total != question_score:
            raise ValueError("AI 统一配分的分问分之和不等于题目分。")
    if total_score != 100:
        raise ValueError("AI 统一配分总分不是 100。")


def _strict_positive_score(value: Any) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError("AI 统一配分包含非整数分值。")
    score = int(value)
    if float(value) != float(score) or score <= 0:
        raise ValueError("AI 统一配分包含非正整数分值。")
    return score


def _score_allocation_structure_summary(
    payload: dict[str, Any],
) -> list[dict[str, Any]]:
    rubric = payload.get("rubric") if isinstance(payload, dict) else None
    questions = rubric.get("questions") if isinstance(rubric, dict) else None
    summary: list[dict[str, Any]] = []
    for question in questions or []:
        if not isinstance(question, dict):
            continue
        parts: list[dict[str, Any]] = []
        for part in question.get("parts") or []:
            if not isinstance(part, dict):
                continue
            steps = [
                {
                    "step_id": str(step.get("step_id") or ""),
                    "core_goal": str(step.get("core_goal") or ""),
                    "required_elements": [
                        str(item) for item in step.get("required_elements") or []
                    ],
                }
                for step in part.get("steps") or []
                if isinstance(step, dict)
            ]
            parts.append(
                {
                    "part_id": str(part.get("part_id") or ""),
                    "response_mode": str(part.get("response_mode") or ""),
                    "steps": steps,
                }
            )
        summary.append(
            {
                "question_id": str(question.get("question_id") or ""),
                "question_type": str(question.get("question_type") or ""),
                "knowledge_name": str(question.get("knowledge_name") or ""),
                "parts": parts,
            }
        )
    return summary


def finalize_completed_grading_config_draft(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    *,
    q_images: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Finish an all-succeeded checkpoint locally without another model request."""
    payload = copy.deepcopy(existing_payload)
    if failed_grading_config_batches(payload):
        raise ValueError("仍有失败批次，不能完成本地发布。")
    _validate_unique_batch_question_ids(question_blocks)
    _apply_local_question_facts(payload, question_blocks)
    normalize_generated_config_schema(payload)
    expected_ids = [str(block.get("question_id") or "").strip() for block in question_blocks]
    _validate_exact_batch_payload(payload, expected_ids)
    force_payload_total_score(payload, target_total=100.0)
    meta = payload.setdefault("meta", {})
    meta["score_allocation_mode"] = "local_normalization"
    meta["score_allocation_ai_success"] = False
    meta["score_allocation_pending"] = False
    _attach_reference_answer_images(payload, q_images)
    refresh_generated_config_quality_warnings(payload)
    validate_generated_config(payload)
    return payload


def _validate_unique_batch_question_ids(question_blocks: list[dict[str, Any]]) -> None:
    ids = [str(block.get("question_id") or "").strip() for block in question_blocks]
    if any(not qid for qid in ids) or len(set(ids)) != len(ids):
        raise ValueError("已确认题目必须包含唯一且非空的题号。")


def _existing_batch_states(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    meta = payload.get("meta") if isinstance(payload, dict) else None
    raw = meta.get("batches") if isinstance(meta, dict) else None
    return {
        str(item.get("batch_id")): copy.deepcopy(item)
        for item in raw or []
        if isinstance(item, dict) and str(item.get("batch_id") or "").strip()
    }


def _attach_batch_generation_meta(
    payload: dict[str, Any],
    batches: list[dict[str, Any]],
    states: dict[str, dict[str, Any]],
    batch_size: int,
) -> None:
    ordered_states: list[dict[str, Any]] = []
    for batch in batches:
        batch_id = str(batch["batch_id"])
        ordered_states.append(
            copy.deepcopy(
                states.get(batch_id)
                or {
                    "batch_id": batch_id,
                    "question_ids": list(batch["question_ids"]),
                    "status": "pending",
                }
            )
        )
    failures = [
        {
            "batch_id": str(item["batch_id"]),
            "question_ids": list(item["question_ids"]),
            "category": str(item.get("category") or "failed"),
            "error": str(item.get("error") or "批次生成失败"),
        }
        for item in ordered_states
        if item.get("status") in {"failed", "pending"}
    ]
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["generation_mode"] = "batched"
    meta["batch_size"] = int(batch_size)
    meta["batch_count"] = len(batches)
    meta["batches"] = ordered_states
    meta["failed_batches"] = failures
    meta["failed_question_ids"] = [
        qid for item in failures for qid in item["question_ids"]
    ]


def _validate_exact_batch_payload(
    payload: dict[str, Any],
    question_ids: list[str],
    *,
    require_canonical_answer: bool = True,
) -> None:
    if not isinstance(payload, dict):
        raise _BatchSchemaMismatch("批次结果顶层不是 JSON 对象。")
    rubric = payload.get("rubric")
    answer_key = payload.get("answer_key")
    questions = rubric.get("questions") if isinstance(rubric, dict) else None
    answers = answer_key.get("questions") if isinstance(answer_key, dict) else None
    if not isinstance(questions, list) or not isinstance(answers, list):
        raise _BatchSchemaMismatch("批次结果缺少 rubric.questions 或 answer_key.questions。")
    rubric_ids = [
        str(item.get("question_id") or "").strip()
        for item in questions if isinstance(item, dict)
    ]
    answer_ids = [
        str(item.get("question_id") or "").strip()
        for item in answers if isinstance(item, dict)
    ]
    if (
        len(rubric_ids) != len(questions)
        or len(answer_ids) != len(answers)
        or rubric_ids != question_ids
        or answer_ids != question_ids
        or len(set(rubric_ids)) != len(rubric_ids)
    ):
        raise _BatchSchemaMismatch("批次结果题号缺失、重复、越界或顺序不一致。")
    for question in questions:
        if not isinstance(question.get("parts"), list) or not question["parts"]:
            raise _BatchSchemaMismatch("批次评分结构不完整。")
    for answer in answers:
        if require_canonical_answer and "canonical_answer" not in answer:
            raise _BatchSchemaMismatch("批次答案结构不完整。")


def _safe_batch_failure_message(exc: Exception) -> str:
    message = str(exc or "").strip()
    if isinstance(exc, _BatchSchemaMismatch):
        return message[:300]
    if "响应字符数" in message and "响应摘要" in message:
        return message[:500]
    status_code = getattr(exc, "status_code", None)
    if isinstance(status_code, int):
        return f"模型请求失败（HTTP {status_code}），未自动重试。"
    return "模型请求或本地解析失败，未自动重试。"


def _build_config_batch_request(
    blocks: list[dict[str, Any]],
    doc_text: str,
    q_images: dict[str, Any] | None,
) -> tuple[str, list[bytes]]:
    import base64

    question_ids = [str(block.get("question_id") or "").strip() for block in blocks]
    contexts: list[dict[str, Any]] = []
    image_blobs: list[bytes] = []
    image_map: list[str] = []
    for block in blocks:
        qid = str(block.get("question_id") or "").strip()
        image_semantic_source = str(block.get("semantic_source") or "").strip() == "images"
        contexts.append(
            {
                "question_id": qid,
                "question_type": str(block.get("question_type") or ""),
                "question_text": "" if image_semantic_source else str(
                    block.get("text") or block.get("question_text") or ""
                ),
                "answer_text": "" if image_semantic_source else str(
                    block.get("answer_text") or block.get("canonical_answer") or ""
                ),
                "analysis": "" if image_semantic_source else str(block.get("analysis") or ""),
            }
        )
        image_data = q_images.get(qid) if isinstance(q_images, dict) else None
        if isinstance(image_data, dict) and image_data.get("question"):
            image_blobs.append(base64.b64decode(str(image_data["question"]), validate=True))
            image_map.append(f"图片{len(image_blobs)}={qid}题目")
            if image_data.get("answer"):
                image_blobs.append(base64.b64decode(str(image_data["answer"]), validate=True))
                image_map.append(f"图片{len(image_blobs)}={qid}答案解析")
        elif not image_semantic_source:
            for blob in _word_block_image_blobs(block, limit=4):
                image_blobs.append(blob)
                image_map.append(f"图片{len(image_blobs)}={qid}内嵌图")
    fallback = str(doc_text or "")[:4000] if any(
        str(block.get("semantic_source") or "").strip() != "images"
        and not item["answer_text"] and not item["analysis"]
        for block, item in zip(blocks, contexts)
    ) else ""
    prompt = (
        "你正在为一个小批次的中学数学题生成可执行评分标准。仅返回一个完整、严格的 JSON 对象。\n"
        f"BATCH_QUESTION_IDS_JSON={json.dumps(question_ids, ensure_ascii=False)}\n"
        "必须完整且仅返回上述题号，并在 rubric.questions 与 answer_key.questions 中按同一顺序各出现一次；"
        "不得遗漏、重复或增加题号。所有 max_score、part_score、step_score 暂设为1，最终总分由本地程序分配。\n"
        "每题必须保留教师确认题型，输出具体 knowledge_id、knowledge_name、parts、steps；"
        "每个 part 必须有 response_mode，每个 step 必须有可核验的 core_goal 和 required_elements。"
        "选择/普通填空只核对最终答案；过程题保留必要过程；作图题输出 visual_requirements。\n"
        "只允许字段 rubric、answer_key、meta 及其既有评分结构；不要 Markdown、解释或续写建议。\n"
        f"图片顺序：{'; '.join(image_map) if image_map else '无'}\n"
        f"本批次结构化内容：{json.dumps(contexts, ensure_ascii=False)}\n"
        f"必要时参考的有限原文：{fallback}"
    )
    return prompt, image_blobs


def failed_grading_config_question_ids(payload: dict[str, Any]) -> list[str]:
    meta = payload.get("meta") if isinstance(payload, dict) else {}
    if not isinstance(meta, dict):
        return []
    raw_ids = meta.get("failed_question_ids")
    if isinstance(raw_ids, list):
        return list(dict.fromkeys(str(qid).strip() for qid in raw_ids if str(qid).strip()))
    failures = meta.get("failed_questions")
    if not isinstance(failures, list):
        return []
    return list(
        dict.fromkeys(
            str(failure.get("question_id") or "").strip()
            for failure in failures
            if isinstance(failure, dict) and str(failure.get("question_id") or "").strip()
        )
    )


def _replace_retry_question_payloads(
    existing_payload: dict[str, Any],
    retry_payload: dict[str, Any],
    successful_qids: set[str],
    question_blocks: list[dict[str, Any]],
) -> None:
    existing_rubric = existing_payload.setdefault("rubric", {})
    existing_answer_key = existing_payload.setdefault("answer_key", {})
    retry_rubric = retry_payload.get("rubric", {})
    retry_answer_key = retry_payload.get("answer_key", {})
    existing_questions = existing_rubric.setdefault("questions", []) if isinstance(existing_rubric, dict) else []
    existing_answers = existing_answer_key.setdefault("questions", []) if isinstance(existing_answer_key, dict) else []
    retry_questions = retry_rubric.get("questions", []) if isinstance(retry_rubric, dict) else []
    retry_answers = retry_answer_key.get("questions", []) if isinstance(retry_answer_key, dict) else []
    if not isinstance(existing_questions, list) or not isinstance(existing_answers, list):
        return

    retry_question_map = {
        str(question.get("question_id") or "").strip(): question
        for question in retry_questions
        if isinstance(question, dict)
    } if isinstance(retry_questions, list) else {}
    retry_answer_map = {
        str(answer.get("question_id") or "").strip(): answer
        for answer in retry_answers
        if isinstance(answer, dict)
    } if isinstance(retry_answers, list) else {}

    existing_questions[:] = [
        question
        for question in existing_questions
        if not isinstance(question, dict) or str(question.get("question_id") or "").strip() not in successful_qids
    ]
    existing_answers[:] = [
        answer
        for answer in existing_answers
        if not isinstance(answer, dict) or str(answer.get("question_id") or "").strip() not in successful_qids
    ]
    for qid in successful_qids:
        if qid in retry_question_map:
            existing_questions.append(copy.deepcopy(retry_question_map[qid]))
        if qid in retry_answer_map:
            existing_answers.append(copy.deepcopy(retry_answer_map[qid]))

    qid_order = {
        str(block.get("question_id") or "").strip(): index
        for index, block in enumerate(question_blocks)
        if str(block.get("question_id") or "").strip()
    }
    existing_questions.sort(
        key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000)
        if isinstance(item, dict) else 10_001
    )
    existing_answers.sort(
        key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000)
        if isinstance(item, dict) else 10_001
    )


def _clear_resolved_generation_warnings(payload: dict[str, Any], resolved_qids: set[str]) -> None:
    meta = payload.get("meta") if isinstance(payload, dict) else {}
    warnings = meta.get("warnings") if isinstance(meta, dict) else None
    if not isinstance(warnings, list):
        return
    generation_markers = (
        "parallel generation failed",
        "placeholder added",
        "unmergeable single-question schema",
    )
    warnings[:] = [
        warning
        for warning in warnings
        if not (
            any(marker in str(warning) for marker in generation_markers)
            and any(qid in str(warning) for qid in resolved_qids)
        )
        and str(warning) != "parallel generation completed with failed question blocks; score allocation paused until retry."
    ]


def retry_failed_grading_config_questions(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    retry_question_ids: Sequence[str] | None = None,
) -> dict[str, Any]:
    all_failed_qids = failed_grading_config_question_ids(existing_payload)
    if not all_failed_qids:
        return copy.deepcopy(existing_payload)

    if retry_question_ids is None:
        failed_qids = all_failed_qids
    else:
        failed_qids = list(
            dict.fromkeys(str(qid).strip() for qid in retry_question_ids if str(qid).strip())
        )
        if not failed_qids:
            raise ValueError("至少选择一道失败题目进行重试。")
        unknown_qids = [qid for qid in failed_qids if qid not in all_failed_qids]
        if unknown_qids:
            raise ValueError("只能重试当前失败题目：" + ", ".join(unknown_qids))

    locked_blocks = [{**block, "question_type_confirmed": True} for block in question_blocks]
    block_map = {
        str(block.get("question_id") or "").strip(): block
        for block in locked_blocks
        if str(block.get("question_id") or "").strip()
    }
    retry_blocks = [block_map[qid] for qid in failed_qids if qid in block_map]
    missing_failures = [
        {
            "question_id": qid,
            "attempts": 0,
            "category": "missing_retry_input",
            "error": "Original confirmed question block is no longer available",
        }
        for qid in failed_qids
        if qid not in block_map
    ]
    if not retry_blocks:
        raise ValueError("失败题目的原始确认数据已丢失，请重新拆分试卷后生成。")

    _validate_image_semantic_inputs(retry_blocks, q_images)
    results, failures, attempt_counts, worker_count, image_semantic_mode = _generate_question_block_results(
        retry_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )
    retry_payload = _merge_single_question_payloads(results, retry_blocks)
    failures = _append_unmergeable_question_failures(retry_payload, retry_blocks, failures, attempt_counts)
    failures.extend(missing_failures)
    selected_failure_map = {
        str(failure.get("question_id") or "").strip(): failure
        for failure in failures
        if str(failure.get("question_id") or "").strip()
    }
    existing_meta = existing_payload.get("meta") if isinstance(existing_payload, dict) else {}
    existing_failures = existing_meta.get("failed_questions") if isinstance(existing_meta, dict) else []
    existing_failure_map = {
        str(failure.get("question_id") or "").strip(): copy.deepcopy(failure)
        for failure in existing_failures
        if isinstance(failure, dict) and str(failure.get("question_id") or "").strip()
    } if isinstance(existing_failures, list) else {}
    failures = []
    for qid in all_failed_qids:
        if qid in selected_failure_map:
            failures.append(selected_failure_map[qid])
        elif qid not in failed_qids:
            failures.append(
                existing_failure_map.get(qid)
                or {
                    "question_id": qid,
                    "attempts": 0,
                    "category": "retry_pending",
                    "error": "Question has not been selected for retry",
                }
            )
    remaining_failed_qids = {
        str(failure.get("question_id") or "").strip()
        for failure in failures
        if str(failure.get("question_id") or "").strip()
    }
    successful_qids = set(failed_qids) - remaining_failed_qids

    merged = copy.deepcopy(existing_payload)
    _replace_retry_question_payloads(merged, retry_payload, successful_qids, locked_blocks)
    _clear_resolved_generation_warnings(merged, successful_qids)
    previous_attempts = merged.get("meta", {}).get("single_question_attempts", {})
    all_attempts = dict(previous_attempts) if isinstance(previous_attempts, dict) else {}
    all_attempts.update(attempt_counts)
    _attach_parallel_generation_meta(merged, locked_blocks, failures, worker_count, all_attempts)
    _ensure_question_blocks_covered(merged, locked_blocks)
    _apply_local_question_facts(merged, locked_blocks)
    normalize_generated_config_schema(merged)

    if failures:
        meta = merged.setdefault("meta", {})
        meta["score_allocation_mode"] = "pending_failed_questions"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = True
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    return retry_grading_config_score_allocation(
        merged,
        locked_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        include_document_text=not image_semantic_mode,
    )


def retry_grading_config_score_allocation(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    include_document_text: bool | None = None,
) -> dict[str, Any]:
    if failed_grading_config_question_ids(existing_payload):
        raise ValueError("仍有失败题目，需先重试失败题目，再进行整体赋分。")
    merged = copy.deepcopy(existing_payload)
    if include_document_text is None:
        include_document_text = not (
            bool(q_images)
            or any(str(block.get("semantic_source") or "").strip() == "images" for block in question_blocks)
        )
    _assign_scores_to_merged_config(
        merged,
        doc_text,
        llm_client,
        model_name,
        report=report,
        include_document_text=include_document_text,
    )
    _ensure_question_blocks_covered(merged, question_blocks)
    _apply_local_question_facts(merged, question_blocks)
    normalize_generated_config_schema(merged)
    meta = merged.setdefault("meta", {})
    meta["score_allocation_mode"] = (
        "dedicated_ai_scoring" if meta.get("score_allocation_ai_success") else "local_score_fallback"
    )
    meta["score_allocation_pending"] = False
    force_payload_total_score(merged, target_total=100.0)
    _attach_reference_answer_images(merged, q_images)
    refresh_generated_config_quality_warnings(merged)
    return merged


def _split_local_question_answer_text(doc_text: str) -> tuple[str, str]:
    text = str(doc_text or "")
    try:
        from question_bank.importers.batch_importer import _split_answer_text

        question_text, answer_text = _split_answer_text(text)
        return str(question_text or ""), str(answer_text or "")
    except Exception:
        lines = text.splitlines()
        for index, line in enumerate(lines):
            if _looks_like_answer_section_heading(line):
                return "\n".join(lines[:index]).strip(), "\n".join(lines[index + 1 :]).strip()
    return text, ""


def _question_type_hints_from_section_headings(doc_text: str) -> dict[str, str]:
    question_text, _ = _split_local_question_answer_text(doc_text)
    current_type = ""
    hints: dict[str, str] = {}
    for line in question_text.splitlines():
        value = str(line or "").strip()
        if "选择题" in value:
            current_type = "choice"
            continue
        if "填空题" in value:
            current_type = "fill_blank"
            continue
        if "证明题" in value:
            current_type = "proof"
            continue
        if "解答题" in value:
            current_type = "comprehensive"
            continue
        number = _extract_question_marker_number(value)
        if number is not None and current_type:
            hints[str(number)] = current_type
    return hints


def _parse_inline_answer_blocks(doc_text: str) -> list[dict[str, Any]] | None:
    """解析"题干 + 【答案】 + 【解析】逐题穿插"格式的试卷（纯本地，不调 AI）。

    返回 None 表示该文本不是内联格式（无【答案】/【解析】标记），交回上层兜底逻辑。
    通过"题号必须单调递增"截断文件后半段重复的答案详解区，避免重复拆题。
    """
    full_text = str(doc_text or "")
    text, answer_section = _split_local_question_answer_text(full_text)
    if ("【答案】" not in text) and ("【解析】" not in text):
        return None

    lines = text.splitlines()
    # 1. 找题号边界：行首 数字. / 数字．，题号单调递增；回退即停止
    markers: list[tuple[int, int]] = []  # (line_index, number)
    last_number = 0
    for idx, line in enumerate(lines):
        stripped = line.strip()
        m = re.match(r"^(\d{1,2})\s*[.．]", stripped)
        if not m:
            continue
        number = int(m.group(1))
        if not (1 <= number <= 99):
            continue
        if markers and number <= last_number:
            break  # 题号回退 → 进入答案详解重复区，停止切题
        markers.append((idx, number))
        last_number = number

    if not markers:
        return None

    answer_section_for_choice = answer_section or _local_answer_section_text(text)
    choice_answers = _extract_choice_answer_sequence(answer_section_for_choice)
    section_hints = _question_type_hints_from_section_headings(text)

    blocks: list[dict[str, Any]] = []
    for pos, (start, number) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        segment_lines = lines[start:end]
        parsed = _parse_inline_segment(
            number,
            segment_lines,
            choice_answers,
            section_type=section_hints.get(str(number), ""),
        )
        if parsed:
            blocks.append(parsed)

    return blocks or None


def _parse_inline_segment(
    number: int,
    segment_lines: list[str],
    choice_answers: dict[str, str],
    *,
    section_type: str = "",
) -> dict[str, Any] | None:
    """把单道题的文本段拆成 题干 / 答案 / 解析。"""
    # 分类收集：题干 / 【答案】 / 【解析】(含【点睛】)
    bucket = "stem"  # stem | answer | analysis
    stem_lines: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []

    for raw in segment_lines:
        line = str(raw or "").strip()
        if not line:
            continue
        if line.startswith("[公式:") or line.startswith("[公式："):
            continue  # 忽略公式占位行
        if "【答案】" in line:
            bucket = "answer"
            after = line.split("【答案】", 1)[1].strip()
            if after:
                answer_lines.append(after)
            continue
        if ("【解析】" in line) or ("【点睛】" in line):
            bucket = "analysis"
            marker = "【解析】" if "【解析】" in line else "【点睛】"
            after = line.split(marker, 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if line.startswith("【") and "】" in line:
            # 其他【…】标记（如【难度】），归入解析区尾部，避免污染题干
            bucket = "analysis"
            after = line.split("】", 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if bucket == "stem":
            stem_lines.append(line)
        elif bucket == "answer":
            answer_lines.append(line)
        else:
            analysis_lines.append(line)

    question_text = _strip_leading_question_number(number, "\n".join(stem_lines).strip())
    answer_raw = "\n".join(answer_lines).strip()
    analysis_raw = "\n".join(analysis_lines).strip()

    if not question_text and not answer_raw and not analysis_raw:
        return None

    num_str = str(number)
    qtype = _infer_local_question_type(
        question_text,
        answer_raw,
        num_str,
        section_type=section_type,
    )
    canonical = _extract_canonical_answer_for_local_question(
        number=num_str,
        qtype=qtype,
        question_text=question_text,
        answer_text=answer_raw,
        choice_answers=choice_answers,
    )
    accepted = _local_accepted_forms(canonical, qtype)
    has_answer = bool(answer_raw or canonical)
    return {
        "question_id": f"Q{number}",
        "text": question_text,
        "question_text": question_text,
        "answer_text": answer_raw,
        "analysis": analysis_raw,
        "question_type": qtype,
        "canonical_answer": canonical,
        "accepted_forms": accepted,
        "local_answer_trusted": has_answer,
        "needs_review": (not bool(canonical)) if qtype in {"choice", "fill_blank"} else False,
    }


def _strip_leading_question_number(number: int, text: str) -> str:
    """去掉题干开头的"N." / "N．" / "N.(本小题X分)"等题号前缀。"""
    value = str(text or "").lstrip()
    value = re.sub(rf"^{number}\s*[.．]\s*", "", value, count=1)
    return value.strip()


def _extract_local_question_blocks(doc_text: str) -> list[dict[str, Any]]:
    normalized_doc_text = "\n".join(
        _normalize_inline_main_question_lines(doc_text)
    )
    inline_blocks = _parse_inline_answer_blocks(normalized_doc_text)
    if inline_blocks:
        return inline_blocks

    parsed_questions: list[Any] = []
    try:
        from question_bank.importers.batch_importer import parse_paper_text

        parsed = parse_paper_text(
            normalized_doc_text,
            source_file="grading_config.docx",
            page_range="document",
        )
        parsed_questions = list(parsed.questions)
    except Exception:
        parsed_questions = []

    # 第二道防线：题号必须单调递增；一旦回退（如卷末答案详解区从1重列）即截断，避免重复拆题。
    if parsed_questions:
        truncated: list[Any] = []
        last_number = 0
        for item in parsed_questions:
            number_str = str(getattr(item, "question_number", "") or "").strip()
            if not number_str.isdigit():
                truncated.append(item)
                continue
            number = int(number_str)
            if truncated and number <= last_number:
                break
            truncated.append(item)
            last_number = number
        parsed_questions = truncated

    answer_section = _local_answer_section_text(normalized_doc_text)
    answer_blocks = _local_answer_blocks(answer_section)
    choice_answers = _extract_choice_answer_sequence(answer_section)
    section_hints = _question_type_hints_from_section_headings(normalized_doc_text)

    blocks: list[dict[str, Any]] = []
    if parsed_questions:
        for item in parsed_questions:
            number = str(getattr(item, "question_number", "") or "").strip()
            if not number.isdigit():
                continue
            qid = f"Q{int(number)}"
            question_text = str(getattr(item, "question_text", "") or "").strip()
            answer_from_map = str(answer_blocks.get(number, "") or "").strip()
            parsed_answer = str(getattr(item, "answer_text", "") or "").strip()
            raw_answer = str(answer_from_map or parsed_answer).strip()
            qtype = _infer_local_question_type(
                question_text,
                raw_answer,
                number,
                section_type=section_hints.get(str(int(number)), ""),
            )
            canonical = _extract_canonical_answer_for_local_question(
                number=number,
                qtype=qtype,
                question_text=question_text,
                answer_text=raw_answer or answer_blocks.get(number, ""),
                choice_answers=choice_answers,
            )
            accepted = _local_accepted_forms(canonical, qtype)
            blocks.append(
                {
                    "question_id": qid,
                    "text": question_text,
                    "question_text": question_text,
                    "answer_text": raw_answer,
                    "question_type": qtype,
                    "canonical_answer": canonical,
                    "accepted_forms": accepted,
                    "local_answer_trusted": bool(answer_from_map)
                    or any(token in raw_answer for token in ["故答案为", "故选", "答案为"]),
                    "needs_review": not bool(canonical) if qtype in {"choice", "fill_blank"} else False,
                }
            )
    else:
        for block in _split_doc_text_into_question_blocks(normalized_doc_text):
            qid = str(block.get("question_id") or "")
            number = qid.removeprefix("Q")
            question_text = str(block.get("text") or "")
            raw_answer = answer_blocks.get(number, "")
            qtype = _infer_local_question_type(
                question_text,
                raw_answer,
                number,
                section_type=section_hints.get(str(int(number)), "") if number.isdigit() else "",
            )
            canonical = _extract_canonical_answer_for_local_question(
                number=number,
                qtype=qtype,
                question_text=question_text,
                answer_text=raw_answer,
                choice_answers=choice_answers,
            )
            block.update(
                {
                    "question_text": question_text,
                    "answer_text": raw_answer,
                    "question_type": qtype,
                    "canonical_answer": canonical,
                    "accepted_forms": _local_accepted_forms(canonical, qtype),
                    "local_answer_trusted": bool(raw_answer)
                    or any(token in raw_answer for token in ["故答案为", "故选", "答案为"]),
                    "needs_review": not bool(canonical) if qtype in {"choice", "fill_blank"} else False,
                }
            )
            blocks.append(block)

    return blocks


def _local_answer_section_text(doc_text: str) -> str:
    text = str(doc_text or "")
    try:
        from question_bank.importers.batch_importer import _split_answer_text

        _, answer_text = _split_answer_text(text)
        if str(answer_text or "").strip():
            return str(answer_text or "")
    except Exception:
        pass

    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if _looks_like_answer_section_heading(line):
            return "\n".join(lines[idx + 1 :]).strip()
    return ""


def _local_answer_blocks(answer_section: str) -> dict[str, str]:
    section = str(answer_section or "")
    if not section.strip():
        return {}
    try:
        from question_bank.importers.batch_importer import _split_numbered_blocks

        blocks = _split_numbered_blocks(section)
        result: dict[str, str] = {}
        for block in blocks:
            number = str(getattr(block, "number", "") or "").strip()
            if number.isdigit() and 1 <= int(number) <= 99:
                result.setdefault(str(int(number)), str(getattr(block, "text", "") or "").strip())
        if result:
            return result
    except Exception:
        pass

    lines = section.splitlines()
    markers: list[tuple[int, str]] = []
    last_number = 0
    for idx, line in enumerate(lines):
        if "声明" in line or "菁优网" in line:
            break
        number = _extract_question_marker_number(line)
        if number is None:
            continue
        if markers and number <= last_number:
            break
        markers.append((idx, str(number)))
        last_number = number

    result: dict[str, str] = {}
    for pos, (start, number) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        block = "\n".join(lines[start:end]).strip()
        if block:
            result[number] = block
    return result


def _extract_choice_answer_sequence(answer_section: str) -> dict[str, str]:
    """Extract choice answers (A/B/C/D) from the answer section in order.

    Supports common Chinese exam answer formats:
    - 故选A / 故选：A
    - 选A / 选：A
    - 答案是A / 答案为A / 答：A
    - 答案选A
    - 独立的单字母行（上下文暗示选择题）
    """
    section = str(answer_section or "")
    found: list[str] = []
    # Broad pattern: any of the common phrasings followed by A/B/C/D
    pattern = (
        r"(?:"
        r"故\s*选"
        r"|选择"
        r"|选\s*[:：]"
        r"|答\s*案\s*(?:是|为|选)?\s*[:：]?"
        r"|答\s*[:：]"
        r"|【答案】"
        r")\s*([A-Da-d])"
    )
    for match in re.finditer(pattern, section):
        found.append(match.group(1).upper())
    if not found:
        # Fallback: numbered answers like "1.A" or "1、A" or "(1)A"
        for match in re.finditer(
            r"(?:^|\n)\s*(?:\(\s*)?\d{1,2}\s*[.．、)）]\s*([A-Da-d])(?:\s|$)", section
        ):
            found.append(match.group(1).upper())
    if not found:
        # Last fallback: solo letter on its own line that looks like a choice answer
        for match in re.finditer(r"(?:^|\n)\s*([A-D])\s*(?:\n|$)", section):
            found.append(match.group(1).upper())
    return {str(index): value for index, value in enumerate(found[:20], start=1)}


def _infer_local_question_type(
    question_text: str,
    answer_text: str,
    number: str,  # noqa: ARG001
    *,
    section_type: str = "",
) -> str:
    """Infer question type, preferring explicit paper-section headings."""
    value = str(question_text or "")
    normalized_section_type = str(section_type or "").strip()
    if normalized_section_type in {"choice", "fill_blank", "proof"}:
        return normalized_section_type

    option_labels = {
        match.group(1).upper()
        for match in re.finditer(r"(?m)^\s*([A-Da-d])\s*(?:[.．、)]|\s{2,})", value)
    }
    compact_options = any(
        re.search(r"[Aa][.．、]?\s*.{0,80}[Bb][.．、]?\s*.{0,80}[Cc][.．、]?\s*.{0,80}[Dd]", line)
        for line in value.splitlines()
    )
    if len(option_labels) >= 3 or compact_options:
        return "choice"
    if re.fullmatch(r"\s*[A-Da-d]\s*", str(answer_text or "")):
        return "choice"

    has_blank = bool(
        re.search(
            r"_{2,}|　{1,}|（\s*）|\(\s*\)|\b填空\b",
            value,
        )
    )
    if has_blank:
        return "fill_blank"

    if any(token in value for token in ["作图", "作出", "画出", "保留作图痕迹"]):
        return "comprehensive"
    if any(token in value for token in ["证明", "理由", "说明", "求证", "全等", "证得"]):
        return "proof"
    has_subparts = bool(re.search(r"[（(]\s*[1-9]\s*[）)]", value))
    if has_subparts or any(token in value for token in ["计算", "求", "解答", "解："]):
        return "calculation"
    if normalized_section_type == "comprehensive":
        return "comprehensive"
    return _infer_question_type_from_block_text(question_text)


def _extract_canonical_answer_for_local_question(
    *,
    number: str,
    qtype: str,
    question_text: str,
    answer_text: str,
    choice_answers: dict[str, str],
) -> str:
    if qtype == "choice":
        # First try the pre-extracted sequence (ordered by question number)
        answer = choice_answers.get(str(int(number))) if str(number or "").isdigit() else ""
        if answer:
            return answer
        # Then try in-line answer patterns
        choice_pattern = (
            r"(?:故\s*选|选择|选\s*[:：]|答\s*案\s*(?:是|为|选)?\s*[:：]?|答\s*[:：]|【答案】)\s*([A-Da-d])"
        )
        match = re.search(choice_pattern, str(answer_text or ""))
        if match:
            return match.group(1).upper()
        # Last: bare single letter
        bare = re.fullmatch(r"\s*([A-Da-d])\s*", str(answer_text or ""))
        return bare.group(1).upper() if bare else ""

    if qtype != "fill_blank":
        return ""

    text = str(answer_text or "")
    # Extended patterns for fill-blank answer extraction
    patterns = [
        r"【答案】\s*([^。\n；;]{1,120})",
        r"故答案为\s*[:：]?\s*([^。\n；;]{1,120})",
        r"答案(?:是|为)\s*[:：]?\s*([^。\n；;]{1,120})",
        r"答案\s*[:：]\s*([^。\n；;]{1,120})",
        r"答\s*[:：]\s*([^。\n；;]{1,120})",
        r"故答案为\s*[:：]?\s*([\s\S]{1,80}?)[。．\n]",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            value = _clean_local_answer_text(match.group(1))
            if value:
                return value

    blank_values = re.findall(r"\u3000\s*([^\u3000\n]{1,40}?)\s*\u3000", text)
    for value in blank_values:
        cleaned = _clean_local_answer_text(value)
        if cleaned:
            return cleaned
    return ""


def _clean_local_answer_text(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^[：:，,、;\s]+|[。．，,、;\s]+$", "", text)
    text = text.replace("°", "").strip() if re.fullmatch(r"[\d.]+°", text) else text
    text = re.sub(r"\s+", " ", text).strip()
    if not text or text in {"解", "如下", "见解析"}:
        return ""
    return text[:120]


def _local_accepted_forms(canonical: str, qtype: str) -> list[str]:
    value = str(canonical or "").strip()
    if not value:
        return []
    if qtype == "choice":
        return [value.upper()]
    candidates = [value]
    for chunk in re.split(r"[；;、，,]|\s+或\s+|\bor\b", value):
        cleaned = _clean_local_answer_text(chunk)
        if cleaned:
            candidates.append(cleaned)
    return merge_equivalent_forms([], *candidates, max_forms=16)


def _apply_local_question_facts(payload: dict[str, Any], question_blocks: list[dict[str, Any]]) -> None:
    if not isinstance(payload, dict):
        return
    rubric = payload.setdefault("rubric", {})
    answer_key = payload.setdefault("answer_key", {})
    if not isinstance(rubric, dict) or not isinstance(answer_key, dict):
        return
    questions = rubric.setdefault("questions", [])
    answers = answer_key.setdefault("questions", [])
    if not isinstance(questions, list) or not isinstance(answers, list):
        return

    facts = {
        str(block.get("question_id") or "").strip(): block
        for block in question_blocks
        if str(block.get("question_id") or "").strip()
    }
    answer_map = {
        str(answer.get("question_id") or "").strip(): answer
        for answer in answers
        if isinstance(answer, dict)
    }

    for question in questions:
        if not isinstance(question, dict):
            continue
        qid = str(question.get("question_id") or "").strip()
        fact = facts.get(qid)
        if not fact:
            continue
        local_type = str(fact.get("question_type") or "").strip()
        type_confirmed = bool(fact.get("question_type_confirmed"))
        if local_type in {"choice", "fill_blank"} or (
            type_confirmed and local_type in {"calculation", "proof", "comprehensive"}
        ):
            question["question_type"] = local_type
            question["grading_mode"] = (
                "direct_answer"
                if local_type in {"choice", "fill_blank"}
                else "deductive_obligation"
            )
        if type_confirmed:
            question["question_type_confirmed"] = True
        image_semantic_source = str(fact.get("semantic_source") or "").strip() == "images"
        if (
            not image_semantic_source
            and str(fact.get("question_text") or "").strip()
            and not str(question.get("stem_summary") or "").strip()
        ):
            question["stem_summary"] = str(fact.get("question_text") or "").strip().splitlines()[0][:120]

        answer = answer_map.get(qid)
        if answer is None:
            answer = {"question_id": qid, "canonical_answer": "", "accepted_forms": [], "method_variants": [], "parts": []}
            answers.append(answer)
            answer_map[qid] = answer
        canonical = str(fact.get("canonical_answer") or "").strip()
        accepted = [str(item).strip() for item in fact.get("accepted_forms") or [] if str(item).strip()]
        if canonical and not image_semantic_source:
            current_canonical = str(answer.get("canonical_answer") or "").strip()
            if bool(fact.get("local_answer_trusted")) or not current_canonical:
                answer["canonical_answer"] = canonical
                base_answer = canonical
            else:
                base_answer = current_canonical
            answer["accepted_forms"] = merge_equivalent_forms(answer.get("accepted_forms"), *accepted, base_answer, max_forms=32)
            parts = answer.get("parts")
            if not isinstance(parts, list) or not parts:
                answer["parts"] = [{"part_id": qid, "answer": base_answer, "analysis": str(fact.get("answer_text") or ""), "step_milestones": []}]
            else:
                for part in parts:
                    if isinstance(part, dict) and not str(part.get("answer") or "").strip():
                        part["answer"] = base_answer


def _attach_reference_answer_images(
    payload: dict[str, Any],
    q_images: dict[str, Any] | None,
) -> None:
    """Persist clean PDF answer and question stem crops for image-aware grading."""
    if not q_images:
        return
    answer_key = payload.get("answer_key") if isinstance(payload, dict) else None
    answer_questions = answer_key.get("questions") if isinstance(answer_key, dict) else None
    rubric = payload.get("rubric") if isinstance(payload, dict) else None
    rubric_questions = rubric.get("questions") if isinstance(rubric, dict) else None
    
    answer_map = {
        str(item.get("question_id") or ""): item
        for item in answer_questions
        if isinstance(item, dict)
    } if isinstance(answer_questions, list) else {}
    
    rubric_map = {
        str(item.get("question_id") or ""): item
        for item in rubric_questions
        if isinstance(item, dict)
    } if isinstance(rubric_questions, list) else {}

    for qid, image_data in q_images.items():
        if not isinstance(image_data, dict):
            continue
        answer_image = str(image_data.get("answer") or "").strip()
        answer_item = answer_map.get(str(qid))
        if answer_image and isinstance(answer_item, dict):
            answer_item["answer_image_base64"] = answer_image
            answer_item["answer_image_role"] = "perfect_standard_answer"
            
        question_image = str(image_data.get("question") or "").strip()
        rubric_item = rubric_map.get(str(qid))
        if question_image and isinstance(rubric_item, dict):
            rubric_item["question_image_base64"] = question_image


def _quality_answer_texts(node: Any) -> list[str]:
    if not isinstance(node, dict):
        return []
    values: list[str] = []
    for key in ("answer", "canonical_answer", "standard_answer", "correct_answer"):
        value = str(node.get(key) or "").strip()
        if value:
            values.append(value)
    accepted = node.get("accepted_forms")
    if isinstance(accepted, list):
        values.extend(str(item).strip() for item in accepted if str(item).strip())
    return values


def _looks_like_serialized_answer_list(value: Any) -> bool:
    return bool(re.fullmatch(r"\s*\[[\s\S]*\]\s*", str(value or "")))


def _looks_like_serialized_knowledge_sequence(value: Any) -> bool:
    return bool(
        re.fullmatch(
            r"\s*(?:\[[\s\S]*\]|\([\s\S]*\))\s*",
            str(value or ""),
        )
    )


def _normalize_serialized_answer_list(value: Any) -> str:
    text = str(value or "").strip()
    if not _looks_like_serialized_answer_list(text):
        return text
    try:
        parsed = ast.literal_eval(text)
    except (SyntaxError, ValueError):
        return text
    if not isinstance(parsed, (list, tuple)) or not parsed:
        return text
    values = [str(item).strip() for item in parsed if str(item).strip()]
    return "、".join(values) if values else text


def _looks_like_garbled_generated_text(value: Any) -> bool:
    text = str(value or "")
    return "\ufffd" in text or "锟" in text or "��" in text


def collect_generated_config_quality_warnings(payload: dict[str, Any]) -> list[str]:
    rubric = payload.get("rubric") if isinstance(payload, dict) else None
    answer_key = payload.get("answer_key") if isinstance(payload, dict) else None
    rubric_questions = rubric.get("questions") if isinstance(rubric, dict) else []
    answer_questions = answer_key.get("questions") if isinstance(answer_key, dict) else []
    if not isinstance(rubric_questions, list):
        return ["[质量检查-阻断] rubric.questions 结构无效"]
    answer_map = {
        str(item.get("question_id") or ""): item
        for item in answer_questions
        if isinstance(item, dict)
    } if isinstance(answer_questions, list) else {}

    warnings: list[str] = []
    for question in rubric_questions:
        if not isinstance(question, dict):
            continue
        qid = str(question.get("question_id") or "未知题号")
        qtype = str(question.get("question_type") or "")
        stem = str(question.get("stem_summary") or "").strip()
        knowledge_id = str(question.get("knowledge_id") or "").strip()
        knowledge_name = str(question.get("knowledge_name") or "").strip()
        if not knowledge_name or knowledge_id in {"", "UNKNOWN"}:
            warnings.append(f"[质量检查-提醒] {qid} 缺少明确知识点")
        if knowledge_name and stem and knowledge_name == stem:
            warnings.append(f"[质量检查-提醒] {qid} 知识点疑似直接复制题干")
        if knowledge_name in {"几何综合", "代数综合", "数学综合", "综合应用", "未知知识点"}:
            warnings.append(f"[质量检查-提醒] {qid} 知识点过于宽泛，应写明具体考查概念")

        answer_item = answer_map.get(qid, {})
        answer_image_present = bool(isinstance(answer_item, dict) and answer_item.get("answer_image_base64"))
        answer_texts = _quality_answer_texts(answer_item)
        answer_parts = answer_item.get("parts") if isinstance(answer_item, dict) else []
        if not isinstance(answer_parts, list):
            answer_parts = []

        knowledge_texts: list[Any] = [knowledge_id, knowledge_name]
        knowledge_ids = question.get("knowledge_ids")
        if isinstance(knowledge_ids, list):
            knowledge_texts.extend(knowledge_ids)
        knowledge_points = question.get("knowledge_points")
        if isinstance(knowledge_points, list):
            for point in knowledge_points:
                if isinstance(point, dict):
                    knowledge_texts.extend(
                        (point.get("knowledge_id"), point.get("knowledge_name"))
                    )
                else:
                    knowledge_texts.append(point)

        text_fields: list[Any] = [stem, *knowledge_texts, *answer_texts]
        parts = question.get("parts")
        if not isinstance(parts, list):
            parts = []
        for part in parts:
            if not isinstance(part, dict):
                continue
            for step in part.get("steps") or []:
                if not isinstance(step, dict):
                    continue
                text_fields.append(step.get("core_goal"))
                required = step.get("required_elements")
                if isinstance(required, list):
                    text_fields.extend(required)
        for answer_part in answer_parts:
            text_fields.extend(_quality_answer_texts(answer_part))
        if any(_looks_like_garbled_generated_text(value) for value in text_fields):
            warnings.append(f"[质量检查-阻断] {qid} 的题干、公式、答案或踩分点中存在疑似乱码")
        if (
            any(_looks_like_serialized_answer_list(value) for value in text_fields)
            or any(
                _looks_like_serialized_knowledge_sequence(value)
                for value in knowledge_texts
            )
        ):
            warnings.append(
                f"[质量检查-阻断] {qid} 的知识点、答案或评分点中混入列表字符串或元组字符串"
            )

        if qtype in {"choice", "fill_blank", "judgement", "true_false", "direct_answer"}:
            if not answer_texts and not answer_image_present:
                warnings.append(f"[质量检查-阻断] {qid} 缺少可评分的标准答案")

        for index, part in enumerate(parts):
            if not isinstance(part, dict):
                continue
            mode = str(part.get("response_mode") or "").strip() or _infer_part_response_mode(question, part)
            steps = [step for step in part.get("steps") or [] if isinstance(step, dict)]
            goals = [str(step.get("core_goal") or "").strip() for step in steps]
            generic_goals = {
                "完成必要的推理或计算步骤",
                "合理的推理过程",
                "正确的结论",
            }
            if qtype in {"choice", "fill_blank", "judgement", "true_false", "direct_answer"} and mode == "process_required":
                warnings.append(f"[质量检查-阻断] {qid} 客观题被错误设置为过程评分")
            if mode == "visual_construction":
                visual_requirements = _string_list(part.get("visual_requirements"))
                meaningful_goals = [goal for goal in goals if goal and goal not in generic_goals]
                if not visual_requirements and not meaningful_goals:
                    warnings.append(f"[质量检查-阻断] {qid} 作图题缺少具体作图要求")
            if mode not in {"exact_objective", "short_answer_points", "visual_construction"}:
                if qtype in {"proof", "calculation", "comprehensive"} and goals and all(goal in generic_goals for goal in goals):
                    warnings.append(f"[质量检查-阻断] {qid} 评分点全部为通用描述，无法执行可靠批改")
                continue
            answer_part = answer_parts[index] if index < len(answer_parts) and isinstance(answer_parts[index], dict) else {}
            if not _quality_answer_texts(answer_part) and not answer_texts and not answer_image_present:
                part_id = str(part.get("part_id") or f"第{index + 1}问")
                warnings.append(f"[质量检查-阻断] {qid}/{part_id} 缺少可评分的标准答案或答案图")

    return list(dict.fromkeys(warnings))


def refresh_generated_config_quality_warnings(payload: dict[str, Any]) -> list[str]:
    meta = payload.setdefault("meta", {}) if isinstance(payload, dict) else {}
    if not isinstance(meta, dict):
        return []
    warnings = meta.get("warnings")
    if not isinstance(warnings, list):
        warnings = []
    warnings = [warning for warning in warnings if not str(warning).startswith("[质量检查-")]
    quality_warnings = collect_generated_config_quality_warnings(payload)
    meta["warnings"] = [*warnings, *quality_warnings]
    return quality_warnings


def _restore_precalibration_objective_answers(payload: dict[str, Any], previous_payload: dict[str, Any]) -> None:
    if not isinstance(payload, dict) or not isinstance(previous_payload, dict):
        return
    previous_questions = previous_payload.get("rubric", {}).get("questions", [])
    previous_answers = previous_payload.get("answer_key", {}).get("questions", [])
    current_answers = payload.get("answer_key", {}).get("questions", [])
    if not isinstance(previous_questions, list) or not isinstance(previous_answers, list) or not isinstance(current_answers, list):
        return

    previous_types = {
        str(question.get("question_id") or "").strip(): str(question.get("question_type") or "").strip()
        for question in previous_questions
        if isinstance(question, dict)
    }
    previous_answer_map = {
        str(answer.get("question_id") or "").strip(): answer
        for answer in previous_answers
        if isinstance(answer, dict)
    }
    for answer in current_answers:
        if not isinstance(answer, dict):
            continue
        qid = str(answer.get("question_id") or "").strip()
        if previous_types.get(qid) not in {"choice", "fill_blank"}:
            continue
        previous = previous_answer_map.get(qid)
        if not isinstance(previous, dict):
            continue
        previous_canonical = str(previous.get("canonical_answer") or "").strip()
        if not previous_canonical:
            continue
        answer["canonical_answer"] = previous_canonical
        answer["accepted_forms"] = merge_equivalent_forms(previous.get("accepted_forms"), previous_canonical, max_forms=32)
        previous_parts = previous.get("parts")
        if isinstance(previous_parts, list) and previous_parts:
            answer["parts"] = previous_parts


def _split_doc_text_into_question_blocks(doc_text: str) -> list[dict[str, str]]:
    lines = _normalize_inline_main_question_lines(doc_text)
    markers: list[tuple[int, str]] = []
    last_number = 0
    for idx, line in enumerate(lines):
        if _looks_like_answer_section_heading(line) and markers:
            break
        number = _extract_question_marker_number(line)
        if number is None:
            continue
        if markers and number <= last_number:
            break
        markers.append((idx, f"Q{number}"))
        last_number = number
    blocks: list[dict[str, str]] = []
    for pos, (start, qid) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        text = "\n".join(lines[start:end]).strip()
        if text:
            blocks.append({"question_id": qid, "text": text})
    return blocks


def _normalize_inline_main_question_lines(doc_text: str) -> list[str]:
    lines: list[str] = []
    current_number: int | None = None
    in_answer_section = False
    raw_lines = [raw_line.rstrip() for raw_line in str(doc_text or "").splitlines()]
    for index, line in enumerate(raw_lines):
        if _looks_like_answer_section_heading(line) and current_number is not None:
            in_answer_section = True
        if in_answer_section:
            lines.append(line)
            continue
        leading_number = _extract_question_marker_number(line)
        if leading_number is not None:
            current_number = leading_number
        if current_number is None:
            lines.append(line)
            continue
        segments, current_number = _split_consecutive_inline_main_questions(
            line,
            current_number=current_number,
            following_number=_next_leading_main_question_number(
                raw_lines,
                after_index=index,
            ),
        )
        lines.extend(segments)
    return lines


def _extract_question_marker_number(line: str) -> int | None:
    value = str(line or "").strip()
    if not value:
        return None
    patterns = [
        r"^(?:Q|q)\s*(\d{1,2})(?:\b|[\s:：.．、)])",
        r"^第\s*(\d{1,2})\s*[题題]",
        r"^(\d{1,2})\s*[.．、)]",
    ]
    for pattern in patterns:
        match = re.match(pattern, value)
        if not match:
            continue
        number = int(match.group(1))
        if 1 <= number <= 99:
            return number
    return None

def _looks_like_answer_section_heading(line: str) -> bool:
    value = str(line or "").strip()
    if not value:
        return False
    lower = value.lower().strip(":： ")
    chinese_markers = [
        "\u53c2\u8003\u7b54\u6848",
        "\u8bd5\u9898\u89e3\u6790",
        "\u7b54\u6848\u4e0e\u89e3\u6790",
    ]
    if any(marker in value for marker in chinese_markers):
        return True
    return lower in {"answer", "answers", "solution", "solutions", "answer key"}

def _build_single_question_image_generation_prompt(
    block: dict[str, Any],
    *,
    has_answer_image: bool,
) -> str:
    qid = str(block.get("question_id") or "").strip()
    confirmed_type = str(block.get("question_type") or "").strip()
    image_description = (
        "图片1是题目原图；图片2是标准答案与解析原图。"
        if has_answer_image
        else "图片1是题目原图；未提供标准答案与解析原图。"
    )
    knowledge_options = "、".join(str(item) for item in KNOWLEDGE_POINT_OPTIONS)
    return (
        "你正在为一道中学数学题生成可执行评分标准。仅返回严格 JSON。\n"
        f"题目ID：{qid}\n"
        f"教师已确认题型：{confirmed_type}\n"
        f"{image_description}\n"
        "图片是唯一权威内容来源。不得参考、猜测或恢复任何 PDF 抽取文字。\n"
        "必须严格保持图片中的点名、线段、公式、运算对象、小问结构和作答要求，不得改写成相似题。\n"
        "分值字段暂时全部设为1；后续会统一分配分值。\n"
        "每个 parts 项必须输出 response_mode："
        "exact_objective（选择/普通填空）、short_answer_points（直接写出、多空或多个结果，按答对数量给分）、"
        "process_required（证明/计算过程）、visual_construction（作图，以答案图为视觉参考）。\n"
        "题目写有“直接写出”“填空”时，不得强制过程；short_answer_points 的每个可独立得分答案应拆成独立踩分点。\n"
        "选择题和普通填空题的评分点只能描述核对最终答案，不能要求推理或计算过程。\n"
        "若一个填空要求列出全部可能答案，必须在 canonical_answer 写全，并输出 match_mode=complete_set、required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；不要把少写答案列为 accepted_forms。\n"
        "作图题不得把答案图臆造为唯一文字答案；应输出 visual_requirements 和必要踩分点。\n"
        "必须输出精炼且具体的 knowledge_name、knowledge_id、knowledge_points；严禁用题干或“几何综合/代数综合/综合应用”充当知识点。\n"
        "每个主观题步骤的 required_elements 必须写出可核验的条件、式子或结论，不得留空或只写通用描述。\n"
        f"知识点参考字典：{knowledge_options}\n"
        "仅为该题输出 rubric.questions 与 answer_key.questions，并保持 question_id 一致。\n"
        "必须严格使用以下字段，不得自行改名或创造 answers、answer_parts、answer_content、desc 等同义字段：\n"
        "{\"rubric\":{\"questions\":[{\"question_id\":\"...\",\"question_type\":\"...\","
        "\"knowledge_name\":\"...\",\"knowledge_id\":\"...\",\"knowledge_points\":["
        "{\"knowledge_id\":\"...\",\"knowledge_name\":\"具体考查内容\"}],\"parts\":["
        "{\"part_id\":\"...\",\"response_mode\":\"...\",\"visual_requirements\":[],\"steps\":["
        "{\"step_id\":\"S1\",\"core_goal\":\"具体踩分点描述\",\"required_elements\":[\"可核验的必要条件或结论\"],"
        "\"allow_alternative_methods\":true}]}]}]},\"answer_key\":{\"questions\":["
        "{\"question_id\":\"...\",\"canonical_answer\":\"...\",\"accepted_forms\":[],\"parts\":["
        "{\"part_id\":\"...\",\"answer\":\"...\",\"accepted_forms\":[],\"analysis\":\"\","
        "\"step_milestones\":[]}]}]}}"
    )


def _build_single_question_generation_prompt(block: dict[str, str], doc_text: str) -> str:
    """Build a prompt for a single question's rubric structure.

    IMPORTANT: Scores are NOT requested here — they are allocated in a dedicated
    second AI call (_assign_scores_to_merged_config) after all questions are merged.
    All score fields must be set to 1 as a placeholder.
    """
    qid = str(block.get("question_id") or "").strip()
    q_text = str(block.get("text") or "").strip()
    answer_text = str(block.get("answer_text") or "").strip()
    analysis_text = str(block.get("analysis") or "").strip()
    confirmed_type = str(block.get("question_type") or "").strip()
    type_is_confirmed = bool(block.get("question_type_confirmed"))

    # Build structured per-question context (replaces the old 6000-char full-doc dump)
    context_parts: list[str] = [f"题目文本：\n{q_text}"]
    if type_is_confirmed and confirmed_type:
        context_parts.append(
            f"教师已确认题型：{confirmed_type}\n"
            "该题型是强约束，必须原样写入 question_type，不得自行改成其他题型。"
        )
    if answer_text:
        context_parts.append(f"参考答案：\n{answer_text}")
    if analysis_text:
        context_parts.append(f"解析/证明过程：\n{analysis_text}")
    if not answer_text and not analysis_text:
        # Fallback: include a limited window of the full doc text to help locate the answer
        context_parts.append(
            "（未能提取到配套答案，以下是试卷原文供参考，请自行定位该题答案区域）：\n"
            + str(doc_text or "")[:4000]
        )
    question_context = "\n\n".join(context_parts)

    return (
        "你正在为【单道】初中数学题目生成评分标准大纲结构。请仅返回严格的 JSON 数据。\n"
        f"题目 ID (QUESTION_ID): {qid}\n"
        "重要：请将所有评分相关的分值字段（包括 max_score, part_score, step_score 等）全部设为占位符 1。"
        "具体分值会在后续步骤中统一分配，请勿尝试在此猜测真实分值。\n"
        "规则：仅为该道题构建相应的 rubric.questions 和 answer_key.questions。\n"
        "如果是选择题或填空题，应给出标准答案以及 accepted_forms（等价接受形式）。\n"
        "如果是计算题、证明题、综合题，必须包含 parts、steps、proof_obligations（证明证据点）、"
        "deduction_policy（扣分策略）、证据要求以及仅写出答案的上限得分（answer_only_max_score）。\n"
        "每个 parts 项必须输出 response_mode：exact_objective、short_answer_points、process_required 或 visual_construction；"
        "不得把大题级过程要求无条件继承给“直接写出”或作图小问。\n"
        "选择题和普通填空题的评分点只能描述核对最终答案，不能要求推理或计算过程。\n"
        "若一个填空要求列出全部可能答案，必须在 canonical_answer 写全，并输出 match_mode=complete_set、required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；不要把少写答案列为 accepted_forms。\n"
        "如果该题包含多个空格、表格单元格或子小问，必须将其拆分为不同的 parts 以给与步骤/部分分。\n"
        "accepted_forms 必须仅包含在数学上完全等价的答案形式。\n\n"
        "必须输出精炼且具体的 knowledge_name、knowledge_id、knowledge_points；严禁用题干或“几何综合/代数综合/综合应用”充当知识点。\n"
        "每个主观题步骤的 required_elements 必须写出可核验的条件、式子或结论，不得留空或只写通用描述。\n\n"
        f"{question_context}"
    )


def _merge_single_question_payloads(
    payloads: list[dict[str, Any] | None],
    question_blocks: list[dict[str, str]] | None = None,
) -> dict[str, Any]:
    merged = {
        "rubric": {"exam_title": "generated", "total_score": 100, "questions": []},
        "answer_key": {"questions": []},
        "meta": {"source": "word_docx_parallel", "warnings": []},
    }
    seen_questions: set[str] = set()
    seen_answers: set[str] = set()
    for index, payload in enumerate(payloads):
        if not isinstance(payload, dict):
            continue
        expected_qid = ""
        if question_blocks and index < len(question_blocks):
            expected_qid = str(question_blocks[index].get("question_id") or "").strip()
        _coerce_single_question_payload_schema(payload, expected_qid, merged["meta"]["warnings"])
        normalize_generated_config_schema(payload)
        rubric = payload.get("rubric") if isinstance(payload, dict) else {}
        answer_key = payload.get("answer_key") if isinstance(payload, dict) else {}
        for question in rubric.get("questions", []) if isinstance(rubric, dict) else []:
            if not isinstance(question, dict):
                continue
            qid = str(question.get("question_id") or "").strip()
            if not qid or qid in seen_questions:
                continue
            seen_questions.add(qid)
            merged["rubric"]["questions"].append(question)
        for answer in answer_key.get("questions", []) if isinstance(answer_key, dict) else []:
            if not isinstance(answer, dict):
                continue
            qid = str(answer.get("question_id") or "").strip()
            if not qid or qid in seen_answers:
                continue
            seen_answers.add(qid)
            merged["answer_key"]["questions"].append(answer)
        meta = payload.get("meta") if isinstance(payload, dict) else {}
        warnings = meta.get("warnings") if isinstance(meta, dict) else []
        if isinstance(warnings, list):
            merged["meta"]["warnings"].extend(str(item) for item in warnings if str(item).strip())
    return merged


def _coerce_single_question_payload_schema(
    payload: dict[str, Any],
    expected_qid: str,
    warnings: list[str],
) -> None:
    """Accept common one-question schemas returned by LLMs before strict normalization."""
    raw_rubric = payload.get("rubric")
    raw_answer_key = payload.get("answer_key")
    rubric = raw_rubric if isinstance(raw_rubric, dict) else {}
    answer_key = raw_answer_key if isinstance(raw_answer_key, dict) else {}
    payload["rubric"] = rubric
    payload["answer_key"] = answer_key

    raw_questions = rubric.get("questions") if isinstance(raw_rubric, dict) else raw_rubric
    questions = _coerce_single_item_list(raw_questions, _looks_like_question_dict)
    rubric["questions"] = questions
    raw_answers = answer_key.get("questions") if isinstance(raw_answer_key, dict) else raw_answer_key
    answers = _coerce_single_item_list(raw_answers, _looks_like_answer_dict)
    answer_key["questions"] = answers

    question = next((item for item in questions if isinstance(item, dict)), None)
    if question is None:
        question = _extract_single_question_candidate(payload, rubric)
        if isinstance(question, dict):
            questions.append(question)
            if expected_qid:
                warnings.append(f"{expected_qid} normalized singular rubric.question schema")

    answer = next((item for item in answers if isinstance(item, dict)), None)
    if answer is None:
        answer = _extract_single_answer_candidate(payload, answer_key, question)
        if isinstance(answer, dict):
            answers.append(answer)
            if expected_qid:
                warnings.append(f"{expected_qid} normalized singular answer_key.question schema")

    if expected_qid:
        for item in [question, answer]:
            if isinstance(item, dict):
                raw_qid = item.get("question_id") or item.get("id") or item.get("number")
                item["question_id"] = _canonical_question_id(raw_qid, expected_qid)

    if question is None and expected_qid:
        keys = ", ".join(sorted(str(key) for key in payload.keys()))
        warnings.append(f"{expected_qid} unmergeable single-question schema; top-level keys: {keys}")


def _coerce_single_item_list(value: Any, predicate: Callable[[Any], bool]) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if isinstance(value, dict):
        if predicate(value):
            return [value]
        return [item for item in value.values() if isinstance(item, dict) and predicate(item)]
    return []


def _extract_single_question_candidate(payload: dict[str, Any], rubric: dict[str, Any]) -> dict[str, Any] | None:
    candidates = [
        rubric.get("question"),
        payload.get("question"),
        payload.get("rubric_question"),
    ]
    if _looks_like_question_dict(rubric):
        candidates.append(rubric)
    if _looks_like_question_dict(payload):
        candidates.append(payload)
    for candidate in candidates:
        if isinstance(candidate, dict):
            return candidate
    return None


def _extract_single_answer_candidate(
    payload: dict[str, Any],
    answer_key: dict[str, Any],
    question: dict[str, Any] | None,
) -> dict[str, Any] | None:
    candidates = [
        answer_key.get("question"),
        payload.get("answer_key_question"),
        payload.get("answer"),
        payload.get("correct_answer"),
    ]
    if _looks_like_answer_dict(answer_key):
        candidates.append(answer_key)
    if isinstance(question, dict) and _looks_like_answer_dict(question):
        candidates.append(question)
    for candidate in candidates:
        if isinstance(candidate, dict):
            return candidate
        if candidate is not None:
            return {"canonical_answer": str(candidate)}
    return None


def _looks_like_question_dict(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    return any(
        key in item
        for key in [
            "question_id",
            "id",
            "number",
            "question_type",
            "type",
            "max_score",
            "score",
            "points",
            "parts",
            "stem_summary",
        ]
    )


def _looks_like_answer_dict(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    return any(
        key in item
        for key in [
            "canonical_answer",
            "answer",
            "correct_answer",
            "direct_answer",
            "accepted_forms",
            "equivalent_answers",
            "parts",
        ]
    )


def _attach_parallel_generation_meta(
    payload: dict[str, Any],
    question_blocks: list[dict[str, str]],
    failures: list[dict[str, Any]],
    worker_count: int,
    attempt_counts: dict[str, int],
) -> None:
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    warnings = meta.setdefault("warnings", [])
    if not isinstance(warnings, list):
        warnings = []
        meta["warnings"] = warnings

    image_split_mode = any(
        str(block.get("semantic_source") or "").strip() == "images"
        for block in question_blocks
        if isinstance(block, dict)
    )
    meta["source"] = "pdf_image_split_parallel" if image_split_mode else "word_rich_split_parallel"
    meta["generation_mode"] = "pdf_image_split_parallel" if image_split_mode else "word_rich_split_parallel"
    meta["question_block_count"] = len(question_blocks)
    failed_question_ids = list(
        dict.fromkeys(
            str(failure.get("question_id") or "").strip()
            for failure in failures
            if str(failure.get("question_id") or "").strip()
        )
    )
    meta["single_question_success_count"] = max(0, len(question_blocks) - len(failed_question_ids))
    meta["single_question_failure_count"] = len(failed_question_ids)
    meta["config_generation_workers"] = int(worker_count)
    meta["single_question_attempts"] = {
        str(qid): int(attempts)
        for qid, attempts in attempt_counts.items()
        if str(qid).strip()
    }
    meta["failed_question_ids"] = failed_question_ids
    meta["failed_questions"] = [dict(failure) for failure in failures]

    warnings[:] = [
        warning
        for warning in warnings
        if str(warning) != "parallel generation completed with failed question blocks; score allocation paused until retry."
    ]
    existing = {str(item) for item in warnings}
    for failure in failures:
        qid = str(failure.get("question_id") or "").strip()
        error = str(failure.get("error") or "").strip()
        attempts = int(failure.get("attempts") or 1)
        warning = f"parallel generation failed for {qid} after {attempts} attempt(s): {error}"
        if warning not in existing:
            warnings.append(warning)
            existing.add(warning)
    if failures:
        summary = "parallel generation completed with failed question blocks; score allocation paused until retry."
        if summary not in existing:
            warnings.append(summary)


def _config_final_calibration_enabled(default: bool = False) -> bool:  # noqa: FBT001,FBT002
    """Final model calibration is OFF by default to avoid extra cost and failure risk.

    Set env var AI_GRADING_CONFIG_FINAL_CALIBRATION=1 to re-enable it.
    """
    raw = str(os.getenv("AI_GRADING_CONFIG_FINAL_CALIBRATION", "") or "").strip().lower()
    if not raw:
        return bool(default)
    if raw in {"0", "false", "no", "off", "disabled"}:
        return False
    return raw in {"1", "true", "yes", "on", "enabled"}


def _ensure_question_blocks_covered(payload: dict[str, Any], question_blocks: list[dict[str, str]]) -> None:
    rubric = payload.setdefault("rubric", {})
    answer_key = payload.setdefault("answer_key", {})
    if not isinstance(rubric, dict) or not isinstance(answer_key, dict):
        return
    questions = rubric.setdefault("questions", [])
    answers = answer_key.setdefault("questions", [])
    if not isinstance(questions, list) or not isinstance(answers, list):
        return

    existing_qids = {
        str(question.get("question_id") or "").strip()
        for question in questions
        if isinstance(question, dict)
    }
    existing_answer_qids = {
        str(answer.get("question_id") or "").strip()
        for answer in answers
        if isinstance(answer, dict)
    }
    warnings = payload.setdefault("meta", {}).setdefault("warnings", [])

    for block in question_blocks:
        qid = str(block.get("question_id") or "").strip()
        if not qid:
            continue
        if qid not in existing_qids:
            questions.append(_placeholder_question_from_block(block))
            existing_qids.add(qid)
            if isinstance(warnings, list):
                warnings.append(f"{qid} placeholder added because per-question generation did not return a mergeable rubric item")
        if qid not in existing_answer_qids:
            answers.append({"question_id": qid, "canonical_answer": "", "accepted_forms": [], "method_variants": [], "parts": []})
            existing_answer_qids.add(qid)

    qid_order = {
        str(block.get("question_id") or "").strip(): index
        for index, block in enumerate(question_blocks)
        if str(block.get("question_id") or "").strip()
    }
    questions.sort(key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000) if isinstance(item, dict) else 10_001)
    answers.sort(key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000) if isinstance(item, dict) else 10_001)


def _placeholder_question_from_block(block: dict[str, str]) -> dict[str, Any]:
    qid = str(block.get("question_id") or "").strip()
    text = str(block.get("text") or "").strip()
    qtype = _infer_question_type_from_block_text(text)
    return {
        "question_id": qid,
        "question_type": qtype,
        "max_score": 1,
        "knowledge_id": "UNKNOWN",
        "stem_summary": text.splitlines()[0][:120] if text else qid,
        "parts": [
            {
                "part_id": qid,
                "part_score": 1,
                "steps": [
                    {
                        "step_id": "S1",
                        "step_score": 1,
                        "core_goal": "Manual rubric confirmation required",
                        "required_elements": ["AI did not return a mergeable question; confirm rubric manually"],
                        "allow_alternative_methods": True,
                    }
                ],
                "presentation_rules": [],
            }
        ],
    }


def _infer_question_type_from_block_text(text: str) -> str:
    value = str(text or "")
    option_labels = {
        match.group(1).upper()
        for match in re.finditer(r"(?m)^\s*([A-Da-d])\s*(?:[.．、)]|\s{2,})", value)
    }
    option_line = any(re.search(r"A.{0,80}B.{0,80}C.{0,80}D", line) for line in value.splitlines())
    if len(option_labels) >= 3 or option_line:
        return "choice"
    has_blank = bool(re.search(r"_{2,}|[ \t]{3,}|　{1,}|（\s*）|\(\s*\)", value))
    if has_blank:
        return "fill_blank"
    if any(token in value for token in ["作图", "作出", "画出", "保留作图痕迹"]):
        return "comprehensive"
    if any(token in value for token in ["证明", "理由", "说明", "求证", "全等", "证得"]):
        return "proof"
    has_subparts = bool(re.search(r"[（(]\s*[1-9]\s*[）)]", value))
    if has_subparts:
        return "calculation"
    if any(token in value.lower() for token in ["proof", "why", "explain", "relationship"]):
        return "comprehensive"
    return "comprehensive"


def _attach_generation_fallback_warning(
    payload: dict[str, Any],
    question_blocks: list[dict[str, str]],
    parallel_error: Exception | None,
) -> None:
    if parallel_error is None:
        return
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    warnings = meta.setdefault("warnings", [])
    if not isinstance(warnings, list):
        warnings = []
        meta["warnings"] = warnings
    meta["source"] = "word_docx_single_fallback"
    meta["question_block_count"] = len(question_blocks)
    meta["parallel_generation_fallback"] = True
    warnings.append(f"parallel generation failed; fell back to single full-document generation: {parallel_error}")


def _assign_scores_to_merged_config(
    merged: dict[str, Any],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    *,
    include_document_text: bool = True,
) -> None:
    """Phase-2 AI call: allocate real scores to all questions so they sum to 100.

    Per-question prompts set all score fields to the placeholder value 1.
    This function sends the merged structure + original Word text to the model,
    which reads score hints like （3分）, 共X分 from the document and distributes
    the remaining points proportionally to question complexity.

    On failure, falls back to local proportional scaling via force_payload_total_score.
    """
    rubric_questions: list[Any] = merged.get("rubric", {}).get("questions", [])
    if not isinstance(rubric_questions, list) or not rubric_questions:
        force_payload_total_score(merged, target_total=100.0)
        return

    # Build a lightweight structure summary: only IDs, types, part/step counts
    structure_summary: list[dict[str, Any]] = []
    for q in rubric_questions:
        if not isinstance(q, dict):
            continue
        parts = q.get("parts") or []
        part_list = []
        for p in parts if isinstance(parts, list) else []:
            if not isinstance(p, dict):
                continue
            steps = p.get("steps") or []
            part_list.append({
                "part_id": str(p.get("part_id") or ""),
                "response_mode": str(p.get("response_mode") or ""),
                "step_ids": [str(s.get("step_id") or "") for s in steps if isinstance(s, dict)],
            })
        structure_summary.append({
            "question_id": str(q.get("question_id") or ""),
            "question_type": str(q.get("question_type") or ""),
            "parts": part_list,
        })

    prompt = (
        "你正在分配试卷各题目的得分。试卷总分必须精确等于 100 分。\n\n"
        "硬性规则（请严格遵守）：\n"
        "1. 所有相同类型（选择题 choice / 填空题 fill_blank）的题目必须共享相同的 max_score（即同类题目的单题分值相同）。\n"
        "2. 对于每道题目：其下属各部分的 part_score 之和必须等于 max_score。对于每个 part：其下属各步骤的 step_score 之和必须等于 part_score。\n"
        "3. 所有分值必须是正整数。\n"
        "4. 试卷所有题目的 max_score 之和必须精确等于 100。\n"
        "5. 将未分配的分数，按照步骤的数量比例分配给主观大题。\n"
        "6. 关键限制 —— 所有题型的单题分值上限均为 18 分：\n"
        "   - 每道独立题目的 max_score 必须 <= 18 分。\n"
        "   - 所有共享相同大题号的子小问（例如 Q10_1 和 Q10_2 共享同一个大题前缀 Q10）的合并总分必须 <= 18 分。\n"
        "   - 若某题原有标注分值超过 18 分，必须封顶为 18 分，并将其余分数分摊给其他题目。\n"
        "   - 你必须严格保持输入中提供的题目结构和 question_id 列表，绝对不允许新增、删除、修改或拆分任何 question_id 题号。\n\n"
        "Return ONLY the following JSON (no markdown, no explanation):\n"
        "{\n"
        "  \"question_scores\": [\n"
        "    {\n"
        "      \"question_id\": \"Q1\",\n"
        "      \"max_score\": 3,\n"
        "      \"parts\": [\n"
        "        {\"part_id\": \"Q1\", \"part_score\": 3,\n"
        "         \"steps\": [{\"step_id\": \"S1\", \"step_score\": 3}]}\n"
        "      ]\n"
        "    }\n"
        "  ]\n"
        "}\n\n"
        f"Word document (for score hints, first 6000 chars):\n{str(doc_text or '')[:6000]}\n\n"
        f"Question structure to score:\n{json.dumps(structure_summary, ensure_ascii=False)}"
    )
    prompt = _build_score_allocation_prompt(
        structure_summary,
        doc_text,
        include_document_text=include_document_text,
    )

    if report:
        q_count = len(structure_summary)
        report(0.88, "AI 赋分", f"正在为合并后的 {q_count} 道题统一分配分值（共 100 分）...")

    try:
        score_data = llm_client.json_from_text(prompt, model=model_name)
        _apply_score_allocation(merged, score_data)
        merged.setdefault("meta", {})["score_allocation_ai_success"] = True
        if report:
            report(0.91, "AI 赋分完成", "已将 AI 分配的分值写入各题评分细则。")
    except Exception as exc:
        if report:
            report(0.91, "AI 赋分失败，使用本地均分兜底", str(exc))
        warnings = merged.setdefault("meta", {}).setdefault("warnings", [])
        if isinstance(warnings, list):
            warnings.append(f"score allocation AI call failed; fell back to local scaling: {exc}")
        merged.setdefault("meta", {})["score_allocation_ai_success"] = False
    # Always do a local normalization pass to guarantee score consistency
    force_payload_total_score(merged, target_total=100.0)


def _build_score_allocation_prompt(
    structure_summary: list[dict[str, Any]],
    doc_text: str,
    *,
    include_document_text: bool,
) -> str:
    source_context = (
        f"\n原始文档文本（仅用于识别原卷分值提示）：\n{str(doc_text or '')[:6000]}\n"
        if include_document_text and str(doc_text or "").strip()
        else "\n本次为图片语义来源，不提供也不得推测 PDF 抽取文字或原卷分值。\n"
    )
    return (
        "请仅为下列已确认题目结构分配分值，总分必须精确等于100。\n"
        "不同题型之间不限制分值高低；相同类型客观题必须同分；所有分值均为正整数；单题不超过18分。\n"
        "保持所有 question_id、part_id、step_id 和小问结构不变。"
        "分值可以不采用原卷分值，但不得改变小问作答要求或 response_mode。\n"
        f"{source_context}"
        "仅返回 JSON：{\"question_scores\":[{\"question_id\":\"Q1\",\"max_score\":1,"
        "\"parts\":[{\"part_id\":\"Q1\",\"part_score\":1,\"steps\":[{\"step_id\":\"S1\",\"step_score\":1}]}]}]}\n"
        f"待分值结构：\n{json.dumps(structure_summary, ensure_ascii=False)}"
    )


def _apply_score_allocation(merged: dict[str, Any], score_data: dict[str, Any]) -> None:
    """Write AI-assigned scores from the scoring step back into the merged rubric."""
    question_scores = score_data.get("question_scores")
    if not isinstance(question_scores, list):
        raise ValueError("score_data must contain a 'question_scores' list")

    rubric_questions: list[Any] = merged.get("rubric", {}).get("questions", [])
    if not isinstance(rubric_questions, list):
        return

    score_map: dict[str, dict[str, Any]] = {
        str(qs.get("question_id") or "").strip(): qs
        for qs in question_scores
        if isinstance(qs, dict) and str(qs.get("question_id") or "").strip()
    }

    for q in rubric_questions:
        if not isinstance(q, dict):
            continue
        qid = str(q.get("question_id") or "").strip()
        score_entry = score_map.get(qid)
        if not score_entry:
            continue

        # Apply max_score
        raw_max = score_entry.get("max_score")
        if raw_max is not None:
            try:
                q["max_score"] = max(1, int(round(float(raw_max))))
            except (TypeError, ValueError):
                pass

        # Apply part and step scores
        parts: list[Any] = q.get("parts") or []
        score_parts: list[Any] = score_entry.get("parts") or []
        if not isinstance(parts, list) or not score_parts:
            continue

        part_score_map: dict[str, dict[str, Any]] = {
            str(sp.get("part_id") or "").strip(): sp
            for sp in score_parts
            if isinstance(sp, dict) and str(sp.get("part_id") or "").strip()
        }

        for part in parts:
            if not isinstance(part, dict):
                continue
            pid = str(part.get("part_id") or "").strip()
            sp = part_score_map.get(pid)
            if not sp:
                continue

            raw_part = sp.get("part_score")
            if raw_part is not None:
                try:
                    part["part_score"] = max(1, int(round(float(raw_part))))
                except (TypeError, ValueError):
                    pass

            steps: list[Any] = part.get("steps") or []
            score_steps: list[Any] = sp.get("steps") or []
            if not isinstance(steps, list) or not score_steps:
                continue

            step_score_map: dict[str, dict[str, Any]] = {
                str(ss.get("step_id") or "").strip(): ss
                for ss in score_steps
                if isinstance(ss, dict) and str(ss.get("step_id") or "").strip()
            }

            for step in steps:
                if not isinstance(step, dict):
                    continue
                sid = str(step.get("step_id") or "").strip()
                ss = step_score_map.get(sid)
                if ss:
                    raw_step = ss.get("step_score")
                    if raw_step is not None:
                        try:
                            step["step_score"] = max(1, int(round(float(raw_step))))
                        except (TypeError, ValueError):
                            pass


def _calibrate_merged_config(
    payload: dict[str, Any],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    *,
    strict: bool = False,
) -> None:
    prompt = (
        "final calibration: You are a rubric calibration assistant. "
        "Check the merged per-question rubric against the full document. "
        "Only fix missing question ids, empty answers, obvious question type errors, and total score allocation. "
        "Do not remove existing proof_obligations, steps, or parts. "
        "accepted_forms must keep only strictly equivalent answers. Return complete JSON with the same schema.\n\n"
        f"Full document text:\n{str(doc_text or '')[:8000]}\n\n"
        f"Merged result:\n{json.dumps(payload, ensure_ascii=False)}"
    )
    try:
        calibrated = llm_client.json_from_text(prompt, model=model_name)
        validate_generated_config(calibrated)
    except Exception as exc:
        if strict:
            raise RuntimeError(f"final calibration failed: {exc}") from exc
        meta = payload.setdefault("meta", {})
        if isinstance(meta, dict):
            warnings = meta.setdefault("warnings", [])
            if isinstance(warnings, list):
                warnings.append(f"final calibration failed; kept merged per-question config: {exc}")
        return
    payload.clear()
    payload.update(calibrated)


def _bounded_int(value: Any, default: int, minimum: int, maximum: int) -> int:
    try:
        resolved = int(value) if value is not None else int(default)
    except (TypeError, ValueError):
        resolved = int(default)
    return max(minimum, min(maximum, resolved))


def refine_grading_config_from_manual_structure(
    payload: dict[str, Any],
    llm_client: LLMClient,
    model_name: str | None = None,
) -> dict[str, Any]:
    """Ask AI to complete answers/rubrics while preserving teacher-edited parts."""
    working_payload = json.loads(json.dumps(payload, ensure_ascii=False))
    validate_generated_config(working_payload)
    prompt = _build_manual_structure_refinement_prompt(working_payload)
    refined = llm_client.json_from_text(prompt, model=model_name)
    try:
        validate_generated_config(refined)
    except Exception:
        _dump_failed_generated_payload(refined)
        raise
    return refined


def _dump_failed_generated_payload(payload: dict[str, Any]) -> None:
    try:
        try:
            from path_manager import get_path_manager
            out_dir = get_path_manager().upload_config_dir
        except Exception:
            out_dir = Path(__file__).resolve().parent / "config" / "uploaded"
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        (out_dir / f"failed_generated_config_normalized_{ts}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except Exception:
        pass


def normalize_generated_config_schema(payload: dict[str, Any]) -> None:
    """Convert common LLM shorthand fields into the strict internal schema."""
    if not isinstance(payload, dict):
        return

    rubric = payload.setdefault("rubric", {})
    answer_key = payload.setdefault("answer_key", {})
    payload.setdefault("meta", {}).setdefault("warnings", [])
    if not isinstance(rubric, dict) or not isinstance(answer_key, dict):
        return

    rubric_questions = rubric.setdefault("questions", [])
    answer_questions = answer_key.setdefault("questions", [])
    if not isinstance(rubric_questions, list):
        rubric["questions"] = []
        rubric_questions = rubric["questions"]
    if not isinstance(answer_questions, list):
        answer_key["questions"] = []
        answer_questions = answer_key["questions"]

    for answer in answer_questions:
        if isinstance(answer, dict):
            answer["question_id"] = _canonical_question_id(
                answer.get("question_id")
                or answer.get("id")
                or answer.get("number")
                or answer.get("棰樺彿")
                or "",
                "",
            )
            _coerce_answer_item_aliases(answer)

    answer_map = {
        _canonical_question_id(q.get("question_id") or q.get("id") or q.get("number") or "", ""): q
        for q in answer_questions
        if isinstance(q, dict)
    }

    for idx, question in enumerate(rubric_questions, start=1):
        if not isinstance(question, dict):
            continue
        qid = _canonical_question_id(
            question.get("question_id")
            or question.get("id")
            or question.get("number")
            or f"Q{idx}",
            f"Q{idx}",
        )
        question["question_id"] = qid
        _promote_nested_question_knowledge(question)
        qtype = _normalize_question_type(
            question.get("question_type")
            or question.get("type")
        )
        question["question_type"] = qtype
        question["max_score"] = _safe_float(
            question.get("max_score")
            or question.get("score")
            or question.get("points"),
            0.0,
        )
        question["knowledge_id"] = question.get("knowledge_id") or question.get("knowledge") or "UNKNOWN"
        question["knowledge_name"] = (
            question.get("knowledge_name")
            or question.get("knowledge_text")
            or question.get("knowledge_label")
            or question.get("knowledge")
            or ""
        )
        _normalize_question_knowledge_fields(question)
        question["stem_summary"] = str(question.get("stem_summary") or question.get("棰樺共鎽樿") or "").strip()
        question["grading_mode"] = str(
            question.get("grading_mode")
            or ("deductive_obligation" if qtype in {"proof", "calculation", "comprehensive"} else "direct_answer")
        )

        answer_item = answer_map.get(qid)
        if not isinstance(answer_item, dict):
            answer_item = {"question_id": qid}
            answer_questions.append(answer_item)
            answer_map[qid] = answer_item
        _normalize_answer_item(answer_item, question)
        if (
            not bool(question.get("question_type_confirmed"))
            and _should_treat_as_direct_answer_question(qtype, question, answer_item)
        ):
            qtype = "fill_blank"
            question["question_type"] = qtype
            question["grading_mode"] = "direct_answer"
        _augment_answer_equivalences(answer_item, qtype)
        _normalize_rubric_question(question, answer_item)
        _align_answer_parts_to_rubric_parts(question, answer_item)
        _align_step_required_elements_with_answer_values(question, answer_item)
        _enforce_objective_question_rules(question, answer_item)
        _ensure_solution_hard_rules(question)

    rubric["total_score"] = _safe_float(
        rubric.get("total_score"),
        sum(_safe_float(q.get("max_score"), 0.0) for q in rubric_questions if isinstance(q, dict)),
    )


def _canonical_question_id(value: Any, fallback: str = "") -> str:
    raw = str(value or "").strip()
    if not raw:
        return str(fallback or "").strip()
    match = re.match(r"^(?:Q|q)\s*(\d{1,2})$", raw)
    if match:
        return f"Q{int(match.group(1))}"
    match = re.match(r"^(\d{1,2})$", raw)
    if match:
        return f"Q{int(match.group(1))}"
    match = re.search(r"(\d{1,2})", raw)
    if re.match(r"^\\D+\\d{1,2}", raw) and match:
        return f"Q{int(match.group(1))}"
    return raw


def _safe_knowledge_sequence(value: Any) -> list[str] | None:
    candidate = value
    if isinstance(value, str):
        text = value.strip()
        if not _looks_like_serialized_knowledge_sequence(text):
            return None
        try:
            candidate = ast.literal_eval(text)
        except (SyntaxError, ValueError):
            return None
    if not isinstance(candidate, (list, tuple)) or not candidate:
        return None

    values: list[str] = []
    for item in candidate:
        if isinstance(item, bool) or not isinstance(item, (str, int, float)):
            return None
        text = str(item).strip()
        if not text:
            return None
        values.append(text)
    return values


def _redundant_knowledge_fields_match_pairs(
    question: dict[str, Any],
    paired_ids: list[str],
    paired_names: list[str],
) -> bool:
    raw_ids = question.get("knowledge_ids")
    if raw_ids not in (None, "", []):
        ids_match = _safe_knowledge_sequence(raw_ids) == paired_ids
        if (
            not ids_match
            and isinstance(raw_ids, (list, tuple))
            and len(raw_ids) == 1
        ):
            ids_match = _safe_knowledge_sequence(raw_ids[0]) == paired_ids
        if not ids_match:
            return False

    raw_points = question.get("knowledge_points")
    if raw_points in (None, "", []):
        return True
    if not isinstance(raw_points, (list, tuple)):
        return False

    expected_names = dict(zip(paired_ids, paired_names))
    canonical_ids: set[str] = set()
    for point in raw_points:
        if isinstance(point, dict):
            raw_id = point.get("knowledge_id") or point.get("id") or ""
            raw_name = (
                point.get("knowledge_name")
                or point.get("name")
                or point.get("label")
                or ""
            )
        else:
            raw_id = point
            raw_name = ""

        point_ids = _safe_knowledge_sequence(raw_id)
        if point_ids is not None:
            if point_ids != paired_ids:
                return False
            if raw_name and _safe_knowledge_sequence(raw_name) != paired_names:
                return False
            continue
        if _looks_like_serialized_knowledge_sequence(raw_id):
            return False
        if raw_name and _looks_like_serialized_knowledge_sequence(raw_name):
            return False

        kid = str(raw_id or "").strip()
        name = str(raw_name or "").strip()
        if not kid or kid not in expected_names:
            return False
        if name and name != expected_names[kid]:
            return False
        canonical_ids.add(kid)

    return not canonical_ids or canonical_ids == set(paired_ids)


def _normalize_question_knowledge_fields(question: dict[str, Any]) -> None:
    raw_primary_id = question.get("knowledge_id") or "UNKNOWN"
    raw_primary_name = question.get("knowledge_name") or ""
    paired_ids = _safe_knowledge_sequence(raw_primary_id)
    paired_names = _safe_knowledge_sequence(raw_primary_name)
    if (
        paired_ids is not None
        and paired_names is not None
        and len(paired_ids) == len(paired_names)
        and len(set(paired_ids)) == len(paired_ids)
        and _redundant_knowledge_fields_match_pairs(
            question,
            paired_ids,
            paired_names,
        )
    ):
        paired_points = [
            {"knowledge_id": kid, "knowledge_name": name}
            for kid, name in zip(paired_ids, paired_names)
        ]
        question["knowledge_points"] = paired_points
        question["knowledge_ids"] = list(paired_ids)
        question["knowledge_id"] = paired_ids[0]
        question["knowledge_name"] = paired_names[0]
        return

    raw_points = question.get("knowledge_points")
    points: list[dict[str, str]] = []
    if isinstance(raw_points, list):
        for item in raw_points:
            if isinstance(item, dict):
                kid = str(item.get("knowledge_id") or item.get("id") or "").strip()
                name = str(item.get("knowledge_name") or item.get("name") or item.get("label") or "").strip()
            else:
                kid = str(item or "").strip()
                name = ""
            if kid:
                points.append({"knowledge_id": kid, "knowledge_name": name})
    elif isinstance(raw_points, str) and raw_points.strip():
        points.append(
            {
                "knowledge_id": str(question.get("knowledge_id") or question.get("knowledge_name") or "UNKNOWN").strip(),
                "knowledge_name": raw_points.strip(),
            }
        )

    raw_ids = question.get("knowledge_ids")
    if isinstance(raw_ids, list):
        for raw_id in raw_ids:
            kid = str(raw_id or "").strip()
            if kid:
                points.append({"knowledge_id": kid, "knowledge_name": ""})

    primary_id = str(raw_primary_id).strip()
    primary_name = str(raw_primary_name).strip()
    if primary_id:
        points.append({"knowledge_id": primary_id, "knowledge_name": primary_name})

    useful_points = [point for point in points if point["knowledge_id"] not in {"", "UNKNOWN"}]
    if useful_points:
        points = useful_points

    normalized: list[dict[str, str]] = []
    for point in points:
        kid = point["knowledge_id"]
        name = point.get("knowledge_name", "")
        if not kid:
            continue
        same_id = [
            index for index, existing in enumerate(normalized)
            if existing["knowledge_id"] == kid
        ]
        if any(normalized[index].get("knowledge_name", "") == name for index in same_id):
            continue
        if not name and same_id:
            continue
        empty_index = next(
            (index for index in same_id if not normalized[index].get("knowledge_name")),
            None,
        )
        if name and empty_index is not None:
            normalized[empty_index]["knowledge_name"] = name
        else:
            normalized.append({"knowledge_id": kid, "knowledge_name": name})

    if not normalized:
        normalized = [{"knowledge_id": "UNKNOWN", "knowledge_name": ""}]

    question["knowledge_points"] = normalized
    question["knowledge_ids"] = list(dict.fromkeys(point["knowledge_id"] for point in normalized))
    question["knowledge_id"] = normalized[0]["knowledge_id"]
    if normalized[0].get("knowledge_name"):
        question["knowledge_name"] = normalized[0]["knowledge_name"]


def normalize_generated_config_knowledge_fields(payload: dict[str, Any]) -> bool:
    """Normalize only rubric knowledge metadata and report whether it changed."""
    rubric = payload.get("rubric") if isinstance(payload, dict) else None
    questions = rubric.get("questions") if isinstance(rubric, dict) else None
    if not isinstance(questions, list):
        return False

    changed = False
    tracked = ("knowledge_id", "knowledge_name", "knowledge_ids", "knowledge_points")
    for question in questions:
        if not isinstance(question, dict):
            continue
        raw_points = question.get("knowledge_points")
        point_ids = [
            str(point.get("knowledge_id") or point.get("id") or "").strip()
            for point in raw_points
            if isinstance(point, dict)
        ] if isinstance(raw_points, list) else []
        sequence_values: list[Any] = [
            question.get("knowledge_id"),
            question.get("knowledge_name"),
        ]
        raw_ids = question.get("knowledge_ids")
        if isinstance(raw_ids, list):
            sequence_values.extend(raw_ids)
        elif raw_ids is not None:
            sequence_values.append(raw_ids)
        if isinstance(raw_points, list):
            for point in raw_points:
                if isinstance(point, dict):
                    sequence_values.extend(
                        (point.get("knowledge_id"), point.get("knowledge_name"))
                    )
                else:
                    sequence_values.append(point)
        needs_compatibility = (
            any(
                isinstance(value, (list, tuple))
                or (
                    isinstance(value, str)
                    and _looks_like_serialized_knowledge_sequence(value)
                )
                for value in sequence_values
            )
            or len([kid for kid in point_ids if kid])
            != len(set(kid for kid in point_ids if kid))
        )
        if not needs_compatibility:
            continue
        before = {key: copy.deepcopy(question.get(key)) for key in tracked}
        _normalize_question_knowledge_fields(question)
        after = {key: question.get(key) for key in tracked}
        changed = changed or before != after
    return changed


def iter_effective_rubric_items(
    payload: Mapping[str, object],
) -> Iterator[tuple[str, Mapping[str, object], Mapping[str, object]]]:
    for item_ref, _parent_ref, raw_question, raw_item in iter_effective_rubric_item_refs(payload):
        yield item_ref, raw_question, raw_item


def iter_effective_rubric_item_refs(
    payload: Mapping[str, object],
) -> Iterator[tuple[str, str, Mapping[str, object], Mapping[str, object]]]:
    rubric = payload.get("rubric") if isinstance(payload.get("rubric"), Mapping) else payload
    questions = rubric.get("questions") if isinstance(rubric, Mapping) else None
    if not isinstance(questions, list):
        return
    for question_index, raw_question in enumerate(questions, start=1):
        if not isinstance(raw_question, Mapping):
            continue
        question_ref = str(
            raw_question.get("question_id")
            or raw_question.get("id")
            or raw_question.get("number")
            or f"Q{question_index}"
        ).strip()
        parts = raw_question.get("parts")
        emitted = False
        if isinstance(parts, list) and parts:
            for part_index, raw_part in enumerate(parts, start=1):
                if not isinstance(raw_part, Mapping):
                    continue
                part_ref = str(
                    raw_part.get("part_id")
                    or raw_part.get("question_id")
                    or f"{question_ref}.{part_index}"
                ).strip()
                emitted = True
                yield part_ref, question_ref, raw_question, raw_part
        if not emitted:
            yield question_ref, question_ref, raw_question, raw_question


def iter_rubric_skill_requests(
    payload: Mapping[str, object],
    *,
    grading_session_id: str,
) -> Iterator[tuple[str, SkillRole, SkillResolutionRequest]]:
    rubric = payload.get("rubric") if isinstance(payload.get("rubric"), Mapping) else payload
    grade = str(
        (rubric.get("grade") if isinstance(rubric, Mapping) else "")
        or payload.get("grade")
        or ""
    ).strip()

    for item_ref, question, item in iter_effective_rubric_items(payload):
        measured = _rubric_knowledge_points(item)
        if not measured and item is not question:
            measured = _rubric_knowledge_points(question)
        supporting = _rubric_supporting_points(item)
        context_text = _rubric_context_text(question, item)
        for role, points in (
            (SkillRole.MEASURED, measured),
            (SkillRole.SUPPORTING, supporting),
        ):
            for raw_id, label in points:
                if not label:
                    continue
                yield (
                    item_ref,
                    role,
                    SkillResolutionRequest(
                        source_type="assessment_item",
                        source_ref=f"{grading_session_id}:{item_ref}",
                        raw_label=label,
                        stable_key_hint=raw_id,
                        grade=grade,
                        question_text=str(
                            question.get("stem_summary")
                            or question.get("question_text")
                            or question.get("text")
                            or ""
                        ),
                        rubric_text=context_text,
                        existing_tags=tuple(
                            value
                            for value in (raw_id, label)
                            if str(value or "").strip()
                        ),
                    ),
                )


def _rubric_knowledge_points(item: Mapping[str, object]) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    raw_points = item.get("knowledge_points")
    if isinstance(raw_points, list):
        for point in raw_points:
            if isinstance(point, Mapping):
                raw_id = str(point.get("knowledge_id") or point.get("id") or "").strip()
                label = str(
                    point.get("knowledge_name")
                    or point.get("name")
                    or point.get("label")
                    or ""
                ).strip()
            else:
                raw_id = ""
                label = str(point or "").strip()
            if label:
                result.append((raw_id, label))
    elif isinstance(raw_points, str) and raw_points.strip():
        result.append((str(item.get("knowledge_id") or "").strip(), raw_points.strip()))
    if not result:
        label = str(
            item.get("knowledge_name")
            or item.get("knowledge_label")
            or item.get("knowledge_text")
            or ""
        ).strip()
        if label:
            result.append((str(item.get("knowledge_id") or "").strip(), label))
    return _unique_rubric_points(result)


def _rubric_supporting_points(item: Mapping[str, object]) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    for field_name in ("prerequisite_points", "supporting_skills"):
        values = item.get(field_name)
        if isinstance(values, str):
            values = [values]
        if not isinstance(values, list):
            continue
        for value in values:
            if isinstance(value, Mapping):
                raw_id = str(value.get("knowledge_id") or value.get("id") or "").strip()
                label = str(
                    value.get("knowledge_name")
                    or value.get("name")
                    or value.get("label")
                    or ""
                ).strip()
            else:
                raw_id = ""
                label = str(value or "").strip()
            if label:
                result.append((raw_id, label))
    return _unique_rubric_points(result)


def _unique_rubric_points(values: list[tuple[str, str]]) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for raw_id, label in values:
        identity = (raw_id.casefold(), label.casefold())
        if identity in seen:
            continue
        seen.add(identity)
        result.append((raw_id, label))
    return result


def _rubric_context_text(
    question: Mapping[str, object],
    item: Mapping[str, object],
) -> str:
    values: list[str] = []
    for source in (question, item):
        for key in ("stem_summary", "knowledge_name", "knowledge_points"):
            value = source.get(key)
            if isinstance(value, str) and value.strip():
                values.append(value.strip())
        steps = source.get("steps")
        if isinstance(steps, list):
            for step in steps:
                if not isinstance(step, Mapping):
                    continue
                for key in ("core_goal", "required_elements"):
                    value = step.get(key)
                    if isinstance(value, list):
                        values.extend(str(item).strip() for item in value if str(item).strip())
                    elif str(value or "").strip():
                        values.append(str(value).strip())
    return "；".join(dict.fromkeys(values))


def _promote_nested_question_knowledge(question: dict[str, Any]) -> None:
    parts = question.get("parts")
    if not isinstance(parts, list):
        return

    current_id = str(question.get("knowledge_id") or "").strip()
    current_name = str(question.get("knowledge_name") or "").strip()
    collected: list[dict[str, str]] = []
    raw_top_points = question.get("knowledge_points")
    if isinstance(raw_top_points, list):
        collected.extend(item for item in raw_top_points if isinstance(item, dict))

    for part in parts:
        if not isinstance(part, dict):
            continue
        part_id = str(part.get("knowledge_id") or part.get("knowledge_name") or "").strip()
        part_name = str(part.get("knowledge_name") or "").strip()
        if part_id and part_name:
            collected.append({"knowledge_id": part_id, "knowledge_name": part_name})
        raw_points = part.get("knowledge_points")
        if isinstance(raw_points, str) and raw_points.strip():
            collected.append(
                {
                    "knowledge_id": part_id or part_name or "DETAIL",
                    "knowledge_name": raw_points.strip(),
                }
            )
        elif isinstance(raw_points, list):
            for raw_point in raw_points:
                if isinstance(raw_point, dict):
                    collected.append(raw_point)
                elif str(raw_point or "").strip():
                    collected.append(
                        {
                            "knowledge_id": part_id or part_name or "DETAIL",
                            "knowledge_name": str(raw_point).strip(),
                        }
                    )

        if (not current_name or current_id in {"", "UNKNOWN"}) and part_name:
            current_name = part_name
            current_id = part_id or part_name

    if current_name:
        question["knowledge_name"] = current_name
    if current_id:
        question["knowledge_id"] = current_id
    if collected:
        question["knowledge_points"] = collected


def force_payload_total_score(payload: dict[str, Any], target_total: float = 100.0) -> None:
    rubric = payload.get("rubric") if isinstance(payload, dict) else None
    if not isinstance(rubric, dict):
        return
    questions = rubric.get("questions")
    if not isinstance(questions, list) or not questions:
        rubric["total_score"] = target_total
        return

    current_total = sum(_safe_float(q.get("max_score"), 0.0) for q in questions if isinstance(q, dict))
    if current_total <= 0:
        each = round(target_total / len(questions), 2)
        for question in questions:
            if isinstance(question, dict):
                _scale_question_to_score(question, each)
        _fix_question_sum(questions, target_total)
    elif abs(current_total - target_total) > 0.01:
        ratio = target_total / current_total
        for question in questions:
            if isinstance(question, dict):
                _scale_question_scores(question, ratio)
        _fix_question_sum(questions, target_total)

    enforce_integer_scores_by_type(
        questions,
        target_total=int(target_total),
        max_question_score=MAX_QUESTION_SCORE,
    )
    # Second pass: fix any +-1 rounding drift produced by integer allocation
    _fix_question_sum(questions, target_total)
    for question in questions:
        if isinstance(question, dict):
            _ensure_solution_hard_rules(question)
    rubric["total_score"] = target_total
    payload.setdefault("meta", {}).setdefault("warnings", [])


def _scale_question_scores(question: dict[str, Any], ratio: float) -> None:
    new_score = round(_safe_float(question.get("max_score"), 0.0) * ratio, 2)
    _scale_question_to_score(question, new_score)


def _scale_question_to_score(question: dict[str, Any], new_score: float) -> None:
    old_score = _safe_float(question.get("max_score"), 0.0)
    question["max_score"] = round(float(new_score), 2)
    parts = question.get("parts")
    if not isinstance(parts, list) or not parts:
        _normalize_rubric_question(question, {})
        return

    ratio = (float(new_score) / old_score) if old_score > 0 else (1.0 / len(parts))
    for part in parts:
        if not isinstance(part, dict):
            continue
        part["part_score"] = round(_safe_float(part.get("part_score"), 0.0) * ratio, 2)
        steps = part.get("steps")
        if isinstance(steps, list):
            for step in steps:
                if isinstance(step, dict):
                    step["step_score"] = round(_safe_float(step.get("step_score"), 0.0) * ratio, 2)
        _force_step_total(part)
    _force_part_total(question)
    _scale_deduction_policy(question, ratio)


def _scale_deduction_policy(question: dict[str, Any], ratio: float) -> None:
    policies = question.get("deduction_policy")
    if isinstance(policies, list):
        for policy in policies:
            if isinstance(policy, dict) and "max_deduction" in policy:
                policy["max_deduction"] = round(_safe_float(policy.get("max_deduction"), 0.0) * ratio, 2)


def _fix_question_sum(questions: list[Any], target_total: float) -> None:
    valid_questions = [q for q in questions if isinstance(q, dict)]
    if not valid_questions:
        return
    total = sum(_safe_float(q.get("max_score"), 0.0) for q in valid_questions)
    diff = round(target_total - total, 2)
    if abs(diff) > 0.001:
        last = valid_questions[-1]
        last["max_score"] = round(_safe_float(last.get("max_score"), 0.0) + diff, 2)
        _force_part_total(last)


def _normalize_answer_item(answer_item: dict[str, Any], rubric_question: dict[str, Any]) -> None:
    _coerce_answer_item_aliases(answer_item)
    qid = str(rubric_question.get("question_id") or answer_item.get("question_id") or "")
    answer_item["question_id"] = qid
    direct_answers = _extract_direct_answer_values(answer_item)
    canonical = (
        answer_item.get("canonical_answer")
        or answer_item.get("answer")
        or answer_item.get("standard_answer")
        or answer_item.get("correct_answer")
        or answer_item.get("绛旀")
        or (direct_answers[0] if direct_answers else "")
        or rubric_question.get("canonical_answer")
        or rubric_question.get("answer")
        or rubric_question.get("correct_answer")
        or ""
    )
    answer_item["canonical_answer"] = str(canonical)
    if bool(answer_item.get("_manual_accepted_forms")):
        answer_item["accepted_forms"] = _string_list(answer_item.get("accepted_forms"))
    else:
        answer_item["accepted_forms"] = _string_list(
            answer_item.get("accepted_forms")
            or answer_item.get("equivalent_answers")
            or answer_item.get("aliases")
            or direct_answers
            or rubric_question.get("accepted_forms")
            or rubric_question.get("equivalent_answers")
            or [canonical]
        )
    answer_item["method_variants"] = _method_variants(
        answer_item.get("method_variants") or rubric_question.get("method_variants") or []
    )
    if not isinstance(answer_item.get("parts"), list) or not answer_item["parts"]:
        answer_item["parts"] = [
            {
                "part_id": qid,
                "answer": str(canonical),
                "analysis": str(answer_item.get("analysis") or rubric_question.get("analysis") or ""),
                "step_milestones": _string_list(answer_item.get("step_milestones") or rubric_question.get("step_milestones")),
            }
        ]
    else:
        for idx, part in enumerate(answer_item["parts"], start=1):
            if not isinstance(part, dict):
                continue
            _coerce_answer_part_aliases(part)
            part_direct_answers = _extract_direct_answer_values(part)
            part["answer_values"] = _string_list(part_direct_answers)
            part.setdefault("part_id", qid if len(answer_item["parts"]) == 1 else f"{qid}({idx})")
            part_answer = (
                part.get("answer")
                or part.get("canonical_answer")
                or part.get("standard_answer")
                or ("；".join(part["answer_values"]) if part["answer_values"] else "")
                or canonical
            )
            part["answer"] = str(part_answer)
            if len(part["answer_values"]) == 1 and not part.get("accepted_forms"):
                part["accepted_forms"] = _string_list(part_direct_answers)
            part.setdefault("analysis", "")
            if not isinstance(part.get("step_milestones"), list):
                part["step_milestones"] = _string_list(part.get("step_milestones"))

    if not str(answer_item.get("canonical_answer") or "").strip():
        qtype = str(rubric_question.get("question_type") or "").strip()
        if qtype in {"choice", "fill_blank", "judgement", "true_false", "direct_answer"}:
            first_part_answer = next(
                (
                    str(part.get("answer") or "").strip()
                    for part in answer_item.get("parts", [])
                    if isinstance(part, dict) and str(part.get("answer") or "").strip()
                ),
                "",
            )
            if first_part_answer:
                answer_item["canonical_answer"] = first_part_answer
                answer_item["accepted_forms"] = _string_list(answer_item.get("accepted_forms")) or [first_part_answer]


def _coerce_answer_item_aliases(answer_item: dict[str, Any]) -> None:
    if not isinstance(answer_item.get("parts"), list) or not answer_item.get("parts"):
        for key in ("answer_parts", "sub_answers", "subquestions"):
            value = answer_item.get(key)
            if isinstance(value, list) and value:
                answer_item["parts"] = value
                break


def _coerce_answer_part_aliases(part: dict[str, Any]) -> None:
    if not str(part.get("answer") or "").strip():
        for key in ("answer_content", "standard_answer", "canonical_answer", "correct_answer"):
            value = part.get(key)
            if value is not None and str(value).strip():
                part["answer"] = str(value).strip()
                break


def _extract_direct_answer_values(item: dict[str, Any]) -> list[Any]:
    values: list[Any] = []
    direct = item.get("direct_answer")
    if isinstance(direct, dict):
        for key in ["accepted_forms", "answers", "answer", "value", "values"]:
            raw = direct.get(key)
            if raw is None:
                continue
            if isinstance(raw, list):
                values.extend(raw)
            else:
                values.append(raw)
    elif isinstance(direct, list):
        values.extend(direct)
    elif direct is not None:
        values.append(direct)
    for key in ("answers", "values"):
        raw = item.get(key)
        if isinstance(raw, list):
            for value in raw:
                if isinstance(value, dict):
                    nested = next(
                        (
                            value.get(alias)
                            for alias in ("value", "answer", "answer_content", "standard_answer", "canonical_answer")
                            if value.get(alias) is not None
                        ),
                        None,
                    )
                    if nested is not None:
                        values.append(nested)
                elif value is not None:
                    values.append(value)
        elif raw is not None:
            values.append(raw)
    answer_content = item.get("answer_content")
    if answer_content is not None:
        values.append(answer_content)
    return values


def _should_treat_as_direct_answer_question(
    qtype: str,
    question: dict[str, Any],
    answer_item: dict[str, Any],
) -> bool:
    normalized_type = str(qtype or "").strip().lower()
    if normalized_type in {"choice", "fill_blank", "judgement", "true_false"}:
        return False
    mode_text = " ".join(
        str(question.get(key) or "")
        for key in ["grading_mode", "scoring_type", "scoring_policy", "scoring_rule", "description"]
    ).lower()
    has_direct_mode = "direct_answer" in mode_text or "鍏ㄥ鍏ㄩ敊" in mode_text
    has_direct_answer = bool(_extract_direct_answer_values(answer_item) or _string_list(answer_item.get("accepted_forms")))
    has_solution_structure = bool(question.get("proof_obligations")) or any(
        isinstance(part, dict) and len(part.get("steps") or []) > 1
        for part in (question.get("parts") if isinstance(question.get("parts"), list) else [])
    )
    return has_direct_answer and has_direct_mode and not has_solution_structure


def _augment_answer_equivalences(answer_item: dict[str, Any], qtype: str) -> None:
    """Add deterministic local equivalents for objective/short-answer items."""
    if qtype == "choice":
        return

    canonical = str(answer_item.get("canonical_answer") or "").strip()
    top_level_manually_edited = bool(answer_item.get("_manual_accepted_forms"))
    parts = answer_item.get("parts")
    subjective_multi_part = (
        qtype in {"proof", "calculation", "comprehensive"}
        and isinstance(parts, list)
        and len([part for part in parts if isinstance(part, dict)]) > 1
    )
    if isinstance(parts, list):
        for part in parts:
            if not isinstance(part, dict):
                continue
            part_answer = str(part.get("answer") or "").strip()
            if part_answer:
                normalized_part_answer = _normalize_serialized_answer_list(part_answer)
                serialized_part_answer = normalized_part_answer != part_answer
                if serialized_part_answer:
                    part["answer"] = normalized_part_answer
                    part_answer = normalized_part_answer
                if not bool(part.get("_manual_accepted_forms")):
                    existing_part_forms = [] if serialized_part_answer else [
                        value
                        for value in _string_list(part.get("accepted_forms"))
                        if not _looks_like_serialized_answer_list(value)
                    ]
                    part["accepted_forms"] = merge_equivalent_forms(existing_part_forms, part_answer, max_forms=16)

    if not top_level_manually_edited:
        existing_top_forms = [] if subjective_multi_part else [
            value
            for value in _string_list(answer_item.get("accepted_forms"))
            if not _looks_like_serialized_answer_list(value)
        ]
        answer_item["accepted_forms"] = merge_equivalent_forms(existing_top_forms, canonical, max_forms=32)


def _record_complete_answer_set_rule(answer_item: dict[str, Any], qtype: str) -> None:
    if qtype != "fill_blank":
        return
    canonical = str(answer_item.get("canonical_answer") or "").strip()
    required_values = complete_answer_set_values(canonical)
    if not required_values:
        return
    rule = {
        "match_mode": "complete_set",
        "required_values": required_values,
        "order_sensitive": False,
        "allow_extra_values": False,
        "partial_credit": False,
    }
    answer_item.update(rule)
    parts = answer_item.get("parts")
    if isinstance(parts, list) and len(parts) == 1 and isinstance(parts[0], dict):
        parts[0].update(rule)


def _sanitize_choice_answer_forms(answer_item: dict[str, Any]) -> None:
    canonical = str(answer_item.get("canonical_answer") or "").strip().upper()
    if not re.fullmatch(r"[A-D]", canonical):
        return
    answer_item["canonical_answer"] = canonical
    answer_item["accepted_forms"] = [canonical]
    parts = answer_item.get("parts")
    if not isinstance(parts, list):
        return
    for part in parts:
        if not isinstance(part, dict):
            continue
        part_answer = str(part.get("answer") or canonical).strip().upper()
        if not re.fullmatch(r"[A-D]", part_answer):
            part_answer = canonical
        part["answer"] = part_answer
        part["accepted_forms"] = [part_answer]


def _normalize_rubric_question(question: dict[str, Any], answer_item: dict[str, Any]) -> None:
    qid = str(question.get("question_id") or "")
    qtype = str(question.get("question_type") or "comprehensive")
    max_score = _safe_float(question.get("max_score"), 0.0)
    _align_solution_parts_with_answer_parts(question, answer_item, qtype, max_score)
    if not isinstance(question.get("parts"), list) or not question["parts"]:
        question["parts"] = [
            {
                "part_id": qid,
                "part_score": max_score,
                "steps": [
                    {
                        "step_id": "S1",
                        "step_score": max_score,
                        "core_goal": _default_core_goal(qtype, answer_item),
                        "required_elements": _default_required_elements(qtype, answer_item),
                        "allow_alternative_methods": qtype not in {"choice"},
                    }
                ],
                "presentation_rules": [],
            }
        ]
    else:
        part_count = len(question["parts"])
        for idx, part in enumerate(question["parts"], start=1):
            if not isinstance(part, dict):
                part = {"part_id": qid if part_count == 1 else f"{qid}({idx})", "answer": str(part)}
                question["parts"][idx - 1] = part
            part_score = _safe_float(
                part.get("part_score")
                or part.get("score")
                or part.get("points"),
                max_score / max(part_count, 1),
            )
            part["part_id"] = str(part.get("part_id") or (qid if part_count == 1 else f"{qid}({idx})"))
            part["part_score"] = part_score
            step_alias = _best_step_alias(part)
            if step_alias is not part.get("steps"):
                part["steps"] = step_alias

            if not isinstance(part.get("steps"), list) or not part["steps"]:
                part["steps"] = [
                    {
                        "step_id": "S1",
                        "step_score": part_score,
                        "core_goal": _default_core_goal(qtype, answer_item),
                        "required_elements": _default_required_elements(qtype, answer_item),
                        "allow_alternative_methods": qtype not in {"choice"},
                    }
                ]
            else:
                for sidx, step in enumerate(part["steps"], start=1):
                    if not isinstance(step, dict):
                        step = {"core_goal": str(step)}
                        part["steps"][sidx - 1] = step
                    step["step_id"] = str(step.get("step_id") or f"S{sidx}")
                    # Avoid overriding explicit 0.0 with fallback by using explicit None checks
                    raw_score = step.get("step_score")
                    if raw_score is None:
                        raw_score = step.get("score")
                    if raw_score is None:
                        raw_score = step.get("points")
                    
                    step["step_score"] = _safe_float(
                        raw_score,
                        part_score / max(len(part["steps"]), 1),
                    )
                    step["core_goal"] = _specific_step_goal(
                        step,
                        _default_core_goal(qtype, answer_item),
                    )
                    required_elements = _string_list(step.get("required_elements"))
                    if not required_elements or all(value in _GENERIC_STEP_GOALS for value in required_elements):
                        goal = str(step.get("core_goal") or "").strip()
                        step["required_elements"] = (
                            [goal]
                            if goal and goal not in _GENERIC_STEP_GOALS
                            else _default_required_elements(qtype, answer_item)
                        )
                    else:
                        step["required_elements"] = required_elements
                    step["allow_alternative_methods"] = bool(step.get("allow_alternative_methods", qtype not in {"choice"}))
            if not isinstance(part.get("presentation_rules"), list):
                part["presentation_rules"] = []
            _force_step_total(part)
    _force_part_total(question)


def _first_list_value(node: dict[str, Any], *keys: str) -> list[Any]:
    for key in keys:
        value = node.get(key)
        if isinstance(value, list) and value:
            return value
    return []


def _best_step_alias(part: dict[str, Any]) -> list[Any]:
    existing = _first_list_value(part, "steps")
    candidates = [
        value
        for key in ("scoring_steps", "criteria", "rubric", "score_points", "points")
        if isinstance((value := part.get(key)), list) and value
    ]
    visual_requirements = _string_list(part.get("visual_requirements"))
    visual_steps = [
        {"core_goal": value, "required_elements": [value]}
        for value in visual_requirements
    ]
    if not existing:
        return next(iter(candidates), visual_steps)
    if not _steps_are_generic(existing):
        return existing
    for candidate in candidates:
        if not _steps_are_generic(candidate):
            return candidate
    return visual_steps or existing


def _steps_are_generic(steps: list[Any]) -> bool:
    descriptions: list[str] = []
    for step in steps:
        if isinstance(step, dict):
            descriptions.append(_specific_step_goal(step, ""))
        else:
            descriptions.append(str(step or "").strip())
    return bool(descriptions) and all(description in _GENERIC_STEP_GOALS for description in descriptions)


_GENERIC_STEP_GOALS = {
    "",
    "完成必要的推理或计算步骤",
    "填写正确或等价的答案",
    "选择正确的选项",
    "合理的推理过程",
    "正确的结论",
}


def _specific_step_goal(step: dict[str, Any], default: str) -> str:
    values = [
        str(step.get(key) or "").strip()
        for key in ("core_goal", "goal", "criterion", "description", "step_description", "desc", "title", "requirement")
    ]
    return next((value for value in values if value and value not in _GENERIC_STEP_GOALS), None) or next(
        (value for value in values if value),
        default,
    )


def _enforce_objective_question_rules(question: dict[str, Any], answer_item: dict[str, Any]) -> None:
    qtype = str(question.get("question_type") or "").strip().lower()
    if qtype not in {"choice", "fill_blank", "judgement", "true_false", "direct_answer"}:
        return

    if qtype == "choice":
        _sanitize_choice_answer_forms(answer_item)
    _record_complete_answer_set_rule(answer_item, qtype)

    top_answers = _string_list(answer_item.get("accepted_forms"))
    canonical = str(answer_item.get("canonical_answer") or "").strip()
    if canonical and canonical not in top_answers:
        top_answers.insert(0, canonical)
    answer_parts = answer_item.get("parts")
    if not isinstance(answer_parts, list):
        answer_parts = []

    parts = question.get("parts")
    if not isinstance(parts, list):
        return
    for index, part in enumerate(parts):
        if not isinstance(part, dict):
            continue
        explicit_mode = str(part.get("response_mode") or "").strip()
        answer_part = answer_parts[index] if index < len(answer_parts) and isinstance(answer_parts[index], dict) else {}
        independent_answer_values = _string_list(answer_part.get("answer_values"))
        is_independent_fill = (
            qtype == "fill_blank"
            and explicit_mode == "short_answer_points"
            and len(independent_answer_values) > 1
        )
        part["response_mode"] = "short_answer_points" if is_independent_fill else "exact_objective"
        part["presentation_rules"] = []
        part["require_final_answer"] = False
        part["answer_only_max_score"] = int(round(_safe_float(part.get("part_score"), 0.0)))

        part_answers: list[str] = []
        if answer_part:
            part_answers = _string_list(answer_part.get("accepted_forms"))
            part_answer = str(answer_part.get("answer") or "").strip()
            if part_answer and part_answer not in part_answers:
                part_answers.insert(0, part_answer)
        required_answers = list(dict.fromkeys([*part_answers, *top_answers]))

        steps = part.get("steps")
        if not isinstance(steps, list) or not steps:
            continue
        if not is_independent_fill:
            first_step = next((step for step in steps if isinstance(step, dict)), {})
            first_step["step_id"] = str(first_step.get("step_id") or "S1")
            first_step["step_score"] = int(round(_safe_float(part.get("part_score"), 0.0)))
            first_step["core_goal"] = _default_core_goal(qtype, answer_item)
            first_step["required_elements"] = required_answers
            first_step["allow_alternative_methods"] = qtype != "choice"
            part["steps"] = [first_step]
            continue
        for step in steps:
            if not isinstance(step, dict):
                continue
            goal = str(step.get("core_goal") or "").strip()
            if not goal or goal == "完成必要的推理或计算步骤":
                step["core_goal"] = _default_core_goal(qtype, answer_item)
            if not _string_list(step.get("required_elements")):
                step["required_elements"] = required_answers
            step["allow_alternative_methods"] = qtype != "choice"

    question["require_final_answer"] = False
    question["answer_only_max_score"] = int(round(_safe_float(question.get("max_score"), 0.0)))
    policies = question.get("deduction_policy")
    if isinstance(policies, list):
        question["deduction_policy"] = [
            policy
            for policy in policies
            if not (
                isinstance(policy, dict)
                and str(policy.get("policy_id") or "") in {"answer_only_process_missing", "core_process_missing"}
            )
        ]
    question["answer_presentation_policy"] = {
        "require_final_answer": False,
        "answer_only_max_score": question["answer_only_max_score"],
        "note": "客观题仅按标准答案或等价答案判分，不要求过程证据。",
    }


def _align_answer_parts_to_rubric_parts(
    question: dict[str, Any],
    answer_item: dict[str, Any],
) -> None:
    rubric_parts = question.get("parts")
    answer_parts = answer_item.get("parts")
    if not isinstance(rubric_parts, list) or not isinstance(answer_parts, list):
        return
    valid_rubric_parts = [part for part in rubric_parts if isinstance(part, dict)]
    valid_answer_parts = [part for part in answer_parts if isinstance(part, dict)]
    if not valid_rubric_parts or len(valid_rubric_parts) != len(valid_answer_parts):
        return

    rubric_ids = [str(part.get("part_id") or "") for part in valid_rubric_parts]
    answer_ids = [str(part.get("part_id") or "") for part in valid_answer_parts]
    if rubric_ids == answer_ids:
        return
    if len(rubric_ids) > 1 and set(rubric_ids) & set(answer_ids):
        return
    for rubric_part, answer_part in zip(valid_rubric_parts, valid_answer_parts):
        answer_part["part_id"] = str(rubric_part.get("part_id") or answer_part.get("part_id") or "")


def _align_step_required_elements_with_answer_values(
    question: dict[str, Any],
    answer_item: dict[str, Any],
) -> None:
    rubric_parts = question.get("parts")
    answer_parts = answer_item.get("parts")
    if not isinstance(rubric_parts, list) or not isinstance(answer_parts, list):
        return
    answer_part_map = {
        str(part.get("part_id") or ""): part
        for part in answer_parts
        if isinstance(part, dict)
    }
    for part_index, rubric_part in enumerate(rubric_parts):
        if not isinstance(rubric_part, dict):
            continue
        answer_part = answer_part_map.get(str(rubric_part.get("part_id") or ""))
        if not isinstance(answer_part, dict) and part_index < len(answer_parts):
            candidate = answer_parts[part_index]
            answer_part = candidate if isinstance(candidate, dict) else None
        if not isinstance(answer_part, dict):
            continue

        raw_answers = answer_part.get("answers")
        answer_values = _string_list(answer_part.get("answer_values"))
        steps = rubric_part.get("steps")
        if not answer_values or not isinstance(steps, list):
            continue

        values_by_id: dict[str, str] = {}
        if isinstance(raw_answers, list):
            for raw_answer in raw_answers:
                if not isinstance(raw_answer, dict):
                    continue
                score_point_id = str(raw_answer.get("score_point_id") or raw_answer.get("step_id") or "").strip()
                value = str(raw_answer.get("value") or raw_answer.get("answer") or "").strip()
                if score_point_id and value:
                    values_by_id[score_point_id] = value

        for step_index, step in enumerate(steps):
            if not isinstance(step, dict):
                continue
            step_key = str(step.get("score_point_id") or step.get("step_id") or "").strip()
            value = values_by_id.get(step_key)
            if not value and step_index < len(answer_values):
                value = answer_values[step_index]
            if value:
                step["required_elements"] = merge_equivalent_forms([], value, max_forms=16)


def _align_solution_parts_with_answer_parts(
    question: dict[str, Any],
    answer_item: dict[str, Any],
    qtype: str,
    max_score: float,
) -> None:
    if qtype not in {"calculation", "proof", "comprehensive"}:
        return
    answer_parts = answer_item.get("parts")
    if not isinstance(answer_parts, list) or len(answer_parts) <= 1:
        return
    rubric_parts = question.get("parts")
    if isinstance(rubric_parts, list) and len(rubric_parts) > 1:
        return
    part_scores = _allocate_scores(max_score, len(answer_parts))
    next_parts: list[dict[str, Any]] = []
    for idx, (answer_part, part_score) in enumerate(zip(answer_parts, part_scores), start=1):
        if not isinstance(answer_part, dict):
            continue
        part_id = str(answer_part.get("part_id") or f"{question.get('question_id')}({idx})")
        answer = str(answer_part.get("answer") or "").strip()
        next_parts.append(
            {
                "part_id": part_id,
                "part_score": part_score,
                "steps": [
                    {
                        "step_id": "S1",
                        "step_score": part_score,
                        "core_goal": f"完成 {part_id} 的答题要求",
                        "required_elements": [answer] if answer else ["合理的推理过程", "正确的结论"],
                        "allow_alternative_methods": True,
                    }
                ],
                "presentation_rules": [],
            }
        )
    if next_parts:
        question["parts"] = next_parts


def _allocate_scores(total: float, count: int) -> list[float]:
    if count <= 0:
        return []
    total_int = int(round(total))
    base = total_int // count
    remainder = total_int - base * count
    return [base + (1 if idx < remainder else 0) for idx in range(count)]


_VALID_RESPONSE_MODES = {
    "exact_objective",
    "short_answer_points",
    "process_required",
    "visual_construction",
}
_NON_PROCESS_RESPONSE_MODES = {
    "exact_objective",
    "short_answer_points",
    "visual_construction",
}


def _infer_part_response_mode(question: dict[str, Any], part: dict[str, Any]) -> str:
    explicit = str(part.get("response_mode") or "").strip().lower()
    aliases = {
        "direct_answer": "short_answer_points",
        "answer_only": "short_answer_points",
        "objective": "exact_objective",
        "construction": "visual_construction",
        "proof": "process_required",
    }
    explicit = aliases.get(explicit, explicit)
    if explicit in _VALID_RESPONSE_MODES:
        return explicit

    qtype = str(question.get("question_type") or "").strip().lower()
    if qtype in {"choice", "fill_blank", "judgement", "true_false", "direct_answer"}:
        return "exact_objective"

    text_parts = [
        str(part.get(key) or "")
        for key in ("core_goal", "description", "part_title", "response_requirement", "answer_requirement")
    ]
    steps = part.get("steps")
    if isinstance(steps, list):
        for step in steps:
            if not isinstance(step, dict):
                continue
            text_parts.extend(
                str(step.get(key) or "")
                for key in ("core_goal", "goal", "criterion", "description")
            )
            required = step.get("required_elements")
            if isinstance(required, list):
                text_parts.extend(str(item or "") for item in required)
    text = " ".join(text_parts)
    if re.search(r"尺规|作图|画出|绘制|保留.{0,6}(?:痕迹|过程)", text):
        return "visual_construction"
    if re.search(r"直接写出|只需.{0,6}答案|无需.{0,6}过程|填空|填写|选择|判断|写出.{0,8}(?:结果|关系|答案|数值)", text):
        return "short_answer_points"
    return "process_required"


def _ensure_solution_hard_rules(question: dict[str, Any]) -> None:
    """Persist process rules without applying them to direct-answer subquestions."""
    qtype = str(question.get("question_type") or "")
    if qtype not in {"proof", "calculation", "comprehensive"}:
        return

    max_score = _safe_float(question.get("max_score"), 0.0)
    default_answer_only = max(1, int(round(max_score * 0.25))) if max_score > 0 else 1

    if "require_final_answer" not in question:
        # Proof and pure calculation questions usually do not need an extra "绛?;
        # general comprehensive answers often do.
        question["require_final_answer"] = qtype == "comprehensive"
    question["require_final_answer"] = bool(question.get("require_final_answer"))

    parts = question.get("parts")
    part_answer_only_total = 0
    has_process_part = False
    if isinstance(parts, list):
        for part in parts:
            if not isinstance(part, dict):
                continue

            part_score = _safe_float(part.get("part_score") or part.get("max_score"), max_score)
            response_mode = _infer_part_response_mode(question, part)
            part["response_mode"] = response_mode

            rules = part.get("presentation_rules")
            if not isinstance(rules, list):
                rules = []
                part["presentation_rules"] = rules
            _remove_rule_by_id(rules, "final_answer_required")

            if response_mode in _NON_PROCESS_RESPONSE_MODES:
                part["require_final_answer"] = False
                part["answer_only_max_score"] = int(round(part_score))
            else:
                has_process_part = True
                part_default_answer_only = max(1, int(round(part_score * 0.25))) if part_score > 0 else 1
                if "require_final_answer" not in part:
                    part["require_final_answer"] = question["require_final_answer"]
                part["require_final_answer"] = bool(part.get("require_final_answer"))
                if "answer_only_max_score" not in part:
                    part["answer_only_max_score"] = part_default_answer_only
                part["answer_only_max_score"] = max(
                    0,
                    min(
                        int(round(_safe_float(part.get("answer_only_max_score"), part_default_answer_only))),
                        int(round(part_score)),
                    ),
                )
                if part["require_final_answer"]:
                    rules.append(
                        {
                            "rule_id": "final_answer_required",
                            "rule": "启用此策略时，必须包含最终答案或明确的结论；如果缺失或不完整则扣分。",
                            "max_deduction": 1,
                        }
                    )
            part_answer_only_total += int(part["answer_only_max_score"])

    if isinstance(parts, list) and parts:
        answer_only_max = part_answer_only_total
    else:
        answer_only_raw = question.get("answer_only_max_score")
        answer_only_max = int(round(_safe_float(answer_only_raw, default_answer_only)))
        has_process_part = True
    if max_score > 0:
        answer_only_max = max(0, min(answer_only_max, int(round(max_score))))
    question["answer_only_max_score"] = answer_only_max

    policies = question.get("deduction_policy")
    if not isinstance(policies, list):
        policies = []
        question["deduction_policy"] = policies

    policies[:] = [
        policy
        for policy in policies
        if not (
            isinstance(policy, dict)
            and str(policy.get("policy_id") or "") in {"answer_only_process_missing", "core_process_missing"}
        )
    ]
    if has_process_part:
        _upsert_policy(
            policies,
            {
                "policy_id": "answer_only_process_missing",
                "issue": f"需要过程的小问仅有最终答案时，整题最高只给 {answer_only_max} 分",
                "max_deduction": max(0, int(round(max_score)) - answer_only_max),
                "severity": "major",
            },
        )
        _upsert_policy(
            policies,
            {
                "policy_id": "core_process_missing",
                "issue": "需要过程的小问缺失关键步骤、证明逻辑或推理链条；扣除对应步骤分",
                "max_deduction": int(round(max_score)),
                "severity": "fatal",
            },
        )

    question["answer_presentation_policy"] = {
        "require_final_answer": question["require_final_answer"],
        "answer_only_max_score": answer_only_max,
        "note": "仅 response_mode=process_required 的小问需要过程证据；其他小问按正确答案项或图形要求给分。",
    }


def _upsert_policy(policies: list[Any], new_policy: dict[str, Any]) -> None:
    policy_id = str(new_policy.get("policy_id") or "")
    for idx, policy in enumerate(policies):
        if isinstance(policy, dict) and str(policy.get("policy_id") or "") == policy_id:
            policies[idx] = {**policy, **new_policy}
            return
    policies.append(new_policy)


def _remove_rule_by_id(rules: list[Any], rule_id: str) -> None:
    rules[:] = [
        rule
        for rule in rules
        if not (isinstance(rule, dict) and str(rule.get("rule_id") or "") == rule_id)
    ]


def _default_core_goal(qtype: str, answer_item: dict[str, Any]) -> str:
    if qtype == "choice":
        return "选择正确的选项"
    if qtype == "fill_blank":
        return "填写正确或等价的答案"
    return "完成必要的推理或计算步骤"


def _default_required_elements(qtype: str, answer_item: dict[str, Any]) -> list[str]:
    canonical = str(answer_item.get("canonical_answer") or "").strip()
    if qtype in {"choice", "fill_blank"} and canonical:
        return [canonical]
    return ["合理的推理过程", "正确的结论"]


def _force_step_total(part: dict[str, Any]) -> None:
    steps = part.get("steps")
    if not isinstance(steps, list) or not steps:
        return
    part_score = _safe_float(part.get("part_score"), 0.0)
    step_sum = sum(_safe_float(step.get("step_score"), 0.0) for step in steps if isinstance(step, dict))
    diff = round(part_score - step_sum, 2)
    if abs(diff) > 0.01 and isinstance(steps[-1], dict):
        steps[-1]["step_score"] = round(_safe_float(steps[-1].get("step_score"), 0.0) + diff, 2)


def _force_part_total(question: dict[str, Any]) -> None:
    parts = question.get("parts")
    if not isinstance(parts, list) or not parts:
        return
    max_score = _safe_float(question.get("max_score"), 0.0)
    part_sum = sum(_safe_float(part.get("part_score"), 0.0) for part in parts if isinstance(part, dict))
    diff = round(max_score - part_sum, 2)
    if abs(diff) > 0.01 and isinstance(parts[-1], dict):
        parts[-1]["part_score"] = round(_safe_float(parts[-1].get("part_score"), 0.0) + diff, 2)
        _force_step_total(parts[-1])


def _safe_float(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def _string_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        if not value.strip():
            return []
        return [item.strip() for item in value.replace(",", ";").split(";") if item.strip()]
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()] if str(value).strip() else []


def _method_variants(value: Any) -> list[dict[str, str]]:
    if not isinstance(value, list):
        return []
    result: list[dict[str, str]] = []
    for idx, item in enumerate(value, start=1):
        if isinstance(item, dict):
            result.append(
                {
                    "name": str(item.get("name") or f"鏂规硶{idx}"),
                    "outline": str(item.get("outline") or item.get("description") or ""),
                }
            )
        elif str(item).strip():
            result.append({"name": f"鏂规硶{idx}", "outline": str(item).strip()})
    return result


def _normalize_question_type(value: Any) -> str:
    raw = str(value or "").strip().lower()
    if not raw:
        return "comprehensive"
    if raw in {"fill_blank", "fill-in", "fill_in_blank", "blank", "short_answer", "short-answer"}:
        return "fill_blank"
    if raw in {"choice", "single_choice", "multiple_choice", "select", "option"}:
        return "choice"
    if raw in {"calculation", "calculate", "solve", "solution", "problem_solving"}:
        return "calculation"
    if raw in {"proof", "prove"}:
        return "proof"
    if raw in {"comprehensive", "subjective", "constructed_response"}:
        return "comprehensive"
    if any(token in raw for token in ["choice", "select", "option"]):
        return "choice"
    if any(token in raw for token in ["blank", "short_answer", "fill"]):
        return "fill_blank"
    if any(token in raw for token in ["proof", "prove"]):
        return "proof"
    if any(token in raw for token in ["calculation", "calculate", "solve", "solution"]):
        return "calculation"
    return "comprehensive"


def _needs_objective_repair(payload: dict[str, Any], doc_text: str) -> bool:
    questions = payload.get("rubric", {}).get("questions", [])
    qtypes = {
        str(q.get("question_type"))
        for q in questions
        if isinstance(q, dict)
    } if isinstance(questions, list) else set()
    text = str(doc_text or "")
    doc_has_choice = any(re.search(r"A.{0,80}B.{0,80}C.{0,80}D", line) for line in text.splitlines())
    doc_has_fill = bool(re.search(r"_{2,}|[ \t]{3,}|　{1,}|（\s*）|\(\s*\)", text))
    if doc_has_choice and "choice" not in qtypes:
        return True
    if doc_has_fill and "fill_blank" not in qtypes:
        return True
    return False


def _build_objective_repair_prompt(doc_text: str, bad_payload: dict[str, Any]) -> str:
    return (
        "你上一次生成的评分标准漏掉了选择题/填空题或客观题，这是严重错误。\n"
        "请重新从 Word 原文中生成完整评分标准，必须包含全部题目。\n\n"
        "修复要求：\n"
        "1) 不要只输出 Q12/Q13 等解答题；必须补齐选择题、填空题、客观题。\n"
        "2) 选择题 question_type=choice；填空题 question_type=fill_blank。\n"
        "3) 客观题也必须在 rubric.questions 和 answer_key.questions 中逐题列出。\n"
        "4) 如果题干前半部分是题目、后半部分是答案解析，请把二者按题号匹配。\n"
        "5) 如果分值没有逐题标明，请根据总分/题型合理分配，并在 meta.warnings 写明。\n"
        "6) 不能再输出“选择题与填空题未包含在详细评分逻辑中”或类似 warning。\n\n"
        "输出格式仍必须与原要求完全一致：严格 JSON 对象，包含 rubric、answer_key、meta。\n\n"
        f"Word 原文：\n{doc_text}\n\n"
        f"上一次错误输出（供你定位遗漏，不要照抄）：\n{json.dumps(bad_payload, ensure_ascii=False)}"
    )


def _build_manual_structure_refinement_prompt(payload: dict[str, Any]) -> str:
    return (
        "您正在对教师编辑过的评分标准进行精修/对齐。请仅返回严格的 JSON 数据。\n"
        "硬性要求：\n"
        "1) 必须保留每一个已有的 rubric.questions[].question_id。\n"
        "2) 必须保留每一个已有的 parts[].part_id；绝对不能合并、删除或修改教师创建 the parts 部分。\n"
        "3) answer_key.questions[].parts 必须通过 part_id 与 rubric 中的 parts 保持一致对齐。\n"
        "4) 补全缺失的答案、accepted_forms、解析、步骤分、证明扣分项和证据链规则。\n"
        "5) 保持整张试卷总分 total_score 和各题 max_score 的总和精确等于 100。\n"
        "6) 对于证明题或计算解答题，按证据步骤步骤分进行细化，并保守地给与仅有答案无过程的得分限制。\n\n"
        f"当前评分标准 JSON：\n{json.dumps(payload, ensure_ascii=False)}"
    )


def validate_generated_config(payload: dict[str, Any]) -> None:
    normalize_generated_config_schema(payload)
    force_payload_total_score(payload, target_total=100.0)
    if not isinstance(payload, dict):
        raise ValueError("Generated config must be a JSON object")
    for top_key in ["rubric", "answer_key", "meta"]:
        if top_key not in payload:
            raise ValueError(f"Generated config is missing top-level field: {top_key}")

    rubric = payload["rubric"]
    answer_key = payload["answer_key"]
    meta = payload["meta"]
    if not isinstance(rubric, dict):
        raise ValueError("rubric must be an object")
    if not isinstance(answer_key, dict):
        raise ValueError("answer_key must be an object")
    if not isinstance(meta, dict):
        raise ValueError("meta must be an object")

    rubric_questions = rubric.get("questions")
    answer_questions = answer_key.get("questions")
    if not isinstance(rubric_questions, list) or not rubric_questions:
        raise ValueError("rubric.questions must be a non-empty list")
    if not isinstance(answer_questions, list) or not answer_questions:
        raise ValueError("answer_key.questions must be a non-empty list")

    rubric_q_map: dict[str, dict[str, Any]] = {}
    scores_by_type: dict[str, float] = {}
    for idx, item in enumerate(rubric_questions, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"rubric.questions[{idx - 1}] must be an object, got {type(item).__name__}: {item!r}")
        for key in ["question_id", "question_type", "max_score", "knowledge_id", "parts"]:
            if key not in item:
                raise ValueError(f"rubric.questions[{idx - 1}] is missing field: {key}")
        qid = str(item["question_id"])
        rubric_q_map[qid] = item
        qtype = str(item.get("question_type") or "comprehensive")
        max_score = float(item["max_score"])
        if not max_score.is_integer():
            raise ValueError(f"rubric.questions[{idx - 1}].max_score must be an integer")
        if max_score > MAX_QUESTION_SCORE:
            raise ValueError(
                f"rubric.questions[{idx - 1}].max_score must not exceed {MAX_QUESTION_SCORE}"
            )
        # 同类同分仅约束客观题（choice/fill_blank/judgement/true_false）；
        # 解答类大题（calculation/proof/comprehensive）允许各题分值不同。
        if _normalize_type(qtype) in OBJECTIVE_TYPES:
            if qtype in scores_by_type and abs(scores_by_type[qtype] - max_score) > 1e-6:
                raise ValueError(f"All questions with question_type={qtype} must use the same max_score")
            scores_by_type[qtype] = max_score

        parts = item.get("parts")
        if not isinstance(parts, list) or not parts:
            raise ValueError(f"rubric.questions[{idx - 1}].parts must be a non-empty list")
        part_total = 0.0
        for pidx, part in enumerate(parts, start=1):
            if not isinstance(part, dict):
                raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}] must be an object")
            for key in ["part_id", "part_score", "steps"]:
                if key not in part:
                    raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}] is missing field: {key}")
            part_score = float(part["part_score"])
            if not part_score.is_integer():
                raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].part_score must be an integer")
            part_total += part_score
            steps = part.get("steps")
            if not isinstance(steps, list) or not steps:
                raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].steps must be a non-empty list")
            step_total = 0.0
            for sidx, step in enumerate(steps, start=1):
                if not isinstance(step, dict):
                    raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].steps[{sidx - 1}] must be an object")
                for key in ["step_id", "step_score", "core_goal", "required_elements", "allow_alternative_methods"]:
                    if key not in step:
                        raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].steps[{sidx - 1}] is missing field: {key}")
                step_score = float(step["step_score"])
                if not step_score.is_integer():
                    raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].steps[{sidx - 1}].step_score must be an integer")
                step_total += step_score
            if abs(step_total - part_score) > 1e-6:
                raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}] step total does not equal part_score")
            if not isinstance(part.get("presentation_rules", []), list):
                raise ValueError(f"rubric.questions[{idx - 1}].parts[{pidx - 1}].presentation_rules must be a list")
        if abs(part_total - max_score) > 1e-6:
            raise ValueError(f"rubric.questions[{idx - 1}] part total does not equal max_score")

    answer_q_map: dict[str, dict[str, Any]] = {}
    for idx, item in enumerate(answer_questions, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"answer_key.questions[{idx - 1}] must be an object")
        for key in ["question_id", "canonical_answer", "accepted_forms", "method_variants", "parts"]:
            if key not in item:
                raise ValueError(f"answer_key.questions[{idx - 1}] is missing field: {key}")
        qid = str(item["question_id"])
        answer_q_map[qid] = item
        if not isinstance(item["accepted_forms"], list):
            raise ValueError(f"answer_key.questions[{idx - 1}].accepted_forms must be a list")
        if not isinstance(item["method_variants"], list):
            raise ValueError(f"answer_key.questions[{idx - 1}].method_variants must be a list")
        parts = item.get("parts")
        if not isinstance(parts, list) or not parts:
            raise ValueError(f"answer_key.questions[{idx - 1}].parts must be a non-empty list")
        for pidx, part in enumerate(parts, start=1):
            if not isinstance(part, dict):
                raise ValueError(f"answer_key.questions[{idx - 1}].parts[{pidx - 1}] must be an object")
            for key in ["part_id", "answer", "analysis", "step_milestones"]:
                if key not in part:
                    raise ValueError(f"answer_key.questions[{idx - 1}].parts[{pidx - 1}] is missing field: {key}")
            if not isinstance(part["step_milestones"], list):
                raise ValueError(f"answer_key.questions[{idx - 1}].parts[{pidx - 1}].step_milestones must be a list")

    if set(rubric_q_map) != set(answer_q_map):
        raise ValueError("rubric.questions and answer_key.questions must have identical question_id sets")
    total_score = sum(float(item.get("max_score") or 0) for item in rubric_q_map.values())
    if abs(total_score - 100.0) > 0.02:
        raise ValueError(f"rubric question total must be 100, got {total_score:.2f}")
    rubric["total_score"] = 100.0


    if "warnings" not in meta or not isinstance(meta["warnings"], list):
        raise ValueError("meta.warnings must exist and be a list")


def save_generated_config(upload_dir: Path, payload: dict[str, Any], ts: str) -> tuple[Path, Path]:
    upload_dir.mkdir(parents=True, exist_ok=True)
    validate_generated_config(payload)
    rubric_path = upload_dir / f"rubric_{ts}.json"
    answer_key_path = upload_dir / f"answer_key_{ts}.json"

    _write_generated_json_atomic(rubric_path, payload["rubric"])
    _write_generated_json_atomic(answer_key_path, payload["answer_key"])

    return rubric_path, answer_key_path


def _write_generated_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.parent / f".{path.name}.{os.urandom(8).hex()}.tmp"
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _build_generation_prompt(doc_text: str, *, include_source_text: bool = True) -> str:
    return (
        "你是中学数学教研评分设计助手。\n"
        "请从给定Word文本中抽取‘题目、分值、答案与解析’，并生成可执行评分标准。\n"
        "核心目标：支持填空题等价答案判分、解答题/证明题按证明义务扣分、不同正确解法给分。\n"
        "硬性总分：本系统所有考试批改统一按 100 分制设计。rubric.total_score 必须等于 100，所有题目 max_score 之和必须等于 100；若原卷不是 100 分制，请按原始分值比例换算。\n"
        "单题上限：所有题型的 question.max_score 均不能超过 18 分。题目可以拆成多个小问 parts 分别赋分，但整道题 max_score 仍不得超过 18。\n"
        "硬性赋分：所有 max_score、part_score、step_score、proof_obligations.weight、deduction_policy.max_deduction 都必须是整数，不能出现 2.5、3.33 这类小数。\n"
        "同类同分（仅客观题）：相同 question_type 的客观题必须分值完全相同。例如所有 choice 题同分，所有 fill_blank 题同分，不能出现有的选择题3分、有的选择题4分。解答类大题（calculation/proof/comprehensive）不要求同类同分，可按题目难度与工作量赋予不同分值。\n"
        "题型细分：解答类题目不要全部写成一种类型，可按实际任务分为 calculation（计算/求解）、proof（证明）、comprehensive（一般综合解答）等多种 question_type。\n"
        "题型纠偏：只要题目要求证明、求证、说明理由、说明某结论成立、补全证明过程、判断并说明、添加条件使结论能够推出，就应标为 proof；如果题目主要要求求角度、求长度、求周长、求面积、求值、计算、化简或解方程，且不要求证明/说明理由，才标为 calculation。comprehensive 用于同时包含证明、计算、作图或开放论述的混合型题。\n"
        "若原卷分值与“100分制、整数、客观题同类同分、单题不超过18分”冲突，请优先按这些规则重新设计赋分。\n"
        "特别要求：证明题/解答题不要把参考答案路径当作唯一标准；应抽象成 proof_obligations（证明义务）和 deduction_policy（扣分规则）。\n"
        "必须覆盖 Word 中出现的全部题目，包括选择题、填空题、判断题、客观题和解答题；禁止只生成大题或只生成答案解析部分。\n"
        "如果 Word 中有选择题/填空题，即使评分逻辑简单，也必须在 rubric.questions 和 answer_key.questions 中逐题列出。\n"
        "选择题/填空题不能写成“仅抽象化处理”或只放在 meta.warnings；如果答案明确，应生成 direct_answer 评分标准。\n"
        "如果信息缺失，不要臆造，把不确定项写入 meta.warnings。\n\n"
        "仅输出严格 JSON，不要 markdown。输出必须尽量压缩为单行 JSON，不要漂亮打印，不要添加解释性文字，避免输出被截断。\n"
        "可以省略空数组字段中的多余说明，但 JSON 顶层和必需字段必须完整闭合。\n"
        "JSON 结构必须是:\n"
        "{\n"
        "  \"rubric\": {\n"
        "    \"exam_title\": \"...\",\n"
        "    \"total_score\": 100,\n"
        "    \"questions\": [\n"
        "      {\n"
        "        \"question_id\": \"Q1\",\n"
        "        \"question_type\": \"fill_blank|choice|calculation|proof|comprehensive\",\n"
        "        \"max_score\": 8,\n"
        "        \"knowledge_id\": \"C2_01\",\n"
        "        \"knowledge_ids\": [\"C2_01\", \"C2_03\"],\n"
        "        \"knowledge_points\": [\n"
        "          {\"knowledge_id\": \"C2_01\", \"knowledge_name\": \"主要知识点名称\"},\n"
        "          {\"knowledge_id\": \"C2_03\", \"knowledge_name\": \"相关知识点名称\"}\n"
        "        ],\n"
        "        \"knowledge_name\": \"简短知识点名称，例如三角形内角和/全等三角形判定/等价代数化简\",\n"
        "        \"stem_summary\": \"题干摘要，用一句话概括本题考查内容\",\n"
        "        \"grading_mode\": \"direct_answer|deductive_obligation\",\n"
        "        \"require_final_answer\": false,\n"
        "        \"answer_only_max_score\": 2,\n"
        "        \"proof_obligations\": [\n"
        "          {\"obligation_id\": \"O1\", \"description\": \"需要证明/建立的关键数学义务\", \"weight\": 3, \"acceptable_evidence\": [\"等价完成方式\"]}\n"
        "        ],\n"
        "        \"deduction_policy\": [\n"
        "          {\"issue\": \"关键逻辑断裂/条件缺失/循环论证\", \"max_deduction\": 2, \"severity\": \"minor|major|fatal\"}\n"
        "        ],\n"
        "        \"parts\": [\n"
        "          {\n"
        "            \"part_id\": \"Q1(1)\",\n"
        "            \"part_score\": 4,\n"
        "            \"steps\": [\n"
        "              {\n"
        "                \"step_id\": \"S1\",\n"
        "                \"step_score\": 2,\n"
        "                \"core_goal\": \"...\",\n"
        "                \"required_elements\": [\"...\"],\n"
        "                \"allow_alternative_methods\": true\n"
        "              }\n"
        "            ],\n"
        "            \"presentation_rules\": [\n"
        "              {\"rule\": \"过程规范/结论完整\", \"max_deduction\": 1}\n"
        "            ]\n"
        "          }\n"
        "        ]\n"
        "      }\n"
        "    ]\n"
        "  },\n"
        "  \"answer_key\": {\n"
        "    \"questions\": [\n"
        "      {\n"
        "        \"question_id\": \"Q1\",\n"
        "        \"canonical_answer\": \"...\",\n"
        "        \"accepted_forms\": [\"...\", \"...\"],\n"
        "        \"method_variants\": [\n"
        "          {\"name\": \"方法A\", \"outline\": \"...\"},\n"
        "          {\"name\": \"方法B\", \"outline\": \"...\"}\n"
        "        ],\n"
        "        \"parts\": [\n"
        "          {\n"
        "            \"part_id\": \"Q1(1)\",\n"
        "            \"answer\": \"...\",\n"
        "            \"analysis\": \"...\",\n"
        "            \"step_milestones\": [\"关键中间结论1\", \"关键中间结论2\"]\n"
        "          }\n"
        "        ]\n"
        "      }\n"
        "    ]\n"
        "  },\n"
        "  \"meta\": {\n"
        "    \"source\": \"word_docx\",\n"
        "    \"warnings\": [\"...\"]\n"
        "  }\n"
        "}\n\n"
        "硬性要求：\n"
        "1) question_id 在 rubric 与 answer_key 中一一对应。\n"
        "1.1) knowledge_id 仅是本试卷内的来源编号，不承担跨系统知识身份；必须同时给出具体 knowledge_name，供统一技能目录归一。\n"
        "1.2) 一道题可以涉及多个知识点。若题目同时考查多个数学概念，必须输出 knowledge_ids 数组和 knowledge_points 数组；knowledge_id 只作为主知识点，取 knowledge_ids 的第一项。\n"
        "2) max_score、part_score、step_score 必须为整数且层级总分一致；不得输出小数。\n"
        "2.1) 整张试卷总分必须严格为 100 分，不能返回 10 分、120 分或其他总分。\n"
        "2.2) 相同 question_type 的客观题（choice/fill_blank）max_score 必须完全一致；解答类大题（calculation/proof/comprehensive）允许不同分值。\n"
        "2.3) 不限制不同题型之间的分值高低关系；仅要求同类型客观题同分。\n"
        "2.4) 任意 question.max_score 必须小于或等于 18；若大题有多问，请在 parts 中拆分小问分值，不要让整题超过 18。\n"
        "3) 对选择题/填空题，question_type 必须分别是 choice 或 fill_blank，parts 可只有一个，steps 可只有一个 direct-answer 步骤。\n"
        "4) 对解答题，必须按题目特征细分为 calculation/proof/comprehensive，必须有 parts 与 steps，不能只给最终答案；同时必须给出 proof_obligations 与 deduction_policy。明确要求证明/求证/说明结论成立的题使用 proof；主要求数值、角度、长度、面积、化简或方程结果的题使用 calculation。\n"
        "4.1) 对 proof 与 calculation，默认 require_final_answer=false，不因未额外写“答”单独扣分；对 comprehensive，默认 require_final_answer=true，未写最终答/结论完整性可小扣分。若题目本身明确要求写结论，可按题意调整。\n"
        "4.2) 对 proof/calculation/comprehensive，必须给出 answer_only_max_score：学生只写最终答案但没有有效过程时最多得分。该分值只能是少量结论分，不得超过本题 max_score 的 30%。\n"
        "5) 对选择题/填空题/direct_answer 题，必须执行全对全错规则：学生答案与 canonical_answer 或 accepted_forms 等价才给该题/该空满分；不符合答案及等价答案时该评分单元必须给 0 分，不得因为接近、猜测方向正确或过程片段给一半分。\n"
        "5.1) 对填空题，必须尽可能提供 accepted_forms（等价表达），例如 1/4、0.25、25% 这类数值等价形式，或 OB=OC、OC=OB、BO=CO 这类几何关系等价写法。\n"
        "6) 对可接受的不同解法，在 method_variants 中列出。\n"
        "7) 若题号或分值不确定，使用连续编号并写入 warnings，但不能因此漏掉题目。\n"
        "8) meta.warnings 不能出现“选择题与填空题未包含在详细评分逻辑中”这类遗漏；遇到这种情况必须回到 questions 中补齐。\n\n"
        "证明题/解答题生成准则：\n"
        "- 答案结论只作为小部分得分依据，过程分主要来自 proof_obligations 是否完成。\n"
        "- 证明题和纯计算题通常不要求额外写“答”；一般综合解答题可要求结论完整，但这应作为 presentation_rules 的小扣分项，不应替代过程分判断。\n"
        "- 只写最终答案但没有有效过程时，只能获得 answer_only_max_score 以内的结论分；不能获得主要过程分。\n"
        "- proof_obligations 描述学生必须完成的数学责任，例如建立辅助条件、证明全等/相似、推出角度/线段关系、说明定理前提等。\n"
        "- deduction_policy 描述扣分项，例如缺少前提、跳步严重、逻辑循环、结论与过程断裂、符号/对象指代不清。\n"
        "- steps 可以对应证明义务的评分权重，但不能要求学生过程与参考答案逐句一致。\n\n"
        + (f"Word原文:\n{doc_text}" if include_source_text else "")
    )



def generate_grading_config_from_text(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
) -> dict[str, Any]:
    return generate_grading_config_from_docx_text(
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
    )


def generate_grading_config_from_images(
    image_blobs: list[bytes],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
) -> dict[str, Any]:
    if not image_blobs:
        raise ValueError("PDF 整卷视觉模式未获得任何页面图片。")
    if report:
        report(0.20, "PDF 整卷视觉单次请求", f"发送 {len(image_blobs)} 张整页图片，一次完成解析与赋分。")
    payload = llm_client.json_from_images_once(
        _build_whole_image_generation_prompt(),
        image_blobs,
        model=model_name,
        extra_kwargs=_config_generation_extra_kwargs(),
        use_config_client=True,
    )
    return _finalize_whole_generation_payload(payload, "whole_pdf_visual_single_request")


# P3-09 compatibility facade. Keep the historical names for callers while the
# active batch workflow and score prompt live under backend.config_generation.
from backend.config_generation.compat import (  # noqa: E402
    failed_grading_config_batches as failed_grading_config_batches,
    failed_grading_config_question_ids as failed_grading_config_question_ids,
    generate_grading_config_in_batches as generate_grading_config_in_batches,
    refine_grading_config_from_manual_structure as refine_grading_config_from_manual_structure,
    retry_failed_grading_config_batches as retry_failed_grading_config_batches,
)
from backend.config_generation.gateway import (  # noqa: E402
    config_generation_extra_kwargs as _config_generation_extra_kwargs,
)
from backend.config_generation.prompts import (  # noqa: E402
    build_manual_structure_refinement_prompt as _build_manual_structure_refinement_prompt,
    build_score_allocation_prompt as _build_score_allocation_prompt,
)


# P3-08 compatibility exports. Parsing now lives in backend.document_parsing;
# callers that historically imported these names keep the same public surface.
from backend.document_parsing import (  # noqa: E402
    extract_docx_text as extract_docx_text,
    image_paths_from_rich_text as _image_paths_from_rich_text,
    infer_question_type_from_text as _infer_question_type_from_block_text,
    parse_plain_question_blocks as preview_question_blocks_from_docx_text,
)
from backend.document_parsing.question_blocks import (  # noqa: E402
    _parse_inline_answer_blocks as _parse_inline_answer_blocks,
)

