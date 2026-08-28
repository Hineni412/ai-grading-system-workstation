from __future__ import annotations

import hashlib
import hmac
import json
import re
import csv
from io import BytesIO, StringIO
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Callable, Protocol
from uuid import uuid4

from .encrypted_database import EncryptedDatabase
from .errors import VaultError
from .existing_student_roster import ExistingStudent, ExistingStudentRosterSource
from .secure_repository import EncryptedObjectRepository
from .subject_canonical import canonical_subjects


_RESULT_STATES = {
    "normal",
    "absent",
    "exempt",
    "missing",
    "incomplete",
    "makeup",
}

# 场次指纹与 metadata_complete 判定共用的 8 个必填元数据字段。
_SESSION_REQUIRED_FIELDS = (
    "title",
    "academic_year",
    "term",
    "grade",
    "exam_type",
    "comparison_series",
    "occurred_on",
    "source_reference",
)

# 允许教师事后更正的场次元数据字段；comparison_series 与 source_reference
# 参与归组与来源追溯，不在更正范围内。
_SESSION_EDITABLE_FIELDS = (
    "title",
    "grade",
    "term",
    "exam_type",
    "occurred_on",
    "academic_year",
)

_SESSION_EDITABLE_LABELS = {
    "title": ("考试名称", 500),
    "grade": ("年级", 40),
    "term": ("学期", 40),
    "exam_type": ("考试类型", 120),
    "occurred_on": ("考试日期", 40),
    "academic_year": ("学年", 40),
}

_SESSION_DELETE_CONFIRMATION_PHRASE = "确认删除本场考试"


def _iso() -> str:
    return datetime.now(UTC).isoformat()


_NAME_HEADER = re.compile(r"姓名|名字")
_SUB_HEADER_LABELS = {
    "得分",
    "成绩",
    "分数",
    "等级",
    "校次",
    "班次",
    "排名",
    "校名",
    "班名",
    "得分率",
    "满分",
}


def _split_header(
    rows: list[list[str]],
) -> tuple[list[str], list[list[str]]]:
    """Locate the real header row and merge a two-level header if present.

    School grade exports often start with a merged title row, followed by a
    group header row (序号/姓名/数学/...) and a sub header row (得分/等级/...).
    The first row containing 姓名/名字 is treated as the header; when the row
    right below it is a sub header row, the two are joined into single labels
    such as "数学-得分". Plain single-header sheets keep the old behavior.
    """
    header_index = 0
    for index, row in enumerate(rows[:10]):
        if any(_NAME_HEADER.search(cell) for cell in row):
            header_index = index
            break
    header_row = rows[header_index]
    data_start = header_index + 1
    sub_row: list[str] | None = None
    if data_start < len(rows):
        candidate = rows[data_start]
        if sum(1 for cell in candidate if cell in _SUB_HEADER_LABELS) >= 2:
            sub_row = candidate
            data_start += 1
    width = max(len(row) for row in rows[header_index:])
    sub_values = set(sub_row) if sub_row is not None else set()

    def _cell(row: list[str], index: int) -> str:
        return row[index] if index < len(row) else ""

    headers: list[str] = []
    group = ""
    for index in range(width):
        top = _cell(header_row, index)
        # Some writers repeat the sub labels (得分/等级/...) inside the merged
        # group row; treat those as blank anchors so the group name carries.
        if top and top not in sub_values:
            group = top
        sub = _cell(sub_row, index) if sub_row is not None else ""
        parts = [part for part in (group, sub) if part]
        # Avoid "姓名-姓名" when a merged label repeats on both levels.
        if len(parts) == 2 and parts[0] == parts[1]:
            parts.pop()
        headers.append("-".join(parts) or f"未命名列 {index + 1}")
    return headers, rows[data_start:]


class AssessmentEvidenceSource(Protocol):
    def read(self, query: dict[str, object]) -> dict[str, object]: ...


class ExistingMathAssessmentAdapter:
    """Read-only boundary around an injected math-summary reader."""

    def __init__(
        self,
        reader: Callable[[dict[str, object]], dict[str, object]],
    ) -> None:
        self._reader = reader

    def read(self, query: dict[str, object]) -> dict[str, object]:
        batch = dict(self._reader(dict(query)))
        batch["source_kind"] = "existing_math"
        batch["read_only"] = True
        return batch


class ConfirmedSpreadsheetAdapter:
    """Turns an already teacher-confirmed preview into a neutral batch."""

    @staticmethod
    def preview(
        *,
        file_name: str,
        content: bytes,
        sheet_name: str | None = None,
    ) -> dict[str, object]:
        if not content or len(content) > 10 * 1024 * 1024:
            raise VaultError(
                "assessment_file_size_invalid",
                "成绩文件不能为空且不能超过 10 MB",
                status_code=422,
            )
        lower_name = file_name.casefold()
        if lower_name.endswith(".csv"):
            decoded = None
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None:
                raise VaultError(
                    "assessment_csv_encoding_invalid",
                    "CSV 文件编码无法识别",
                    status_code=422,
                )
            rows = [
                [str(cell).strip() for cell in row]
                for row in csv.reader(StringIO(decoded))
                if any(str(cell).strip() for cell in row)
            ]
            sheets = ["CSV"]
            selected_sheet = "CSV"
        elif lower_name.endswith(".xlsx"):
            try:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    BytesIO(content),
                    read_only=True,
                    data_only=True,
                )
                sheets = list(workbook.sheetnames)
                selected_sheet = sheet_name or sheets[0]
                if selected_sheet not in sheets:
                    raise VaultError(
                        "assessment_sheet_not_found",
                        "选择的工作表不存在",
                        status_code=422,
                    )
                worksheet = workbook[selected_sheet]
                # Some tools write a wrong declared dimension (e.g. A1:A1);
                # rescan so trailing columns are not silently dropped.
                worksheet.reset_dimensions()
                rows = [
                    [
                        "" if cell is None else str(cell).strip()
                        for cell in row
                    ]
                    for row in worksheet.iter_rows(
                        min_row=1,
                        max_row=5002,
                        values_only=True,
                    )
                    if any(cell is not None and str(cell).strip() for cell in row)
                ]
                workbook.close()
            except VaultError:
                raise
            except Exception as exc:
                raise VaultError(
                    "assessment_xlsx_invalid",
                    "XLSX 文件无法读取或已经损坏",
                    status_code=422,
                ) from exc
        else:
            raise VaultError(
                "assessment_file_type_invalid",
                "只支持 CSV 或 XLSX 成绩文件",
                status_code=422,
            )
        if not rows:
            raise VaultError(
                "assessment_sheet_empty",
                "成绩工作表为空",
                status_code=422,
            )
        headers, data_rows = _split_header(rows)
        width = len(headers)
        preview_rows = [
            {
                headers[index]: (
                    row[index] if index < len(row) else ""
                )
                for index in range(width)
            }
            for row in data_rows[:5000]
        ]
        return {
            "file_name": file_name,
            "sheet_names": sheets,
            "selected_sheet": selected_sheet,
            "headers": headers,
            "rows": preview_rows,
            "preview_row_count": len(preview_rows),
            "truncated": len(data_rows) > 5000,
            "raw_file_retained": False,
            "temporary_file_created": False,
        }

    def read(self, query: dict[str, object]) -> dict[str, object]:
        if not bool(query.get("teacher_confirmed")):
            raise VaultError(
                "assessment_preview_confirmation_required",
                "成绩预览必须由教师确认后才能进入证据库",
                status_code=422,
            )
        return {
            "source_kind": "confirmed_spreadsheet",
            "read_only": True,
            "teacher_confirmed": True,
            "source_label": str(
                query.get("source_label") or "教师确认的表格预览"
            ),
            "assessments": list(query.get("assessments") or []),
        }


class AssessmentEvidenceService:
    def __init__(
        self,
        database: EncryptedDatabase,
        repository: EncryptedObjectRepository,
        key_provider: Callable[[str], bytes],
        subject_ensurer: Callable[..., str] | None = None,
        roster_source: ExistingStudentRosterSource | None = None,
        projections: Any | None = None,
    ) -> None:
        self.database = database
        self.repository = repository
        self._key_provider = key_provider
        self._subject_ensurer = subject_ensurer
        self._roster_source = roster_source
        self.projections = projections

    def confirm_batch(
        self,
        *,
        token: str,
        operation_id: str,
        batch: dict[str, object],
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.batch.confirm")
        if replay is not None:
            return replay
        source_kind = str(batch.get("source_kind") or "")
        if source_kind not in {"existing_math", "confirmed_spreadsheet"}:
            raise VaultError(
                "assessment_source_invalid",
                "学业证据来源无效",
                status_code=422,
            )
        if not bool(batch.get("teacher_confirmed")):
            raise VaultError(
                "assessment_confirmation_required",
                "学业证据必须经过教师确认",
                status_code=422,
            )
        assessments = list(batch.get("assessments") or [])
        if not assessments:
            raise VaultError(
                "assessment_batch_empty",
                "没有可确认的考试证据",
                status_code=422,
            )
        source_signature = {
            "source_kind": source_kind,
            "source_label": str(batch.get("source_label") or ""),
            "assessments": assessments,
        }
        fingerprint = hmac.new(
            vmk,
            json.dumps(
                source_signature,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8"),
            hashlib.sha256,
        ).digest()
        with closing(self.database.connect()) as connection:
            duplicate = connection.execute(
                """
                SELECT import_id FROM source_fingerprints
                WHERE source_fingerprint = ?
                """,
                (fingerprint,),
            ).fetchone()
        if duplicate is not None:
            return {
                "import_id": str(duplicate["import_id"]),
                "duplicate": True,
                "created_assessments": 0,
                "created_results": 0,
                "model_enabled": False,
                "physical_request_count": 0,
            }
        import_id = uuid4().hex
        import_object_id = f"assessment-import-{import_id}"
        timestamp = _iso()
        created_results = 0
        with closing(self.database.connect()) as connection:
            with connection:
                self.repository.put(
                    connection,
                    vmk=vmk,
                    object_id=import_object_id,
                    object_type="assessment_import",
                    payload={
                        "source_kind": source_kind,
                        "source_label": str(
                            batch.get("source_label") or ""
                        ),
                        "teacher_confirmed": True,
                        "read_only_source": True,
                        "raw_file_retained": False,
                        "model_enabled": False,
                        "physical_request_count": 0,
                    },
                )
                connection.execute(
                    """
                    INSERT INTO assessment_imports (
                        import_id, payload_object_id, source_kind,
                        status, created_at
                    ) VALUES (?, ?, ?, 'confirmed', ?)
                    """,
                    (
                        import_id,
                        import_object_id,
                        source_kind,
                        timestamp,
                    ),
                )
                connection.execute(
                    """
                    INSERT INTO source_fingerprints (
                        source_fingerprint, import_id, created_at
                    ) VALUES (?, ?, ?)
                    """,
                    (fingerprint, import_id, timestamp),
                )
                for assessment in assessments:
                    created_results += self._insert_assessment(
                        connection,
                        vmk=vmk,
                        import_id=import_id,
                        assessment=dict(assessment),
                    )
                result = {
                    "import_id": import_id,
                    "duplicate": False,
                    "created_assessments": len(assessments),
                    "created_results": created_results,
                    "model_enabled": False,
                    "physical_request_count": 0,
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.batch.confirm",
                    result,
                )
        return result

    def list_subject_evidence(
        self,
        *,
        token: str,
        subject_id: str,
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        with closing(self.database.connect()) as connection:
            self._subject_exists(connection, subject_id)
            rows = connection.execute(
                """
                SELECT sr.*, a.occurred_on, a.payload_object_id AS assessment_object,
                       ai.source_kind, ai.payload_object_id AS import_object,
                       ev.evidence_version_id, ev.version,
                       ev.payload_object_id AS evidence_object, ev.state
                FROM subject_results sr
                JOIN assessments a ON a.assessment_id = sr.assessment_id
                JOIN assessment_imports ai ON ai.import_id = a.import_id
                JOIN evidence_versions ev ON ev.result_id = sr.result_id
                WHERE sr.subject_id = ? AND ev.state = 'active'
                ORDER BY a.occurred_on, sr.created_at
                """,
                (subject_id,),
            ).fetchall()
            items = [
                self._evidence_from_row(connection, vmk, row)
                for row in rows
            ]
        return {
            "subject_id": subject_id,
            "items": items,
            "model_enabled": False,
            "physical_request_count": 0,
        }

    def compare(
        self,
        *,
        token: str,
        older_evidence_version_id: str,
        newer_evidence_version_id: str,
    ) -> dict[str, object]:
        older = self.get_evidence(
            token=token,
            evidence_version_id=older_evidence_version_id,
        )
        newer = self.get_evidence(
            token=token,
            evidence_version_id=newer_evidence_version_id,
        )
        if older["subject_id"] != newer["subject_id"]:
            raise VaultError(
                "assessment_subject_mismatch",
                "只能比较同一学生的证据",
                status_code=422,
            )
        status, limitations = self._comparability(older, newer)
        delta = None
        if (
            status == "directly_comparable"
            and
            older["result_state"] in {"normal", "makeup"}
            and newer["result_state"] in {"normal", "makeup"}
            and older.get("score") is not None
            and newer.get("score") is not None
        ):
            delta = float(newer["score"]) - float(older["score"])
        return {
            "older_evidence_version_id": older_evidence_version_id,
            "newer_evidence_version_id": newer_evidence_version_id,
            "comparability": status,
            "limitations": limitations,
            "score_delta": delta,
        }

    def trend(
        self,
        *,
        token: str,
        subject_id: str,
        subject_name: str,
    ) -> dict[str, object]:
        items = [
            item
            for item in self.list_subject_evidence(
                token=token,
                subject_id=subject_id,
            )["items"]
            if item["subject_name"] == subject_name
        ]
        comparisons = [
            self._comparability(items[index - 1], items[index])
            for index in range(1, len(items))
        ]
        trend_allowed = (
            len(items) >= 3
            and all(status == "directly_comparable" for status, _ in comparisons)
        )
        return {
            "subject_id": subject_id,
            "subject_name": subject_name,
            "evidence_count": len(items),
            "trend_allowed": trend_allowed,
            "conclusion": (
                "口径相近的证据达到三次，可描述变化趋势"
                if trend_allowed
                else "证据不足或口径不同，不能称为趋势"
            ),
            "comparison_basis": [
                {
                    "from": items[index - 1]["evidence_version_id"],
                    "to": items[index]["evidence_version_id"],
                    "comparability": comparisons[index - 1][0],
                    "limitations": comparisons[index - 1][1],
                }
                for index in range(1, len(items))
            ],
        }

    def get_evidence(
        self,
        *,
        token: str,
        evidence_version_id: str,
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        with closing(self.database.connect()) as connection:
            row = connection.execute(
                """
                SELECT sr.*, a.occurred_on,
                       a.payload_object_id AS assessment_object,
                       ai.source_kind, ai.payload_object_id AS import_object,
                       ev.evidence_version_id, ev.version,
                       ev.payload_object_id AS evidence_object, ev.state
                FROM evidence_versions ev
                JOIN subject_results sr ON sr.result_id = ev.result_id
                JOIN assessments a ON a.assessment_id = sr.assessment_id
                JOIN assessment_imports ai ON ai.import_id = a.import_id
                WHERE ev.evidence_version_id = ?
                """,
                (evidence_version_id,),
            ).fetchone()
            if row is None:
                raise VaultError(
                    "assessment_evidence_not_found",
                    "学业证据不存在",
                    status_code=404,
                )
            return self._evidence_from_row(connection, vmk, row)

    def supersede_evidence(
        self,
        *,
        token: str,
        operation_id: str,
        evidence_version_id: str,
        reason: str,
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.supersede")
        if replay is not None:
            return replay
        clean_reason = self._text(reason, "证据修订原因", 1000)
        with closing(self.database.connect()) as connection:
            with connection:
                row = connection.execute(
                    """
                    SELECT * FROM evidence_versions
                    WHERE evidence_version_id = ?
                    """,
                    (evidence_version_id,),
                ).fetchone()
                if row is None:
                    raise VaultError(
                        "assessment_evidence_not_found",
                        "学业证据不存在",
                        status_code=404,
                    )
                payload, revision = self.repository.get(
                    connection,
                    vmk=vmk,
                    object_id=str(row["payload_object_id"]),
                )
                payload["superseded_reason"] = clean_reason
                payload["superseded_at"] = _iso()
                self.repository.put(
                    connection,
                    vmk=vmk,
                    object_id=str(row["payload_object_id"]),
                    object_type="assessment_evidence_version",
                    payload=payload,
                    expected_revision=revision,
                )
                connection.execute(
                    """
                    UPDATE evidence_versions SET state = 'superseded'
                    WHERE evidence_version_id = ?
                    """,
                    (evidence_version_id,),
                )
                connection.execute(
                    """
                    UPDATE attention_cards SET state = 'invalidated',
                        updated_at = ?
                    WHERE evidence_version_id = ? AND state = 'draft'
                    """,
                    (_iso(), evidence_version_id),
                )
                result = {
                    "evidence_version_id": evidence_version_id,
                    "superseded": True,
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.supersede",
                    result,
                )
        return result

    def update_session_metadata(
        self,
        *,
        token: str,
        operation_id: str,
        session_id: str,
        fields: dict[str, object],
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.session_update")
        if replay is not None:
            return replay
        updates = self._validated_session_updates(fields)
        with closing(self.database.connect()) as connection:
            with connection:
                row = connection.execute(
                    """
                    SELECT session_id, payload_object_id
                    FROM assessment_sessions
                    WHERE session_id = ?
                    """,
                    (session_id,),
                ).fetchone()
                if row is None:
                    raise VaultError(
                        "assessment_session_not_found",
                        "成绩场次不存在",
                        status_code=404,
                    )
                payload, revision = self.repository.get(
                    connection,
                    vmk=vmk,
                    object_id=str(row["payload_object_id"]),
                )
                merged = {**payload, **updates}
                merged["metadata_complete"] = all(
                    str(merged.get(key) or "").strip()
                    for key in _SESSION_REQUIRED_FIELDS
                )
                fingerprint = self._session_fingerprint(vmk, merged)
                conflict = connection.execute(
                    """
                    SELECT session_id FROM assessment_sessions
                    WHERE source_fingerprint = ? AND session_id != ?
                    """,
                    (fingerprint, session_id),
                ).fetchone()
                if conflict is not None:
                    raise VaultError(
                        "assessment_session_conflict",
                        "已存在相同场次，可删除本场次后重新上传",
                        status_code=409,
                    )
                self.repository.put(
                    connection,
                    vmk=vmk,
                    object_id=str(row["payload_object_id"]),
                    object_type="assessment_session",
                    payload=merged,
                    expected_revision=revision,
                )
                connection.execute(
                    """
                    UPDATE assessment_sessions
                    SET source_fingerprint = ?, metadata_complete = ?,
                        updated_at = ?
                    WHERE session_id = ?
                    """,
                    (
                        fingerprint,
                        int(bool(merged["metadata_complete"])),
                        _iso(),
                        session_id,
                    ),
                )
                result = {
                    "session_id": session_id,
                    "updated": True,
                    "session": {
                        key: merged.get(key)
                        for key in _SESSION_REQUIRED_FIELDS
                    }
                    | {"metadata_complete": bool(merged["metadata_complete"])},
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.session_update",
                    result,
                )
        return result

    def update_session_max_scores(
        self,
        *,
        token: str,
        operation_id: str,
        session_id: str,
        max_scores: dict[str, float],
        participant_count: int | None = None,
    ) -> dict[str, object]:
        """更正场次满分与年级人数：满分改考试/证据两级，人数改全部排名上下文。"""
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.session_max_scores")
        if replay is not None:
            return replay
        cleaned = self._validated_max_score_updates(max_scores, participant_count)
        with closing(self.database.connect()) as connection:
            with connection:
                self._session_row(connection, session_id)
                updated = self._apply_session_max_scores(
                    connection,
                    vmk,
                    session_id=session_id,
                    max_scores=cleaned,
                    participant_count=participant_count,
                )
                result = {
                    "session_id": session_id,
                    "updated": True,
                    "subjects": sorted(updated),
                    "participant_count": participant_count,
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.session_max_scores",
                    result,
                )
        return result

    def update_global_max_scores(
        self,
        *,
        token: str,
        operation_id: str,
        max_scores: dict[str, float],
        participant_count: int | None = None,
    ) -> dict[str, object]:
        """全年级统一更正：对所有 active 场次按归一科目名覆盖满分与年级人数。

        覆盖语义：有值即覆盖（包括已有值）。max_scores 的键是归一展示名，
        每场先用归一映射反查该场次的原始列名再更新；该场无此科目则跳过，
        实际覆盖的场次数计入返回。
        """
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.global_max_scores")
        if replay is not None:
            return replay
        cleaned = self._validated_max_score_updates(max_scores, participant_count)
        with closing(self.database.connect()) as connection:
            with connection:
                session_ids = [
                    str(row["session_id"])
                    for row in connection.execute(
                        """
                        SELECT session_id FROM assessment_sessions
                        WHERE state = 'active'
                        ORDER BY session_id
                        """
                    ).fetchall()
                ]
                sessions_updated = 0
                subject_hits = {name: 0 for name in cleaned}
                for session_id in session_ids:
                    name_map = canonical_subjects(
                        self._session_subject_names(connection, vmk, session_id)
                    )
                    translated = {
                        raw: cleaned[canonical]
                        for raw, canonical in name_map.items()
                        if canonical in cleaned
                    }
                    if not translated and participant_count is None:
                        continue
                    updated = self._apply_session_max_scores(
                        connection,
                        vmk,
                        session_id=session_id,
                        max_scores=translated,
                        participant_count=participant_count,
                    )
                    sessions_updated += 1
                    for canonical in {name_map[raw] for raw in updated}:
                        subject_hits[canonical] += 1
                result = {
                    "sessions_updated": sessions_updated,
                    "subjects": subject_hits,
                    "participant_count": participant_count,
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.global_max_scores",
                    result,
                )
        return result

    @staticmethod
    def _validated_max_score_updates(
        max_scores: dict[str, float],
        participant_count: int | None,
    ) -> dict[str, float]:
        """满分/年级人数校验：空更新、非正人数、非正满分都在这里拒绝。

        「总分」键被忽略：总分满分由读取层按其余展示科目满分之和派生，
        不再单独维护；旧客户端仍带该键时按未填写处理，不报错。
        """
        max_scores = {
            name: value
            for name, value in max_scores.items()
            if str(name).strip() != "总分"
        }
        if not max_scores and participant_count is None:
            raise VaultError(
                "assessment_session_update_empty",
                "没有需要保存的更正内容",
                status_code=422,
            )
        if participant_count is not None and (
            isinstance(participant_count, bool)
            or not isinstance(participant_count, int)
            or participant_count <= 0
        ):
            raise VaultError(
                "assessment_participant_count_invalid",
                "年级人数必须是正整数",
                status_code=422,
            )
        cleaned: dict[str, float] = {}
        for name, value in max_scores.items():
            subject_name = str(name).strip()
            try:
                number = float(value)
            except (TypeError, ValueError) as exc:
                raise VaultError(
                    "assessment_max_score_invalid",
                    "满分必须是正数",
                    status_code=422,
                ) from exc
            if not number > 0:
                raise VaultError(
                    "assessment_max_score_invalid",
                    "满分必须是正数",
                    status_code=422,
                )
            cleaned[subject_name] = number
        return cleaned

    def _session_subject_names(
        self,
        connection: Any,
        vmk: bytes,
        session_id: str,
    ) -> set[str]:
        """场次内全部 active 考试的原始科目名（未归一）。"""
        rows = connection.execute(
            """
            SELECT DISTINCT a.payload_object_id AS assessment_object
            FROM assessment_session_members m
            JOIN evidence_versions ev
              ON ev.evidence_version_id = m.evidence_version_id
            JOIN subject_results sr ON sr.result_id = ev.result_id
            JOIN assessments a ON a.assessment_id = sr.assessment_id
            WHERE m.session_id = ? AND ev.state = 'active'
            """,
            (session_id,),
        ).fetchall()
        names: set[str] = set()
        for row in rows:
            payload, _revision = self.repository.get(
                connection,
                vmk=vmk,
                object_id=str(row["assessment_object"]),
            )
            subject_name = str(payload.get("subject_name") or "")
            if subject_name:
                names.add(subject_name)
        return names

    def _apply_session_max_scores(
        self,
        connection: Any,
        vmk: bytes,
        *,
        session_id: str,
        max_scores: dict[str, float],
        participant_count: int | None,
    ) -> list[str]:
        """单场应用满分/人数更正，返回实际更新的原始科目名。"""
        member_rows = connection.execute(
            """
            SELECT ev.payload_object_id AS evidence_object,
                   a.assessment_id,
                   a.payload_object_id AS assessment_object
            FROM assessment_session_members m
            JOIN evidence_versions ev
              ON ev.evidence_version_id = m.evidence_version_id
            JOIN subject_results sr ON sr.result_id = ev.result_id
            JOIN assessments a ON a.assessment_id = sr.assessment_id
            WHERE m.session_id = ? AND ev.state = 'active'
            """,
            (session_id,),
        ).fetchall()
        assessments: dict[str, dict[str, Any]] = {}
        for row in member_rows:
            entry = assessments.setdefault(
                str(row["assessment_id"]),
                {
                    "assessment_object": str(row["assessment_object"]),
                    "evidence_objects": [],
                },
            )
            entry["evidence_objects"].append(
                str(row["evidence_object"])
            )
        payloads: dict[str, tuple[dict[str, Any], int, str]] = {}
        session_subjects: set[str] = set()
        for assessment_id, entry in assessments.items():
            payload, revision = self.repository.get(
                connection,
                vmk=vmk,
                object_id=str(entry["assessment_object"]),
            )
            subject_name = str(payload.get("subject_name") or "")
            payloads[assessment_id] = (payload, revision, subject_name)
            if subject_name:
                session_subjects.add(subject_name)
        unknown = sorted(set(max_scores) - session_subjects)
        if unknown:
            raise VaultError(
                "assessment_subject_not_in_session",
                "场次不包含学科："
                + "、".join(unknown)
                + "；本场次学科："
                + "、".join(sorted(session_subjects)),
                status_code=422,
            )
        updated: list[str] = []
        for assessment_id, (
            payload,
            revision,
            subject_name,
        ) in payloads.items():
            if subject_name not in max_scores:
                continue
            new_max = max_scores[subject_name]
            payload["max_score"] = new_max
            self.repository.put(
                connection,
                vmk=vmk,
                object_id=str(
                    assessments[assessment_id]["assessment_object"]
                ),
                object_type="assessment",
                payload=payload,
                expected_revision=revision,
            )
            for evidence_object in assessments[assessment_id][
                "evidence_objects"
            ]:
                evidence, evidence_revision = self.repository.get(
                    connection,
                    vmk=vmk,
                    object_id=str(evidence_object),
                )
                evidence["max_score"] = new_max
                self.repository.put(
                    connection,
                    vmk=vmk,
                    object_id=str(evidence_object),
                    object_type="assessment_evidence_version",
                    payload=evidence,
                    expected_revision=evidence_revision,
                )
            updated.append(subject_name)
        if participant_count is not None:
            # 年级人数写进该场次全部 active 证据对应的排名上下文。
            rank_rows = connection.execute(
                """
                SELECT rc.payload_object_id AS rank_object
                FROM assessment_session_members m
                JOIN evidence_versions ev
                  ON ev.evidence_version_id = m.evidence_version_id
                JOIN rank_contexts rc ON rc.result_id = ev.result_id
                WHERE m.session_id = ? AND ev.state = 'active'
                """,
                (session_id,),
            ).fetchall()
            for rank_row in rank_rows:
                rank_object = str(rank_row["rank_object"])
                rank_context, rank_revision = self.repository.get(
                    connection,
                    vmk=vmk,
                    object_id=rank_object,
                )
                rank_context["participant_count"] = participant_count
                self.repository.put(
                    connection,
                    vmk=vmk,
                    object_id=rank_object,
                    object_type="rank_context",
                    payload=rank_context,
                    expected_revision=rank_revision,
                )
        return updated

    def preview_delete_session(
        self,
        *,
        token: str,
        session_id: str,
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        with closing(self.database.connect()) as connection:
            row = self._session_row(connection, session_id)
            payload, _revision = self.repository.get(
                connection,
                vmk=vmk,
                object_id=str(row["payload_object_id"]),
            )
            impact = self._session_impact(connection, session_id)
        counts = {
            "results": len(impact["result_ids"]),
            "assessments": len(impact["assessment_ids"]),
            "imports": len(impact["import_ids"]),
            "attention_cards": len(impact["attention_card_ids"]),
        }
        version_payload = {
            "session_id": session_id,
            "evidence_version_ids": sorted(impact["evidence_version_ids"]),
            "result_ids": sorted(impact["result_ids"]),
            "assessment_ids": sorted(impact["assessment_ids"]),
            "import_ids": sorted(impact["import_ids"]),
            "attention_card_ids": sorted(impact["attention_card_ids"]),
            "counts": counts,
        }
        preview_version = hashlib.sha256(
            json.dumps(
                version_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        return {
            "session_id": session_id,
            "title": str(payload.get("title") or ""),
            "counts": counts,
            "preview_version": preview_version,
            "confirmation_phrase": _SESSION_DELETE_CONFIRMATION_PHRASE,
        }

    def delete_session(
        self,
        *,
        token: str,
        operation_id: str,
        session_id: str,
        preview_version: str | None = None,
        confirmation_phrase: str,
    ) -> dict[str, object]:
        self._key_provider(token)
        if re.fullmatch(r"[A-Za-z0-9_-]{8,128}", operation_id) is None:
            raise VaultError(
                "vault_operation_id_invalid",
                "操作编号无效，请刷新页面后重试",
                status_code=422,
            )
        replay = self._idempotent(operation_id, "evidence.session_delete")
        if replay is not None:
            return replay
        preview = self.preview_delete_session(
            token=token,
            session_id=session_id,
        )
        if (
            preview_version is not None
            and preview_version != preview["preview_version"]
        ):
            raise VaultError(
                "assessment_session_delete_preview_changed",
                "删除影响已经变化，请重新查看并确认",
                status_code=409,
                details=preview,
            )
        live_database = self.database
        snapshot = live_database.snapshot_bytes()
        with TemporaryDirectory(
            prefix=".session-delete-candidate-",
            dir=live_database.root,
        ) as candidate_root:
            candidate_database = live_database.isolated_copy(
                snapshot,
                root=Path(candidate_root),
            )
            self.database = candidate_database
            try:
                result = self._delete_session_once(
                    token=token,
                    operation_id=operation_id,
                    session_id=session_id,
                    confirmation_phrase=confirmation_phrase,
                )
                candidate_snapshot = candidate_database.snapshot_bytes()
            finally:
                self.database = live_database
        transaction: Path | None = None
        tombstoned: list[dict[str, object]] = []
        try:
            if self.projections is not None:
                with closing(live_database.connect()) as connection:
                    rows = connection.execute(
                        """
                        SELECT g.* FROM sensitive_work_groups g
                        WHERE g.source_kind = 'attention_followup'
                          AND g.source_id IN (
                            SELECT attention_card_id FROM attention_cards
                            WHERE evidence_version_id IN (
                                SELECT evidence_version_id
                                FROM assessment_session_members
                                WHERE session_id = ?
                            )
                        )
                        """,
                        (session_id,),
                    ).fetchall()
                    tombstoned = [dict(row) for row in rows]
                for group in tombstoned:
                    self.projections.tombstone(
                        token=token,
                        group_id=str(group["group_id"]),
                    )
            transaction = live_database.create_subject_delete_transaction(
                operation_id=operation_id,
                database_snapshot=snapshot,
            )
            live_database.replace_from_snapshot_atomically(
                candidate_snapshot
            )
            live_database.commit_subject_delete_transaction(transaction)
            return result
        except Exception:
            if transaction is not None and transaction.exists():
                live_database.recover_interrupted_operations()
            if self.projections is not None:
                for group in tombstoned:
                    try:
                        self.projections.upsert(
                            token=token,
                            source_kind=str(group["source_kind"]),
                            source_id=str(group["source_id"]),
                            occurrence_id=str(group["occurrence_id"] or "") or None,
                            state=str(group["state"]),
                            due_date=None
                            if group["due_date"] is None
                            else str(group["due_date"]),
                        )
                    except Exception:
                        pass
            raise

    def _delete_session_once(
        self,
        *,
        token: str,
        operation_id: str,
        session_id: str,
        confirmation_phrase: str,
    ) -> dict[str, object]:
        vmk = self._key_provider(token)
        replay = self._idempotent(operation_id, "evidence.session_delete")
        if replay is not None:
            return replay
        if confirmation_phrase != _SESSION_DELETE_CONFIRMATION_PHRASE:
            raise VaultError(
                "assessment_session_delete_confirmation_required",
                f"请输入“{_SESSION_DELETE_CONFIRMATION_PHRASE}”后再删除",
                status_code=422,
            )
        with closing(self.database.connect()) as connection:
            with connection:
                row = self._session_row(connection, session_id)
                impact = self._session_impact(connection, session_id)
                object_ids = set(impact["object_ids"])
                object_ids.add(str(row["payload_object_id"]))

                def _placeholders(values: list[str]) -> str:
                    return ",".join("?" for _ in values)

                # 被删关注卡片的敏感事项投影按 attention_card_id 关联，
                # 与 delete_subject 同法：在事务内删除投影组（outbox 随之
                # 级联删除），跨库墓碑由 delete_session 在落盘前处理。
                projection_groups = connection.execute(
                    """
                    SELECT g.group_id FROM sensitive_work_groups g
                    WHERE g.source_kind = 'attention_followup'
                      AND g.source_id IN (
                        SELECT attention_card_id FROM attention_cards
                        WHERE evidence_version_id IN (
                            SELECT evidence_version_id
                            FROM assessment_session_members
                            WHERE session_id = ?
                        )
                    )
                    """,
                    (session_id,),
                ).fetchall()
                projection_group_ids = [
                    str(item[0]) for item in projection_groups
                ]
                if projection_group_ids:
                    group_placeholders = _placeholders(projection_group_ids)
                    object_ids.update(
                        str(item[0])
                        for item in connection.execute(
                            f"""
                            SELECT envelope_object_id
                            FROM sensitive_work_projection_outbox
                            WHERE group_id IN ({group_placeholders})
                            """,
                            projection_group_ids,
                        ).fetchall()
                    )
                    connection.execute(
                        f"""
                        DELETE FROM sensitive_work_groups
                        WHERE group_id IN ({group_placeholders})
                        """,
                        projection_group_ids,
                    )
                evidence_ids = impact["evidence_version_ids"]
                if evidence_ids:
                    evidence_placeholders = _placeholders(evidence_ids)
                    connection.execute(
                        f"""
                        DELETE FROM attention_card_evidence_links
                        WHERE evidence_version_id IN ({evidence_placeholders})
                        """,
                        evidence_ids,
                    )
                    connection.execute(
                        f"""
                        DELETE FROM attention_cards
                        WHERE evidence_version_id IN ({evidence_placeholders})
                        """,
                        evidence_ids,
                    )
                    connection.execute(
                        f"""
                        DELETE FROM evidence_versions
                        WHERE evidence_version_id IN ({evidence_placeholders})
                        """,
                        evidence_ids,
                    )
                result_ids = impact["result_ids"]
                if result_ids:
                    result_placeholders = _placeholders(result_ids)
                    connection.execute(
                        f"""
                        DELETE FROM rank_contexts
                        WHERE result_id IN ({result_placeholders})
                        """,
                        result_ids,
                    )
                    connection.execute(
                        f"""
                        DELETE FROM subject_results
                        WHERE result_id IN ({result_placeholders})
                        """,
                        result_ids,
                    )
                assessment_ids = impact["assessment_ids"]
                if assessment_ids:
                    connection.execute(
                        f"""
                        DELETE FROM assessments
                        WHERE assessment_id IN ({_placeholders(assessment_ids)})
                        """,
                        assessment_ids,
                    )
                connection.execute(
                    """
                    DELETE FROM assessment_session_members
                    WHERE session_id = ?
                    """,
                    (session_id,),
                )
                # payload_object_id 外键是 RESTRICT，必须先删场次行再删对象。
                connection.execute(
                    "DELETE FROM assessment_sessions WHERE session_id = ?",
                    (session_id,),
                )
                import_ids = impact["import_ids"]
                if import_ids:
                    import_placeholders = _placeholders(import_ids)
                    connection.execute(
                        f"""
                        DELETE FROM source_fingerprints
                        WHERE import_id IN ({import_placeholders})
                        """,
                        import_ids,
                    )
                    connection.execute(
                        f"""
                        DELETE FROM assessment_imports
                        WHERE import_id IN ({import_placeholders})
                        """,
                        import_ids,
                    )
                connection.executemany(
                    "DELETE FROM encrypted_objects WHERE object_id = ?",
                    [(object_id,) for object_id in sorted(object_ids)],
                )
                result = {
                    "session_id": session_id,
                    "deleted": True,
                    "counts": {
                        "results": len(result_ids),
                        "assessments": len(assessment_ids),
                        "imports": len(import_ids),
                        "attention_cards": len(impact["attention_card_ids"]),
                    },
                }
                self._remember(
                    connection,
                    operation_id,
                    "evidence.session_delete",
                    result,
                )
        return result

    def _session_row(self, connection: Any, session_id: str) -> Any:
        row = connection.execute(
            """
            SELECT session_id, payload_object_id
            FROM assessment_sessions
            WHERE session_id = ?
            """,
            (session_id,),
        ).fetchone()
        if row is None:
            raise VaultError(
                "assessment_session_not_found",
                "成绩场次不存在",
                status_code=404,
            )
        return row

    def _session_impact(
        self,
        connection: Any,
        session_id: str,
    ) -> dict[str, Any]:
        """收集删除场次会牵连的全部行标识与加密对象标识。"""
        evidence_version_ids = [
            str(row[0])
            for row in connection.execute(
                """
                SELECT evidence_version_id
                FROM assessment_session_members
                WHERE session_id = ?
                """,
                (session_id,),
            ).fetchall()
        ]
        result_ids: list[str] = []
        attention_card_ids: list[str] = []
        object_ids: set[str] = set()
        if evidence_version_ids:
            placeholders = ",".join("?" for _ in evidence_version_ids)
            evidence_rows = connection.execute(
                f"""
                SELECT result_id, payload_object_id
                FROM evidence_versions
                WHERE evidence_version_id IN ({placeholders})
                """,
                evidence_version_ids,
            ).fetchall()
            result_ids = sorted({str(row["result_id"]) for row in evidence_rows})
            object_ids.update(
                str(row["payload_object_id"]) for row in evidence_rows
            )
            attention_rows = connection.execute(
                f"""
                SELECT attention_card_id, payload_object_id
                FROM attention_cards
                WHERE evidence_version_id IN ({placeholders})
                """,
                evidence_version_ids,
            ).fetchall()
            attention_card_ids = [
                str(row["attention_card_id"]) for row in attention_rows
            ]
            object_ids.update(
                str(row["payload_object_id"]) for row in attention_rows
            )
        assessment_ids: list[str] = []
        if result_ids:
            placeholders = ",".join("?" for _ in result_ids)
            result_rows = connection.execute(
                f"""
                SELECT assessment_id, payload_object_id
                FROM subject_results
                WHERE result_id IN ({placeholders})
                """,
                result_ids,
            ).fetchall()
            assessment_ids = sorted(
                {str(row["assessment_id"]) for row in result_rows}
            )
            object_ids.update(
                str(row["payload_object_id"]) for row in result_rows
            )
            object_ids.update(
                str(row[0])
                for row in connection.execute(
                    f"""
                    SELECT payload_object_id FROM rank_contexts
                    WHERE result_id IN ({placeholders})
                    """,
                    result_ids,
                ).fetchall()
            )
        import_ids: list[str] = []
        if assessment_ids:
            placeholders = ",".join("?" for _ in assessment_ids)
            assessment_rows = connection.execute(
                f"""
                SELECT import_id, payload_object_id
                FROM assessments
                WHERE assessment_id IN ({placeholders})
                """,
                assessment_ids,
            ).fetchall()
            object_ids.update(
                str(row["payload_object_id"]) for row in assessment_rows
            )
            # 一个导入批次的全部考试都随本场次删除时才连带删除批次，
            # 否则保留，避免误删其他场次的成绩。
            for import_id in sorted(
                {str(row["import_id"]) for row in assessment_rows}
            ):
                shared = connection.execute(
                    f"""
                    SELECT 1 FROM assessments
                    WHERE import_id = ?
                      AND assessment_id NOT IN ({placeholders})
                    LIMIT 1
                    """,
                    [import_id, *assessment_ids],
                ).fetchone()
                if shared is None:
                    import_ids.append(import_id)
            if import_ids:
                import_placeholders = ",".join("?" for _ in import_ids)
                object_ids.update(
                    str(row[0])
                    for row in connection.execute(
                        f"""
                        SELECT payload_object_id FROM assessment_imports
                        WHERE import_id IN ({import_placeholders})
                        """,
                        import_ids,
                    ).fetchall()
                )
        return {
            "evidence_version_ids": evidence_version_ids,
            "result_ids": result_ids,
            "assessment_ids": assessment_ids,
            "import_ids": import_ids,
            "attention_card_ids": attention_card_ids,
            "object_ids": object_ids,
        }

    def _insert_assessment(
        self,
        connection: Any,
        *,
        vmk: bytes,
        import_id: str,
        assessment: dict[str, object],
    ) -> int:
        title = self._text(assessment.get("title"), "考试名称", 500)
        subject_name = self._text(
            assessment.get("subject_name"),
            "学科名称",
            120,
        )
        occurred_on = self._text(
            assessment.get("occurred_on"),
            "考试日期",
            40,
        )
        try:
            datetime.fromisoformat(occurred_on)
        except ValueError as exc:
            raise VaultError(
                "assessment_date_invalid",
                "考试日期无效",
                status_code=422,
            ) from exc
        results = list(assessment.get("results") or [])
        if not results:
            raise VaultError(
                "assessment_results_empty",
                "考试没有学生结果",
                status_code=422,
            )
        assessment_id = uuid4().hex
        object_id = f"assessment-{assessment_id}"
        assessment_payload = {
            "title": title,
            "subject_name": subject_name,
            "max_score": self._optional_number(assessment.get("max_score")),
            "rank_scope": str(assessment.get("rank_scope") or "") or None,
            "participant_count": self._optional_integer(
                assessment.get("participant_count")
            ),
            "assessment_nature": str(
                assessment.get("assessment_nature") or ""
            ) or None,
            "rank_origin": str(assessment.get("rank_origin") or "") or None,
            "cohort_key": str(assessment.get("cohort_key") or "") or None,
            "ranking_rule_version": str(
                assessment.get("ranking_rule_version") or ""
            ) or None,
        }
        self.repository.put(
            connection,
            vmk=vmk,
            object_id=object_id,
            object_type="assessment",
            payload=assessment_payload,
        )
        connection.execute(
            """
            INSERT INTO assessments (
                assessment_id, import_id, payload_object_id,
                occurred_on, created_at
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (assessment_id, import_id, object_id, occurred_on, _iso()),
        )
        session_id = self._ensure_session(
            connection,
            vmk=vmk,
            assessment=assessment,
            fallback_title=title,
            fallback_date=occurred_on,
        )
        measure_role = str(assessment.get("measure_role") or "subject_score")
        if measure_role not in {"subject_score", "total_score"}:
            raise VaultError("assessment_measure_role_invalid", "学业指标类型无效", status_code=422)
        measure_key = hmac.new(
            vmk,
            f"academic-measure|{subject_name}".encode("utf-8"),
            hashlib.sha256,
        ).digest()
        for result in results:
            evidence_id = self._insert_result(
                connection,
                vmk=vmk,
                assessment_id=assessment_id,
                assessment_payload=assessment_payload,
                occurred_on=occurred_on,
                session_id=session_id,
                measure_role=measure_role,
                measure_key=measure_key,
                result=dict(result),
            )
            connection.execute(
                """
                INSERT INTO assessment_session_members (
                    session_id, evidence_version_id, measure_role,
                    measure_key, created_at
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (session_id, evidence_id, measure_role, measure_key, _iso()),
            )
        return len(results)

    def _insert_result(
        self,
        connection: Any,
        *,
        vmk: bytes,
        assessment_id: str,
        assessment_payload: dict[str, object],
        occurred_on: str,
        session_id: str,
        measure_role: str,
        measure_key: bytes,
        result: dict[str, object],
    ) -> str:
        subject_id = str(result.get("subject_id") or "")
        if not subject_id:
            subject_id = self._ensure_subject_from_identity(
                connection, vmk=vmk, result=result
            )
        self._subject_exists(connection, subject_id)
        result_state = str(result.get("result_state") or "")
        if result_state not in _RESULT_STATES:
            raise VaultError(
                "assessment_result_state_invalid",
                "成绩状态无效",
                status_code=422,
            )
        score = self._optional_number(result.get("score"))
        if result_state in {"normal", "makeup"} and score is None:
            raise VaultError(
                "assessment_score_required",
                "正常成绩和补考成绩必须保留数值，0 分也是有效成绩",
                status_code=422,
            )
        if result_state not in {"normal", "makeup"} and score is not None:
            raise VaultError(
                "assessment_score_state_conflict",
                "缺考、免考、缺失或未完成不能伪装成数值成绩",
                status_code=422,
            )
        result_id = uuid4().hex
        result_object_id = f"subject-result-{result_id}"
        result_payload = {
            "score": score,
            "result_state": result_state,
            "teacher_note": str(result.get("teacher_note") or "") or None,
            "grade_level": self._optional_grade_level(result.get("grade_level")),
        }
        self.repository.put(
            connection,
            vmk=vmk,
            object_id=result_object_id,
            object_type="subject_result",
            payload=result_payload,
        )
        connection.execute(
            """
            INSERT INTO subject_results (
                result_id, assessment_id, subject_id,
                payload_object_id, result_state, created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                result_id,
                assessment_id,
                subject_id,
                result_object_id,
                result_state,
                _iso(),
            ),
        )
        rank = self._optional_integer(result.get("rank"))
        class_rank = self._optional_integer(result.get("class_rank"))
        if rank is not None or class_rank is not None:
            rank_context_id = uuid4().hex
            rank_object_id = f"rank-context-{rank_context_id}"
            self.repository.put(
                connection,
                vmk=vmk,
                object_id=rank_object_id,
                object_type="rank_context",
                payload={
                    "rank": rank,
                    "class_rank": class_rank,
                    "rank_scope": assessment_payload["rank_scope"],
                    "participant_count": assessment_payload[
                        "participant_count"
                    ],
                    "rank_origin": str(
                        result.get("rank_origin")
                        or assessment_payload.get("rank_origin")
                        or ""
                    ) or None,
                    "cohort_key": str(
                        result.get("cohort_key")
                        or assessment_payload.get("cohort_key")
                        or ""
                    ) or None,
                    "ranking_rule_version": str(
                        result.get("ranking_rule_version")
                        or assessment_payload.get("ranking_rule_version")
                        or ""
                    ) or None,
                },
            )
            connection.execute(
                """
                INSERT INTO rank_contexts (
                    rank_context_id, result_id,
                    payload_object_id, created_at
                ) VALUES (?, ?, ?, ?)
                """,
                (rank_context_id, result_id, rank_object_id, _iso()),
            )
        evidence_id = uuid4().hex
        evidence_object_id = f"evidence-version-{evidence_id}"
        self.repository.put(
            connection,
            vmk=vmk,
            object_id=evidence_object_id,
            object_type="assessment_evidence_version",
            payload={
                **assessment_payload,
                **result_payload,
                "occurred_on": occurred_on,
                "session_id": session_id,
                "measure_role": measure_role,
                "measure_key": measure_key.hex(),
                "source_version": hmac.new(
                    vmk,
                    json.dumps(
                        {
                            "assessment": assessment_payload,
                            "result": result_payload,
                            "occurred_on": occurred_on,
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                        separators=(",", ":"),
                    ).encode("utf-8"),
                    hashlib.sha256,
                ).hexdigest(),
                "teacher_confirmed_at": _iso(),
                "source_snapshot": True,
                "source_write_back": False,
            },
        )
        connection.execute(
            """
            INSERT INTO evidence_versions (
                evidence_version_id, result_id, version,
                payload_object_id, state, created_at
            ) VALUES (?, ?, 1, ?, 'active', ?)
            """,
            (evidence_id, result_id, evidence_object_id, _iso()),
        )
        return evidence_id

    def _ensure_session(
        self,
        connection: Any,
        *,
        vmk: bytes,
        assessment: dict[str, object],
        fallback_title: str,
        fallback_date: str,
    ) -> str:
        supplied = dict(assessment.get("session") or {})
        required = (
            "title",
            "academic_year",
            "term",
            "grade",
            "exam_type",
            "comparison_series",
            "occurred_on",
            "source_reference",
        )
        metadata_complete = all(str(supplied.get(key) or "").strip() for key in required)
        payload = {
            "title": str(supplied.get("title") or fallback_title),
            "academic_year": supplied.get("academic_year"),
            "term": supplied.get("term"),
            "grade": supplied.get("grade"),
            "exam_type": supplied.get("exam_type"),
            "comparison_series": supplied.get("comparison_series"),
            "occurred_on": str(supplied.get("occurred_on") or fallback_date),
            "source_reference": supplied.get("source_reference"),
            "teacher_confirmed_at": str(supplied.get("teacher_confirmed_at") or _iso()),
            "raw_file_retained": False,
            "metadata_complete": metadata_complete,
        }
        fingerprint = self._session_fingerprint(vmk, payload)
        row = connection.execute(
            "SELECT session_id FROM assessment_sessions WHERE source_fingerprint = ?",
            (fingerprint,),
        ).fetchone()
        if row is not None:
            return str(row["session_id"])
        session_id = uuid4().hex
        object_id = f"assessment-session-{session_id}"
        self.repository.put(
            connection,
            vmk=vmk,
            object_id=object_id,
            object_type="assessment_session",
            payload=payload,
        )
        connection.execute(
            """
            INSERT INTO assessment_sessions (
                session_id, payload_object_id, source_fingerprint, state,
                metadata_complete, created_at, updated_at
            ) VALUES (?, ?, ?, 'active', ?, ?, ?)
            """,
            (session_id, object_id, fingerprint, int(metadata_complete), _iso(), _iso()),
        )
        return session_id

    def _evidence_from_row(
        self,
        connection: Any,
        vmk: bytes,
        row: Any,
    ) -> dict[str, object]:
        payload, _ = self.repository.get(
            connection,
            vmk=vmk,
            object_id=str(row["evidence_object"]),
        )
        rank_row = connection.execute(
            """
            SELECT payload_object_id FROM rank_contexts
            WHERE result_id = ?
            """,
            (str(row["result_id"]),),
        ).fetchone()
        rank_context = None
        if rank_row is not None:
            rank_context, _ = self.repository.get(
                connection,
                vmk=vmk,
                object_id=str(rank_row["payload_object_id"]),
            )
        return {
            "evidence_version_id": str(row["evidence_version_id"]),
            "version": int(row["version"]),
            "state": str(row["state"]),
            "result_id": str(row["result_id"]),
            "subject_id": str(row["subject_id"]),
            "result_state": str(row["result_state"]),
            "source_kind": str(row["source_kind"]),
            **payload,
            "rank_context": rank_context,
        }

    @staticmethod
    def _comparability(
        older: dict[str, object],
        newer: dict[str, object],
    ) -> tuple[str, list[str]]:
        limitations: list[str] = []
        if (
            older.get("subject_name") != newer.get("subject_name")
            or older.get("max_score") is None
            or newer.get("max_score") is None
            or not older.get("assessment_nature")
            or not newer.get("assessment_nature")
        ):
            return "insufficient_information", ["学科、满分或考试性质不完整"]
        special_states = {
            str(older.get("result_state")),
            str(newer.get("result_state")),
        }
        if special_states - {"normal", "makeup"}:
            return "not_comparable", ["存在缺考、免考、缺失或未完成"]
        if older.get("assessment_nature") != newer.get("assessment_nature"):
            return "not_comparable", ["考试性质不同"]
        if older.get("max_score") != newer.get("max_score"):
            limitations.append("满分不同")
        if older.get("rank_scope") != newer.get("rank_scope"):
            limitations.append("排名范围不同")
        if older.get("participant_count") != newer.get("participant_count"):
            limitations.append("参考人数不同")
        if "makeup" in special_states:
            limitations.append("包含补考")
        if limitations:
            return "reference_only", limitations
        return "directly_comparable", []

    def _ensure_subject_from_identity(
        self,
        connection: Any,
        *,
        vmk: bytes,
        result: dict[str, object],
    ) -> str:
        identity = result.get("subject_identity")
        if not isinstance(identity, dict):
            raise VaultError(
                "support_subject_not_found",
                "学生支持对象不存在",
                status_code=404,
            )
        display_name = str(identity.get("display_name") or "").strip()
        if not display_name:
            raise VaultError(
                "assessment_subject_identity_invalid",
                "新建学生档案必须提供姓名",
                status_code=422,
            )
        if self._subject_ensurer is None:
            raise VaultError(
                "assessment_subject_ensurer_unavailable",
                "当前部署不支持随成绩登记新建学生档案",
                status_code=422,
            )
        class_label = str(identity.get("class_label") or "").strip() or None
        student_code = str(identity.get("student_code") or "").strip() or None
        # 优先按花名册身份建档（班级+姓名唯一命中，或姓名全校唯一），
        # 保证「一个学生自始至终只有一条档案」；花名册查不到才退回临时身份。
        roster_student = self._match_roster_student(class_label, display_name)
        if roster_student is not None:
            return self._subject_ensurer(
                connection,
                vmk=vmk,
                source_student_id=roster_student.source_key,
                display_name=roster_student.display_name,
                class_label=roster_student.class_label,
                student_code=roster_student.student_code,
            )
        return self._subject_ensurer(
            connection,
            vmk=vmk,
            source_student_id=(
                f"evidence-upload|{class_label or ''}|{student_code or display_name}"
            ),
            display_name=display_name,
            class_label=class_label,
            student_code=student_code,
        )

    def _match_roster_student(
        self,
        class_label: str | None,
        display_name: str,
    ) -> ExistingStudent | None:
        if self._roster_source is None:
            return None
        students, _revision = self._roster_source.snapshot()
        if class_label:
            matches = [
                student
                for student in students
                if student.class_label == class_label
                and student.display_name == display_name
            ]
            if len(matches) == 1:
                return matches[0]
            if matches:
                return None
        matches = [
            student for student in students if student.display_name == display_name
        ]
        return matches[0] if len(matches) == 1 else None

    @staticmethod
    def _subject_exists(connection: Any, subject_id: str) -> None:
        if connection.execute(
            "SELECT 1 FROM student_subject_links WHERE subject_id = ?",
            (subject_id,),
        ).fetchone() is None:
            raise VaultError(
                "support_subject_not_found",
                "学生支持对象不存在",
                status_code=404,
            )

    def _idempotent(
        self,
        operation_id: str,
        operation_type: str,
    ) -> dict[str, object] | None:
        with closing(self.database.connect()) as connection:
            row = connection.execute(
                """
                SELECT operation_type, result_json FROM idempotency_ledger
                WHERE operation_id = ?
                """,
                (operation_id,),
            ).fetchone()
        if row is None:
            return None
        if str(row["operation_type"]) != operation_type:
            raise VaultError(
                "vault_operation_conflict",
                "此操作编号已经用于另一项操作",
                status_code=409,
            )
        return dict(json.loads(str(row["result_json"])))

    @staticmethod
    def _remember(
        connection: Any,
        operation_id: str,
        operation_type: str,
        result: dict[str, object],
    ) -> None:
        connection.execute(
            """
            INSERT INTO idempotency_ledger (
                operation_id, operation_type, result_json, created_at
            ) VALUES (?, ?, ?, ?)
            """,
            (
                operation_id,
                operation_type,
                json.dumps(result, sort_keys=True),
                _iso(),
            ),
        )

    @staticmethod
    def _session_fingerprint(vmk: bytes, payload: dict[str, object]) -> bytes:
        """场次归并指纹：只覆盖 _SESSION_REQUIRED_FIELDS 的 8 个元数据字段。

        teacher_confirmed_at、raw_file_retained、metadata_complete 不参与
        指纹：同一场考试分多次（分学科）确认时，每次确认时刻各不相同，
        但成绩必须归并到同一 session。登记（_ensure_session）与元数据
        更正（update_session_metadata）必须使用同一算法。
        """
        metadata = {key: payload.get(key) for key in _SESSION_REQUIRED_FIELDS}
        signature = json.dumps(metadata, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return hmac.new(vmk, f"assessment-session|{signature}".encode("utf-8"), hashlib.sha256).digest()

    @classmethod
    def _validated_session_updates(
        cls,
        fields: dict[str, object],
    ) -> dict[str, str]:
        unknown = sorted(set(fields) - set(_SESSION_EDITABLE_FIELDS))
        if unknown:
            raise VaultError(
                "assessment_session_field_invalid",
                f"场次字段不可更正：{', '.join(unknown)}",
                status_code=422,
            )
        updates: dict[str, str] = {}
        for key, value in fields.items():
            label, maximum = _SESSION_EDITABLE_LABELS[key]
            clean = cls._text(value, label, maximum)
            if key == "occurred_on":
                try:
                    datetime.fromisoformat(clean)
                except ValueError as exc:
                    raise VaultError(
                        "assessment_date_invalid",
                        "考试日期无效",
                        status_code=422,
                    ) from exc
            updates[key] = clean
        return updates

    @staticmethod
    def _text(value: object, label: str, maximum: int) -> str:
        clean = str(value or "").strip()
        if not clean or len(clean) > maximum:
            raise VaultError(
                "assessment_text_invalid",
                f"{label}不能为空且不能超过 {maximum} 个字符",
                status_code=422,
            )
        return clean

    @staticmethod
    def _optional_grade_level(value: object) -> str | None:
        """等级原样保存：trim 后非空才存，超长拒绝；缺省为 None。"""
        if value is None:
            return None
        clean = str(value).strip()
        if not clean:
            return None
        if len(clean) > 20:
            raise VaultError(
                "assessment_grade_level_invalid",
                "等级不能超过 20 个字符",
                status_code=422,
            )
        return clean

    @staticmethod
    def _optional_number(value: object) -> float | None:
        if value is None or value == "":
            return None
        try:
            return float(value)
        except (TypeError, ValueError) as exc:
            raise VaultError(
                "assessment_number_invalid",
                "成绩或满分必须是数字",
                status_code=422,
            ) from exc

    @staticmethod
    def _optional_integer(value: object) -> int | None:
        if value is None or value == "":
            return None
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise VaultError(
                "assessment_integer_invalid",
                "排名或人数必须是整数",
                status_code=422,
            ) from exc
        if parsed < 0:
            raise VaultError(
                "assessment_integer_invalid",
                "排名或人数不能为负数",
                status_code=422,
            )
        return parsed


__all__ = [
    "AssessmentEvidenceService",
    "AssessmentEvidenceSource",
    "ConfirmedSpreadsheetAdapter",
    "ExistingMathAssessmentAdapter",
]
