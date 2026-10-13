"""智学网 Excel 工作簿解析：识别「小题得分明细」与「客观题作答明细」两种表。

只解析单元格文本，不在这里做名单匹配与分值检查——那些属于 planning。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_SHEET_NAME_IDENTITY = {"姓名"}
_EXAM_NO_HEADERS = {"准考证号", "考号", "考生号", "学籍号"}
_SUMMARY_HEADERS = {
    "客观分": "objective",
    "主观分": "subjective",
    "总分": "total",
    "单选题": "single_choice_total",
    "多选题": "multi_choice_total",
    "填空题": "fill_total",
    "解答题": "solution_total",
    "客观题": "objective",
}
_KEY_ROW_MARKER = "客观题答案"
_ABSENT_MARKERS = {"缺考", "缺", "未考", "—", "-", "--", "/"}

_CLASS_HINT_RE = re.compile(r"(\d+)\s*班")
_SUB_COL_RE = re.compile(r"^(\d+)\s*[(（]\s*(\d+)\s*[)）]$")
_COMBINED_COL_RE = re.compile(r"^(\d+)\s*[-~～—–]+\s*(\d+)$")
_QUESTION_COL_RE = re.compile(r"^\d{1,3}$")
_OPTION_LETTERS_RE = re.compile(r"^[A-Fa-f]{1,6}$")

_HEADER_SEARCH_ROWS = 5


@dataclass(frozen=True, slots=True)
class ParsedColumn:
    index: int
    header: str
    role: str  # identity|summary|question|sub|combined|unknown
    key: str = ""
    question_no: int | None = None
    sub_index: int | None = None
    combined_end: int | None = None


@dataclass(slots=True)
class ParsedRow:
    excel_row: int
    cells: dict[int, str]


@dataclass(slots=True, eq=False)
class ParsedSheet:
    name: str
    kind: str  # scores|options|unrecognized
    title: str = ""
    class_hint: str | None = None
    header_row: int = 0
    columns: list[ParsedColumn] = field(default_factory=list)
    rows: list[ParsedRow] = field(default_factory=list)
    answer_key_row: int | None = None
    answer_key_cells: dict[int, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


@dataclass(slots=True)
class ParsedUpload:
    upload_id: str
    filename: str
    sheets: list[ParsedSheet]


def normalize_name(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def class_digits(value: Any) -> str | None:
    """Extract the class number from '9', '9班', '八年级9班', '八(9)班'."""
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"(\d+)", text)
    return match.group(1) if match else None


def parse_score(value: Any) -> float | None:
    """Numeric score or None for blank/absent markers."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if not text or text in _ABSENT_MARKERS:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def is_absent_marker(value: Any) -> bool:
    return str(value or "").strip() in _ABSENT_MARKERS


def normalize_option_letters(value: Any) -> str | None:
    text = str(value or "").strip().upper()
    if not text:
        return None
    if not _OPTION_LETTERS_RE.fullmatch(text):
        return None
    return "".join(sorted(set(text)))


def _classify_column(header: str) -> ParsedColumn:
    clean = str(header or "").strip()
    if not clean:
        return ParsedColumn(0, header, "unknown")
    if clean in _SHEET_NAME_IDENTITY:
        return ParsedColumn(0, clean, "identity", key="name")
    if clean in _EXAM_NO_HEADERS:
        return ParsedColumn(0, clean, "identity", key="exam_no")
    if clean == "班级":
        return ParsedColumn(0, clean, "identity", key="class")
    if clean in ("序号", "座号", "科目", "场次"):
        return ParsedColumn(0, clean, "identity", key="other")
    summary_key = _SUMMARY_HEADERS.get(clean)
    if summary_key is not None:
        return ParsedColumn(0, clean, "summary", key=summary_key)
    if any(
        token in clean
        for token in ("排名", "名次", "占比", "得分率", "学情", "均分")
    ):
        return ParsedColumn(0, clean, "summary", key="other")
    if _QUESTION_COL_RE.fullmatch(clean):
        return ParsedColumn(0, clean, "question", question_no=int(clean))
    sub = _SUB_COL_RE.match(clean)
    if sub is not None:
        return ParsedColumn(
            0,
            clean,
            "sub",
            question_no=int(sub.group(1)),
            sub_index=int(sub.group(2)),
        )
    combined = _COMBINED_COL_RE.match(clean)
    if combined is not None and int(combined.group(1)) < int(combined.group(2)):
        return ParsedColumn(
            0,
            clean,
            "combined",
            question_no=int(combined.group(1)),
            combined_end=int(combined.group(2)),
        )
    return ParsedColumn(0, clean, "unknown")


def _row_texts(row: tuple[Any, ...]) -> list[str]:
    return [str(cell).strip() if cell is not None else "" for cell in row]


def _find_header(rows: list[tuple[int, list[str]]]) -> int:
    for index, (_excel_row, texts) in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        if "姓名" in texts:
            return index
    return -1


def _first_text(texts: list[str]) -> str:
    for text in texts:
        if text:
            return text
    return ""


def parse_sheet(name: str, raw_rows: list[tuple[int, list[str]]]) -> ParsedSheet:
    header_index = _find_header(raw_rows)
    if header_index < 0:
        return ParsedSheet(name=name, kind="unrecognized")
    header_excel_row, header_texts = raw_rows[header_index]
    title = ""
    for _excel_row, texts in raw_rows[:header_index]:
        first = _first_text(texts)
        if first:
            title = first
            break
    hint_match = _CLASS_HINT_RE.search(title)
    columns: list[ParsedColumn] = []
    for index, header in enumerate(header_texts):
        column = _classify_column(header)
        columns.append(
            ParsedColumn(
                index=index,
                header=header,
                role=column.role,
                key=column.key,
                question_no=column.question_no,
                sub_index=column.sub_index,
                combined_end=column.combined_end,
            )
        )
    data_rows = [
        ParsedRow(excel_row=excel_row, cells={i: t for i, t in enumerate(texts) if t != ""})
        for excel_row, texts in raw_rows[header_index + 1 :]
        if any(texts)
    ]
    answer_key_row: int | None = None
    answer_key_cells: dict[int, str] = {}
    for row in data_rows:
        if _first_text(
            [row.cells.get(index, "") for index in range(len(header_texts))]
        ) == _KEY_ROW_MARKER:
            answer_key_row = row.excel_row
            answer_key_cells = dict(row.cells)
            break
    kind = "options" if answer_key_row is not None else "scores"
    rows = [row for row in data_rows if row.excel_row != answer_key_row]
    return ParsedSheet(
        name=name,
        kind=kind,
        title=title,
        class_hint=hint_match.group(1) if hint_match else None,
        header_row=header_excel_row,
        columns=columns,
        rows=rows,
        answer_key_row=answer_key_row,
        answer_key_cells=answer_key_cells,
    )


def parse_upload(upload_id: str, filename: str, path: Path) -> ParsedUpload:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[ParsedSheet] = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            raw_rows: list[tuple[int, list[str]]] = [
                (index, _row_texts(row))
                for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
            ]
            sheets.append(parse_sheet(str(sheet_name), raw_rows))
        return ParsedUpload(upload_id=upload_id, filename=filename, sheets=sheets)
    finally:
        workbook.close()


def sheet_inventory(sheets: list[ParsedSheet]) -> list[dict[str, Any]]:
    return [
        {
            "name": sheet.name,
            "kind": sheet.kind,
            "class_hint": sheet.class_hint,
            "row_count": len(sheet.rows),
        }
        for sheet in sheets
    ]
