"""考试场次分析数据装配：数据库 → 报告/页面共用的数据模型。

本模块只负责把评分库、评分依据与题库侧信息装配成 SessionAnalysisData；
HTML 渲染、AI 叙述与报告导出仍留在报告导出模块，
内容生成模型配置解析在 backend/model_profiles/content_generation.py。
"""

from __future__ import annotations

import base64
import html
import io
import json
import re
import sqlite3
import statistics
import zipfile
from dataclasses import dataclass, field, replace
from datetime import datetime
from pathlib import Path
from typing import Any

from PIL import Image

from backend.scan_grading.answer_key_utils import answer_forms_map
from backend.analytics.service import load_session_score_type_maps
from backend.repositories.access import (
    GradingRepositoryAccess,
    as_grading_repositories,
)
from backend.scan_grading.grading_completeness import resolve_grading_completeness
from path_manager import resolve_stored_file_path
from backend.question_id_contract import (
    QuestionIdContractError,
    canonicalize_question_document,
    question_id_coordinates,
    resolve_known_question_id,
)
from backend.reporting.report import natural_question_order

SMALL_SAMPLE_LIMIT = 10


# 客观题批改记录里的内部调试码 → 家长可读文本；
# 口径与前端 frontend/src/utils/grading-reasons.ts 保持一致。
_OBJECTIVE_ANSWER_PREFIX = "objective_answer="
_REASON_CODE_LABELS = {
    "objective_api_disabled": "客观题识别未启用，已转教师复核",
    "objective_api_not_configured": "客观题识别模型配置不完整，已转教师复核",
    "objective_paper_model_failed": "客观题识别请求失败，已转教师复核",
    "objective_paper_region_failed": "无法读取选填题作答区域",
    "objective_region_not_found": "未找到此题的有效作答区域",
    "missing_question_result": "AI 未返回此题的识别结果",
    "missing_answer_field": "AI 未返回此题的作答内容",
    "duplicate_question_result": "AI 返回了重复的识别结果",
    "paper_key_mismatch": "识别结果与当前答卷不一致",
    "low_confidence": "作答辨识度较低，需要教师复核",
    "needs_review": "AI 建议教师复核",
    "objective_needs_review": "客观题识别结果需要教师复核",
    "objective_score_uncertain": "答案识别存在不确定性",
    "prompt_injection_or_score_bait": "作答区出现与答题无关的批改指令",
    "discarded_answer_only": "只识别到已经涂抹或作废的答案",
    "assignment_changed": "批改期间答卷匹配发生变化",
    "missing_detail_question_ids": "AI 未完整返回所有小问的评分结果",
    "duplicate_detail_question_id": "AI 重复返回了同一小问",
    "unexpected_detail_question_id": "AI 返回了不属于本题的小问",
    "no_numeric_value": "未能可靠识别填写的数值，需要教师确认",
    "multiple options selected": "识别到选择了多个选项",
    # 阅卷模型自由发挥的英文 error_summary 代码（历史数据实测值，见错因体系改造方案）。
    "blank_or_no_valid_work": "未识别到有效作答内容",
    "no_valid_answer": "未识别到有效作答内容",
    "insufficient_work_shown": "作答过程不完整",
    "incomplete_process": "作答过程不完整",
    "missing_steps": "作答缺少关键步骤",
    "unclear_or_missing_process": "作答过程不完整",
    "formatting_issue": "作答书写不规范",
    "not_simplified": "结果未化到最简",
    "review_required": "需要教师复核确认",
    "calculation_error": "计算有误",
    "sign_error": "符号处理有误",
    "incorrect_reasoning": "推理过程有误",
    "incorrect_reason": "所给理由有误",
    "concept_misunderstanding": "概念理解有误",
    "conceptual_error": "概念理解有误",
}

_CJK_RE = re.compile(r"[\u3400-\u9fff]")


def has_cjk(text: str) -> bool:
    return bool(_CJK_RE.search(text))

# 分数段按满分等比缩放（满分 100 时即方案 §3 的固定分段）。
BAND_CUTOFFS = (85, 70, 60, 40, 0)
_PASS_CUTOFF = 60


@dataclass
class QuestionInfo:
    question_id: str
    question_type: str
    max_score: float
    stem_summary: str
    canonical_answer: str
    attempts: int = 0
    score_sum: float = 0.0
    question_text: str = ""
    question_markup: str = ""
    reference_analysis: str = ""
    reference_images: list[tuple[str, bytes]] = field(default_factory=list)

    @property
    def class_rate(self) -> float | None:
        full = self.max_score * self.attempts
        if full <= 0:
            return None
        return self.score_sum / full

    @property
    def class_avg(self) -> float | None:
        if self.attempts <= 0:
            return None
        return self.score_sum / self.attempts


@dataclass
class StudentQuestionRecord:
    question_id: str
    score: float
    max_score: float
    deduction_reason: str
    error_category: str
    error_summary: str
    secondary_errors: list[str] = field(default_factory=list)
    student_answer: str = ""
    evidence_steps: list[str] = field(default_factory=list)
    missing_steps: list[str] = field(default_factory=list)
    teacher_confirmed: bool = False
    teacher_comment: str = ""
    # 独立扣分的评分步骤（沿用前步错误的不在列）；教师锁定时取自教师逐步改分。
    failed_steps: list[dict[str, Any]] = field(default_factory=list)
    # 新教师分步记录即使没有可整理错因，也不能回退为整题 AI 错因。
    teacher_step_reviewed: bool = False

    @property
    def lost(self) -> bool:
        return self.max_score > 0 and self.score < self.max_score - 1e-6

    @property
    def lost_points(self) -> float:
        return max(0.0, self.max_score - self.score) if self.lost else 0.0


@dataclass
class StudentReportData:
    result_id: int
    student_id: int
    student_code: str
    student_name: str
    class_name: str
    student_score: float
    needs_review: bool
    graded_at: str
    rank: int = 0
    records: list[StudentQuestionRecord] = field(default_factory=list)
    material_notes: list[str] = field(default_factory=list)
    knowledge_mastery: dict[str, dict[str, Any]] = field(default_factory=dict)

    @property
    def lost_points_total(self) -> float:
        return sum(record.lost_points for record in self.records)


@dataclass
class SessionAnalysisData:
    session_id: int
    session_name: str
    subject: str
    full_score: float
    graded_at: str
    present: int
    roster_absent: int
    small_sample: bool
    questions: list[QuestionInfo]
    students: list[StudentReportData]
    skipped: list[dict[str, Any]]
    stats: dict[str, Any]
    knowledge_backfill: dict[str, list[dict[str, str]]]
    attendance_by_class: dict[str, dict[str, int]] = field(default_factory=dict)
    class_name: str | None = None
    rubric: dict[str, Any] = field(default_factory=dict)
    question_assessments: dict[str, dict[str, Any]] = field(default_factory=dict)
    knowledge_structure: dict[str, Any] = field(default_factory=dict)


def assemble_session_analysis(
    db: GradingRepositoryAccess | Any,
    session_id: int,
    *,
    data_root: Path | None = None,
    page_only: bool = False,
    include_answer_evidence: bool = False,
    include_knowledge: bool = False,
) -> SessionAnalysisData:
    """轻量页面按需读取标签或作答文字，不重新识别原卷。"""
    repositories = as_grading_repositories(db)
    resolved_data_root = data_root or infer_data_root(repositories.db_path)
    snapshot = repositories.reports.get_session_report_snapshot(
        int(session_id),
        question_bank_path=None if page_only and not include_knowledge else question_bank_db_path(repositories.db_path),
    )
    session_row = repositories.sessions.get_grading_session(int(session_id)) or {}
    rubric = load_rubric(session_row, resolved_data_root)
    score_map, type_map = load_session_score_type_maps(
        session_row or None,
        data_root=resolved_data_root,
    )
    stem_map = _stem_summary_map(rubric)
    answer_map = _canonical_answer_map(session_row, resolved_data_root)

    attendance_by_identity = {
        (
            str(row.get("student_code") or ""),
            str(row.get("class_name") or "未分班"),
        ): str(row.get("attendance_status") or "").strip()
        for row in snapshot.attendance
    }
    present_count = sum(
        1
        for row in snapshot.attendance
        if str(row.get("attendance_status") or "").strip() == "present"
    )
    roster_absent = sum(
        1
        for row in snapshot.attendance
        if str(row.get("attendance_status") or "").strip() in {"absent", "scan_issue"}
    )

    details_by_result: dict[int, list[dict[str, Any]]] = {}
    for detail in snapshot.details:
        details_by_result.setdefault(int(detail.get("result_id") or 0), []).append(
            detail
        )

    locks_by_student_question = {
        (
            int(lock["student_id"]),
            resolve_known_question_id(str(lock["question_id"]), score_map)
            or str(lock["question_id"]),
        ): lock
        for lock in snapshot.locks
    }
    students: list[StudentReportData] = []
    skipped: list[dict[str, Any]] = []
    covered_identities: set[tuple[str, str]] = set()
    for result in snapshot.results:
        identity = (
            str(result.get("student_code") or ""),
            str(result.get("class_name") or "未分班"),
        )
        covered_identities.add(identity)
        label = {
            "student_id": int(result.get("student_id") or 0),
            "class_name": identity[1],
            "student_code": identity[0],
            "student_name": str(result.get("student_name") or ""),
        }
        attendance_status = attendance_by_identity.get(identity)
        if attendance_status in {"absent", "scan_issue"}:
            skipped.append(
                {**label, "reason": "缺考" if attendance_status == "absent" else "扫描异常"}
            )
            continue
        result_id = int(result.get("result_id") or 0)
        completeness = resolve_grading_completeness(
            result.get("raw_json"),
            rubric=rubric if isinstance(rubric, dict) else None,
            details=details_by_result.get(result_id, []),
        )
        status = (
            str(completeness.get("status") or "").strip()
            if isinstance(completeness, dict)
            else ""
        )
        if status != "complete":
            missing = (
                completeness.get("missing_question_ids")
                if isinstance(completeness, dict)
                else []
            )
            missing_text = (
                "、".join(str(item) for item in missing if str(item).strip())
                if isinstance(missing, list)
                else ""
            )
            skipped.append(
                {
                    **label,
                    "reason": (
                        f"批改不完整（缺失题目：{missing_text}）"
                        if missing_text
                        else "批改结果无效或不完整"
                    ),
                }
            )
            continue
        raw_payload = result.get("raw_json")
        if isinstance(raw_payload, str):
            try:
                raw_payload = json.loads(raw_payload)
            except json.JSONDecodeError:
                raw_payload = {}
        students.append(
            StudentReportData(
                result_id=result_id,
                student_id=int(result.get("student_id") or 0),
                student_code=identity[0],
                student_name=str(result.get("student_name") or ""),
                class_name=identity[1],
                student_score=float(result.get("student_score") or 0),
                needs_review=bool(result.get("needs_human_review")),
                graded_at=str(result.get("graded_at") or "")[:10],
                records=_merge_student_records(
                    repositories,
                    result_id,
                    details_by_result.get(result_id, []),
                    score_map,
                    student_answers=_student_answer_map(result.get("raw_json")) if not page_only or include_answer_evidence else None,
                    grading_evidence=_grading_detail_map(result.get("raw_json")) if not page_only or include_answer_evidence else None,
                    include_secondary_errors=not page_only,
                    teacher_reviews=(
                        raw_payload.get("teacher_reviews")
                        if isinstance(raw_payload, dict)
                        else None
                    ),
                    teacher_locks={
                        question_id: lock
                        for (student_id, question_id), lock in locks_by_student_question.items()
                        if student_id == int(result.get("student_id") or 0)
                    },
                    rubric=rubric if isinstance(rubric, dict) else None,
                ),
            )
        )
    for student in students:
        for record in student.records:
            lock = locks_by_student_question.get((student.student_id, record.question_id))
            if lock is not None:
                record.teacher_confirmed = True
                record.teacher_comment = _sanitize_grading_text(lock.get("deduction_reason"))
        # Match the review queue's per-item rules; the original whole-paper AI
        # flag remains historical after a teacher confirms a flagged answer.
        from backend.review.service import _is_substantive_review_reason
        from backend.scan_grading.grading_completeness import is_objective_detail
        pending_review = False
        for detail in details_by_result.get(student.result_id, []):
            qid = resolve_known_question_id(str(detail.get("question_id") or ""), score_map) or str(detail.get("question_id") or "")
            lock = locks_by_student_question.get((student.student_id, qid))
            if lock is not None:
                pending_review |= abs(float(lock.get("max_score") or 0) - float(score_map.get(qid) or 0)) > 1e-6
            else:
                pending_review |= _is_substantive_review_reason(
                    str(detail.get("deduction_reason") or ""),
                    str(detail.get("error_category") or ""),
                    detail.get("confidence_score"),
                    objective=is_objective_detail(detail),
                )
        student.needs_review = pending_review
    # 名册中缺考且无结果的学生也要进入未生成清单。
    attendance_labels = {"absent": "缺考", "scan_issue": "扫描异常"}
    for row in snapshot.attendance:
        status = str(row.get("attendance_status") or "").strip()
        identity = (
            str(row.get("student_code") or ""),
            str(row.get("class_name") or "未分班"),
        )
        if status in attendance_labels and identity not in covered_identities:
            skipped.append(
                {
                    "class_name": identity[1],
                    "student_code": identity[0],
                    "student_name": str(row.get("student_name") or ""),
                    "student_id": int(row.get("student_id") or 0),
                    "reason": attendance_labels[status],
                }
            )

    # 名次：按总分降序，同分同名次（method="min"）。
    ordered = sorted(
        students,
        key=lambda item: (-item.student_score, item.student_code, item.student_name),
    )
    previous_score: float | None = None
    previous_rank = 0
    for index, student in enumerate(ordered, start=1):
        if previous_score is None or student.student_score < previous_score - 1e-9:
            previous_rank = index
            previous_score = student.student_score
        student.rank = previous_rank
    students = ordered

    questions: list[QuestionInfo] = []
    for qid in natural_question_order([str(qid) for qid in score_map]):
        questions.append(
            QuestionInfo(
                question_id=qid,
                question_type=str(type_map.get(qid) or ""),
                max_score=float(score_map.get(qid) or 0),
                stem_summary=stem_map.get(qid) or "",
                canonical_answer=answer_map.get(qid) or "",
            )
        )
    info_by_qid = {info.question_id: info for info in questions}
    for student in students:
        for record in student.records:
            info = info_by_qid.get(record.question_id)
            if info is None:
                continue
            info.attempts += 1
            info.score_sum += min(record.score, record.max_score or record.score)

    # 大题满分与实际评分的小问同时存在时，只展示真正参与评分的行。
    scored_parents = {
        parent_question_id(info.question_id)
        for info in questions
        if info.attempts and parent_question_id(info.question_id) != info.question_id
    }
    questions = [
        info for info in questions
        if info.attempts or info.question_id not in scored_parents
    ]
    attendance_by_class: dict[str, dict[str, int]] = {}
    for row in snapshot.attendance:
        counts = attendance_by_class.setdefault(
            str(row.get("class_name") or "未分班"), {"present": 0, "absent": 0},
        )
        status = str(row.get("attendance_status") or "").strip()
        counts["present"] += int(status == "present")
        counts["absent"] += int(status in {"absent", "scan_issue"})

    scores = [student.student_score for student in students]
    stats = _score_distribution(scores, _full_score(rubric, score_map, snapshot))
    session_name = (
        str(session_row.get("session_name") or "").strip() or f"考试批改_{session_id}"
    )
    present = present_count if snapshot.attendance else len(students)
    return SessionAnalysisData(
        session_id=int(session_id),
        session_name=session_name,
        subject=str(rubric.get("subject") or "数学").strip() or "数学",
        full_score=_full_score(rubric, score_map, snapshot),
        graded_at=max((student.graded_at for student in students), default=""),
        present=present,
        roster_absent=roster_absent,
        small_sample=present < SMALL_SAMPLE_LIMIT,
        questions=questions,
        students=students,
        skipped=skipped,
        stats=stats,
        knowledge_backfill=snapshot.knowledge_backfill,
        question_assessments=snapshot.question_assessments,
        attendance_by_class=attendance_by_class,
        rubric=rubric,
    )


def split_session_analysis_by_class(
    data: SessionAnalysisData,
) -> dict[str, SessionAnalysisData]:
    """个人报告、班级页面与班级导出共用的本班统计和同分名次。"""
    names = {student.class_name for student in data.students}
    names.update(item["class_name"] for item in data.skipped)
    names.update(data.attendance_by_class)
    groups: dict[str, SessionAnalysisData] = {}
    for name in sorted(names, key=lambda value: [
        (0, int(part)) if part.isdigit() else (1, part)
        for part in re.split(r"(\d+)", value)
    ]):
        students = sorted(
            (replace(student, material_notes=list(student.material_notes))
             for student in data.students if student.class_name == name),
            key=lambda student: (-student.student_score, student.student_code, student.student_name),
        )
        previous_score = None
        rank = 0
        for index, student in enumerate(students, 1):
            if previous_score is None or student.student_score < previous_score - 1e-9:
                rank, previous_score = index, student.student_score
            student.rank = rank
        questions = []
        for question in data.questions:
            records = [record for student in students for record in student.records
                       if record.question_id == question.question_id]
            questions.append(replace(
                question, attempts=len(records),
                score_sum=sum(min(record.score, record.max_score or record.score) for record in records),
            ))
        skipped = [item for item in data.skipped if item["class_name"] == name]
        counts = data.attendance_by_class.get(name)
        present = counts["present"] if counts else len(students)
        absent = counts["absent"] if counts else sum(item["reason"] in {"缺考", "扫描异常"} for item in skipped)
        groups[name] = replace(
            data, class_name=name, students=students, questions=questions, skipped=skipped,
            present=present, roster_absent=absent, small_sample=present < SMALL_SAMPLE_LIMIT,
            stats=_score_distribution([student.student_score for student in students], data.full_score),
            graded_at=max((student.graded_at for student in students), default=""),
            attendance_by_class={name: counts} if counts else {},
        )
    return groups


def infer_data_root(db_path: Path) -> Path | None:
    path = Path(db_path)
    return path.parent.parent if path.parent.name == "databases" else None


def question_bank_db_path(db_path: Path) -> Path | None:
    path = Path(db_path)
    data_root = path.parent.parent if path.parent.name == "databases" else path.parent
    candidate = data_root / "databases" / "question_bank.db"
    return candidate if candidate.is_file() else None


def load_rubric(session_row: dict[str, Any], data_root: Path | None) -> dict[str, Any]:
    rubric_path = resolve_stored_file_path(
        session_row.get("rubric_path"),
        data_root=data_root,
    )
    try:
        rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
        if not isinstance(rubric, dict):
            return {}
        return canonicalize_question_document(rubric)
    except (
        OSError,
        UnicodeDecodeError,
        json.JSONDecodeError,
        QuestionIdContractError,
    ):
        return {}


def _stem_summary_map(rubric: dict[str, Any]) -> dict[str, str]:
    stems: dict[str, str] = {}
    questions = rubric.get("questions") if isinstance(rubric, dict) else []
    if not isinstance(questions, list):
        return stems
    for question in questions:
        if not isinstance(question, dict):
            continue
        question_id = str(question.get("question_id") or "").strip()
        stem = str(question.get("stem_summary") or "").strip()
        if question_id and stem:
            stems[question_id] = stem
        parts = question.get("parts")
        if not isinstance(parts, list):
            continue
        for part in parts:
            if not isinstance(part, dict):
                continue
            part_id = str(part.get("part_id") or "").strip()
            part_stem = str(part.get("stem_summary") or "").strip()
            # 小问缺题干摘要时回退大问摘要，仍没有则渲染层用题号替代（方案 §9）。
            if part_id and (part_stem or stem):
                stems[part_id] = part_stem or stem
    return stems


# 标准答案展示收敛参数：判分用的等价变体常含标点/单位差异的重复全文，
# 报告展示前做规范化去重、限量、限长。
_ANSWER_FORM_MAX_LEN = 120
_ANSWER_DISPLAY_MAX_FORMS = 3
_ANSWER_DISPLAY_MAX_CHARS = 400

_ANSWER_NORMALIZE_PAIRS = (
    ("：", ":"),
    ("，", ","),
    ("。", "."),
    ("（", "("),
    ("）", ")"),
    ("＝", "="),
    ("；", ";"),
    ("°", "度"),
)


def _normalize_answer_form(text: str) -> str:
    normalized = re.sub(r"\s+", "", str(text))
    for full, half in _ANSWER_NORMALIZE_PAIRS:
        normalized = normalized.replace(full, half)
    return normalized.lower()


def _display_answer_text(forms: list[str]) -> str:
    """把判分用的全部等价变体收敛为家长可读的标准答案文本。"""
    seen: set[str] = set()
    unique: list[str] = []
    for form in forms:
        text = str(form or "").strip()
        key = _normalize_answer_form(text)
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(text)
    if not unique:
        return ""
    # 长文本（证明过程等）只保留代表形式；短答案最多并列 3 个。
    if any(len(text) > _ANSWER_FORM_MAX_LEN for text in unique):
        unique = unique[:1]
    else:
        unique = unique[:_ANSWER_DISPLAY_MAX_FORMS]
    display = " 或 ".join(unique)
    if len(display) > _ANSWER_DISPLAY_MAX_CHARS:
        display = display[:_ANSWER_DISPLAY_MAX_CHARS].rstrip() + " ……"
    return display


def _canonical_answer_map(
    session_row: dict[str, Any],
    data_root: Path | None,
) -> dict[str, str]:
    answer_path = resolve_stored_file_path(
        session_row.get("answer_key_path"),
        data_root=data_root,
    )
    try:
        answer_key = json.loads(answer_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return {
        qid: display
        for qid, forms in answer_forms_map(answer_key).items()
        if (display := _display_answer_text(forms))
    }


def _full_score(
    rubric: dict[str, Any],
    score_map: dict[str, float],
    snapshot: Any,
) -> float:
    try:
        total = float(rubric.get("total_score") or 0)
    except (TypeError, ValueError):
        total = 0.0
    if total > 0:
        return total
    # 只加总大题满分，避免小问与大问重复计数。
    parent_total = 0.0
    for qid, max_score in score_map.items():
        coordinates = question_id_coordinates(qid)
        if coordinates is not None and coordinates[1] is not None:
            continue
        parent_total += float(max_score or 0)
    if parent_total > 0:
        return parent_total
    return max(
        (float(result.get("total_score") or 0) for result in snapshot.results),
        default=0.0,
    )


def _score_distribution(scores: list[float], full_score: float) -> dict[str, Any]:
    if not scores:
        return {
            "avg": None,
            "median": None,
            "max": None,
            "min": None,
            "pass_rate": None,
            "excellent_count": 0,
            "bands": [],
        }
    scale = full_score / 100 if full_score > 0 else 1.0
    bands: list[dict[str, Any]] = []
    for index, cutoff in enumerate(BAND_CUTOFFS):
        lower = cutoff * scale
        upper = full_score if index == 0 else BAND_CUTOFFS[index - 1] * scale
        members = [
            score for score in scores
            if lower <= score and (score <= upper if index == 0 else score < upper)
        ]
        bands.append(
            {
                "label": f"{fmt_num(lower)} – {'不足' if index else ''}{fmt_num(upper)} 分",
                "count": len(members),
                "ratio": len(members) / len(scores),
            }
        )
    passing = [score for score in scores if score >= _PASS_CUTOFF * scale]
    return {
        "avg": round(sum(scores) / len(scores), 2),
        "median": round(statistics.median(scores), 2),
        "max": max(scores),
        "min": min(scores),
        "pass_rate": round(len(passing) / len(scores), 4),
        "excellent_count": len(
            [score for score in scores if score >= BAND_CUTOFFS[0] * scale]
        ),
        "bands": bands,
    }


def _sanitize_grading_text(value: object) -> str:
    """批改记录文本转家长可读口径：内部调试码翻译，无中文的文本丢弃。

    阅卷模型偶发用英文写扣分原因；报告面向家长，没有中文信息的文本
    不如不显示（其余中文字段仍可说明情况）。
    """
    text = str(value or "").strip()
    if not text:
        return ""
    lowered = text.lower()
    if lowered.startswith(_OBJECTIVE_ANSWER_PREFIX):
        answer = text[len(_OBJECTIVE_ANSWER_PREFIX):].strip()
        return f"作答识别为「{answer}」，与参考答案不符" if answer else ""
    direct = _REASON_CODE_LABELS.get(lowered)
    if direct is not None:
        return direct
    for code, label in _REASON_CODE_LABELS.items():
        if code in lowered:
            return label
    return text if has_cjk(text) else ""


def _format_secondary_error(item: object) -> str:
    """次要错因可能是 dict（category/summary/evidence），格式化为可读文本。"""
    if isinstance(item, dict):
        summary = _sanitize_grading_text(item.get("summary"))
        category = _sanitize_grading_text(item.get("category"))
        evidence = str(item.get("evidence") or "").strip()
        text = summary or category
        if evidence and evidence not in text:
            text = f"{text}（{evidence}）" if text else evidence
        return text
    return _sanitize_grading_text(item)


# 学生作答识别文本在 raw_json grading_details 里的兼容键（与 ai_grader 一致）。
_STUDENT_ANSWER_KEYS = ("observed_answer", "student_answer", "answer_observed", "raw_answer", "recognized_answer")
_STUDENT_ANSWER_MAX_LEN = 2000


def _grading_detail_map(raw_json: Any) -> dict[str, dict[str, Any]]:
    """兼容整卷与混合批改保留的作答证据，不根据分数推断空白。"""
    parsed = raw_json
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except json.JSONDecodeError:
            return {}
    if not isinstance(parsed, dict):
        return {}
    result: dict[str, dict[str, Any]] = {}
    metadata = parsed.get("detail_metadata")
    if isinstance(metadata, dict):
        for qid, item in metadata.items():
            if isinstance(item, dict):
                result[str(qid)] = dict(item)
    details = parsed.get("grading_details")
    for item in details if isinstance(details, list) else []:
        if isinstance(item, dict) and (qid := str(item.get("question_id") or "").strip()):
            result.setdefault(qid, {}).update(item)
    return result


def _student_answer_map(raw_json: Any) -> dict[str, str]:
    """从批改结果提取作答文字；图文报告同时使用原卷核对识别结果。"""
    answers: dict[str, str] = {}
    for qid, item in _grading_detail_map(raw_json).items():
        text = ""
        for key in _STUDENT_ANSWER_KEYS:
            value = item.get(key)
            if value:
                text = str(value).strip()
                break
        if text:
            # 兼容旧批改记录把同一段 observed_answer / evidence_steps 拼接两次。
            for evidence in _evidence_texts(item.get("evidence_steps")):
                doubled = f"{evidence} {evidence}"
                if text == doubled:
                    text = evidence
                    break
            answers[qid] = text[:_STUDENT_ANSWER_MAX_LEN]
    return answers


_STEP_UNIT_FIELDS = (
    "step_id",
    "part_id",
    "achievement",
    "score_awarded",
    "reason",
    "missing_or_error",
    "student_evidence",
)


def _independently_failed_steps(assessments: Any) -> list[dict[str, Any]]:
    """独立扣分的步骤：非满分/等价且未标记沿用前步错误；只保留整理需要的字段。"""
    if not isinstance(assessments, list):
        return []
    failed: list[dict[str, Any]] = []
    for item in assessments:
        if not isinstance(item, dict):
            continue
        if str(item.get("achievement") or "").strip().lower() in {"full", "equivalent"}:
            continue
        if item.get("carried_error_from"):
            continue
        if item.get("deduction_source") == "teacher":
            note = str(item.get("teacher_note") or "").strip()
            if not note:
                continue
            item = {**item, "reason": note}
        entry = {
            key: value
            for key in _STEP_UNIT_FIELDS
            if (value := item.get(key)) is not None
            and value != ""
            and (not isinstance(value, (list, dict)) or value)
        }
        if entry:
            failed.append(entry)
    return failed


def _teacher_failed_steps(
    review: Any,
    lock: Any,
    score: float,
    rubric: Any,
    question_id: str,
) -> list[dict[str, Any]]:
    records = _validated_teacher_steps(review, lock, score, rubric, question_id)
    return _independently_failed_steps(records)


def _validated_teacher_steps(
    review: Any, lock: Any, score: float, rubric: Any, question_id: str,
) -> list[dict[str, Any]] | None:
    """教师逐步改分记录仅在对应当前最终分锁时可信；任一校验失败即不使用。"""
    from backend.scan_grading.solution_answer_guard import rubric_scoring_unit_steps

    if not isinstance(review, dict) or not isinstance(lock, dict):
        return None
    if (
        review.get("revision") != lock.get("revision")
        or review.get("scan_batch_id") != lock.get("scan_batch_id")
    ):
        return None
    records = review.get("steps")
    expected = rubric_scoring_unit_steps(
        rubric if isinstance(rubric, dict) else {}, question_id
    )
    if not isinstance(records, list) or not expected or len(records) != len(expected):
        return None
    try:
        locked_score = float(review.get("score_awarded"))
    except (TypeError, ValueError):
        return None
    if locked_score != score:
        return None
    normalized = []
    for step in expected:
        matches = [
            record
            for record in records
            if isinstance(record, dict)
            and record.get("step_id") == step.get("step_id")
            and (not step.get("part_id") or record.get("part_id") == step.get("part_id"))
        ]
        if len(matches) != 1:
            return None
        record = matches[0]
        try:
            awarded = float(record.get("score_awarded"))
            maximum = float(step.get("step_score"))
        except (TypeError, ValueError):
            return None
        if maximum <= 0 or not awarded.is_integer() or awarded < 0 or awarded > maximum:
            return None
        if record.get("max_score") != maximum or set(
            record.get("evidence_point_ids") or []
        ) != set(step.get("evidence_point_ids") or []):
            return None
        normalized.append(
            {**record, "achievement": "full" if awarded == maximum else "none"}
        )
    if sum(float(record["score_awarded"]) for record in normalized) != score:
        return None
    return normalized


def _merge_student_records(
    repositories: GradingRepositoryAccess,
    result_id: int,
    details: list[dict[str, Any]],
    score_map: dict[str, float],
    student_answers: dict[str, str] | None = None,
    grading_evidence: dict[str, dict[str, Any]] | None = None,
    include_secondary_errors: bool = True,
    teacher_reviews: dict[str, Any] | None = None,
    teacher_locks: dict[str, dict[str, Any]] | None = None,
    rubric: dict[str, Any] | None = None,
) -> list[StudentQuestionRecord]:
    """把同一学生同一题的多条明细合并为一条（兼容小问拆行存储）。"""
    secondary_map: dict[str, list[str]] = {}
    try:
        for row in repositories.results.get_result_details(result_id) if include_secondary_errors else []:
            qid = str(row.get("question_id") or "").strip()
            canonical = resolve_known_question_id(qid, score_map) or qid
            errors = row.get("secondary_errors")
            if canonical and isinstance(errors, list):
                merged = secondary_map.setdefault(canonical, [])
                for item in errors:
                    text = _format_secondary_error(item)
                    if text and text not in merged:
                        merged.append(text)
    except Exception:
        secondary_map = {}

    merged: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for detail in details:
        raw_qid = str(detail.get("question_id") or "").strip()
        if not raw_qid:
            continue
        qid = resolve_known_question_id(raw_qid, score_map) or raw_qid
        if qid not in merged:
            merged[qid] = {
                "score": 0.0,
                "deduction_reasons": [],
                "error_categories": [],
                "error_summaries": [],
                "objective_observed": "",
            }
            order.append(qid)
        bucket = merged[qid]
        bucket["score"] += float(detail.get("score_awarded") or 0)
        for field_name, bucket_key in (
            ("deduction_reason", "deduction_reasons"),
            ("error_category", "error_categories"),
            ("error_summary", "error_summaries"),
        ):
            raw_text = str(detail.get(field_name) or "").strip()
            # 客观题的作答识别结果只存在于调试串里，先取出再清洗。
            if (
                field_name == "deduction_reason"
                and raw_text.lower().startswith(_OBJECTIVE_ANSWER_PREFIX)
                and not bucket["objective_observed"]
            ):
                bucket["objective_observed"] = raw_text.split("=", 1)[1].strip()
            text = _sanitize_grading_text(raw_text)
            if text and text not in bucket[bucket_key]:
                bucket[bucket_key].append(text)

    answers = {
        resolve_known_question_id(qid, score_map) or qid: text
        for qid, text in (student_answers or {}).items()
    }
    evidence_by_qid = {
        resolve_known_question_id(qid, score_map) or qid: item
        for qid, item in (grading_evidence or {}).items()
    }
    reviews_by_qid = {
        resolve_known_question_id(str(qid), score_map) or str(qid): item
        for qid, item in (teacher_reviews or {}).items()
        if isinstance(item, dict)
    }
    records: list[StudentQuestionRecord] = []
    for qid in order:
        bucket = merged[qid]
        max_score = float(score_map.get(qid) or 0)
        score = min(bucket["score"], max_score) if max_score > 0 else bucket["score"]
        evidence = evidence_by_qid.get(qid, {})
        lock = (teacher_locks or {}).get(qid)
        teacher_steps = None
        if lock is not None:
            # 教师锁定题不使用 AI 步骤；校验失败的教师步骤按整题处理。
            teacher_steps = _validated_teacher_steps(
                reviews_by_qid.get(qid), lock, bucket["score"], rubric, qid
            )
            failed_steps = _independently_failed_steps(teacher_steps)
        else:
            failed_steps = _independently_failed_steps(evidence.get("step_assessments"))
        records.append(
            StudentQuestionRecord(
                question_id=qid,
                score=score,
                max_score=max_score,
                deduction_reason="；".join(bucket["deduction_reasons"]),
                error_category="；".join(bucket["error_categories"]),
                error_summary="；".join(bucket["error_summaries"]),
                secondary_errors=secondary_map.get(qid, []),
                student_answer=(
                    answers.get(qid) or bucket["objective_observed"] or ""
                ),
                evidence_steps=_evidence_texts(evidence.get("evidence_steps")),
                missing_steps=_evidence_texts(evidence.get("missing_steps")),
                failed_steps=failed_steps,
                teacher_step_reviewed=bool(teacher_steps) and all(
                    step.get("deduction_source") in {"none", "ai", "teacher"}
                    for step in teacher_steps
                ),
            )
        )
    order_index = {
        qid: index for index, qid in enumerate(natural_question_order(order))
    }
    records.sort(
        key=lambda record: order_index.get(record.question_id, len(order_index))
    )
    return records


def _docx_question_figures(source: Any) -> dict[str, list[bytes]]:
    """兼容原题重复出现在解析页的 Word：只取题干之后、解析开始之前的图。"""
    if getattr(source, "suffix", None) != ".docx":
        return {}
    from docx import Document

    from backend.document_parsing.question_blocks import (
        _extract_question_marker_number,
        _looks_like_answer_section_heading,
    )
    document = Document(io.BytesIO(source.private_source_bytes))
    stems = {
        str(item.get("question_id")): re.sub(r"\s+", "", str(item.get("question_text") or item.get("text") or ""))[:18]
        for item in source.private_blocks if isinstance(item, dict)
    }
    images: dict[str, list[bytes]] = {}
    current = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        number = _extract_question_marker_number(text)
        if number is not None:
            qid = f"Q{number}"
            stem = stems.get(qid)
            # 已从前面的题干取得图时，不再重复补入解析页中的题干副本。
            current = qid if qid not in images and stem and stem in re.sub(r"\s+", "", text) else None
        elif _looks_like_answer_section_heading(text) or re.match(
            r"^(?:【|\[)?(?:分析|解答|解析|答案|点评|解：|证明：)", text,
        ):
            current = None
        if current is None:
            continue
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = document.part.related_parts.get(relationship)
            if part is None:
                continue
            blob = part.blob
            try:
                with Image.open(io.BytesIO(blob)) as figure:
                    figure.verify()
            except (OSError, ValueError):
                continue
            bucket = images.setdefault(current, [])
            if blob not in bucket:
                bucket.append(blob)
    return images


def enrich_personal_questions(
    repositories: GradingRepositoryAccess,
    data: SessionAnalysisData,
    data_root: Path | None,
    *,
    include_images: bool = True,
) -> None:
    """读取本场评分依据和已绑定原题；考试错因整理仅复用文字投影。"""
    from backend.config_workspace.sources import ConfigSourceService
    from backend.document_parsing.question_blocks import rich_text_for_model

    session = repositories.sessions.get_grading_session(data.session_id) or {}
    root = data_root or infer_data_root(repositories.db_path)
    rubric = load_rubric(session, root)
    try:
        answer_key = json.loads(resolve_stored_file_path(
            session.get("answer_key_path"), data_root=root,
        ).read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        answer_key = {}

    def question_map(document: Any) -> dict[str, dict[str, Any]]:
        result: dict[str, dict[str, Any]] = {}
        for item in document.get("questions", []) if isinstance(document, dict) else []:
            if not isinstance(item, dict):
                continue
            qid = str(item.get("question_id") or "")
            result[qid] = item
            for part in item.get("parts", []) if isinstance(item.get("parts"), list) else []:
                if isinstance(part, dict):
                    raw_id = str(part.get("part_id") or part.get("question_id") or "")
                    canonical = resolve_known_question_id(raw_id, known_ids) or raw_id
                    result[canonical] = {**item, **part} if canonical == qid else part
        return result

    known_ids = {info.question_id: info.max_score for info in data.questions}
    rubric_items = question_map(rubric)
    answer_items = question_map(answer_key)
    source = None
    # Active uploads may be uncommitted candidates.  Only use the source whose
    # existing publication identity matches this exam's bound original paper.
    bound_source = str(session.get("source_paper_sha256") or "").strip()
    if bound_source and root is not None:
        try:
            candidate = ConfigSourceService(root / "config" / "uploaded").load_active_record(
                session_id=data.session_id,
            )
            if candidate.sha256 == bound_source:
                source = candidate
        except (OSError, ValueError, RuntimeError):
            pass
    source_items = {
        str(item.get("question_id") or ""): item
        for item in (source.private_blocks if source is not None else [])
        if isinstance(item, dict)
    }
    source_figures = _docx_question_figures(source) if source is not None and include_images else {}
    # Use the native formula alongside the exact importer's HTML projection.
    # This preserves root indices and mixed fractions that plain text loses.
    source_formulas = _source_formula_markup(source) if source is not None else {}
    formula_pattern = re.compile('|'.join(re.escape(key) for key in sorted(source_formulas, key=len, reverse=True))) if source_formulas else None

    def texts(item: dict[str, Any], fields: tuple[str, ...]) -> str:
        return "\n".join(dict.fromkeys(
            text for key in fields
            if (text := rich_text_for_model(item.get(key) or ""))
        ))

    for info in data.questions:
        parent = parent_question_id(info.question_id)
        original = source_items.get(parent, {})
        rubric_item = rubric_items.get(info.question_id, rubric_items.get(parent, {}))
        answer_item = answer_items.get(info.question_id, answer_items.get(parent, {}))
        # These fields are alternative representations of the same stem.
        stem = next((value for key in ("question_html", "question_text", "text")
                     if (value := texts(original, (key,)))), "")
        if not stem:
            stem = texts(rubric_items.get(parent, {}), ("question_text", "text", "stem", "question_html"))
        options = rubric_items.get(parent, {}).get("options")
        if options:
            option_text = json.dumps(options, ensure_ascii=False) if isinstance(options, (dict, list)) else str(options)
            stem = "\n".join(filter(None, (stem, rich_text_for_model(option_text))))
        part_text = texts(rubric_item, ("question_text", "text", "stem_summary"))
        if parent != info.question_id and part_text and part_text not in stem:
            stem = "\n".join(filter(None, (stem, f"本小问：{part_text}")))
        info.question_text = stem
        info.question_markup = str(original.get("question_html") or "")
        if formula_pattern and info.question_markup and 'data-latex=' not in info.question_markup:
            info.question_markup = formula_pattern.sub(lambda match: source_formulas[match.group(0)], info.question_markup)
        info.reference_analysis = texts(answer_item, ("analysis", "full_answer", "explanation"))
        if not info.reference_analysis:
            info.reference_analysis = texts(original, ("analysis_html", "analysis", "answer_html", "answer_text"))
        if source is not None and include_images:
            assets = source.private_question_images.get(parent, {})
            for role in ("question", "answer"):
                encoded = assets.get(role)
                values = [encoded] if isinstance(encoded, str) else encoded if isinstance(encoded, list) else []
                for value in values:
                    try:
                        blob = base64.b64decode(value, validate=True)
                        if blob:
                            info.reference_images.append((role, blob))
                    except (ValueError, TypeError):
                        continue
            if not any(role == "question" for role, _blob in info.reference_images):
                info.reference_images.extend(("question", blob) for blob in source_figures.get(parent, []))


def _source_formula_markup(source: Any) -> dict[str, str]:
    if getattr(source, 'suffix', None) != '.docx':
        return {}
    from docx import Document

    from question_bank.importers.docx_importer import _math_text
    from question_bank.services.inline_math import omml_parts
    try:
        document = Document(io.BytesIO(source.private_source_bytes))
        variants: dict[str, set[tuple[str, str]]] = {}
        for node in document.element.xpath('.//m:oMath'):
            formula = omml_parts(node)
            original = _math_text(node)
            if formula and original and (len(original) > 2 or any(c in original for c in '√∛')):
                variants.setdefault(original, set()).add(formula)
        return {original: report_math_span(*next(iter(values)))
                for original, values in variants.items() if len(values) == 1}
    except (OSError, ValueError, zipfile.BadZipFile):
        return {}


def enrich_personal_knowledge(
    repositories: GradingRepositoryAccess, data: SessionAnalysisData,
    data_root: Path | None, *, student_ids: set[int] | None = None,
) -> None:
    """Freeze the existing semester diagnosis once for the whole export batch."""
    from integration.diagnosis_profile_service import DiagnosisProfileService
    from question_bank.current_knowledge import (
        CurrentKnowledgeResolver,
        CurrentKnowledgeUnavailable,
    )
    from question_bank.solution_evidence.part_assessments import reading
    root = data_root or infer_data_root(repositories.db_path)
    path = root / 'databases' / 'question_bank.db'
    snapshot = {'as_of': datetime.now().strftime('%Y-%m-%d %H:%M'), 'catalog': [], 'associations': [],
                'note': '当前掌握度暂不可用；保留本卷考查范围与得分，不以得分率代替掌握度。'}
    data.knowledge_structure = snapshot
    selected = [student for student in data.students if student_ids is None or student.student_id in student_ids]
    for student in selected:
        student.knowledge_mastery = {}
    if not path.is_file() or not selected:
        return
    session = repositories.sessions.get_grading_session(data.session_id) or {}
    volume = str(session.get('curriculum_volume_id') or '').strip()
    if not volume:
        snapshot['note'] = '这场考试尚未关联教学学期，暂不展示当前掌握度；下方保留本卷考查范围与得分。'
        return
    try:
        resolver = CurrentKnowledgeResolver.from_active_database(path)
        # Exact current identities only; never join two nodes by a short label.
        for entries in data.knowledge_backfill.values():
            for entry in entries:
                identities = resolver.resolve(entry.get('stable_key') or entry.get('path') or '')
                if len(identities) == 1:
                    entry['stable_key'] = identities[0].stable_key
        with reading(path) as connection:
            service = DiagnosisProfileService(repositories.db_path, path, grading_db=repositories,
                                              question_bank_connection=connection, data_root=root)
            profile = service.build_tag_profiles(
                scope={'mode': 'selected', 'student_ids': [str(student.student_id) for student in selected], 'use_historical_fallback': False},
                exam_scope={'mode': 'semester', 'curriculum_volume_id': volume},
            )
        snapshot.update(catalog=profile.get('knowledge_catalog') or [],
                        associations=profile.get('knowledge_associations') or [])
        profiles = {str(item['student_id']): item for item in profile.get('students', [])}
        for student in selected:
            student.knowledge_mastery = {
                str(item['knowledge_key']): item
                for item in profiles.get(str(student.student_id), {}).get('weak_points', [])
                # A failed mastery calculation can leave exam-only diagnostic
                # rates. Those are not current mastery and must not be relabeled.
                if 'effective_weight' in item and 'direct_evidence_count' in item
            }
        if snapshot['catalog']:
            snapshot['note'] = '仅展开本卷涉及的知识与技能。当前掌握度综合本学期考试、训练、题目难度与时间；本次考试得分单独列出。'
    except (CurrentKnowledgeUnavailable, OSError, sqlite3.Error, ValueError, TypeError):
        # A missing graph must never prevent delivery of the scored report.
        return


def _evidence_texts(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if isinstance(item, str) and item.strip()]


def parent_question_id(question_id: str) -> str:
    coordinates = question_id_coordinates(question_id)
    if coordinates is None:
        return str(question_id or "").strip()
    return f"Q{coordinates[0]}"


def fmt_num(value: object) -> str:
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return str(value or "")
    if number.is_integer():
        return str(int(number))
    return f"{number:g}"


def esc(value: object) -> str:
    return html.escape(str(value or ""), quote=True)


def report_math_span(tex: str, fallback: str, *, display: bool = False) -> str:
    return (f'<span class="qm{" qm-display" if display else ""}" data-latex="{esc(tex)}"'
            f'{" data-display=\"true\"" if display else ""}>{esc(fallback)}</span>')
