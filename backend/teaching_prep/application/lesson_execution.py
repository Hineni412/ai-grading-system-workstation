"""本机执行已确认的改编计划并做机械审计。

流程：从资源包解析主课件源文件 → 用题库现取插入题内容 → python-pptx
在隔离副本上执行操作单 → 机械审计（上下标、页数红线、页码核对）→
把产物与报告写入 pptx_local_outputs。全程本机完成，不调用模型。
"""

from __future__ import annotations

import html
import re
import uuid
from pathlib import Path
from typing import Any

from backend.teaching_prep.application.lesson_audit import run_full_audit
from backend.teaching_prep.application.pptx_editor import (
    PptxEditorError,
    apply_plan,
)
from backend.teaching_prep.application.question_selection import (
    fetch_question_content,
)
from backend.teaching_prep.domain.errors import (
    TeachingPrepConflictError,
    TeachingPrepNotFoundError,
    TeachingPrepValidationError,
)


def _reference_ppt_source(
    database: Any,
    pack_payload: dict[str, Any],
) -> tuple[str, str, int]:
    """返回主课件 (源文件路径, 材料版本 id, 源页数)。"""

    version_id = ""
    for material in pack_payload.get("materials") or []:
        if not isinstance(material, dict):
            continue
        if str(material.get("purpose") or "") != "reference_ppt":
            continue
        candidate = str(material.get("material_version_id") or "").strip()
        if candidate:
            version_id = candidate
            break
    if not version_id:
        raise TeachingPrepValidationError(
            "resource pack does not contain a bound reference PPT"
        )
    with database.connect() as connection:
        row = connection.execute(
            """
            SELECT location.local_path, location.availability,
                   version.unit_count
            FROM material_locations AS location
            JOIN material_versions AS version ON version.id = location.version_id
            WHERE location.version_id = ?
            """,
            (version_id,),
        ).fetchone()
    if row is None:
        raise TeachingPrepConflictError(
            "reference PPT file location is missing"
        )
    if str(row["availability"]) != "available":
        raise TeachingPrepConflictError(
            "reference PPT file is not currently available"
        )
    local_path = Path(str(row["local_path"]))
    if not local_path.is_file():
        raise TeachingPrepConflictError(
            "reference PPT file cannot be found on disk"
        )
    return str(local_path), version_id, int(row["unit_count"] or 0)


def _enrich_question_operations(
    operations: list[dict[str, Any]],
    question_bank_path: Path | None,
) -> tuple[list[dict[str, Any]], list[int]]:
    """把计划中的插题操作补上题干、答案与配图路径。"""

    if not question_bank_path:
        raise TeachingPrepValidationError(
            "question bank is unavailable for question page execution"
        )
    question_ids: list[int] = []
    for operation in operations:
        if str(operation.get("kind")) != "add_question_slide":
            continue
        details = operation.get("details")
        question = (
            details.get("question")
            if isinstance(details, dict)
            else None
        )
        if isinstance(question, dict) and question.get("question_id"):
            question_ids.append(int(question["question_id"]))
    content = fetch_question_content(question_bank_path, question_ids)
    enriched: list[dict[str, Any]] = []
    missing: list[int] = []
    for operation in operations:
        if str(operation.get("kind")) != "add_question_slide":
            enriched.append(operation)
            continue
        details = dict(operation.get("details") or {})
        question = dict(details.get("question") or {})
        question_id = int(question.get("question_id") or 0)
        body = content.get(question_id)
        if body is None:
            missing.append(question_id)
            enriched.append(operation)
            continue
        question.update(
            {
                "stem_html": body["stem_html"],
                "answer_html": body["answer_html"],
                "question_type": body["question_type"],
                "image_paths": body["image_paths"],
            }
        )
        details["question"] = question
        enriched.append({**operation, "details": details})
    if missing:
        raise TeachingPrepConflictError(
            "inserted questions are missing from the question bank: "
            + ",".join(str(q) for q in missing)
        )
    return enriched, question_ids


def execute_confirmed_plan(
    *,
    database: Any,
    outputs_root: Path,
    outputs_repo: Any,
    plan_payload: dict[str, Any],
    plan_id: str,
    draft_id: str,
    pack_id: str,
    pack_payload: dict[str, Any],
    lesson_node_id: str,
    question_bank_path: Path | None,
) -> dict[str, Any]:
    """执行一张已确认计划，返回产物行（含审计报告）。"""

    operations = plan_payload.get("operations")
    if not isinstance(operations, list) or not operations:
        raise TeachingPrepValidationError("slide plan has no operations")

    source_path, version_id, source_page_count = _reference_ppt_source(
        database,
        pack_payload,
    )
    enriched, _question_ids = _enrich_question_operations(
        operations,
        question_bank_path,
    )

    output_id = uuid.uuid4().hex
    output_relpath = f"outputs/{output_id}.pptx"
    output_path = outputs_root / f"{output_id}.pptx"
    classroom = pack_payload.get("classroom")
    lesson_kind = (
        str(classroom.get("lesson_type") or "review")
        if isinstance(classroom, dict)
        else "review"
    )

    try:
        report = apply_plan(
            source_path=source_path,
            output_path=output_path,
            operations=enriched,
        )
    except PptxEditorError as exc:
        raise TeachingPrepValidationError(str(exc)) from exc

    audit = run_full_audit(
        output_path,
        inserted_question_pages=report.get("inserted_question_pages", []),
        lesson_kind=lesson_kind,
    )

    import json

    source_name = Path(source_path).name
    row = outputs_repo.create(
        id=output_id,
        lesson_node_id=lesson_node_id,
        slide_plan_id=plan_id,
        draft_id=draft_id,
        resource_pack_id=pack_id,
        output_relpath=output_relpath,
        output_filename=output_path.name,
        output_sha256=report["output_sha256"],
        source_material_version_id=version_id,
        source_file_name=source_name,
        source_page_count=report["source_page_count"],
        final_page_count=report["final_page_count"],
        lesson_kind=lesson_kind,
        execution_report_json=json.dumps(
            report, ensure_ascii=False, sort_keys=True
        ),
        audit_report_json=json.dumps(
            audit, ensure_ascii=False, sort_keys=True
        ),
        inserted_question_pages_json=json.dumps(
            report.get("inserted_question_pages", []),
            ensure_ascii=False,
            sort_keys=True,
        ),
    )
    return row


def _docx_runs_from_html(paragraph: Any, raw: str) -> None:
    """把轻量 HTML 题干写入 docx 段落（支持 sup/sub/u）。"""

    pattern = re.compile(r"<(/?)(sup|sub|u|br)\s*/?>", re.IGNORECASE)
    baseline = None
    underline = False
    position = 0

    def _add(text: str) -> None:
        text = re.sub(r"<[^>]+>", "", text)
        if not text:
            return
        run = paragraph.add_run(text)
        if baseline is not None:
            run.font.superscript = baseline > 0
            run.font.subscript = baseline < 0
        if underline:
            run.font.underline = True

    for match in pattern.finditer(raw):
        _add(html.unescape(raw[position : match.start()]))
        closing, tag = match.group(1), match.group(2).lower()
        if tag == "br":
            _add("\n")
        elif tag == "sup":
            baseline = None if closing else 1
        elif tag == "sub":
            baseline = None if closing else -1
        elif tag == "u":
            underline = not bool(closing)
        position = match.end()
    _add(html.unescape(raw[position:]))


def generate_worksheet(
    *,
    database: Any,
    outputs_root: Path,
    outputs_repo: Any,
    output_id: str,
    question_bank_path: Path,
) -> dict[str, Any]:
    """为已生成的成片生成配套学案 DOCX（紧凑题单，无解答留白）。"""

    from docx import Document
    from docx.shared import Inches as _Inches, Pt as _Pt

    row = outputs_repo.get(output_id)
    if str(row.get("worksheet_relpath") or ""):
        return row
    inserted = outputs_repo.inserted_question_pages(row)
    if not inserted:
        raise TeachingPrepValidationError(
            "worksheet requires at least one inserted question page"
        )
    ordered_ids = [
        int(item["question_id"])
        for item in sorted(inserted, key=lambda x: int(x["final_position"]))
    ]
    content = fetch_question_content(question_bank_path, ordered_ids)
    final_page_count = int(row["final_page_count"])

    document = Document()
    style = document.styles["Normal"]
    style.font.size = _Pt(10.5)
    section = document.sections[0]
    section.left_margin = _Inches(0.6)
    section.right_margin = _Inches(0.6)
    section.top_margin = _Inches(0.6)
    section.bottom_margin = _Inches(0.6)

    header = document.add_paragraph()
    header_run = header.add_run("课堂学案（配套改编课件）")
    header_run.bold = True
    header_run.font.size = _Pt(14)

    for question_id in ordered_ids:
        body = content.get(int(question_id))
        if body is None:
            continue
        page_number = next(
            (
                int(item["final_position"])
                for item in inserted
                if int(item["question_id"]) == int(question_id)
            ),
            None,
        )
        title = document.add_paragraph()
        title_run = title.add_run(
            f"第 {page_number} 页对应题（{body['question_type']}）"
        )
        title_run.bold = True
        title_run.font.size = _Pt(11)

        stem = document.add_paragraph()
        _docx_runs_from_html(stem, body["stem_html"])

        for image_path in body["image_paths"][:3]:
            if not Path(str(image_path)).is_file():
                continue
            paragraph = document.add_paragraph()
            try:
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=_Inches(4.5))
            except Exception:
                paragraph.add_run("（配图缺失）")

    if final_page_count:
        footer = document.add_paragraph()
        footer_run = footer.add_run(
            f"共 {len(ordered_ids)} 题；页码对应改编后课件第 1-{final_page_count} 页。"
        )
        footer_run.font.size = _Pt(9)

    output_docx_id = uuid.uuid4().hex
    worksheet_path = outputs_root / f"{output_docx_id}.docx"
    document.save(str(worksheet_path))

    updated = outputs_repo.attach_worksheet(
        output_id,
        worksheet_relpath=f"outputs/{worksheet_path.name}",
        worksheet_filename=f"worksheet-{output_id}.docx",
    )
    _ = database
    return updated


def worksheet_file(
    outputs_root: Path,
    row: dict[str, Any],
) -> tuple[Path, str]:
    relpath = str(row.get("worksheet_relpath") or "")
    if not relpath:
        raise TeachingPrepNotFoundError("worksheet was not generated")
    base = outputs_root.parent.resolve()
    candidate = (base / relpath).resolve()
    if not str(candidate).startswith(str(base)):
        raise TeachingPrepValidationError("worksheet path is invalid")
    if not candidate.is_file():
        raise TeachingPrepNotFoundError("worksheet file is missing")
    return candidate, str(row.get("worksheet_filename") or "worksheet.docx")
