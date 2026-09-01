"""改编副本的本机机械审计。

包含三类检查，全部在本地完成、不调用模型：
1. 上下标审计：发现整句/整行被误标为上标或下标的文本片段；
2. 课时容量红线：对照 40 分钟课时的页数上限给出提示；
3. 页码核对：题页插入位置与最终页序的双向核对数据。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.oxml.ns import qn

_REVIEW_PAGE_LIMIT = 18
_BASELINE_RUN_TEXT_LIMIT = 8
_BASELINE_CJK_LIMIT = 4
_CJK_RANGES = (
    (0x4E00, 0x9FFF),
    (0x3400, 0x4DBF),
)


def _count_cjk(text: str) -> int:
    total = 0
    for char in text:
        code = ord(char)
        if any(start <= code <= end for start, end in _CJK_RANGES):
            total += 1
    return total


def audit_superscript_subscript(pptx_path: str | Path) -> dict[str, Any]:
    """扫描全部文本 run，报告可疑的整句上下标残留。

    规则：带 baseline 属性的 run，若其文本较长（>=8 个字符）或包含
    较多中文（>=4 个汉字），判定为可疑残留——正常数学上下标只会覆盖
    少数字符（如 x 的平方、脚标序号）。
    """

    path = Path(pptx_path)
    presentation = Presentation(str(path))
    findings: list[dict[str, Any]] = []
    scanned_runs = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    scanned_runs += 1
                    r_pr = run._r.find(qn("a:rPr"))
                    if r_pr is None:
                        continue
                    baseline_raw = r_pr.get("baseline")
                    if not baseline_raw:
                        continue
                    try:
                        baseline = int(baseline_raw)
                    except (TypeError, ValueError):
                        continue
                    if baseline == 0:
                        continue
                    text = run.text or ""
                    suspicious = (
                        len(text.strip()) >= _BASELINE_RUN_TEXT_LIMIT
                        or _count_cjk(text) >= _BASELINE_CJK_LIMIT
                    )
                    if suspicious:
                        findings.append(
                            {
                                "page": slide_index,
                                "baseline": baseline,
                                "kind": "superscript" if baseline > 0 else "subscript",
                                "text": text.strip()[:60],
                            }
                        )
    return {
        "scanned_runs": scanned_runs,
        "finding_count": len(findings),
        "findings": findings[:50],
        "passed": not findings,
    }


def audit_page_budget(
    final_page_count: int,
    *,
    lesson_kind: str = "review",
) -> dict[str, Any]:
    """页数红线：40 分钟课时约 16–18 页，超出时给出提示。"""

    limit = _REVIEW_PAGE_LIMIT if lesson_kind != "new" else 14
    return {
        "final_page_count": final_page_count,
        "limit": limit,
        "over_limit": final_page_count > limit,
        "suggestion": (
            f"当前 {final_page_count} 页超过 {limit} 页上限，建议在对照页优先"
            "删除重复类型题页或标注可跳过页。"
            if final_page_count > limit
            else "页数在课时容量红线内。"
        ),
    }


def cross_check_question_pages(
    inserted_question_pages: list[dict[str, Any]],
    *,
    final_page_count: int,
) -> dict[str, Any]:
    """核对每道插入题的最终页码是否落在成片范围内且互不冲突。"""

    problems: list[str] = []
    seen_pages: dict[int, int] = {}
    for entry in inserted_question_pages:
        position = entry.get("final_position")
        question_id = entry.get("question_id")
        if not isinstance(position, int) or not 1 <= position <= final_page_count:
            problems.append(
                f"题 {question_id} 的页码 {position} 超出成片范围（1-{final_page_count}）。"
            )
            continue
        if position in seen_pages:
            problems.append(
                f"题 {seen_pages[position]} 与题 {question_id} 落在同一页 "
                f"{position}，插入位置冲突。"
            )
        else:
            seen_pages[position] = int(question_id or 0)
    return {
        "checked": len(inserted_question_pages),
        "pages": dict(sorted(seen_pages.items())),
        "problem_count": len(problems),
        "problems": problems,
        "passed": not problems,
    }


def run_full_audit(
    pptx_path: str | Path,
    *,
    inserted_question_pages: list[dict[str, Any]] | None = None,
    lesson_kind: str = "review",
) -> dict[str, Any]:
    """对成片副本执行全部机械审计，返回统一报告。"""

    superscript = audit_superscript_subscript(pptx_path)
    presentation = Presentation(str(pptx_path))
    final_page_count = len(presentation.slides._sldIdLst)  # noqa: SLF001
    page_budget = audit_page_budget(final_page_count, lesson_kind=lesson_kind)
    question_pages = cross_check_question_pages(
        inserted_question_pages or [],
        final_page_count=final_page_count,
    )
    return {
        "superscript_subscript": superscript,
        "page_budget": page_budget,
        "question_pages": question_pages,
        "passed": (
            superscript["passed"]
            and question_pages["passed"]
        ),
    }
