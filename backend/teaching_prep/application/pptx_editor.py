"""用 python-pptx 在隔离副本上执行课件改编计划。

输入是改编计划的操作单（与 slide_plans 契约一致），输出是新的 PPTX 文件
与逐操作执行报告。源文件只读，不调用模型，不访问网络。

执行分两个阶段：先在抽象页序上按计划顺序完成全部操作（删除、调序、隐藏、
插入），再一次性物化到 PPTX（先加页、后删页、再整体重排），避免部件重名
与索引漂移。
"""

from __future__ import annotations

import hashlib
import html
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

_MAX_SOURCE_BYTES = 256 * 1024 * 1024
_BASELINE_SUPERSCRIPT = 30000
_BASELINE_SUBSCRIPT = -25000
_ANSWER_GRAY = RGBColor(0x59, 0x59, 0x59)

_TAG_PATTERN = re.compile(r"<(/?)(sup|sub|u|br)\s*/?>", re.IGNORECASE)
_ANY_TAG_PATTERN = re.compile(r"<[^>]+>")
_ENTITY_PATTERN = re.compile(r"&[a-z]+;|\&#\d+;")


class PptxEditorError(ValueError):
    """计划无法在副本上安全执行。"""


@dataclass
class _RunSpec:
    text: str
    baseline: int | None = None
    underline: bool = False


@dataclass
class _ParagraphSpec:
    runs: list[_RunSpec] = field(default_factory=list)


def _decode_entities(text: str) -> str:
    return _ENTITY_PATTERN.sub(lambda m: html.unescape(m.group(0)), text)


def _parse_rich_text(raw: str) -> list[_ParagraphSpec]:
    """把题干中的轻量 HTML（sup/sub/u/br）解析为带上下标信息的段落。"""

    paragraphs: list[_ParagraphSpec] = []
    current = _ParagraphSpec()
    baseline: int | None = None
    underline = False
    position = 0
    for match in _TAG_PATTERN.finditer(raw):
        text = _decode_entities(raw[position : match.start()])
        if text:
            current.runs.append(
                _RunSpec(text=text, baseline=baseline, underline=underline)
            )
        closing, tag = match.group(1), match.group(2).lower()
        if tag == "br":
            paragraphs.append(current)
            current = _ParagraphSpec()
        elif tag == "sup":
            baseline = None if closing else _BASELINE_SUPERSCRIPT
        elif tag == "sub":
            baseline = None if closing else _BASELINE_SUBSCRIPT
        elif tag == "u":
            underline = not bool(closing)
        position = match.end()
    tail = _decode_entities(raw[position:])
    if tail:
        current.runs.append(
            _RunSpec(text=tail, baseline=baseline, underline=underline)
        )
    if current.runs:
        paragraphs.append(current)
    cleaned: list[_ParagraphSpec] = []
    for paragraph in paragraphs:
        runs = [run for run in paragraph.runs if run.text]
        if runs:
            cleaned.append(_ParagraphSpec(runs=runs))
    return cleaned or [_ParagraphSpec(runs=[_RunSpec(text="")])]


def _strip_markup(raw: object) -> str:
    text = str(raw or "")
    text = _TAG_PATTERN.sub("", text)
    text = _ANY_TAG_PATTERN.sub("", text)
    return _decode_entities(text).strip()


def _load_presentation(source: Path) -> Presentation:
    if not source.is_file():
        raise PptxEditorError("源课件文件不存在。")
    if source.stat().st_size > _MAX_SOURCE_BYTES:
        raise PptxEditorError("源课件文件超过大小限制。")
    return Presentation(str(source))


def _sld_id_list(presentation: Presentation) -> Any:
    return presentation.slides._sldIdLst


def _sld_id_entries(presentation: Presentation) -> list[Any]:
    return list(_sld_id_list(presentation))


def _drop_slide_element(presentation: Presentation, entry: Any) -> None:
    presentation.part.drop_rel(entry.rId)
    _sld_id_list(presentation).remove(entry)


def _blank_layout(presentation: Presentation, fallback_layout: Any) -> Any:
    for layout in presentation.slide_layouts:
        if not list(layout.placeholders):
            return layout
    return fallback_layout


def _set_run(run: Any, spec: _RunSpec) -> None:
    run.text = spec.text
    if spec.baseline is not None:
        run.font._rPr.set("baseline", str(spec.baseline))
    if spec.underline:
        run.font.underline = True


def _fill_text_frame(
    frame: Any, paragraphs: list[_ParagraphSpec], *, size_pt: float
) -> None:
    frame.word_wrap = True
    first = True
    for paragraph_spec in paragraphs:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        if not paragraph_spec.runs:
            paragraph_spec.runs = [_RunSpec(text="")]
        for run_spec in paragraph_spec.runs:
            run = paragraph.add_run()
            _set_run(run, run_spec)
            run.font.size = Pt(size_pt)


def _add_question_slide(
    presentation: Presentation,
    *,
    layout: Any,
    title: str,
    stem_paragraphs: list[_ParagraphSpec],
    image_paths: list[str],
    answer_paragraphs: list[_ParagraphSpec] | None,
) -> Any:
    slide = presentation.slides.add_slide(layout)
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    margin = Inches(0.45)
    content_width = slide_width - margin * 2

    title_box = slide.shapes.add_textbox(margin, margin, content_width, Inches(0.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    run = title_frame.paragraphs[0].add_run()
    _set_run(run, _RunSpec(text=title))
    run.font.size = Pt(24)
    run.font.bold = True

    cursor_y = margin + Inches(0.7)
    answer_height = Inches(0.85) if answer_paragraphs else Inches(0)
    body_height = slide_height - cursor_y - margin - answer_height

    stem_height = body_height if not image_paths else int(body_height * 0.45)
    stem_box = slide.shapes.add_textbox(margin, cursor_y, content_width, stem_height)
    _fill_text_frame(stem_box.text_frame, stem_paragraphs, size_pt=16)

    if image_paths:
        image_y = cursor_y + stem_height + Inches(0.2)
        image_area = slide_height - image_y - margin - answer_height
        per_image = max(Inches(0.6), int(image_area / max(1, len(image_paths))))
        for image_path in image_paths[:4]:
            try:
                picture = slide.shapes.add_picture(
                    str(image_path), margin, image_y, width=content_width
                )
            except Exception:
                note_box = slide.shapes.add_textbox(
                    margin, image_y, content_width, Inches(0.4)
                )
                note_box.text_frame.text = "（本页配图缺失）"
                continue
            if picture.height > per_image:
                scale = per_image / picture.height
                picture.height = per_image
                picture.width = int(picture.width * scale)
            image_y = image_y + picture.height + Inches(0.1)

    if answer_paragraphs:
        answer_box = slide.shapes.add_textbox(
            margin,
            slide_height - margin - answer_height,
            content_width,
            answer_height,
        )
        _fill_text_frame(answer_box.text_frame, answer_paragraphs, size_pt=12)
        for paragraph in answer_box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = _ANSWER_GRAY
    return slide


def _sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _operation_page(operation: dict[str, Any]) -> int:
    target = operation.get("target")
    if not isinstance(target, dict):
        return 0
    try:
        return int(target.get("generated_page_number") or 0)
    except (TypeError, ValueError):
        return 0


def apply_plan(
    *,
    source_path: str | Path,
    output_path: str | Path,
    operations: list[dict[str, Any]],
) -> dict[str, Any]:
    """在隔离副本上执行操作单，返回执行报告。

    页码均指源文件的 1 起页码。delete/reorder/hide 作用于原页；add_slide /
    add_question_slide 在指定页后插入新页；modify_text_box 默认转人工。
    """

    source = Path(source_path)
    output = Path(output_path)
    if not source.is_file():
        raise PptxEditorError("源课件文件不存在。")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, output)

    presentation = _load_presentation(output)
    source_page_count = len(_sld_id_entries(presentation))
    if source_page_count == 0:
        raise PptxEditorError("源课件没有任何页面。")

    report: list[dict[str, Any]] = []

    def _base_entry(operation: dict[str, Any]) -> dict[str, Any]:
        return {
            "kind": str(operation.get("kind") or ""),
            "source_page": _operation_page(operation),
            "target_key": str(operation.get("target_key") or ""),
            "status": "skipped",
            "detail": "",
        }

    # 第一阶段：在抽象页序上按计划顺序执行。
    # 页序令牌：("page", 源页码) 或 ("new", 插入序号)。
    abstract_order: list[tuple[str, int]] = [
        ("page", page) for page in range(1, source_page_count + 1)
    ]
    new_slide_payloads: dict[int, dict[str, Any]] = {}
    new_slide_kind: dict[int, str] = {}
    new_slide_counter = 0
    insert_chain: dict[int, int] = {}
    deleted_pages: set[int] = set()

    for operation in operations:
        kind = str(operation.get("kind") or "")
        entry = _base_entry(operation)
        details = operation.get("details")
        details = details if isinstance(details, dict) else {}
        page = entry["source_page"]

        def _page_token_index(page_number: int) -> int | None:
            for index, token in enumerate(abstract_order):
                if token == ("page", page_number):
                    return index
            return None

        if kind == "delete_slide":
            if page in deleted_pages:
                entry.update(status="skipped", detail="页面已删除。")
            elif _page_token_index(page) is not None:
                abstract_order.remove(("page", page))
                deleted_pages.add(page)
                entry.update(status="applied", detail="已删除该页。")
            else:
                entry.update(status="failed", detail="页码超出范围。")
        elif kind == "reorder_slide":
            try:
                target_position = int(details.get("target_position") or 0)
            except (TypeError, ValueError):
                target_position = 0
            index = _page_token_index(page)
            if index is None or target_position <= 0:
                entry.update(status="failed", detail="页码或目标位置无效。")
            else:
                abstract_order.pop(index)
                insert_at = min(target_position - 1, len(abstract_order))
                abstract_order.insert(insert_at, ("page", page))
                entry.update(
                    status="applied", detail=f"移动到第 {insert_at + 1} 页。"
                )
        elif kind == "hide_slide":
            if _page_token_index(page) is None:
                entry.update(status="failed", detail="页码超出范围。")
            else:
                entry.update(status="applied", detail="已隐藏该页。")
        elif kind in ("add_slide", "add_question_slide"):
            try:
                anchor = int(details.get("anchor_page") or page or 0)
            except (TypeError, ValueError):
                anchor = 0
            new_slide_counter += 1
            token = ("new", new_slide_counter)
            if kind == "add_question_slide":
                question = details.get("question")
                question = question if isinstance(question, dict) else {}
                new_slide_payloads[new_slide_counter] = question
            else:
                new_slide_payloads[new_slide_counter] = {
                    "text": _strip_markup(operation.get("reason")) or "新增内容"
                }
            new_slide_kind[new_slide_counter] = kind
            insert_index: int | None = None
            if anchor:
                for index, item in enumerate(abstract_order):
                    if item == ("page", anchor):
                        insert_index = index + 1
                        # 同一锚点连续插入时，按计划顺序排在上一题之后。
                        while insert_index < len(abstract_order) and (
                            abstract_order[insert_index][0] == "new"
                            and insert_chain.get(anchor)
                            == abstract_order[insert_index][1]
                        ):
                            insert_index += 1
                        break
            if insert_index is None:
                abstract_order.append(token)
            else:
                abstract_order.insert(insert_index, token)
            insert_chain[anchor] = new_slide_counter
            entry.update(
                status="applied",
                detail="已插入新页。" if kind == "add_slide" else "已插入题页。",
                new_token=new_slide_counter,
            )
        elif kind == "modify_text_box":
            entry.update(status="manual", detail="文本替换需教师人工完成。")
        else:
            entry.update(status="skipped", detail="未支持的操作类型。")
        report.append(entry)

    # 第二阶段：物化。先加页（页号顺延，避免部件重名），再删页，再重排。
    slide_layout_for_page: dict[int, Any] = {}
    page_tokens = {token[1] for token in abstract_order if token[0] == "page"}
    for page in sorted(page_tokens):
        element_index = page - 1 - sum(1 for p in deleted_pages if p < page)
        elements = _sld_id_entries(presentation)
        if 0 <= element_index < len(elements):
            slide_layout_for_page[page] = presentation.slides[
                element_index
            ].slide_layout

    new_elements: dict[int, Any] = {}
    for token in [t for t in abstract_order if t[0] == "new"]:
        counter = token[1]
        payload = new_slide_payloads.get(counter, {})
        fallback_layout = next(
            (
                slide_layout_for_page[p]
                for p in sorted(slide_layout_for_page)
            ),
            presentation.slide_layouts[-1],
        )
        layout = _blank_layout(presentation, fallback_layout)
        if new_slide_kind[counter] == "add_question_slide":
            image_paths = [
                str(item)
                for item in (payload.get("image_paths") or [])
                if str(item) and Path(str(item)).is_file()
            ]
            answer_html = str(payload.get("answer_html") or "")
            element_slide = _add_question_slide(
                presentation,
                layout=layout,
                title=_strip_markup(payload.get("title") or "课堂练习"),
                stem_paragraphs=_parse_rich_text(
                    str(payload.get("stem_html") or "")
                ),
                image_paths=image_paths,
                answer_paragraphs=(
                    _parse_rich_text(answer_html) if answer_html else None
                ),
            )
        else:
            element_slide = presentation.slides.add_slide(layout)
            box = element_slide.shapes.add_textbox(
                Inches(0.5),
                Inches(0.5),
                presentation.slide_width - Inches(1.0),
                Inches(1.2),
            )
            _fill_text_frame(
                box.text_frame,
                [_ParagraphSpec(runs=[_RunSpec(text=str(payload.get("text") or ""))])],
                size_pt=18,
            )
        new_elements[counter] = _sld_id_entries(presentation)[-1]

    # 删除页
    elements = _sld_id_entries(presentation)
    original_elements = elements[:source_page_count]
    for page in sorted(deleted_pages):
        if 1 <= page <= len(original_elements):
            _drop_slide_element(presentation, original_elements[page - 1])

    # 重排：把 sldIdLst 排成抽象页序。
    sld_list = _sld_id_list(presentation)
    ordered_elements: list[Any] = []
    for kind_token, value in abstract_order:
        if kind_token == "page":
            ordered_elements.append(original_elements[value - 1])
        else:
            ordered_elements.append(new_elements[value])
    for element in ordered_elements:
        sld_list.remove(element)
    for element in ordered_elements:
        sld_list.append(element)

    # 隐藏页：show 属性写在 slide part 根元素（p:sld）上。
    for entry in [e for e in report if e["kind"] == "hide_slide"]:
        page = entry["source_page"]
        if ("page", page) in abstract_order:
            final_index = abstract_order.index(("page", page))
            presentation.slides[final_index]._element.set("show", "0")

    presentation.save(str(output))
    try:
        verification = _load_presentation(output)
        final_page_count = len(_sld_id_entries(verification))
    except Exception as exc:  # pragma: no cover - 防御性
        raise PptxEditorError(f"副本保存后无法重新打开：{exc}") from exc

    final_position_by_question: list[dict[str, Any]] = []
    for index, token in enumerate(abstract_order, start=1):
        if token[0] == "new":
            payload = new_slide_payloads.get(token[1], {})
            if new_slide_kind.get(token[1]) == "add_question_slide":
                final_position_by_question.append(
                    {
                        "question_id": payload.get("question_id"),
                        "final_position": index,
                    }
                )

    return {
        "source_path": str(source),
        "output_path": str(output),
        "source_page_count": source_page_count,
        "final_page_count": final_page_count,
        "operations": report,
        "final_page_order": [
            token[1] if token[0] == "page" else f"new:{token[1]}"
            for token in abstract_order
        ],
        "inserted_question_pages": final_position_by_question,
        "output_sha256": _sha256_of(output),
    }
