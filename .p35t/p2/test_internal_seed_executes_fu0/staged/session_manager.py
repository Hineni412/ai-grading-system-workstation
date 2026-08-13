from __future__ import annotations

import copy
import hashlib
import io
import json
import os
import re
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator, Mapping, Sequence
from xml.etree import ElementTree

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from equivalence_engine import merge_equivalent_forms
from llm_client import LLMClient


DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS = 600.0
DEFAULT_CONFIG_GENERATION_RETRY_DELAYS = (2.0, 6.0)
DEFAULT_CONFIG_GENERATION_BATCH_SIZE = 3
MAX_WORD_IMAGES_PER_QUESTION = 16
MAX_WORD_IMAGES_WHOLE_DOCUMENT = 64
MAX_WORD_IMAGE_BYTES_PER_QUESTION = 24 * 1024 * 1024
MAX_WORD_IMAGE_BYTES_WHOLE_DOCUMENT = 48 * 1024 * 1024


class _QuestionGenerationRequestError(RuntimeError):
    def __init__(self, question_id: str, attempts: int, category: str, original: Exception) -> None:
        super().__init__(str(original))
        self.question_id = question_id
        self.attempts = attempts
        self.category = category
        self.original = original


def _config_generation_extra_kwargs() -> dict[str, float]:
    raw_timeout = os.getenv("AI_GRADING_CONFIG_TIMEOUT_SECONDS")
    try:
        timeout = float(raw_timeout) if raw_timeout else DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS
    except (TypeError, ValueError):
        timeout = DEFAULT_CONFIG_GENERATION_TIMEOUT_SECONDS
    return {"timeout": max(120.0, timeout)}


def save_uploaded_json(upload_dir: Path, file_name: str, file_bytes: bytes) -> Path:
    upload_dir.mkdir(parents=True, exist_ok=True)
    output_path = upload_dir / file_name

    parsed = json.load(io.BytesIO(file_bytes))
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)

    return output_path


def extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    seen_lines: set[str] = set()

    for block in _iter_doc_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text and text not in seen_lines:
                lines.append(text)
                seen_lines.add(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    cell_text = "\n".join(
                        para.text.strip()
                        for para in cell.paragraphs
                        if para.text.strip()
                    )
                    if cell_text:
                        cells.append(cell_text)
                row_text = " | ".join(cells)
                if row_text and row_text not in seen_lines:
                    lines.append(row_text)
                    seen_lines.add(row_text)

    # Some exam files put questions in Word text boxes/shapes. python-docx
    # does not expose those as normal paragraphs, so read raw XML text too.
    for text in _extract_docx_xml_text(file_bytes):
        if text and text not in seen_lines:
            lines.append(text)
            seen_lines.add(text)

    if not lines:
        raise ValueError("Word 鏂囨。鏈В鏋愬埌鏈夋晥鏂囨湰")

    return "\n".join(lines)


def _iter_doc_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _extract_docx_xml_text(file_bytes: bytes) -> list[str]:
    """Extract text from XML streams: main doc, headers/footers, footnotes, and text boxes."""
    xml_names = [
        "word/document.xml",
        "word/footnotes.xml",
        "word/endnotes.xml",
    ]
    result: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            xml_names.extend(
                name
                for name in archive.namelist()
                if re.match(r"word/(header|footer)\d+\.xml$", name)
            )
            for name in xml_names:
                if name not in archive.namelist():
                    continue
                raw_xml = archive.read(name)
                root = ElementTree.fromstring(raw_xml)
                # ① Extract OMML math formula as readable placeholder before iterating paragraphs
                # OMML uses namespace http://schemas.openxmlformats.org/officeDocument/2006/math
                for omath in root.iter():
                    tag = omath.tag
                    if not (tag.endswith("}oMath") or tag.endswith("}oMathPara")):
                        continue
                    # Collect all literal text runs inside the formula
                    math_chunks: list[str] = []
                    for node in omath.iter():
                        if node.tag.endswith("}t") and node.text:
                            math_chunks.append(node.text)
                        elif node.tag.endswith("}r") and node.text:
                            math_chunks.append(node.text)
                    formula_text = "".join(math_chunks).strip()
                    if formula_text:
                        placeholder = f"[公式: {formula_text}]"
                    else:
                        placeholder = "[数学公式]"
                    if placeholder not in result:
                        result.append(placeholder)

                # ② Extract paragraphs from main body AND from text boxes (txbxContent)
                for paragraph in root.iter():
                    if not paragraph.tag.endswith("}p"):
                        continue
                    chunks: list[str] = []
                    for node in paragraph.iter():
                        if node.tag.endswith("}t") and node.text:
                            chunks.append(node.text)
                        elif node.tag.endswith("}tab"):
                            chunks.append("\t")
                        elif node.tag.endswith("}br"):
                            chunks.append("\n")
                    text = re.sub(r"[ \t]+", " ", "".join(chunks)).strip()
                    if text and text not in result:
                        result.append(text)
    except Exception:
        return []
    return result


def _aligned_whole_generation_rules() -> str:
    return (
        "\n\n与分题生成模式一致的硬性规则：\n"
        "1. 每个 parts 项必须输出 response_mode：exact_objective、short_answer_points、"
        "process_required 或 visual_construction。\n"
        "2. 选择题和普通填空题只按最终答案判分，不得要求推理或计算过程。\n"
        "3. 要求列出全部可能答案的填空题必须使用 match_mode=complete_set，并输出 required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；少写、错写、多写均不得分。\n"
        "4. 知识点与题目标签由题库程序单独维护，不要输出任何 knowledge 字段。\n"
        "5. 主观题评分点必须写出可核验的必要条件、式子或结论，不得只写通用描述。\n"
        "6. 作图题应输出 visual_requirements 和必要踩分点，不得把答案图臆造为唯一文字答案。\n"
        "7. 总分严格为100；单题不超过18分；相同类型客观题必须同分，其他题型不要求同分。\n"
    )


def _build_whole_text_generation_prompt(
    doc_text: str,
    *,
    image_map: Sequence[str] = (),
    omitted_image_count: int = 0,
) -> str:
    image_context = ""
    if image_map:
        image_context = (
            "\n\nWord 关联图片说明：后附图片按下列顺序补充原文中的题图、答案图和解析图；"
            "图片与全文文字共同构成输入，不得忽略图片中的条件、标注、表格或公式。\n"
            f"图片顺序：{'; '.join(image_map)}"
        )
    if omitted_image_count:
        image_context += (
            f"\n安全载入提示：另有 {int(omitted_image_count)} 张关联图片因单次请求数量、"
            "总体积或文件可读性限制未附带；不得臆造这些图片中的内容。"
        )
    return (
        "这是 Word 整卷单次请求。你必须在本次响应中一次完成所有题目的解析、评分点生成与赋分；"
        "不要建议后续补充请求。\n"
        + _build_generation_prompt("", include_source_text=False)
        + _aligned_whole_generation_rules()
        + image_context
        + f"\nWord 解析文本：\n{doc_text}"
    )


def _build_whole_image_generation_prompt() -> str:
    return (
        "这是 PDF 整卷视觉单次请求。后续附带的图片按顺序对应整份试卷及答案/解析的各页原图。"
        "图片是唯一权威内容来源，不得参考、猜测或恢复任何 PDF 抽取文字。"
        "你必须在本次响应中一次完成所有题目的解析、评分点生成与赋分；不要建议后续补充请求。\n\n"
        + _build_generation_prompt("", include_source_text=False)
        + _aligned_whole_generation_rules()
    )


def _finalize_whole_generation_payload(
    payload: dict[str, Any],
    generation_mode: str,
) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("AI 整卷生成未返回有效 JSON 对象。")
    normalize_new_generated_config_payload(payload)
    force_payload_total_score(payload, target_total=100.0)
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["generation_mode"] = generation_mode
    meta["score_allocation_mode"] = "single_request_local_normalization"
    meta["single_request"] = True
    refresh_generated_config_quality_warnings(payload)
    try:
        validate_generated_config(payload)
    except Exception:
        _dump_failed_generated_payload(payload)
        raise
    return payload


def generate_grading_config_from_docx_text(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
    *,
    question_blocks: Sequence[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    image_blobs, image_map, omitted_image_count = (
        _word_document_image_request_assets(question_blocks or ())
    )
    if report:
        detail = (
            f"发送整份 Word 文本及 {len(image_blobs)} 张题目/答案关联图片，"
            "一次完成解析与赋分。"
            if image_blobs
            else "发送整份 Word 文本，一次完成解析与赋分。"
        )
        if omitted_image_count:
            detail += f" 另有 {omitted_image_count} 张图片受安全上限或可读性限制未载入。"
        report(0.18, "Word 整卷单次请求", detail)
    prompt = _build_whole_text_generation_prompt(
        doc_text,
        image_map=image_map,
        omitted_image_count=omitted_image_count,
    )
    if image_blobs:
        payload = llm_client.json_from_images_once(
            prompt,
            image_blobs,
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
            use_config_client=True,
        )
    else:
        payload = llm_client.json_from_text_once(
            prompt,
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
        )
    finalized = _finalize_whole_generation_payload(
        payload,
        "whole_word_text_single_request",
    )
    meta = finalized.setdefault("meta", {})
    if isinstance(meta, dict):
        meta["word_source_image_count"] = len(image_blobs)
        meta["word_source_images_omitted"] = int(omitted_image_count)
    return finalized


def generate_grading_config_from_docx_text_legacy(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    if report:
        report(0.18, "旧版整卷生成", "直接把整份 Word 文本发送给模型生成评分标准。")
    prompt = _build_generation_prompt(doc_text)
    payload = llm_client.json_from_text(prompt, model=model_name, extra_kwargs=_config_generation_extra_kwargs())
    normalize_new_generated_config_payload(payload)
    validate_generated_config(payload)
    if _needs_objective_repair(payload, doc_text):
        if report:
            report(0.70, "旧版整卷修复", "检测到客观题可能缺失，正在请求模型修复。")
        payload = llm_client.json_from_text(
            _build_objective_repair_prompt(doc_text, payload),
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
        )
        normalize_new_generated_config_payload(payload)
        validate_generated_config(payload)
    return payload


def _config_generation_retry_delays() -> tuple[float, ...]:
    raw = str(os.getenv("AI_GRADING_CONFIG_RETRY_DELAYS") or "").strip()
    if not raw:
        return DEFAULT_CONFIG_GENERATION_RETRY_DELAYS
    try:
        values = tuple(max(0.0, float(item.strip())) for item in raw.split(",") if item.strip())
    except ValueError:
        return DEFAULT_CONFIG_GENERATION_RETRY_DELAYS
    return values or DEFAULT_CONFIG_GENERATION_RETRY_DELAYS


def _is_transient_config_generation_error(exc: Exception) -> bool:
    status_code = getattr(exc, "status_code", None)
    if status_code == 429 or (isinstance(status_code, int) and 500 <= status_code <= 599):
        return True
    class_name = exc.__class__.__name__.lower()
    if any(token in class_name for token in ("apiconnection", "apitimeout", "ratelimit", "timeout", "connection")):
        return True
    message = str(exc).lower()
    return any(
        token in message
        for token in (
            "connection reset",
            "connection aborted",
            "connection error",
            "temporarily unavailable",
            "temporary upstream",
            "timed out",
            "timeout",
            "rate limit",
            "too many requests",
            "service unavailable",
            "bad gateway",
            "gateway timeout",
        )
    )


def _call_question_generation_with_retry(
    question_id: str,
    request: Callable[[], dict[str, Any]],
) -> tuple[dict[str, Any], int]:
    retry_delays = _config_generation_retry_delays()
    attempts = 0
    while True:
        attempts += 1
        try:
            return request(), attempts
        except Exception as exc:
            transient = _is_transient_config_generation_error(exc)
            retry_index = attempts - 1
            if not transient or retry_index >= len(retry_delays):
                category = "transient_network" if transient else "non_retryable"
                raise _QuestionGenerationRequestError(question_id, attempts, category, exc) from exc
            time.sleep(retry_delays[retry_index])


def _word_block_image_candidates(
    block: dict[str, Any],
) -> list[tuple[str, str]]:
    candidates: list[tuple[str, str]] = []
    image_paths = block.get("image_paths")
    if isinstance(image_paths, list):
        candidates.extend(
            ("question", str(path).strip())
            for path in image_paths
            if str(path).strip()
        )
    for role, key in (
        ("question", "question_html"),
        ("answer", "answer_html"),
        ("analysis", "analysis_html"),
    ):
        candidates.extend(
            (role, raw_path)
            for raw_path in _image_paths_from_rich_text(
                str(block.get(key) or "")
            )
        )
    return candidates


def _word_block_image_assets(
    block: dict[str, Any],
    *,
    limit: int = MAX_WORD_IMAGES_PER_QUESTION,
    max_total_bytes: int = MAX_WORD_IMAGE_BYTES_PER_QUESTION,
) -> tuple[list[dict[str, Any]], int]:
    max_images = max(0, int(limit))
    byte_budget = max(0, int(max_total_bytes))
    assets: list[dict[str, Any]] = []
    seen: set[str] = set()
    role_ordinals = {"question": 0, "answer": 0, "analysis": 0}
    total_bytes = 0
    omitted_count = 0
    for role, raw_path in _word_block_image_candidates(block):
        candidates = [Path(raw_path)]
        if not Path(raw_path).is_absolute():
            candidates.append(Path.cwd() / raw_path)
        path = next((candidate for candidate in candidates if candidate.is_file()), None)
        if path is None:
            omitted_count += 1
            continue
        resolved = str(path.resolve())
        if resolved in seen:
            continue
        seen.add(resolved)
        try:
            file_size = int(path.stat().st_size)
        except OSError:
            omitted_count += 1
            continue
        if (
            len(assets) >= max_images
            or file_size <= 0
            or file_size > byte_budget
            or total_bytes + file_size > byte_budget
        ):
            omitted_count += 1
            continue
        try:
            content = path.read_bytes()
        except OSError:
            omitted_count += 1
            continue
        if (
            not content
            or len(content) > byte_budget
            or total_bytes + len(content) > byte_budget
        ):
            omitted_count += 1
            continue
        total_bytes += len(content)
        role_ordinals[role] += 1
        assets.append(
            {
                "role": role,
                "ordinal": role_ordinals[role],
                "content": content,
            }
        )
    return assets, omitted_count


def _word_block_image_blobs(
    block: dict[str, Any],
    *,
    limit: int = 8,
) -> list[bytes]:
    assets, _omitted_count = _word_block_image_assets(
        block,
        limit=limit,
        max_total_bytes=MAX_WORD_IMAGE_BYTES_PER_QUESTION,
    )
    return [
        bytes(asset["content"])
        for asset in assets
        if isinstance(asset.get("content"), (bytes, bytearray))
    ]


def _word_document_image_request_assets(
    question_blocks: Sequence[dict[str, Any]],
) -> tuple[list[bytes], list[str], int]:
    image_blobs: list[bytes] = []
    image_labels: list[list[str]] = []
    digest_to_index: dict[str, int] = {}
    total_bytes = 0
    omitted_count = 0
    role_labels = {
        "question": "题目图",
        "answer": "答案图",
        "analysis": "解析图",
    }

    for block_index, block in enumerate(question_blocks, start=1):
        if not isinstance(block, dict):
            continue
        question_id = str(block.get("question_id") or "").strip() or (
            f"第{block_index}题"
        )
        assets, block_omitted = _word_block_image_assets(
            block,
            limit=MAX_WORD_IMAGES_PER_QUESTION,
            max_total_bytes=MAX_WORD_IMAGE_BYTES_PER_QUESTION,
        )
        omitted_count += block_omitted
        for asset in assets:
            content = asset.get("content")
            if not isinstance(content, (bytes, bytearray)):
                omitted_count += 1
                continue
            blob = bytes(content)
            role = str(asset.get("role") or "question")
            ordinal = int(asset.get("ordinal") or 1)
            label = (
                f"{question_id}{role_labels.get(role, '关联图')}{ordinal}"
            )
            digest = hashlib.sha256(blob).hexdigest()
            existing_index = digest_to_index.get(digest)
            if existing_index is not None:
                if label not in image_labels[existing_index]:
                    image_labels[existing_index].append(label)
                continue
            if (
                len(image_blobs) >= MAX_WORD_IMAGES_WHOLE_DOCUMENT
                or total_bytes + len(blob)
                > MAX_WORD_IMAGE_BYTES_WHOLE_DOCUMENT
            ):
                omitted_count += 1
                continue
            digest_to_index[digest] = len(image_blobs)
            image_blobs.append(blob)
            image_labels.append([label])
            total_bytes += len(blob)

    image_map = [
        f"图片{index}={'、'.join(labels)}"
        for index, labels in enumerate(image_labels, start=1)
    ]
    return image_blobs, image_map, omitted_count


def _generate_question_block_results(
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any] | None], list[dict[str, Any]], dict[str, int], int, bool]:
    max_workers = _bounded_int(
        os.getenv("AI_GRADING_CONFIG_WORKERS") or os.getenv("AI_GRADING_MAX_WORKERS"),
        8,
        1,
        1000,
    )
    worker_count = min(max_workers, len(question_blocks))
    results: list[dict[str, Any] | None] = [None] * len(question_blocks)
    failures: list[dict[str, Any]] = []
    attempt_counts: dict[str, int] = {}
    image_semantic_mode = bool(q_images) or any(
        str(block.get("semantic_source") or "").strip() == "images"
        for block in question_blocks
    )

    if report:
        qids = ", ".join(str(block.get("question_id") or "") for block in question_blocks if block.get("question_id"))
        report(0.18, "Split paper", f"Parsed {len(question_blocks)} questions; workers={worker_count}; qids: {qids}")

    import base64

    with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="config-q") as executor:
        future_map = {}
        for index, block in enumerate(question_blocks):
            qid = str(block.get("question_id") or f"Q{index + 1}")
            image_data = q_images.get(qid) if isinstance(q_images, dict) else None
            has_question_images = bool(isinstance(image_data, dict) and image_data.get("question"))
            image_semantic_source = str(block.get("semantic_source") or "").strip() == "images"
            word_image_blobs = (
                _word_block_image_blobs(block)
                if not has_question_images and not image_semantic_source
                else []
            )
            prompt = (
                _build_single_question_image_generation_prompt(
                    block,
                    has_answer_image=bool(isinstance(image_data, dict) and image_data.get("answer")),
                )
                if has_question_images or image_semantic_source
                else _build_single_question_generation_prompt(block, doc_text)
            )
            if has_question_images and isinstance(image_data, dict):
                img_list = [base64.b64decode(image_data["question"], validate=True)]
                if image_data.get("answer"):
                    img_list.append(base64.b64decode(image_data["answer"], validate=True))

                def request(prompt: str = prompt, img_list: list[bytes] = img_list) -> dict[str, Any]:
                    return llm_client.json_from_images(prompt, img_list, model=model_name)
            elif word_image_blobs:
                def request(prompt: str = prompt, img_list: list[bytes] = word_image_blobs) -> dict[str, Any]:
                    return llm_client.json_from_images(prompt, img_list, model=model_name)
            else:
                def request(prompt: str = prompt) -> dict[str, Any]:
                    return llm_client.json_from_text(prompt, model=model_name)

            future_map[executor.submit(_call_question_generation_with_retry, qid, request)] = index

        if report:
            qids = ", ".join(str(block.get("question_id") or "") for block in question_blocks if block.get("question_id"))
            report(0.20, "Submit per-question requests", f"Submitted {len(future_map)} question requests; waiting for model: {qids}")
        completed = 0
        total = len(question_blocks)
        for future in as_completed(future_map):
            index = future_map[future]
            qid = str(question_blocks[index].get("question_id") or f"Q{index + 1}")
            try:
                results[index], attempt_counts[qid] = future.result()
            except _QuestionGenerationRequestError as exc:
                attempt_counts[qid] = exc.attempts
                failures.append(
                    {
                        "question_id": qid,
                        "attempts": exc.attempts,
                        "category": exc.category,
                        "error": str(exc.original),
                    }
                )
            except Exception as exc:
                attempt_counts[qid] = 1
                failures.append(
                    {
                        "question_id": qid,
                        "attempts": 1,
                        "category": "unexpected",
                        "error": str(exc),
                    }
                )
            completed += 1
            if report:
                progress = 0.18 + 0.7 * (completed / total)
                report(progress, "Parse question", f"Question {index + 1} parsed (progress: {completed}/{total})")

    return results, failures, attempt_counts, worker_count, image_semantic_mode


def _append_unmergeable_question_failures(
    merged: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    attempt_counts: dict[str, int],
) -> list[dict[str, Any]]:
    questions = merged.get("rubric", {}).get("questions", [])
    merged_qids = {
        str(question.get("question_id") or "").strip()
        for question in questions
        if isinstance(question, dict)
    } if isinstance(questions, list) else set()
    failed_qids = {str(item.get("question_id") or "").strip() for item in failures}
    for block in question_blocks:
        qid = str(block.get("question_id") or "").strip()
        if not qid or qid in merged_qids or qid in failed_qids:
            continue
        failures.append(
            {
                "question_id": qid,
                "attempts": int(attempt_counts.get(qid, 1)),
                "category": "unmergeable_schema",
                "error": "AI returned a result that could not be merged into the rubric schema",
            }
        )
    return failures


def _generate_grading_config_by_question_blocks(
    question_blocks: list[dict[str, str]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    _validate_image_semantic_inputs(question_blocks, q_images)
    results, failures, attempt_counts, worker_count, image_semantic_mode = _generate_question_block_results(
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )

    if report:
        report(0.90, "Assemble rubric", "All questions parsed; assembling final scoring rubric...")

    merged = _merge_single_question_payloads(results, question_blocks)
    failures = _append_unmergeable_question_failures(merged, question_blocks, failures, attempt_counts)
    _attach_parallel_generation_meta(merged, question_blocks, failures, worker_count, attempt_counts)
    _ensure_question_blocks_covered(merged, question_blocks)
    _apply_local_question_facts(merged, question_blocks)
    normalize_new_generated_config_payload(merged)

    if failures:
        meta = merged.setdefault("meta", {})
        meta["score_allocation_mode"] = "pending_failed_questions"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = True
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    # ① Phase 2: dedicated score-allocation AI call
    # All per-question prompts set score placeholders (=1); this step assigns real scores.
    return retry_grading_config_score_allocation(
        merged,
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        include_document_text=not image_semantic_mode,
    )


def _validate_image_semantic_inputs(
    question_blocks: list[dict[str, Any]],
    q_images: dict[str, Any] | None,
) -> None:
    image_qids = [
        str(block.get("question_id") or "").strip()
        for block in question_blocks
        if isinstance(block, dict) and str(block.get("semantic_source") or "").strip() == "images"
    ]
    if not image_qids:
        return
    if not isinstance(q_images, dict) or not q_images:
        raise ValueError("PDF 图片裁题状态已丢失，请重新点击“① 拆分试卷”后再生成评分标准。")

    import base64

    invalid_qids: list[str] = []
    for qid in image_qids:
        image_data = q_images.get(qid)
        question_image = image_data.get("question") if isinstance(image_data, dict) else None
        if not isinstance(question_image, str) or not question_image.strip():
            invalid_qids.append(qid)
            continue
        try:
            base64.b64decode(question_image, validate=True)
            answer_image = image_data.get("answer")
            if answer_image:
                base64.b64decode(str(answer_image), validate=True)
        except (ValueError, TypeError):
            invalid_qids.append(qid)
    if invalid_qids:
        raise ValueError(
            "PDF 图片裁题数据无效，未退回纯文本模式。请重新拆题："
            + ", ".join(invalid_qids)
        )


def preview_question_blocks_from_docx_text(doc_text: str) -> list[dict[str, Any]]:
    """本地拆题预览，不调用任何 AI。

    供 UI 在发送 AI 之前展示拆题结果，让人工确认每道题拆分是否正确、答案是否靠谱、
    题型判断是否准确（可在预览中删题或切换题型）。直接复用本地拆题逻辑。
    """
    return _extract_local_question_blocks(doc_text)


_INLINE_IMAGE_MARKER = re.compile(r"\[\[IMAGE:(?P<path>.+?)\]\]")
_INLINE_MAIN_QUESTION_MARKER = re.compile(
    r"(?P<prefix>[。！？!?．.][ \t\r\n]*)"
    r"(?P<number>\d{1,2})[ \t]*[.．、](?![ \t]*\d)[ \t]*"
)


class _ControlledDocxWriteError(RuntimeError):
    pass


def preview_question_blocks_from_docx_bytes(
    file_bytes: bytes,
    *,
    fallback_doc_text: str = "",
    temporary_root: str | Path | None = None,
    asset_root: str | Path | None = None,
    register_created_file: Callable[[Path], None] | None = None,
    write_created_file: Callable[[Path, bytes], None] | None = None,
) -> list[dict[str, Any]]:
    """富文本拆题预览（复用题库 import_docx），保留公式 HTML 与图片，不调用 AI。

    解析失败时回退到纯文本拆题（preview_question_blocks_from_docx_text）。
    """
    from backend.document_parsing import parse_docx_question_blocks

    return parse_docx_question_blocks(
        file_bytes,
        fallback_doc_text=fallback_doc_text,
        temporary_root=(
            temporary_root
            if temporary_root is not None
            else _resolve_upload_config_dir()
        ),
        asset_root=asset_root,
        register_created_file=register_created_file,
        write_created_file=write_created_file,
    )


def _extract_rich_question_blocks(
    file_bytes: bytes,
    *,
    temporary_root: str | Path | None = None,
    asset_root: str | Path | None = None,
    register_created_file: Callable[[Path], None] | None = None,
    write_created_file: Callable[[Path, bytes], None] | None = None,
) -> list[dict[str, Any]] | None:
    """用题库 import_docx + map_rich_content_by_number 提取每题富文本块。"""
    from question_bank.importers.docx_importer import import_docx
    from question_bank.importers.batch_importer import map_rich_content_by_number

    tmp_dir = (
        Path(temporary_root)
        if temporary_root is not None
        else Path(_resolve_upload_config_dir())
    )
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    tmp_path = tmp_dir / f"_rich_split_{ts}.docx"

    def controlled_writer(path: Path, content: bytes) -> None:
        assert write_created_file is not None
        try:
            write_created_file(path, content)
        except Exception as error:
            raise _ControlledDocxWriteError() from error

    extracted = import_docx(
        io.BytesIO(file_bytes),
        source_name=tmp_path,
        asset_root=asset_root,
        asset_root_is_output_dir=write_created_file is not None,
        register_created_file=register_created_file,
        write_created_file=(controlled_writer if write_created_file is not None else None),
    )
    content = map_rich_content_by_number(
        _split_inline_main_question_paragraphs(extracted.rich_paragraphs),
        source_file=str(tmp_path),
    )

    question_map = content.get("question") if isinstance(content, dict) else {}
    answer_map = content.get("answer") if isinstance(content, dict) else {}
    if not isinstance(question_map, dict) or not question_map:
        return None

    # 题号按数值升序
    def _num_key(k: str) -> int:
        try:
            return int(str(k).strip())
        except (TypeError, ValueError):
            return 999

    choice_answers = _extract_choice_answer_sequence(
        _rich_blocks_plain_text(_flatten_answer_blocks(answer_map))
    )

    blocks: list[dict[str, Any]] = []
    for num_key in sorted(question_map.keys(), key=_num_key):
        number = _num_key(num_key)
        if not (1 <= number <= 99):
            continue
        q_blocks = question_map.get(num_key) or []
        a_blocks = answer_map.get(num_key) if isinstance(answer_map, dict) else None
        block = _parse_rich_question_blocks(number, q_blocks, a_blocks, choice_answers)
        if block:
            blocks.append(block)
    return blocks or None


def _split_inline_main_question_paragraphs(
    rich_paragraphs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Split only consecutive main-question markers embedded in DOCX paragraphs."""
    normalized: list[dict[str, Any]] = []
    current_number: int | None = None
    in_answer_section = False
    plain_paragraphs = [
        _strip_inline_html(
            _INLINE_IMAGE_MARKER.sub("", str(paragraph.get("text") or ""))
        ).strip()
        for paragraph in rich_paragraphs
    ]

    for index, paragraph in enumerate(rich_paragraphs):
        text = str(paragraph.get("text") or "").strip()
        plain_text = plain_paragraphs[index]
        if _looks_like_answer_section_heading(plain_text):
            in_answer_section = True
        if in_answer_section or not text:
            normalized.append(paragraph)
            continue

        leading_number = _extract_question_marker_number(plain_text)
        if leading_number is not None:
            current_number = leading_number
        if current_number is None:
            normalized.append(paragraph)
            continue

        text_without_images = _INLINE_IMAGE_MARKER.sub("", text).strip()
        segments, final_number = _split_consecutive_inline_main_questions(
            text_without_images,
            current_number=current_number,
            following_number=_next_leading_main_question_number(
                plain_paragraphs,
                after_index=index,
            ),
        )
        if len(segments) == 1:
            normalized.append(paragraph)
            continue

        segments = _attach_inline_images_to_source_segments(
            paragraph,
            segments=segments,
            current_number=current_number,
            final_number=final_number,
        )
        for segment in segments:
            normalized.append({**paragraph, "text": segment})
        current_number = final_number

    return normalized


def _split_consecutive_inline_main_questions(
    text: str,
    *,
    current_number: int,
    following_number: int | None,
) -> tuple[list[str], int]:
    split_offsets: list[int] = []
    expected_number = current_number + 1
    for marker in _INLINE_MAIN_QUESTION_MARKER.finditer(text):
        marker_number = int(marker.group("number"))
        if marker_number != expected_number:
            continue
        split_offsets.append(marker.start("number"))
        expected_number += 1
    if not split_offsets or expected_number != following_number:
        return [text], current_number

    boundaries = [0, *split_offsets, len(text)]
    segments = [
        text[start:end].strip()
        for start, end in zip(boundaries, boundaries[1:])
        if text[start:end].strip()
    ]
    return segments, expected_number - 1


def _attach_inline_images_to_source_segments(
    paragraph: dict[str, Any],
    *,
    segments: list[str],
    current_number: int,
    final_number: int,
) -> list[str]:
    fallback_paths = _image_paths_from_rich_text(
        str(paragraph.get("text") or "")
    )

    def attach_to_last(paths: list[str]) -> list[str]:
        assigned = list(segments)
        for path in paths:
            assigned[-1] = f"{assigned[-1].rstrip()}\n[[IMAGE:{path}]]"
        return assigned

    image_relationships = paragraph.get("image_relationships")
    raw_xml = str(paragraph.get("xml") or "")
    if not isinstance(image_relationships, dict) or not image_relationships or not raw_xml:
        return attach_to_last(fallback_paths)

    try:
        root = ElementTree.fromstring(raw_xml)
    except ElementTree.ParseError:
        return attach_to_last(fallback_paths)

    text_parts: list[str] = []
    positioned_images: list[tuple[int, str]] = []
    text_length = 0
    for element in root.iter():
        local_name = str(element.tag).split("}")[-1]
        if local_name == "t":
            value = str(element.text or "")
            text_parts.append(value)
            text_length += len(value)
            continue
        if local_name != "blip":
            continue
        relationship_id = next(
            (
                str(value)
                for key, value in element.attrib.items()
                if str(key).split("}")[-1] == "embed"
            ),
            "",
        )
        image_path = str(image_relationships.get(relationship_id) or "").strip()
        if image_path:
            positioned_images.append((text_length, image_path))

    if not positioned_images:
        return attach_to_last(fallback_paths)

    source_text = "".join(text_parts)
    expected_number = current_number + 1
    split_offsets: list[int] = []
    for marker in _INLINE_MAIN_QUESTION_MARKER.finditer(source_text):
        marker_number = int(marker.group("number"))
        if marker_number != expected_number:
            continue
        split_offsets.append(marker.start("number"))
        expected_number += 1
        if marker_number == final_number:
            break
    if len(split_offsets) != len(segments) - 1:
        return attach_to_last(fallback_paths)

    assigned = list(segments)
    positioned_paths: set[str] = set()
    for image_offset, image_path in positioned_images:
        segment_index = sum(image_offset >= offset for offset in split_offsets)
        assigned[segment_index] = (
            f"{assigned[segment_index].rstrip()}\n[[IMAGE:{image_path}]]"
        )
        positioned_paths.add(image_path)
    for image_path in fallback_paths:
        if image_path not in positioned_paths:
            assigned[-1] = f"{assigned[-1].rstrip()}\n[[IMAGE:{image_path}]]"
    return assigned


def _next_leading_main_question_number(
    texts: list[str],
    *,
    after_index: int,
) -> int | None:
    for text in texts[after_index + 1 :]:
        if _looks_like_answer_section_heading(text):
            return None
        number = _extract_question_marker_number(text)
        if number is not None:
            return number
    return None


def _flatten_answer_blocks(answer_map: Any) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    if isinstance(answer_map, dict):
        for _k, blocks in answer_map.items():
            if isinstance(blocks, list):
                out.extend(b for b in blocks if isinstance(b, dict))
    return out


def _rich_block_text(block: Any) -> str:
    if isinstance(block, dict):
        return str(block.get("text") or "")
    return str(block or "")


def _rich_blocks_plain_text(blocks: list[Any]) -> str:
    parts = [_rich_block_text(b) for b in blocks]
    return _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", "\n".join(parts)))


def _strip_inline_html(value: str) -> str:
    """去掉公式 HTML 标签，得到可用于答案/题型推断的纯文本。"""
    text = re.sub(r"</?(?:sub|sup|u|table|tbody|tr|td|th|br)\b[^>]*>", "", str(value or ""), flags=re.IGNORECASE)
    return text


def _image_paths_from_rich_text(value: str) -> list[str]:
    seen: list[str] = []
    for m in _INLINE_IMAGE_MARKER.finditer(str(value or "")):
        p = m.group("path").strip()
        if p and p not in seen:
            seen.append(p)
    return seen


def _parse_answer_section_blocks(answer_blocks: list[Any]) -> tuple[str, str]:
    """从卷末"参考答案"块中提取（最终答案, 解答过程）。

    答案块通常重列了题干/选项，需跳过；只保留 【分析】/【解答】 后、【点评】 前的解答内容，
    并优先抽取"故选X""故答案为：X"作为最终答案。
    """
    in_solution = False
    solution_lines: list[str] = []
    final_answer = ""
    for blk in answer_blocks:
        text = _rich_block_text(blk)
        for raw in (text.splitlines() if "\n" in text else [text]):
            line = str(raw or "").strip()
            if not line:
                continue
            if "【点评】" in line or "【点睛】" in line:
                in_solution = False
                continue
            if "【解答】" in line or "【分析】" in line or "【解析】" in line:
                in_solution = True
                marker = next(m for m in ("【解答】", "【分析】", "【解析】") if m in line)
                after = line.split(marker, 1)[1].strip()
                if after:
                    solution_lines.append(after)
                continue
            if line.startswith("【") and "】" in line:
                # 其他标记（如【答案】）后内容并入解答
                in_solution = True
                after = line.split("】", 1)[1].strip()
                if after:
                    solution_lines.append(after)
                continue
            if in_solution:
                solution_lines.append(line)

    solution_text = "\n".join(solution_lines).strip()
    # 抽取最终答案：故选X / 故答案为：X
    plain = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", solution_text))
    m = re.search(r"故\s*选\s*[:：]?\s*([A-Da-d]+)", plain)
    if m:
        final_answer = m.group(1).upper()
    else:
        m = re.search(r"故\s*答\s*案\s*为\s*[:：]?\s*([^。\n．]+)", plain)
        if m:
            final_answer = m.group(1).strip().rstrip("．.")
    return final_answer, solution_text


def _parse_rich_question_blocks(
    number: int,
    question_blocks: list[Any],
    answer_blocks: list[Any] | None,
    choice_answers: dict[str, str],
) -> dict[str, Any] | None:
    """把单题的富文本块（题干区，含内联【答案】【解析】）拆成 题干/答案/解析 富文本。"""
    # 题干区逐行送入分桶（沿用纯文本分桶逻辑），保留 HTML 公式与 [[IMAGE:]] 标记
    lines: list[str] = []
    for blk in question_blocks:
        text = _rich_block_text(blk)
        lines.extend(text.splitlines() if "\n" in text else [text])

    bucket = "stem"
    stem_lines: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []
    for raw in lines:
        line = str(raw or "").strip()
        if not line:
            continue
        if "【答案】" in line:
            bucket = "answer"
            after = line.split("【答案】", 1)[1].strip()
            if after:
                answer_lines.append(after)
            continue
        if ("【解析】" in line) or ("【点睛】" in line):
            bucket = "analysis"
            marker = "【解析】" if "【解析】" in line else "【点睛】"
            after = line.split(marker, 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if line.startswith("【") and "】" in line:
            bucket = "analysis"
            after = line.split("】", 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if bucket == "stem":
            stem_lines.append(line)
        elif bucket == "answer":
            answer_lines.append(line)
        else:
            analysis_lines.append(line)

    # 若题干区没有内联答案，但存在独立答案块（卷末"参考答案"区），从中解析。
    # 答案块结构通常为：[重列题干/选项] 【分析】… 【解答】…故选C/故答案为：X 【点评】…
    # 需跳过重列的题干/选项，只取标记后的解答与最终答案。
    if answer_blocks and not answer_lines:
        ans_final, ana_text = _parse_answer_section_blocks(answer_blocks)
        if ans_final:
            answer_lines.append(ans_final)
        if ana_text and not analysis_lines:
            analysis_lines.append(ana_text)

    question_html = _strip_leading_question_number(number, "\n".join(stem_lines).strip())
    answer_html = "\n".join(answer_lines).strip()
    analysis_html = "\n".join(analysis_lines).strip()

    image_paths = _image_paths_from_rich_text(question_html)

    question_text = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", question_html)).strip()
    answer_text = _strip_inline_html(_INLINE_IMAGE_MARKER.sub("", answer_html)).strip()

    if not question_text and not answer_text and not analysis_html and not image_paths:
        return None

    num_str = str(number)
    qtype = _infer_local_question_type(question_text, answer_text, num_str)
    canonical = _extract_canonical_answer_for_local_question(
        number=num_str,
        qtype=qtype,
        question_text=question_text,
        answer_text=answer_text,
        choice_answers=choice_answers,
    )
    accepted = _local_accepted_forms(canonical, qtype)
    has_answer = bool(answer_text or canonical)
    return {
        "question_id": f"Q{number}",
        "text": question_text,
        "question_text": question_text,
        "question_html": question_html,
        "answer_text": answer_text,
        "answer_html": answer_html,
        "analysis": _strip_inline_html(analysis_html).strip(),
        "analysis_html": analysis_html,
        "image_paths": image_paths,
        "question_type": qtype,
        "canonical_answer": canonical,
        "accepted_forms": accepted,
        "local_answer_trusted": has_answer,
        "needs_review": (not bool(canonical)) if qtype in {"choice", "fill_blank"} else False,
    }


def _resolve_upload_config_dir() -> str:
    try:
        from path_manager import PathManager  # type: ignore

        return str(PathManager().upload_config_dir)
    except Exception:
        return str(Path.cwd() / "user_data" / "config" / "uploaded")


def generate_grading_config_from_confirmed_blocks(
    confirmed_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, str] = None,
) -> dict[str, Any]:
    """用人工确认后的题块直接进入单题并发生成 + 统一赋分管线，跳过重新拆题。

    confirmed_blocks 应来自 preview_question_blocks_from_docx_text 的输出，
    经人工删除/题型修正后传入。
    """
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in confirmed_blocks
    ]
    return _generate_grading_config_by_question_blocks(
        locked_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )


class _BatchSchemaMismatch(ValueError):
    pass


def generate_grading_config_in_batches(
    confirmed_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    batch_size: int = DEFAULT_CONFIG_GENERATION_BATCH_SIZE,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Generate one rubric in small sequential batches with no hidden model calls."""
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in confirmed_blocks
    ]
    if not locked_blocks:
        raise ValueError("至少需要一道已确认题目。")
    clean_size = max(1, min(20, int(batch_size)))
    return _run_config_generation_batches(
        existing_payload=None,
        question_blocks=locked_blocks,
        doc_text=doc_text,
        llm_client=llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        batch_size=clean_size,
        retry_question_ids=None,
        checkpoint=checkpoint,
    )


def retry_failed_grading_config_batches(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    retry_question_ids: Sequence[str] | None = None,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    meta = existing_payload.get("meta") if isinstance(existing_payload, dict) else None
    failed_batches = meta.get("failed_batches") if isinstance(meta, dict) else None
    if not isinstance(failed_batches, list) or not failed_batches:
        return _score_completed_batch_draft_once(
            existing_payload,
            question_blocks,
            doc_text,
            llm_client,
            model_name=model_name,
            report=report,
            q_images=q_images,
            checkpoint=checkpoint,
        )

    failed_by_id: dict[str, list[str]] = {}
    for item in failed_batches:
        if not isinstance(item, dict):
            continue
        batch_id = str(item.get("batch_id") or "").strip()
        raw_ids = item.get("question_ids")
        ids = [str(qid).strip() for qid in raw_ids or [] if str(qid).strip()]
        if batch_id and ids:
            failed_by_id[batch_id] = ids
    if not failed_by_id:
        raise ValueError("失败批次记录已损坏，请重新开始分批生成。")

    if retry_question_ids is None:
        selected_question_ids = {
            qid for ids in failed_by_id.values() for qid in ids
        }
    else:
        selected_ids = list(
            dict.fromkeys(str(qid).strip() for qid in retry_question_ids if str(qid).strip())
        )
        selected_set = set(selected_ids)
        selected_batch_ids = [
            batch_id
            for batch_id, ids in failed_by_id.items()
            if set(ids).issubset(selected_set)
        ]
        expected_set = {
            qid for batch_id in selected_batch_ids for qid in failed_by_id[batch_id]
        }
        if not selected_set or selected_set != expected_set:
            raise ValueError("只能按完整失败批次重试，不能只选择批次中的部分题目。")
        selected_question_ids = expected_set

    existing_size = meta.get("batch_size") if isinstance(meta, dict) else None
    try:
        batch_size = int(existing_size)
    except (TypeError, ValueError):
        batch_size = DEFAULT_CONFIG_GENERATION_BATCH_SIZE
    locked_blocks = [
        {**block, "question_type_confirmed": True}
        for block in question_blocks
    ]
    return _run_config_generation_batches(
        existing_payload=existing_payload,
        question_blocks=locked_blocks,
        doc_text=doc_text,
        llm_client=llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        batch_size=batch_size,
        retry_question_ids=selected_question_ids,
        checkpoint=checkpoint,
    )


def failed_grading_config_batches(payload: dict[str, Any]) -> list[dict[str, Any]]:
    meta = payload.get("meta") if isinstance(payload, dict) else None
    value = meta.get("failed_batches") if isinstance(meta, dict) else None
    return copy.deepcopy(value) if isinstance(value, list) else []


def _run_config_generation_batches(
    *,
    existing_payload: dict[str, Any] | None,
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None,
    report: Any,
    q_images: dict[str, Any] | None,
    batch_size: int,
    retry_question_ids: set[str] | None,
    checkpoint: Callable[[dict[str, Any]], None] | None,
) -> dict[str, Any]:
    _validate_unique_batch_question_ids(question_blocks)
    _validate_image_semantic_inputs(question_blocks, q_images)
    batch_groups: list[list[dict[str, Any]]] = []
    objective_group: list[dict[str, Any]] = []
    for block in question_blocks:
        question_type = str(block.get("question_type") or "").strip().lower()
        if question_type in {"choice", "fill_blank"}:
            objective_group.append(block)
            if len(objective_group) >= batch_size:
                batch_groups.append(objective_group)
                objective_group = []
            continue
        if objective_group:
            batch_groups.append(objective_group)
            objective_group = []
        batch_groups.append([block])
    if objective_group:
        batch_groups.append(objective_group)

    batches = [
        {
            "batch_id": f"B{index + 1:03d}",
            "question_ids": [
                str(block.get("question_id") or "").strip() for block in group
            ],
            "blocks": group,
        }
        for index, group in enumerate(batch_groups)
    ]
    merged = copy.deepcopy(existing_payload) if existing_payload is not None else {
        "rubric": {"exam_title": "generated", "total_score": 100, "questions": []},
        "answer_key": {"questions": []},
        "meta": {"warnings": []},
    }
    states = _existing_batch_states(merged)
    if retry_question_ids is None:
        targets = list(batches)
    else:
        targets = [
            batch
            for batch in batches
            if set(batch["question_ids"]).issubset(retry_question_ids)
        ]
        targeted_question_ids = {
            qid for batch in targets for qid in batch["question_ids"]
        }
        if targeted_question_ids != retry_question_ids:
            raise ValueError("旧草稿的失败题目不能安全映射到新的分批规则。")
    total_targets = len(targets)
    for completed, batch in enumerate(targets, start=1):
        batch_id = str(batch["batch_id"])
        question_ids = list(batch["question_ids"])
        if report:
            progress = 0.10 + 0.75 * ((completed - 1) / max(1, total_targets))
            report(
                progress,
                "分批生成评分标准",
                f"正在生成批次 {batch_id}（{', '.join(question_ids)}）",
            )
        try:
            prompt, image_blobs = _build_config_batch_request(
                list(batch["blocks"]), doc_text, q_images
            )
            if image_blobs:
                batch_payload = llm_client.json_from_images_once(
                    prompt,
                    image_blobs,
                    model=model_name,
                    extra_kwargs=_config_generation_extra_kwargs(),
                    use_config_client=True,
                )
            else:
                batch_payload = llm_client.json_from_text_once(
                    prompt,
                    model=model_name,
                    extra_kwargs=_config_generation_extra_kwargs(),
                )
            _validate_exact_batch_payload(
                batch_payload,
                question_ids,
                require_canonical_answer=False,
            )
            normalize_new_generated_config_payload(batch_payload)
            _validate_exact_batch_payload(batch_payload, question_ids)
            _replace_retry_question_payloads(
                merged,
                batch_payload,
                set(question_ids),
                question_blocks,
            )
            batch_meta = batch_payload.get("meta") if isinstance(batch_payload, dict) else None
            repair = batch_meta.get("local_json_repair") if isinstance(batch_meta, dict) else None
            state: dict[str, Any] = {
                "batch_id": batch_id,
                "question_ids": question_ids,
                "status": "succeeded",
            }
            if isinstance(repair, dict):
                state["local_json_repair"] = {
                    "repaired": bool(repair.get("repaired")),
                    "operations": [str(item) for item in repair.get("operations") or []],
                    "response_chars": int(repair.get("response_chars") or 0),
                    "response_sha256": str(repair.get("response_sha256") or ""),
                }
            states[batch_id] = state
        except Exception as exc:  # each explicit submission makes exactly one request
            if isinstance(exc, _BatchSchemaMismatch):
                category = "schema_mismatch"
            elif "响应字符数" in str(exc) and "响应摘要" in str(exc):
                category = "invalid_json"
            elif _is_transient_config_generation_error(exc):
                category = "transient_network"
            else:
                category = "model_request"
            states[batch_id] = {
                "batch_id": batch_id,
                "question_ids": question_ids,
                "status": "failed",
                "category": category,
                "error": _safe_batch_failure_message(exc),
            }
        _attach_batch_generation_meta(merged, batches, states, batch_size)
        if checkpoint:
            checkpoint(copy.deepcopy(merged))

    _apply_local_question_facts(merged, question_blocks)
    normalize_new_generated_config_payload(merged)
    _attach_batch_generation_meta(merged, batches, states, batch_size)
    failed = failed_grading_config_batches(merged)
    meta = merged.setdefault("meta", {})
    if failed:
        meta["score_allocation_mode"] = "pending_failed_batches"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = False
        meta["score_allocation_failed"] = False
        meta.pop("score_allocation_error", None)
        meta.pop("score_allocation_failure_category", None)
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    merged = _score_completed_batch_draft_once(
        merged,
        question_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        checkpoint=checkpoint,
    )
    if report:
        report(0.92, "分批生成完成", f"{len(batches)} 个批次和整卷 AI 统一配分均已完成。")
    return merged


def _score_completed_batch_draft_once(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    checkpoint: Callable[[dict[str, Any]], None] | None = None,
) -> dict[str, Any]:
    """Allocate scores once after every generation batch has succeeded."""
    if failed_grading_config_batches(existing_payload):
        raise ValueError("仍有失败批次，不能进行整卷 AI 统一配分。")
    payload = copy.deepcopy(existing_payload)
    meta = payload.setdefault("meta", {})
    if (
        isinstance(meta, dict)
        and bool(meta.get("score_allocation_ai_success"))
        and not bool(meta.get("score_allocation_pending"))
    ):
        return payload
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["score_allocation_mode"] = "dedicated_ai_scoring"
    meta["score_allocation_ai_success"] = False
    meta["score_allocation_pending"] = True
    meta["score_allocation_failed"] = False
    meta.pop("score_allocation_error", None)
    meta.pop("score_allocation_failure_category", None)
    if checkpoint:
        checkpoint(copy.deepcopy(payload))

    structure_summary = _score_allocation_structure_summary(payload)
    question_ids = [str(item.get("question_id") or "") for item in structure_summary]
    prompt = _build_score_allocation_prompt(
        structure_summary,
        doc_text,
        include_document_text=not (
            bool(q_images)
            or any(
                str(block.get("semantic_source") or "").strip() == "images"
                for block in question_blocks
            )
        ),
    )
    prompt = (
        f"SCORE_QUESTION_IDS_JSON={json.dumps(question_ids, ensure_ascii=False)}\n"
        + prompt
    )
    if report:
        report(
            0.88,
            "AI 统一配分",
            f"正在根据 {len(question_ids)} 道题的完整评分步骤统一配置 100 分。",
        )
    score_repair: dict[str, Any] | None = None
    try:
        score_data = llm_client.json_from_text_once(
            prompt,
            model=model_name,
            extra_kwargs=_config_generation_extra_kwargs(),
        )
        score_meta = score_data.get("meta") if isinstance(score_data, dict) else None
        raw_score_repair = (
            score_meta.get("local_json_repair")
            if isinstance(score_meta, dict)
            else None
        )
        if isinstance(raw_score_repair, dict):
            score_repair = {
                "repaired": bool(raw_score_repair.get("repaired")),
                "operations": [
                    str(item) for item in raw_score_repair.get("operations") or []
                ],
                "response_chars": int(raw_score_repair.get("response_chars") or 0),
                "response_sha256": str(
                    raw_score_repair.get("response_sha256") or ""
                ),
            }
        _validate_exact_score_allocation_payload(score_data, structure_summary)
        _apply_score_allocation(payload, score_data)
    except Exception as exc:
        meta["score_allocation_failed"] = True
        meta["score_allocation_failure_category"] = (
            "transient_network"
            if _is_transient_config_generation_error(exc)
            else "model_request"
        )
        meta["score_allocation_error"] = _safe_score_allocation_failure_message(exc)
        if report:
            report(
                0.91,
                "AI 统一配分失败",
                "评分标准批次已保存在本机；没有自动重试，也没有使用本地分值替代。",
            )
        if checkpoint:
            checkpoint(copy.deepcopy(payload))
        return payload
    payload = finalize_completed_grading_config_draft(
        payload,
        question_blocks,
        q_images=q_images,
    )
    meta = payload.setdefault("meta", {})
    meta["score_allocation_mode"] = "dedicated_ai_scoring"
    meta["score_allocation_ai_success"] = True
    meta["score_allocation_pending"] = False
    meta["score_allocation_failed"] = False
    meta.pop("score_allocation_error", None)
    meta.pop("score_allocation_failure_category", None)
    if score_repair is not None:
        meta["score_allocation_local_json_repair"] = score_repair
    else:
        meta.pop("score_allocation_local_json_repair", None)
    if checkpoint:
        checkpoint(copy.deepcopy(payload))
    return payload


def _safe_score_allocation_failure_message(exc: Exception) -> str:
    status_code = getattr(exc, "status_code", None)
    if isinstance(status_code, int):
        return f"AI 统一配分失败（HTTP {status_code}），未自动重试。"
    return "AI 统一配分或本地校验失败，未自动重试。"








def finalize_completed_grading_config_draft(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    *,
    q_images: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Finish an all-succeeded checkpoint locally without another model request."""
    payload = copy.deepcopy(existing_payload)
    if failed_grading_config_batches(payload):
        raise ValueError("仍有失败批次，不能完成本地发布。")
    _validate_unique_batch_question_ids(question_blocks)
    _apply_local_question_facts(payload, question_blocks)
    normalize_new_generated_config_payload(payload)
    expected_ids = [str(block.get("question_id") or "").strip() for block in question_blocks]
    _validate_exact_batch_payload(payload, expected_ids)
    force_payload_total_score(payload, target_total=100.0)
    meta = payload.setdefault("meta", {})
    meta["score_allocation_mode"] = "local_normalization"
    meta["score_allocation_ai_success"] = False
    meta["score_allocation_pending"] = False
    _attach_reference_answer_images(payload, q_images)
    refresh_generated_config_quality_warnings(payload)
    validate_generated_config(payload)
    return payload


def _validate_unique_batch_question_ids(question_blocks: list[dict[str, Any]]) -> None:
    ids = [str(block.get("question_id") or "").strip() for block in question_blocks]
    if any(not qid for qid in ids) or len(set(ids)) != len(ids):
        raise ValueError("已确认题目必须包含唯一且非空的题号。")


def _existing_batch_states(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    meta = payload.get("meta") if isinstance(payload, dict) else None
    raw = meta.get("batches") if isinstance(meta, dict) else None
    return {
        str(item.get("batch_id")): copy.deepcopy(item)
        for item in raw or []
        if isinstance(item, dict) and str(item.get("batch_id") or "").strip()
    }


def _attach_batch_generation_meta(
    payload: dict[str, Any],
    batches: list[dict[str, Any]],
    states: dict[str, dict[str, Any]],
    batch_size: int,
) -> None:
    ordered_states: list[dict[str, Any]] = []
    for batch in batches:
        batch_id = str(batch["batch_id"])
        ordered_states.append(
            copy.deepcopy(
                states.get(batch_id)
                or {
                    "batch_id": batch_id,
                    "question_ids": list(batch["question_ids"]),
                    "status": "pending",
                }
            )
        )
    failures = [
        {
            "batch_id": str(item["batch_id"]),
            "question_ids": list(item["question_ids"]),
            "category": str(item.get("category") or "failed"),
            "error": str(item.get("error") or "批次生成失败"),
        }
        for item in ordered_states
        if item.get("status") in {"failed", "pending"}
    ]
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    meta["generation_mode"] = "batched"
    meta["batch_size"] = int(batch_size)
    meta["batch_count"] = len(batches)
    meta["batches"] = ordered_states
    meta["failed_batches"] = failures
    meta["failed_question_ids"] = [
        qid for item in failures for qid in item["question_ids"]
    ]


def _validate_exact_batch_payload(
    payload: dict[str, Any],
    question_ids: list[str],
    *,
    require_canonical_answer: bool = True,
) -> None:
    if not isinstance(payload, dict):
        raise _BatchSchemaMismatch("批次结果顶层不是 JSON 对象。")
    rubric = payload.get("rubric")
    answer_key = payload.get("answer_key")
    questions = rubric.get("questions") if isinstance(rubric, dict) else None
    answers = answer_key.get("questions") if isinstance(answer_key, dict) else None
    if not isinstance(questions, list) or not isinstance(answers, list):
        raise _BatchSchemaMismatch("批次结果缺少 rubric.questions 或 answer_key.questions。")
    rubric_ids = [
        str(item.get("question_id") or "").strip()
        for item in questions if isinstance(item, dict)
    ]
    answer_ids = [
        str(item.get("question_id") or "").strip()
        for item in answers if isinstance(item, dict)
    ]
    if (
        len(rubric_ids) != len(questions)
        or len(answer_ids) != len(answers)
        or rubric_ids != question_ids
        or answer_ids != question_ids
        or len(set(rubric_ids)) != len(rubric_ids)
    ):
        raise _BatchSchemaMismatch("批次结果题号缺失、重复、越界或顺序不一致。")
    for question in questions:
        if not isinstance(question.get("parts"), list) or not question["parts"]:
            raise _BatchSchemaMismatch("批次评分结构不完整。")
    for answer in answers:
        if require_canonical_answer and "canonical_answer" not in answer:
            raise _BatchSchemaMismatch("批次答案结构不完整。")


def _safe_batch_failure_message(exc: Exception) -> str:
    message = str(exc or "").strip()
    if isinstance(exc, _BatchSchemaMismatch):
        return message[:300]
    if "响应字符数" in message and "响应摘要" in message:
        return message[:500]
    status_code = getattr(exc, "status_code", None)
    if isinstance(status_code, int):
        return f"模型请求失败（HTTP {status_code}），未自动重试。"
    return "模型请求或本地解析失败，未自动重试。"


def _build_config_batch_request(
    blocks: list[dict[str, Any]],
    doc_text: str,
    q_images: dict[str, Any] | None,
) -> tuple[str, list[bytes]]:
    import base64

    question_ids = [str(block.get("question_id") or "").strip() for block in blocks]
    contexts: list[dict[str, Any]] = []
    image_blobs: list[bytes] = []
    image_map: list[str] = []
    for block in blocks:
        qid = str(block.get("question_id") or "").strip()
        image_semantic_source = str(block.get("semantic_source") or "").strip() == "images"
        contexts.append(
            {
                "question_id": qid,
                "question_type": str(block.get("question_type") or ""),
                "question_text": "" if image_semantic_source else str(
                    block.get("text") or block.get("question_text") or ""
                ),
                "answer_text": "" if image_semantic_source else str(
                    block.get("answer_text") or block.get("canonical_answer") or ""
                ),
                "analysis": "" if image_semantic_source else str(block.get("analysis") or ""),
            }
        )
        image_data = q_images.get(qid) if isinstance(q_images, dict) else None
        if isinstance(image_data, dict) and image_data.get("question"):
            image_blobs.append(base64.b64decode(str(image_data["question"]), validate=True))
            image_map.append(f"图片{len(image_blobs)}={qid}题目")
            if image_data.get("answer"):
                image_blobs.append(base64.b64decode(str(image_data["answer"]), validate=True))
                image_map.append(f"图片{len(image_blobs)}={qid}答案解析")
        elif not image_semantic_source:
            for blob in _word_block_image_blobs(block, limit=4):
                image_blobs.append(blob)
                image_map.append(f"图片{len(image_blobs)}={qid}内嵌图")
    fallback = str(doc_text or "")[:4000] if any(
        str(block.get("semantic_source") or "").strip() != "images"
        and not item["answer_text"] and not item["analysis"]
        for block, item in zip(blocks, contexts)
    ) else ""
    prompt = (
        "你正在为一个小批次的中学数学题生成可执行评分标准。仅返回一个完整、严格的 JSON 对象。\n"
        f"BATCH_QUESTION_IDS_JSON={json.dumps(question_ids, ensure_ascii=False)}\n"
        "必须完整且仅返回上述题号，并在 rubric.questions 与 answer_key.questions 中按同一顺序各出现一次；"
        "不得遗漏、重复或增加题号。所有 max_score、part_score、step_score 暂设为1，最终总分由本地程序分配。\n"
        "每题必须保留教师确认题型并输出完整 parts、steps；不要输出任何 knowledge 字段；"
        "每个 part 必须有 response_mode，每个 step 必须有可核验的 core_goal 和 required_elements。"
        "选择/普通填空只核对最终答案；过程题保留必要过程；作图题输出 visual_requirements。\n"
        "只允许字段 rubric、answer_key、meta 及其既有评分结构；不要 Markdown、解释或续写建议。\n"
        f"图片顺序：{'; '.join(image_map) if image_map else '无'}\n"
        f"本批次结构化内容：{json.dumps(contexts, ensure_ascii=False)}\n"
        f"必要时参考的有限原文：{fallback}"
    )
    return prompt, image_blobs


def failed_grading_config_question_ids(payload: dict[str, Any]) -> list[str]:
    meta = payload.get("meta") if isinstance(payload, dict) else {}
    if not isinstance(meta, dict):
        return []
    raw_ids = meta.get("failed_question_ids")
    if isinstance(raw_ids, list):
        return list(dict.fromkeys(str(qid).strip() for qid in raw_ids if str(qid).strip()))
    failures = meta.get("failed_questions")
    if not isinstance(failures, list):
        return []
    return list(
        dict.fromkeys(
            str(failure.get("question_id") or "").strip()
            for failure in failures
            if isinstance(failure, dict) and str(failure.get("question_id") or "").strip()
        )
    )


def _replace_retry_question_payloads(
    existing_payload: dict[str, Any],
    retry_payload: dict[str, Any],
    successful_qids: set[str],
    question_blocks: list[dict[str, Any]],
) -> None:
    existing_rubric = existing_payload.setdefault("rubric", {})
    existing_answer_key = existing_payload.setdefault("answer_key", {})
    retry_rubric = retry_payload.get("rubric", {})
    retry_answer_key = retry_payload.get("answer_key", {})
    existing_questions = existing_rubric.setdefault("questions", []) if isinstance(existing_rubric, dict) else []
    existing_answers = existing_answer_key.setdefault("questions", []) if isinstance(existing_answer_key, dict) else []
    retry_questions = retry_rubric.get("questions", []) if isinstance(retry_rubric, dict) else []
    retry_answers = retry_answer_key.get("questions", []) if isinstance(retry_answer_key, dict) else []
    if not isinstance(existing_questions, list) or not isinstance(existing_answers, list):
        return

    retry_question_map = {
        str(question.get("question_id") or "").strip(): question
        for question in retry_questions
        if isinstance(question, dict)
    } if isinstance(retry_questions, list) else {}
    retry_answer_map = {
        str(answer.get("question_id") or "").strip(): answer
        for answer in retry_answers
        if isinstance(answer, dict)
    } if isinstance(retry_answers, list) else {}

    existing_questions[:] = [
        question
        for question in existing_questions
        if not isinstance(question, dict) or str(question.get("question_id") or "").strip() not in successful_qids
    ]
    existing_answers[:] = [
        answer
        for answer in existing_answers
        if not isinstance(answer, dict) or str(answer.get("question_id") or "").strip() not in successful_qids
    ]
    for qid in successful_qids:
        if qid in retry_question_map:
            existing_questions.append(copy.deepcopy(retry_question_map[qid]))
        if qid in retry_answer_map:
            existing_answers.append(copy.deepcopy(retry_answer_map[qid]))

    qid_order = {
        str(block.get("question_id") or "").strip(): index
        for index, block in enumerate(question_blocks)
        if str(block.get("question_id") or "").strip()
    }
    existing_questions.sort(
        key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000)
        if isinstance(item, dict) else 10_001
    )
    existing_answers.sort(
        key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000)
        if isinstance(item, dict) else 10_001
    )


def _clear_resolved_generation_warnings(payload: dict[str, Any], resolved_qids: set[str]) -> None:
    meta = payload.get("meta") if isinstance(payload, dict) else {}
    warnings = meta.get("warnings") if isinstance(meta, dict) else None
    if not isinstance(warnings, list):
        return
    generation_markers = (
        "parallel generation failed",
        "placeholder added",
        "unmergeable single-question schema",
    )
    warnings[:] = [
        warning
        for warning in warnings
        if not (
            any(marker in str(warning) for marker in generation_markers)
            and any(qid in str(warning) for qid in resolved_qids)
        )
        and str(warning) != "parallel generation completed with failed question blocks; score allocation paused until retry."
    ]


def retry_failed_grading_config_questions(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    retry_question_ids: Sequence[str] | None = None,
) -> dict[str, Any]:
    all_failed_qids = failed_grading_config_question_ids(existing_payload)
    if not all_failed_qids:
        return copy.deepcopy(existing_payload)

    if retry_question_ids is None:
        failed_qids = all_failed_qids
    else:
        failed_qids = list(
            dict.fromkeys(str(qid).strip() for qid in retry_question_ids if str(qid).strip())
        )
        if not failed_qids:
            raise ValueError("至少选择一道失败题目进行重试。")
        unknown_qids = [qid for qid in failed_qids if qid not in all_failed_qids]
        if unknown_qids:
            raise ValueError("只能重试当前失败题目：" + ", ".join(unknown_qids))

    locked_blocks = [{**block, "question_type_confirmed": True} for block in question_blocks]
    block_map = {
        str(block.get("question_id") or "").strip(): block
        for block in locked_blocks
        if str(block.get("question_id") or "").strip()
    }
    retry_blocks = [block_map[qid] for qid in failed_qids if qid in block_map]
    missing_failures = [
        {
            "question_id": qid,
            "attempts": 0,
            "category": "missing_retry_input",
            "error": "Original confirmed question block is no longer available",
        }
        for qid in failed_qids
        if qid not in block_map
    ]
    if not retry_blocks:
        raise ValueError("失败题目的原始确认数据已丢失，请重新拆分试卷后生成。")

    _validate_image_semantic_inputs(retry_blocks, q_images)
    results, failures, attempt_counts, worker_count, image_semantic_mode = _generate_question_block_results(
        retry_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
    )
    retry_payload = _merge_single_question_payloads(results, retry_blocks)
    failures = _append_unmergeable_question_failures(retry_payload, retry_blocks, failures, attempt_counts)
    failures.extend(missing_failures)
    selected_failure_map = {
        str(failure.get("question_id") or "").strip(): failure
        for failure in failures
        if str(failure.get("question_id") or "").strip()
    }
    existing_meta = existing_payload.get("meta") if isinstance(existing_payload, dict) else {}
    existing_failures = existing_meta.get("failed_questions") if isinstance(existing_meta, dict) else []
    existing_failure_map = {
        str(failure.get("question_id") or "").strip(): copy.deepcopy(failure)
        for failure in existing_failures
        if isinstance(failure, dict) and str(failure.get("question_id") or "").strip()
    } if isinstance(existing_failures, list) else {}
    failures = []
    for qid in all_failed_qids:
        if qid in selected_failure_map:
            failures.append(selected_failure_map[qid])
        elif qid not in failed_qids:
            failures.append(
                existing_failure_map.get(qid)
                or {
                    "question_id": qid,
                    "attempts": 0,
                    "category": "retry_pending",
                    "error": "Question has not been selected for retry",
                }
            )
    remaining_failed_qids = {
        str(failure.get("question_id") or "").strip()
        for failure in failures
        if str(failure.get("question_id") or "").strip()
    }
    successful_qids = set(failed_qids) - remaining_failed_qids

    merged = copy.deepcopy(existing_payload)
    _replace_retry_question_payloads(merged, retry_payload, successful_qids, locked_blocks)
    _clear_resolved_generation_warnings(merged, successful_qids)
    previous_attempts = merged.get("meta", {}).get("single_question_attempts", {})
    all_attempts = dict(previous_attempts) if isinstance(previous_attempts, dict) else {}
    all_attempts.update(attempt_counts)
    _attach_parallel_generation_meta(merged, locked_blocks, failures, worker_count, all_attempts)
    _ensure_question_blocks_covered(merged, locked_blocks)
    _apply_local_question_facts(merged, locked_blocks)
    normalize_new_generated_config_payload(merged)

    if failures:
        meta = merged.setdefault("meta", {})
        meta["score_allocation_mode"] = "pending_failed_questions"
        meta["score_allocation_ai_success"] = False
        meta["score_allocation_pending"] = True
        _attach_reference_answer_images(merged, q_images)
        refresh_generated_config_quality_warnings(merged)
        return merged

    return retry_grading_config_score_allocation(
        merged,
        locked_blocks,
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        q_images=q_images,
        include_document_text=not image_semantic_mode,
    )


def retry_grading_config_score_allocation(
    existing_payload: dict[str, Any],
    question_blocks: list[dict[str, Any]],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    q_images: dict[str, Any] | None = None,
    *,
    include_document_text: bool | None = None,
) -> dict[str, Any]:
    if failed_grading_config_question_ids(existing_payload):
        raise ValueError("仍有失败题目，需先重试失败题目，再进行整体赋分。")
    merged = copy.deepcopy(existing_payload)
    if include_document_text is None:
        include_document_text = not (
            bool(q_images)
            or any(str(block.get("semantic_source") or "").strip() == "images" for block in question_blocks)
        )
    _assign_scores_to_merged_config(
        merged,
        doc_text,
        llm_client,
        model_name,
        report=report,
        include_document_text=include_document_text,
    )
    _ensure_question_blocks_covered(merged, question_blocks)
    _apply_local_question_facts(merged, question_blocks)
    normalize_new_generated_config_payload(merged)
    meta = merged.setdefault("meta", {})
    meta["score_allocation_mode"] = (
        "dedicated_ai_scoring" if meta.get("score_allocation_ai_success") else "local_score_fallback"
    )
    meta["score_allocation_pending"] = False
    force_payload_total_score(merged, target_total=100.0)
    _attach_reference_answer_images(merged, q_images)
    refresh_generated_config_quality_warnings(merged)
    return merged


def _split_local_question_answer_text(doc_text: str) -> tuple[str, str]:
    text = str(doc_text or "")
    try:
        from question_bank.importers.batch_importer import _split_answer_text

        question_text, answer_text = _split_answer_text(text)
        return str(question_text or ""), str(answer_text or "")
    except Exception:
        lines = text.splitlines()
        for index, line in enumerate(lines):
            if _looks_like_answer_section_heading(line):
                return "\n".join(lines[:index]).strip(), "\n".join(lines[index + 1 :]).strip()
    return text, ""


def _question_type_hints_from_section_headings(doc_text: str) -> dict[str, str]:
    question_text, _ = _split_local_question_answer_text(doc_text)
    current_type = ""
    hints: dict[str, str] = {}
    for line in question_text.splitlines():
        value = str(line or "").strip()
        if "选择题" in value:
            current_type = "choice"
            continue
        if "填空题" in value:
            current_type = "fill_blank"
            continue
        if "证明题" in value:
            current_type = "proof"
            continue
        if "解答题" in value:
            current_type = "comprehensive"
            continue
        number = _extract_question_marker_number(value)
        if number is not None and current_type:
            hints[str(number)] = current_type
    return hints


def _parse_inline_answer_blocks(doc_text: str) -> list[dict[str, Any]] | None:
    """解析"题干 + 【答案】 + 【解析】逐题穿插"格式的试卷（纯本地，不调 AI）。

    返回 None 表示该文本不是内联格式（无【答案】/【解析】标记），交回上层兜底逻辑。
    通过"题号必须单调递增"截断文件后半段重复的答案详解区，避免重复拆题。
    """
    full_text = str(doc_text or "")
    text, answer_section = _split_local_question_answer_text(full_text)
    if ("【答案】" not in text) and ("【解析】" not in text):
        return None

    lines = text.splitlines()
    # 1. 找题号边界：行首 数字. / 数字．，题号单调递增；回退即停止
    markers: list[tuple[int, int]] = []  # (line_index, number)
    last_number = 0
    for idx, line in enumerate(lines):
        stripped = line.strip()
        m = re.match(r"^(\d{1,2})\s*[.．]", stripped)
        if not m:
            continue
        number = int(m.group(1))
        if not (1 <= number <= 99):
            continue
        if markers and number <= last_number:
            break  # 题号回退 → 进入答案详解重复区，停止切题
        markers.append((idx, number))
        last_number = number

    if not markers:
        return None

    answer_section_for_choice = answer_section or _local_answer_section_text(text)
    choice_answers = _extract_choice_answer_sequence(answer_section_for_choice)
    section_hints = _question_type_hints_from_section_headings(text)

    blocks: list[dict[str, Any]] = []
    for pos, (start, number) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        segment_lines = lines[start:end]
        parsed = _parse_inline_segment(
            number,
            segment_lines,
            choice_answers,
            section_type=section_hints.get(str(number), ""),
        )
        if parsed:
            blocks.append(parsed)

    return blocks or None


def _parse_inline_segment(
    number: int,
    segment_lines: list[str],
    choice_answers: dict[str, str],
    *,
    section_type: str = "",
) -> dict[str, Any] | None:
    """把单道题的文本段拆成 题干 / 答案 / 解析。"""
    # 分类收集：题干 / 【答案】 / 【解析】(含【点睛】)
    bucket = "stem"  # stem | answer | analysis
    stem_lines: list[str] = []
    answer_lines: list[str] = []
    analysis_lines: list[str] = []

    for raw in segment_lines:
        line = str(raw or "").strip()
        if not line:
            continue
        if line.startswith("[公式:") or line.startswith("[公式："):
            continue  # 忽略公式占位行
        if "【答案】" in line:
            bucket = "answer"
            after = line.split("【答案】", 1)[1].strip()
            if after:
                answer_lines.append(after)
            continue
        if ("【解析】" in line) or ("【点睛】" in line):
            bucket = "analysis"
            marker = "【解析】" if "【解析】" in line else "【点睛】"
            after = line.split(marker, 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if line.startswith("【") and "】" in line:
            # 其他【…】标记（如【难度】），归入解析区尾部，避免污染题干
            bucket = "analysis"
            after = line.split("】", 1)[1].strip()
            if after:
                analysis_lines.append(after)
            continue
        if bucket == "stem":
            stem_lines.append(line)
        elif bucket == "answer":
            answer_lines.append(line)
        else:
            analysis_lines.append(line)

    question_text = _strip_leading_question_number(number, "\n".join(stem_lines).strip())
    answer_raw = "\n".join(answer_lines).strip()
    analysis_raw = "\n".join(analysis_lines).strip()

    if not question_text and not answer_raw and not analysis_raw:
        return None

    num_str = str(number)
    qtype = _infer_local_question_type(
        question_text,
        answer_raw,
        num_str,
        section_type=section_type,
    )
    canonical = _extract_canonical_answer_for_local_question(
        number=num_str,
        qtype=qtype,
        question_text=question_text,
        answer_text=answer_raw,
        choice_answers=choice_answers,
    )
    accepted = _local_accepted_forms(canonical, qtype)
    has_answer = bool(answer_raw or canonical)
    return {
        "question_id": f"Q{number}",
        "text": question_text,
        "question_text": question_text,
        "answer_text": answer_raw,
        "analysis": analysis_raw,
        "question_type": qtype,
        "canonical_answer": canonical,
        "accepted_forms": accepted,
        "local_answer_trusted": has_answer,
        "needs_review": (not bool(canonical)) if qtype in {"choice", "fill_blank"} else False,
    }


def _strip_leading_question_number(number: int, text: str) -> str:
    """去掉题干开头的"N." / "N．" / "N.(本小题X分)"等题号前缀。"""
    value = str(text or "").lstrip()
    value = re.sub(rf"^{number}\s*[.．]\s*", "", value, count=1)
    return value.strip()


def _extract_local_question_blocks(doc_text: str) -> list[dict[str, Any]]:
    normalized_doc_text = "\n".join(
        _normalize_inline_main_question_lines(doc_text)
    )
    inline_blocks = _parse_inline_answer_blocks(normalized_doc_text)
    if inline_blocks:
        return inline_blocks

    parsed_questions: list[Any] = []
    try:
        from question_bank.importers.batch_importer import parse_paper_text

        parsed = parse_paper_text(
            normalized_doc_text,
            source_file="grading_config.docx",
            page_range="document",
        )
        parsed_questions = list(parsed.questions)
    except Exception:
        parsed_questions = []

    # 第二道防线：题号必须单调递增；一旦回退（如卷末答案详解区从1重列）即截断，避免重复拆题。
    if parsed_questions:
        truncated: list[Any] = []
        last_number = 0
        for item in parsed_questions:
            number_str = str(getattr(item, "question_number", "") or "").strip()
            if not number_str.isdigit():
                truncated.append(item)
                continue
            number = int(number_str)
            if truncated and number <= last_number:
                break
            truncated.append(item)
            last_number = number
        parsed_questions = truncated

    answer_section = _local_answer_section_text(normalized_doc_text)
    answer_blocks = _local_answer_blocks(answer_section)
    choice_answers = _extract_choice_answer_sequence(answer_section)
    section_hints = _question_type_hints_from_section_headings(normalized_doc_text)

    blocks: list[dict[str, Any]] = []
    if parsed_questions:
        for item in parsed_questions:
            number = str(getattr(item, "question_number", "") or "").strip()
            if not number.isdigit():
                continue
            qid = f"Q{int(number)}"
            question_text = str(getattr(item, "question_text", "") or "").strip()
            answer_from_map = str(answer_blocks.get(number, "") or "").strip()
            parsed_answer = str(getattr(item, "answer_text", "") or "").strip()
            raw_answer = str(answer_from_map or parsed_answer).strip()
            qtype = _infer_local_question_type(
                question_text,
                raw_answer,
                number,
                section_type=section_hints.get(str(int(number)), ""),
            )
            canonical = _extract_canonical_answer_for_local_question(
                number=number,
                qtype=qtype,
                question_text=question_text,
                answer_text=raw_answer or answer_blocks.get(number, ""),
                choice_answers=choice_answers,
            )
            accepted = _local_accepted_forms(canonical, qtype)
            blocks.append(
                {
                    "question_id": qid,
                    "text": question_text,
                    "question_text": question_text,
                    "answer_text": raw_answer,
                    "question_type": qtype,
                    "canonical_answer": canonical,
                    "accepted_forms": accepted,
                    "local_answer_trusted": bool(answer_from_map)
                    or any(token in raw_answer for token in ["故答案为", "故选", "答案为"]),
                    "needs_review": not bool(canonical) if qtype in {"choice", "fill_blank"} else False,
                }
            )
    else:
        for block in _split_doc_text_into_question_blocks(normalized_doc_text):
            qid = str(block.get("question_id") or "")
            number = qid.removeprefix("Q")
            question_text = str(block.get("text") or "")
            raw_answer = answer_blocks.get(number, "")
            qtype = _infer_local_question_type(
                question_text,
                raw_answer,
                number,
                section_type=section_hints.get(str(int(number)), "") if number.isdigit() else "",
            )
            canonical = _extract_canonical_answer_for_local_question(
                number=number,
                qtype=qtype,
                question_text=question_text,
                answer_text=raw_answer,
                choice_answers=choice_answers,
            )
            block.update(
                {
                    "question_text": question_text,
                    "answer_text": raw_answer,
                    "question_type": qtype,
                    "canonical_answer": canonical,
                    "accepted_forms": _local_accepted_forms(canonical, qtype),
                    "local_answer_trusted": bool(raw_answer)
                    or any(token in raw_answer for token in ["故答案为", "故选", "答案为"]),
                    "needs_review": not bool(canonical) if qtype in {"choice", "fill_blank"} else False,
                }
            )
            blocks.append(block)

    return blocks


def _local_answer_section_text(doc_text: str) -> str:
    text = str(doc_text or "")
    try:
        from question_bank.importers.batch_importer import _split_answer_text

        _, answer_text = _split_answer_text(text)
        if str(answer_text or "").strip():
            return str(answer_text or "")
    except Exception:
        pass

    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if _looks_like_answer_section_heading(line):
            return "\n".join(lines[idx + 1 :]).strip()
    return ""


def _local_answer_blocks(answer_section: str) -> dict[str, str]:
    section = str(answer_section or "")
    if not section.strip():
        return {}
    try:
        from question_bank.importers.batch_importer import _split_numbered_blocks

        blocks = _split_numbered_blocks(section)
        result: dict[str, str] = {}
        for block in blocks:
            number = str(getattr(block, "number", "") or "").strip()
            if number.isdigit() and 1 <= int(number) <= 99:
                result.setdefault(str(int(number)), str(getattr(block, "text", "") or "").strip())
        if result:
            return result
    except Exception:
        pass

    lines = section.splitlines()
    markers: list[tuple[int, str]] = []
    last_number = 0
    for idx, line in enumerate(lines):
        if "声明" in line or "菁优网" in line:
            break
        number = _extract_question_marker_number(line)
        if number is None:
            continue
        if markers and number <= last_number:
            break
        markers.append((idx, str(number)))
        last_number = number

    result: dict[str, str] = {}
    for pos, (start, number) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        block = "\n".join(lines[start:end]).strip()
        if block:
            result[number] = block
    return result


def _extract_choice_answer_sequence(answer_section: str) -> dict[str, str]:
    """Extract choice answers (A/B/C/D) from the answer section in order.

    Supports common Chinese exam answer formats:
    - 故选A / 故选：A
    - 选A / 选：A
    - 答案是A / 答案为A / 答：A
    - 答案选A
    - 独立的单字母行（上下文暗示选择题）
    """
    section = str(answer_section or "")
    found: list[str] = []
    # Broad pattern: any of the common phrasings followed by A/B/C/D
    pattern = (
        r"(?:"
        r"故\s*选"
        r"|选择"
        r"|选\s*[:：]"
        r"|答\s*案\s*(?:是|为|选)?\s*[:：]?"
        r"|答\s*[:：]"
        r"|【答案】"
        r")\s*([A-Da-d])"
    )
    for match in re.finditer(pattern, section):
        found.append(match.group(1).upper())
    if not found:
        # Fallback: numbered answers like "1.A" or "1、A" or "(1)A"
        for match in re.finditer(
            r"(?:^|\n)\s*(?:\(\s*)?\d{1,2}\s*[.．、)）]\s*([A-Da-d])(?:\s|$)", section
        ):
            found.append(match.group(1).upper())
    if not found:
        # Last fallback: solo letter on its own line that looks like a choice answer
        for match in re.finditer(r"(?:^|\n)\s*([A-D])\s*(?:\n|$)", section):
            found.append(match.group(1).upper())
    return {str(index): value for index, value in enumerate(found[:20], start=1)}


def _infer_local_question_type(
    question_text: str,
    answer_text: str,
    number: str,  # noqa: ARG001
    *,
    section_type: str = "",
) -> str:
    """Infer question type, preferring explicit paper-section headings."""
    value = str(question_text or "")
    normalized_section_type = str(section_type or "").strip()
    if normalized_section_type in {"choice", "fill_blank", "proof"}:
        return normalized_section_type

    option_labels = {
        match.group(1).upper()
        for match in re.finditer(r"(?m)^\s*([A-Da-d])\s*(?:[.．、)]|\s{2,})", value)
    }
    compact_options = any(
        re.search(r"[Aa][.．、]?\s*.{0,80}[Bb][.．、]?\s*.{0,80}[Cc][.．、]?\s*.{0,80}[Dd]", line)
        for line in value.splitlines()
    )
    if len(option_labels) >= 3 or compact_options:
        return "choice"
    if re.fullmatch(r"\s*[A-Da-d]\s*", str(answer_text or "")):
        return "choice"

    has_blank = bool(
        re.search(
            r"_{2,}|　{1,}|（\s*）|\(\s*\)|\b填空\b",
            value,
        )
    )
    if has_blank:
        return "fill_blank"

    if any(token in value for token in ["作图", "作出", "画出", "保留作图痕迹"]):
        return "comprehensive"
    if any(token in value for token in ["证明", "理由", "说明", "求证", "全等", "证得"]):
        return "proof"
    has_subparts = bool(re.search(r"[（(]\s*[1-9]\s*[）)]", value))
    if has_subparts or any(token in value for token in ["计算", "求", "解答", "解："]):
        return "calculation"
    if normalized_section_type == "comprehensive":
        return "comprehensive"
    return _infer_question_type_from_block_text(question_text)


def _extract_canonical_answer_for_local_question(
    *,
    number: str,
    qtype: str,
    question_text: str,
    answer_text: str,
    choice_answers: dict[str, str],
) -> str:
    if qtype == "choice":
        # First try the pre-extracted sequence (ordered by question number)
        answer = choice_answers.get(str(int(number))) if str(number or "").isdigit() else ""
        if answer:
            return answer
        # Then try in-line answer patterns
        choice_pattern = (
            r"(?:故\s*选|选择|选\s*[:：]|答\s*案\s*(?:是|为|选)?\s*[:：]?|答\s*[:：]|【答案】)\s*([A-Da-d])"
        )
        match = re.search(choice_pattern, str(answer_text or ""))
        if match:
            return match.group(1).upper()
        # Last: bare single letter
        bare = re.fullmatch(r"\s*([A-Da-d])\s*", str(answer_text or ""))
        return bare.group(1).upper() if bare else ""

    if qtype != "fill_blank":
        return ""

    text = str(answer_text or "")
    # Extended patterns for fill-blank answer extraction
    patterns = [
        r"【答案】\s*([^。\n；;]{1,120})",
        r"故答案为\s*[:：]?\s*([^。\n；;]{1,120})",
        r"答案(?:是|为)\s*[:：]?\s*([^。\n；;]{1,120})",
        r"答案\s*[:：]\s*([^。\n；;]{1,120})",
        r"答\s*[:：]\s*([^。\n；;]{1,120})",
        r"故答案为\s*[:：]?\s*([\s\S]{1,80}?)[。．\n]",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            value = _clean_local_answer_text(match.group(1))
            if value:
                return value

    blank_values = re.findall(r"\u3000\s*([^\u3000\n]{1,40}?)\s*\u3000", text)
    for value in blank_values:
        cleaned = _clean_local_answer_text(value)
        if cleaned:
            return cleaned
    return ""


def _clean_local_answer_text(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^[：:，,、;\s]+|[。．，,、;\s]+$", "", text)
    text = text.replace("°", "").strip() if re.fullmatch(r"[\d.]+°", text) else text
    text = re.sub(r"\s+", " ", text).strip()
    if not text or text in {"解", "如下", "见解析"}:
        return ""
    return text[:120]


def _local_accepted_forms(canonical: str, qtype: str) -> list[str]:
    value = str(canonical or "").strip()
    if not value:
        return []
    if qtype == "choice":
        return [value.upper()]
    candidates = [value]
    for chunk in re.split(r"[；;、，,]|\s+或\s+|\bor\b", value):
        cleaned = _clean_local_answer_text(chunk)
        if cleaned:
            candidates.append(cleaned)
    return merge_equivalent_forms([], *candidates, max_forms=16)




















def _restore_precalibration_objective_answers(payload: dict[str, Any], previous_payload: dict[str, Any]) -> None:
    if not isinstance(payload, dict) or not isinstance(previous_payload, dict):
        return
    previous_questions = previous_payload.get("rubric", {}).get("questions", [])
    previous_answers = previous_payload.get("answer_key", {}).get("questions", [])
    current_answers = payload.get("answer_key", {}).get("questions", [])
    if not isinstance(previous_questions, list) or not isinstance(previous_answers, list) or not isinstance(current_answers, list):
        return

    previous_types = {
        str(question.get("question_id") or "").strip(): str(question.get("question_type") or "").strip()
        for question in previous_questions
        if isinstance(question, dict)
    }
    previous_answer_map = {
        str(answer.get("question_id") or "").strip(): answer
        for answer in previous_answers
        if isinstance(answer, dict)
    }
    for answer in current_answers:
        if not isinstance(answer, dict):
            continue
        qid = str(answer.get("question_id") or "").strip()
        if previous_types.get(qid) not in {"choice", "fill_blank"}:
            continue
        previous = previous_answer_map.get(qid)
        if not isinstance(previous, dict):
            continue
        previous_canonical = str(previous.get("canonical_answer") or "").strip()
        if not previous_canonical:
            continue
        answer["canonical_answer"] = previous_canonical
        answer["accepted_forms"] = merge_equivalent_forms(previous.get("accepted_forms"), previous_canonical, max_forms=32)
        previous_parts = previous.get("parts")
        if isinstance(previous_parts, list) and previous_parts:
            answer["parts"] = previous_parts


def _split_doc_text_into_question_blocks(doc_text: str) -> list[dict[str, str]]:
    lines = _normalize_inline_main_question_lines(doc_text)
    markers: list[tuple[int, str]] = []
    last_number = 0
    for idx, line in enumerate(lines):
        if _looks_like_answer_section_heading(line) and markers:
            break
        number = _extract_question_marker_number(line)
        if number is None:
            continue
        if markers and number <= last_number:
            break
        markers.append((idx, f"Q{number}"))
        last_number = number
    blocks: list[dict[str, str]] = []
    for pos, (start, qid) in enumerate(markers):
        end = markers[pos + 1][0] if pos + 1 < len(markers) else len(lines)
        text = "\n".join(lines[start:end]).strip()
        if text:
            blocks.append({"question_id": qid, "text": text})
    return blocks


def _normalize_inline_main_question_lines(doc_text: str) -> list[str]:
    lines: list[str] = []
    current_number: int | None = None
    in_answer_section = False
    raw_lines = [raw_line.rstrip() for raw_line in str(doc_text or "").splitlines()]
    for index, line in enumerate(raw_lines):
        if _looks_like_answer_section_heading(line) and current_number is not None:
            in_answer_section = True
        if in_answer_section:
            lines.append(line)
            continue
        leading_number = _extract_question_marker_number(line)
        if leading_number is not None:
            current_number = leading_number
        if current_number is None:
            lines.append(line)
            continue
        segments, current_number = _split_consecutive_inline_main_questions(
            line,
            current_number=current_number,
            following_number=_next_leading_main_question_number(
                raw_lines,
                after_index=index,
            ),
        )
        lines.extend(segments)
    return lines


def _extract_question_marker_number(line: str) -> int | None:
    value = str(line or "").strip()
    if not value:
        return None
    patterns = [
        r"^(?:Q|q)\s*(\d{1,2})(?:\b|[\s:：.．、)])",
        r"^第\s*(\d{1,2})\s*[题題]",
        r"^(\d{1,2})\s*[.．、)]",
    ]
    for pattern in patterns:
        match = re.match(pattern, value)
        if not match:
            continue
        number = int(match.group(1))
        if 1 <= number <= 99:
            return number
    return None

def _looks_like_answer_section_heading(line: str) -> bool:
    value = str(line or "").strip()
    if not value:
        return False
    lower = value.lower().strip(":： ")
    chinese_markers = [
        "\u53c2\u8003\u7b54\u6848",
        "\u8bd5\u9898\u89e3\u6790",
        "\u7b54\u6848\u4e0e\u89e3\u6790",
    ]
    if any(marker in value for marker in chinese_markers):
        return True
    return lower in {"answer", "answers", "solution", "solutions", "answer key"}

def _build_single_question_image_generation_prompt(
    block: dict[str, Any],
    *,
    has_answer_image: bool,
) -> str:
    qid = str(block.get("question_id") or "").strip()
    confirmed_type = str(block.get("question_type") or "").strip()
    image_description = (
        "图片1是题目原图；图片2是标准答案与解析原图。"
        if has_answer_image
        else "图片1是题目原图；未提供标准答案与解析原图。"
    )
    return (
        "你正在为一道中学数学题生成可执行评分标准。仅返回严格 JSON。\n"
        f"题目ID：{qid}\n"
        f"教师已确认题型：{confirmed_type}\n"
        f"{image_description}\n"
        "图片是唯一权威内容来源。不得参考、猜测或恢复任何 PDF 抽取文字。\n"
        "必须严格保持图片中的点名、线段、公式、运算对象、小问结构和作答要求，不得改写成相似题。\n"
        "分值字段暂时全部设为1；后续会统一分配分值。\n"
        "每个 parts 项必须输出 response_mode："
        "exact_objective（选择/普通填空）、short_answer_points（直接写出、多空或多个结果，按答对数量给分）、"
        "process_required（证明/计算过程）、visual_construction（作图，以答案图为视觉参考）。\n"
        "题目写有“直接写出”“填空”时，不得强制过程；short_answer_points 的每个可独立得分答案应拆成独立踩分点。\n"
        "选择题和普通填空题的评分点只能描述核对最终答案，不能要求推理或计算过程。\n"
        "若一个填空要求列出全部可能答案，必须在 canonical_answer 写全，并输出 match_mode=complete_set、required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；不要把少写答案列为 accepted_forms。\n"
        "作图题不得把答案图臆造为唯一文字答案；应输出 visual_requirements 和必要踩分点。\n"
        "知识点与题目标签由题库程序单独维护，不要输出任何 knowledge 字段。\n"
        "每个主观题步骤的 required_elements 必须写出可核验的条件、式子或结论，不得留空或只写通用描述。\n"
        "仅为该题输出 rubric.questions 与 answer_key.questions，并保持 question_id 一致。\n"
        "必须严格使用以下字段，不得自行改名或创造 answers、answer_parts、answer_content、desc 等同义字段：\n"
        "{\"rubric\":{\"questions\":[{\"question_id\":\"...\",\"question_type\":\"...\",\"parts\":["
        "{\"part_id\":\"...\",\"response_mode\":\"...\",\"visual_requirements\":[],\"steps\":["
        "{\"step_id\":\"S1\",\"core_goal\":\"具体踩分点描述\",\"required_elements\":[\"可核验的必要条件或结论\"],"
        "\"allow_alternative_methods\":true}]}]}]},\"answer_key\":{\"questions\":["
        "{\"question_id\":\"...\",\"canonical_answer\":\"...\",\"accepted_forms\":[],\"parts\":["
        "{\"part_id\":\"...\",\"answer\":\"...\",\"accepted_forms\":[],\"analysis\":\"\","
        "\"step_milestones\":[]}]}]}}"
    )


def _build_single_question_generation_prompt(block: dict[str, str], doc_text: str) -> str:
    """Build a prompt for a single question's rubric structure.

    IMPORTANT: Scores are NOT requested here — they are allocated in a dedicated
    second AI call (_assign_scores_to_merged_config) after all questions are merged.
    All score fields must be set to 1 as a placeholder.
    """
    qid = str(block.get("question_id") or "").strip()
    q_text = str(block.get("text") or "").strip()
    answer_text = str(block.get("answer_text") or "").strip()
    analysis_text = str(block.get("analysis") or "").strip()
    confirmed_type = str(block.get("question_type") or "").strip()
    type_is_confirmed = bool(block.get("question_type_confirmed"))

    # Build structured per-question context (replaces the old 6000-char full-doc dump)
    context_parts: list[str] = [f"题目文本：\n{q_text}"]
    if type_is_confirmed and confirmed_type:
        context_parts.append(
            f"教师已确认题型：{confirmed_type}\n"
            "该题型是强约束，必须原样写入 question_type，不得自行改成其他题型。"
        )
    if answer_text:
        context_parts.append(f"参考答案：\n{answer_text}")
    if analysis_text:
        context_parts.append(f"解析/证明过程：\n{analysis_text}")
    if not answer_text and not analysis_text:
        # Fallback: include a limited window of the full doc text to help locate the answer
        context_parts.append(
            "（未能提取到配套答案，以下是试卷原文供参考，请自行定位该题答案区域）：\n"
            + str(doc_text or "")[:4000]
        )
    question_context = "\n\n".join(context_parts)

    return (
        "你正在为【单道】初中数学题目生成评分标准大纲结构。请仅返回严格的 JSON 数据。\n"
        f"题目 ID (QUESTION_ID): {qid}\n"
        "重要：请将所有评分相关的分值字段（包括 max_score, part_score, step_score 等）全部设为占位符 1。"
        "具体分值会在后续步骤中统一分配，请勿尝试在此猜测真实分值。\n"
        "规则：仅为该道题构建相应的 rubric.questions 和 answer_key.questions。\n"
        "如果是选择题或填空题，应给出标准答案以及 accepted_forms（等价接受形式）。\n"
        "如果是计算题、证明题、综合题，必须包含 parts、steps、proof_obligations（证明证据点）、"
        "deduction_policy（扣分策略）、证据要求以及仅写出答案的上限得分（answer_only_max_score）。\n"
        "每个 parts 项必须输出 response_mode：exact_objective、short_answer_points、process_required 或 visual_construction；"
        "不得把大题级过程要求无条件继承给“直接写出”或作图小问。\n"
        "选择题和普通填空题的评分点只能描述核对最终答案，不能要求推理或计算过程。\n"
        "若一个填空要求列出全部可能答案，必须在 canonical_answer 写全，并输出 match_mode=complete_set、required_values、"
        "order_sensitive=false、allow_extra_values=false、partial_credit=false；不要把少写答案列为 accepted_forms。\n"
        "如果该题包含多个空格、表格单元格或子小问，必须将其拆分为不同的 parts 以给与步骤/部分分。\n"
        "accepted_forms 必须仅包含在数学上完全等价的答案形式。\n\n"
        "知识点与题目标签由题库程序单独维护，不要输出任何 knowledge 字段。\n"
        "每个主观题步骤的 required_elements 必须写出可核验的条件、式子或结论，不得留空或只写通用描述。\n\n"
        f"{question_context}"
    )


def _merge_single_question_payloads(
    payloads: list[dict[str, Any] | None],
    question_blocks: list[dict[str, str]] | None = None,
) -> dict[str, Any]:
    merged = {
        "rubric": {"exam_title": "generated", "total_score": 100, "questions": []},
        "answer_key": {"questions": []},
        "meta": {"source": "word_docx_parallel", "warnings": []},
    }
    seen_questions: set[str] = set()
    seen_answers: set[str] = set()
    for index, payload in enumerate(payloads):
        if not isinstance(payload, dict):
            continue
        expected_qid = ""
        if question_blocks and index < len(question_blocks):
            expected_qid = str(question_blocks[index].get("question_id") or "").strip()
        _coerce_single_question_payload_schema(payload, expected_qid, merged["meta"]["warnings"])
        normalize_new_generated_config_payload(payload)
        rubric = payload.get("rubric") if isinstance(payload, dict) else {}
        answer_key = payload.get("answer_key") if isinstance(payload, dict) else {}
        for question in rubric.get("questions", []) if isinstance(rubric, dict) else []:
            if not isinstance(question, dict):
                continue
            qid = str(question.get("question_id") or "").strip()
            if not qid or qid in seen_questions:
                continue
            seen_questions.add(qid)
            merged["rubric"]["questions"].append(question)
        for answer in answer_key.get("questions", []) if isinstance(answer_key, dict) else []:
            if not isinstance(answer, dict):
                continue
            qid = str(answer.get("question_id") or "").strip()
            if not qid or qid in seen_answers:
                continue
            seen_answers.add(qid)
            merged["answer_key"]["questions"].append(answer)
        meta = payload.get("meta") if isinstance(payload, dict) else {}
        warnings = meta.get("warnings") if isinstance(meta, dict) else []
        if isinstance(warnings, list):
            merged["meta"]["warnings"].extend(str(item) for item in warnings if str(item).strip())
    return merged


def _coerce_single_question_payload_schema(
    payload: dict[str, Any],
    expected_qid: str,
    warnings: list[str],
) -> None:
    """Accept common one-question schemas returned by LLMs before strict normalization."""
    raw_rubric = payload.get("rubric")
    raw_answer_key = payload.get("answer_key")
    rubric = raw_rubric if isinstance(raw_rubric, dict) else {}
    answer_key = raw_answer_key if isinstance(raw_answer_key, dict) else {}
    payload["rubric"] = rubric
    payload["answer_key"] = answer_key

    raw_questions = rubric.get("questions") if isinstance(raw_rubric, dict) else raw_rubric
    questions = _coerce_single_item_list(raw_questions, _looks_like_question_dict)
    rubric["questions"] = questions
    raw_answers = answer_key.get("questions") if isinstance(raw_answer_key, dict) else raw_answer_key
    answers = _coerce_single_item_list(raw_answers, _looks_like_answer_dict)
    answer_key["questions"] = answers

    question = next((item for item in questions if isinstance(item, dict)), None)
    if question is None:
        question = _extract_single_question_candidate(payload, rubric)
        if isinstance(question, dict):
            questions.append(question)
            if expected_qid:
                warnings.append(f"{expected_qid} normalized singular rubric.question schema")

    answer = next((item for item in answers if isinstance(item, dict)), None)
    if answer is None:
        answer = _extract_single_answer_candidate(payload, answer_key, question)
        if isinstance(answer, dict):
            answers.append(answer)
            if expected_qid:
                warnings.append(f"{expected_qid} normalized singular answer_key.question schema")

    if expected_qid:
        for item in [question, answer]:
            if isinstance(item, dict):
                raw_qid = item.get("question_id") or item.get("id") or item.get("number")
                item["question_id"] = _canonical_question_id(raw_qid, expected_qid)

    if question is None and expected_qid:
        keys = ", ".join(sorted(str(key) for key in payload.keys()))
        warnings.append(f"{expected_qid} unmergeable single-question schema; top-level keys: {keys}")


def _coerce_single_item_list(value: Any, predicate: Callable[[Any], bool]) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if isinstance(value, dict):
        if predicate(value):
            return [value]
        return [item for item in value.values() if isinstance(item, dict) and predicate(item)]
    return []


def _extract_single_question_candidate(payload: dict[str, Any], rubric: dict[str, Any]) -> dict[str, Any] | None:
    candidates = [
        rubric.get("question"),
        payload.get("question"),
        payload.get("rubric_question"),
    ]
    if _looks_like_question_dict(rubric):
        candidates.append(rubric)
    if _looks_like_question_dict(payload):
        candidates.append(payload)
    for candidate in candidates:
        if isinstance(candidate, dict):
            return candidate
    return None


def _extract_single_answer_candidate(
    payload: dict[str, Any],
    answer_key: dict[str, Any],
    question: dict[str, Any] | None,
) -> dict[str, Any] | None:
    candidates = [
        answer_key.get("question"),
        payload.get("answer_key_question"),
        payload.get("answer"),
        payload.get("correct_answer"),
    ]
    if _looks_like_answer_dict(answer_key):
        candidates.append(answer_key)
    if isinstance(question, dict) and _looks_like_answer_dict(question):
        candidates.append(question)
    for candidate in candidates:
        if isinstance(candidate, dict):
            return candidate
        if candidate is not None:
            return {"canonical_answer": str(candidate)}
    return None


def _looks_like_question_dict(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    return any(
        key in item
        for key in [
            "question_id",
            "id",
            "number",
            "question_type",
            "type",
            "max_score",
            "score",
            "points",
            "parts",
            "stem_summary",
        ]
    )


def _looks_like_answer_dict(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    return any(
        key in item
        for key in [
            "canonical_answer",
            "answer",
            "correct_answer",
            "direct_answer",
            "accepted_forms",
            "equivalent_answers",
            "parts",
        ]
    )


def _attach_parallel_generation_meta(
    payload: dict[str, Any],
    question_blocks: list[dict[str, str]],
    failures: list[dict[str, Any]],
    worker_count: int,
    attempt_counts: dict[str, int],
) -> None:
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    warnings = meta.setdefault("warnings", [])
    if not isinstance(warnings, list):
        warnings = []
        meta["warnings"] = warnings

    image_split_mode = any(
        str(block.get("semantic_source") or "").strip() == "images"
        for block in question_blocks
        if isinstance(block, dict)
    )
    meta["source"] = "pdf_image_split_parallel" if image_split_mode else "word_rich_split_parallel"
    meta["generation_mode"] = "pdf_image_split_parallel" if image_split_mode else "word_rich_split_parallel"
    meta["question_block_count"] = len(question_blocks)
    failed_question_ids = list(
        dict.fromkeys(
            str(failure.get("question_id") or "").strip()
            for failure in failures
            if str(failure.get("question_id") or "").strip()
        )
    )
    meta["single_question_success_count"] = max(0, len(question_blocks) - len(failed_question_ids))
    meta["single_question_failure_count"] = len(failed_question_ids)
    meta["config_generation_workers"] = int(worker_count)
    meta["single_question_attempts"] = {
        str(qid): int(attempts)
        for qid, attempts in attempt_counts.items()
        if str(qid).strip()
    }
    meta["failed_question_ids"] = failed_question_ids
    meta["failed_questions"] = [dict(failure) for failure in failures]

    warnings[:] = [
        warning
        for warning in warnings
        if str(warning) != "parallel generation completed with failed question blocks; score allocation paused until retry."
    ]
    existing = {str(item) for item in warnings}
    for failure in failures:
        qid = str(failure.get("question_id") or "").strip()
        error = str(failure.get("error") or "").strip()
        attempts = int(failure.get("attempts") or 1)
        warning = f"parallel generation failed for {qid} after {attempts} attempt(s): {error}"
        if warning not in existing:
            warnings.append(warning)
            existing.add(warning)
    if failures:
        summary = "parallel generation completed with failed question blocks; score allocation paused until retry."
        if summary not in existing:
            warnings.append(summary)


def _config_final_calibration_enabled(default: bool = False) -> bool:  # noqa: FBT001,FBT002
    """Final model calibration is OFF by default to avoid extra cost and failure risk.

    Set env var AI_GRADING_CONFIG_FINAL_CALIBRATION=1 to re-enable it.
    """
    raw = str(os.getenv("AI_GRADING_CONFIG_FINAL_CALIBRATION", "") or "").strip().lower()
    if not raw:
        return bool(default)
    if raw in {"0", "false", "no", "off", "disabled"}:
        return False
    return raw in {"1", "true", "yes", "on", "enabled"}


def _ensure_question_blocks_covered(payload: dict[str, Any], question_blocks: list[dict[str, str]]) -> None:
    rubric = payload.setdefault("rubric", {})
    answer_key = payload.setdefault("answer_key", {})
    if not isinstance(rubric, dict) or not isinstance(answer_key, dict):
        return
    questions = rubric.setdefault("questions", [])
    answers = answer_key.setdefault("questions", [])
    if not isinstance(questions, list) or not isinstance(answers, list):
        return

    existing_qids = {
        str(question.get("question_id") or "").strip()
        for question in questions
        if isinstance(question, dict)
    }
    existing_answer_qids = {
        str(answer.get("question_id") or "").strip()
        for answer in answers
        if isinstance(answer, dict)
    }
    warnings = payload.setdefault("meta", {}).setdefault("warnings", [])

    for block in question_blocks:
        qid = str(block.get("question_id") or "").strip()
        if not qid:
            continue
        if qid not in existing_qids:
            questions.append(_placeholder_question_from_block(block))
            existing_qids.add(qid)
            if isinstance(warnings, list):
                warnings.append(f"{qid} placeholder added because per-question generation did not return a mergeable rubric item")
        if qid not in existing_answer_qids:
            answers.append({"question_id": qid, "canonical_answer": "", "accepted_forms": [], "method_variants": [], "parts": []})
            existing_answer_qids.add(qid)

    qid_order = {
        str(block.get("question_id") or "").strip(): index
        for index, block in enumerate(question_blocks)
        if str(block.get("question_id") or "").strip()
    }
    questions.sort(key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000) if isinstance(item, dict) else 10_001)
    answers.sort(key=lambda item: qid_order.get(str(item.get("question_id") or "").strip(), 10_000) if isinstance(item, dict) else 10_001)


def _placeholder_question_from_block(block: dict[str, str]) -> dict[str, Any]:
    qid = str(block.get("question_id") or "").strip()
    text = str(block.get("text") or "").strip()
    qtype = _infer_question_type_from_block_text(text)
    return {
        "question_id": qid,
        "question_type": qtype,
        "max_score": 1,
        "stem_summary": text.splitlines()[0][:120] if text else qid,
        "parts": [
            {
                "part_id": qid,
                "part_score": 1,
                "steps": [
                    {
                        "step_id": "S1",
                        "step_score": 1,
                        "core_goal": "Manual rubric confirmation required",
                        "required_elements": ["AI did not return a mergeable question; confirm rubric manually"],
                        "allow_alternative_methods": True,
                    }
                ],
                "presentation_rules": [],
            }
        ],
    }


def _infer_question_type_from_block_text(text: str) -> str:
    value = str(text or "")
    option_labels = {
        match.group(1).upper()
        for match in re.finditer(r"(?m)^\s*([A-Da-d])\s*(?:[.．、)]|\s{2,})", value)
    }
    option_line = any(re.search(r"A.{0,80}B.{0,80}C.{0,80}D", line) for line in value.splitlines())
    if len(option_labels) >= 3 or option_line:
        return "choice"
    has_blank = bool(re.search(r"_{2,}|[ \t]{3,}|　{1,}|（\s*）|\(\s*\)", value))
    if has_blank:
        return "fill_blank"
    if any(token in value for token in ["作图", "作出", "画出", "保留作图痕迹"]):
        return "comprehensive"
    if any(token in value for token in ["证明", "理由", "说明", "求证", "全等", "证得"]):
        return "proof"
    has_subparts = bool(re.search(r"[（(]\s*[1-9]\s*[）)]", value))
    if has_subparts:
        return "calculation"
    if any(token in value.lower() for token in ["proof", "why", "explain", "relationship"]):
        return "comprehensive"
    return "comprehensive"


def _attach_generation_fallback_warning(
    payload: dict[str, Any],
    question_blocks: list[dict[str, str]],
    parallel_error: Exception | None,
) -> None:
    if parallel_error is None:
        return
    meta = payload.setdefault("meta", {})
    if not isinstance(meta, dict):
        payload["meta"] = meta = {}
    warnings = meta.setdefault("warnings", [])
    if not isinstance(warnings, list):
        warnings = []
        meta["warnings"] = warnings
    meta["source"] = "word_docx_single_fallback"
    meta["question_block_count"] = len(question_blocks)
    meta["parallel_generation_fallback"] = True
    warnings.append(f"parallel generation failed; fell back to single full-document generation: {parallel_error}")


def _assign_scores_to_merged_config(
    merged: dict[str, Any],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    *,
    include_document_text: bool = True,
) -> None:
    """Phase-2 AI call: allocate real scores to all questions so they sum to 100.

    Per-question prompts set all score fields to the placeholder value 1.
    This function sends the merged structure + original Word text to the model,
    which reads score hints like （3分）, 共X分 from the document and distributes
    the remaining points proportionally to question complexity.

    On failure, falls back to local proportional scaling via force_payload_total_score.
    """
    rubric_questions: list[Any] = merged.get("rubric", {}).get("questions", [])
    if not isinstance(rubric_questions, list) or not rubric_questions:
        force_payload_total_score(merged, target_total=100.0)
        return

    # Build a lightweight structure summary: only IDs, types, part/step counts
    structure_summary: list[dict[str, Any]] = []
    for q in rubric_questions:
        if not isinstance(q, dict):
            continue
        parts = q.get("parts") or []
        part_list = []
        for p in parts if isinstance(parts, list) else []:
            if not isinstance(p, dict):
                continue
            steps = p.get("steps") or []
            part_list.append({
                "part_id": str(p.get("part_id") or ""),
                "response_mode": str(p.get("response_mode") or ""),
                "step_ids": [str(s.get("step_id") or "") for s in steps if isinstance(s, dict)],
            })
        structure_summary.append({
            "question_id": str(q.get("question_id") or ""),
            "question_type": str(q.get("question_type") or ""),
            "parts": part_list,
        })

    prompt = (
        "你正在分配试卷各题目的得分。试卷总分必须精确等于 100 分。\n\n"
        "硬性规则（请严格遵守）：\n"
        "1. 所有相同类型（选择题 choice / 填空题 fill_blank）的题目必须共享相同的 max_score（即同类题目的单题分值相同）。\n"
        "2. 对于每道题目：其下属各部分的 part_score 之和必须等于 max_score。对于每个 part：其下属各步骤的 step_score 之和必须等于 part_score。\n"
        "3. 所有分值必须是正整数。\n"
        "4. 试卷所有题目的 max_score 之和必须精确等于 100。\n"
        "5. 将未分配的分数，按照步骤的数量比例分配给主观大题。\n"
        "6. 关键限制 —— 所有题型的单题分值上限均为 18 分：\n"
        "   - 每道独立题目的 max_score 必须 <= 18 分。\n"
        "   - 所有共享相同大题号的子小问（例如 Q10_1 和 Q10_2 共享同一个大题前缀 Q10）的合并总分必须 <= 18 分。\n"
        "   - 若某题原有标注分值超过 18 分，必须封顶为 18 分，并将其余分数分摊给其他题目。\n"
        "   - 你必须严格保持输入中提供的题目结构和 question_id 列表，绝对不允许新增、删除、修改或拆分任何 question_id 题号。\n\n"
        "Return ONLY the following JSON (no markdown, no explanation):\n"
        "{\n"
        "  \"question_scores\": [\n"
        "    {\n"
        "      \"question_id\": \"Q1\",\n"
        "      \"max_score\": 3,\n"
        "      \"parts\": [\n"
        "        {\"part_id\": \"Q1\", \"part_score\": 3,\n"
        "         \"steps\": [{\"step_id\": \"S1\", \"step_score\": 3}]}\n"
        "      ]\n"
        "    }\n"
        "  ]\n"
        "}\n\n"
        f"Word document (for score hints, first 6000 chars):\n{str(doc_text or '')[:6000]}\n\n"
        f"Question structure to score:\n{json.dumps(structure_summary, ensure_ascii=False)}"
    )
    prompt = _build_score_allocation_prompt(
        structure_summary,
        doc_text,
        include_document_text=include_document_text,
    )

    if report:
        q_count = len(structure_summary)
        report(0.88, "AI 赋分", f"正在为合并后的 {q_count} 道题统一分配分值（共 100 分）...")

    try:
        score_data = llm_client.json_from_text(prompt, model=model_name)
        _apply_score_allocation(merged, score_data)
        merged.setdefault("meta", {})["score_allocation_ai_success"] = True
        if report:
            report(0.91, "AI 赋分完成", "已将 AI 分配的分值写入各题评分细则。")
    except Exception as exc:
        if report:
            report(0.91, "AI 赋分失败，使用本地均分兜底", str(exc))
        warnings = merged.setdefault("meta", {}).setdefault("warnings", [])
        if isinstance(warnings, list):
            warnings.append(f"score allocation AI call failed; fell back to local scaling: {exc}")
        merged.setdefault("meta", {})["score_allocation_ai_success"] = False
    # Always do a local normalization pass to guarantee score consistency
    force_payload_total_score(merged, target_total=100.0)


def _build_score_allocation_prompt(
    structure_summary: list[dict[str, Any]],
    doc_text: str,
    *,
    include_document_text: bool,
) -> str:
    source_context = (
        f"\n原始文档文本（仅用于识别原卷分值提示）：\n{str(doc_text or '')[:6000]}\n"
        if include_document_text and str(doc_text or "").strip()
        else "\n本次为图片语义来源，不提供也不得推测 PDF 抽取文字或原卷分值。\n"
    )
    return (
        "请仅为下列已确认题目结构分配分值，总分必须精确等于100。\n"
        "不同题型之间不限制分值高低；相同类型客观题必须同分；所有分值均为正整数；单题不超过18分。\n"
        "保持所有 question_id、part_id、step_id 和小问结构不变。"
        "分值可以不采用原卷分值，但不得改变小问作答要求或 response_mode。\n"
        f"{source_context}"
        "仅返回 JSON：{\"question_scores\":[{\"question_id\":\"Q1\",\"max_score\":1,"
        "\"parts\":[{\"part_id\":\"Q1\",\"part_score\":1,\"steps\":[{\"step_id\":\"S1\",\"step_score\":1}]}]}]}\n"
        f"待分值结构：\n{json.dumps(structure_summary, ensure_ascii=False)}"
    )




def _calibrate_merged_config(
    payload: dict[str, Any],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    *,
    strict: bool = False,
) -> None:
    prompt = (
        "final calibration: You are a rubric calibration assistant. "
        "Check the merged per-question rubric against the full document. "
        "Only fix missing question ids, empty answers, obvious question type errors, and total score allocation. "
        "Do not remove existing proof_obligations, steps, or parts. "
        "accepted_forms must keep only strictly equivalent answers. Return complete JSON with the same schema.\n\n"
        f"Full document text:\n{str(doc_text or '')[:8000]}\n\n"
        f"Merged result:\n{json.dumps(payload, ensure_ascii=False)}"
    )
    try:
        calibrated = llm_client.json_from_text(prompt, model=model_name)
        validate_generated_config(calibrated)
    except Exception as exc:
        if strict:
            raise RuntimeError(f"final calibration failed: {exc}") from exc
        meta = payload.setdefault("meta", {})
        if isinstance(meta, dict):
            warnings = meta.setdefault("warnings", [])
            if isinstance(warnings, list):
                warnings.append(f"final calibration failed; kept merged per-question config: {exc}")
        return
    payload.clear()
    payload.update(calibrated)


def _bounded_int(value: Any, default: int, minimum: int, maximum: int) -> int:
    try:
        resolved = int(value) if value is not None else int(default)
    except (TypeError, ValueError):
        resolved = int(default)
    return max(minimum, min(maximum, resolved))


def refine_grading_config_from_manual_structure(
    payload: dict[str, Any],
    llm_client: LLMClient,
    model_name: str | None = None,
) -> dict[str, Any]:
    """Ask AI to complete answers/rubrics while preserving teacher-edited parts."""
    working_payload = json.loads(json.dumps(payload, ensure_ascii=False))
    validate_generated_config(working_payload)
    prompt = _build_manual_structure_refinement_prompt(working_payload)
    refined = llm_client.json_from_text(prompt, model=model_name)
    try:
        validate_generated_config(refined)
    except Exception:
        _dump_failed_generated_payload(refined)
        raise
    return refined


def _dump_failed_generated_payload(payload: dict[str, Any]) -> None:
    try:
        try:
            from path_manager import get_path_manager
            out_dir = get_path_manager().upload_config_dir
        except Exception:
            out_dir = Path(__file__).resolve().parent / "config" / "uploaded"
        out_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        (out_dir / f"failed_generated_config_normalized_{ts}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except Exception:
        pass














def iter_effective_rubric_items(
    payload: Mapping[str, object],
) -> Iterator[tuple[str, Mapping[str, object], Mapping[str, object]]]:
    for item_ref, _parent_ref, raw_question, raw_item in iter_effective_rubric_item_refs(payload):
        yield item_ref, raw_question, raw_item


def iter_effective_rubric_item_refs(
    payload: Mapping[str, object],
) -> Iterator[tuple[str, str, Mapping[str, object], Mapping[str, object]]]:
    rubric = payload.get("rubric") if isinstance(payload.get("rubric"), Mapping) else payload
    questions = rubric.get("questions") if isinstance(rubric, Mapping) else None
    if not isinstance(questions, list):
        return
    for question_index, raw_question in enumerate(questions, start=1):
        if not isinstance(raw_question, Mapping):
            continue
        question_ref = str(
            raw_question.get("question_id")
            or raw_question.get("id")
            or raw_question.get("number")
            or f"Q{question_index}"
        ).strip()
        parts = raw_question.get("parts")
        emitted = False
        if isinstance(parts, list) and parts:
            for part_index, raw_part in enumerate(parts, start=1):
                if not isinstance(raw_part, Mapping):
                    continue
                part_ref = str(
                    raw_part.get("part_id")
                    or raw_part.get("question_id")
                    or f"{question_ref}.{part_index}"
                ).strip()
                emitted = True
                yield part_ref, question_ref, raw_question, raw_part
        if not emitted:
            yield question_ref, question_ref, raw_question, raw_question


def _needs_objective_repair(payload: dict[str, Any], doc_text: str) -> bool:
    questions = payload.get("rubric", {}).get("questions", [])
    qtypes = {
        str(q.get("question_type"))
        for q in questions
        if isinstance(q, dict)
    } if isinstance(questions, list) else set()
    text = str(doc_text or "")
    doc_has_choice = any(re.search(r"A.{0,80}B.{0,80}C.{0,80}D", line) for line in text.splitlines())
    doc_has_fill = bool(re.search(r"_{2,}|[ \t]{3,}|　{1,}|（\s*）|\(\s*\)", text))
    if doc_has_choice and "choice" not in qtypes:
        return True
    if doc_has_fill and "fill_blank" not in qtypes:
        return True
    return False


def _build_objective_repair_prompt(doc_text: str, bad_payload: dict[str, Any]) -> str:
    return (
        "你上一次生成的评分标准漏掉了选择题/填空题或客观题，这是严重错误。\n"
        "请重新从 Word 原文中生成完整评分标准，必须包含全部题目。\n\n"
        "修复要求：\n"
        "1) 不要只输出 Q12/Q13 等解答题；必须补齐选择题、填空题、客观题。\n"
        "2) 选择题 question_type=choice；填空题 question_type=fill_blank。\n"
        "3) 客观题也必须在 rubric.questions 和 answer_key.questions 中逐题列出。\n"
        "4) 如果题干前半部分是题目、后半部分是答案解析，请把二者按题号匹配。\n"
        "5) 如果分值没有逐题标明，请根据总分/题型合理分配，并在 meta.warnings 写明。\n"
        "6) 不能再输出“选择题与填空题未包含在详细评分逻辑中”或类似 warning。\n\n"
        "输出格式仍必须与原要求完全一致：严格 JSON 对象，包含 rubric、answer_key、meta。\n\n"
        f"Word 原文：\n{doc_text}\n\n"
        f"上一次错误输出（供你定位遗漏，不要照抄）：\n{json.dumps(bad_payload, ensure_ascii=False)}"
    )


def _build_manual_structure_refinement_prompt(payload: dict[str, Any]) -> str:
    return (
        "您正在对教师编辑过的评分标准进行精修/对齐。请仅返回严格的 JSON 数据。\n"
        "硬性要求：\n"
        "1) 必须保留每一个已有的 rubric.questions[].question_id。\n"
        "2) 必须保留每一个已有的 parts[].part_id；绝对不能合并、删除或修改教师创建的 parts[] 小问结构。\n"
        "3) answer_key.questions[].parts 必须通过 part_id 与 rubric 中的 parts 保持一致对齐。\n"
        "4) 补全缺失的答案、accepted_forms、解析、步骤分、证明扣分项和证据链规则。\n"
        "5) 保持整张试卷总分 total_score 和各题 max_score 的总和精确等于 100。\n"
        "6) 对于证明题或计算解答题，按证据步骤步骤分进行细化，并保守地给与仅有答案无过程的得分限制。\n\n"
        f"当前评分标准 JSON：\n{json.dumps(payload, ensure_ascii=False)}"
    )




def save_generated_config(upload_dir: Path, payload: dict[str, Any], ts: str) -> tuple[Path, Path]:
    upload_dir.mkdir(parents=True, exist_ok=True)
    validate_generated_config(payload)
    rubric_path = upload_dir / f"rubric_{ts}.json"
    answer_key_path = upload_dir / f"answer_key_{ts}.json"

    _write_generated_json_atomic(rubric_path, payload["rubric"])
    _write_generated_json_atomic(answer_key_path, payload["answer_key"])

    return rubric_path, answer_key_path


def _write_generated_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.parent / f".{path.name}.{os.urandom(8).hex()}.tmp"
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _build_generation_prompt(doc_text: str, *, include_source_text: bool = True) -> str:
    return (
        "你是中学数学教研评分设计助手。\n"
        "请从给定Word文本中抽取‘题目、分值、答案与解析’，并生成可执行评分标准。\n"
        "核心目标：支持填空题等价答案判分、解答题/证明题按证明义务扣分、不同正确解法给分。\n"
        "硬性总分：本系统所有考试批改统一按 100 分制设计。rubric.total_score 必须等于 100，所有题目 max_score 之和必须等于 100；若原卷不是 100 分制，请按原始分值比例换算。\n"
        "单题上限：所有题型的 question.max_score 均不能超过 18 分。题目可以拆成多个小问 parts 分别赋分，但整道题 max_score 仍不得超过 18。\n"
        "硬性赋分：所有 max_score、part_score、step_score、proof_obligations.weight、deduction_policy.max_deduction 都必须是整数，不能出现 2.5、3.33 这类小数。\n"
        "同类同分（仅客观题）：相同 question_type 的客观题必须分值完全相同。例如所有 choice 题同分，所有 fill_blank 题同分，不能出现有的选择题3分、有的选择题4分。解答类大题（calculation/proof/comprehensive）不要求同类同分，可按题目难度与工作量赋予不同分值。\n"
        "题型细分：解答类题目不要全部写成一种类型，可按实际任务分为 calculation（计算/求解）、proof（证明）、comprehensive（一般综合解答）等多种 question_type。\n"
        "题型纠偏：只要题目要求证明、求证、说明理由、说明某结论成立、补全证明过程、判断并说明、添加条件使结论能够推出，就应标为 proof；如果题目主要要求求角度、求长度、求周长、求面积、求值、计算、化简或解方程，且不要求证明/说明理由，才标为 calculation。comprehensive 用于同时包含证明、计算、作图或开放论述的混合型题。\n"
        "若原卷分值与“100分制、整数、客观题同类同分、单题不超过18分”冲突，请优先按这些规则重新设计赋分。\n"
        "特别要求：证明题/解答题不要把参考答案路径当作唯一标准；应抽象成 proof_obligations（证明义务）和 deduction_policy（扣分规则）。\n"
        "必须覆盖 Word 中出现的全部题目，包括选择题、填空题、判断题、客观题和解答题；禁止只生成大题或只生成答案解析部分。\n"
        "如果 Word 中有选择题/填空题，即使评分逻辑简单，也必须在 rubric.questions 和 answer_key.questions 中逐题列出。\n"
        "选择题/填空题不能写成“仅抽象化处理”或只放在 meta.warnings；如果答案明确，应生成 direct_answer 评分标准。\n"
        "如果信息缺失，不要臆造，把不确定项写入 meta.warnings。\n\n"
        "仅输出严格 JSON，不要 markdown。输出必须尽量压缩为单行 JSON，不要漂亮打印，不要添加解释性文字，避免输出被截断。\n"
        "可以省略空数组字段中的多余说明，但 JSON 顶层和必需字段必须完整闭合。\n"
        "JSON 结构必须是:\n"
        "{\n"
        "  \"rubric\": {\n"
        "    \"exam_title\": \"...\",\n"
        "    \"total_score\": 100,\n"
        "    \"questions\": [\n"
        "      {\n"
        "        \"question_id\": \"Q1\",\n"
        "        \"question_type\": \"fill_blank|choice|calculation|proof|comprehensive\",\n"
        "        \"max_score\": 8,\n"
        "        \"stem_summary\": \"题干摘要，用一句话概括本题考查内容\",\n"
        "        \"grading_mode\": \"direct_answer|deductive_obligation\",\n"
        "        \"require_final_answer\": false,\n"
        "        \"answer_only_max_score\": 2,\n"
        "        \"proof_obligations\": [\n"
        "          {\"obligation_id\": \"O1\", \"description\": \"需要证明/建立的关键数学义务\", \"weight\": 3, \"acceptable_evidence\": [\"等价完成方式\"]}\n"
        "        ],\n"
        "        \"deduction_policy\": [\n"
        "          {\"issue\": \"关键逻辑断裂/条件缺失/循环论证\", \"max_deduction\": 2, \"severity\": \"minor|major|fatal\"}\n"
        "        ],\n"
        "        \"parts\": [\n"
        "          {\n"
        "            \"part_id\": \"Q1(P1)\",\n"
        "            \"part_score\": 4,\n"
        "            \"steps\": [\n"
        "              {\n"
        "                \"step_id\": \"S1\",\n"
        "                \"step_score\": 2,\n"
        "                \"core_goal\": \"...\",\n"
        "                \"required_elements\": [\"...\"],\n"
        "                \"allow_alternative_methods\": true\n"
        "              }\n"
        "            ],\n"
        "            \"presentation_rules\": [\n"
        "              {\"rule\": \"过程规范/结论完整\", \"max_deduction\": 1}\n"
        "            ]\n"
        "          }\n"
        "        ]\n"
        "      }\n"
        "    ]\n"
        "  },\n"
        "  \"answer_key\": {\n"
        "    \"questions\": [\n"
        "      {\n"
        "        \"question_id\": \"Q1\",\n"
        "        \"canonical_answer\": \"...\",\n"
        "        \"accepted_forms\": [\"...\", \"...\"],\n"
        "        \"method_variants\": [\n"
        "          {\"name\": \"方法A\", \"outline\": \"...\"},\n"
        "          {\"name\": \"方法B\", \"outline\": \"...\"}\n"
        "        ],\n"
        "        \"parts\": [\n"
        "          {\n"
        "            \"part_id\": \"Q1(P1)\",\n"
        "            \"answer\": \"...\",\n"
        "            \"analysis\": \"...\",\n"
        "            \"step_milestones\": [\"关键中间结论1\", \"关键中间结论2\"]\n"
        "          }\n"
        "        ]\n"
        "      }\n"
        "    ]\n"
        "  },\n"
        "  \"meta\": {\n"
        "    \"source\": \"word_docx\",\n"
        "    \"warnings\": [\"...\"]\n"
        "  }\n"
        "}\n\n"
        "硬性要求：\n"
        "1) question_id 在 rubric 与 answer_key 中一一对应。\n"
        "1.1) 知识点与题目标签由题库程序单独维护，不要输出任何 knowledge 字段。\n"
        "2) max_score、part_score、step_score 必须为整数且层级总分一致；不得输出小数。\n"
        "2.1) 整张试卷总分必须严格为 100 分，不能返回 10 分、120 分或其他总分。\n"
        "2.2) 相同 question_type 的客观题（choice/fill_blank）max_score 必须完全一致；解答类大题（calculation/proof/comprehensive）允许不同分值。\n"
        "2.3) 不限制不同题型之间的分值高低关系；仅要求同类型客观题同分。\n"
        "2.4) 任意 question.max_score 必须小于或等于 18；若大题有多问，请在 parts 中拆分小问分值，不要让整题超过 18。\n"
        "3) 对选择题/填空题，question_type 必须分别是 choice 或 fill_blank，parts 可只有一个，steps 可只有一个 direct-answer 步骤。\n"
        "4) 对解答题，必须按题目特征细分为 calculation/proof/comprehensive，必须有 parts 与 steps，不能只给最终答案；同时必须给出 proof_obligations 与 deduction_policy。明确要求证明/求证/说明结论成立的题使用 proof；主要求数值、角度、长度、面积、化简或方程结果的题使用 calculation。\n"
        "4.1) 对 proof 与 calculation，默认 require_final_answer=false，不因未额外写“答”单独扣分；对 comprehensive，默认 require_final_answer=true，未写最终答/结论完整性可小扣分。若题目本身明确要求写结论，可按题意调整。\n"
        "4.2) 对 proof/calculation/comprehensive，必须给出 answer_only_max_score：学生只写最终答案但没有有效过程时最多得分。该分值只能是少量结论分，不得超过本题 max_score 的 30%。\n"
        "5) 对选择题/填空题/direct_answer 题，必须执行全对全错规则：学生答案与 canonical_answer 或 accepted_forms 等价才给该题/该空满分；不符合答案及等价答案时该评分单元必须给 0 分，不得因为接近、猜测方向正确或过程片段给一半分。\n"
        "5.1) 对填空题，必须尽可能提供 accepted_forms（等价表达），例如 1/4、0.25、25% 这类数值等价形式，或 OB=OC、OC=OB、BO=CO 这类几何关系等价写法。\n"
        "6) 对可接受的不同解法，在 method_variants 中列出。\n"
        "7) 若题号或分值不确定，使用连续编号并写入 warnings，但不能因此漏掉题目。\n"
        "8) meta.warnings 不能出现“选择题与填空题未包含在详细评分逻辑中”这类遗漏；遇到这种情况必须回到 questions 中补齐。\n\n"
        "证明题/解答题生成准则：\n"
        "- 答案结论只作为小部分得分依据，过程分主要来自 proof_obligations 是否完成。\n"
        "- 证明题和纯计算题通常不要求额外写“答”；一般综合解答题可要求结论完整，但这应作为 presentation_rules 的小扣分项，不应替代过程分判断。\n"
        "- 只写最终答案但没有有效过程时，只能获得 answer_only_max_score 以内的结论分；不能获得主要过程分。\n"
        "- proof_obligations 描述学生必须完成的数学责任，例如建立辅助条件、证明全等/相似、推出角度/线段关系、说明定理前提等。\n"
        "- deduction_policy 描述扣分项，例如缺少前提、跳步严重、逻辑循环、结论与过程断裂、符号/对象指代不清。\n"
        "- steps 可以对应证明义务的评分权重，但不能要求学生过程与参考答案逐句一致。\n\n"
        + (f"Word原文:\n{doc_text}" if include_source_text else "")
    )



def generate_grading_config_from_text(
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
    *,
    question_blocks: Sequence[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return generate_grading_config_from_docx_text(
        doc_text,
        llm_client,
        model_name=model_name,
        report=report,
        question_blocks=question_blocks,
    )


def generate_grading_config_from_images(
    image_blobs: list[bytes],
    doc_text: str,
    llm_client: LLMClient,
    model_name: str | None = None,
    report: Any = None,
) -> dict[str, Any]:
    if not image_blobs:
        raise ValueError("PDF 整卷视觉模式未获得任何页面图片。")
    if report:
        report(0.20, "PDF 整卷视觉单次请求", f"发送 {len(image_blobs)} 张整页图片，一次完成解析与赋分。")
    payload = llm_client.json_from_images_once(
        _build_whole_image_generation_prompt(),
        image_blobs,
        model=model_name,
        extra_kwargs=_config_generation_extra_kwargs(),
        use_config_client=True,
    )
    return _finalize_whole_generation_payload(payload, "whole_pdf_visual_single_request")


# P3-10 compatibility exports. The implementation lives in focused policy
# modules; historical callers keep resolving the same session_manager names.
from backend.config_generation import local_facts as _config_local_facts  # noqa: E402
from backend.config_generation import normalization as _config_normalization  # noqa: E402
from backend.config_generation import quality as _config_quality  # noqa: E402
from backend.config_generation import score_allocation as _config_score_allocation  # noqa: E402

for _compat_module in (
    _config_normalization,
    _config_score_allocation,
    _config_local_facts,
    _config_quality,
):
    for _compat_name in _compat_module.SESSION_MANAGER_COMPAT_EXPORTS:
        globals()[_compat_name] = getattr(_compat_module, _compat_name)
del _compat_module, _compat_name

# P3-09 compatibility facade. Keep the historical names for callers while the
# active batch workflow and score prompt live under backend.config_generation.
from backend.config_generation.compat import (  # noqa: E402
    failed_grading_config_batches as failed_grading_config_batches,
    failed_grading_config_question_ids as failed_grading_config_question_ids,
    generate_grading_config_in_batches as generate_grading_config_in_batches,
    refine_grading_config_from_manual_structure as refine_grading_config_from_manual_structure,
    retry_failed_grading_config_batches as retry_failed_grading_config_batches,
)
from backend.config_generation.gateway import (  # noqa: E402
    config_generation_extra_kwargs as _config_generation_extra_kwargs,
)
from backend.config_generation.prompts import (  # noqa: E402
    build_manual_structure_refinement_prompt as _build_manual_structure_refinement_prompt,
    build_score_allocation_prompt as _build_score_allocation_prompt,
)


# P3-08 compatibility exports. Parsing now lives in backend.document_parsing;
# callers that historically imported these names keep the same public surface.
from backend.document_parsing import (  # noqa: E402
    extract_docx_text as extract_docx_text,
    image_paths_from_rich_text as _image_paths_from_rich_text,
    infer_question_type_from_text as _infer_question_type_from_block_text,
    parse_plain_question_blocks as preview_question_blocks_from_docx_text,
)
from backend.document_parsing.question_blocks import (  # noqa: E402
    _parse_inline_answer_blocks as _parse_inline_answer_blocks,
)
