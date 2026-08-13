from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


AUTHOR_NOTE_PREFIX = "【作者核实】"
NUMBER_PATTERN = re.compile(r"[+-]?\d+(?:\.\d+)?%?")


def validate_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
    if bad_member is not None:
        raise AssertionError(f"{path.name}: corrupt member {bad_member}")


def nonempty_paragraphs(document: Document) -> list[str]:
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def table_text(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def table_numbers(table) -> list[str]:
    return NUMBER_PATTERN.findall(
        "\n".join(cell.text for row in table.rows for cell in row.cells)
    )


def source_data_numbers(table, body_table_index: int) -> list[str]:
    included_columns = None
    if body_table_index == 4:  # Table 5: exclude narrative interpretation column.
        included_columns = set(range(6))
    elif body_table_index == 5:  # Table 6: compare count/rate/adjustment fields only.
        included_columns = set(range(5))

    texts = []
    for row in table.rows:
        for column_index, cell in enumerate(row.cells):
            if included_columns is None or column_index in included_columns:
                texts.append(cell.text)
    return NUMBER_PATTERN.findall("\n".join(texts))


def reference_texts(document: Document) -> list[str]:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    heading = next(
        index for index, text in enumerate(paragraphs) if text.strip() == "参考文献"
    )
    return [text for text in paragraphs[heading + 1 :] if text.strip()]


def ordered_blocks(document: Document) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text or text.startswith(AUTHOR_NOTE_PREFIX):
                continue
            blocks.append(("p", text))
        elif tag == "tbl":
            table = Table(child, document)
            blocks.append(("tbl", table_text(table)))
    return blocks


def tracked_counts(path: Path) -> tuple[int, int, bool]:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        settings_xml = archive.read("word/settings.xml")
    return (
        document_xml.count(b"<w:ins"),
        document_xml.count(b"<w:del"),
        b"<w:trackRevisions" in settings_xml,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--clean", type=Path, required=True)
    parser.add_argument("--tracked", type=Path, required=True)
    parser.add_argument("--accepted", type=Path, required=True)
    args = parser.parse_args()

    for path in (args.source, args.clean, args.tracked, args.accepted):
        validate_zip(path)

    source = Document(str(args.source))
    clean = Document(str(args.clean))
    tracked = Document(str(args.tracked))
    accepted = Document(str(args.accepted))

    clean_text = "\n".join(paragraph.text for paragraph in clean.paragraphs)
    first_paragraph = nonempty_paragraphs(clean)[0]
    assert first_paragraph == "初中数学纸笔测验中两种 AI 阅卷模式的题型误差与教师复核"
    assert "基本信息" not in clean_text
    assert "论文自检报告正文" not in clean_text
    assert not re.search(r"\b1[3-9]\d{9}\b", clean_text)
    assert not re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", clean_text)

    assert len(source.tables) == 7
    assert len(clean.tables) == 7
    for index, original_table in enumerate(source.tables[1:]):
        clean_table = clean.tables[index]
        assert source_data_numbers(
            original_table,
            index,
        ) == source_data_numbers(clean_table, index), (
            f"Numeric table content changed in table {index + 1}"
        )

    assert reference_texts(source) == reference_texts(clean)
    assert len(clean.inline_shapes) == 4
    assert "图 5 基于题型证据结构的AI阅卷两层分流与反馈循环机制" not in clean_text
    assert "表 7 本次题型结果与暂定处理方式" in clean_text

    forbidden_phrases = [
        "题型证据结构—输入方式—误差机制—教师裁决",
        "单生完整证据包",
        "同题多生局部证据包",
        "两级评价责任配置",
        "本研究的核心贡献不在于",
    ]
    for phrase in forbidden_phrases:
        assert phrase not in clean_text

    critical_values = [
        "78",
        "1404",
        "4.49",
        "5.85",
        "0.972",
        "0.944",
        "+0.32",
        "-3.16",
        "24.4%",
        "24.1%",
        "18 条",
        "11 条",
    ]
    for value in critical_values:
        assert value in clean_text, f"Missing critical value: {value}"

    insertions, deletions, tracking_enabled = tracked_counts(args.tracked)
    assert insertions > 0
    assert deletions > 0
    assert tracking_enabled

    assert ordered_blocks(clean) == ordered_blocks(accepted)
    assert [table_text(table) for table in clean.tables] == [
        table_text(table) for table in accepted.tables
    ]

    print(
        {
            "source_tables": len(source.tables),
            "clean_tables": len(clean.tables),
            "clean_images": len(clean.inline_shapes),
            "tracked_insertions": insertions,
            "tracked_deletions": deletions,
            "references": len(reference_texts(clean)),
            "accepted_matches_clean": True,
        }
    )


if __name__ == "__main__":
    main()
