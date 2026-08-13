from __future__ import annotations

import argparse
import datetime as dt
import difflib
import importlib.util
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.shared import RGBColor


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"
BODY_START_BLOCK = 50
DELETE_BLOCKS = {129, 130}  # Original Figure 5 and its caption.
TABLE_ANCHOR_BLOCK = 128
TABLE_GEOMETRY_HELPER = Path(
    r"C:\Users\89418\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.727.11326\skills\documents\scripts\table_geometry.py"
)
REPLACEMENT_TABLE_CAPTION = "表 7 本次题型结果与暂定处理方式"
REPLACEMENT_TABLE_ROWS = [
    ["本次证据", "暂定处理方式", "结论边界"],
    [
        "Q1—Q5：两模式 MAE 均为 0.45（1.3%）",
        "批量识别；多选、涂改等异常作答由教师核对",
        "仅适用于本次选择题",
    ],
    [
        "Q6—Q9：整卷 MAE 为 1.29（4.0%），混合为 3.96（12.4%）",
        "核对等价表达、作废答案和替换答案",
        "尚不能确定固定输入方式",
    ],
    [
        "Q10：整卷和混合 MAE 分别为 2.68（24.4%）和 2.65（24.1%）",
        "教师逐项复核并确定最终得分；AI 仅整理图像和提示疑似痕迹",
        "本次证据最明确的处理建议",
    ],
    [
        "Q11：混合 MAE 为 0.60（5.4%），整卷为 1.00（9.1%）",
        "后续试卷继续检验局部呈现方式",
        "不据此宣称稳定优势",
    ],
    [
        "Q12：整卷 MAE 为 0.81（7.3%），混合为 1.97（17.9%）",
        "后续试卷继续检验单份完整答卷的呈现方式",
        "不据此宣称稳定优势",
    ],
    [
        "整卷、混合分别标记 18 条和 11 条；复核后调整 1 条和 4 条",
        "系统标记只作为追加复核提示",
        "本研究未验证其检出能力",
    ],
]
TABLE_CELL_REPLACEMENTS = [
    (1, 0, 4, "教师需查看并确认的内容"),
    (1, 1, 4, "查看多选、涂改或异常标记，确认有效选项"),
    (1, 2, 4, "核对非标准表达是否等价，确认涂改后保留的答案"),
    (1, 3, 3, "标出识别到的点、线及疑似位置"),
    (1, 3, 4, "逐项核对点线位置及几何关系是否符合评分细则"),
    (1, 4, 4, "核对分段讨论及各步骤是否符合评分细则"),
    (1, 5, 3, "列出识别到的条件、推理步骤及缺失环节"),
    (1, 5, 4, "核对条件调用、推理衔接和结论是否完整"),
    (2, 0, 2, "核查所用记录"),
    (2, 2, 3, "整卷：单名学生整页或整卷图像；混合：同题多名学生作答区域拼接"),
    (2, 3, 3, "日志保存两种模式实际使用的模型名称"),
    (2, 5, 1, "统一答案、评分细则和结构化提示词"),
    (2, 5, 3, "提示词包含作废内容处理、步骤提取和提示注入防护要求"),
    (2, 6, 1, "模型输出“需要复核”字段，系统据此筛选记录供教师查看"),
    (2, 6, 3, "该字段由模型报告置信度触发；置信度不作为正确概率"),
    (5, 1, 6, "两模式 MAE 相同"),
    (5, 2, 6, "本样本中整卷 MAE 较低"),
    (5, 3, 6, "两模式相对 MAE 均约为 24%"),
    (5, 4, 6, "本样本中混合 MAE 较低"),
    (5, 5, 6, "本样本中整卷 MAE 较低"),
    (6, 0, 2, "“需要复核”=是的日志数"),
    (6, 0, 3, "上述日志占比"),
    (6, 0, 4, "上述日志中发生分数调整的数量"),
    (6, 0, 5, "上述日志所在题号"),
    (6, 0, 6, "补充说明"),
    (6, 1, 5, "Q10、Q11、Q12"),
    (6, 1, 6, "18 条标记中 1 条分数下调"),
    (6, 2, 5, "Q3—Q9"),
    (6, 2, 6, "11 条标记中 4 条分数上调；Q10、Q12 无标记"),
]
AUTHOR_CHECK_NOTE = (
    "【作者核实】图 4 题面显示 Q10“本小题 9 分”，而表 1、表 5 按 11 分计算；"
    "图内“偏宽/偏严”的判断也未附案例级分数。请回查原卷、评分细则和复核日志，"
    "再决定是否调整分值、相对 MAE 及图内标签。"
)


def local_name(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def paragraph_text(paragraph_element) -> str:
    return "".join(node.text or "" for node in paragraph_element.iter(qn("w:t")))


def first_run_properties(paragraph_element):
    for run in paragraph_element.iter(qn("w:r")):
        run_properties = run.find(qn("w:rPr"))
        if run_properties is not None:
            return deepcopy(run_properties)
    return None


def add_text_node(parent, tag: str, text: str) -> None:
    node = OxmlElement(tag)
    if text[:1].isspace() or text[-1:].isspace() or "  " in text:
        node.set(f"{{{XML_NS}}}space", "preserve")
    node.text = text
    parent.append(node)


def make_run(text: str, run_properties=None, deleted: bool = False):
    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(deepcopy(run_properties))
    add_text_node(run, "w:delText" if deleted else "w:t", text)
    return run


def clear_paragraph_content(paragraph_element) -> None:
    paragraph_properties = paragraph_element.find(qn("w:pPr"))
    for child in list(paragraph_element):
        if child is not paragraph_properties:
            paragraph_element.remove(child)


def set_clean_paragraph_text(paragraph_element, new_text: str) -> None:
    run_properties = first_run_properties(paragraph_element)
    clear_paragraph_content(paragraph_element)
    if new_text:
        paragraph_element.append(make_run(new_text, run_properties))


class ChangeId:
    def __init__(self, start: int = 1):
        self.value = start

    def next(self) -> str:
        current = self.value
        self.value += 1
        return str(current)


def make_change_wrapper(
    kind: str,
    text: str,
    change_ids: ChangeId,
    author: str,
    when: str,
    run_properties=None,
):
    wrapper = OxmlElement(f"w:{kind}")
    wrapper.set(qn("w:id"), change_ids.next())
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), when)
    wrapper.append(make_run(text, run_properties, deleted=(kind == "del")))
    return wrapper


def set_tracked_paragraph_text(
    paragraph_element,
    new_text: str,
    change_ids: ChangeId,
    author: str,
    when: str,
) -> None:
    old_text = paragraph_text(paragraph_element)
    if old_text == new_text:
        return

    run_properties = first_run_properties(paragraph_element)
    clear_paragraph_content(paragraph_element)
    matcher = difflib.SequenceMatcher(a=old_text, b=new_text, autojunk=False)
    for operation, a_start, a_end, b_start, b_end in matcher.get_opcodes():
        old_segment = old_text[a_start:a_end]
        new_segment = new_text[b_start:b_end]
        if operation == "equal":
            if old_segment:
                paragraph_element.append(make_run(old_segment, run_properties))
        elif operation == "delete":
            if old_segment:
                paragraph_element.append(
                    make_change_wrapper(
                        "del",
                        old_segment,
                        change_ids,
                        author,
                        when,
                        run_properties,
                    )
                )
        elif operation == "insert":
            if new_segment:
                paragraph_element.append(
                    make_change_wrapper(
                        "ins",
                        new_segment,
                        change_ids,
                        author,
                        when,
                        run_properties,
                    )
                )
        elif operation == "replace":
            if old_segment:
                paragraph_element.append(
                    make_change_wrapper(
                        "del",
                        old_segment,
                        change_ids,
                        author,
                        when,
                        run_properties,
                    )
                )
            if new_segment:
                paragraph_element.append(
                    make_change_wrapper(
                        "ins",
                        new_segment,
                        change_ids,
                        author,
                        when,
                        run_properties,
                    )
                )


def track_delete_paragraph_content(
    paragraph_element,
    change_ids: ChangeId,
    author: str,
    when: str,
) -> None:
    paragraph_properties = paragraph_element.find(qn("w:pPr"))
    content = [
        child
        for child in list(paragraph_element)
        if child is not paragraph_properties
    ]
    for child in content:
        paragraph_element.remove(child)
        wrapper = OxmlElement("w:del")
        wrapper.set(qn("w:id"), change_ids.next())
        wrapper.set(qn("w:author"), author)
        wrapper.set(qn("w:date"), when)
        wrapper.append(child)
        paragraph_element.append(wrapper)


def track_insert_paragraph_content(
    paragraph_element,
    change_ids: ChangeId,
    author: str,
    when: str,
) -> None:
    paragraph_properties = paragraph_element.find(qn("w:pPr"))
    content = [
        child
        for child in list(paragraph_element)
        if child is not paragraph_properties
    ]
    for child in content:
        paragraph_element.remove(child)
        wrapper = OxmlElement("w:ins")
        wrapper.set(qn("w:id"), change_ids.next())
        wrapper.set(qn("w:author"), author)
        wrapper.set(qn("w:date"), when)
        wrapper.append(child)
        paragraph_element.append(wrapper)


def enable_tracking(document: Document) -> None:
    settings = document.settings.element
    if settings.find(qn("w:trackRevisions")) is None:
        settings.insert(0, OxmlElement("w:trackRevisions"))


def load_table_geometry_module():
    spec = importlib.util.spec_from_file_location(
        "codex_table_geometry",
        TABLE_GEOMETRY_HELPER,
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Unable to load table geometry helper")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def style_table_text(table) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    run.bold = row_index == 0
                    run.font.size = Pt(8.5)
                    run.font.name = "Times New Roman"
                    run_properties = run._element.get_or_add_rPr()
                    fonts = run_properties.find(qn("w:rFonts"))
                    if fonts is None:
                        fonts = OxmlElement("w:rFonts")
                        run_properties.insert(0, fonts)
                    fonts.set(qn("w:ascii"), "Times New Roman")
                    fonts.set(qn("w:hAnsi"), "Times New Roman")
                    fonts.set(qn("w:eastAsia"), "宋体")

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)


def insert_replacement_table(
    document: Document,
    anchor_element,
    caption_template,
    *,
    tracked: bool,
    change_ids: ChangeId | None = None,
    author: str = "",
    when: str = "",
) -> None:
    caption = deepcopy(caption_template)
    set_clean_paragraph_text(caption, REPLACEMENT_TABLE_CAPTION)

    table = document.add_table(
        rows=len(REPLACEMENT_TABLE_ROWS),
        cols=len(REPLACEMENT_TABLE_ROWS[0]),
    )
    table.style = "Table Grid"
    for row_index, row_values in enumerate(REPLACEMENT_TABLE_ROWS):
        for column_index, value in enumerate(row_values):
            table.cell(row_index, column_index).text = value

    geometry = load_table_geometry_module()
    content_width = geometry.section_content_width_dxa(document.sections[0])
    widths = geometry.column_widths_from_weights(
        [1.18, 1.05, 0.92],
        total_width_dxa=content_width,
    )
    geometry.apply_table_geometry(
        table,
        widths,
        table_width_dxa=content_width,
    )
    style_table_text(table)

    anchor_element.addnext(caption)
    caption.addnext(table._tbl)

    if tracked:
        if change_ids is None:
            raise ValueError("Tracked insertion requires change ids")
        track_insert_paragraph_content(caption, change_ids, author, when)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    track_insert_paragraph_content(
                        paragraph._p,
                        change_ids,
                        author,
                        when,
                    )


def insert_tracked_author_note(
    document: Document,
    anchor_element,
    change_ids: ChangeId,
    author: str,
    when: str,
) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(AUTHOR_CHECK_NOTE)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(192, 0, 0)
    run.font.name = "Times New Roman"
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    anchor_element.addnext(paragraph._p)
    track_insert_paragraph_content(
        paragraph._p,
        change_ids,
        author=author,
        when=when,
    )


def apply_clean_table_cell_replacements(document: Document) -> None:
    tables = list(document.tables)
    for table_index, row_index, column_index, new_text in TABLE_CELL_REPLACEMENTS:
        paragraph = tables[table_index].cell(row_index, column_index).paragraphs[0]
        set_clean_paragraph_text(paragraph._p, new_text)


def apply_tracked_table_cell_replacements(
    document: Document,
    change_ids: ChangeId,
    author: str,
    when: str,
) -> None:
    tables = list(document.tables)
    for table_index, row_index, column_index, new_text in TABLE_CELL_REPLACEMENTS:
        paragraph = tables[table_index].cell(row_index, column_index).paragraphs[0]
        set_tracked_paragraph_text(
            paragraph._p,
            new_text,
            change_ids,
            author=author,
            when=when,
        )


def apply_clean_changes(
    source: Path,
    output: Path,
    replacements: dict[int, str],
) -> None:
    document = Document(str(source))
    apply_clean_table_cell_replacements(document)
    body = document.element.body
    children = list(body.iterchildren())
    anchor_element = children[TABLE_ANCHOR_BLOCK]
    caption_template = children[130]
    for index, child in enumerate(children):
        if local_name(child) == "sectPr":
            continue
        if index < BODY_START_BLOCK or index in DELETE_BLOCKS:
            body.remove(child)
            continue
        if index in replacements:
            if local_name(child) != "p":
                raise ValueError(f"Block {index} is not a paragraph")
            set_clean_paragraph_text(child, replacements[index])
    insert_replacement_table(
        document,
        anchor_element,
        caption_template,
        tracked=False,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def apply_tracked_changes(
    source: Path,
    output: Path,
    replacements: dict[int, str],
    author: str,
) -> None:
    document = Document(str(source))
    enable_tracking(document)
    body = document.element.body
    children = list(body.iterchildren())
    anchor_element = children[TABLE_ANCHOR_BLOCK]
    caption_template = children[130]
    change_ids = ChangeId()
    when = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    apply_tracked_table_cell_replacements(
        document,
        change_ids,
        author=author,
        when=when,
    )

    for index, child in enumerate(children):
        if local_name(child) == "sectPr":
            continue
        if index < BODY_START_BLOCK:
            body.remove(child)
            continue
        if index in DELETE_BLOCKS:
            if local_name(child) != "p":
                raise ValueError(f"Deleted block {index} is not a paragraph")
            track_delete_paragraph_content(
                child,
                change_ids,
                author=author,
                when=when,
            )
            continue
        if index in replacements:
            if local_name(child) != "p":
                raise ValueError(f"Block {index} is not a paragraph")
            set_tracked_paragraph_text(
                child,
                replacements[index],
                change_ids,
                author=author,
                when=when,
            )

    insert_replacement_table(
        document,
        anchor_element,
        caption_template,
        tracked=True,
        change_ids=change_ids,
        author=author,
        when=when,
    )
    insert_tracked_author_note(
        document,
        children[118],
        change_ids,
        author=author,
        when=when,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))


def load_replacements(path: Path) -> dict[int, str]:
    data = json.loads(path.read_text(encoding="utf-8"))
    replacements = {int(key): value for key, value in data.items()}
    invalid = sorted(index for index in replacements if index < BODY_START_BLOCK)
    if invalid:
        raise ValueError(f"Replacement blocks precede manuscript body: {invalid}")
    return replacements


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--replacements", type=Path, required=True)
    parser.add_argument("--clean-out", type=Path, required=True)
    parser.add_argument("--tracked-out", type=Path, required=True)
    parser.add_argument("--author", default="Codex逐句修订")
    args = parser.parse_args()

    replacements = load_replacements(args.replacements)
    apply_clean_changes(args.source, args.clean_out, replacements)
    apply_tracked_changes(
        args.source,
        args.tracked_out,
        replacements,
        author=args.author,
    )
    print(f"clean={args.clean_out}")
    print(f"tracked={args.tracked_out}")
    print(f"paragraph_replacements={len(replacements)}")


if __name__ == "__main__":
    main()
