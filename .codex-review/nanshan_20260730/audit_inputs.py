from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from lxml import etree


PDF_PATH = Path(r"C:\Users\89418\Desktop\《南山教育》稿件模板.pdf")
CLEAN_PATH = Path(
    r"D:\AI阅卷系统_工作机版_v1.5.0\论文修订输出"
    r"\初中数学AI阅卷论文_逐句修订净稿.docx"
)
SOURCE_PATH = Path(
    r"C:\Users\89418\Desktop\论文材料"
    r"\初中数学纸笔测验中AI阅卷的题型适配与教师复核机制"
    r"——基于两种多模态批改模式的探索性案例研究.docx"
)
OUT_DIR = Path(r"D:\AI阅卷系统_工作机版_v1.5.0\.codex-review\nanshan_20260730")


def length_value(value):
    if value is None:
        return None
    return {
        "emu": int(value),
        "pt": round(value.pt, 3),
        "cm": round(value.cm, 3),
    }


def paragraph_info(paragraph, index):
    p_format = paragraph.paragraph_format
    runs = []
    for run in paragraph.runs:
        rpr = run._element.rPr
        east_asia = None
        if rpr is not None and rpr.rFonts is not None:
            east_asia = rpr.rFonts.get(qn("w:eastAsia"))
        runs.append(
            {
                "text": run.text,
                "font": run.font.name,
                "east_asia": east_asia,
                "size_pt": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
                "superscript": run.font.superscript,
            }
        )
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "alignment": str(paragraph.alignment),
        "left_indent": length_value(p_format.left_indent),
        "right_indent": length_value(p_format.right_indent),
        "first_line_indent": length_value(p_format.first_line_indent),
        "space_before": length_value(p_format.space_before),
        "space_after": length_value(p_format.space_after),
        "line_spacing": (
            round(p_format.line_spacing, 3)
            if isinstance(p_format.line_spacing, float)
            else length_value(p_format.line_spacing)
        ),
        "line_spacing_rule": str(p_format.line_spacing_rule),
        "keep_with_next": p_format.keep_with_next,
        "page_break_before": p_format.page_break_before,
        "runs": runs,
    }


def table_info(table, index):
    rows = []
    for row in table.rows:
        rows.append([cell.text for cell in row.cells])
    return {
        "index": index,
        "style": table.style.name if table.style else None,
        "rows": rows,
    }


def inspect_docx(path: Path):
    document = Document(path)
    sections = []
    for index, section in enumerate(document.sections):
        sections.append(
            {
                "index": index,
                "page_width": length_value(section.page_width),
                "page_height": length_value(section.page_height),
                "top_margin": length_value(section.top_margin),
                "bottom_margin": length_value(section.bottom_margin),
                "left_margin": length_value(section.left_margin),
                "right_margin": length_value(section.right_margin),
                "header_distance": length_value(section.header_distance),
                "footer_distance": length_value(section.footer_distance),
                "orientation": str(section.orientation),
            }
        )
    shapes = []
    for index, shape in enumerate(document.inline_shapes):
        shapes.append(
            {
                "index": index,
                "width": length_value(shape.width),
                "height": length_value(shape.height),
                "type": str(shape.type),
            }
        )
    return {
        "path": str(path),
        "core_properties": {
            "title": document.core_properties.title,
            "subject": document.core_properties.subject,
            "author": document.core_properties.author,
            "last_modified_by": document.core_properties.last_modified_by,
            "keywords": document.core_properties.keywords,
        },
        "sections": sections,
        "paragraphs": [
            paragraph_info(paragraph, index)
            for index, paragraph in enumerate(document.paragraphs)
        ],
        "tables": [
            table_info(table, index) for index, table in enumerate(document.tables)
        ],
        "inline_shapes": shapes,
        "style_counts": Counter(
            paragraph.style.name if paragraph.style else ""
            for paragraph in document.paragraphs
        ),
        "all_story_text": extract_all_story_text(path),
    }


def extract_all_story_text(path: Path):
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    }
    parts = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            try:
                root = etree.fromstring(archive.read(name))
            except etree.XMLSyntaxError:
                continue
            texts = root.xpath("//w:t/text() | //w:delText/text()", namespaces=namespaces)
            joined = "\n".join(text for text in texts if text)
            if joined:
                parts[name] = joined
    return parts


def inspect_pdf(path: Path):
    pages = []
    font_counts = Counter()
    size_counts = Counter()
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            chars = page.chars
            for char in chars:
                font_counts[char.get("fontname")] += 1
                size = char.get("size")
                if size is not None:
                    size_counts[round(float(size), 2)] += 1
            pages.append(
                {
                    "page_number": page_number,
                    "width": page.width,
                    "height": page.height,
                    "text": page.extract_text(x_tolerance=2, y_tolerance=3) or "",
                    "words": page.extract_words(
                        x_tolerance=2,
                        y_tolerance=3,
                        extra_attrs=["fontname", "size"],
                    ),
                }
            )
    return {
        "path": str(path),
        "pages": pages,
        "font_counts": font_counts,
        "size_counts": size_counts,
    }


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pdf_info = inspect_pdf(PDF_PATH)
    clean_info = inspect_docx(CLEAN_PATH)
    source_info = inspect_docx(SOURCE_PATH)

    (OUT_DIR / "template_pdf.json").write_text(
        json.dumps(pdf_info, ensure_ascii=False, indent=2, default=dict),
        encoding="utf-8",
    )
    (OUT_DIR / "clean_docx.json").write_text(
        json.dumps(clean_info, ensure_ascii=False, indent=2, default=dict),
        encoding="utf-8",
    )
    (OUT_DIR / "source_docx.json").write_text(
        json.dumps(source_info, ensure_ascii=False, indent=2, default=dict),
        encoding="utf-8",
    )

    print(
        {
            "pdf_pages": len(pdf_info["pages"]),
            "clean_paragraphs": len(clean_info["paragraphs"]),
            "clean_tables": len(clean_info["tables"]),
            "clean_images": len(clean_info["inline_shapes"]),
            "source_paragraphs": len(source_info["paragraphs"]),
            "source_tables": len(source_info["tables"]),
        }
    )


if __name__ == "__main__":
    main()
