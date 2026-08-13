from __future__ import annotations

import importlib.util
import posixpath
import re
from pathlib import Path, PurePosixPath
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree


INPUT_PATH = Path(
    r"D:\AI阅卷系统_工作机版_v1.5.0\论文修订输出"
    r"\初中数学AI阅卷论文_逐句修订净稿.docx"
)
WORK_DIR = Path(
    r"D:\AI阅卷系统_工作机版_v1.5.0\.codex-review\nanshan_20260730"
)
DRAFT_PATH = WORK_DIR / "nanshan_before_media_cleanup.docx"
OUTPUT_PATH = WORK_DIR / "nanshan_before_scrub.docx"
TABLE_GEOMETRY_HELPER = Path(
    r"C:\Users\89418\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.727.11326\skills\documents\scripts\table_geometry.py"
)

TITLE = "初中数学AI阅卷的题型误差与教师复核——基于78份七年级纸笔答卷的实践分析"
AUTHOR_LINE = "深圳大学附属教育集团外国语中学　　贾浩然"
ABSTRACT = (
    "为判断AI阅卷在初中数学纸笔测验中的实际适用范围，以某校七年级78份配对答卷为样本，"
    "比较整卷批改和混合分批批改两种方式，并对照教师手改分分析总分误差、题型误差和复核日志。"
    "原卷为总分57分的校内非标准化试卷，比较分析统一采用系统中的100分制评分口径。结果显示："
    "整卷模式流程最终分的MAE为4.49，整体偏高2.67分；混合模式流程最终分的MAE为5.85，总分"
    "差异虽不显著，但不同题型的正负偏差存在抵消。两种模式在Q10作图题上的相对MAE均约为"
    "24%，系统“需要复核”标记也未覆盖全部高误差题目。由此可见，AI阅卷不能只依据总分相关性"
    "或模型置信度判断是否可用。选择、填空、计算、证明和作图题应分别设置复核要求，其中作图题"
    "仍应由教师逐项核对并确定得分。"
)
KEYWORDS = "人工智能　　初中数学　　智能阅卷　　题型误差　　教师复核"
CAPTION_TEXTS = {
    "表1 试卷题型结构与评分任务特征",
    "表2 AI 阅卷的图像输入、模型参数与日志记录设置",
    "表3 AI 原始分与流程最终分相对于教师手改分的总分差异及一致性",
    "图1 流程最终分与教师手改分的 Bland-Altman 一致性分析",
    "表4 排除 Q10 作图题前后流程最终分与教师手改分的差异及一致性",
    "图2 排除 Q10 前后流程最终分的 Bias 变化",
    "表5 不同题型下两种模式流程最终分的 MAE、相对 MAE 与 Bias",
    "图3 不同题型下两种 AI 批改模式的标准化 MAE",
    "图4 Q10 作图题的两个脱敏作答案例（原卷题面9分；统计按100分制计11分）",
    "表6 两种模式中“需要复核”字段及相应改分记录",
    "表7 本次题型结果与暂定处理方式",
}

CONTACT_FIELDS = [
    ("作者姓名：", "贾浩然"),
    ("作者单位：", "深圳大学附属教育集团外国语中学"),
    (
        "通讯地址：",
        "广东省深圳市南山区桃源街道留仙大道3633号深圳大学附属教育集团外国语中学"
        "（新校区）（邮编：518071）",
    ),
    ("联系电话：", "13201469126"),
    ("QQ或微信：", "894189001（QQ）"),
    ("电子邮箱：", "894189001@qq.com"),
]

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
BODY_SIZE = 10.5
HEADING1_SIZE = 12.0
TABLE_SIZE = 7.5
TITLE_SIZE = 22.0

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def load_table_geometry_module():
    spec = importlib.util.spec_from_file_location(
        "codex_table_geometry_nanshan",
        TABLE_GEOMETRY_HELPER,
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Unable to load table geometry helper")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    return Paragraph(new_element, paragraph._parent)


def clear_paragraph_content(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph_content(paragraph)
    paragraph.add_run(text)


def set_labeled_text(paragraph: Paragraph, label: str, content: str) -> None:
    clear_paragraph_content(paragraph)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    set_run_font(label_run, HEADING_FONT, HEADING_FONT, BODY_SIZE)
    content_run = paragraph.add_run(content)
    set_run_font(content_run, BODY_FONT, LATIN_FONT, BODY_SIZE)


def set_run_font(
    run,
    east_asia: str,
    latin: str,
    size_pt: float,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)


def reset_paragraph_format(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = False
    fmt.keep_together = False
    fmt.page_break_before = False


def set_multiple_spacing(paragraph: Paragraph, value: float = 1.25) -> None:
    paragraph.paragraph_format.line_spacing = value
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def apply_font_to_paragraph(
    paragraph: Paragraph,
    east_asia: str,
    latin: str,
    size_pt: float,
    *,
    bold: bool | None = None,
) -> None:
    for run in paragraph.runs:
        set_run_font(
            run,
            east_asia,
            latin,
            size_pt,
            bold=bold if bold is not None else run.bold,
        )


def add_named_style(document: Document, name: str):
    if name in document.styles:
        return document.styles[name]
    return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def configure_named_styles(document: Document) -> None:
    specifications = {
        "Nanshan Title": (HEADING_FONT, HEADING_FONT, TITLE_SIZE, True),
        "Nanshan Author": (BODY_FONT, LATIN_FONT, BODY_SIZE, False),
        "Nanshan Body": (BODY_FONT, LATIN_FONT, BODY_SIZE, False),
        "Nanshan Heading 1": (HEADING_FONT, HEADING_FONT, HEADING1_SIZE, False),
        "Nanshan Heading 2": (HEADING_FONT, HEADING_FONT, BODY_SIZE, False),
        "Nanshan Heading 3": (BODY_FONT, LATIN_FONT, BODY_SIZE, True),
        "Nanshan Heading 4": (BODY_FONT, LATIN_FONT, BODY_SIZE, False),
        "Nanshan Caption": (HEADING_FONT, HEADING_FONT, TABLE_SIZE, True),
        "Nanshan Reference": (BODY_FONT, LATIN_FONT, TABLE_SIZE, False),
        "Nanshan Contact": (BODY_FONT, LATIN_FONT, BODY_SIZE, False),
    }
    for name, (east_asia, latin, size, bold) in specifications.items():
        style = add_named_style(document, name)
        style.font.name = latin
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), east_asia)
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)
        rfonts.set(qn("w:cs"), latin)


def has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def has_page_break(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:br[@w:type='page']"))


def replace_body_references(text: str) -> str:
    text = text.replace("Bland–Altman", "Bland-Altman")
    text = re.sub(r"(表|图)\s+(\d+)", r"\1\2", text)
    if text not in CAPTION_TEXTS:
        text = re.sub(r"(表|图)(\d+)\s+(?=[\u4e00-\u9fff])", r"\1\2", text)
    return text


def find_reference_paragraphs(document: Document) -> list[Paragraph]:
    return [
        paragraph
        for paragraph in document.paragraphs
        if re.match(r"^\[\d+\]\s*", paragraph.text)
    ]


def normalize_reference_texts(document: Document) -> None:
    for paragraph in find_reference_paragraphs(document):
        text = re.sub(r"^(\[\d+\])\s+", r"\1", paragraph.text)
        if text.startswith(("[19]", "[20]", "[21]")):
            text = text.replace("[J]. arXiv preprint arXiv:", "[EB/OL]. arXiv:")
        set_plain_text(paragraph, text)


def apply_scoring_scale_notes(document: Document) -> None:
    # Table 1 and Table 5 use the 100-point analysis rubric. The short parenthetical
    # keeps the original 9-point paper label visible without changing any statistics.
    document.tables[0].cell(3, 1).text = "作图题，11分（原卷9分）"
    document.tables[4].cell(3, 0).text = "Q10作图题（原卷9分）"


def superscript_body_citations(paragraph: Paragraph) -> int:
    text = paragraph.text
    pattern = re.compile(r"\[\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*\]")
    matches = list(pattern.finditer(text))
    if not matches:
        return 0
    clear_paragraph_content(paragraph)
    cursor = 0
    for match in matches:
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, BODY_FONT, LATIN_FONT, BODY_SIZE)
        citation = paragraph.add_run(match.group(0))
        set_run_font(citation, BODY_FONT, LATIN_FONT, TABLE_SIZE)
        citation.font.superscript = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, BODY_FONT, LATIN_FONT, BODY_SIZE)
    return len(matches)


def format_title(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Title"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    apply_font_to_paragraph(
        paragraph,
        HEADING_FONT,
        HEADING_FONT,
        TITLE_SIZE,
        bold=True,
    )


def format_author(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Author"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    set_multiple_spacing(paragraph)
    apply_font_to_paragraph(paragraph, BODY_FONT, LATIN_FONT, BODY_SIZE)


def format_abstract_or_keywords(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Body"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
    set_multiple_spacing(paragraph)


def format_body(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Body"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
    set_multiple_spacing(paragraph)
    apply_font_to_paragraph(paragraph, BODY_FONT, LATIN_FONT, BODY_SIZE)


def format_heading(paragraph: Paragraph, level: int) -> None:
    paragraph.style = f"Nanshan Heading {level}"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    size = HEADING1_SIZE if level == 1 else BODY_SIZE
    paragraph.paragraph_format.first_line_indent = Pt(size * 2)
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 3)
    paragraph.paragraph_format.keep_with_next = True
    set_multiple_spacing(paragraph)
    if level in (1, 2):
        apply_font_to_paragraph(
            paragraph,
            HEADING_FONT,
            HEADING_FONT,
            size,
            bold=False,
        )
    elif level == 3:
        apply_font_to_paragraph(
            paragraph,
            BODY_FONT,
            LATIN_FONT,
            size,
            bold=True,
        )
    else:
        apply_font_to_paragraph(
            paragraph,
            BODY_FONT,
            LATIN_FONT,
            size,
            bold=False,
        )


def format_caption(paragraph: Paragraph, is_table: bool) -> None:
    paragraph.style = "Nanshan Caption"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = is_table
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    apply_font_to_paragraph(
        paragraph,
        HEADING_FONT,
        HEADING_FONT,
        TABLE_SIZE,
        bold=True,
    )


def format_image_paragraph(paragraph: Paragraph) -> None:
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)


def format_reference_heading(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Body"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    set_multiple_spacing(paragraph)
    apply_font_to_paragraph(
        paragraph,
        BODY_FONT,
        LATIN_FONT,
        BODY_SIZE,
        bold=False,
    )


def format_reference(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Reference"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    apply_font_to_paragraph(
        paragraph,
        BODY_FONT,
        LATIN_FONT,
        TABLE_SIZE,
        bold=False,
    )


def format_contact(paragraph: Paragraph) -> None:
    paragraph.style = "Nanshan Contact"
    reset_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_multiple_spacing(paragraph)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        tag = qn(f"w:{edge}")
        element = tc_mar.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_border(parent_pr, edge: str, *, value: str, size: int = 0) -> None:
    borders_tag = "w:tblBorders" if parent_pr.tag == qn("w:tblPr") else "w:tcBorders"
    borders = parent_pr.find(qn(borders_tag))
    if borders is None:
        borders = OxmlElement(borders_tag)
        parent_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), value)
    if value == "single":
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def apply_three_line_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    for edge in ("left", "right", "insideH", "insideV"):
        set_border(tbl_pr, edge, value="nil")
    set_border(tbl_pr, "top", value="single", size=8)
    set_border(tbl_pr, "bottom", value="single", size=8)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        set_border(tc_pr, "bottom", value="single", size=6)


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def remove_fixed_row_height(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for element in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(element)


def set_xml_run_size(element, half_points: int) -> None:
    for run in element.xpath(".//w:r"):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            run.insert(0, rpr)
        for tag in ("w:sz", "w:szCs"):
            size = rpr.find(qn(tag))
            if size is None:
                size = OxmlElement(tag)
                rpr.append(size)
            size.set(qn("w:val"), str(half_points))


def format_tables(document: Document) -> None:
    geometry = load_table_geometry_module()
    content_width = geometry.section_content_width_dxa(document.sections[0])
    weights = [
        [1.0, 1.9, 2.2, 2.4, 3.1],
        [1.2, 2.5, 2.0, 3.0],
        [1.8, 2.0, 1.0, 1.0, 1.0, 1.0, 2.1, 1.0, 1.2],
        [1.6, 1.4, 1.0, 1.0, 1.0, 2.0, 1.0, 1.0],
        [2.1, 0.8, 1.3, 1.0, 1.3, 1.0, 2.5],
        [0.8, 0.9, 1.5, 1.1, 1.8, 1.4, 2.7],
        [1.2, 1.05, 0.92],
    ]
    if len(document.tables) != len(weights):
        raise AssertionError(
            f"Expected {len(weights)} tables, found {len(document.tables)}"
        )
    for table_index, (table, table_weights) in enumerate(
        zip(document.tables, weights)
    ):
        widths = geometry.column_widths_from_weights(
            table_weights,
            total_width_dxa=content_width,
        )
        geometry.apply_table_geometry(
            table,
            widths,
            table_width_dxa=content_width,
        )
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        apply_three_line_borders(table)
        mark_repeat_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            remove_fixed_row_height(row)
            prevent_row_split(row)
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    reset_paragraph_format(paragraph)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.line_spacing_rule = (
                        WD_LINE_SPACING.SINGLE
                    )
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        or len(paragraph.text.strip()) <= 14
                        or column_index < 2
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    apply_font_to_paragraph(
                        paragraph,
                        HEADING_FONT if row_index == 0 else BODY_FONT,
                        HEADING_FONT if row_index == 0 else LATIN_FONT,
                        TABLE_SIZE,
                        bold=True if row_index == 0 else False,
                    )
                set_xml_run_size(cell._tc, int(TABLE_SIZE * 2))


def append_contact_fields(
    document: Document, anchor: Paragraph
) -> list[Paragraph]:
    spacer = insert_paragraph_after(anchor)
    current = spacer
    contacts: list[Paragraph] = []
    for label, value in CONTACT_FIELDS:
        current = insert_paragraph_after(current)
        clear_paragraph_content(current)
        label_run = current.add_run(label)
        set_run_font(
            label_run,
            HEADING_FONT,
            HEADING_FONT,
            BODY_SIZE,
            bold=True,
        )
        value_run = current.add_run(value)
        set_run_font(value_run, BODY_FONT, LATIN_FONT, BODY_SIZE)
        contacts.append(current)
    return contacts


def relationship_source_part(rels_name: str) -> str | None:
    path = PurePosixPath(rels_name)
    if path.name == ".rels":
        return None
    if path.parent.name != "_rels" or not path.name.endswith(".rels"):
        return None
    source_name = path.name[: -len(".rels")]
    return str(path.parent.parent / source_name)


def remove_orphan_media(source: Path, destination: Path) -> dict[str, int]:
    with ZipFile(source) as archive:
        infos = archive.infolist()
        payloads = {info.filename: archive.read(info.filename) for info in infos}

    removed_relationships = 0
    used_media: set[str] = set()
    updated_payloads = dict(payloads)
    for name, data in list(payloads.items()):
        if not name.endswith(".rels"):
            continue
        source_part = relationship_source_part(name)
        try:
            rel_root = etree.fromstring(data)
        except etree.XMLSyntaxError:
            continue
        source_xml = payloads.get(source_part or "")
        used_ids: set[str] = set()
        if source_xml:
            try:
                source_root = etree.fromstring(source_xml)
                for element in source_root.iter():
                    for attribute, value in element.attrib.items():
                        if etree.QName(attribute).namespace == R_NS:
                            used_ids.add(value)
            except etree.XMLSyntaxError:
                pass

        changed = False
        source_dir = posixpath.dirname(source_part or "")
        for relationship in list(rel_root):
            if not relationship.tag.endswith("Relationship"):
                continue
            rel_type = relationship.get("Type", "")
            if not rel_type.endswith("/image"):
                continue
            rel_id = relationship.get("Id", "")
            if source_part and rel_id not in used_ids:
                rel_root.remove(relationship)
                removed_relationships += 1
                changed = True
                continue
            target = relationship.get("Target")
            if target and not target.startswith(("http://", "https://")):
                used_media.add(
                    posixpath.normpath(posixpath.join(source_dir, target))
                )
        if changed:
            updated_payloads[name] = etree.tostring(
                rel_root,
                xml_declaration=True,
                encoding="UTF-8",
                standalone=True,
            )

    removed_media = 0
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(destination, "w", compression=ZIP_DEFLATED) as output:
        for info in infos:
            name = info.filename
            if name.startswith("word/media/") and name not in used_media:
                removed_media += 1
                continue
            output.writestr(info, updated_payloads[name])
    return {
        "removed_image_relationships": removed_relationships,
        "removed_media_parts": removed_media,
    }


def build() -> dict[str, object]:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    document = Document(INPUT_PATH)
    original = list(document.paragraphs)
    if len(original) < 113:
        raise AssertionError(f"Unexpected paragraph count: {len(original)}")
    configure_named_styles(document)

    set_plain_text(original[0], TITLE)
    remove_paragraph(original[1])
    set_labeled_text(original[2], "【摘  要】", ABSTRACT)
    set_labeled_text(original[3], "【关键词】", KEYWORDS)
    for paragraph in original[4:7]:
        remove_paragraph(paragraph)

    author = insert_paragraph_after(original[0])
    set_plain_text(author, AUTHOR_LINE)

    # The Q10 image is retained. Its printed 9-point label belongs to the original
    # 57-point paper, while the comparison analysis uses the system's 100-point
    # scoring rubric in which Q10 is assigned 11 points.

    replacements = {
        24: (
            "研究纳入某校七年级两个教学班一次阶段测试的78份答卷。纳入条件是同一份答卷同时具有"
            "任课教师手改记录、整卷AI批改记录和混合分批AI批改记录，三类记录可按同一份答卷一一"
            "对应。测试基于北师大版七年级下册相关内容，覆盖轴对称、角平分线性质、最短路径、一次"
            "函数图像应用和几何证明。整卷与混合模式均按每份答卷18个题目/小问评分单元保存日志，"
            "各形成78×18=1404条学生—题目/小问记录。研究文本中的姓名、班级及学号均已匿名化。"
            "试卷题型及评分任务见表1。"
        ),
        29: (
            "两种模式按以下方式组织输入图像。整卷模式将单名学生的整卷或整页图像输入"
            "doubao-seed-2.0-pro，原答卷中的题目位置、跨区域作答和步骤顺序随页面一并呈现。"
            "混合分批模式先裁切同一道题的学生作答区域，再将多名学生的同题区域拼接；选择、填空题"
            "每批最多15人，解答题每批最多4人，拼接完成后输入参数相同的doubao-seed-2.0-pro。"
            "混合模式不是对每名学生逐题单独调用模型，而是在一次调用中处理同题的多名学生作答，"
            "其设计目的是限制单次输入范围，并减少调用批次以控制调用时间和token支出。本文比较的"
            "是上述两种输入组织方式，现有结果未单独报告调用时间或token成本的比较数据。具体设置"
            "见表2。"
        ),
        33: (
            "研究区分三类分数。“教师手改分”是任课教师依据统一评分细则在纸笔批改后登记的分数；"
            "“AI原始分”是模型首次输出的分数；“流程最终分”是完成本轮系统复核流程后导出的成绩，"
            "可能沿用AI原始分，也可能在教师查看后调整。原卷为校内非标准化阶段测试，卷面各题标注"
            "分值合计57分；为便于系统统一评分，本研究的比较分析采用100分制评分细则，Q10在该"
            "口径下满分为11分，而图4保留的原卷题面仍显示9分。本文表1、表3—表5、表7以及各项"
            "误差指标均采用100分制口径。由于未设置第二评分者，教师手改分只作为本次课堂评分的专业"
            "参照，不被视为学生表现的绝对真值。本文据此考察两种AI批改结果与该教师手改分的差异"
            "及一致程度。"
        ),
        41: (
            "表3显示，整卷模式流程最终分的MAE为4.49、RMSE为5.90、ICC为0.972，但平均高于"
            "教师手改分2.67分，配对差异显著。混合模式流程最终分的MAE为5.85、RMSE为7.95、"
            "ICC为0.944，平均低于教师手改分1.48分，配对差异不显著。在本样本中，整卷模式的"
            "绝对误差较小，但存在整体偏高；混合模式未检出显著的总分差异，也不能据此解释为没有"
            "偏差。表3未对两种模式各指标之差进行直接检验，因此不据这些数值判定一种模式总体优于"
            "另一种模式。两种模式最终分与教师手改分的一致性分布见图1。"
        ),
        47: (
            "表4显示，排除Q10后，整卷模式的Bias由+2.67收窄至+0.32，MAE由4.49降至2.67，"
            "配对差异不再显著；整卷总分偏高与Q10的正向偏差密切相关。表5显示，混合模式在Q10上"
            "的正向偏差与填空题、选择题和Q11的负向偏差部分抵消。混合模式的Bias由-1.48变为"
            "-3.16，且配对差异转为显著。因而，Q10是否计入会明显改变两种模式的总分Bias及配对"
            "检验结果；总分层面“未检出显著差异”不能被解释为各题型均无系统性偏高或偏低。"
            "Bias变化见图2。"
        ),
        59: (
            "表6显示，整卷模式在1404条日志中标记18条“需要复核”，标记率为1.28%；其中1条记录"
            "的流程最终分下调，其余被标记记录未发生分数调整。混合模式标记11条，标记率为0.78%；"
            "其中4条记录的流程最终分均上调，其余被标记记录未发生分数调整。整卷模式的标记主要位"
            "于Q10、Q11和Q12，混合模式的标记主要位于Q3—Q9，Q10和Q12未被混合模式标记。表5中"
            "混合模式在Q10和Q12的相对MAE分别为24.1%和17.9%，可见模型置信度触发的标记分布与"
            "题型相对误差的分布并不一致。本文没有“每条记录是否存在评分误差”与“是否被标记”的完整"
            "交叉数据，因此不能据表6计算该字段的敏感度、漏标率或误标率。"
        ),
        62: (
            "图4 Q10 作图题的两个脱敏作答案例"
            "（原卷题面9分；统计按100分制计11分）"
        ),
        63: (
            "图4保留了原卷题面“本小题9分”的标注；本文统计采用AI阅卷流程中的100分制评分"
            "细则，Q10按11分计。两份脱敏作答中，可以直接观察到多条相交线段、点名和辅助作图"
            "痕迹，部分线条或标注较浅，不同线条在网格区域内相互重叠。图下“整卷偏宽与混合偏严”"
            "“浅色作图痕迹识别风险”是复核时对现象的简写，并非由这两个案例单独验证的因果结论。"
            "图中未列出两个案例的教师手改分、AI原始分、流程最终分及逐项模型输出，因此不能仅凭"
            "作答图像确定模型遗漏了哪一条线、采用了何种判断路径，也不能把评分分歧确定归因于整卷"
            "页面、裁切缩放或多名学生拼接。结合表5中整卷和混合模式在Q10上分别为24.4%和24.1%"
            "的相对MAE，这两个案例只能说明：识别到点线痕迹与正确确认点线归属、交点位置及几何"
            "关系并不是同一个评分步骤；案例本身不能证明某一种输入方式必然造成偏宽或偏严。"
        ),
        70: (
            "表5中，Q10在整卷和混合模式下的标准化MAE分别为24.4%和24.1%，明显高于本试卷其他"
            "题型；排除Q10后，整卷模式的Bias又由+2.67降至+0.32。图4中可以看到点线重叠、浅色"
            "线条和文字标注等复核难点；图下“偏宽/偏严”等标签只作现象提示，不能替代案例级分数"
            "证据。两种模式使用同一次300 dpi扫描，且均未启用图像增强，这使扫描来源和名义分辨率"
            "保持一致，但不能据此排除图像因素；手写深浅、线条重叠、裁切和缩放仍可能影响作答内容"
            "是否清楚。对Q10这类题，教师仍需按评分细则逐项核对点线位置和几何关系。AI可以协助"
            "呈现原图、裁切图和疑似痕迹，不宜单独决定最终得分。"
        ),
        72: (
            "表6显示，整卷模式在1404条日志中标记18条需复核记录，标记率为1.28%，复核后仅有1条"
            "下调；混合模式标记11条，标记率为0.78%，其中4条复核后上调。更值得注意的是，混合"
            "模式的标记主要分布在Q3—Q9，却没有覆盖表5中误差较高的Q10和Q12。本研究没有对模型"
            "报告的置信度进行概率校准，也没有对同一图像进行重复调用，因此该字段不能解释为答对"
            "概率，也不能用于判断输出是否稳定。实际使用时，它只能作为追加复核的提示。就本次结果"
            "而言，Q10应直接进入教师复核，Q6—Q9需重点核对等价表达和作废答案，Q11、Q12则需查看"
            "关键步骤或证明过程是否完整；系统标记只是在这些检查之外增加一条提醒。各题型的暂定"
            "处理方式见表7。"
        ),
        77: (
            "本研究只分析了同一学校两个教学班的一次测试，共78份完整配对答卷，表5中的每类题型又"
            "只包含有限题目；图4的两个Q10案例是从分歧记录中目的性选取的，不能代表该题所有错误。"
            "教师手改没有设置第二评分者，研究者同时参与系统设计、课堂评分和结果解释，因此教师分"
            "也不能视为不受影响的绝对标准。模型未对同一图像重复调用，top_p、top_k采用服务端"
            "默认参数，表6中的标记率也不足以评价复核提示的检出能力。后续研究需要在不同学校和"
            "试卷中预先设定待检验的题型处理方式，加入独立教师双评分和分歧仲裁，并分别改变裁切"
            "范围、拼图人数和重复调用次数，检验本次观察到的差异能否再次出现。"
        ),
    }
    for index, replacement in replacements.items():
        set_plain_text(original[index], replacement)

    apply_scoring_scale_notes(document)

    for paragraph in list(document.paragraphs):
        text = replace_body_references(paragraph.text)
        if text != paragraph.text:
            set_plain_text(paragraph, text)

    normalize_reference_texts(document)
    references = find_reference_paragraphs(document)
    if len(references) != 30:
        raise AssertionError(f"Expected 30 references, found {len(references)}")
    contacts = append_contact_fields(document, references[-1])

    # Remove layout-only empty paragraphs; image-bearing paragraphs are retained.
    for paragraph in list(document.paragraphs):
        if (
            not paragraph.text.strip()
            and not has_drawing(paragraph)
            and not has_page_break(paragraph)
            and paragraph not in contacts
        ):
            remove_paragraph(paragraph)

    in_references = False
    contact_labels = tuple(label for label, _ in CONTACT_FIELDS)
    citation_count = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == TITLE:
            format_title(paragraph)
            continue
        if text == AUTHOR_LINE:
            format_author(paragraph)
            continue
        if text.startswith("【摘  要】") or text.startswith("【关键词】"):
            format_abstract_or_keywords(paragraph)
            continue
        if text.startswith(contact_labels):
            format_contact(paragraph)
            continue
        if text == "参考文献":
            in_references = True
            format_reference_heading(paragraph)
            continue
        if in_references:
            if re.match(r"^\[\d+\]", text):
                format_reference(paragraph)
            continue
        if has_drawing(paragraph):
            format_image_paragraph(paragraph)
            continue
        if re.match(r"^[一二三四五六七八九十]+、", text):
            format_heading(paragraph, 1)
        elif re.match(r"^（[一二三四五六七八九十]+）", text):
            format_heading(paragraph, 2)
        elif re.match(r"^\d+\.", text):
            format_heading(paragraph, 3)
        elif re.match(r"^（\d+）", text):
            format_heading(paragraph, 4)
        elif text in CAPTION_TEXTS and text.startswith("表"):
            format_caption(paragraph, is_table=True)
        elif text in CAPTION_TEXTS and text.startswith("图"):
            format_caption(paragraph, is_table=False)
        else:
            format_body(paragraph)
            citation_count += superscript_body_citations(paragraph)

    format_tables(document)

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # The template does not prescribe margins; retain the manuscript's existing
    # A4 margins rather than guessing from the cropped PDF example.

    document.core_properties.title = TITLE
    document.core_properties.subject = "《南山教育》投稿格式适配稿"
    document.core_properties.keywords = "人工智能；初中数学；智能阅卷；题型误差；教师复核"
    document.save(DRAFT_PATH)
    media_report = remove_orphan_media(DRAFT_PATH, OUTPUT_PATH)
    return {
        "output": str(OUTPUT_PATH),
        "paragraphs": len(Document(OUTPUT_PATH).paragraphs),
        "tables": len(Document(OUTPUT_PATH).tables),
        "images": len(Document(OUTPUT_PATH).inline_shapes),
        "references": len(find_reference_paragraphs(Document(OUTPUT_PATH))),
        "superscript_citations": citation_count,
        **media_report,
    }


if __name__ == "__main__":
    print(build())
