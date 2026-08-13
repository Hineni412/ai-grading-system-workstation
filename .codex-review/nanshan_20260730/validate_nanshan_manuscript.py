from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree


EXPECTED_SOURCE_SHA256 = (
    "CF9171BF33C666C96F54E9A25B1298D4CE0B4D160C60E7ACE978E454E09936E8"
)
EXPECTED_TITLE = "初中数学AI阅卷的题型误差与教师复核——基于78份七年级纸笔答卷的实践分析"
CITATION_PATTERN = re.compile(r"\[\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*\]")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def test_zip(path: Path) -> None:
    with ZipFile(path) as archive:
        bad = archive.testzip()
        if bad is not None:
            raise AssertionError(f"Corrupt DOCX member: {bad}")


def east_asia_font(run) -> str | None:
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(qn("w:eastAsia"))


def table_texts(document: Document) -> list[list[list[str]]]:
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]


def media_hashes(path: Path) -> dict[str, str]:
    hashes = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/"):
                hashes[name] = hashlib.sha256(archive.read(name)).hexdigest()
    return hashes


def check_three_line_tables(document: Document) -> None:
    for index, table in enumerate(document.tables, start=1):
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            raise AssertionError(f"Table {index} has no direct table borders")
        values = {
            edge: (
                borders.find(qn(f"w:{edge}")).get(qn("w:val"))
                if borders.find(qn(f"w:{edge}")) is not None
                else None
            )
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV")
        }
        assert values["top"] == "single", (index, values)
        assert values["bottom"] == "single", (index, values)
        for edge in ("left", "right", "insideH", "insideV"):
            assert values[edge] == "nil", (index, values)


def expand_citation(text: str) -> set[int]:
    numbers: set[int] = set()
    for part in text.strip("[]").split(","):
        if "-" in part:
            start_text, end_text = part.split("-", 1)
            start, end = int(start_text), int(end_text)
            assert start <= end
            numbers.update(range(start, end + 1))
        else:
            numbers.add(int(part))
    return numbers


def metadata_is_scrubbed(path: Path) -> None:
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        core = etree.fromstring(archive.read("docProps/core.xml"))
        ns = {
            "dc": "http://purl.org/dc/elements/1.1/",
            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        }
        creator = core.find("dc:creator", ns)
        modified_by = core.find("cp:lastModifiedBy", ns)
        assert creator is None or not (creator.text or "").strip()
        assert modified_by is None or not (modified_by.text or "").strip()
        assert "docProps/custom.xml" not in names
        for name in names:
            if name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(archive.read(name))
                for element in root.iter():
                    assert not any(
                        etree.QName(attribute).localname.startswith("rsid")
                        for attribute in element.attrib
                    ), f"Revision session id remains in {name}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--final", required=True, type=Path)
    args = parser.parse_args()

    assert sha256(args.source) == EXPECTED_SOURCE_SHA256
    test_zip(args.source)
    test_zip(args.final)

    source = Document(args.source)
    final = Document(args.final)
    paragraphs = [paragraph for paragraph in final.paragraphs if paragraph.text.strip()]
    assert len(final.tables) == 7
    assert len(final.inline_shapes) == 4
    expected_tables = table_texts(source)
    expected_tables[0][3][1] = "作图题，11分（原卷9分）"
    expected_tables[4][3][0] = "Q10作图题（原卷9分）"
    assert expected_tables == table_texts(final)

    assert paragraphs[0].text == EXPECTED_TITLE
    assert paragraphs[0].style.name == "Nanshan Title"
    assert paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.font.size and abs(run.font.size.pt - 22.0) < 0.01 for run in paragraphs[0].runs)
    assert all(east_asia_font(run) == "黑体" for run in paragraphs[0].runs)

    assert paragraphs[1].text == "深圳大学附属教育集团外国语中学　　贾浩然"
    assert paragraphs[1].style.name == "Nanshan Author"
    assert paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraphs[2].text.startswith("【摘  要】")
    assert paragraphs[3].text.startswith("【关键词】")
    assert "人工智能　　初中数学　　智能阅卷　　题型误差　　教师复核" in paragraphs[3].text
    for paragraph in (paragraphs[2], paragraphs[3]):
        assert paragraph.paragraph_format.first_line_indent
        assert abs(paragraph.paragraph_format.first_line_indent.pt - 21.0) < 0.01
        assert abs(float(paragraph.paragraph_format.line_spacing) - 1.25) < 0.001

    full_text = "\n".join(paragraph.text for paragraph in final.paragraphs)
    assert "Title:" not in full_text
    assert "Abstract:" not in full_text
    assert "Keywords:" not in full_text
    assert "图4 Q10 作图题的两个脱敏作答案例（原卷题面9分；统计按100分制计11分）" in full_text
    assert "原卷为校内非标准化阶段测试，卷面各题标注分值合计57分" in full_text
    assert "Q10在该口径下满分为11分" in full_text
    assert "并非由这两个案例单独验证的因果结论" in full_text
    assert "图下“偏宽/偏严”等标签只作现象提示" in full_text
    assert "表 1" not in full_text
    assert "图 1" not in full_text
    assert "Bland–Altman" not in full_text

    reference_heading_index = next(
        index
        for index, paragraph in enumerate(final.paragraphs)
        if paragraph.text.strip() == "参考文献"
    )
    body_paragraphs = final.paragraphs[:reference_heading_index]
    citation_runs = []
    for paragraph in body_paragraphs:
        for run in paragraph.runs:
            if CITATION_PATTERN.fullmatch(run.text):
                citation_runs.append(run)
    assert len(citation_runs) == 28
    assert all(run.font.superscript is True for run in citation_runs)
    cited_numbers: set[int] = set()
    for run in citation_runs:
        cited_numbers.update(expand_citation(run.text))
    assert cited_numbers == set(range(1, 31))

    references = [
        paragraph
        for paragraph in final.paragraphs[reference_heading_index + 1 :]
        if re.match(r"^\[\d+\]", paragraph.text)
    ]
    assert len(references) == 30
    for expected, paragraph in enumerate(references, start=1):
        assert paragraph.text.startswith(f"[{expected}]")
        assert not paragraph.text.startswith(f"[{expected}] ")
        assert paragraph.style.name == "Nanshan Reference"
        assert all(run.font.superscript is not True for run in paragraph.runs)
        assert all(
            run.font.size and abs(run.font.size.pt - 7.5) < 0.01
            for run in paragraph.runs
        )
    for index in (19, 20, 21):
        assert "[EB/OL]. arXiv:" in references[index - 1].text

    captions = [
        paragraph
        for paragraph in final.paragraphs
        if paragraph.style.name == "Nanshan Caption"
    ]
    table_captions = [p for p in captions if p.text.startswith("表")]
    figure_captions = [p for p in captions if p.text.startswith("图")]
    assert len(table_captions) == 7
    assert len(figure_captions) == 4
    for paragraph in captions:
        assert paragraph.style.name == "Nanshan Caption"
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert all(
            run.font.size and abs(run.font.size.pt - 7.5) < 0.01
            for run in paragraph.runs
        )
        assert all(east_asia_font(run) == "黑体" for run in paragraph.runs)

    level_one = [
        paragraph
        for paragraph in final.paragraphs
        if re.match(r"^[一二三四五六七八九十]+、", paragraph.text)
    ]
    level_two = [
        paragraph
        for paragraph in final.paragraphs
        if re.match(r"^（[一二三四五六七八九十]+）", paragraph.text)
    ]
    assert len(level_one) == 6
    assert len(level_two) >= 10
    assert all(p.style.name == "Nanshan Heading 1" for p in level_one)
    assert all(p.style.name == "Nanshan Heading 2" for p in level_two)

    for label, value in (
        ("作者姓名：", "贾浩然"),
        ("作者单位：", "深圳大学附属教育集团外国语中学"),
        ("联系电话：", "13201469126"),
        ("QQ或微信：", "894189001（QQ）"),
        ("电子邮箱：", "894189001@qq.com"),
    ):
        assert f"{label}{value}" in full_text
    assert "518071" in full_text

    source_media = media_hashes(args.source)
    final_media = media_hashes(args.final)
    assert set(final_media) == {
        "word/media/image1.png",
        "word/media/image2.png",
        "word/media/image3.png",
        "word/media/image4.png",
    }
    for name, digest in final_media.items():
        assert source_media[name] == digest

    check_three_line_tables(final)
    metadata_is_scrubbed(args.final)

    with ZipFile(args.final) as archive:
        names = set(archive.namelist())
        assert not any(name.startswith("word/comments") for name in names)
        document_xml = archive.read("word/document.xml")
        document_root = etree.fromstring(document_xml)
        assert not document_root.xpath(
            "//w:ins | //w:del",
            namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )

    section = final.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.01
    assert abs(section.page_height.cm - 29.7) < 0.01

    print(
        {
            "source_unchanged": True,
            "paragraphs": len(final.paragraphs),
            "tables": len(final.tables),
            "figures": len(final.inline_shapes),
            "references": len(references),
            "superscript_citation_groups": len(citation_runs),
            "cited_reference_numbers": len(cited_numbers),
            "table_geometry_ready_for_audit": True,
            "metadata_scrubbed": True,
            "q10_scoring_scale_explained": True,
            "qq_contact_present": True,
        }
    )


if __name__ == "__main__":
    main()
