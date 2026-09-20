from __future__ import annotations

import io
import json
import os
import re
import tempfile
from dataclasses import asdict, replace
from pathlib import Path, PurePosixPath
from typing import Iterable, Mapping

import fitz
from docx import Document

from .adapters import (
    ImagePreprocessor,
    LocalOcrAdapter,
    OcrLine,
    MineruOcrAdapter,
    SafeImagePreprocessor,
)
from .contracts import (
    AnswerLink,
    ApplyReviewCommand,
    DocumentSnapshot,
    DocumentState,
    ExportReceipt,
    ExportState,
    ExportWordCommand,
    LayoutBlock,
    LayoutBlockKind,
    ManualQuestionRegion,
    PageSnapshot,
    PrepareSourceCommand,
    PublishCommand,
    PublishReceipt,
    QuestionDraft,
    RecognitionSource,
    ReviewState,
    SourceRegion,
    TextLayerState,
    canonical_hash,
    make_document_snapshot,
    make_question_draft,
    sha256_bytes,
    snapshot_from_payload,
    snapshot_to_payload,
)
from .math_omml import build_math_expression
from .word_renderer import SharedWordQuestionRenderer, WordStyleProfile, validate_docx


_QUESTION_MARKER = re.compile(
    r"^\s*(?:第\s*)?(?P<number>\d{1,3})(?:\s*题|[.．、)）:]|\s)",
    re.IGNORECASE,
)
_ANSWER_MARKER = re.compile(r"^\s*(?:答案|解答|解析)\s*[:：]?", re.IGNORECASE)
_MATH_RUN = re.compile(r"\$\$(.+?)\$\$|\$(.+?)\$", re.DOTALL)


class DocumentPipelineError(RuntimeError):
    pass


class DocumentPipelineConflict(DocumentPipelineError):
    pass


class DocumentPipelineValidationError(DocumentPipelineError):
    pass


class QuestionDocumentPipeline:
    """Stable Interface for source preparation, review, publication and Word export.

    Every filesystem root is injected.  The module has no default connection to
    application data and therefore cannot accidentally inspect a teacher's files.
    """

    def __init__(
        self,
        workspace_root: str | Path,
        *,
        ocr_adapter: LocalOcrAdapter | None = None,
        image_preprocessor: ImagePreprocessor | None = None,
        publisher: object | None = None,
    ) -> None:
        root = Path(workspace_root).resolve()
        root.mkdir(parents=True, exist_ok=True)
        self.workspace_root = root
        self.ocr_adapter = ocr_adapter or MineruOcrAdapter()
        self.image_preprocessor = image_preprocessor or SafeImagePreprocessor()
        self.publisher = publisher

    def prepare_source(self, command: PrepareSourceCommand) -> DocumentSnapshot:
        command_fingerprint = self._prepare_fingerprint(command)
        manifest_path = self._manifest_path(command.operation_id)
        if manifest_path.is_file():
            envelope = self._read_envelope(manifest_path)
            if envelope.get("command_fingerprint") != command_fingerprint:
                raise DocumentPipelineConflict(
                    "operation_id is already bound to a different source or parser command"
                )
            return snapshot_from_payload(envelope["snapshot"])

        source_asset = self._source_asset_path(command)
        self._write_content_addressed(source_asset, command.source.content)
        relative_source = self._relative(source_asset)

        if command.source.media_type == "application/pdf":
            pages = self._prepare_pdf(command)
        elif command.source.media_type in {"image/jpeg", "image/png"}:
            pages = self._prepare_image(command)
        else:  # guarded by DocumentSource; keeps the boundary explicit.
            raise DocumentPipelineValidationError("unsupported document media type")
        if not pages:
            raise DocumentPipelineValidationError("document contains no page")

        questions = self._build_questions(pages, command.manual_questions)
        state = (
            DocumentState.REVIEW_REQUIRED
            if any(question.review_state != ReviewState.TEACHER_VERIFIED for question in questions)
            or any(
                page.text_layer_state != TextLayerState.EMBEDDED
                or page.transform.requires_review
                for page in pages
            )
            else DocumentState.READY
        )
        snapshot = make_document_snapshot(
            document_id=f"document-{command.source.sha256[:24]}",
            operation_id=command.operation_id,
            source_id=command.source.source_id,
            source_filename=command.source.filename,
            media_type=command.source.media_type,
            source_sha256=command.source.sha256,
            source_asset=relative_source,
            parser_profile=command.parser_profile,
            parser_version=command.parser_version,
            state=state,
            pages=pages,
            questions=questions,
        )
        self._write_envelope(manifest_path, command_fingerprint, snapshot)
        return snapshot

    def apply_review(self, command: ApplyReviewCommand) -> DocumentSnapshot:
        manifest_path = self._manifest_path(command.operation_id)
        envelope = self._read_envelope(manifest_path)
        snapshot = snapshot_from_payload(envelope["snapshot"])
        self._assert_revision(snapshot, command.expected_snapshot_revision)

        decisions = {decision.question_id: decision for decision in command.decisions}
        existing_ids = {question.question_id for question in snapshot.questions}
        unknown = sorted(set(decisions) - existing_ids)
        if unknown:
            raise DocumentPipelineValidationError(
                f"review contains unknown question IDs: {', '.join(unknown)}"
            )

        reviewed: list[QuestionDraft] = []
        for question in snapshot.questions:
            decision = decisions.get(question.question_id)
            if decision is None:
                reviewed.append(question)
                continue
            regions = decision.question_regions or question.source_regions
            self._validate_regions(snapshot.pages, regions)
            text = (
                decision.question_text.strip()
                if decision.question_text is not None
                else self._text_for_regions(snapshot.pages, regions) or question.question_text
            )
            answer_regions = (
                decision.answer_regions
                if decision.answer_regions is not None
                else question.answer.source_regions
            )
            self._validate_regions(snapshot.pages, answer_regions)
            answer_text = (
                decision.answer_text.strip()
                if decision.answer_text is not None
                else self._text_for_regions(snapshot.pages, answer_regions)
                or question.answer.text
            )
            answer_status = decision.answer_status or question.answer.status
            if answer_status == ReviewState.MISSING:
                answer_text = ""
                answer_regions = ()
            answer = AnswerLink(
                status=answer_status,
                text=answer_text,
                source_regions=tuple(answer_regions),
                basis=("teacher_review",),
                confidence=1.0 if answer_status == ReviewState.TEACHER_VERIFIED else 0.0,
            )
            math_items = self._math_expressions(
                question.question_id,
                text,
                tuple(regions),
                snapshot.pages,
                confidence=1.0,
                teacher_verified=decision.confirm_content,
            )
            content_changed = (
                text != question.question_text
                or tuple(regions) != question.source_regions
            )
            issues: list[str] = (
                []
                if decision.confirm_content
                else [
                    item
                    for item in question.blocking_issues
                    if item
                    not in {"question_text_missing", "math_render_review_required"}
                ]
            )
            if not text:
                issues.append("question_text_missing")
            if any(not item.omml and not item.fallback_asset for item in math_items):
                issues.append("math_render_review_required")
            if content_changed and not decision.confirm_content:
                issues.append("teacher_content_confirmation_required")
            issues = list(dict.fromkeys(issues))
            review_state = question.review_state
            if decision.confirm_content and not issues:
                review_state = ReviewState.TEACHER_VERIFIED
            elif decision.confirm_content:
                review_state = ReviewState.CANDIDATE
            elif content_changed:
                review_state = ReviewState.CANDIDATE
            reviewed.append(
                make_question_draft(
                    question_id=question.question_id,
                    question_number=question.question_number,
                    question_text=text,
                    source_regions=tuple(regions),
                    answer=answer,
                    math_expressions=math_items,
                    review_state=review_state,
                    blocking_issues=tuple(issues),
                    manual_regions=question.manual_regions
                    or decision.question_regions is not None,
                )
            )

        state = (
            DocumentState.READY
            if reviewed
            and all(item.review_state == ReviewState.TEACHER_VERIFIED for item in reviewed)
            else DocumentState.REVIEW_REQUIRED
        )
        updated = make_document_snapshot(
            document_id=snapshot.document_id,
            operation_id=snapshot.operation_id,
            source_id=snapshot.source_id,
            source_filename=snapshot.source_filename,
            media_type=snapshot.media_type,
            source_sha256=snapshot.source_sha256,
            source_asset=snapshot.source_asset,
            parser_profile=snapshot.parser_profile,
            parser_version=snapshot.parser_version,
            state=state,
            pages=snapshot.pages,
            questions=tuple(reviewed),
        )
        self._write_envelope(manifest_path, str(envelope["command_fingerprint"]), updated)
        return updated

    def publish_questions(self, command: PublishCommand) -> PublishReceipt:
        if self.publisher is None or not hasattr(self.publisher, "publish"):
            raise DocumentPipelineValidationError("a publication adapter was not configured")
        snapshot = self.load_snapshot(command.operation_id)
        self._assert_revision(snapshot, command.expected_snapshot_revision)
        selected = self._select_questions(snapshot, command.question_ids)
        for question in selected:
            if question.review_state != ReviewState.TEACHER_VERIFIED:
                raise DocumentPipelineValidationError(
                    f"question {question.question_id} has not been verified by a teacher"
                )
            if any(not item.omml and not item.fallback_asset for item in question.math_expressions):
                raise DocumentPipelineValidationError(
                    f"question {question.question_id} contains unrenderable math"
                )
        return self.publisher.publish(snapshot, command, selected)

    def export_word(
        self,
        command: ExportWordCommand,
        *,
        style: WordStyleProfile | None = None,
    ) -> ExportReceipt:
        snapshot = self.load_snapshot(command.operation_id)
        self._assert_revision(snapshot, command.expected_snapshot_revision)
        if command.render_mode == "faithful":
            raise DocumentPipelineValidationError(
                "faithful Word rendering is only available for an unchanged DOCX source"
            )
        selected = self._select_questions(snapshot, command.question_ids)
        if not selected:
            raise DocumentPipelineValidationError("export contains no question")

        profile = style or WordStyleProfile()
        renderer = SharedWordQuestionRenderer(
            style=profile,
            asset_resolver=self.resolve_asset,
        )
        document = Document()
        document.add_heading(command.title, level=0)
        fallbacks = []
        for question in selected:
            document.add_heading(f"{question.question_number}.", level=2)
            fallbacks.extend(
                renderer.add_text(
                    document,
                    question.question_text,
                    question_id=question.question_id,
                    expressions=question.math_expressions,
                )
            )
            if (
                command.include_answers
                and question.answer.status == ReviewState.TEACHER_VERIFIED
            ):
                document.add_paragraph("答案：")
                fallbacks.extend(
                    renderer.add_text(
                        document,
                        question.answer.text,
                        question_id=f"{question.question_id}-answer",
                    )
                )

        output = Path(command.output_path).resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}-",
            suffix=".docx",
            dir=output.parent,
            delete=False,
        ) as handle:
            staging = Path(handle.name)
        try:
            document.save(staging)
            errors = validate_docx(staging)
            if errors:
                raise DocumentPipelineValidationError("; ".join(errors))
            os.replace(staging, output)
        finally:
            if staging.exists():
                staging.unlink()
        artifact_hash = sha256_bytes(output.read_bytes())
        return ExportReceipt(
            operation_id=snapshot.operation_id,
            snapshot_revision=snapshot.snapshot_revision,
            state=(
                ExportState.COMPLETE_WITH_FALLBACKS
                if fallbacks
                else ExportState.COMPLETE
            ),
            artifact_path=str(output),
            artifact_sha256=artifact_hash,
            question_ids=tuple(item.question_id for item in selected),
            content_revisions=tuple(item.content_revision for item in selected),
            style_profile_sha256=profile.sha256,
            converter_version="python-docx+restricted-math-omml-v1",
            fallbacks=tuple(fallbacks),
            validation_errors=(),
        )

    def load_snapshot(self, operation_id: str) -> DocumentSnapshot:
        return snapshot_from_payload(
            self._read_envelope(self._manifest_path(operation_id))["snapshot"]
        )

    def resolve_asset(self, relative_path: str) -> Path | None:
        try:
            pure = PurePosixPath(str(relative_path or ""))
            if pure.is_absolute() or ".." in pure.parts or not pure.parts:
                return None
            candidate = (self.workspace_root / Path(*pure.parts)).resolve()
            candidate.relative_to(self.workspace_root)
            return candidate
        except (OSError, ValueError):
            return None

    def _prepare_pdf(self, command: PrepareSourceCommand) -> tuple[PageSnapshot, ...]:
        try:
            document = fitz.open(stream=command.source.content, filetype="pdf")
        except Exception as exc:
            raise DocumentPipelineValidationError("PDF source could not be opened") from exc
        pages: list[PageSnapshot] = []
        try:
            for page_index, page in enumerate(document):
                page_number = page_index + 1
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                png_bytes = pixmap.tobytes("png")
                asset = self._page_asset_path(command.operation_id, page_number)
                self._write_content_addressed(asset, png_bytes)
                raw_blocks = self._embedded_blocks(page, page_number)
                if raw_blocks:
                    blocks = self._ordered_blocks(raw_blocks, page_number)
                    layer = TextLayerState.EMBEDDED
                else:
                    blocks = self._ocr_blocks(png_bytes, page_number)
                    layer = TextLayerState.LOCAL_OCR if blocks and any(
                        item.text for item in blocks
                    ) else TextLayerState.EMPTY
                pages.append(
                    PageSnapshot(
                        page_number=page_number,
                        width=float(page.rect.width),
                        height=float(page.rect.height),
                        rendered_asset=self._relative(asset),
                        rendered_sha256=sha256_bytes(png_bytes),
                        text_layer_state=layer,
                        blocks=blocks,
                        rendered_width=int(pixmap.width),
                        rendered_height=int(pixmap.height),
                        source_to_rendered=(
                            float(pixmap.width) / max(float(page.rect.width), 1.0),
                            0.0,
                            0.0,
                            float(pixmap.height) / max(float(page.rect.height), 1.0),
                            0.0,
                            0.0,
                        ),
                    )
                )
        finally:
            document.close()
        return tuple(pages)

    def _prepare_image(self, command: PrepareSourceCommand) -> tuple[PageSnapshot, ...]:
        prepared = self.image_preprocessor.prepare(
            command.source.content,
            media_type=command.source.media_type,
        )
        asset = self._page_asset_path(command.operation_id, 1)
        self._write_content_addressed(asset, prepared.png_bytes)
        blocks = self._ocr_blocks(prepared.png_bytes, 1)
        return (
            PageSnapshot(
                page_number=1,
                width=float(prepared.width),
                height=float(prepared.height),
                rendered_asset=self._relative(asset),
                rendered_sha256=sha256_bytes(prepared.png_bytes),
                text_layer_state=(
                    TextLayerState.LOCAL_OCR
                    if blocks and any(item.text for item in blocks)
                    else TextLayerState.EMPTY
                ),
                blocks=blocks,
                transform=prepared.transform,
                rendered_width=prepared.width,
                rendered_height=prepared.height,
            ),
        )

    def _embedded_blocks(self, page, page_number: int) -> list[Mapping[str, object]]:
        width = max(float(page.rect.width), 1.0)
        height = max(float(page.rect.height), 1.0)
        result: list[Mapping[str, object]] = []
        for item in page.get_text("blocks"):
            if len(item) < 5:
                continue
            text = str(item[4] or "").strip()
            if not text:
                continue
            x0, y0, x1, y1 = (float(value) for value in item[:4])
            polygon = self._rect_polygon(x0 / width, y0 / height, x1 / width, y1 / height)
            result.append(
                {
                    "text": text,
                    "polygon": polygon,
                    "source": RecognitionSource.EMBEDDED_TEXT,
                    "confidence": 1.0,
                    "engine_version": f"pymupdf/{fitz.VersionBind}",
                    "requires_review": False,
                }
            )
        return result

    def _ocr_blocks(self, png_bytes: bytes, page_number: int) -> tuple[LayoutBlock, ...]:
        try:
            lines = self.ocr_adapter.recognize(png_bytes, page_number=page_number)
        except Exception:
            lines = ()
        raw: list[Mapping[str, object]] = [
            {
                "text": line.text,
                "polygon": line.polygon,
                "source": RecognitionSource.LOCAL_OCR,
                "confidence": line.confidence,
                "engine_version": line.engine_version,
                "requires_review": True,
            }
            for line in lines
        ]
        if not raw:
            raw.append(
                {
                    "text": "",
                    "polygon": self._rect_polygon(0.0, 0.0, 1.0, 1.0),
                    "source": RecognitionSource.LOCAL_OCR,
                    "confidence": 0.0,
                    "engine_version": "local-ocr/unavailable",
                    "requires_review": True,
                }
            )
        return self._ordered_blocks(raw, page_number)

    def _ordered_blocks(
        self,
        raw_blocks: Iterable[Mapping[str, object]],
        page_number: int,
    ) -> tuple[LayoutBlock, ...]:
        items = list(raw_blocks)

        def bounds(item: Mapping[str, object]) -> tuple[float, float, float, float]:
            points = tuple(item["polygon"])  # type: ignore[arg-type]
            xs = [float(point[0]) for point in points]
            ys = [float(point[1]) for point in points]
            return min(xs), min(ys), max(xs), max(ys)

        bounded = [(item, bounds(item)) for item in items]
        has_left = any(box[2] <= 0.62 for _, box in bounded)
        has_right = any(box[0] >= 0.38 for _, box in bounded)
        two_columns = has_left and has_right

        def order_key(pair: tuple[Mapping[str, object], tuple[float, float, float, float]]):
            _, box = pair
            x0, y0, x1, _ = box
            width = x1 - x0
            if not two_columns:
                return (0, y0, x0)
            if width >= 0.7:
                return (0 if y0 < 0.2 else 3, y0, x0)
            center = (x0 + x1) / 2
            return (1 if center < 0.5 else 2, y0, x0)

        ordered = sorted(bounded, key=order_key)
        blocks: list[LayoutBlock] = []
        for index, (raw, _) in enumerate(ordered):
            text = str(raw.get("text") or "").strip()
            region = SourceRegion(
                page_number=page_number,
                polygon=tuple(raw["polygon"]),  # type: ignore[arg-type]
                source=RecognitionSource(raw["source"]),
                confidence=float(raw["confidence"]),
            )
            blocks.append(
                LayoutBlock(
                    block_id=f"p{page_number}-b{index + 1}",
                    page_number=page_number,
                    kind=self._block_kind(text),
                    text=text,
                    region=region,
                    reading_order=index,
                    source=RecognitionSource(raw["source"]),
                    confidence=float(raw["confidence"]),
                    engine_version=str(raw["engine_version"]),
                    requires_review=bool(raw.get("requires_review")),
                )
            )
        return tuple(blocks)

    def _build_questions(
        self,
        pages: tuple[PageSnapshot, ...],
        manual_questions: tuple[ManualQuestionRegion, ...],
    ) -> tuple[QuestionDraft, ...]:
        if manual_questions:
            return tuple(self._manual_question(pages, item) for item in manual_questions)

        all_blocks = [block for page in pages for block in page.blocks]
        marker_indexes = [
            index for index, block in enumerate(all_blocks) if block.kind == LayoutBlockKind.QUESTION_MARKER
        ]
        groups: list[list[LayoutBlock]] = []
        if marker_indexes:
            for marker_position, start in enumerate(marker_indexes):
                end = marker_indexes[marker_position + 1] if marker_position + 1 < len(marker_indexes) else len(all_blocks)
                groups.append(all_blocks[start:end])
        elif all_blocks:
            groups.append(all_blocks)

        questions: list[QuestionDraft] = []
        for index, blocks in enumerate(groups):
            text = "\n".join(block.text for block in blocks if block.text).strip()
            match = _QUESTION_MARKER.match(text)
            number = match.group("number") if match else str(index + 1)
            regions = tuple(block.region for block in blocks)
            issues = ["teacher_content_confirmation_required"]
            if not text:
                issues.append("question_text_missing")
            if any(block.requires_review for block in blocks):
                issues.append("ocr_review_required")
            if not marker_indexes:
                issues.append("question_boundary_not_detected")
            question_id = f"q-{index + 1}"
            questions.append(
                make_question_draft(
                    question_id=question_id,
                    question_number=number,
                    question_text=text,
                    source_regions=regions,
                    answer=AnswerLink(status=ReviewState.MISSING),
                    math_expressions=self._math_expressions(
                        question_id,
                        text,
                        regions,
                        pages,
                        confidence=min((block.confidence for block in blocks), default=0.0),
                    ),
                    review_state=ReviewState.CANDIDATE,
                    blocking_issues=tuple(dict.fromkeys(issues)),
                )
            )
        return tuple(questions)

    def _manual_question(
        self,
        pages: tuple[PageSnapshot, ...],
        manual: ManualQuestionRegion,
    ) -> QuestionDraft:
        self._validate_regions(pages, manual.question_regions)
        self._validate_regions(pages, manual.answer_regions)
        text = self._text_for_regions(pages, manual.question_regions)
        answer_text = manual.answer_text or self._text_for_regions(pages, manual.answer_regions)
        answer_status = (
            ReviewState.CANDIDATE if answer_text or manual.answer_regions else ReviewState.MISSING
        )
        issues = ["teacher_content_confirmation_required"]
        selected_blocks = self._blocks_for_regions(pages, manual.question_regions)
        if not text:
            issues.append("question_text_missing")
        if any(block.requires_review for block in selected_blocks):
            issues.append("ocr_review_required")
        return make_question_draft(
            question_id=manual.question_id,
            question_number=manual.question_number,
            question_text=text,
            source_regions=manual.question_regions,
            answer=AnswerLink(
                status=answer_status,
                text=answer_text,
                source_regions=manual.answer_regions,
                basis=("teacher_region",) if manual.answer_regions else (),
                confidence=1.0 if manual.answer_text else 0.0,
            ),
            math_expressions=self._math_expressions(
                manual.question_id,
                text,
                manual.question_regions,
                pages,
                confidence=min((block.confidence for block in selected_blocks), default=0.0),
            ),
            review_state=ReviewState.CANDIDATE,
            blocking_issues=tuple(dict.fromkeys(issues)),
            manual_regions=True,
        )

    def _math_expressions(
        self,
        question_id: str,
        text: str,
        regions: tuple[SourceRegion, ...],
        pages: tuple[PageSnapshot, ...],
        *,
        confidence: float,
        teacher_verified: bool = False,
    ) -> tuple:
        items = []
        fallback_asset = None
        fallback_sha256 = None
        if regions:
            page = pages[regions[0].page_number - 1]
            fallback_asset = page.rendered_asset
            fallback_sha256 = page.rendered_sha256
        for index, match in enumerate(_MATH_RUN.finditer(text)):
            source = match.group(1) if match.group(1) is not None else match.group(2)
            items.append(
                build_math_expression(
                    expression_id=f"{question_id}-math-{index + 1}",
                    source=source,
                    source_regions=regions,
                    confidence=confidence,
                    review_state=(
                        ReviewState.TEACHER_VERIFIED
                        if teacher_verified
                        or (
                            confidence >= 0.9
                        and all(
                            region.source == RecognitionSource.EMBEDDED_TEXT
                            for region in regions
                        )
                        )
                        else ReviewState.CANDIDATE
                    ),
                    fallback_asset=fallback_asset,
                    fallback_sha256=fallback_sha256,
                )
            )
        return tuple(items)

    def _text_for_regions(
        self,
        pages: tuple[PageSnapshot, ...],
        regions: Iterable[SourceRegion],
    ) -> str:
        fragments: list[str] = []
        for region in regions:
            blocks = self._blocks_for_region(pages, region)
            fragment = "\n".join(block.text for block in blocks if block.text).strip()
            if fragment:
                fragments.append(fragment)
        return "\n".join(fragments).strip()

    def _blocks_for_regions(
        self,
        pages: tuple[PageSnapshot, ...],
        regions: Iterable[SourceRegion],
    ) -> tuple[LayoutBlock, ...]:
        found: list[LayoutBlock] = []
        seen: set[str] = set()
        for region in regions:
            for block in self._blocks_for_region(pages, region):
                if block.block_id not in seen:
                    found.append(block)
                    seen.add(block.block_id)
        return tuple(found)

    def _blocks_for_region(
        self,
        pages: tuple[PageSnapshot, ...],
        region: SourceRegion,
    ) -> tuple[LayoutBlock, ...]:
        page = pages[region.page_number - 1]
        return tuple(
            block
            for block in page.blocks
            if self._regions_intersect(block.region, region)
        )

    @staticmethod
    def _regions_intersect(first: SourceRegion, second: SourceRegion) -> bool:
        ax0, ay0, ax1, ay1 = first.bbox
        bx0, by0, bx1, by1 = second.bbox
        center_x = (ax0 + ax1) / 2
        center_y = (ay0 + ay1) / 2
        if bx0 <= center_x <= bx1 and by0 <= center_y <= by1:
            return True
        overlap_w = max(0.0, min(ax1, bx1) - max(ax0, bx0))
        overlap_h = max(0.0, min(ay1, by1) - max(ay0, by0))
        block_area = max((ax1 - ax0) * (ay1 - ay0), 1e-9)
        return overlap_w * overlap_h / block_area >= 0.15

    @staticmethod
    def _validate_regions(
        pages: tuple[PageSnapshot, ...],
        regions: Iterable[SourceRegion],
    ) -> None:
        page_count = len(pages)
        for region in regions:
            if region.page_number > page_count:
                raise DocumentPipelineValidationError(
                    f"source region page {region.page_number} is outside the document"
                )

    @staticmethod
    def _block_kind(text: str) -> LayoutBlockKind:
        if _ANSWER_MARKER.match(text):
            return LayoutBlockKind.ANSWER_MARKER
        if _QUESTION_MARKER.match(text):
            return LayoutBlockKind.QUESTION_MARKER
        if text.startswith("$") and text.endswith("$"):
            return LayoutBlockKind.FORMULA
        return LayoutBlockKind.TEXT

    @staticmethod
    def _rect_polygon(
        x0: float, y0: float, x1: float, y1: float
    ) -> tuple[tuple[float, float], ...]:
        values = [max(0.0, min(1.0, value)) for value in (x0, y0, x1, y1)]
        left, top, right, bottom = values
        if right <= left:
            right = min(1.0, left + 1e-6)
            left = max(0.0, right - 1e-6)
        if bottom <= top:
            bottom = min(1.0, top + 1e-6)
            top = max(0.0, bottom - 1e-6)
        return ((left, top), (right, top), (right, bottom), (left, bottom))

    def _select_questions(
        self,
        snapshot: DocumentSnapshot,
        question_ids: tuple[str, ...],
    ) -> tuple[QuestionDraft, ...]:
        if not question_ids:
            return snapshot.questions
        lookup = {item.question_id: item for item in snapshot.questions}
        missing = [item for item in question_ids if item not in lookup]
        if missing:
            raise DocumentPipelineValidationError(
                f"unknown question IDs: {', '.join(missing)}"
            )
        return tuple(lookup[item] for item in question_ids)

    @staticmethod
    def _assert_revision(snapshot: DocumentSnapshot, expected: str) -> None:
        if snapshot.snapshot_revision != expected:
            raise DocumentPipelineConflict(
                "document snapshot changed; reload it before applying this operation"
            )

    def _source_asset_path(self, command: PrepareSourceCommand) -> Path:
        suffix = {
            "application/pdf": ".pdf",
            "image/jpeg": ".jpg",
            "image/png": ".png",
        }[command.source.media_type]
        return self.workspace_root / "assets" / "sources" / f"{command.source.sha256}{suffix}"

    def _page_asset_path(self, operation_id: str, page_number: int) -> Path:
        return (
            self.workspace_root
            / "operations"
            / self._operation_folder(operation_id)
            / "pages"
            / f"page-{page_number:04d}.png"
        )

    def _manifest_path(self, operation_id: str) -> Path:
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.:-]{0,127}", str(operation_id)):
            raise DocumentPipelineValidationError("operation_id is invalid")
        return (
            self.workspace_root
            / "operations"
            / self._operation_folder(operation_id)
            / "snapshot.json"
        )

    @staticmethod
    def _operation_folder(operation_id: str) -> str:
        return f"operation-{canonical_hash({'operation_id': operation_id})[:24]}"

    def _relative(self, path: Path) -> str:
        candidate = path.resolve()
        candidate.relative_to(self.workspace_root)
        return candidate.relative_to(self.workspace_root).as_posix()

    @staticmethod
    def _prepare_fingerprint(command: PrepareSourceCommand) -> str:
        return canonical_hash(
            {
                "operation_id": command.operation_id,
                "source_id": command.source.source_id,
                "source_filename": command.source.filename,
                "source_sha256": command.source.sha256,
                "media_type": command.source.media_type,
                "parser_profile": command.parser_profile,
                "parser_version": command.parser_version,
                "manual_questions": [asdict(item) for item in command.manual_questions],
            }
        )

    def _write_content_addressed(self, path: Path, content: bytes) -> None:
        if path.is_file():
            if sha256_bytes(path.read_bytes()) != sha256_bytes(content):
                raise DocumentPipelineConflict("content-addressed asset hash mismatch")
            return
        self._atomic_write(path, content)

    def _write_envelope(
        self,
        path: Path,
        command_fingerprint: str,
        snapshot: DocumentSnapshot,
    ) -> None:
        payload = json.dumps(
            {
                "contract_version": 1,
                "command_fingerprint": command_fingerprint,
                "snapshot": snapshot_to_payload(snapshot),
            },
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        self._atomic_write(path, payload)

    @staticmethod
    def _read_envelope(path: Path) -> dict:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise DocumentPipelineValidationError("document snapshot is missing or invalid") from exc
        if payload.get("contract_version") != 1 or not isinstance(payload.get("snapshot"), dict):
            raise DocumentPipelineValidationError("document snapshot contract version is invalid")
        return payload

    @staticmethod
    def _atomic_write(path: Path, content: bytes) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.name}-",
            suffix=".tmp",
            dir=path.parent,
            delete=False,
        ) as handle:
            staging = Path(handle.name)
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        try:
            os.replace(staging, path)
        finally:
            if staging.exists():
                staging.unlink()


__all__ = [
    "DocumentPipelineConflict",
    "DocumentPipelineError",
    "DocumentPipelineValidationError",
    "QuestionDocumentPipeline",
]
