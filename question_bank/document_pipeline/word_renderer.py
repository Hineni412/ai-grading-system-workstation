from __future__ import annotations

import logging
import posixpath
import re
import zipfile
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from docx import Document as WordDocument
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .contracts import FormulaFallback, MathExpression, canonical_hash
from .math_omml import build_math_expression

_MATH_RUN = re.compile(r"\$\$(.+?)\$\$|\$(.+?)\$", re.DOTALL)
_LEADING_QUESTION_NUMBER = re.compile(
    r"^\s*(?:第\s*)?\d{1,3}\s*(?:[.．、]|题)[ \t]*"
)
_REL_EMBED_ATTR = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
LOGGER = logging.getLogger(__name__)

# 无 DPI 元数据时按 150dpi 折算（MinerU 页面裁图约 200dpi，屏幕图约 96dpi，
# 取中间值让小题图按接近原稿的尺寸落版，而不是统一拉满页宽）。
_DEFAULT_IMAGE_DPI = 150.0
_MAX_IMAGE_HEIGHT_INCHES = 4.5


def natural_image_width_inches(
    path: Path,
    *,
    max_width_inches: float,
    max_height_inches: float | None = None,
) -> float:
    """按图片自然像素与 DPI 折算显示宽度， capped by max_width_inches 与高度上限。"""
    try:
        from PIL import Image

        height_cap = (
            _MAX_IMAGE_HEIGHT_INCHES
            if max_height_inches is None
            else float(max_height_inches)
        )
        with Image.open(path) as image:
            width_px, height_px = image.size
            dpi = image.info.get("dpi")
            if isinstance(dpi, (tuple, list)) and dpi and dpi[0]:
                dpi_x = float(dpi[0])
            elif isinstance(dpi, (int, float)) and dpi:
                dpi_x = float(dpi)
            else:
                dpi_x = _DEFAULT_IMAGE_DPI
            if dpi_x <= 0 or width_px <= 0 or height_px <= 0:
                return max_width_inches
            width_inches = width_px / dpi_x
            height_inches = height_px / dpi_x
            if height_inches > height_cap:
                width_inches *= height_cap / height_inches
            return min(width_inches, max_width_inches)
    except Exception:
        return max_width_inches


def add_floating_picture(
    paragraph,
    image_path,
    *,
    width_inches: float,
    align: str = "right",
    descr: str | None = None,
) -> bool:
    """把图片以浮动锚定（wp:anchor，方形环绕）挂到既有段落上。

    先按内联图插入拿到 extent/docPr 等子树，再把 wp:inline 改写成
    wp:anchor；子元素顺序必须是 simplePos, positionH, positionV, extent,
    effectExtent, wrapSquare, docPr, cNvGraphicFramePr, graphic。
    """
    run = paragraph.add_run()
    try:
        run.add_picture(str(image_path), width=Inches(width_inches))
    except Exception:
        return False
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline")) if drawing is not None else None
    if inline is None:
        return False
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    graphic_frame = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    if descr and doc_pr is not None:
        doc_pr.set("descr", descr)

    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "114300",
        "distR": "114300",
        "simplePos": "0",
        "relativeHeight": "251658240",
        "behindDoc": "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "0",
    }.items():
        anchor.set(key, value)
    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "margin")
    horizontal = OxmlElement("wp:align")
    horizontal.text = align
    position_h.append(horizontal)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "paragraph")
    offset = OxmlElement("wp:posOffset")
    offset.text = "0"
    position_v.append(offset)
    effect_extent = OxmlElement("wp:effectExtent")
    for key in ("l", "t", "r", "b"):
        effect_extent.set(key, "0")
    wrap_square = OxmlElement("wp:wrapSquare")
    # 图在右时文字只绕左侧；左对齐时反之。
    wrap_square.set("wrapText", "left" if align == "right" else "right")

    anchor.append(simple_pos)
    anchor.append(position_h)
    anchor.append(position_v)
    if extent is not None:
        anchor.append(extent)
    anchor.append(effect_extent)
    anchor.append(wrap_square)
    for child in (doc_pr, graphic_frame, graphic):
        if child is not None:
            anchor.append(child)
    drawing.replace(inline, anchor)
    return True


@dataclass(frozen=True, slots=True)
class WordStyleProfile:
    body_font: str = "宋体"
    body_font_ascii: str = "Times New Roman"
    body_size_pt: float = 10.5
    line_spacing: float = 1.1
    formula_size_pt: float = 11.0
    image_width_inches: float = 4.8
    keep_question_together: bool = True
    content_width_dxa: int = 9865

    def __post_init__(self) -> None:
        if not str(self.body_font or "").strip() or not str(self.body_font_ascii or "").strip():
            raise ValueError("Word fonts must be non-empty")
        if not 6.0 <= float(self.body_size_pt) <= 48.0:
            raise ValueError("body_size_pt is outside the supported range")
        if not 6.0 <= float(self.formula_size_pt) <= 48.0:
            raise ValueError("formula_size_pt is outside the supported range")
        if not 0.8 <= float(self.line_spacing) <= 3.0:
            raise ValueError("line_spacing is outside the supported range")
        if not 0.2 <= float(self.image_width_inches) <= 8.0:
            raise ValueError("image_width_inches is outside the supported range")
        if not 3000 <= int(self.content_width_dxa) <= 16000:
            raise ValueError("content_width_dxa is outside the supported range")

    @property
    def sha256(self) -> str:
        return canonical_hash(asdict(self))

    @classmethod
    def from_export_config(cls, config: object | None) -> WordStyleProfile:
        if config is None:
            return cls()
        page_width_cm = (
            29.7
            if str(getattr(config, "orientation", "portrait")).casefold()
            == "landscape"
            else 21.0
        )
        content_width_cm = max(
            page_width_cm
            - float(getattr(config, "margin_left_cm", 1.45))
            - float(getattr(config, "margin_right_cm", 1.45)),
            5.3,
        )
        return cls(
            body_font=str(getattr(config, "body_font", "宋体")),
            body_font_ascii=str(getattr(config, "body_font_ascii", "Times New Roman")),
            body_size_pt=float(getattr(config, "body_size_pt", 10.5)),
            line_spacing=float(getattr(config, "line_spacing", 1.1)),
            content_width_dxa=int(round(content_width_cm * 567)),
        )


@dataclass(frozen=True, slots=True)
class RichBlockRenderResult:
    appended: bool
    embedded_assets: tuple[str, ...] = ()
    inline_prefix_applied: bool = False
    compacted_image_count: int = 0


class SharedWordQuestionRenderer:
    """One semantic math renderer shared by ordinary and personalized papers."""

    def __init__(
        self,
        *,
        style: WordStyleProfile | None = None,
        asset_resolver: Callable[[str], Path | None] | None = None,
    ) -> None:
        self.style = style or WordStyleProfile()
        self.asset_resolver = asset_resolver

    def add_rich_blocks(
        self,
        document,
        blocks: Sequence[Mapping[str, object]],
        *,
        strip_leading_number: bool = False,
        inline_prefix: str = "",
        align_standalone_images_right: bool = True,
        compact_standalone_images_with_text: bool = False,
    ) -> RichBlockRenderResult:
        """Append frozen Word blocks without flattening their native semantics.

        The blocks are preflighted as one unit. If any XML or required image is
        unavailable, nothing is appended and the caller can safely use its
        plain-text fallback without mixing duplicate or partial content.
        """

        prepared: list[tuple[object, dict[str, Path]]] = []
        embedded_assets: set[str] = set()
        has_unrenderable_text_block = False
        try:
            for index, block in enumerate(blocks):
                xml = str(block.get("xml") or "").strip()
                if not xml:
                    if str(block.get("text") or "").strip():
                        has_unrenderable_text_block = True
                    continue
                element = parse_xml(xml)
                if strip_leading_number and index == 0:
                    _strip_leading_question_number(element)
                _style_rich_paragraphs(element, line_spacing=self.style.line_spacing)
                _cap_option_picture_height(element)
                if align_standalone_images_right:
                    _align_standalone_picture_paragraphs_right(element)
                relationships = block.get("image_relationships")
                relationship_values = (
                    relationships if isinstance(relationships, Mapping) else {}
                )
                resolved_relationships: dict[str, Path] = {}
                required_relationships = {
                    str(child.get(_REL_EMBED_ATTR))
                    for child in element.iter()
                    if child.get(_REL_EMBED_ATTR)
                }
                for relationship_id in required_relationships:
                    source = str(relationship_values.get(relationship_id) or "").strip()
                    if not source:
                        raise ValueError(
                            f"rich Word image relationship is missing: {relationship_id}"
                        )
                    resolved = (
                        self.asset_resolver(source)
                        if self.asset_resolver
                        else Path(source)
                    )
                    if resolved is None or not Path(resolved).is_file():
                        raise ValueError("rich Word image asset is unavailable")
                    resolved_relationships[relationship_id] = Path(resolved)
                prepared.append((element, resolved_relationships))
            if not prepared or has_unrenderable_text_block:
                return RichBlockRenderResult(appended=False)
        except Exception:
            LOGGER.exception("Failed to preflight rich Word blocks")
            return RichBlockRenderResult(appended=False)

        inline_prefix_applied = False
        if inline_prefix:
            inline_prefix_applied = _prepend_inline_prefix(
                [element for element, _relationships in prepared],
                inline_prefix,
                style=self.style,
            )
            if not inline_prefix_applied:
                prefix_paragraph = OxmlElement("w:p")
                _style_rich_paragraphs(
                    prefix_paragraph,
                    line_spacing=self.style.line_spacing,
                )
                _prepend_inline_prefix(
                    [prefix_paragraph],
                    inline_prefix,
                    style=self.style,
                )
                prepared.insert(0, (prefix_paragraph, {}))
                inline_prefix_applied = True
        _keep_text_with_following_picture(
            [element for element, _relationships in prepared]
        )
        compacted_image_count = 0
        if compact_standalone_images_with_text:
            prepared, compacted_image_count = _compact_image_pairs(
                prepared,
                content_width_dxa=self.style.content_width_dxa,
            )
        prepared, _gridded_option_count = _grid_option_picture_groups(
            prepared,
            content_width_dxa=self.style.content_width_dxa,
        )

        try:
            for element, relationships in prepared:
                relationship_ids: dict[str, str] = {}
                for old_relationship_id, resolved in relationships.items():
                    new_relationship_id, _ = document.part.get_or_add_image(
                        str(resolved)
                    )
                    relationship_ids[old_relationship_id] = new_relationship_id
                    embedded_assets.add(str(resolved))
                for child in element.iter():
                    old_relationship_id = child.get(_REL_EMBED_ATTR)
                    if old_relationship_id in relationship_ids:
                        child.set(
                            _REL_EMBED_ATTR,
                            relationship_ids[old_relationship_id],
                        )
            for element, _relationships in prepared:
                _append_to_document_body(document, element)
        except Exception:
            LOGGER.exception("Failed to append rich Word blocks")
            return RichBlockRenderResult(appended=False)
        return RichBlockRenderResult(
            appended=True,
            embedded_assets=tuple(sorted(embedded_assets)),
            inline_prefix_applied=inline_prefix_applied,
            compacted_image_count=compacted_image_count,
        )

    def add_text(
        self,
        document,
        text: str,
        *,
        question_id: str,
        expressions: Iterable[MathExpression] = (),
        inline_prefix: str = "",
    ) -> tuple[FormulaFallback, ...]:
        expression_pool = list(expressions)
        fallbacks: list[FormulaFallback] = []
        lines = str(text or "").splitlines() or [""]
        expression_index = 0
        for line in lines:
            paragraph = document.add_paragraph()
            self._style_paragraph(paragraph)
            if inline_prefix:
                self._add_text_run(paragraph, inline_prefix)
                inline_prefix = ""
            cursor = 0
            for match in _MATH_RUN.finditer(line):
                if match.start() > cursor:
                    self._add_text_run(paragraph, line[cursor : match.start()])
                raw_latex = match.group(1) if match.group(1) is not None else match.group(2)
                display = match.group(1) is not None
                expression = self._expression_for(
                    expression_pool,
                    raw_latex,
                    question_id=question_id,
                    index=expression_index,
                )
                expression_index += 1
                target = paragraph
                if display:
                    target = document.add_paragraph()
                    target.alignment = 1
                    self._style_paragraph(target)
                fallback = self._append_expression(
                    target,
                    expression,
                    question_id=question_id,
                )
                if fallback is not None:
                    fallbacks.append(fallback)
                cursor = match.end()
            if cursor < len(line):
                self._add_text_run(paragraph, line[cursor:])
        return tuple(fallbacks)

    def add_images(self, document, image_paths: Iterable[str]) -> None:
        for value in image_paths:
            resolved = self.asset_resolver(str(value)) if self.asset_resolver else Path(str(value))
            if resolved is None or not resolved.is_file():
                document.add_paragraph("[题图不可用]")
                continue
            try:
                width = natural_image_width_inches(
                    resolved, max_width_inches=self.style.image_width_inches
                )
                document.add_picture(str(resolved), width=Inches(width))
            except Exception:
                document.add_paragraph("[题图无法插入]")

    def add_to_paragraph(
        self,
        paragraph,
        text: str,
        *,
        question_id: str,
        expressions: Iterable[MathExpression] = (),
    ) -> tuple[FormulaFallback, ...]:
        """Render text and editable OMML into an existing paragraph or table cell."""

        pool = list(expressions)
        fallbacks: list[FormulaFallback] = []
        cursor = 0
        expression_index = 0
        value = str(text or "")
        self._style_paragraph(paragraph)
        for match in _MATH_RUN.finditer(value):
            if match.start() > cursor:
                self._add_text_run(paragraph, value[cursor : match.start()])
            raw_latex = match.group(1) if match.group(1) is not None else match.group(2)
            expression = self._expression_for(
                pool,
                raw_latex,
                question_id=question_id,
                index=expression_index,
            )
            expression_index += 1
            fallback = self._append_expression(
                paragraph,
                expression,
                question_id=question_id,
            )
            if fallback is not None:
                fallbacks.append(fallback)
            cursor = match.end()
        if cursor < len(value):
            self._add_text_run(paragraph, value[cursor:])
        return tuple(fallbacks)

    def _expression_for(
        self,
        pool: list[MathExpression],
        raw_latex: str,
        *,
        question_id: str,
        index: int,
    ) -> MathExpression:
        clean = str(raw_latex or "").strip()
        for expression in pool:
            if expression.restricted_latex.strip("$").strip() == clean:
                return expression
        return build_math_expression(
            expression_id=f"{_safe_question_id(question_id)}-math-{index + 1}",
            source=clean,
        )

    def _append_expression(
        self,
        paragraph,
        expression: MathExpression,
        *,
        question_id: str,
    ) -> FormulaFallback | None:
        if expression.omml:
            try:
                paragraph._p.append(parse_xml(expression.omml))
                return None
            except Exception:
                reason = "validated OMML could not be inserted"
        else:
            reason = expression.validation_error or "expression is not supported"
        fallback_path = expression.fallback_asset
        resolved = (
            self.asset_resolver(fallback_path)
            if fallback_path and self.asset_resolver
            else Path(fallback_path) if fallback_path else None
        )
        # 回退文本剥掉 $ 定界符；空表达式（形如 $$）不输出任何内容。
        fallback_text = expression.restricted_latex.replace("$", "").strip()
        if resolved is not None and resolved.is_file():
            try:
                paragraph.add_run().add_picture(
                    str(resolved), width=Inches(min(3.0, self.style.image_width_inches))
                )
            except Exception:
                if fallback_text:
                    paragraph.add_run(f"${fallback_text}$")
                reason = f"{reason}; fallback image could not be inserted"
        elif fallback_text:
            paragraph.add_run(f"${fallback_text}$")
            if fallback_path:
                reason = f"{reason}; fallback image is unavailable"
        return FormulaFallback(
            question_id=_safe_question_id(question_id),
            expression_id=expression.expression_id,
            reason=reason,
            asset_path=fallback_path,
        )

    def _style_paragraph(self, paragraph) -> None:
        paragraph.paragraph_format.line_spacing = self.style.line_spacing
        paragraph.paragraph_format.widow_control = True
        if self.style.keep_question_together:
            paragraph.paragraph_format.keep_together = True

    def _add_text_run(self, paragraph, value: str) -> None:
        if not value:
            return
        run = paragraph.add_run(value)
        run.font.name = self.style.body_font_ascii
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.style.body_font)
        run.font.size = Pt(self.style.body_size_pt)


def validate_docx(path: str | Path) -> tuple[str, ...]:
    candidate = Path(path)
    errors: list[str] = []
    if not candidate.is_file() or candidate.suffix.casefold() != ".docx":
        return ("DOCX artifact is missing",)
    try:
        with zipfile.ZipFile(candidate) as archive:
            names = set(archive.namelist())
            for required in ("[Content_Types].xml", "word/document.xml"):
                if required not in names:
                    errors.append(f"DOCX member is missing: {required}")
            document_xml = archive.read("word/document.xml") if "word/document.xml" in names else b""
            if b"<m:oMath" in document_xml and b"officeDocument/2006/math" not in document_xml:
                errors.append("OMML namespace is missing")
            bad_members = [name for name in names if name.startswith("/") or ".." in Path(name).parts]
            if bad_members:
                errors.append("DOCX contains an unsafe member path")
            for relationship_name in sorted(
                name for name in names if name.endswith(".rels")
            ):
                try:
                    root = ElementTree.fromstring(archive.read(relationship_name))
                except ElementTree.ParseError:
                    errors.append(f"DOCX relationship XML is invalid: {relationship_name}")
                    continue
                relationship_path = PurePosixPath(relationship_name)
                owner_directory = relationship_path.parent.parent.as_posix()
                for relationship in root:
                    if relationship.attrib.get("TargetMode") == "External":
                        continue
                    target = relationship.attrib.get("Target", "")
                    resolved = posixpath.normpath(
                        posixpath.join(owner_directory, target)
                    )
                    escapes_package = (
                        resolved == ".."
                        or resolved.startswith("../")
                        or resolved.startswith("/")
                    )
                    package_target = resolved.removeprefix("./")
                    if escapes_package or package_target not in names:
                        errors.append(
                            "DOCX relationship target is missing: "
                            f"{package_target or target}"
                        )
    except (OSError, zipfile.BadZipFile, KeyError):
        errors.append("DOCX artifact cannot be reopened")
    if not errors:
        try:
            WordDocument(candidate)
        except Exception:
            errors.append("DOCX artifact cannot be reopened by python-docx")
    return tuple(errors)


def _safe_question_id(value: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9_.:-]+", "-", str(value or "").strip()).strip("-")
    return (clean or "question")[:128]


def rich_block_text(
    blocks: Sequence[Mapping[str, object]],
    *,
    strip_leading_number: bool = False,
) -> str | None:
    """Return the visible native Word text, or None when rich fallback is required."""

    elements: list[object] = []
    try:
        for index, block in enumerate(blocks):
            xml = str(block.get("xml") or "").strip()
            if not xml:
                if str(block.get("text") or "").strip():
                    return None
                continue
            element = parse_xml(xml)
            if strip_leading_number and index == 0:
                _strip_leading_question_number(element)
            elements.append(element)
    except Exception:
        return None
    if not elements:
        return None
    text_tags = {qn("w:t"), qn("m:t")}
    return "".join(
        child.text or ""
        for element in elements
        for child in element.iter()
        if child.tag in text_tags
    )


def compact_source_label(
    source: object = "",
    *,
    year: object = "",
    district: object = "",
    exam_type: object = "",
) -> str:
    """Reduce verbose paper provenance to a compact, factual inline label."""

    text = " ".join(
        part
        for part in (
            str(source or "").strip(),
            str(year or "").strip(),
            str(district or "").strip(),
            str(exam_type or "").strip(),
        )
        if part
    )
    years = re.findall(r"(?:19|20)\d{2}", text)
    year_label = years[-1] if years else ""

    city = ""
    if "深圳" in text:
        city = "深圳"
    else:
        city_match = re.search(
            r"(?:省|自治区)([^省市区县]{2,8})市",
            text,
        ) or re.search(r"([^省市区县]{2,8})市", text)
        if city_match:
            city = city_match.group(1).strip()

    phase = next(
        (
            label
            for token, label in (
                ("期末", "期末"),
                ("期中", "期中"),
                ("一模", "一模"),
                ("二模", "二模"),
                ("三模", "三模"),
                ("月考", "月考"),
                ("中考", "中考"),
                ("联考", "联考"),
            )
            if token in text
        ),
        "",
    )
    location_phase = f"{city}{phase}"
    parts = [part for part in (year_label, location_phase) if part]
    return f"（{'·'.join(parts)}）" if parts else ""


_SCORE_PREFIX = re.compile(r"^\s*[（(]\s*(\d+)\s*分")


def answer_space_lines(
    question_type: object,
    question_text: object = None,
) -> int:
    """Reserved handwriting lines: selection none; big questions by score.

    题型标注不可靠（大量解答题被标成填空），题干开头的（N 分）更可信：
    6 分及以上按分值留空间（约 1 行/分，封顶 10 行）。
    """
    value = str(question_type or "").casefold()
    if any(token in value for token in ("选择", "choice")):
        return 0
    match = _SCORE_PREFIX.match(str(question_text or ""))
    score = int(match.group(1)) if match else 0
    if score >= 6:
        return min(score, 10)
    if any(token in value for token in ("填空", "fill")):
        return 0
    return 3


def add_answer_space(
    document,
    *,
    question_paragraphs: Sequence[object] = (),
    minimum_lines: int = 3,
    line_height_mm: float = 9.0,
    content_width_dxa: int = 9865,
) -> None:
    """Reserve handwriting space; a trailing figure moves beside it.

    题末有独立图片时，图片排右、作答区排左侧整列，高度约为
    普通留白的一半（学生直接在图旁作答，省页且可见）。
    """

    if int(minimum_lines) <= 0:
        return
    paragraphs = list(question_paragraphs)
    trailing_picture = (
        paragraphs[-1]
        if paragraphs and _is_standalone_picture_paragraph(paragraphs[-1]._p)
        else None
    )
    text_paragraphs = [
        paragraph
        for paragraph in paragraphs
        if _paragraph_visible_text(paragraph._p)
    ]
    if text_paragraphs:
        text_paragraphs[-1].paragraph_format.keep_with_next = True
    blank = OxmlElement("w:p")
    _style_rich_paragraphs(blank, line_spacing=1.1)
    if trailing_picture is not None:
        minimum_height_mm = float(minimum_lines) * float(line_height_mm) / 2
        widths = _split_widths(content_width_dxa, right_fraction=0.42)
        _fit_picture_to_width(
            trailing_picture._p,
            max_width_dxa=widths[1] - 160,
        )
        table = _build_borderless_layout_table(
            columns=((blank,), (trailing_picture._p,)),
            widths=widths,
            minimum_height_mm=minimum_height_mm,
        )
    else:
        minimum_height_mm = float(minimum_lines) * float(line_height_mm)
        table = _build_borderless_layout_table(
            columns=((blank,),),
            widths=(int(content_width_dxa),),
            minimum_height_mm=minimum_height_mm,
        )
    _append_to_document_body(document, table)


def _is_option_picture_paragraph(element) -> bool:
    if element.tag != qn("w:p") or not _paragraph_has_picture(element):
        return False
    return bool(_OPTION_LETTER_PREFIX.match(_paragraph_visible_text(element)))


def _grid_option_picture_groups(
    prepared: list[tuple[object, dict[str, Path]]],
    *,
    content_width_dxa: int,
) -> tuple[list[tuple[object, dict[str, Path]]], int]:
    """Pack consecutive A./B./C./D. picture options into two-column rows."""
    result: list[tuple[object, dict[str, Path]]] = []
    gridded = 0
    index = 0
    while index < len(prepared):
        if not _is_option_picture_paragraph(prepared[index][0]):
            result.append(prepared[index])
            index += 1
            continue
        run: list[tuple[object, dict[str, Path]]] = []
        while index < len(prepared) and _is_option_picture_paragraph(
            prepared[index][0]
        ):
            run.append(prepared[index])
            index += 1
        if len(run) < 2:
            result.extend(run)
            continue
        widths = _split_widths(content_width_dxa, right_fraction=0.5)
        for start in range(0, len(run), 2):
            pair = run[start : start + 2]
            left_element, left_relationships = pair[0]
            if len(pair) == 2:
                right_element, right_relationships = pair[1]
                relationships = {**left_relationships, **right_relationships}
                columns = ((left_element,), (right_element,))
            else:
                relationships = dict(left_relationships)
                columns = ((left_element,), (OxmlElement("w:p"),))
            table = _build_borderless_layout_table(
                columns=columns,
                widths=widths,
            )
            result.append((table, relationships))
            gridded += 1
            # 相邻表格之间必须有段落分隔，否则 Word/WPS 会把它们合并成一个表格。
            result.append((OxmlElement("w:p"), {}))
    return result, gridded


def _compact_image_pairs(
    prepared: list[tuple[object, dict[str, Path]]],
    *,
    content_width_dxa: int,
) -> tuple[list[tuple[object, dict[str, Path]]], int]:
    result = list(prepared)
    compacted = 0
    index = 1
    while index < len(result):
        previous, previous_relationships = result[index - 1]
        picture, picture_relationships = result[index]
        if (
            previous.tag == qn("w:p")
            and _paragraph_visible_text(previous)
            and not _paragraph_has_picture(previous)
            and picture.tag == qn("w:p")
            and _is_standalone_picture_paragraph(picture)
        ):
            widths = _split_widths(content_width_dxa, right_fraction=0.32)
            _fit_picture_to_width(
                picture,
                max_width_dxa=min(
                    widths[1] - 160,
                    _OBJECTIVE_PICTURE_MAX_WIDTH_DXA,
                ),
            )
            table = _build_borderless_layout_table(
                columns=((previous,), (picture,)),
                widths=widths,
            )
            result[index - 1 : index + 1] = [
                (
                    table,
                    {**previous_relationships, **picture_relationships},
                )
            ]
            compacted += 1
            continue
        index += 1
    return result, compacted


def _split_widths(total_width: int, *, right_fraction: float) -> tuple[int, int]:
    right = int(round(int(total_width) * float(right_fraction)))
    return int(total_width) - right, right


def _build_borderless_layout_table(
    *,
    columns: Sequence[Sequence[object]],
    widths: Sequence[int],
    minimum_height_mm: float | None = None,
):
    if len(columns) != len(widths) or not columns:
        raise ValueError("layout table columns and widths must match")
    total_width = sum(int(width) for width in widths)
    table = OxmlElement("w:tbl")
    properties = OxmlElement("w:tblPr")
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(total_width))
    table_width.set(qn("w:type"), "dxa")
    properties.append(table_width)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    properties.append(borders)
    table.append(properties)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(width)))
        grid.append(column)
    table.append(grid)
    row = OxmlElement("w:tr")
    row_properties = OxmlElement("w:trPr")
    row_properties.append(OxmlElement("w:cantSplit"))
    if minimum_height_mm is not None:
        height = OxmlElement("w:trHeight")
        height.set(
            qn("w:val"),
            str(int(round(float(minimum_height_mm) * 56.7))),
        )
        height.set(qn("w:hRule"), "atLeast")
        row_properties.append(height)
    row.append(row_properties)
    for column_index, (elements, width) in enumerate(zip(columns, widths, strict=True)):
        cell = OxmlElement("w:tc")
        cell_properties = OxmlElement("w:tcPr")
        cell_width = OxmlElement("w:tcW")
        cell_width.set(qn("w:w"), str(int(width)))
        cell_width.set(qn("w:type"), "dxa")
        cell_properties.append(cell_width)
        vertical = OxmlElement("w:vAlign")
        vertical.set(qn("w:val"), "top")
        cell_properties.append(vertical)
        margins = OxmlElement("w:tcMar")
        for edge, value in (
            ("top", 0),
            ("left", 0 if column_index == 0 else 80),
            ("bottom", 0),
            ("right", 80 if column_index < len(columns) - 1 else 0),
        ):
            margin = OxmlElement(f"w:{edge}")
            margin.set(qn("w:w"), str(value))
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        cell_properties.append(margins)
        cell.append(cell_properties)
        appended = False
        for element in elements:
            cell.append(element)
            appended = True
        if not appended:
            cell.append(OxmlElement("w:p"))
        row.append(cell)
    table.append(row)
    return table


def _paragraph_visible_text(paragraph) -> str:
    return "".join(
        child.text or ""
        for child in paragraph.iter()
        if child.tag in {qn("w:t"), qn("m:t")}
    ).strip()


def _paragraph_has_picture(paragraph) -> bool:
    return any(
        child.tag in {qn("w:drawing"), qn("w:pict")}
        for child in paragraph.iter()
    )


def _is_standalone_picture_paragraph(paragraph) -> bool:
    return _paragraph_has_picture(paragraph) and not _paragraph_visible_text(paragraph)


def _fit_picture_to_width(paragraph, *, max_width_dxa: int) -> None:
    # Source Word paragraph indents consume width inside the new picture cell.
    # Its layout is now defined by the cell, so retain no inherited indentation.
    for properties in paragraph.iter(qn("w:pPr")):
        indentation = properties.find(qn("w:ind"))
        if indentation is not None:
            properties.remove(indentation)
    max_width_emu = max(int(max_width_dxa), 1) * 635
    for extent in paragraph.iter():
        if extent.tag != qn("wp:extent"):
            continue
        try:
            width = int(extent.get("cx") or 0)
            height = int(extent.get("cy") or 0)
        except ValueError:
            continue
        if width <= max_width_emu or width <= 0 or height <= 0:
            continue
        ratio = max_width_emu / width
        new_height = int(round(height * ratio))
        extent.set("cx", str(max_width_emu))
        extent.set("cy", str(new_height))
        for child in paragraph.iter():
            if child.tag == qn("a:ext"):
                child.set("cx", str(max_width_emu))
                child.set("cy", str(new_height))


def _style_rich_paragraphs(element, *, line_spacing: float) -> None:
    line_value = str(int(round(float(line_spacing) * 240)))
    for paragraph in element.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph.insert(0, properties)
        spacing = properties.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            properties.append(spacing)
        spacing.set(qn("w:line"), line_value)
        spacing.set(qn("w:lineRule"), "auto")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        if properties.find(qn("w:widowControl")) is None:
            properties.append(OxmlElement("w:widowControl"))


_OPTION_LETTER_PREFIX = re.compile(r"^\s*[A-DＡ-Ｄ][.、．]")
_OPTION_PICTURE_MAX_HEIGHT_MM = 22.0
# 客观题配图上限：能看清即可，不超过 55mm，避免占掉作答空间。
_OBJECTIVE_PICTURE_MAX_WIDTH_DXA = 3119


def _cap_option_picture_height(element) -> None:
    """选项（A./B./C./D.）行内的图片限高，避免选项图撑满整行。"""
    max_height_emu = int(_OPTION_PICTURE_MAX_HEIGHT_MM * 36000)
    for paragraph in element.iter(qn("w:p")):
        visible_text = "".join(
            child.text or ""
            for child in paragraph.iter()
            if child.tag in {qn("w:t"), qn("m:t")}
        ).strip()
        if not _OPTION_LETTER_PREFIX.match(visible_text):
            continue
        for extent in paragraph.iter(qn("wp:extent")):
            try:
                width = int(extent.get("cx") or 0)
                height = int(extent.get("cy") or 0)
            except ValueError:
                continue
            if height <= max_height_emu or height <= 0 or width <= 0:
                continue
            ratio = max_height_emu / height
            new_width = int(round(width * ratio))
            extent.set("cx", str(new_width))
            extent.set("cy", str(max_height_emu))
            for child in paragraph.iter():
                if child.tag == qn("a:ext"):
                    child.set("cx", str(new_width))
                    child.set("cy", str(max_height_emu))


def _align_standalone_picture_paragraphs_right(element) -> None:
    for paragraph in element.iter(qn("w:p")):
        has_picture = any(
            child.tag in {qn("w:drawing"), qn("w:pict")}
            for child in paragraph.iter()
        )
        visible_text = "".join(
            child.text or ""
            for child in paragraph.iter()
            if child.tag in {qn("w:t"), qn("m:t")}
        ).strip()
        if not has_picture or visible_text:
            continue
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph.insert(0, properties)
        alignment = properties.find(qn("w:jc"))
        if alignment is None:
            alignment = OxmlElement("w:jc")
            properties.append(alignment)
        alignment.set(qn("w:val"), "right")
        if properties.find(qn("w:keepLines")) is None:
            properties.append(OxmlElement("w:keepLines"))


def _prepend_inline_prefix(
    elements: Sequence[object],
    value: str,
    *,
    style: WordStyleProfile,
) -> bool:
    for element in elements:
        if element.tag != qn("w:p"):
            continue
        has_picture = any(
            child.tag in {qn("w:drawing"), qn("w:pict")}
            for child in element.iter()
        )
        visible_text = "".join(
            child.text or ""
            for child in element.iter()
            if child.tag in {qn("w:t"), qn("m:t")}
        ).strip()
        if has_picture and not visible_text:
            continue
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), style.body_font_ascii)
        fonts.set(qn("w:hAnsi"), style.body_font_ascii)
        fonts.set(qn("w:eastAsia"), style.body_font)
        run_properties.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(int(round(style.body_size_pt * 2))))
        run_properties.append(size)
        run.append(run_properties)
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = str(value)
        run.append(text)
        insertion_index = 1 if len(element) and element[0].tag == qn("w:pPr") else 0
        element.insert(insertion_index, run)
        return True
    return False


def _keep_text_with_following_picture(elements: Sequence[object]) -> None:
    for index, element in enumerate(elements):
        if element.tag != qn("w:p"):
            continue
        has_picture = any(
            child.tag in {qn("w:drawing"), qn("w:pict")}
            for child in element.iter()
        )
        if not has_picture:
            continue
        # 只锁定图片前的最后一个文本段落：图与引言不分页即可；
        # 整题加 keepNext 会把不可拆的长链整体推到下一页，造成空白首页。
        text_paragraphs = [
            paragraph
            for previous in elements[:index]
            for paragraph in previous.iter(qn("w:p"))
            if _paragraph_visible_text(paragraph)
        ]
        if not text_paragraphs:
            continue
        paragraph = text_paragraphs[-1]
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph.insert(0, properties)
        if properties.find(qn("w:keepNext")) is None:
            properties.append(OxmlElement("w:keepNext"))


def _append_to_document_body(document, element) -> None:
    body = document._body._element
    if len(body) and str(body[-1].tag).endswith("}sectPr"):
        body.insert(len(body) - 1, element)
    else:
        body.append(element)


def _strip_leading_question_number(element) -> None:
    for child in element.iter():
        if child.tag != qn("w:t"):
            continue
        text = child.text or ""
        if not text:
            continue
        stripped = _LEADING_QUESTION_NUMBER.sub("", text, count=1)
        if stripped != text:
            child.text = stripped.lstrip()
        return


__all__ = [
    "add_answer_space",
    "add_floating_picture",
    "answer_space_lines",
    "compact_source_label",
    "RichBlockRenderResult",
    "SharedWordQuestionRenderer",
    "WordStyleProfile",
    "rich_block_text",
    "validate_docx",
]
