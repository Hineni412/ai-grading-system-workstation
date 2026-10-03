from __future__ import annotations

import hashlib
import logging
import re
from collections.abc import Mapping
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from question_bank.document_pipeline.contracts import FormulaFallback, MathExpression
from question_bank.document_pipeline.legacy_exports import (
    PublishedMathMetadata,
    data_root_for_database,
    published_math_metadata,
    save_validated_legacy_export,
)
from question_bank.document_pipeline.word_renderer import (
    SharedWordQuestionRenderer,
    WordStyleProfile,
    add_answer_space,
    add_floating_picture,
    answer_space_lines,
    natural_image_width_inches,
)
from question_bank.exporters.base_exporter import (
    _resolve_image_path,
    _visible_length,
    add_noborder_table,
    apply_exporter_layout,
    extract_and_format_options,
)
from question_bank.exporters.export_config import ExportConfig
from question_bank.services.assembly_basket_state import SectionSpec
from question_bank.services.file_cache import cached_processed_image_digest
from question_bank.services.question_read_service import QuestionBankReadService
from question_bank.services.rich_content_service import (
    strip_question_source_score,
    strip_question_source_score_blocks,
)

LOGGER = logging.getLogger(__name__)
IMAGE_MARKER_PATTERN = re.compile(
    r"\[\[IMAGE:(?P<path>[^\]|]+?)(?:\|caption=(?P<caption>[^\]]*))?\]\]"
)
LEADING_QUESTION_NUMBER_PATTERN = re.compile(r"^\s*(?:第\s*)?\d{1,3}\s*(?:[.．、]|题)[ \t]*")

_CN_SECTION_NUMS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
                     "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八", "十九", "二十"]


def _image_paths_from_text(value: object) -> list[str]:
    return [match.group("path").strip() for match in IMAGE_MARKER_PATTERN.finditer(str(value or ""))]


def _canonical_type_group(qtype: str | None) -> str:
    qtype = str(qtype or "").strip()
    if qtype in ["choice", "single_choice", "multiple_choice", "multi_choice", "选择题", "多选题"]:
        return "选择题"
    if qtype in ["fill_blank", "blank", "填空题"]:
        return "填空题"
    return "解答题"


def _section_index_label(index: int) -> str:
    if 0 <= index < len(_CN_SECTION_NUMS):
        return _CN_SECTION_NUMS[index]
    return str(index + 1)


def _render_question_body(
    document: Document,
    question: dict[str, Any],
    index: int,
    config: ExportConfig,
    metadata_by_id: dict[int, PublishedMathMetadata],
    data_root: Path,
    fallbacks: list[FormulaFallback],
    *,
    trailing_blank: int = 0,
    include_answer_space: bool = True,
    question_notes: Mapping[int, str] | None = None,
) -> None:
    """渲染单个题目的正文（题号 + 富文本/纯文本 + 图片）。

    抽出来供 sections 分支复用，消除原三段式里重复的渲染块。
    """
    note = (question_notes or {}).get(int(question['id']))
    if note:
        run = document.add_paragraph().add_run(note)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string('666666')
    inline_prefix = _question_prefix(index)
    question_paragraph_start = len(document.paragraphs)
    minimum_lines = answer_space_lines(
        question.get("question_type"),
        question.get("question_text"),
    )
    style_profile = WordStyleProfile.from_export_config(config)
    metadata = metadata_by_id[int(question["id"])]
    rich_content = metadata.rich_content
    appended, embedded_paths = _add_rich_blocks(
        document,
        strip_question_source_score_blocks(
            _rich_blocks(rich_content, "question_blocks"),
            question_number=str(question.get("question_number") or ""),
        ),
        strip_leading_number=True,
        inline_prefix=inline_prefix,
        config=config,
        compact_standalone_images_with_text=(minimum_lines == 0),
        data_root=data_root,
    )

    all_image_paths = _dedupe_paths([
        *(question.get("image_paths") or []),
        *_image_paths_from_text(question.get("question_text") or "")
    ])
    resolved_all_paths = []
    for p in all_image_paths:
        rp = _resolve_image_path(p, data_root=data_root)
        if rp:
            resolved_all_paths.append(str(rp))
    # Legacy image_paths also contains pictures from the answer. Their rich
    # ownership is authoritative, even when this export hides answers.
    excluded_paths = embedded_paths | _rich_image_paths(rich_content, "answer_blocks", data_root)
    missing_images = _unembedded_image_paths(resolved_all_paths, excluded_paths)

    if not appended:
        _add_text_and_images(
            document,
            strip_question_source_score(
                question.get("question_text") or "",
                question_number=str(question.get("question_number") or ""),
            ),
            extra_image_paths=question.get("image_paths") or [],
            strip_leading_number=True,
            config=config,
            question_id=f"question-{int(question['id'])}",
            expressions=metadata.expressions,
            fallback_sink=fallbacks,
            data_root=data_root,
            inline_prefix=inline_prefix,
        )
    elif missing_images:
        _add_images(document, missing_images)
    if minimum_lines and include_answer_space:
        add_answer_space(
            document,
            question_paragraphs=document.paragraphs[question_paragraph_start:],
            minimum_lines=minimum_lines,
            content_width_dxa=style_profile.content_width_dxa,
        )
    else:
        for _ in range(trailing_blank):
            document.add_paragraph("")


def _render_answer_body(
    document: Document,
    question: dict[str, Any],
    index: int,
    config: ExportConfig,
    metadata_by_id: dict[int, PublishedMathMetadata],
    data_root: Path,
    fallbacks: list[FormulaFallback],
    *,
    trailing_blank: int = 1,
) -> None:
    """渲染单个题目的答案正文。"""
    inline_prefix = _question_prefix(index)
    metadata = metadata_by_id[int(question["id"])]
    rich_content = metadata.rich_content
    appended, embedded_paths = _add_rich_blocks(
        document,
        _rich_blocks(rich_content, "answer_blocks"),
        strip_leading_number=True,
        inline_prefix=inline_prefix,
        config=config,
        data_root=data_root,
    )

    all_image_paths = _dedupe_paths([*_image_paths_from_text(question.get("answer_text") or "")])
    resolved_all_paths = []
    for p in all_image_paths:
        rp = _resolve_image_path(p, data_root=data_root)
        if rp:
            resolved_all_paths.append(str(rp))
    missing_images = _unembedded_image_paths(resolved_all_paths, embedded_paths)

    if not appended:
        _add_text_and_images(
            document,
            question.get("answer_text") or "暂无答案",
            strip_leading_number=True,
            config=config,
            question_id=f"question-{int(question['id'])}-answer",
            fallback_sink=fallbacks,
            data_root=data_root,
            inline_prefix=inline_prefix,
        )
    elif missing_images:
        _add_images(document, missing_images)
    for _ in range(trailing_blank):
        document.add_paragraph("")


def _add_choice_answer_table(document: Document, choice_questions: list[tuple[int, dict[str, Any]]]) -> None:
    if not choice_questions:
        return

    document.add_paragraph("选择题答题区：")
    num_cols = len(choice_questions)

    table = document.add_table(rows=2, cols=num_cols + 1)
    table.style = 'Table Grid'

    table.cell(0, 0).paragraphs[0].text = "题号"
    table.cell(1, 0).paragraphs[0].text = "答案"

    for i, (index, _) in enumerate(choice_questions, start=1):
        table.cell(0, i).paragraphs[0].text = str(index)
        table.cell(1, i).paragraphs[0].text = ""

    document.add_paragraph("")



def export_question_paper_docx(
    db_path: str | Path,
    question_ids: list[int],
    output_dir: str | Path,
    *,
    title: str,
    include_answer: bool = False,
    ensure_previews: bool = True,
    grouped_by_type: bool = False,
    header_text: str | None = None,
    sections: list[SectionSpec] | None = None,
    config: ExportConfig | None = None,
    include_answer_space: bool = True,
    include_student_fields: bool = True,
    page_header_text: str | None = None,
    question_notes: Mapping[int, str] | None = None,
) -> Path:
    service = QuestionBankReadService(Path(db_path))
    questions = service.get_questions_for_export(question_ids)
    questions = [question for question in questions if question is not None]
    if not questions:
        raise ValueError("试题篮为空，无法导出 Word")

    output_path = _output_path(output_dir, title)
    document = Document()
    active_config = config or ExportConfig()
    apply_exporter_layout(document, active_config)
    data_root = data_root_for_database(db_path)
    style_profile = WordStyleProfile.from_export_config(active_config)
    metadata_by_id = {
        int(question["id"]): published_math_metadata(
            int(question["id"]),
            data_root=data_root,
            question_text=str(question.get("question_text") or ""),
            answer_text=str(question.get("answer_text") or ""),
        )
        for question in questions
    }
    fallbacks: list[FormulaFallback] = []

    if page_header_text:
        document.sections[0].header.paragraphs[0].text = page_header_text

    if header_text:
        header_p = document.add_paragraph()
        header_p.alignment = 1 # Center
        hrun = header_p.add_run(header_text.strip())
        hrun.font.size = Pt(14)
        hrun.bold = True

    document.add_heading(title.strip() or _default_title(), level=0)
    if include_student_fields:
        document.add_paragraph("姓名：________________    班级：________________    日期：________________")
    document.add_paragraph("")


    indexed_questions = list(enumerate(questions, start=1))

    if sections:
        # AI 智能分类编排：按外部传入的 sections 渲染
        question_by_id_docx = {int(q["id"]): q for q in questions}
        next_index = 1
        # 为答案区准备：每个 section 的 [(index, question)] 列表
        sections_indexed: list[tuple[str, list[tuple[int, dict[str, Any]]]]] = []

        for sec_idx, sec in enumerate(sections):
            sec_questions: list[tuple[int, dict[str, Any]]] = []
            for qid in sec.question_ids:
                q = question_by_id_docx.get(int(qid))
                if not q:
                    continue
                sec_questions.append((next_index, q))
                next_index += 1
            if not sec_questions:
                continue
            heading = f"{_section_index_label(sec_idx)}、{sec.title.strip()}"
            document.add_heading(heading, level=1)

            # 选择题答题区表格：段内若含选择题则触发
            choice_in_sec = [(idx, q) for idx, q in sec_questions if _canonical_type_group(q.get("question_type")) == "选择题"]
            if choice_in_sec and include_answer_space:
                _add_choice_answer_table(document, choice_in_sec)

            for idx, q in sec_questions:
                _render_question_body(
                    document,
                    q,
                    idx,
                    active_config,
                    metadata_by_id,
                    data_root,
                    fallbacks,
                    trailing_blank=0,
                    include_answer_space=include_answer_space,
                    question_notes=question_notes,
                )

            sections_indexed.append((sec.title.strip(), sec_questions))

        # 答案区：同样按 sections 分段
        if include_answer:
            document.add_page_break()
            document.add_heading("答案", level=1)
            for sec_idx, (sec_title, sec_qs) in enumerate(sections_indexed):
                heading = f"{_section_index_label(sec_idx)}、{sec_title} 答案"
                document.add_heading(heading, level=2)
                for idx, q in sec_qs:
                    _render_answer_body(
                        document,
                        q,
                        idx,
                        active_config,
                        metadata_by_id,
                        data_root,
                        fallbacks,
                        trailing_blank=1,
                    )

        save_validated_legacy_export(
            document,
            output_path,
            operation_namespace="ordinary-word",
            question_ids=tuple(f"question-{int(item['id'])}" for item in questions),
            content_revisions=tuple(
                metadata_by_id[int(item["id"])].content_revision for item in questions
            ),
            style=style_profile,
            fallbacks=fallbacks,
        )
        return output_path

    if grouped_by_type:
        choices, blanks, solutions = _group_indexed_questions(questions, active_config.numbering_mode)

        if choices:
            document.add_heading("一、选择题", level=1)
            if include_answer_space:
                _add_choice_answer_table(document, choices)
            for index, question in choices:
                _render_question_body(
                    document, question, index, active_config,
                    metadata_by_id, data_root, fallbacks, trailing_blank=0,
                    include_answer_space=include_answer_space,
                    question_notes=question_notes,
                )

        if blanks:
            heading_title = "二、填空题" if choices else "一、填空题"
            document.add_heading(heading_title, level=1)
            for index, question in blanks:
                _render_question_body(
                    document, question, index, active_config,
                    metadata_by_id, data_root, fallbacks, trailing_blank=0,
                    include_answer_space=include_answer_space,
                    question_notes=question_notes,
                )

        if solutions:
            if choices and blanks:
                heading_title = "三、解答题"
            elif choices or blanks:
                heading_title = "二、解答题"
            else:
                heading_title = "一、解答题"
            document.add_heading(heading_title, level=1)
            for index, question in solutions:
                _render_question_body(
                    document, question, index, active_config,
                    metadata_by_id, data_root, fallbacks, trailing_blank=0,
                    include_answer_space=include_answer_space,
                    question_notes=question_notes,
                )
    else:
        for index, question in indexed_questions:
            _render_question_body(
                document, question, index, active_config,
                metadata_by_id, data_root, fallbacks, trailing_blank=0,
                    include_answer_space=include_answer_space,
                    question_notes=question_notes,
            )

    if include_answer:
        document.add_page_break()
        document.add_heading("答案", level=1)

        if grouped_by_type:
            choices, blanks, solutions = _group_indexed_questions(questions, active_config.numbering_mode)
            for heading, group in (
                ("选择题答案", choices),
                ("填空题答案", blanks),
                ("解答题答案", solutions),
            ):
                if not group:
                    continue
                document.add_heading(heading, level=2)
                for index, question in group:
                    _render_answer_body(
                        document, question, index, active_config,
                        metadata_by_id, data_root, fallbacks, trailing_blank=1,
                    )
        else:
            for index, question in indexed_questions:
                _render_answer_body(
                    document, question, index, active_config,
                    metadata_by_id, data_root, fallbacks,
                    trailing_blank=1 if index < len(questions) else 0,
                )

    save_validated_legacy_export(
        document,
        output_path,
        operation_namespace="ordinary-word",
        question_ids=tuple(f"question-{int(item['id'])}" for item in questions),
        content_revisions=tuple(
            metadata_by_id[int(item["id"])].content_revision for item in questions
        ),
        style=style_profile,
        fallbacks=fallbacks,
    )
    return output_path


def _group_indexed_questions(
    questions: list[dict[str, object]],
    numbering_mode: str = "global",
) -> tuple[list[tuple[int, dict[str, object]]], list[tuple[int, dict[str, object]]], list[tuple[int, dict[str, object]]]]:
    raw_choices: list[dict[str, object]] = []
    raw_blanks: list[dict[str, object]] = []
    raw_solutions: list[dict[str, object]] = []
    for question in questions:
        group = _canonical_type_group(question.get("question_type"))
        if group == "选择题":
            raw_choices.append(question)
        elif group == "填空题":
            raw_blanks.append(question)
        else:
            raw_solutions.append(question)

    grouped: list[list[tuple[int, dict[str, object]]]] = []
    next_index = 1
    for raw_group in (raw_choices, raw_blanks, raw_solutions):
        if numbering_mode == "per_section":
            next_index = 1
        numbered_group = []
        for question in raw_group:
            numbered_group.append((next_index, question))
            next_index += 1
        grouped.append(numbered_group)
    return grouped[0], grouped[1], grouped[2]


def _rich_blocks(rich_content: dict[str, object] | None, key: str) -> list[dict[str, object]]:
    if not isinstance(rich_content, dict):
        return []
    blocks = rich_content.get(key)
    if not isinstance(blocks, list):
        return []
    return [block for block in blocks if isinstance(block, dict)]


def _add_rich_blocks(
    document: Document,
    blocks: list[dict[str, object]],
    *,
    strip_leading_number: bool = False,
    inline_prefix: str = "",
    config: ExportConfig | None = None,
    compact_standalone_images_with_text: bool = False,
    data_root: Path | None = None,
) -> tuple[bool, set[str]]:
    renderer = SharedWordQuestionRenderer(
        style=WordStyleProfile.from_export_config(config or ExportConfig()),
        asset_resolver=lambda value: _resolve_image_path(value, data_root=data_root),
    )
    result = renderer.add_rich_blocks(
        document,
        blocks,
        strip_leading_number=strip_leading_number,
        inline_prefix=inline_prefix,
        compact_standalone_images_with_text=compact_standalone_images_with_text,
    )
    return result.appended, set(result.embedded_assets)


def _unembedded_image_paths(paths, embedded_paths) -> list[str]:
    """Supplemental assets are not a second copy of explicit rich drawings.

    Keep all occurrences inside the rich content. Compare full pixels without
    resampling; unknown formats use byte equality only.
    """
    if not embedded_paths:
        return [str(path) for path in paths]

    def compute(path: Path) -> str:
        try:
            with Image.open(path) as image:
                if image.format not in {"PNG", "JPEG"}:
                    return "bytes:" + hashlib.sha256(path.read_bytes()).hexdigest()
                digest = hashlib.sha256(str(image.size).encode())
                digest.update(image.convert("RGBA").tobytes())
                digest.update(image.info.get("icc_profile") or b"")
                digest.update(str(image.getexif().get(274, 1)).encode())
                return "pixels:" + digest.hexdigest()
        except (OSError, ValueError):
            return "bytes:" + hashlib.sha256(path.read_bytes()).hexdigest()

    def identity(value):
        path = Path(value)
        return cached_processed_image_digest(path, variant="export-full-pixels-v1", compute=compute)

    seen = {identity(path) for path in embedded_paths if Path(path).is_file()}
    remaining = []
    for value in paths:
        digest = identity(value)
        if digest not in seen:
            remaining.append(str(value))
    return remaining


def _rich_image_paths(rich_content, key, data_root) -> set[str]:
    paths = set()
    for block in _rich_blocks(rich_content, key):
        for value in (block.get("image_relationships") or {}).values():
            path = _resolve_image_path(value, data_root=data_root)
            if path is not None:
                paths.add(str(path))
    return paths


_MD_TABLE_SEP_CELL = re.compile(r":?-{2,}:?")


def _iter_render_items(lines: list[str]):
    """连续的 Markdown 管道行合并为 ('table', rows)，其余行原样为 ('line', text)。"""
    table_rows: list[list[str]] = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 3:
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if cells and all(_MD_TABLE_SEP_CELL.fullmatch(cell) for cell in cells):
                continue
            table_rows.append(cells)
            continue
        if table_rows:
            yield "table", table_rows
            table_rows = []
        if re.fullmatch(r"\$+", stripped):
            # 识别残留的孤立公式定界符行，不落为文字。
            continue
        yield "line", line
    if table_rows:
        yield "table", table_rows


def _add_markdown_table(
    document: Document,
    rows: list[list[str]],
    *,
    render_cell,
    suffix_prefix: str,
) -> None:
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return
    cols = max(len(row) for row in rows)
    if len(rows) == 1:
        # 单行管道行多为图注/并排布局，不要边框。
        table = add_noborder_table(document, 1, cols)
    else:
        table = document.add_table(rows=len(rows), cols=cols)
        try:
            table.style = "Table Grid"
        except Exception:
            pass
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            value = row[col_index] if col_index < len(row) else ""
            render_cell(
                cell.paragraphs[0],
                value,
                suffix=f"{suffix_prefix}-r{row_index + 1}c{col_index + 1}",
            )


_IMG_SENTINEL = re.compile(r"⟦IMG(?P<idx>\d+)⟧")


def _add_text_and_images(
    document: Document,
    text: object,
    *,
    extra_image_paths: list[object] | None = None,
    strip_leading_number: bool = False,
    config: ExportConfig | None = None,
    question_id: str = "ordinary-export",
    expressions: tuple[MathExpression, ...] = (),
    fallback_sink: list[FormulaFallback] | None = None,
    data_root: str | Path | None = None,
    inline_prefix: str = "",
) -> None:
    active_config = config or ExportConfig()
    raw_text = str(text or "")
    # 图片标记换成短哨兵，保住图片在行文中的原位置（避免全部堆到题末）。
    marker_entries = [
        (match.group("path").strip(), (match.group("caption") or "").strip())
        for match in IMAGE_MARKER_PATTERN.finditer(raw_text)
    ]
    marker_paths = [path for path, _caption in marker_entries]
    marker_captions = {
        index: caption
        for index, (_path, caption) in enumerate(marker_entries)
        if caption
    }
    marker_index = 0

    def _next_sentinel(match: re.Match[str]) -> str:
        nonlocal marker_index
        marker_index += 1
        return f"⟦IMG{marker_index - 1}⟧"

    sentinel_text = IMAGE_MARKER_PATTERN.sub(_next_sentinel, raw_text)
    clean_text = sentinel_text.strip()
    if strip_leading_number:
        clean_text = LEADING_QUESTION_NUMBER_PATTERN.sub("", clean_text, count=1).lstrip()

    rendered_images: set[str] = set()
    # 题干足够长时，第一张图浮动锚定到题目首段右侧环绕；短题干浮动只会
    # 压住下一题，仍按独立段落插图。浮动锚点目标段是题目正文第一段。
    float_first = (
        bool(marker_paths)
        and len(re.sub(r"\s+", "", _IMG_SENTINEL.sub("", clean_text))) >= 40
    )
    floating_image: tuple[Path, str] | None = None
    first_body_paragraph = len(document.paragraphs)

    if clean_text:
        renderer = SharedWordQuestionRenderer(
            style=WordStyleProfile.from_export_config(active_config),
            asset_resolver=lambda value: _resolve_image_path(value, data_root=data_root),
        )

        prefix_pending = inline_prefix

        def render_line(value: str, *, suffix: str) -> None:
            nonlocal prefix_pending
            fallbacks = renderer.add_text(
                document,
                value,
                question_id=question_id,
                expressions=expressions,
                inline_prefix=prefix_pending,
            )
            prefix_pending = ""
            if fallback_sink is not None:
                fallback_sink.extend(fallbacks)
            for fallback in fallbacks:
                LOGGER.warning(
                    "Word formula fallback for %s: %s",
                    fallback.expression_id,
                    fallback.reason,
                )

        def emit_sentinel_image(index: int) -> None:
            nonlocal floating_image
            if 0 <= index < len(marker_paths):
                path = marker_paths[index]
                caption = marker_captions.get(index, "")
                if index == 0 and float_first:
                    resolved = _resolve_image_path(path, data_root=data_root)
                    if resolved is not None:
                        rendered_images.add(path)
                        floating_image = (resolved, caption)
                        return
                rendered_images.add(path)
                _add_images(
                    document,
                    [path],
                    data_root=data_root,
                    captions={path: caption} if caption else None,
                )

        def render_flow(value: str, *, suffix: str) -> None:
            """渲染含 ⟦IMGn⟧ 哨兵的行：文字走公式渲染，哨兵处就地插图。"""
            cursor = 0
            for match in _IMG_SENTINEL.finditer(value):
                segment = value[cursor : match.start()].strip()
                if segment:
                    render_line(segment, suffix=suffix)
                emit_sentinel_image(int(match.group("idx")))
                cursor = match.end()
            tail = value[cursor:].strip()
            if tail:
                render_line(tail, suffix=suffix)

        def render_cell(
            paragraph,
            value: str,
            *,
            suffix: str,
            image_option_cell: bool = False,
        ) -> None:
            # 单元格内的 ⟦IMGn⟧ 哨兵就地转成内嵌小图；图片选项单元格
            # 限高 22mm、限宽 1.6 英寸，保证 1×4 行内不顶破单元格。
            picture_max_width = 1.6 if image_option_cell else 2.5
            picture_max_height = 0.87 if image_option_cell else None
            cursor = 0
            for match in _IMG_SENTINEL.finditer(value):
                segment = value[cursor : match.start()].strip()
                if segment:
                    _render_cell_text(paragraph, segment, suffix=suffix)
                index = int(match.group("idx"))
                if 0 <= index < len(marker_paths):
                    path_text = marker_paths[index]
                    resolved = _resolve_image_path(path_text, data_root=data_root)
                    if resolved is not None:
                        rendered_images.add(path_text)
                        try:
                            paragraph.add_run().add_picture(
                                str(resolved),
                                width=Inches(
                                    natural_image_width_inches(
                                        resolved,
                                        max_width_inches=picture_max_width,
                                        max_height_inches=picture_max_height,
                                    )
                                ),
                            )
                        except Exception:
                            pass
                cursor = match.end()
            tail = value[cursor:].strip()
            if tail:
                _render_cell_text(paragraph, tail, suffix=suffix)

        def _render_cell_text(paragraph, value: str, *, suffix: str) -> None:
            fallbacks = renderer.add_to_paragraph(
                paragraph,
                value,
                question_id=question_id,
                expressions=expressions,
            )
            if fallback_sink is not None:
                fallback_sink.extend(fallbacks)
            for fallback in fallbacks:
                LOGGER.warning(
                    "Word formula fallback for %s: %s",
                    fallback.expression_id,
                    fallback.reason,
                )

        # Check if we can extract options
        stem, options = extract_and_format_options(clean_text)
        if options:
            # Render stem
            for line_index, (kind, item) in enumerate(
                _iter_render_items(stem.splitlines())
            ):
                if kind == "table":
                    _add_markdown_table(
                        document, item, render_cell=render_cell,
                        suffix_prefix=f"stem-{line_index + 1}",
                    )
                elif item.strip():
                    render_flow(item.strip(), suffix=f"stem-{line_index + 1}")
            # Render options in table/columns
            image_options = len(options) >= 3 and all(
                _IMG_SENTINEL.fullmatch(opt_text.strip())
                for _, opt_text in options
            )
            max_len = max(
                _visible_length(opt_text) for _, opt_text in options
            )
            if image_options:
                # 图片选项一律单行铺开（不足 4 个留空单元格），底边对齐。
                table = add_noborder_table(document, 1, max(len(options), 4))
                for i, (label, opt_text) in enumerate(options):
                    cell = table.cell(0, i)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
                    p = cell.paragraphs[0]
                    render_cell(
                        p,
                        f"{label}. {opt_text}",
                        suffix=f"option-{i + 1}",
                        image_option_cell=True,
                    )
            elif float_first:
                # 浮动图不环绕表格：改段落排布，短选项一行两项。
                # 哨兵 0 在题尾（选项区之后）时，图片在收尾阶段仍会浮动，
                # 这里同样按浮动布局排选项，避免表格右行被图覆盖。
                if max_len <= 12:
                    for pair_start in range(0, len(options), 2):
                        p = document.add_paragraph()
                        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.1))
                        label, opt_text = options[pair_start]
                        _render_cell_text(
                            p,
                            f"{label}. {opt_text}",
                            suffix=f"option-{pair_start + 1}",
                        )
                        if pair_start + 1 < len(options):
                            p.add_run("\t")
                            label, opt_text = options[pair_start + 1]
                            _render_cell_text(
                                p,
                                f"{label}. {opt_text}",
                                suffix=f"option-{pair_start + 2}",
                            )
                else:
                    for option_index, (label, opt_text) in enumerate(options):
                        render_flow(
                            f"{label}. {opt_text}",
                            suffix=f"option-{option_index + 1}",
                        )
            elif max_len <= 6:
                table = add_noborder_table(document, 1, 4)
                for i, (label, opt_text) in enumerate(options):
                    cell = table.cell(0, i)
                    p = cell.paragraphs[0]
                    render_cell(p, f"{label}. {opt_text}", suffix=f"option-{i + 1}")
            elif max_len <= 15:
                table = add_noborder_table(document, 2, 2)
                opt_coords = [(0, 0), (0, 1), (1, 0), (1, 1)]
                for idx, (label, opt_text) in enumerate(options):
                    r_idx, c_idx = opt_coords[idx]
                    cell = table.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    render_cell(p, f"{label}. {opt_text}", suffix=f"option-{idx + 1}")
            else:
                for option_index, (label, opt_text) in enumerate(options):
                    render_flow(
                        f"{label}. {opt_text}",
                        suffix=f"option-{option_index + 1}",
                    )
        else:
            for line_index, (kind, item) in enumerate(
                _iter_render_items(clean_text.splitlines())
            ):
                if kind == "table":
                    _add_markdown_table(
                        document, item, render_cell=render_cell,
                        suffix_prefix=f"line-{line_index + 1}",
                    )
                elif item.strip():
                    render_flow(item.strip(), suffix=f"line-{line_index + 1}")

    elif inline_prefix:
        paragraph = document.add_paragraph(inline_prefix)
        paragraph.paragraph_format.line_spacing = active_config.line_spacing
        paragraph.paragraph_format.keep_with_next = True

    # 哨兵 0 落在选项区或题尾而未原位渲染时，仍把首图浮动到题干首段右侧。
    if (
        float_first
        and floating_image is None
        and marker_paths
        and marker_paths[0] not in rendered_images
    ):
        resolved = _resolve_image_path(marker_paths[0], data_root=data_root)
        if resolved is not None:
            rendered_images.add(marker_paths[0])
            floating_image = (resolved, marker_captions.get(0, ""))

    if floating_image is not None and first_body_paragraph < len(document.paragraphs):
        floating_path, floating_caption = floating_image
        width_inches = natural_image_width_inches(
            floating_path,
            max_width_inches=2.8,
            max_height_inches=3.0,
        )
        anchored = add_floating_picture(
            document.paragraphs[first_body_paragraph],
            floating_path,
            width_inches=width_inches,
            descr=floating_caption or None,
        )
        if not anchored:
            _add_images(
                document,
                [str(floating_path)],
                data_root=data_root,
                captions={str(floating_path): floating_caption}
                if floating_caption
                else None,
            )
        # 浮动图挂在首段；题目末段不要 keep_with_next 黏住下一题。
        if document.paragraphs:
            document.paragraphs[-1].paragraph_format.keep_with_next = False

    # 只有未被原位渲染的图片（外部 image_paths / 未命中的哨兵）才补到题末。
    image_paths = _dedupe_paths(
        [
            *[str(p) for p in (extra_image_paths or [])],
            *[p for p in marker_paths if p not in rendered_images],
        ]
    )
    image_paths = [p for p in image_paths if p not in rendered_images]
    leftover_captions = {
        marker_paths[index]: caption
        for index, caption in marker_captions.items()
        if marker_paths[index] in image_paths
    }
    if image_paths and document.paragraphs:
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    _add_images(
        document,
        image_paths,
        data_root=data_root,
        captions=leftover_captions or None,
    )


_MAX_IMAGE_WIDTH_INCHES = 4.8


def _add_images(
    document: Document,
    image_paths: list[object],
    *,
    data_root: str | Path | None = None,
    captions: Mapping[str, str] | None = None,
) -> None:
    for image_path in image_paths:
        path = _resolve_image_path(str(image_path), data_root=data_root)
        if path is None:
            continue
        try:
            width_inches = natural_image_width_inches(
                path, max_width_inches=_MAX_IMAGE_WIDTH_INCHES
            )
            document.add_picture(str(path), width=Inches(width_inches))
            paragraph = document.paragraphs[-1]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.keep_together = True
            caption = str((captions or {}).get(str(image_path)) or "").strip()
            if caption:
                paragraph.paragraph_format.keep_with_next = True
                caption_paragraph = document.add_paragraph()
                caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_run = caption_paragraph.add_run(caption)
                caption_run.font.size = Pt(9)
        except Exception:
            continue


def _dedupe_paths(paths: list[object]) -> list[str]:
    deduped: list[str] = []
    for path in paths:
        text = str(path or "").strip()
        if text and text not in deduped:
            deduped.append(text)
    return deduped


def _output_path(output_dir: str | Path, title: str) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    safe_title = _safe_filename(title) or "question_paper"
    return output_dir / f"{safe_title}_{timestamp}.docx"


def _default_title() -> str:
    return f"{datetime.now().strftime('%Y-%m-%d')}习题"


def _safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\s]+', "_", str(value or "")).strip("._")


def _question_prefix(index: int) -> str:
    return f"{index}. "


__all__ = ["export_question_paper_docx"]
