from __future__ import annotations

import hashlib
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from question_bank.exporters.export_config import ExportConfig
from question_bank.services.asset_path_service import resolve_question_bank_asset_path

LOGGER = logging.getLogger(__name__)


def _resolve_image_path(
    saved_path_str: str,
    *,
    data_root: str | Path | None = None,
) -> Path | None:
    if not saved_path_str:
        return None
    resolved = resolve_question_bank_asset_path(
        saved_path_str,
        data_root=data_root,
        search_subdirs=("question_bank/extracted_images", "question_bank/previews"),
    )
    return resolved if resolved.is_file() else None


def add_page_number_to_paragraph(paragraph) -> None:
    # Add page number XML fields: "第 X 页 / 共 Y 页"
    paragraph.alignment = 1  # Center
    paragraph.add_run("第 ")

    # PAGE field
    page_run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    page_run._r.append(fldChar1)
    page_run._r.append(instrText)
    page_run._r.append(fldChar2)
    page_run._r.append(fldChar3)

    paragraph.add_run(" 页 / 共 ")

    # NUMPAGES field
    numpages_run = paragraph.add_run()
    fldChar1_num = OxmlElement('w:fldChar')
    fldChar1_num.set(qn('w:fldCharType'), 'begin')
    instrText_num = OxmlElement('w:instrText')
    instrText_num.set(qn('xml:space'), 'preserve')
    instrText_num.text = "NUMPAGES"
    fldChar2_num = OxmlElement('w:fldChar')
    fldChar2_num.set(qn('w:fldCharType'), 'separate')
    fldChar3_num = OxmlElement('w:fldChar')
    fldChar3_num.set(qn('w:fldCharType'), 'end')
    numpages_run._r.append(fldChar1_num)
    numpages_run._r.append(instrText_num)
    numpages_run._r.append(fldChar2_num)
    numpages_run._r.append(fldChar3_num)

    paragraph.add_run(" 页")


def apply_exporter_layout(document, config: ExportConfig) -> None:
    # Margins setup
    for section in document.sections:
        section.top_margin = Cm(config.margin_top_cm)
        section.bottom_margin = Cm(config.margin_bottom_cm)
        section.left_margin = Cm(config.margin_left_cm)
        section.right_margin = Cm(config.margin_right_cm)

        # Enable page numbers in footer
        if config.show_page_numbers:
            footer = section.footer
            footer_p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            footer_p.text = ""  # Clear existing
            add_page_number_to_paragraph(footer_p)

    # Normal Style settings
    normal_style = document.styles["Normal"]
    normal_style.font.name = config.body_font_ascii
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), config.body_font)
    normal_style.font.size = Pt(config.body_size_pt)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.line_spacing = config.line_spacing


def latex_to_png(latex_str: str, config: ExportConfig) -> Path | None:
    latex_str = latex_str.strip()
    if not latex_str:
        return None

    # Compute hash of LaTeX for caching
    h = hashlib.md5(latex_str.encode("utf-8")).hexdigest()
    cache_dir = Path(config.latex_cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    output_path = cache_dir / f"latex_{h}.png"

    if output_path.exists():
        return output_path

    # Try 1: Standalone compiler if pdflatex is in PATH
    if shutil.which("pdflatex"):
        try:
            if _render_via_pdflatex(latex_str, output_path, dpi=config.latex_dpi):
                return output_path
        except Exception as exc:
            LOGGER.warning("pdflatex compilation failed for latex: %s, error: %s", latex_str, exc)

    # Try 2: Matplotlib mathtext fallback
    try:
        if _render_via_matplotlib(latex_str, output_path, dpi=config.latex_dpi):
            return output_path
    except Exception as exc:
        LOGGER.warning("Matplotlib fallback latex compilation failed: %s, error: %s", latex_str, exc)

    return None


def _render_via_pdflatex(latex_str: str, output_path: Path, dpi: int) -> bool:
    clean_latex = latex_str.strip()
    # If it has math mode delimiters, remove them since we wrap them ourselves in standalone
    if clean_latex.startswith("$$") and clean_latex.endswith("$$"):
        clean_latex = clean_latex[2:-2].strip()
    elif clean_latex.startswith("$") and clean_latex.endswith("$"):
        clean_latex = clean_latex[1:-1].strip()

    tex_template = r"""\documentclass[preview,border=1pt,varwidth]{standalone}
\usepackage{amsmath}
\usepackage{amssymb}
\begin{document}
$""" + clean_latex + r"""$
\end{document}
"""

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        tex_file = tmpdir_path / "temp.tex"
        tex_file.write_text(tex_template, encoding="utf-8")

        res = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "temp.tex"],
            cwd=str(tmpdir_path),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=10,
        )
        if res.returncode != 0:
            return False

        pdf_file = tmpdir_path / "temp.pdf"
        if not pdf_file.exists():
            return False

        # Convert PDF to PNG using PyMuPDF (fitz)
        import fitz
        with fitz.open(pdf_file) as doc:
            page = doc[0]
            zoom = dpi / 72.0
            mat = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat)
            pix.save(str(output_path))
            return True


def _render_via_matplotlib(latex_str: str, output_path: Path, dpi: int) -> bool:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt

    clean_latex = latex_str.strip()
    if not clean_latex.startswith("$") and not clean_latex.endswith("$"):
        clean_latex = f"${clean_latex}$"

    fig = plt.figure(figsize=(0.1, 0.1))
    fig.text(0.5, 0.5, clean_latex, ha='center', va='center', fontsize=11)

    # Save figure as transparent png
    fig.savefig(str(output_path), dpi=dpi, bbox_inches='tight', pad_inches=0.02, transparent=True)
    plt.close(fig)
    return True


import re


def parse_latex_runs(text: str) -> list[tuple[str, str]]:
    runs = []
    last_idx = 0
    pattern = re.compile(r'\$\$(.*?)\$\$|\$(.*?)\$', re.DOTALL)
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            runs.append(("text", text[last_idx:start]))

        block, inline = match.groups()
        if block is not None:
            runs.append(("block_math", block))
        else:
            runs.append(("inline_math", inline))
        last_idx = end

    if last_idx < len(text):
        runs.append(("text", text[last_idx:]))
    return runs


def add_paragraph_with_latex(document, text: str, config: ExportConfig) -> None:
    runs = parse_latex_runs(text)
    p = document.add_paragraph()
    # Keep paragraph heading or layout together if possible
    p.paragraph_format.widow_control = True

    for run_type, content in runs:
        if run_type == "text":
            p.add_run(content)
        elif run_type == "inline_math":
            png_path = latex_to_png(content, config)
            if png_path and png_path.exists():
                try:
                    from PIL import Image
                    with Image.open(png_path) as img:
                        w_px, h_px = img.size
                    w_pt = (w_px / config.latex_dpi) * 72 * 0.85

                    run = p.add_run()
                    run.add_picture(str(png_path), width=Pt(w_pt))
                except Exception as exc:
                    LOGGER.warning("Failed to insert inline math image: %s", exc)
                    p.add_run(f"${content}$")
            else:
                p.add_run(f"${content}$")
        elif run_type == "block_math":
            p_block = document.add_paragraph()
            p_block.alignment = 1  # Center
            p_block.paragraph_format.widow_control = True
            png_path = latex_to_png(content, config)
            if png_path and png_path.exists():
                try:
                    from PIL import Image
                    with Image.open(png_path) as img:
                        w_px, h_px = img.size
                    w_pt = (w_px / config.latex_dpi) * 72
                    w_pt = min(w_pt, 400.0)

                    run = p_block.add_run()
                    run.add_picture(str(png_path), width=Pt(w_pt))
                except Exception:
                    p_block.add_run(f"$$\n{content}\n$$")
            else:
                p_block.add_run(f"$$\n{content}\n$$")


def add_noborder_table(document, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tblPr.remove(borders)

    new_borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        new_borders.append(border)
    tblPr.append(new_borders)
    return table


def _visible_length(text: str) -> int:
    """选项布局用的可见长度：LaTeX 源码折成渲染后的近似字符数。"""
    value = re.sub(r"\$", "", str(text or ""))
    value = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"\1/\2", value)
    value = re.sub(r"\\[A-Za-z]+", "", value)
    value = re.sub(r"[{}^_]", "", value)
    value = re.sub(r"\s+", "", value)
    return len(value)


_OPTION_LABEL = r'(?<![A-Za-z])[A-D](?:[\.．、\)]|）)'
# 教辅把答案印在选项区末尾，形如“（B）”。
_EMBEDDED_OPTION_ANSWER = re.compile(r'[（(]\s*[A-D]\s*[)）]')
_TRAILING_OPTION_ANSWER = re.compile(r'[（(]\s*[A-D]\s*[)）]\s*$')


def extract_and_format_options(text: str) -> tuple[str, list[tuple[str, str]]]:
    # 标签前只要不是字母或数字即视为选项区起点（允许紧贴中文，如“根是A.”）。
    pattern = re.compile(r'(?<![A-Za-z0-9])(A)(?:[\.．、\)]|）)\s*')
    match = pattern.search(text)
    if not match:
        return text, []

    start_idx = match.start(1)
    stem = text[:start_idx].strip()
    # 先去掉选项区末尾的内嵌答案，避免被拆成幻影选项。
    options_text = _TRAILING_OPTION_ANSWER.sub("", text[start_idx:])

    # 选项值允许跨行，但在下一个标签或 [[IMAGE:]] 标记行之前停下，
    # 否则末选项遇到跟随的图片标记行会整体匹配失败而丢失。
    opt_pattern = re.compile(
        r'(?<![A-Za-z])([A-D])(?:[\.．、\)]|）)\s*(.*?)(?='
        + _OPTION_LABEL + r'|\n\[\[IMAGE:|\n⟦IMG\d+⟧|$)',
        flags=re.S,
    )
    matches = opt_pattern.findall(options_text)

    # 选项标签必须递增；递增不上的视为内嵌答案等噪声。
    ordered: list[tuple[str, str]] = []
    for label, value in matches:
        if ordered and label <= ordered[-1][0]:
            continue
        ordered.append((label, value.strip()))

    # OCR 偶发丢一个选项；3 个标签也按选项区拆分。
    if len(ordered) < 3:
        return text, []

    opts = ordered[:4]
    if opts:
        label, value = opts[-1]
        opts[-1] = (label, _EMBEDDED_OPTION_ANSWER.sub("", value).strip())
    return stem, opts
