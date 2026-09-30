from __future__ import annotations

import json
import logging
import re
import sqlite3
from collections import defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches

from question_bank.document_pipeline.contracts import FormulaFallback, MathExpression
from question_bank.document_pipeline.legacy_exports import (
    data_root_for_database,
    published_math_metadata,
    save_validated_legacy_export,
)
from question_bank.document_pipeline.word_renderer import (
    SharedWordQuestionRenderer,
    WordStyleProfile,
)
from question_bank.exporters.base_exporter import (
    _resolve_image_path,
    apply_exporter_layout,
)
from question_bank.exporters.export_config import ExportConfig
from question_bank.services.question_frequency_service import (
    FrequencyMetrics,
    QuestionFrequencyService,
    frequency_summary,
)

LOGGER = logging.getLogger(__name__)

STAGE_ORDER = ["基础回补", "方法形成", "典型模型", "综合提升", "压轴迁移"]
STAGE_TITLES = {
    "基础回补": "一、基础回补",
    "方法形成": "二、方法形成",
    "典型模型": "三、典型模型",
    "综合提升": "四、综合提升",
    "压轴迁移": "五、压轴迁移",
}
AUDIENCE_LABELS = {
    "student": "学生版",
    "teacher": "教师版",
}


@dataclass(frozen=True)
class ExportQuestion:
    question_id: int
    source_paper: str
    question_number: str
    question_text: str
    answer_text: str = ""
    page_range: str = ""
    image_paths: list[str] = field(default_factory=list)
    difficulty: str = ""
    frequency: str = ""
    knowledge_points: list[str] = field(default_factory=list)
    method_tags: list[str] = field(default_factory=list)
    error_prone_points: list[str] = field(default_factory=list)
    model_tags: list[str] = field(default_factory=list)
    recommend_reason: str = ""
    suggested_order: int = 0
    training_stage: str = "基础回补"
    task_item_code: str = ""
    content_revision: str = ""
    math_expressions: tuple[MathExpression, ...] = ()


def export_training_docx(
    db_path: str | Path,
    recommendations: Iterable[Mapping[str, Any]],
    output_dir: str | Path,
    *,
    audience: str,
    display_name: str | None = None,
    student_id: str | None = None,
    class_id: str | None = None,
    use_real_name: bool = False,
    task_code: str | None = None,
    variant_code: str | None = None,
    config: ExportConfig | None = None,
) -> Path:
    """Export selected recommendations to a student or teacher DOCX file."""

    audience = _normalize_audience(audience)
    items = load_export_questions(db_path, recommendations)
    if not items:
        raise ValueError("没有可导出的推荐题目")

    output_path = _output_path(
        output_dir,
        audience=audience,
        title_name=resolve_title_name(
            display_name=display_name,
            student_id=student_id,
            class_id=class_id,
            use_real_name=use_real_name,
        ),
        suffix=".docx",
    )

    document = Document()
    active_config = config or ExportConfig()
    apply_exporter_layout(document, active_config)
    data_root = data_root_for_database(db_path)
    style_profile = WordStyleProfile.from_export_config(active_config)
    renderer = SharedWordQuestionRenderer(
        style=style_profile,
        asset_resolver=lambda value: _resolve_image_path(value, data_root=data_root),
    )

    title_name = resolve_title_name(
        display_name=display_name,
        student_id=student_id,
        class_id=class_id,
        use_real_name=use_real_name,
    )
    document.add_heading(f"{title_name} 专项训练", level=0)
    document.add_paragraph("说明：根据最近一次考试薄弱点生成")
    document.add_paragraph(f"版本：{AUDIENCE_LABELS[audience]}")
    if audience == "teacher" and task_code:
        document.add_paragraph(f"训练任务：{task_code}")
    if audience == "teacher" and variant_code:
        document.add_paragraph(f"训练版本：{variant_code}")

    fallbacks: list[FormulaFallback] = []
    for stage, stage_items in _group_by_stage(items).items():
        document.add_heading(STAGE_TITLES.get(stage, stage), level=1)
        for item in stage_items:
            fallbacks.extend(
                _add_question_to_docx(
                    document,
                    item,
                    include_teacher_fields=audience == "teacher",
                    renderer=renderer,
                )
            )

    save_validated_legacy_export(
        document,
        output_path,
        operation_namespace="training-word",
        question_ids=tuple(f"question-{item.question_id}" for item in items),
        content_revisions=tuple(item.content_revision for item in items),
        style=style_profile,
        fallbacks=fallbacks,
    )
    return output_path


def load_export_questions(
    db_path: str | Path,
    recommendations: Iterable[Mapping[str, Any]],
) -> list[ExportQuestion]:
    recommendation_list = [dict(item) for item in recommendations]
    if not recommendation_list:
        return []
    db_path = Path(db_path)
    if not db_path.exists():
        raise FileNotFoundError(f"题库数据库不存在：{db_path}")

    question_ids = [int(item["question_id"]) for item in recommendation_list if item.get("question_id")]
    if not question_ids:
        return []

    details = _load_question_details(db_path, question_ids)
    tags = _load_question_tags(db_path, question_ids)
    frequencies = QuestionFrequencyService(db_path).metrics_for_questions(question_ids)
    data_root = data_root_for_database(db_path)
    result: list[ExportQuestion] = []
    for recommendation in recommendation_list:
        question_id = int(recommendation.get("question_id") or 0)
        snapshot = recommendation.get("question_snapshot")
        detail = dict(snapshot) if isinstance(snapshot, Mapping) else details.get(question_id)
        if not detail:
            continue
        metadata = published_math_metadata(
            question_id,
            data_root=data_root,
            question_text=_text(_row_value(detail, "question_text")),
            answer_text=_text(_row_value(detail, "answer_text")),
        )
        tag_map = _merged_tags(tags.get(question_id, {}), detail)
        result.append(
            ExportQuestion(
                question_id=question_id,
                source_paper=_source_display(detail, recommendation),
                question_number=_text(_row_value(detail, "question_number") or recommendation.get("question_number")),
                question_text=_text(_row_value(detail, "question_text")),
                answer_text=_text(_row_value(detail, "answer_text")),
                page_range=_text(_row_value(detail, "page_range")),
                image_paths=_parse_image_paths(_row_value(detail, "image_paths")),
                difficulty=_text(_row_value(detail, "difficulty") or recommendation.get("difficulty")),
                frequency=_frequency_display(recommendation, frequencies.get(question_id)),
                knowledge_points=_unique_strings(
                    [*tag_map.get("knowledge_point", []), *recommendation.get("knowledge_points", [])]
                ),
                method_tags=_unique_strings([*tag_map.get("method", []), *recommendation.get("method_tags", [])]),
                error_prone_points=_unique_strings(tag_map.get("error_type", [])),
                model_tags=_unique_strings(tag_map.get("model", [])),
                recommend_reason=_text(recommendation.get("recommend_reason")),
                suggested_order=_int(recommendation.get("suggested_order")),
                training_stage=_text(recommendation.get("training_stage")) or "基础回补",
                task_item_code=_text(recommendation.get("task_item_code")),
                content_revision=metadata.content_revision,
                math_expressions=metadata.expressions,
            )
        )
    return sorted(result, key=lambda item: (STAGE_ORDER.index(item.training_stage) if item.training_stage in STAGE_ORDER else 99, item.suggested_order, item.question_id))


def resolve_title_name(
    *,
    display_name: str | None = None,
    student_id: str | None = None,
    class_id: str | None = None,
    use_real_name: bool = False,
) -> str:
    if use_real_name and _text(display_name):
        return _text(display_name)
    return _text(student_id) or _text(class_id) or "学生A"


def teaching_tip(item: ExportQuestion) -> str:
    stage_tips = {
        "基础回补": "先让学生复述核心概念和基本关系，再完成同类基础变式。",
        "方法形成": "讲评时突出解题入口和方法选择，要求学生用一句话说明为什么这样做。",
        "典型模型": "引导学生标注模型特征，并和已学典型题建立对应关系。",
        "综合提升": "拆分条件链，先找关键中间量，再整合多个知识点完成推理。",
        "压轴迁移": "保留探索时间，重点追问辅助构造、分类边界和可迁移策略。",
    }
    tip = stage_tips.get(item.training_stage, "讲评时关注学生的卡点，并补充一题同类变式。")
    if item.error_prone_points:
        tip += f" 易错提醒：{', '.join(item.error_prone_points)}。"
    return tip


def _add_question_to_docx(
    document: Document,
    item: ExportQuestion,
    *,
    include_teacher_fields: bool,
    renderer: SharedWordQuestionRenderer,
) -> tuple[FormulaFallback, ...]:
    fallbacks: list[FormulaFallback] = []
    heading = f"{item.suggested_order or item.question_id}. 来源：{item.source_paper} 第{item.question_number}题"
    document.add_paragraph(heading)
    fallbacks.extend(
        _report_formula_fallbacks(
            renderer.add_text(
            document,
            item.question_text,
            question_id=f"training-{item.question_id}",
            expressions=item.math_expressions,
            )
        )
    )
    renderer.add_images(document, item.image_paths)
    if include_teacher_fields:
        if item.task_item_code:
            document.add_paragraph(f"任务题码：{item.task_item_code}")
        document.add_paragraph("答案：")
        fallbacks.extend(
            _report_formula_fallbacks(
                renderer.add_text(
                document,
                item.answer_text or "（暂无答案）",
                question_id=f"training-{item.question_id}-answer",
                )
            )
        )
        document.add_paragraph(f"知识点：{_join_or_dash(item.knowledge_points)}")
        document.add_paragraph(f"方法标签：{_join_or_dash(item.method_tags)}")
        document.add_paragraph(f"难度：{item.difficulty or '-'}")
        document.add_paragraph(f"考频：{item.frequency or '仅期中、期末、中考试题计算'}")
        document.add_paragraph(f"推荐原因：{item.recommend_reason or '-'}")
        document.add_paragraph(f"教学提示：{teaching_tip(item)}")
    else:
        document.add_paragraph("答题区：")
        for _ in range(4):
            document.add_paragraph("____________________________________________________________")
    return tuple(fallbacks)


def _report_formula_fallbacks(
    fallbacks: tuple[FormulaFallback, ...],
) -> tuple[FormulaFallback, ...]:
    for fallback in fallbacks:
        LOGGER.warning(
            "Word formula fallback for %s: %s",
            fallback.expression_id,
            fallback.reason,
        )
    return fallbacks


def _add_images(document: Document, image_paths: list[str]) -> None:
    for image_path in image_paths:
        resolved = _resolve_image_path(image_path)
        if not resolved or not resolved.exists():
            document.add_paragraph(f"图像：{image_path}（未找到）")
            continue
        try:
            document.add_picture(str(resolved), width=Inches(4.8))
        except Exception:
            document.add_paragraph(f"图像：{image_path}（无法插入）")


def _group_by_stage(items: list[ExportQuestion]) -> dict[str, list[ExportQuestion]]:
    grouped: dict[str, list[ExportQuestion]] = {}
    by_stage: defaultdict[str, list[ExportQuestion]] = defaultdict(list)
    for item in items:
        by_stage[item.training_stage].append(item)
    for stage in STAGE_ORDER:
        if by_stage.get(stage):
            grouped[stage] = by_stage[stage]
    for stage, stage_items in by_stage.items():
        if stage not in grouped:
            grouped[stage] = stage_items
    return grouped


def _load_question_details(db_path: Path, question_ids: list[int]) -> dict[int, sqlite3.Row]:
    placeholders = ", ".join("?" for _ in question_ids)
    db_uri = f"{db_path.resolve().as_uri()}?mode=ro"
    with sqlite3.connect(db_uri, uri=True) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            f"""
            SELECT
                q.id,
                q.question_number,
                q.question_text,
                q.answer_text,
                q.source_file AS question_source_file,
                q.page_range,
                q.image_paths,
                q.difficulty,
                p.title AS paper_title,
                p.source_file AS paper_source_file,
                p.year,
                p.district,
                p.exam_type
            FROM questions q
            LEFT JOIN papers p ON p.id = q.paper_id
            WHERE q.id IN ({placeholders})
              AND COALESCE(q.is_deleted, 0) = 0
            """,
            question_ids,
        ).fetchall()
    return {int(row["id"]): row for row in rows}


def _load_question_tags(db_path: Path, question_ids: list[int]) -> dict[int, dict[str, list[str]]]:
    placeholders = ", ".join("?" for _ in question_ids)
    db_uri = f"{db_path.resolve().as_uri()}?mode=ro"
    tags: dict[int, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    with sqlite3.connect(db_uri, uri=True) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            f"""
            SELECT question_id, tag_type, tag_value
            FROM question_tags
            WHERE question_id IN ({placeholders})
            ORDER BY id ASC
            """,
            question_ids,
        ).fetchall()
    for row in rows:
        question_id = int(row["question_id"])
        tag_type = _text(row["tag_type"])
        tag_value = _text(row["tag_value"])
        if tag_type and tag_value and tag_value not in tags[question_id][tag_type]:
            tags[question_id][tag_type].append(tag_value)
    return {question_id: dict(tag_map) for question_id, tag_map in tags.items()}


def _source_display(detail: Mapping[str, Any] | sqlite3.Row, recommendation: Mapping[str, Any]) -> str:
    year = _text(_row_value(detail, "year"))
    region_exam = f"{_text(_row_value(detail, 'district'))}{_text(_row_value(detail, 'exam_type'))}"
    source = " ".join(part for part in [year, region_exam] if part)
    return source or _text(_row_value(detail, "paper_title")) or _text(recommendation.get("source_paper")) or _text(_row_value(detail, "paper_source_file")) or _text(_row_value(detail, "question_source_file"))


def _output_path(output_dir: str | Path, *, audience: str, title_name: str, suffix: str) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    safe_title = _safe_filename(title_name) or "student"
    path = output_dir / f"{safe_title}_training_{audience}_{timestamp}{suffix}"
    counter = 1
    while path.exists():
        path = output_dir / f"{safe_title}_training_{audience}_{timestamp}_{counter}{suffix}"
        counter += 1
    return path


def _parse_image_paths(value: object) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return _unique_strings(value)
    text = _text(value)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return [text] if text else []
    if isinstance(parsed, list):
        return _unique_strings(parsed)
    if isinstance(parsed, str):
        return [parsed]
    return []


def _normalize_audience(audience: str) -> str:
    audience = _text(audience).lower()
    if audience not in AUDIENCE_LABELS:
        raise ValueError("audience 必须是 student 或 teacher")
    return audience


def _safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\s]+', "_", value).strip("._")


def _join_or_dash(values: list[str]) -> str:
    return "、".join(values) if values else "-"


def _unique_strings(values: Iterable[object]) -> list[str]:
    result: list[str] = []
    for value in values:
        text = _text(value)
        if text and text not in result:
            result.append(text)
    return result


def _int(value: object) -> int:
    try:
        return int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 0


def _text(value: object) -> str:
    return str(value or "").strip()


def _merged_tags(
    current: Mapping[str, list[str]],
    detail: Mapping[str, Any] | sqlite3.Row,
) -> dict[str, list[str]]:
    result = {key: list(values) for key, values in current.items()}
    snapshot_tags = _row_value(detail, "tags")
    if isinstance(snapshot_tags, list):
        for tag in snapshot_tags:
            if not isinstance(tag, Mapping):
                continue
            tag_type = _text(tag.get("tag_type"))
            tag_value = _text(tag.get("tag_value"))
            values = result.setdefault(tag_type, [])
            if tag_type and tag_value and tag_value not in values:
                values.append(tag_value)
    return result


def _frequency_display(
    recommendation: Mapping[str, Any],
    current: FrequencyMetrics | None,
) -> str:
    snapshot = recommendation.get("frequency")
    if isinstance(snapshot, Mapping):
        fields = FrequencyMetrics.__dataclass_fields__
        values = {key: snapshot[key] for key in fields if key in snapshot}
        try:
            return frequency_summary(FrequencyMetrics(**values))
        except (TypeError, ValueError):
            pass
    return frequency_summary(current)


def _row_value(row: Mapping[str, Any] | sqlite3.Row, key: str) -> object:
    return row[key] if key in row.keys() else None
