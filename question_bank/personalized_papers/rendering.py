from __future__ import annotations

import hashlib
import hmac
import json
import os
import re
import shutil
import subprocess
from collections.abc import Mapping, Sequence
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from typing import Any, Protocol

import cv2
import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image
from question_bank.document_pipeline import (
    SharedWordQuestionRenderer,
    WordStyleProfile,
    add_answer_space,
    answer_space_lines,
    build_math_expression,
    rich_block_text,
)
from question_bank.document_pipeline.contracts import math_expression_from_payload


LAYOUT_VERSION = "personalized-paper-school-a4-v3"
PAGE_IDENTITY_VERSION = "P4P2"
LEGACY_PAGE_IDENTITY_VERSION = "P4P1"
DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PDF_MEDIA_TYPE = "application/pdf"

_INK = RGBColor(30, 41, 59)
_BLUE = RGBColor(30, 84, 120)
_MUTED = RGBColor(100, 116, 139)
_FONT_LATIN = "Calibri"
_FONT_EAST_ASIA = "Microsoft YaHei"
_MATH_RUN = re.compile(r"\$\$(.+?)\$\$|\$(.+?)\$", re.DOTALL)


class PaperRenderError(RuntimeError):
    pass


class PdfConversionAdapter(Protocol):
    def convert(self, source_docx: Path, output_pdf: Path) -> None: ...


class OfficePdfConverter:
    """Convert a reviewed DOCX through WPS/Word or a local LibreOffice binary."""

    def __init__(self, *, timeout_seconds: int = 120) -> None:
        self.timeout_seconds = int(timeout_seconds)

    def convert(self, source_docx: Path, output_pdf: Path) -> None:
        source = Path(source_docx).resolve()
        destination = Path(output_pdf).resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)
        if os.name == "nt":
            script = Path(__file__).with_name("office_to_pdf.ps1")
            pwsh = shutil.which("pwsh") or r"C:\Program Files\PowerShell\7\pwsh.exe"
            if Path(pwsh).exists() and script.is_file():
                completed = subprocess.run(
                    [
                        str(pwsh),
                        "-NoLogo",
                        "-NoProfile",
                        "-File",
                        str(script),
                        "-InputPath",
                        str(source),
                        "-OutputPath",
                        str(destination),
                    ],
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=self.timeout_seconds,
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                )
                if (
                    completed.returncode == 0
                    and destination.is_file()
                    and destination.stat().st_size > 0
                ):
                    return
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if executable:
            completed = subprocess.run(
                [
                    executable,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(destination.parent),
                    str(source),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=self.timeout_seconds,
            )
            produced = destination.parent / f"{source.stem}.pdf"
            if completed.returncode == 0 and produced.is_file():
                if produced != destination:
                    os.replace(produced, destination)
                if destination.stat().st_size > 0:
                    return
        raise PaperRenderError("no compatible DOCX to PDF converter is available")


def render_review_docx(
    snapshot: Mapping[str, Any],
    *,
    data_root: Path,
    output_path: Path,
) -> tuple[dict[str, Any], ...]:
    document = Document()
    _configure_document(document)
    student = _mapping(snapshot.get("student"))
    items = _mappings(snapshot.get("items"))
    paper_id = str(snapshot["paper_instance_id"])
    formula_fallbacks: list[dict[str, Any]] = []
    renderer = SharedWordQuestionRenderer(
        style=WordStyleProfile(
            body_font=_FONT_EAST_ASIA,
            body_font_ascii=_FONT_LATIN,
            body_size_pt=11,
            formula_size_pt=11,
            line_spacing=1.1,
        ),
        asset_resolver=lambda value: _controlled_asset_path(data_root, value),
    )

    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(
        header,
        (
            f"{student.get('student_name') or student.get('student_code') or student.get('student_id')}"
            "  ·  个性化训练卷"
        ),
        size=9,
        color=_MUTED,
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    _add_run(title, "个性化训练卷", size=18, bold=True, color=_BLUE)

    identity = document.add_paragraph()
    identity.paragraph_format.space_before = Pt(1)
    identity.paragraph_format.space_after = Pt(7)
    identity.paragraph_format.keep_with_next = True
    _add_run(
        identity,
        (
            f"姓名：{student.get('student_name') or student.get('student_code') or '学生'}"
            f"    班级：{student.get('class_id') or '________'}"
            "    日期：____年__月__日"
        ),
        size=10.5,
        color=_INK,
    )
    _add_hidden_run(identity, paper_id)
    _paragraph_bottom_border(identity, color="94A3B8", size="5")

    for index, item in enumerate(items, start=1):
        question = _mapping(item.get("question_snapshot"))
        context = _mapping(question.get("tagging_context"))
        question_type = str(context.get("question_type") or "")
        inline_prefix = f"{index}. "
        question_paragraph_start = len(document.paragraphs)
        rich_result = renderer.add_rich_blocks(
            document,
            _mappings(question.get("rich_question_blocks")),
            strip_leading_number=True,
            inline_prefix=inline_prefix,
            compact_standalone_images_with_text=(
                answer_space_lines(question_type) == 0
            ),
        )
        if not rich_result.appended:
            body = document.add_paragraph()
            body.paragraph_format.keep_with_next = True
            body.paragraph_format.space_after = Pt(0)
            body.paragraph_format.line_spacing = 1.1
            _add_run(body, inline_prefix, size=11, bold=True, color=_INK)
            expressions = tuple(
                math_expression_from_payload(value)
                for value in _mappings(question.get("math_expressions"))
            )
            if not expressions:
                images = _mappings(question.get("images"))
                fallback_asset = images[0].get("asset_path") if images else None
                fallback_sha256 = images[0].get("sha256") if images else None
                expressions = tuple(
                    build_math_expression(
                        expression_id=f"p4-{question.get('question_id') or index}-math-{expression_index}",
                        source=match.group(1) if match.group(1) is not None else match.group(2),
                        fallback_asset=str(fallback_asset) if fallback_asset else None,
                        fallback_sha256=str(fallback_sha256) if fallback_sha256 else None,
                    )
                    for expression_index, match in enumerate(
                        _MATH_RUN.finditer(str(context.get("question_text") or "")),
                        start=1,
                    )
                )
            if any(
                not expression.omml and not expression.fallback_asset
                for expression in expressions
            ):
                raise PaperRenderError(
                    "unsupported formula has no governed source image fallback"
                )
            formula_fallbacks.extend(
                asdict(fallback)
                for fallback in renderer.add_to_paragraph(
                    body,
                    str(context.get("question_text") or ""),
                    question_id=str(
                        question.get("question_id")
                        or item.get("question_id")
                        or index
                    ),
                    expressions=expressions,
                )
            )
        image_paragraph_start = len(document.paragraphs)
        _add_question_images(
            document,
            question,
            data_root=Path(data_root),
            skip_assets=set(rich_result.embedded_assets),
        )
        for image_paragraph in document.paragraphs[image_paragraph_start:]:
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            image_paragraph.paragraph_format.keep_together = True
        question_paragraphs = document.paragraphs[question_paragraph_start:]
        minimum_lines = answer_space_lines(question_type)
        if minimum_lines:
            add_answer_space(
                document,
                question_paragraphs=question_paragraphs,
                minimum_lines=minimum_lines,
                content_width_dxa=renderer.style.content_width_dxa,
            )
        marker_target = _last_document_body_paragraph(document)
        if marker_target is None:
            marker_target = document.add_paragraph()
            marker_target.paragraph_format.space_after = Pt(0)
        _add_hidden_run(marker_target, f"任务题码：{item['task_item_code']}")

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(footer, "第 ", size=8, color=_MUTED)
    _add_field(footer, "PAGE")
    _add_run(footer, " 页 / 共 ", size=8, color=_MUTED)
    _add_field(footer, "NUMPAGES")
    _add_run(footer, " 页", size=8, color=_MUTED)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    if not output.is_file() or output.stat().st_size <= 0:
        raise PaperRenderError("review DOCX was not created")
    return tuple(formula_fallbacks)


def stamp_frozen_pdf(
    source_pdf: Path,
    output_pdf: Path,
    *,
    paper_instance_id: str,
    paper_batch_id: str,
    series_version: int,
    student_name: str,
    student_code: str,
    class_id: str,
    signing_secret: str,
    reviewed_docx_sha256: str,
    layout_version: str = LAYOUT_VERSION,
) -> tuple[dict[str, Any], ...]:
    source = fitz.open(source_pdf)
    try:
        if source.page_count <= 0:
            raise PaperRenderError("converted PDF has no pages")
        total_pages = source.page_count
        pages: list[dict[str, Any]] = []
        for page_number, page in enumerate(source, start=1):
            signature = page_signature(
                signing_secret=signing_secret,
                paper_batch_id=paper_batch_id,
                paper_instance_id=paper_instance_id,
                series_version=series_version,
                page_number=page_number,
                total_pages=total_pages,
                reviewed_docx_sha256=reviewed_docx_sha256,
                layout_version=layout_version,
            )
            identity = page_identity(
                paper_instance_id=paper_instance_id,
                series_version=series_version,
                page_number=page_number,
                total_pages=total_pages,
                signature=signature,
            )
            qr = _qr_png(identity)
            width = float(page.rect.width)
            height = float(page.rect.height)
            footer_top = max(height - 72.0, 0.0)
            page.draw_rect(
                fitz.Rect(22.0, footer_top - 3.0, width - 18.0, height - 8.0),
                color=(1, 1, 1),
                fill=(1, 1, 1),
                overlay=True,
            )
            visible_name = str(student_name or student_code or "Student").strip()[:40]
            visible_class = str(class_id or "-").strip()[:30]
            visible_identity = (
                f"姓名：{visible_name}  班级：{visible_class}  "
                f"第 {page_number} 页 / 共 {total_pages} 页"
            )
            visible_font = "helv"
            font_path = _visible_identity_font()
            if font_path is not None:
                visible_font = "p4-cjk"
                page.insert_font(fontname=visible_font, fontfile=str(font_path))
            else:
                visible_identity = (
                    f"Student: {student_code or paper_instance_id[:12]}  "
                    f"Class: {visible_class}  Page {page_number}/{total_pages}"
                )
            page.insert_text(
                fitz.Point(26.0, footer_top + 13.0),
                visible_identity,
                fontsize=8.0,
                fontname=visible_font,
                color=(0.25, 0.32, 0.4),
                overlay=True,
            )
            page.insert_text(
                fitz.Point(26.0, footer_top + 27.0),
                f"Identity {signature[:16]}  Instance {paper_instance_id[:12]}  Layout {layout_version}",
                fontsize=6.5,
                fontname="helv",
                color=(0.35, 0.4, 0.47),
                overlay=True,
            )
            page.insert_image(
                fitz.Rect(width - 70.0, footer_top, width - 18.0, footer_top + 52.0),
                stream=qr,
                overlay=True,
            )
            pages.append(
                {
                    "page_number": page_number,
                    "total_pages": total_pages,
                    "page_identity": identity,
                    "page_signature": signature,
                    "layout_version": layout_version,
                    "width_points": width,
                    "height_points": height,
                }
            )
        output = Path(output_pdf)
        output.parent.mkdir(parents=True, exist_ok=True)
        source.save(output, garbage=4, deflate=True)
    finally:
        source.close()

    pdf_sha256 = _file_sha256(output_pdf)
    with fitz.open(output_pdf) as frozen:
        if frozen.page_count != len(pages):
            raise PaperRenderError("frozen PDF page count changed")
        finalized: list[dict[str, Any]] = []
        for page, item in zip(frozen, pages, strict=True):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            finalized.append(
                {
                    **item,
                    "page_content_hash": hashlib.sha256(
                        pixmap.samples
                    ).hexdigest(),
                    "pdf_sha256": pdf_sha256,
                }
            )
    return tuple(finalized)


def _visible_identity_font() -> Path | None:
    candidates = (
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    return next((path for path in candidates if path.is_file()), None)


def page_identity(
    *,
    paper_instance_id: str,
    series_version: int,
    page_number: int,
    total_pages: int,
    signature: str,
) -> str:
    # P4P2 stays inside QR's compact alphanumeric mode. The previous
    # lowercase, pipe-delimited byte payload produced data-dependent OpenCV
    # decode failures even before printing.
    return ":".join(
        (
            PAGE_IDENTITY_VERSION,
            paper_instance_id.upper(),
            str(int(series_version)),
            str(int(page_number)),
            str(int(total_pages)),
            signature.upper(),
        )
    )


def page_signature(
    *,
    signing_secret: str,
    paper_batch_id: str,
    paper_instance_id: str,
    series_version: int,
    page_number: int,
    total_pages: int,
    reviewed_docx_sha256: str,
    layout_version: str,
    identity_version: str = PAGE_IDENTITY_VERSION,
) -> str:
    message = "|".join(
        (
            identity_version,
            paper_batch_id,
            paper_instance_id,
            str(int(series_version)),
            str(int(page_number)),
            str(int(total_pages)),
            reviewed_docx_sha256,
            layout_version,
        )
    )
    return hmac.new(
        bytes.fromhex(signing_secret),
        message.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()[:32]


def decode_page_identity(value: str) -> dict[str, Any]:
    raw = str(value or "")
    if raw.startswith(f"{PAGE_IDENTITY_VERSION}:"):
        parts = raw.split(":")
    elif raw.startswith(f"{LEGACY_PAGE_IDENTITY_VERSION}|"):
        parts = raw.split("|")
    else:
        raise ValueError("page identity is invalid")
    if len(parts) != 6 or parts[0] not in {
        PAGE_IDENTITY_VERSION,
        LEGACY_PAGE_IDENTITY_VERSION,
    }:
        raise ValueError("page identity is invalid")
    instance_id, version, page, total, signature = parts[1:]
    if (
        len(instance_id) != 64
        or len(signature) != 32
    ):
        raise ValueError("page identity is invalid")
    try:
        version_number = int(version)
        page_number = int(page)
        total_pages = int(total)
    except ValueError as exc:
        raise ValueError("page identity is invalid") from exc
    if (
        version_number <= 0
        or page_number <= 0
        or total_pages <= 0
        or page_number > total_pages
    ):
        raise ValueError("page identity is invalid")
    return {
        "identity_version": parts[0],
        "paper_instance_id": instance_id.casefold(),
        "series_version": version_number,
        "page_number": page_number,
        "total_pages": total_pages,
        "page_signature": signature.casefold(),
    }


def inspect_docx(
    source: Path,
    *,
    paper_instance_id: str,
    task_item_codes: Sequence[str],
    question_snapshots: Sequence[Mapping[str, Any]],
) -> None:
    try:
        document = Document(source)
    except Exception as exc:  # noqa: BLE001
        raise PaperRenderError("reviewed file is not a readable DOCX") from exc
    text_tags = {qn("w:t"), qn("m:t")}
    text = "".join(
        child.text or ""
        for child in document._body._element.iter()  # noqa: SLF001
        if child.tag in text_tags
    )
    normalized = _normalized_text(text)
    instance_marker = _normalized_text(paper_instance_id)
    if not instance_marker or instance_marker not in normalized:
        raise PaperRenderError(
            "reviewed DOCX no longer matches the frozen question list"
        )
    _assert_ordered_text(normalized, task_item_codes)
    question_texts = []
    for question in question_snapshots:
        rich_text = rich_block_text(
            _mappings(question.get("rich_question_blocks")),
            strip_leading_number=True,
        )
        context = _mapping(question.get("tagging_context"))
        question_texts.append(
            rich_text
            if rich_text is not None
            else str(context.get("question_text") or "")
        )
    _assert_ordered_text(normalized, question_texts)


def pdf_page_count(path: Path) -> int:
    try:
        with fitz.open(path) as document:
            return int(document.page_count)
    except Exception as exc:  # noqa: BLE001
        raise PaperRenderError("PDF output is invalid") from exc


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(28)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    section.start_type = WD_SECTION.NEW_PAGE
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = _FONT_LATIN
    normal.font.size = Pt(11)
    normal.font.color.rgb = _INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_EAST_ASIA)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, _BLUE, 18, 10),
        ("Heading 2", 13, _BLUE, 14, 7),
        ("Heading 3", 12, _INK, 10, 5),
    ):
        style = styles[name]
        style.font.name = _FONT_LATIN
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_EAST_ASIA)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_question_images(
    document: Document,
    question: Mapping[str, Any],
    *,
    data_root: Path,
    skip_assets: set[str] | None = None,
) -> None:
    images = _mappings(question.get("images"))
    for image in images:
        if image.get("role") != "question":
            continue
        relative = Path(str(image.get("asset_path") or ""))
        resolved = (data_root / relative).resolve()
        root = data_root.resolve()
        if root not in resolved.parents or not resolved.is_file():
            raise PaperRenderError("frozen question image is missing")
        if str(resolved) in (skip_assets or set()):
            continue
        payload = resolved.read_bytes()
        try:
            document.add_picture(BytesIO(payload), width=Mm(120))
        except Exception:
            with Image.open(BytesIO(payload)) as source:
                converted = BytesIO()
                source.convert("RGB").save(converted, format="PNG")
                converted.seek(0)
                document.add_picture(converted, width=Mm(120))


def _add_run(
    paragraph,
    text: str,
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
) -> None:
    run = paragraph.add_run(text)
    run.font.name = _FONT_LATIN
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_EAST_ASIA)


def _add_hidden_run(paragraph, text: str) -> None:
    run = paragraph.add_run(str(text or ""))
    run.font.hidden = True


def _last_document_body_paragraph(document: Document) -> Paragraph | None:
    for element in reversed(document._body._element):  # noqa: SLF001
        if element.tag == qn("w:p"):
            return Paragraph(element, document._body)  # noqa: SLF001
        if element.tag == qn("w:tbl"):
            paragraphs = element.xpath(".//w:p")
            if paragraphs:
                return Paragraph(paragraphs[-1], document._body)  # noqa: SLF001
    return None


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, code, separate, text, end))


def _paragraph_bottom_border(paragraph, *, color: str, size: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _qr_png(payload: str) -> bytes:
    params = cv2.QRCodeEncoder_Params()
    params.correction_level = cv2.QRCodeEncoder_CORRECT_LEVEL_M
    # A fixed, slightly roomier symbol avoids OpenCV's data-dependent AUTO
    # version choices while the P4P2 payload remains comfortably within
    # capacity.
    params.version = 6
    encoder = cv2.QRCodeEncoder_create(params)
    image = encoder.encode(payload)
    image = cv2.copyMakeBorder(
        image,
        4,
        4,
        4,
        4,
        cv2.BORDER_CONSTANT,
        value=255,
    )
    image = cv2.resize(
        image,
        None,
        fx=5,
        fy=5,
        interpolation=cv2.INTER_NEAREST,
    )
    success, encoded = cv2.imencode(".png", image)
    if not success:
        raise PaperRenderError("page identity QR could not be encoded")
    return bytes(encoded)


def _all_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _normalized_text(value: object) -> str:
    return "".join(str(value or "").split())


def _assert_ordered_text(
    haystack: str,
    values: Sequence[str],
) -> None:
    cursor = 0
    for value in values:
        marker = _normalized_text(value)
        position = haystack.find(marker, cursor)
        if not marker or position < 0:
            raise PaperRenderError(
                "reviewed DOCX no longer matches the frozen question list"
            )
        cursor = position + len(marker)


def _mapping(value: object) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _controlled_asset_path(data_root: Path, value: str) -> Path | None:
    candidate = Path(str(value or ""))
    if candidate.is_absolute() or ".." in candidate.parts:
        return None
    root = Path(data_root).resolve()
    resolved = (root / candidate).resolve()
    if root != resolved and root not in resolved.parents:
        return None
    return resolved if resolved.is_file() else None


def _mappings(value: object) -> list[dict[str, Any]]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


__all__ = [
    "DOCX_MEDIA_TYPE",
    "LAYOUT_VERSION",
    "OfficePdfConverter",
    "PDF_MEDIA_TYPE",
    "PAGE_IDENTITY_VERSION",
    "PaperRenderError",
    "PdfConversionAdapter",
    "decode_page_identity",
    "inspect_docx",
    "page_signature",
    "pdf_page_count",
    "render_review_docx",
    "stamp_frozen_pdf",
]
