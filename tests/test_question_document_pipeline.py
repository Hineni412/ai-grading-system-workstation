from __future__ import annotations

import io
import json
import sqlite3
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

import fitz
import pytest
from docx import Document
from PIL import Image

from question_bank.database.schema import initialize_database
from question_bank.document_pipeline import (
    ApplyReviewCommand,
    DocumentPipelineConflict,
    DocumentSource,
    DocumentState,
    ExportState,
    ExportWordCommand,
    ManualQuestionRegion,
    OcrLine,
    PrepareSourceCommand,
    PublishCommand,
    QuestionBankPublicationAdapter,
    QuestionDocumentPipeline,
    QuestionReviewDecision,
    RecognitionSource,
    ReviewState,
    SharedWordQuestionRenderer,
    SourceRegion,
    TextLayerState,
    validate_docx,
)
from question_bank.document_pipeline.contracts import sha256_bytes
from question_bank.document_pipeline.legacy_exports import (
    export_receipt_path,
    load_legacy_export_receipt,
)
from question_bank.document_pipeline.math_omml import build_math_expression
from question_bank.exporters.docx_exporter import export_training_docx
from question_bank.exporters.paper_docx_exporter import export_question_paper_docx
from question_bank.importers.pdf_importer import import_pdf


def _region(
    page_number: int,
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    *,
    source: RecognitionSource = RecognitionSource.MANUAL,
    confidence: float = 1.0,
) -> SourceRegion:
    return SourceRegion(
        page_number=page_number,
        polygon=((x0, y0), (x1, y0), (x1, y1), (x0, y1)),
        source=source,
        confidence=confidence,
    )


def _text_pdf(*page_lines: tuple[float, float, str]) -> bytes:
    document = fitz.open()
    page = document.new_page(width=600, height=800)
    for x, y, text in page_lines:
        page.insert_text((x, y), text, fontsize=12)
    content = document.tobytes()
    document.close()
    return content


def _scanned_pdf(page_count: int = 1) -> bytes:
    document = fitz.open()
    for _ in range(page_count):
        page = document.new_page(width=600, height=800)
        image = Image.new("RGB", (600, 800), "white")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        page.insert_image(page.rect, stream=buffer.getvalue())
    content = document.tobytes()
    document.close()
    return content


@dataclass
class FakeOcr:
    lines_by_page: dict[int, tuple[OcrLine, ...]]
    calls: list[int]

    def recognize(self, png_bytes: bytes, *, page_number: int) -> tuple[OcrLine, ...]:
        assert png_bytes.startswith(b"\x89PNG")
        self.calls.append(page_number)
        return self.lines_by_page.get(page_number, ())


class ExplodingOcr:
    def recognize(self, png_bytes: bytes, *, page_number: int) -> tuple[OcrLine, ...]:
        raise AssertionError(f"OCR must not run for text PDF page {page_number}")


def _line(
    text: str,
    box: tuple[float, float, float, float],
    *,
    confidence: float = 0.45,
) -> OcrLine:
    x0, y0, x1, y1 = box
    return OcrLine(
        text=text,
        polygon=((x0, y0), (x1, y0), (x1, y1), (x0, y1)),
        confidence=confidence,
        engine_version="fake-ocr/1",
    )


def test_text_pdf_keeps_page_snapshot_and_never_calls_ocr(tmp_path: Path) -> None:
    content = _text_pdf((40, 80, "1. Solve $x^2+1$"), (40, 120, "show work"))
    pipeline = QuestionDocumentPipeline(tmp_path / "workspace", ocr_adapter=ExplodingOcr())

    snapshot = pipeline.prepare_source(
        PrepareSourceCommand(
            operation_id="text-pdf-1",
            source=DocumentSource(
                source_id="source-text-1",
                filename="synthetic.pdf",
                media_type="application/pdf",
                content=content,
            ),
        )
    )

    assert snapshot.pages[0].text_layer_state == TextLayerState.EMBEDDED
    assert snapshot.pages[0].rendered_asset.endswith("page-0001.png")
    assert snapshot.pages[0].rendered_width > snapshot.pages[0].width
    assert snapshot.pages[0].source_to_rendered[0] > 1.0
    assert pipeline.resolve_asset(snapshot.pages[0].rendered_asset).is_file()
    assert snapshot.questions[0].source_regions[0].coordinate_space == "normalized"
    assert snapshot.questions[0].math_expressions[0].omml.startswith("<m:oMath")

    pdf_path = tmp_path / "source.pdf"
    pdf_path.write_bytes(content)
    imported = import_pdf(
        pdf_path,
        document_pipeline=pipeline,
        operation_id="text-pdf-importer",
    )
    assert imported.needs_ocr is False
    assert imported.document_snapshot is not None
    assert len(imported.image_paths) == 1


def test_scan_double_column_cross_page_and_low_confidence_require_review(
    tmp_path: Path,
) -> None:
    ocr = FakeOcr(
        {
            1: (
                _line("2. right column", (0.60, 0.10, 0.94, 0.18)),
                _line("1. left start $x^2$", (0.05, 0.10, 0.42, 0.18)),
                _line("left detail", (0.05, 0.20, 0.42, 0.28)),
            ),
            2: (
                _line("unrelated left", (0.05, 0.10, 0.42, 0.18)),
                _line("continued on page two", (0.60, 0.10, 0.94, 0.18)),
            ),
        },
        [],
    )
    pipeline = QuestionDocumentPipeline(tmp_path / "workspace", ocr_adapter=ocr)
    snapshot = pipeline.prepare_source(
        PrepareSourceCommand(
            operation_id="scan-cross-page",
            source=DocumentSource(
                source_id="source-scan-1",
                filename="scan.pdf",
                media_type="application/pdf",
                content=_scanned_pdf(2),
            ),
            manual_questions=(
                ManualQuestionRegion(
                    question_id="manual-q1",
                    question_number="1",
                    question_regions=(
                        _region(1, 0.0, 0.0, 0.50, 0.40),
                        _region(2, 0.50, 0.0, 1.0, 0.40),
                    ),
                ),
            ),
        )
    )

    assert ocr.calls == [1, 2]
    assert [item.text for item in snapshot.pages[0].blocks] == [
        "1. left start $x^2$",
        "left detail",
        "2. right column",
    ]
    assert snapshot.questions[0].question_text == (
        "1. left start $x^2$\nleft detail\ncontinued on page two"
    )
    assert snapshot.state == DocumentState.REVIEW_REQUIRED
    assert "ocr_review_required" in snapshot.questions[0].blocking_issues
    assert snapshot.questions[0].math_expressions[0].review_state == ReviewState.CANDIDATE
    assert all(page.text_layer_state == TextLayerState.LOCAL_OCR for page in snapshot.pages)


def test_rotated_photo_uses_orientation_interface_and_manual_review(tmp_path: Path) -> None:
    image = Image.new("RGB", (120, 60), "white")
    exif = Image.Exif()
    exif[274] = 6
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", exif=exif)
    ocr = FakeOcr({1: (_line("1. rotated scan", (0.1, 0.1, 0.9, 0.2)),)}, [])
    pipeline = QuestionDocumentPipeline(tmp_path / "workspace", ocr_adapter=ocr)

    snapshot = pipeline.prepare_source(
        PrepareSourceCommand(
            operation_id="rotated-photo",
            source=DocumentSource(
                source_id="source-photo-1",
                filename="photo.jpg",
                media_type="image/jpeg",
                content=buffer.getvalue(),
            ),
        )
    )

    page = snapshot.pages[0]
    assert page.transform.rotation_degrees == 90
    assert page.transform.requires_review is True
    assert page.width == 60
    assert page.height == 120
    assert snapshot.state == DocumentState.REVIEW_REQUIRED


def test_reviewed_math_exports_editable_omml_and_reports_original_image_fallback(
    tmp_path: Path,
) -> None:
    ocr = FakeOcr({1: (_line("1. placeholder", (0.05, 0.05, 0.95, 0.25)),)}, [])
    workspace = tmp_path / "workspace"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ocr)
    prepared = pipeline.prepare_source(
        PrepareSourceCommand(
            operation_id="word-math",
            source=DocumentSource(
                source_id="source-math-1",
                filename="scan.pdf",
                media_type="application/pdf",
                content=_scanned_pdf(),
            ),
        )
    )
    reviewed = pipeline.apply_review(
        ApplyReviewCommand(
            operation_id=prepared.operation_id,
            expected_snapshot_revision=prepared.snapshot_revision,
            decisions=(
                QuestionReviewDecision(
                    question_id=prepared.questions[0].question_id,
                    question_text=r"Solve $x^2+1$ and preserve $\unknown{x}$.",
                    answer_status=ReviewState.MISSING,
                    confirm_content=True,
                ),
            ),
        )
    )
    assert reviewed.state == DocumentState.READY
    assert reviewed.questions[0].math_expressions[0].review_state == ReviewState.TEACHER_VERIFIED
    assert reviewed.questions[0].math_expressions[1].fallback_asset

    output = tmp_path / "math-export.docx"
    receipt = pipeline.export_word(
        ExportWordCommand(
            operation_id=reviewed.operation_id,
            expected_snapshot_revision=reviewed.snapshot_revision,
            output_path=str(output),
            question_ids=(reviewed.questions[0].question_id,),
        )
    )
    assert receipt.state == ExportState.COMPLETE_WITH_FALLBACKS
    assert len(receipt.fallbacks) == 1
    assert receipt.fallbacks[0].asset_path == reviewed.pages[0].rendered_asset
    assert validate_docx(output) == ()
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        assert b"<m:oMath" in document_xml
        assert any(name.startswith("word/media/") for name in archive.namelist())


def _review_single_question(
    pipeline: QuestionDocumentPipeline,
    *,
    operation_id: str,
) -> tuple:
    prepared = pipeline.prepare_source(
        PrepareSourceCommand(
            operation_id=operation_id,
            source=DocumentSource(
                source_id=f"source-{operation_id}",
                filename="paper.pdf",
                media_type="application/pdf",
                content=_text_pdf((40, 80, "1. Solve $x^2$")),
            ),
        )
    )
    reviewed = pipeline.apply_review(
        ApplyReviewCommand(
            operation_id=operation_id,
            expected_snapshot_revision=prepared.snapshot_revision,
            decisions=(
                QuestionReviewDecision(
                    question_id=prepared.questions[0].question_id,
                    confirm_content=True,
                    answer_status=ReviewState.MISSING,
                ),
            ),
        )
    )
    return prepared, reviewed


def test_publication_commits_only_after_sidecars_exist_and_is_idempotent(
    tmp_path: Path,
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    workspace = tmp_path / "pipeline"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(pipeline, operation_id="publication-ok")
    adapter = QuestionBankPublicationAdapter(
        db_path=db_path,
        data_root=tmp_path,
        pipeline_workspace=workspace,
    )
    pipeline.publisher = adapter
    command = PublishCommand(
        operation_id=reviewed.operation_id,
        expected_snapshot_revision=reviewed.snapshot_revision,
        question_ids=(reviewed.questions[0].question_id,),
    )

    first = pipeline.publish_questions(command)
    second = pipeline.publish_questions(command)
    assert first == second
    assert first.source_revision == reviewed.source_revision
    assert first.asset_paths
    assert len(first.questions) == 1
    bank_id = first.questions[0].bank_question_id
    sidecar = tmp_path / "question_bank" / "rich_content" / f"question_{bank_id}.json"
    assert sidecar.is_file()
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    assert payload["content_revision"] == reviewed.questions[0].content_revision
    with sqlite3.connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM papers").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0] == 1
        assert conn.execute(
            "SELECT state FROM question_document_publications WHERE operation_id = ?",
            (reviewed.operation_id,),
        ).fetchone()[0] == "published"
    sidecar.unlink()
    assert adapter.reconcile(reviewed.operation_id) == "recovery_required"
    with sqlite3.connect(db_path) as conn:
        assert conn.execute(
            "SELECT state FROM question_document_publications WHERE operation_id = ?",
            (reviewed.operation_id,),
        ).fetchone()[0] == "recovery_required"


def test_publication_reconcile_finishes_an_interrupted_staged_file_move(
    tmp_path: Path,
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    workspace = tmp_path / "pipeline"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(
        pipeline,
        operation_id="publication-recover-staged",
    )
    adapter = QuestionBankPublicationAdapter(
        db_path=db_path,
        data_root=tmp_path,
        pipeline_workspace=workspace,
    )
    pipeline.publisher = adapter
    receipt = pipeline.publish_questions(
        PublishCommand(
            operation_id=reviewed.operation_id,
            expected_snapshot_revision=reviewed.snapshot_revision,
            question_ids=(reviewed.questions[0].question_id,),
        )
    )
    manifest_path = adapter._manifest_path(reviewed.operation_id)
    manifest = adapter._read_json(manifest_path)
    sidecar_record = next(
        item for item in manifest["staged_files"] if item["kind"] == "sidecar"
    )
    final_sidecar = adapter._resolve_data_path(sidecar_record["final_path"])
    staged_sidecar = adapter._resolve_data_path(sidecar_record["staged_path"])
    staged_sidecar.parent.mkdir(parents=True, exist_ok=True)
    final_sidecar.replace(staged_sidecar)
    manifest["state"] = "publishing"
    adapter._write_validated_manifest(manifest_path, manifest)

    assert adapter.reconcile(reviewed.operation_id) == "published"
    assert final_sidecar.is_file()
    assert not adapter._staging_root(reviewed.operation_id).exists()
    assert pipeline.publish_questions(
        PublishCommand(
            operation_id=reviewed.operation_id,
            expected_snapshot_revision=reviewed.snapshot_revision,
            question_ids=(reviewed.questions[0].question_id,),
        )
    ) == receipt


def test_publication_rolls_back_database_when_sidecar_write_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    workspace = tmp_path / "pipeline"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(pipeline, operation_id="publication-fail")
    adapter = QuestionBankPublicationAdapter(
        db_path=db_path,
        data_root=tmp_path,
        pipeline_workspace=workspace,
    )
    pipeline.publisher = adapter

    def fail_sidecar(*args, **kwargs):
        raise OSError("synthetic sidecar failure")

    monkeypatch.setattr(adapter, "_write_owned_sidecar", fail_sidecar)
    with pytest.raises(OSError, match="synthetic sidecar failure"):
        pipeline.publish_questions(
            PublishCommand(
                operation_id=reviewed.operation_id,
                expected_snapshot_revision=reviewed.snapshot_revision,
                question_ids=(reviewed.questions[0].question_id,),
            )
        )
    with sqlite3.connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM papers").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM question_document_publications"
        ).fetchone()[0] == 0


def test_publication_copy_failure_leaves_no_orphaned_final_assets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    workspace = tmp_path / "pipeline"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(pipeline, operation_id="publication-copy-fail")
    adapter = QuestionBankPublicationAdapter(
        db_path=db_path,
        data_root=tmp_path,
        pipeline_workspace=workspace,
    )
    pipeline.publisher = adapter
    real_copy = adapter._copy_content_addressed
    copy_count = 0

    def fail_second_copy(source: Path, destination: Path, expected_hash: str) -> None:
        nonlocal copy_count
        copy_count += 1
        if copy_count == 2:
            raise OSError("synthetic second asset failure")
        real_copy(source, destination, expected_hash)

    monkeypatch.setattr(adapter, "_copy_content_addressed", fail_second_copy)
    with pytest.raises(OSError, match="synthetic second asset failure"):
        pipeline.publish_questions(
            PublishCommand(
                operation_id=reviewed.operation_id,
                expected_snapshot_revision=reviewed.snapshot_revision,
                question_ids=(reviewed.questions[0].question_id,),
            )
        )

    assert not list((tmp_path / "question_bank" / "raw_papers").glob("*"))
    assert not list((tmp_path / "question_bank" / "document_pages").glob("*"))
    with sqlite3.connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM papers").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0] == 0


def test_publication_manifest_failure_leaves_no_orphaned_final_assets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    workspace = tmp_path / "pipeline"
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(
        pipeline,
        operation_id="publication-manifest-fail",
    )
    adapter = QuestionBankPublicationAdapter(
        db_path=db_path,
        data_root=tmp_path,
        pipeline_workspace=workspace,
    )
    pipeline.publisher = adapter

    def fail_manifest(*_args, **_kwargs) -> None:
        raise OSError("synthetic manifest failure")

    monkeypatch.setattr(adapter, "_write_json", fail_manifest)
    with pytest.raises(OSError, match="synthetic manifest failure"):
        pipeline.publish_questions(
            PublishCommand(
                operation_id=reviewed.operation_id,
                expected_snapshot_revision=reviewed.snapshot_revision,
                question_ids=(reviewed.questions[0].question_id,),
            )
        )

    assert not list((tmp_path / "question_bank" / "raw_papers").glob("*"))
    assert not list((tmp_path / "question_bank" / "document_pages").glob("*"))
    with sqlite3.connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM papers").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0] == 0


def test_ordinary_and_personalized_exports_share_editable_omml_renderer(
    tmp_path: Path,
) -> None:
    db_path = tmp_path / "databases" / "question_bank.db"
    initialize_database(db_path)
    with sqlite3.connect(db_path) as conn:
        paper_id = conn.execute(
            "INSERT INTO papers (title, import_status) VALUES ('Synthetic', 'success')"
        ).lastrowid
        question_id = conn.execute(
            """
            INSERT INTO questions (
                paper_id, question_number, question_text, answer_text, difficulty
            ) VALUES (?, '1', 'Solve $x^2+1$ and $\\unknown{x}$.', 'Answer $x=1$.', '3')
            """,
            (paper_id,),
        ).lastrowid
        conn.commit()

    fallback_image = tmp_path / "question_bank" / "document_pages" / "fallback.png"
    fallback_image.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (320, 80), "white")
    image.save(fallback_image, format="PNG")
    fallback_hash = sha256_bytes(fallback_image.read_bytes())
    expressions = (
        build_math_expression(
            expression_id="legacy-valid-math",
            source="x^2+1",
        ),
        build_math_expression(
            expression_id="legacy-fallback-math",
            source=r"\unknown{x}",
            fallback_asset="question_bank/document_pages/fallback.png",
            fallback_sha256=fallback_hash,
        ),
    )
    rich_path = (
        tmp_path
        / "question_bank"
        / "rich_content"
        / f"question_{int(question_id)}.json"
    )
    rich_path.parent.mkdir(parents=True, exist_ok=True)
    rich_path.write_text(
        json.dumps(
            {
                "version": 3,
                "question_id": int(question_id),
                "content_revision": "a" * 64,
                "math_expressions": [asdict(item) for item in expressions],
                "question_blocks": [],
                "answer_blocks": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    ordinary = export_question_paper_docx(
        db_path,
        [int(question_id)],
        tmp_path / "ordinary",
        title="ordinary",
        ensure_previews=False,
    )
    personalized = export_training_docx(
        db_path,
        [{"question_id": int(question_id), "training_stage": "基础回补"}],
        tmp_path / "personalized",
        audience="student",
        student_id="synthetic-student",
    )
    for artifact in (ordinary, personalized):
        assert validate_docx(artifact) == ()
        assert export_receipt_path(artifact).is_file()
        receipt = load_legacy_export_receipt(artifact)
        assert receipt.state == ExportState.COMPLETE_WITH_FALLBACKS
        assert receipt.content_revisions == ("a" * 64,)
        assert len(receipt.fallbacks) == 1
        assert receipt.fallbacks[0].asset_path == (
            "question_bank/document_pages/fallback.png"
        )
        with zipfile.ZipFile(artifact) as archive:
            assert b"<m:oMath" in archive.read("word/document.xml")
            assert any(name.startswith("word/media/") for name in archive.namelist())


def test_compact_picture_cell_drops_source_indentation_before_fitting(tmp_path: Path) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    image = tmp_path / 'synthetic-wide-picture.png'
    Image.new('RGB', (310, 149), 'black').save(image)
    source = Document()
    text = source.add_paragraph('合成题：根据图中三角形求长度。')
    picture = source.add_paragraph()
    picture.paragraph_format.left_indent = Pt(60)
    picture._p.pPr.find(qn('w:ind')).set(qn('w:leftChars'), '130')
    picture.add_run().add_picture(str(image), width=Mm(80))
    relationship_id = picture._p.xpath('.//a:blip')[0].get(qn('r:embed'))
    result = Document()
    rendered = SharedWordQuestionRenderer().add_rich_blocks(result, [
        {'xml':text._p.xml}, {'xml':picture._p.xml, 'image_relationships':{relationship_id:str(image)}}
    ], compact_standalone_images_with_text=True)
    assert rendered.appended and rendered.compacted_image_count == 1
    cell = result.tables[0].cell(0,1)
    assert not cell._tc.xpath('.//w:ind')
    extent = cell._tc.xpath('.//wp:extent')[0]
    assert int(extent.get('cx')) <= int(cell.width) - 80*635
    assert float(extent.get('cx')) / float(extent.get('cy')) == pytest.approx(310/149, rel=.001)


def test_shared_renderer_handles_one_hundred_synthetic_exports(tmp_path: Path) -> None:
    renderer = SharedWordQuestionRenderer()
    for index in range(100):
        document = Document()
        assert renderer.add_text(
            document,
            f"Synthetic {index}: $x^{index % 5 + 1}+1$",
            question_id=f"batch-{index}",
        ) == ()
        output = tmp_path / f"batch-{index:03d}.docx"
        document.save(output)
        assert validate_docx(output) == ()


def test_prepare_operation_id_is_idempotent_but_rejects_changed_source(
    tmp_path: Path,
) -> None:
    pipeline = QuestionDocumentPipeline(tmp_path / "workspace", ocr_adapter=ExplodingOcr())
    command = PrepareSourceCommand(
        operation_id="stable-operation",
        source=DocumentSource(
            source_id="stable-source",
            filename="one.pdf",
            media_type="application/pdf",
            content=_text_pdf((40, 80, "1. one")),
        ),
    )
    assert pipeline.prepare_source(command) == pipeline.prepare_source(command)
    with pytest.raises(DocumentPipelineConflict):
        pipeline.prepare_source(
            PrepareSourceCommand(
                operation_id="stable-operation",
                source=DocumentSource(
                    source_id="stable-source",
                    filename="two.pdf",
                    media_type="application/pdf",
                    content=_text_pdf((40, 80, "1. two")),
                ),
            )
        )


def test_edit_after_review_invalidates_confirmation_and_old_revision(tmp_path: Path) -> None:
    pipeline = QuestionDocumentPipeline(tmp_path / "workspace", ocr_adapter=ExplodingOcr())
    _, reviewed = _review_single_question(pipeline, operation_id="review-invalidation")
    changed = pipeline.apply_review(
        ApplyReviewCommand(
            operation_id=reviewed.operation_id,
            expected_snapshot_revision=reviewed.snapshot_revision,
            decisions=(
                QuestionReviewDecision(
                    question_id=reviewed.questions[0].question_id,
                    question_text="1. changed but not confirmed",
                    confirm_content=False,
                ),
            ),
        )
    )
    assert changed.questions[0].review_state == ReviewState.CANDIDATE
    assert "teacher_content_confirmation_required" in changed.questions[0].blocking_issues
    with pytest.raises(DocumentPipelineConflict):
        pipeline.apply_review(
            ApplyReviewCommand(
                operation_id=reviewed.operation_id,
                expected_snapshot_revision=reviewed.snapshot_revision,
                decisions=(
                    QuestionReviewDecision(
                        question_id=reviewed.questions[0].question_id,
                        confirm_content=True,
                    ),
                ),
            )
        )
