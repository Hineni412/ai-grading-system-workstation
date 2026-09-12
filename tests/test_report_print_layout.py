from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from db_manager import DBManager
from report import ReportGenerator


RUBRIC = {
    "questions": [
        {
            "question_id": "Q1",
            "question_type": "calculation",
            "max_score": 10,
            "knowledge_points": [
                {"knowledge_id": "K1", "knowledge_name": "有理数运算"}
            ],
            "parts": [],
        },
        {
            "question_id": "Q2",
            "question_type": "choice",
            "max_score": 10,
            "knowledge_points": [
                {"knowledge_id": "K2", "knowledge_name": "代数式"}
            ],
            "parts": [],
        },
    ]
}


def _complete_payload() -> str:
    return json.dumps(
        {
            "grading_completeness": {
                "status": "complete",
                "missing_question_ids": [],
                "duplicate_question_ids": [],
                "unexpected_question_ids": [],
                "score_out_of_range": [],
                "affected_major_question_ids": [],
            }
        },
        ensure_ascii=False,
    )


def _seed_print_report(tmp_path: Path) -> Path:
    databases_dir = tmp_path / "databases"
    databases_dir.mkdir(exist_ok=True)
    db_path = databases_dir / "grading.db"
    rubric_path = tmp_path / "rubric.json"
    rubric_path.write_text(json.dumps(RUBRIC, ensure_ascii=False), encoding="utf-8")
    DBManager(db_path).initialize()
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO grading_sessions (
                id, session_name, rubric_path, answer_key_path
            ) VALUES (1, '七年级期末测试', ?, '')
            """,
            (str(rubric_path),),
        )
        for student_id in range(1, 12):
            conn.execute(
                """
                INSERT INTO students (id, student_code, name, class_name)
                VALUES (?, ?, ?, '一班')
                """,
                (
                    student_id,
                    f"S{student_id:03d}",
                    f"学生{student_id:02d}",
                ),
            )
            is_complete = student_id <= 10
            conn.execute(
                """
                INSERT INTO exam_papers (
                    id, session_id, student_id, front_image, back_image,
                    match_status, processing_status
                ) VALUES (?, 1, ?, 'front.png', 'back.png', 'matched', 'graded')
                """,
                (student_id, student_id),
            )
            total_score = float(22 - student_id) if is_complete else 7.0
            raw_json = _complete_payload() if is_complete else "{}"
            conn.execute(
                """
                INSERT INTO session_results (
                    id, session_id, student_id, paper_id, total_score,
                    student_score, needs_human_review, raw_json, graded_at
                ) VALUES (?, 1, ?, ?, 20, ?, 0, ?, '2026-07-28 08:00:00')
                """,
                (student_id, student_id, student_id, total_score, raw_json),
            )
            q1_score = 10.0 if student_id == 1 else max(0.0, 10.0 - student_id)
            conn.execute(
                """
                INSERT INTO session_details (
                    result_id, question_id, score_awarded, deduction_reason,
                    knowledge_ids, error_category, error_summary
                ) VALUES (?, 'Q1', ?, '计算过程有误', '["K1"]', '计算错误', '符号错误')
                """,
                (student_id, q1_score),
            )
            if is_complete:
                conn.execute(
                    """
                    INSERT INTO session_details (
                        result_id, question_id, score_awarded, deduction_reason,
                        knowledge_ids, error_category, error_summary
                    ) VALUES (?, 'Q2', 10, '', '["K2"]', '', '')
                    """,
                    (student_id,),
                )
        conn.commit()
    return db_path


def _table_rows(sheet, header_row: int) -> list[dict[str, object]]:
    headers = [
        str(sheet.cell(header_row, column).value or "")
        for column in range(1, sheet.max_column + 1)
    ]
    return [
        {
            headers[column - 1]: sheet.cell(row, column).value
            for column in range(1, sheet.max_column + 1)
        }
        for row in range(header_row + 1, sheet.max_row + 1)
        if any(
            sheet.cell(row, column).value is not None
            for column in range(1, sheet.max_column + 1)
        )
    ]


def test_score_excel_separates_official_statistics_from_print_name_hiding(
    tmp_path: Path,
) -> None:
    db_path = _seed_print_report(tmp_path)

    export_path = ReportGenerator(db_path, tmp_path / "reports").export_session(
        1,
        score_excel_options={
            "hide_bottom_enabled": True,
            "hide_bottom_n": 8,
            "manual_hidden_student_ids": [2],
        },
    )

    workbook = load_workbook(export_path, data_only=False)
    assert workbook.sheetnames == [
        "考试总览",
        "班级成绩总表",
        "小题分析打印",
        "成绩与小题明细",
        "AI与人工分对比",
        "错因明细",
        "知识点分析",
        "缺考与异常",
    ]

    overview = workbook["考试总览"]
    assert overview["B3"].value == 10
    assert overview["D3"].value == 1
    assert overview["F3"].value == 16.5

    analysis_rows = _table_rows(workbook["小题分析打印"], 5)
    q1 = next(row for row in analysis_rows if row["题号"] == "Q1")
    assert q1["统计人数"] == 10
    assert q1["失分人数"] == 9
    assert q1["隐藏姓名数"] == 9
    assert q1["显示姓名数"] == 0
    assert q1["失分同学"] == "姓名已隐藏"

    detail_headers = [
        workbook["成绩与小题明细"].cell(1, column).value
        for column in range(1, workbook["成绩与小题明细"].max_column + 1)
    ]
    assert "班级均分" not in detail_headers
    assert "批改完整性" in detail_headers
    assert workbook["班级成绩总表"]["A3"].value == "班级均分"

    exception_values = [
        value
        for row in workbook["缺考与异常"].iter_rows(values_only=True)
        for value in row
    ]
    assert "学生11" in exception_values
    assert "批改不完整" in exception_values


def test_score_excel_uses_full_cell_borders_and_requested_sorting(
    tmp_path: Path,
) -> None:
    db_path = _seed_print_report(tmp_path)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            UPDATE session_results
            SET student_score = CASE
                WHEN id = 1 THEN 12
                WHEN id = 10 THEN 21
                ELSE student_score
            END
            WHERE session_id = 1
            """
        )
        conn.execute(
            """
            UPDATE session_details
            SET score_awarded = 10
            WHERE question_id = 'Q1' AND result_id <= 10
            """
        )
        conn.execute(
            """
            UPDATE session_details
            SET score_awarded = 0
            WHERE question_id = 'Q2' AND result_id <= 10
            """
        )
        conn.commit()

    export_path = ReportGenerator(
        db_path,
        tmp_path / "reports",
    ).export_session(1)
    workbook = load_workbook(export_path, data_only=False)

    score_sheet = workbook["班级成绩总表"]
    score_rows = _table_rows(score_sheet, 4)
    score_values = [float(row["学生得分"]) for row in score_rows]
    assert score_values == sorted(score_values, reverse=True)

    analysis_rows = _table_rows(workbook["小题分析打印"], 5)
    assert [row["题号"] for row in analysis_rows] == ["Q2", "Q1"]
    assert [
        float(row["全班得分率"])
        for row in analysis_rows
    ] == sorted(float(row["全班得分率"]) for row in analysis_rows)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(
            min_row=1,
            max_row=sheet.max_row,
            min_col=1,
            max_col=sheet.max_column,
        ):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                for side in ("left", "right", "top", "bottom"):
                    assert getattr(cell.border, side).style == "thin", (
                        sheet.title, cell.coordinate, side,
                    )
        for merged in sheet.merged_cells.ranges:
            for column in range(merged.min_col, merged.max_col + 1):
                assert sheet.cell(merged.min_row, column).border.top.style == "thin"
                assert sheet.cell(merged.max_row, column).border.bottom.style == "thin"
            for row_index in range(merged.min_row, merged.max_row + 1):
                assert sheet.cell(row_index, merged.min_col).border.left.style == "thin"
                assert sheet.cell(row_index, merged.max_col).border.right.style == "thin"
