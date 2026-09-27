from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches
from PIL import Image

from backend.config_workspace import formula_preview


def _omath(inner: str) -> str:
    return f'<m:oMath {nsdecls("m")}>{inner}</m:oMath>'


_SQRT_37 = _omath(
    "<m:rad><m:deg/><m:e><m:r><m:t>37</m:t></m:r></m:e></m:rad>"
)
_FRACTION_1_5 = _omath(
    "<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
    "<m:den><m:r><m:t>5</m:t></m:r></m:den></m:f>"
)


def _docx_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_formula_questions() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("2. 计算 ")
    paragraph._p.append(parse_xml(_SQRT_37))  # noqa: SLF001 - fixture needs exact Word OMML.
    paragraph.add_run(" 的值。")
    fraction = document.add_paragraph()
    fraction.add_run("3. 化简 ")
    fraction._p.append(parse_xml(_FRACTION_1_5))  # noqa: SLF001
    fraction.add_run(" 的结果。")
    document.add_paragraph("4. 没有公式的普通题目。")
    return _docx_bytes(document)


def test_paragraph_xml_renders_formula_spans() -> None:
    docx = _docx_with_formula_questions()
    index = formula_preview.load_paragraph_xml_index(docx)

    rendered = {
        text: formula_preview.preview_html_for_text(index, text)
        for text in [
            "2. 计算 √(37) 的值。",
            "3. 化简 (1)/(5) 的结果。",
        ]
    }

    assert 'class="qm"' in rendered["2. 计算 √(37) 的值。"]
    assert "\\sqrt{37}" in rendered["2. 计算 √(37) 的值。"]
    assert 'data-latex' in rendered["2. 计算 √(37) 的值。"]
    assert "\\frac{1}{5}" in rendered["3. 化简 (1)/(5) 的结果。"]


def test_stripped_question_number_matches_and_removes_prefix() -> None:
    docx = _docx_with_formula_questions()
    index = formula_preview.load_paragraph_xml_index(docx)

    rendered = formula_preview.preview_html_for_text(index, "计算 √(37) 的值。")

    assert rendered
    assert "\\sqrt{37}" in rendered
    plain = re.sub(r"<[^>]+>", "", rendered)
    assert not plain.lstrip().startswith("2")


def test_mismatched_text_falls_back_to_no_html() -> None:
    docx = _docx_with_formula_questions()
    index = formula_preview.load_paragraph_xml_index(docx)

    assert formula_preview.preview_html_for_text(index, "完全不同的题目文本") == ""
    assert formula_preview.preview_html_for_text({}, "计算 √(37) 的值。") == ""
    assert formula_preview.preview_html_for_text(index, "") == ""


def test_xml_index_is_cached_per_source_revision() -> None:
    docx = _docx_with_formula_questions()
    calls: list[bytes] = []
    original = formula_preview.load_paragraph_xml_index

    def counting(payload: bytes) -> dict[str, str]:
        calls.append(payload)
        return original(payload)

    formula_preview._INDEX_CACHE.clear()  # noqa: SLF001 - isolates cache state for the assertion.
    formula_preview.load_paragraph_xml_index = counting
    try:
        first = formula_preview.paragraph_xml_index(7, "a" * 32, "b" * 64, docx)
        second = formula_preview.paragraph_xml_index(7, "a" * 32, "b" * 64, docx)
        other_revision = formula_preview.paragraph_xml_index(
            7, "a" * 32, "c" * 64, docx
        )
    finally:
        formula_preview.load_paragraph_xml_index = original

    assert first is second
    assert len(calls) == 2
    assert other_revision == first


def test_drawing_renders_nothing_without_relationship_urls() -> None:
    image = io.BytesIO()
    Image.new("RGB", (24, 18), "navy").save(image, format="PNG")
    document = Document()
    paragraph = document.add_paragraph("5. 如图，求面积。")
    paragraph.add_run().add_picture(
        io.BytesIO(image.getvalue()), width=Inches(0.25)
    )
    docx = _docx_bytes(document)

    index = formula_preview.load_paragraph_xml_index(docx)
    rendered = formula_preview.preview_html_for_text(
        index, "如图，求面积。"
    )

    assert rendered
    assert "<img" not in rendered
