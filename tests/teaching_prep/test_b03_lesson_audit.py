from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from backend.teaching_prep.application.lesson_audit import (
    audit_page_budget,
    audit_superscript_subscript,
    cross_check_question_pages,
    run_full_audit,
)

_RPR = "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"


def _deck_with_baseline_run(
    path: Path,
    *,
    baseline: str | None,
    text: str,
    slides: int = 1,
) -> Path:
    """构造带可疑上下标 run 的课件：baseline 30000 上标 / -25000 下标。"""
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for index in range(slides):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(9), Inches(1)
        )
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"第{index + 1}页正常文字"
        run.font.size = Pt(16)
        if baseline is not None and index == 0:
            bad = box.text_frame.paragraphs[0].add_run()
            bad.text = text
            bad.font.size = Pt(16)
            bad.font._rPr.set("baseline", baseline)
    presentation.save(str(path))
    return path


def test_long_superscript_sentence_is_flagged(tmp_path: Path) -> None:
    deck = _deck_with_baseline_run(
        tmp_path / "bad.pptx",
        baseline="30000",
        text="这一整句都被误标成了上标文字",
    )

    result = audit_superscript_subscript(deck)

    assert result["passed"] is False
    assert result["finding_count"] == 1
    finding = result["findings"][0]
    assert finding["page"] == 1
    assert finding["baseline"] == 30000
    assert finding["kind"] == "superscript"
    assert "误标成了上标" in finding["text"]


def test_long_subscript_run_is_flagged_and_short_marker_is_not(
    tmp_path: Path,
) -> None:
    bad = _deck_with_baseline_run(
        tmp_path / "sub.pptx",
        baseline="-25000",
        text="a long english subscript sentence",
    )
    result = audit_superscript_subscript(bad)
    assert result["passed"] is False
    assert result["findings"][0]["kind"] == "subscript"

    legit = _deck_with_baseline_run(
        tmp_path / "legit.pptx",
        baseline="30000",
        text="2",
    )
    result = audit_superscript_subscript(legit)
    assert result["passed"] is True
    assert result["finding_count"] == 0
    assert result["scanned_runs"] > 0

    plain = _deck_with_baseline_run(
        tmp_path / "plain.pptx",
        baseline=None,
        text="",
    )
    assert audit_superscript_subscript(plain)["passed"] is True


def test_page_budget_red_line() -> None:
    within = audit_page_budget(18)
    assert within["over_limit"] is False
    assert within["limit"] == 18

    over = audit_page_budget(19)
    assert over["over_limit"] is True
    assert "19 页" in over["suggestion"]


def test_question_page_cross_check_detects_out_of_range_and_conflict() -> None:
    consistent = cross_check_question_pages(
        [
            {"question_id": 1, "final_position": 2},
            {"question_id": 2, "final_position": 4},
        ],
        final_page_count=5,
    )
    assert consistent["passed"] is True
    assert consistent["pages"] == {2: 1, 4: 2}

    out_of_range = cross_check_question_pages(
        [{"question_id": 1, "final_position": 9}],
        final_page_count=5,
    )
    assert out_of_range["passed"] is False
    assert out_of_range["problem_count"] == 1
    assert "超出成片范围" in out_of_range["problems"][0]

    conflict = cross_check_question_pages(
        [
            {"question_id": 1, "final_position": 2},
            {"question_id": 2, "final_position": 2},
        ],
        final_page_count=5,
    )
    assert conflict["passed"] is False
    assert "同一页" in conflict["problems"][0]


def test_run_full_audit_combines_all_three_checks(tmp_path: Path) -> None:
    clean = _deck_with_baseline_run(
        tmp_path / "clean.pptx",
        baseline=None,
        text="",
        slides=3,
    )
    report = run_full_audit(
        clean,
        inserted_question_pages=[{"question_id": 7, "final_position": 2}],
    )
    assert report["passed"] is True
    assert report["page_budget"]["final_page_count"] == 3
    assert report["question_pages"]["checked"] == 1
    assert report["superscript_subscript"]["passed"] is True

    dirty = _deck_with_baseline_run(
        tmp_path / "dirty.pptx",
        baseline="30000",
        text="整句上标残留需要教师复核",
        slides=3,
    )
    report = run_full_audit(
        dirty,
        inserted_question_pages=[{"question_id": 7, "final_position": 99}],
    )
    assert report["passed"] is False
    assert report["superscript_subscript"]["finding_count"] == 1
    assert report["question_pages"]["problem_count"] == 1
