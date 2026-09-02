from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from backend.teaching_prep.application.pptx_editor import (
    PptxEditorError,
    apply_plan,
)


def _deck(path: Path, titles: list[str] | None = None) -> Path:
    """生成小样课件：每页一个带标题文本框，便于断言未触及页内容不变。"""
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for index, title in enumerate(titles or ["第一页", "第二页", "第三页"], 1):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(9), Inches(1)
        )
        box.text_frame.text = title or f"第{index}页"
    presentation.save(str(path))
    return path


def _slide_texts(path: Path) -> list[str]:
    presentation = Presentation(str(path))
    texts = []
    for slide in presentation.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        texts.append("\n".join(parts))
    return texts


def _part_names(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _operation(
    kind: str,
    page: int | None,
    *,
    details: dict[str, object] | None = None,
    reason: str = "合成操作",
) -> dict[str, object]:
    return {
        "kind": kind,
        "target": (
            {"generated_page_number": page} if page is not None else {}
        ),
        "target_key": f"synthetic:{kind}:{page}",
        "reason": reason,
        "details": details or {},
    }


def test_delete_reorder_hide_and_question_insert_on_isolated_copy(
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "out" / "candidate.pptx"
    source_before = _sha256(source)
    operations = [
        _operation("delete_slide", 2),
        _operation("reorder_slide", 3, details={"target_position": 1}),
        _operation("hide_slide", 3),
        _operation(
            "add_question_slide",
            None,
            details={
                "anchor_page": 1,
                "question": {
                    "question_id": 101,
                    "title": "课堂练习",
                    "stem_html": "求解 x<sup>2</sup> = 4",
                    "answer_html": "x = ±2",
                },
            },
        ),
    ]

    report = apply_plan(
        source_path=source,
        output_path=output,
        operations=operations,
    )

    assert [entry["status"] for entry in report["operations"]] == [
        "applied",
        "applied",
        "applied",
        "applied",
    ]
    assert report["source_page_count"] == 3
    assert report["final_page_count"] == 3
    assert report["final_page_order"] == [3, 1, "new:1"]
    assert report["inserted_question_pages"] == [
        {"question_id": 101, "final_position": 3}
    ]
    assert report["output_sha256"] == _sha256(output)
    assert _sha256(source) == source_before

    texts = _slide_texts(output)
    assert "第三页" in texts[0]
    assert "第一页" in texts[1]
    assert "求解 x2 = 4" in texts[2].replace("\n", " ") or "求解" in texts[2]

    # 隐藏标记落在重排后的第一页（源第 3 页）。
    presentation = Presentation(str(output))
    assert presentation.slides[0]._element.get("show") == "0"
    assert presentation.slides[1]._element.get("show") is None

    # 部件无重名：删页 + 插页后包内路径必须唯一。
    names = _part_names(output)
    assert len(names) == len(set(names))

    # 题页上标保留：x 的平方 run 带 baseline=30000。
    question_slide = presentation.slides[2]
    baselines = []
    for shape in question_slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                r_pr = run._r.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"
                )
                baselines.append(
                    (run.text, r_pr.get("baseline") if r_pr is not None else None)
                )
    assert ("2", "30000") in baselines


def test_plain_add_slide_uses_reason_text(tmp_path: Path) -> None:
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "candidate.pptx"

    report = apply_plan(
        source_path=source,
        output_path=output,
        operations=[
            _operation(
                "add_slide",
                1,
                reason="在本页之后新增一页小结。",
            )
        ],
    )

    assert report["final_page_count"] == 4
    assert report["final_page_order"] == [1, "new:1", 2, 3]
    assert "在本页之后新增一页小结。" in _slide_texts(output)[1]


def test_unsupported_invalid_and_repeated_operations_are_reported(
    tmp_path: Path,
) -> None:
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "candidate.pptx"

    report = apply_plan(
        source_path=source,
        output_path=output,
        operations=[
            _operation("modify_text_box", 1),
            _operation("unknown_kind", 1),
            _operation("delete_slide", 99),
            _operation("delete_slide", 2),
            _operation("delete_slide", 2),
        ],
    )

    statuses = {
        (entry["kind"], entry["source_page"]): entry["status"]
        for entry in report["operations"]
    }
    assert statuses[("modify_text_box", 1)] == "manual"
    assert statuses[("unknown_kind", 1)] == "skipped"
    assert statuses[("delete_slide", 99)] == "failed"
    assert statuses[("delete_slide", 2)] == "skipped"  # 第二次删除被跳过
    assert report["operations"][3]["status"] == "applied"
    assert report["final_page_count"] == 2


def test_missing_source_fails_closed_and_empty_plan_is_a_copy(
    tmp_path: Path,
) -> None:
    with pytest.raises(PptxEditorError):
        apply_plan(
            source_path=tmp_path / "missing.pptx",
            output_path=tmp_path / "candidate.pptx",
            operations=[],
        )

    source = _deck(tmp_path / "source.pptx")
    # 空操作单合法：不改页序，产出与源内容一致。
    report = apply_plan(
        source_path=source,
        output_path=tmp_path / "copy.pptx",
        operations=[],
    )
    assert report["final_page_count"] == 3
    assert _slide_texts(tmp_path / "copy.pptx") == _slide_texts(source)
