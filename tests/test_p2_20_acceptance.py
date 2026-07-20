from __future__ import annotations

import hashlib
import json
import socket
import sqlite3
import tempfile
import subprocess
import textwrap
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from db_manager import DBManager
from tools import p2_20_acceptance as acceptance


def _git(repo: Path, *args: str) -> str:
    completed = subprocess.run(
        ["git", *args],
        cwd=repo,
        capture_output=True,
        text=True,
        check=True,
    )
    return completed.stdout.strip()


def _committed_source_repo(tmp_path: Path) -> tuple[Path, str]:
    repo = tmp_path / "source"
    repo.mkdir()
    _git(repo, "init", "-q")
    _git(repo, "config", "user.email", "p2-20@example.invalid")
    _git(repo, "config", "user.name", "P2-20 Test")
    (repo / "safe.txt").write_text("safe", encoding="utf-8")
    config_dir = repo / "config"
    config_dir.mkdir()
    (config_dir / "app_config.yaml").write_text(
        "DATA_DIR: user_data\nLOGS_DIR: logs\nVERSION: test\n",
        encoding="utf-8",
    )
    protected = repo / "user_data"
    protected.mkdir()
    (protected / "secret.txt").write_text("must not stage", encoding="utf-8")
    _git(repo, "add", ".")
    _git(repo, "commit", "-qm", "fixture")
    return repo, _git(repo, "rev-parse", "HEAD")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _authorized_source_data(tmp_path: Path, *, paper_count: int = 3) -> Path:
    data_root = tmp_path / "source-user-data"
    database_dir = data_root / "databases"
    exam_dir = data_root / "exams" / "session_1"
    template_dir = data_root / "templates" / "session_1"
    database_dir.mkdir(parents=True)
    exam_dir.mkdir(parents=True)
    template_dir.mkdir(parents=True)
    source_paper = exam_dir / "source.docx"
    source_paper.write_bytes(b"authorized exam source")
    (template_dir / "template_source_full_class.pdf").write_bytes(
        b"%PDF-1.4 authorized template"
    )
    database = database_dir / "grading_system.db"
    connection = sqlite3.connect(database)
    connection.executescript(
        """
        CREATE TABLE grading_sessions (
            id INTEGER PRIMARY KEY,
            session_name TEXT NOT NULL,
            source_paper_path TEXT,
            is_deleted INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE students (
            id INTEGER PRIMARY KEY,
            student_code TEXT NOT NULL,
            name TEXT NOT NULL,
            class_name TEXT
        );
        CREATE TABLE exam_papers (
            id INTEGER PRIMARY KEY,
            session_id INTEGER NOT NULL,
            front_image TEXT,
            back_image TEXT,
            student_id INTEGER,
            match_status TEXT,
            processing_status TEXT
        );
        CREATE TABLE session_results (
            id INTEGER PRIMARY KEY,
            session_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            paper_id INTEGER NOT NULL
        );
        CREATE TABLE session_templates (
            id INTEGER PRIMARY KEY,
            session_id INTEGER NOT NULL,
            is_confirmed INTEGER NOT NULL
        );
        """
    )
    connection.execute(
        "INSERT INTO grading_sessions(id, session_name, source_paper_path) VALUES(1, ?, ?)",
        ("0609", str(source_paper)),
    )
    connection.execute(
        "INSERT INTO session_templates(id, session_id, is_confirmed) VALUES(1, 1, 1)"
    )
    for index in range(1, paper_count + 1):
        front = exam_dir / f"student-{index}-front.jpg"
        back = exam_dir / f"student-{index}-back.jpg"
        front.write_bytes(f"front-{index}".encode())
        back.write_bytes(f"back-{index}".encode())
        connection.execute(
            "INSERT INTO students(id, student_code, name, class_name) VALUES(?, ?, ?, ?)",
            (index, f"S{index:03}", f"学生{index}", "七年级一班"),
        )
        connection.execute(
            """
            INSERT INTO exam_papers(
                id, session_id, front_image, back_image, student_id,
                match_status, processing_status
            ) VALUES(?, 1, ?, ?, ?, 'matched', 'graded')
            """,
            (100 + index, str(front), str(back), index),
        )
        connection.execute(
            """
            INSERT INTO session_results(id, session_id, student_id, paper_id)
            VALUES(?, 1, ?, ?)
            """,
            (200 + index, index, 100 + index),
        )
    connection.commit()
    connection.close()
    return data_root


def _prepared_report_workspace(
    tmp_path: Path,
) -> tuple[Path, Path, DBManager, int]:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "machine-config" / "api_profiles.json"
    profile_path.parent.mkdir()
    profile_path.write_text(
        json.dumps(
            [
                {
                    "provider": "custom-openai-compatible",
                    "base_url": "https://grading.example.invalid/v1",
                    "api_key": "grading-secret",
                    "grading_model": "grading-model",
                    "config_model": "config-model",
                }
            ]
        ),
        encoding="utf-8",
    )
    acceptance.prepare_runtime_profile(
        workspace,
        source_profile_path=profile_path,
        proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
        max_forwarded_requests=4,
    )

    data_root = workspace / "acceptance_data"
    config_dir = data_root / "config" / "uploaded"
    config_dir.mkdir(parents=True)
    rubric = config_dir / "rubric.json"
    answer = config_dir / "answer.json"
    rubric.write_text(
        json.dumps(
            {
                "total_score": 100,
                "questions": [
                    {"question_id": "Q1", "max_score": 40},
                    {"question_id": "Q2", "max_score": 60},
                ],
            }
        ),
        encoding="utf-8",
    )
    answer.write_text(
        json.dumps(
            {
                "questions": [
                    {"question_id": "Q1", "canonical_answer": "A"},
                    {"question_id": "Q2", "canonical_answer": "B"},
                ]
            }
        ),
        encoding="utf-8",
    )
    (data_root / "databases").mkdir()
    db = DBManager(data_root / "databases" / "grading_system.db")
    db.initialize()
    session_id = db.create_grading_session(
        "P2-20 synthetic acceptance",
        str(rubric),
        str(answer),
    )
    return workspace, source_data, db, session_id


def _seed_confirmed_template(
    workspace: Path,
    db: DBManager,
    session_id: int,
    *,
    regions: list[dict[str, object]] | None = None,
) -> int:
    template_dir = (
        workspace
        / "acceptance_data"
        / "templates"
        / f"session_{session_id}"
    )
    template_dir.mkdir(parents=True)
    paths = {
        "front": template_dir / "front.png",
        "back": template_dir / "back.png",
        "analysis": template_dir / "analysis.json",
        "config": template_dir / "template.json",
        "regions": template_dir / "regions.json",
    }
    paths["front"].write_bytes(b"synthetic-front")
    paths["back"].write_bytes(b"synthetic-back")
    for key in ("analysis", "config", "regions"):
        paths[key].write_text("{}", encoding="utf-8")
    template_id = db.activate_session_template(
        session_id,
        front_template_path=str(paths["front"]),
        back_template_path=str(paths["back"]),
        ai_analysis_path=str(paths["analysis"]),
        template_config_path=str(paths["config"]),
        regions_path=str(paths["regions"]),
    )
    token = db.replace_answer_regions_atomic(
        session_id,
        template_id,
        regions
        or [
            {
                "region_uuid": "region-q1",
                "page": "front",
                "region_order": 1,
                "x": 1,
                "y": 1,
                "w": 10,
                "h": 10,
                "mapped_question_id": "Q1",
                "mapping_status": "manual",
            },
            {
                "region_uuid": "region-q2",
                "page": "back",
                "region_order": 2,
                "x": 2,
                "y": 2,
                "w": 10,
                "h": 10,
                "mapped_question_id": "Q2",
                "mapping_status": "manual",
            },
        ],
        confirmed=True,
    )
    assert db.mark_region_snapshot_complete(
        session_id,
        expected_token=token,
    )
    return template_id


def _seed_complete_grading(
    workspace: Path,
    db: DBManager,
    session_id: int,
) -> None:
    scans_dir = workspace / "acceptance_data" / "scans"
    scans_dir.mkdir(parents=True)
    database = workspace / "acceptance_data" / "databases" / "grading_system.db"
    with sqlite3.connect(database) as connection:
        for index in range(1, 4):
            front = scans_dir / f"paper-{index}-front.jpg"
            back = scans_dir / f"paper-{index}-back.jpg"
            front.write_bytes(f"front-{index}".encode())
            back.write_bytes(f"back-{index}".encode())
            student_id = int(
                connection.execute(
                    """
                    INSERT INTO students (student_code, name, class_name)
                    VALUES (?, ?, ?)
                    """,
                    (f"S{index:03}", f"Student {index}", "Class A"),
                ).lastrowid
            )
            paper_id = int(
                connection.execute(
                    """
                    INSERT INTO exam_papers (
                        session_id, front_image, back_image, ocr_name,
                        student_id, match_status, processing_status
                    ) VALUES (?, ?, ?, ?, ?, 'matched', 'graded')
                    """,
                    (
                        session_id,
                        str(front),
                        str(back),
                        f"Student {index}",
                        student_id,
                    ),
                ).lastrowid
            )
            result_id = int(
                connection.execute(
                    """
                    INSERT INTO session_results (
                        session_id, student_id, paper_id, total_score,
                        student_score, needs_human_review, raw_json
                    ) VALUES (?, ?, ?, 100, 70, 0, ?)
                    """,
                    (
                        session_id,
                        student_id,
                        paper_id,
                        json.dumps(
                            {
                                "grading_completeness": {
                                    "status": "complete",
                                    "missing_question_ids": [],
                                    "duplicate_question_ids": [],
                                    "unexpected_question_ids": [],
                                }
                            }
                        ),
                    ),
                ).lastrowid
            )
            connection.executemany(
                """
                INSERT INTO session_details (
                    result_id, question_id, score_awarded,
                    deduction_reason, knowledge_id, confidence_score
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                [
                    (result_id, "Q1", 30, "", "K1", 95),
                    (result_id, "Q2", 40, "", "K2", 95),
                ],
            )
        connection.commit()
    db.update_session_status(session_id, "completed")


def _write_matching_report(workspace: Path) -> Path:
    reports_dir = workspace / "acceptance_data" / "reports"
    reports_dir.mkdir(parents=True)
    report_path = reports_dir / "session-report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩与小题明细"
    sheet.append(["学号", "总分"])
    for index in range(1, 4):
        sheet.append([f"S{index:03}", 70])
    workbook.save(report_path)
    workbook.close()
    return report_path


def test_consistency_report_accepts_safe_created_session_stage(
    tmp_path: Path,
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)

    report = acceptance.build_consistency_report(
        workspace,
        stage="session",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report == {
        "package": "P2-20",
        "stage": "session",
        "ok": True,
        "source_unchanged": True,
        "session_id": session_id,
        "session_status": "created",
        "counts": {
            "papers": 0,
            "results": 0,
            "details": 0,
        },
    }
    evidence = json.loads(
        (
            workspace
            / "acceptance_evidence"
            / "consistency-session.json"
        ).read_text(encoding="utf-8")
    )
    assert evidence == report
    encoded = json.dumps(report)
    assert str(workspace) not in encoded
    assert "P2-20 synthetic acceptance" not in encoded


def test_consistency_report_accepts_matching_config_without_content_or_paths(
    tmp_path: Path,
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)

    report = acceptance.build_consistency_report(
        workspace,
        stage="config",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report["stage"] == "config"
    assert report["ok"] is True
    assert report["source_unchanged"] is True
    assert report["question_count"] == 2
    assert len(str(report["config_sha256"])) == 64
    assert report["counts"] == {"papers": 0, "results": 0, "details": 0}
    encoded = json.dumps(report)
    assert str(workspace) not in encoded
    assert "canonical_answer" not in encoded
    assert '"A"' not in encoded
    assert (
        json.loads(
            (
                workspace
                / "acceptance_evidence"
                / "consistency-config.json"
            ).read_text(encoding="utf-8")
        )
        == report
    )


def test_consistency_report_rejects_conflicting_config_question_sets(
    tmp_path: Path,
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)
    answer_path = (
        workspace
        / "acceptance_data"
        / "config"
        / "uploaded"
        / "answer.json"
    )
    answer_path.write_text(
        json.dumps(
            {
                "questions": [
                    {"question_id": "Q1", "canonical_answer": "A"},
                    {"question_id": "Q3", "canonical_answer": "C"},
                ]
            }
        ),
        encoding="utf-8",
    )

    with pytest.raises(
        acceptance.AcceptanceError,
        match="configuration question identities conflict",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="config",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_rejects_config_path_outside_acceptance_data(
    tmp_path: Path,
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)
    outside_rubric = tmp_path / "outside-rubric.json"
    outside_rubric.write_text(
        json.dumps(
            {
                "questions": [
                    {"question_id": "Q1", "max_score": 40},
                    {"question_id": "Q2", "max_score": 60},
                ]
            }
        ),
        encoding="utf-8",
    )
    database = workspace / "acceptance_data" / "databases" / "grading_system.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE grading_sessions SET rubric_path = ? WHERE id = ?",
            (str(outside_rubric), session_id),
        )
        connection.commit()

    with pytest.raises(
        acceptance.AcceptanceError,
        match="outside the acceptance workspace",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="config",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_accepts_confirmed_template_and_regions(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)

    report = acceptance.build_consistency_report(
        workspace,
        stage="template",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report["stage"] == "template"
    assert report["question_count"] == 2
    assert report["region_count"] == 2
    assert report["template_confirmed"] is True
    assert len(str(report["template_sha256"])) == 64
    encoded = json.dumps(report)
    assert str(workspace) not in encoded
    assert "region-q1" not in encoded
    assert "Q1" not in encoded


def test_consistency_report_rejects_pending_region_snapshot(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    template_id = _seed_confirmed_template(workspace, db, session_id)
    db.replace_answer_regions_atomic(
        session_id,
        template_id,
        [
            {
                "region_uuid": "pending-region",
                "mapped_question_id": "Q1",
            }
        ],
        confirmed=True,
    )

    with pytest.raises(
        acceptance.AcceptanceError,
        match="region snapshot is incomplete",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="template",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_accepts_part_regions_and_optional_name_region(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    config_dir = workspace / "acceptance_data" / "config" / "uploaded"
    (config_dir / "rubric.json").write_text(
        json.dumps(
            {
                "questions": [
                    {"question_id": "Q1", "max_score": 40},
                    {
                        "question_id": "Q2",
                        "max_score": 60,
                        "parts": [
                            {"part_id": "Q2(1)", "part_score": 30},
                            {"part_id": "Q2(2)", "part_score": 30},
                        ],
                    },
                ]
            }
        ),
        encoding="utf-8",
    )
    _seed_confirmed_template(
        workspace,
        db,
        session_id,
        regions=[
            {
                "region_uuid": "region-name",
                "mapped_question_id": "__student_name__",
                "mapping_status": "manual",
            },
            {
                "region_uuid": "region-q1",
                "mapped_question_id": "Q1",
                "mapping_status": "manual",
            },
            {
                "region_uuid": "region-q2-p1",
                "mapped_question_id": "Q2(P1)",
                "mapping_status": "manual",
            },
            {
                "region_uuid": "region-q2-p2",
                "mapped_question_id": "Q2(P2)",
                "mapping_status": "manual",
            },
        ],
    )

    report = acceptance.build_consistency_report(
        workspace,
        stage="template",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report["region_count"] == 4
    assert report["template_confirmed"] is True


def test_consistency_report_accepts_three_complete_grading_results(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)

    report = acceptance.build_consistency_report(
        workspace,
        stage="grading",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report["stage"] == "grading"
    assert report["grading_complete"] is True
    assert report["counts"] == {"papers": 3, "results": 3, "details": 6}
    encoded = json.dumps(report)
    assert "Student" not in encoded
    assert "S001" not in encoded
    assert str(workspace) not in encoded


def test_consistency_report_rejects_partial_grading(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)
    database = workspace / "acceptance_data" / "databases" / "grading_system.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE exam_papers SET processing_status = 'pending' WHERE id = 2"
        )
        connection.commit()

    with pytest.raises(
        acceptance.AcceptanceError,
        match="grading is incomplete or conflicted",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="grading",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_rejects_unresolved_teacher_review(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)
    database = workspace / "acceptance_data" / "databases" / "grading_system.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE session_details SET confidence_score = 50 WHERE id = 1"
        )
        connection.commit()

    with pytest.raises(
        acceptance.AcceptanceError,
        match="teacher review is incomplete",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="review",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_accepts_completed_teacher_review(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)

    report = acceptance.build_consistency_report(
        workspace,
        stage="review",
        session_id=session_id,
        source_data_root=source_data,
    )

    assert report["stage"] == "review"
    assert report["teacher_review_complete"] is True
    assert report["pending_review_count"] == 0


def test_consistency_report_accepts_report_matching_reviewed_scores(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)
    report_path = _write_matching_report(workspace)

    report = acceptance.build_consistency_report(
        workspace,
        stage="report",
        session_id=session_id,
        source_data_root=source_data,
        report_file=report_path,
    )

    assert report["stage"] == "report"
    assert report["report_matches_results"] is True
    assert report["report_row_count"] == 3
    assert len(str(report["report_sha256"])) == 64
    encoded = json.dumps(report)
    assert "S001" not in encoded
    assert "Student" not in encoded
    assert str(workspace) not in encoded


def test_consistency_report_rejects_report_score_conflict(
    tmp_path: Path,
) -> None:
    workspace, source_data, db, session_id = _prepared_report_workspace(tmp_path)
    _seed_confirmed_template(workspace, db, session_id)
    _seed_complete_grading(workspace, db, session_id)
    report_path = _write_matching_report(workspace)
    from openpyxl import load_workbook

    workbook = load_workbook(report_path)
    workbook["成绩与小题明细"]["B2"] = 69
    workbook.save(report_path)
    workbook.close()

    with pytest.raises(
        acceptance.AcceptanceError,
        match="report conflicts with reviewed results",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="report",
            session_id=session_id,
            source_data_root=source_data,
            report_file=report_path,
        )


def test_consistency_report_rechecks_authorized_source_fingerprint(
    tmp_path: Path,
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)
    source_database = source_data / "databases" / "grading_system.db"
    with sqlite3.connect(source_database) as connection:
        connection.execute("UPDATE students SET name = 'changed' WHERE id = 1")
        connection.commit()

    with pytest.raises(
        acceptance.AcceptanceError,
        match="authorized source changed after preparation",
    ):
        acceptance.build_consistency_report(
            workspace,
            stage="session",
            session_id=session_id,
            source_data_root=source_data,
        )


def test_consistency_report_cli_emits_only_safe_json(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workspace, source_data, _db, session_id = _prepared_report_workspace(tmp_path)

    exit_code = acceptance.main(
        [
            "report",
            "--workspace",
            str(workspace),
            "--stage",
            "session",
            "--session-id",
            str(session_id),
            "--source-data-root",
            str(source_data),
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["stage"] == "session"
    assert payload["ok"] is True
    assert str(workspace) not in json.dumps(payload)


def test_validate_workspace_accepts_only_empty_directory_below_system_temp(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "p2-20-workspace"

    assert acceptance.validate_workspace(workspace) == workspace.resolve()

    workspace.mkdir()
    (workspace / "unexpected.txt").write_text("occupied", encoding="utf-8")
    with pytest.raises(acceptance.AcceptanceError, match="empty"):
        acceptance.validate_workspace(workspace)


def test_validate_workspace_rejects_temp_root_repository_and_user_data_paths(
    tmp_path: Path,
) -> None:
    repo_root = Path(acceptance.__file__).resolve().parents[1]

    with pytest.raises(acceptance.AcceptanceError, match="strictly below"):
        acceptance.validate_workspace(Path(tempfile.gettempdir()))
    with pytest.raises(acceptance.AcceptanceError, match="system temporary"):
        acceptance.validate_workspace(repo_root / "output" / "p2-20")
    with pytest.raises(acceptance.AcceptanceError, match="user_data"):
        acceptance.validate_workspace(tmp_path / "user_data" / "p2-20")


def test_prepare_code_workspace_stages_exact_commit_without_user_data(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"

    metadata = acceptance.prepare_code_workspace(
        source_sha,
        workspace,
        repo_root=repo,
    )

    assert metadata == {
        "package": "P2-20",
        "source_sha": source_sha,
        "state": "code_prepared",
        "authorized_session_id": None,
    }
    assert (workspace / "safe.txt").read_text(encoding="utf-8") == "safe"
    assert not (workspace / "user_data").exists()
    assert acceptance.load_metadata(workspace) == metadata


def test_prepare_code_workspace_rejects_symbolic_ref(tmp_path: Path) -> None:
    repo, _source_sha = _committed_source_repo(tmp_path)

    with pytest.raises(acceptance.AcceptanceError, match="40-character"):
        acceptance.prepare_code_workspace(
            "HEAD",
            tmp_path / "prepared",
            repo_root=repo,
        )


def test_prepare_authorized_inputs_copies_only_selected_exam_materials(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    source_database = source_data / "databases" / "grading_system.db"
    database_hash = _sha256(source_database)

    metadata = acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )

    assert metadata["state"] == "inputs_prepared"
    assert metadata["authorized_session_id"] == 1
    assert metadata["paper_count"] == 3
    assert _sha256(source_database) == database_hash
    assert not (workspace / "acceptance_inputs" / "databases").exists()
    assert (workspace / "acceptance_inputs" / "source-paper.docx").is_file()
    assert (workspace / "acceptance_inputs" / "template.pdf").is_file()
    assert sorted(
        path.name
        for path in (workspace / "acceptance_inputs" / "answer-sheets").iterdir()
    ) == [
        "answer-01-back.jpg",
        "answer-01-front.jpg",
        "answer-02-back.jpg",
        "answer-02-front.jpg",
        "answer-03-back.jpg",
        "answer-03-front.jpg",
    ]
    roster = (workspace / "acceptance_inputs" / "roster.csv").read_text(
        encoding="utf-8-sig"
    )
    assert roster.splitlines() == [
        "student_code,name,class_name",
        "S001,学生1,七年级一班",
        "S002,学生2,七年级一班",
        "S003,学生3,七年级一班",
    ]
    backup = workspace / "authorized-inputs-backup.zip"
    with zipfile.ZipFile(backup) as archive:
        assert "manifest.json" in archive.namelist()
        assert "roster.csv" in archive.namelist()
        assert not any("database" in name for name in archive.namelist())


def test_prepare_authorized_inputs_fails_closed_on_wrong_exam_or_too_few_papers(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path, paper_count=2)

    with pytest.raises(acceptance.AcceptanceError, match="authorized exam"):
        acceptance.prepare_authorized_inputs(
            workspace,
            source_data_root=source_data,
            session_id=1,
            expected_session_name="not-authorized",
            paper_limit=2,
        )
    with pytest.raises(acceptance.AcceptanceError, match="eligible answer sheets"):
        acceptance.prepare_authorized_inputs(
            workspace,
            source_data_root=source_data,
            session_id=1,
            expected_session_name="0609",
            paper_limit=3,
        )
    assert acceptance.load_metadata(workspace)["state"] == "code_prepared"


def test_prepare_runtime_profile_keeps_only_bounded_config_and_grading_secrets(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "machine-config" / "api_profiles.json"
    profile_path.parent.mkdir()
    profile_path.write_text(
        json.dumps(
            [
                {"name": "old", "api_key": "old-secret"},
                {
                    "name": "active",
                    "provider": "custom-openai-compatible",
                    "base_url": "https://grading.example.invalid/v1",
                    "api_key": "grading-secret",
                    "grading_model": "grading-model",
                    "ocr_model": "ocr-model",
                    "config_provider": "custom-openai-compatible",
                    "config_base_url": "https://config.example.invalid/v1",
                    "config_api_key": "config-secret",
                    "config_model": "config-model",
                    "objective_enabled": True,
                    "objective_api_key": "objective-secret",
                    "tagging_api_key": "tagging-secret",
                },
            ]
        ),
        encoding="utf-8",
    )
    profile_hash = _sha256(profile_path)

    metadata = acceptance.prepare_runtime_profile(
        workspace,
        source_profile_path=profile_path,
        proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
        max_forwarded_requests=4,
    )

    assert metadata["state"] == "runtime_ready"
    assert metadata["model_request_budget"] == 4
    assert _sha256(profile_path) == profile_hash
    runtime_profiles = json.loads(
        (
            workspace / "acceptance_config" / "api_profiles.json"
        ).read_text(encoding="utf-8")
    )
    assert len(runtime_profiles) == 1
    runtime_profile = runtime_profiles[0]
    assert runtime_profile["name"] == "P2-20 Acceptance"
    assert runtime_profile["base_url"].startswith("http://127.0.0.1:8120/")
    assert runtime_profile["config_base_url"] == runtime_profile["base_url"]
    assert runtime_profile["api_key"] == "acceptance-proxy"
    assert runtime_profile["config_api_key"] == "acceptance-proxy"
    assert runtime_profile["ocr_model"] == "p2-20-local-ocr"
    assert runtime_profile["objective_enabled"] is False
    assert runtime_profile["grading_max_workers"] == 1
    assert runtime_profile["llm_grading_max_retries"] == 0
    assert runtime_profile["llm_config_generation_max_retries"] == 0
    assert runtime_profile["llm_recognition_max_retries"] == 0
    assert not any("tagging" in key for key in runtime_profile)
    assert "objective_api_key" not in runtime_profile
    upstream = json.loads(
        (workspace / "acceptance_config" / "upstream.json").read_text(
            encoding="utf-8"
        )
    )
    assert upstream == {
        "grading": {
            "base_url": "https://grading.example.invalid/v1",
            "api_key": "grading-secret",
            "model": "grading-model",
        },
        "config": {
            "base_url": "https://config.example.invalid/v1",
            "api_key": "config-secret",
            "model": "config-model",
        },
        "max_forwarded_requests": 4,
    }


def test_prepare_runtime_profile_supports_manual_counting_without_a_hard_limit(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "machine-config" / "api_profiles.json"
    profile_path.parent.mkdir()
    profile_path.write_text(
        json.dumps(
            [
                {
                    "name": "active",
                    "provider": "custom-openai-compatible",
                    "base_url": "https://apic1.ohmycdn.com/v1",
                    "api_key": "shared-secret",
                    "grading_model": "grading-model",
                    "config_model": "config-model",
                }
            ]
        ),
        encoding="utf-8",
    )

    metadata = acceptance.prepare_runtime_profile(
        workspace,
        source_profile_path=profile_path,
        proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
        max_forwarded_requests=None,
        initial_forwarded_requests=9,
        direct_diagnostic_requests=1,
    )

    assert metadata["model_request_budget"] is None
    assert metadata["model_request_count_baseline"] == 9
    assert metadata["direct_diagnostic_requests"] == 1
    upstream = json.loads(
        (workspace / "acceptance_config" / "upstream.json").read_text(
            encoding="utf-8"
        )
    )
    assert upstream["max_forwarded_requests"] is None
    assert upstream["initial_forwarded_requests"] == 9
    assert upstream["direct_diagnostic_requests"] == 1
    assert upstream["manual_request_control"] == {
        "default_route_id": "ohmygpt",
        "routes": {
            "ohmycdn": "https://apic1.ohmycdn.com/v1",
            "ohmygpt": "https://api.ohmygpt.com/v1",
        },
    }
    serialized = json.dumps(upstream, ensure_ascii=False)
    assert "shared-secret" in serialized
    assert "acceptance-proxy" not in serialized


def test_prepare_runtime_profile_rejects_unsupported_or_unconfigured_provider(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "profiles.json"
    profile_path.write_text(
        json.dumps(
            [
                {
                    "name": "active",
                    "provider": "unsupported",
                    "base_url": "",
                    "api_key": "",
                    "grading_model": "",
                    "config_model": "",
                }
            ]
        ),
        encoding="utf-8",
    )

    with pytest.raises(acceptance.AcceptanceError, match="OpenAI-compatible"):
        acceptance.prepare_runtime_profile(
            workspace,
            source_profile_path=profile_path,
            proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
            max_forwarded_requests=4,
        )
    assert acceptance.load_metadata(workspace)["state"] == "inputs_prepared"


@pytest.mark.parametrize("budget", [4, 5, 6, 9])
def test_model_budget_proxy_stubs_ocr_and_hard_stops_at_explicit_budget(
    tmp_path: Path,
    budget: int,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": budget,
            }
        ),
        encoding="utf-8",
    )
    forwarded: list[dict[str, object]] = []

    async def fake_forward(
        upstream: dict[str, str],
        payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        forwarded.append({"upstream": upstream, "payload": payload})
        return (
            200,
            {"content-type": "application/json"},
            json.dumps(
                {
                    "id": "bounded",
                    "choices": [
                        {"message": {"role": "assistant", "content": "ok"}}
                    ],
                }
            ).encode(),
        )

    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=fake_forward,
    )
    client = TestClient(app)
    headers = {"Authorization": "Bearer acceptance-proxy"}

    ocr = client.post(
        "/v1/chat/completions",
        headers=headers,
        json={"model": "p2-20-local-ocr", "messages": []},
    )
    assert ocr.status_code == 200
    assert ocr.json()["choices"][0]["message"]["content"] == "NOT_FOUND"
    assert forwarded == []

    for _ in range(budget):
        response = client.post(
            "/v1/chat/completions",
            headers=headers,
            json={"model": "shared-model", "messages": []},
        )
        assert response.status_code == 200
    blocked = client.post(
        "/v1/chat/completions",
        headers=headers,
        json={"model": "shared-model", "messages": []},
    )
    assert blocked.status_code == 429
    assert len(forwarded) == budget
    assert all(
        item["upstream"]["api_key"] in {"grading-secret", "config-secret"}
        for item in forwarded
    )
    state = json.loads(
        (tmp_path / "budget-state.json").read_text(encoding="utf-8")
    )
    assert state == {
        "forwarded_requests": budget,
        "local_ocr_requests": 1,
        "max_forwarded_requests": budget,
    }
    assert "secret" not in blocked.text


def test_manual_model_control_counts_past_nine_and_freezes_each_selected_route(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    forwarded_routes: list[str] = []

    async def fake_forward(
        upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        forwarded_routes.append(upstream["base_url"])
        return (
            200,
            {"content-type": "application/json"},
            b'{"choices":[{"finish_reason":"stop","message":{"content":"ok"}}]}',
        )

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=fake_forward,
    )
    app = FastAPI()
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    client = TestClient(app)
    control_headers = {"X-Acceptance-Control-Token": "expected-token"}
    proxy_headers = {"Authorization": "Bearer acceptance-proxy"}

    for offset in range(11):
        route_id = "ohmycdn" if offset % 2 else "ohmygpt"
        token = f"{offset + 1:032x}"
        permit = client.post(
            "/acceptance-control/model-permit",
            headers=control_headers,
            json={
                "action": "q11",
                "route_id": route_id,
                "token": token,
                "source_job_id": 5,
            },
        )
        assert permit.status_code == 200
        response = client.post(
            "/acceptance-llm/v1/chat/completions",
            headers=proxy_headers,
            json={
                "model": "config-model",
                "messages": [
                    {
                        "role": "user",
                        "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]\nprivate prompt',
                    }
                ],
            },
        )
        assert response.status_code == 200

    assert forwarded_routes == [
        (
            "https://apic1.ohmycdn.com/v1"
            if offset % 2
            else "https://api.ohmygpt.com/v1"
        )
        for offset in range(11)
    ]
    state_response = client.get(
        "/acceptance-control/model-state",
        headers=control_headers,
    )
    assert state_response.status_code == 200
    state = state_response.json()
    assert state["mode"] == "manual_unbounded"
    assert state["forwarded_requests"] == 20
    assert state["direct_diagnostic_requests"] == 1
    assert state["total_external_requests"] == 21
    assert state["max_forwarded_requests"] is None
    assert state["in_flight_request"] is None
    assert state["active_permit"] is None
    assert state["outcome_counts"]["success"] == 11
    serialized = state_response.text
    assert "secret" not in serialized
    assert "private prompt" not in serialized


def test_manual_model_control_requires_one_matching_permit_and_an_allowed_route(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    forwards = 0

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        nonlocal forwards
        forwards += 1
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=fake_forward,
    )
    app = FastAPI()
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    client = TestClient(app)
    control_headers = {"X-Acceptance-Control-Token": "expected-token"}
    proxy_headers = {"Authorization": "Bearer acceptance-proxy"}

    q11_payload = {
        "model": "config-model",
        "messages": [
            {
                "role": "user",
                "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]\nprivate',
            }
        ],
    }
    q12_payload = {
        "model": "config-model",
        "messages": [
            {
                "role": "user",
                "content": 'BATCH_QUESTION_IDS_JSON=["Q12"]\nprivate',
            }
        ],
    }

    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json=q11_payload,
    ).status_code == 428
    invalid_route = client.post(
        "/acceptance-control/model-permit",
        headers=control_headers,
        json={
            "action": "q11",
            "route_id": "custom",
            "token": "1" * 32,
            "source_job_id": 5,
        },
    )
    assert invalid_route.status_code == 409
    assert "https://" not in invalid_route.text

    assert client.post(
        "/acceptance-control/model-permit",
        headers=control_headers,
        json={
            "action": "q11",
            "route_id": "ohmycdn",
            "token": "2" * 32,
            "source_job_id": 5,
        },
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json=q12_payload,
    ).status_code == 428
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json=q11_payload,
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json=q11_payload,
    ).status_code == 428

    state = client.get(
        "/acceptance-control/model-state",
        headers=control_headers,
    ).json()
    assert forwards == 1
    assert state["forwarded_requests"] == 10
    assert state["outcome_counts"]["success"] == 1
    assert state["active_permit"] is None


def test_manual_continuation_page_is_local_explicit_and_contains_no_secret(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-private-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-private-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=lambda *_args: None,
    )
    app = FastAPI()
    acceptance.mount_manual_continuation(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
        session_id=1,
    )

    response = TestClient(app).get("/manual")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    assert "仅请求 Q11" in response.text
    assert "仅请求 Q12" in response.text
    assert "AI 统一配分" in response.text
    assert "已等待" in response.text
    assert "https://api.ohmygpt.com/v1" in response.text
    assert "https://apic1.ohmycdn.com/v1" in response.text
    assert "不设总次数上限" in response.text
    assert "程序不会自动重试" in response.text
    assert "最近一次" in response.text
    assert "grading-private-secret" not in response.text
    assert "config-private-secret" not in response.text
    assert str(tmp_path) not in response.text
    assert TestClient(app).get("/acceptance-control/manual-state").status_code == 404


def test_manual_unbounded_proxy_rejects_sensitive_request_capture(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )

    with pytest.raises(
        acceptance.AcceptanceError,
        match="manual mode does not allow model request capture",
    ):
        acceptance.create_model_budget_proxy(
            upstream_path,
            state_path=tmp_path / "budget-state.json",
            request_capture_path=tmp_path / "request-capture.json",
            forwarder=lambda *_args: None,
        )

    assert not (tmp_path / "budget-state.json").exists()
    assert not (tmp_path / "request-capture.json").exists()


def test_manual_action_is_stage_bound_and_replays_the_same_click_token(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    forwarded = 0

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        nonlocal forwarded
        forwarded += 1
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=fake_forward,
    )

    class Job:
        def __init__(
            self,
            job_id: int,
            *,
            failed_ids: list[str],
        ) -> None:
            self.id = job_id
            self.status = "succeeded"
            self.progress = 1.0
            self.stage = "partial"
            self.detail = ""
            self.result = {
                "outcome": "partial",
                "failed_question_ids": failed_ids,
                "failed_batches": [
                    {"batch_id": "B005", "question_ids": ["Q11"]},
                    {"batch_id": "B006", "question_ids": ["Q12"]},
                ],
            }

    class Manager:
        def __init__(self) -> None:
            self.jobs = {5: Job(5, failed_ids=["Q11", "Q12"])}

        def list(
            self,
            *,
            statuses: tuple[str, ...] = (),
            **_kwargs: object,
        ) -> tuple[list[Job], int]:
            jobs = [
                job
                for job in sorted(
                    self.jobs.values(), key=lambda item: item.id, reverse=True
                )
                if not statuses or job.status in statuses
            ]
            return jobs[:1], len(jobs)

        def get(self, job_id: int) -> Job | None:
            return self.jobs.get(job_id)

    manager = Manager()
    submitted: list[dict[str, object]] = []

    def submit_action(**kwargs: object) -> dict[str, object]:
        submitted.append(dict(kwargs))
        manager.jobs[6] = Job(6, failed_ids=["Q11", "Q12"])
        return {"id": 6, "status": "queued"}

    app = FastAPI()
    app.state.job_manager = manager
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_manual_continuation(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
        session_id=1,
        submit_action=submit_action,
    )
    client = TestClient(app)
    headers = {
        "Content-Type": "application/json",
        "X-Acceptance-Control-Token": "expected-token",
    }
    token = "a" * 32

    wrong_stage = client.post(
        "/acceptance-control/manual-action",
        headers=headers,
        json={
            "action": "q12",
            "route_id": "ohmycdn",
            "source_job_id": 5,
            "token": "b" * 32,
        },
    )
    assert wrong_stage.status_code == 409

    first = client.post(
        "/acceptance-control/manual-action",
        headers=headers,
        json={
            "action": "q11",
            "route_id": "ohmycdn",
            "source_job_id": 5,
            "token": token,
        },
    )
    duplicate = client.post(
        "/acceptance-control/manual-action",
        headers=headers,
        json={
            "action": "q11",
            "route_id": "ohmycdn",
            "source_job_id": 5,
            "token": token,
        },
    )
    assert first.status_code == 200
    assert duplicate.status_code == 200
    assert first.json() == duplicate.json() == {"id": 6, "status": "queued"}
    assert len(submitted) == 1

    q11 = client.post(
        "/acceptance-llm/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={
            "model": "config-model",
            "messages": [
                {
                    "role": "user",
                    "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]\nprivate',
                }
            ],
        },
    )
    automatic_score = client.post(
        "/acceptance-llm/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={
            "model": "config-model",
            "messages": [
                {
                    "role": "user",
                    "content": (
                        "SCORE_QUESTION_IDS_JSON="
                        + json.dumps([f"Q{number}" for number in range(1, 13)])
                    ),
                }
            ],
        },
    )
    assert q11.status_code == 200
    assert automatic_score.status_code == 428
    assert forwarded == 1


def test_manual_action_uses_the_real_config_retry_api_contract(
    tmp_path: Path,
) -> None:
    from backend.jobs.manager import JobManager
    from backend.jobs.store import JobStore

    database_path = tmp_path / "grading.db"
    db = DBManager(database_path)
    db.initialize()
    rubric = tmp_path / "rubric.json"
    answer = tmp_path / "answer.json"
    rubric.write_text("{}", encoding="utf-8")
    answer.write_text("{}", encoding="utf-8")
    session_id = db.create_grading_session(
        "P2-20 manual continuation",
        str(rubric),
        str(answer),
    )
    store = JobStore(database_path)
    source = store.create_job(
        "config_generation",
        {
            "session_id": session_id,
            "generation_mode": "batched",
            "input_id": "input-1",
            "source_id": "a" * 32,
            "source_revision": "b" * 64,
        },
    )
    partial = {
        "outcome": "partial",
        "failed_question_ids": ["Q11", "Q12"],
        "failed_batches": [
            {"batch_id": "B005", "question_ids": ["Q11"]},
            {"batch_id": "B006", "question_ids": ["Q12"]},
        ],
        "score_allocation_pending": False,
    }
    store.finish(source.id, "succeeded", result=partial)
    manager = JobManager(store, max_workers=1, cleanup_interrupted=False)
    manager.register("config_generation", lambda _context: partial)

    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=lambda *_args: None,
    )
    app = FastAPI()
    app.state.job_manager = manager

    class Paths:
        db_path = database_path

    app.state.path_manager = Paths()
    acceptance.mount_manual_continuation(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
        session_id=session_id,
    )
    client = TestClient(app)
    token = "f" * 32
    try:
        response = client.post(
            "/acceptance-control/manual-action",
            headers={"X-Acceptance-Control-Token": "expected-token"},
            json={
                "action": "q11",
                "route_id": "ohmycdn",
                "source_job_id": source.id,
                "token": token,
            },
        )
        assert response.status_code == 200
        job_id = int(response.json()["id"])
        manager.wait(job_id, timeout=5)
        created = manager.get(job_id)
        assert created is not None
        assert created.payload["session_id"] == session_id
        assert created.payload["source_job_id"] == source.id
        assert created.payload["retry_question_ids"] == ["Q11"]
        assert created.payload["client_request_token"] == token
        assert proxy.state.model_control.snapshot()["active_permit"] == {
            "action": "q11",
            "route_id": "ohmycdn",
            "source_job_id": source.id,
        }
    finally:
        proxy.state.model_control.revoke_pending_permit(token)
        manager.shutdown()


def test_q12_and_score_allocation_each_require_their_own_manual_click(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    forwards: list[str] = []

    async def fake_forward(
        upstream: dict[str, str],
        payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        serialized = json.dumps(payload)
        action = "q12" if "BATCH_QUESTION_IDS_JSON" in serialized else "score"
        forwards.append(f"{upstream['_route_id']}:{action}")
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=fake_forward,
    )

    class Job:
        def __init__(
            self,
            job_id: int,
            *,
            failed_ids: list[str],
            score_pending: bool,
            outcome: str = "partial",
        ) -> None:
            self.id = job_id
            self.status = "succeeded"
            self.progress = 1.0
            self.stage = "complete" if outcome == "complete" else "partial"
            self.detail = ""
            self.result = {
                "outcome": outcome,
                "failed_question_ids": failed_ids,
                "failed_batches": [
                    {"batch_id": "B006", "question_ids": ["Q12"]}
                ]
                if failed_ids
                else [],
                "score_allocation_pending": score_pending,
            }

    class Manager:
        def __init__(self) -> None:
            self.jobs = {
                5: Job(
                    5,
                    failed_ids=["Q12"],
                    score_pending=False,
                )
            }

        def list(
            self,
            *,
            statuses: tuple[str, ...] = (),
            **_kwargs: object,
        ) -> tuple[list[Job], int]:
            jobs = [
                job
                for job in sorted(
                    self.jobs.values(), key=lambda item: item.id, reverse=True
                )
                if not statuses or job.status in statuses
            ]
            return jobs[:1], len(jobs)

        def get(self, job_id: int) -> Job | None:
            return self.jobs.get(job_id)

    manager = Manager()

    def submit_action(**kwargs: object) -> dict[str, object]:
        if kwargs["action"] == "q12":
            manager.jobs[6] = Job(
                6,
                failed_ids=[],
                score_pending=True,
            )
            return {"id": 6, "status": "queued"}
        manager.jobs[7] = Job(
            7,
            failed_ids=[],
            score_pending=False,
            outcome="complete",
        )
        return {"id": 7, "status": "queued"}

    app = FastAPI()
    app.state.job_manager = manager
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_manual_continuation(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
        session_id=1,
        submit_action=submit_action,
    )
    client = TestClient(app)
    control_headers = {"X-Acceptance-Control-Token": "expected-token"}
    proxy_headers = {"Authorization": "Bearer acceptance-proxy"}
    score_marker = (
        "SCORE_QUESTION_IDS_JSON="
        + json.dumps([f"Q{number}" for number in range(1, 13)])
    )

    assert client.post(
        "/acceptance-control/manual-action",
        headers=control_headers,
        json={
            "action": "q12",
            "route_id": "ohmycdn",
            "source_job_id": 5,
            "token": "1" * 32,
        },
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json={
            "model": "config-model",
            "messages": [
                {
                    "role": "user",
                    "content": 'BATCH_QUESTION_IDS_JSON=["Q12"]',
                }
            ],
        },
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json={
            "model": "config-model",
            "messages": [{"role": "user", "content": score_marker}],
        },
    ).status_code == 428
    state_after_q12 = client.get(
        "/acceptance-control/manual-state",
        headers=control_headers,
    ).json()
    assert state_after_q12["available_action"] == "score"
    assert state_after_q12["forwarded_requests"] == 10

    assert client.post(
        "/acceptance-control/manual-action",
        headers=control_headers,
        json={
            "action": "score",
            "route_id": "ohmygpt",
            "source_job_id": 6,
            "token": "2" * 32,
        },
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers=proxy_headers,
        json={
            "model": "config-model",
            "messages": [{"role": "user", "content": score_marker}],
        },
    ).status_code == 200

    assert forwards == ["ohmycdn:q12", "ohmygpt:score"]
    assert proxy.state.model_control.snapshot()["forwarded_requests"] == 11


def test_manual_model_control_allows_only_one_concurrent_permit_consumer(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    entered = threading.Event()
    release = threading.Event()
    forwards = 0

    async def blocking_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        nonlocal forwards
        forwards += 1
        entered.set()
        assert release.wait(timeout=5)
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=blocking_forward,
    )
    app = FastAPI()
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    control = TestClient(app)
    assert control.post(
        "/acceptance-control/model-permit",
        headers={"X-Acceptance-Control-Token": "expected-token"},
        json={
            "action": "q11",
            "route_id": "ohmygpt",
            "token": "c" * 32,
            "source_job_id": 5,
        },
    ).status_code == 200
    payload = {
        "model": "config-model",
        "messages": [
            {
                "role": "user",
                "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]\nprivate',
            }
        ],
    }
    headers = {"Authorization": "Bearer acceptance-proxy"}

    with ThreadPoolExecutor(max_workers=1) as executor:
        first = executor.submit(
            TestClient(app).post,
            "/acceptance-llm/v1/chat/completions",
            headers=headers,
            json=payload,
        )
        assert entered.wait(timeout=5)
        second = TestClient(app).post(
            "/acceptance-llm/v1/chat/completions",
            headers=headers,
            json=payload,
        )
        release.set()
        first_response = first.result(timeout=5)

    assert first_response.status_code == 200
    assert second.status_code == 428
    assert forwards == 1
    assert control.get(
        "/acceptance-control/model-state",
        headers={"X-Acceptance-Control-Token": "expected-token"},
    ).json()["forwarded_requests"] == 10


def test_manual_model_control_counts_transport_failure_then_allows_manual_retry(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    attempts = 0

    async def flaky_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            try:
                raise httpx.ReadTimeout("private transport detail")
            except httpx.ReadTimeout as exc:
                raise acceptance.AcceptanceError(
                    "approved model request failed at the provider"
                ) from exc
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        forwarder=flaky_forward,
    )
    app = FastAPI()
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    client = TestClient(app)
    control_headers = {"X-Acceptance-Control-Token": "expected-token"}
    payload = {
        "model": "config-model",
        "messages": [
            {
                "role": "user",
                "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]\nprivate',
            }
        ],
    }

    statuses: list[int] = []
    for index, route_id in enumerate(("ohmygpt", "ohmycdn"), start=1):
        assert client.post(
            "/acceptance-control/model-permit",
            headers=control_headers,
            json={
                "action": "q11",
                "route_id": route_id,
                "token": f"{index:032x}",
                "source_job_id": 5,
            },
        ).status_code == 200
        statuses.append(
            client.post(
                "/acceptance-llm/v1/chat/completions",
                headers={"Authorization": "Bearer acceptance-proxy"},
                json=payload,
            ).status_code
        )

    state = client.get(
        "/acceptance-control/model-state",
        headers=control_headers,
    ).json()
    assert statuses == [502, 200]
    assert attempts == 2
    assert state["forwarded_requests"] == 11
    assert state["outcome_counts"] == {
        "success": 1,
        "upstream_http_error": 0,
        "transport_error": 1,
        "unknown": 0,
    }
    assert state["last_request"]["route_id"] == "ohmycdn"


def test_manual_model_control_recovers_an_interrupted_request_as_unknown(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )

    class SimulatedProcessExit(BaseException):
        pass

    async def interrupted_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        raise SimulatedProcessExit()

    state_path = tmp_path / "budget-state.json"
    first_proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=state_path,
        forwarder=interrupted_forward,
    )
    first_app = FastAPI()
    first_app.mount("/acceptance-llm", first_proxy)
    acceptance.mount_acceptance_model_control(
        first_app,
        controller=first_proxy.state.model_control,
        control_token="expected-token",
    )
    first_client = TestClient(first_app)
    assert first_client.post(
        "/acceptance-control/model-permit",
        headers={"X-Acceptance-Control-Token": "expected-token"},
        json={
            "action": "q11",
            "route_id": "ohmygpt",
            "token": "d" * 32,
            "source_job_id": 5,
        },
    ).status_code == 200
    with pytest.raises(SimulatedProcessExit):
        first_client.post(
            "/acceptance-llm/v1/chat/completions",
            headers={"Authorization": "Bearer acceptance-proxy"},
            json={
                "model": "config-model",
                "messages": [
                    {
                        "role": "user",
                        "content": 'BATCH_QUESTION_IDS_JSON=["Q11"]',
                    }
                ],
            },
        )

    persisted = json.loads(state_path.read_text(encoding="utf-8"))
    assert persisted["forwarded_requests"] == 10
    assert persisted["in_flight_request"]["number"] == 10

    recovered_proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=state_path,
        forwarder=interrupted_forward,
    )
    recovered_app = FastAPI()
    acceptance.mount_acceptance_model_control(
        recovered_app,
        controller=recovered_proxy.state.model_control,
        control_token="expected-token",
    )
    recovered = TestClient(recovered_app).get(
        "/acceptance-control/model-state",
        headers={"X-Acceptance-Control-Token": "expected-token"},
    ).json()

    assert recovered["forwarded_requests"] == 10
    assert recovered["outcome_counts"]["unknown"] == 1
    assert recovered["in_flight_request"] is None
    assert recovered["last_request"]["number"] == 10
    assert recovered["last_request"]["outcome"] == "unknown"


def test_manual_model_control_rejects_corrupted_persisted_permission(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    state_path = tmp_path / "budget-state.json"
    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=state_path,
        forwarder=lambda *_args: None,
    )
    app = FastAPI()
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    assert TestClient(app).post(
        "/acceptance-control/model-permit",
        headers={"X-Acceptance-Control-Token": "expected-token"},
        json={
            "action": "q11",
            "route_id": "ohmygpt",
            "token": "e" * 32,
            "source_job_id": 5,
        },
    ).status_code == 200
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["active_permit"]["route_id"] = "unlisted"
    state_path.write_text(json.dumps(state), encoding="utf-8")

    with pytest.raises(
        acceptance.AcceptanceError,
        match="manual model request state",
    ):
        acceptance.create_model_budget_proxy(
            upstream_path,
            state_path=state_path,
            forwarder=lambda *_args: None,
        )


def test_manual_model_diagnostic_records_the_selected_route_without_content(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "grading-private-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://unused.example.invalid/v1",
                    "api_key": "config-private-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": None,
                "initial_forwarded_requests": 9,
                "direct_diagnostic_requests": 1,
                "manual_request_control": {
                    "default_route_id": "ohmygpt",
                    "routes": {
                        "ohmycdn": "https://apic1.ohmycdn.com/v1",
                        "ohmygpt": "https://api.ohmygpt.com/v1",
                    },
                },
            }
        ),
        encoding="utf-8",
    )
    diagnostic_path = tmp_path / "model-proxy.jsonl"

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        return 200, {"content-type": "application/json"}, b"{}"

    proxy = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=diagnostic_path,
        forwarder=fake_forward,
    )
    app = FastAPI()
    app.mount("/acceptance-llm", proxy)
    acceptance.mount_acceptance_model_control(
        app,
        controller=proxy.state.model_control,
        control_token="expected-token",
    )
    client = TestClient(app)
    assert client.post(
        "/acceptance-control/model-permit",
        headers={"X-Acceptance-Control-Token": "expected-token"},
        json={
            "action": "q11",
            "route_id": "ohmycdn",
            "token": "9" * 32,
            "source_job_id": 5,
        },
    ).status_code == 200
    assert client.post(
        "/acceptance-llm/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={
            "model": "config-model",
            "messages": [
                {
                    "role": "user",
                    "content": (
                        'BATCH_QUESTION_IDS_JSON=["Q11"]\n'
                        "private prompt body"
                    ),
                }
            ],
        },
    ).status_code == 200

    event = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert event["request_number"] == 10
    assert event["max_forwarded_requests"] is None
    assert event["route_id"] == "ohmycdn"
    assert event["upstream_host"] == "apic1.ohmycdn.com"
    serialized = diagnostic_path.read_text(encoding="utf-8")
    assert "private prompt body" not in serialized
    assert "grading-private-secret" not in serialized
    assert "config-private-secret" not in serialized


@pytest.mark.parametrize(
    (
        "forward_result",
        "raised_cause",
        "expected_outcome",
        "expected_status",
        "expected_error_type",
    ),
    [
        (
            (
                502,
                {"content-type": "text/plain"},
                b"provider-private-response",
            ),
            None,
            "upstream_http_error",
            502,
            "",
        ),
        (
            None,
            httpx.ReadTimeout(
                "transport-private-detail",
                request=httpx.Request(
                    "POST", "https://provider.example.invalid/v1/chat/completions"
                ),
            ),
            "transport_error",
            None,
            "ReadTimeout",
        ),
    ],
)
def test_model_budget_proxy_diagnostic_log_distinguishes_502_source_without_content(
    tmp_path: Path,
    forward_result: tuple[int, dict[str, str], bytes] | None,
    raised_cause: Exception | None,
    expected_outcome: str,
    expected_status: int | None,
    expected_error_type: str,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-super-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-super-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": 1,
            }
        ),
        encoding="utf-8",
    )
    diagnostic_path = tmp_path / "model-proxy.jsonl"
    payload = {
        "model": "shared-model",
        "messages": [{"role": "user", "content": "private-q12-prompt"}],
        "image_url": "data:image/png;base64,private-image",
    }

    async def fake_forward(
        upstream: dict[str, str],
        request_payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        assert upstream["api_key"] in {"grading-super-secret", "config-super-secret"}
        assert request_payload == payload
        if raised_cause is not None:
            try:
                raise raised_cause
            except Exception as exc:
                raise acceptance.AcceptanceError(
                    "approved model request failed at the provider"
                ) from exc
        assert forward_result is not None
        return forward_result

    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=diagnostic_path,
        forwarder=fake_forward,
    )
    response = TestClient(app).post(
        "/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json=payload,
    )

    assert response.status_code == 502
    event = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    expected_payload = json.dumps(
        payload,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    expected_response = forward_result[2] if forward_result is not None else b""
    assert event["schema_version"] == 1
    assert event["event"] == "model_proxy_forward"
    assert event["outcome"] == expected_outcome
    assert event["upstream_status_code"] == expected_status
    assert event["transport_error_type"] == expected_error_type
    assert event["model"] == "shared-model"
    assert event["request_number"] == 1
    assert event["max_forwarded_requests"] == 1
    assert event["elapsed_ms"] >= 0
    assert event["payload_bytes"] == len(expected_payload)
    assert event["payload_sha256"] == hashlib.sha256(expected_payload).hexdigest()
    assert event["response_bytes"] == len(expected_response)
    assert event["response_sha256"] == (
        hashlib.sha256(expected_response).hexdigest() if expected_response else ""
    )
    assert set(event) == {
        "schema_version",
        "event",
        "recorded_at_unix_ms",
        "outcome",
        "upstream_status_code",
        "transport_error_type",
        "elapsed_ms",
        "model",
        "request_number",
        "max_forwarded_requests",
        "payload_bytes",
        "payload_sha256",
        "response_bytes",
        "response_sha256",
        "finish_reason",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
    }
    serialized = diagnostic_path.read_text(encoding="utf-8")
    assert "super-secret" not in serialized
    assert "private-q12-prompt" not in serialized
    assert "private-image" not in serialized
    assert "provider-private-response" not in serialized
    assert "transport-private-detail" not in serialized
    assert str(tmp_path) not in serialized


def test_model_budget_proxy_diagnostic_log_summarizes_success_without_response_text(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": 1,
            }
        ),
        encoding="utf-8",
    )
    response_body = json.dumps(
        {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {"role": "assistant", "content": "private-answer"},
                }
            ],
            "usage": {
                "prompt_tokens": 11,
                "completion_tokens": 7,
                "total_tokens": 18,
            },
        }
    ).encode()

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        return 200, {"content-type": "application/json"}, response_body

    diagnostic_path = tmp_path / "model-proxy.jsonl"
    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=diagnostic_path,
        forwarder=fake_forward,
    )
    response = TestClient(app).post(
        "/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={
            "model": "shared-model",
            "messages": [{"role": "user", "content": "private-prompt"}],
        },
    )

    assert response.status_code == 200
    assert response.json()["choices"][0]["message"]["content"] == "private-answer"
    event = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert event["outcome"] == "success"
    assert event["upstream_status_code"] == 200
    assert event["transport_error_type"] == ""
    assert event["finish_reason"] == "stop"
    assert event["prompt_tokens"] == 11
    assert event["completion_tokens"] == 7
    assert event["total_tokens"] == 18
    assert event["response_bytes"] == len(response_body)
    assert len(event["response_sha256"]) == 64
    serialized = diagnostic_path.read_text(encoding="utf-8")
    assert "private-prompt" not in serialized
    assert "private-answer" not in serialized
    assert "grading-secret" not in serialized
    assert "config-secret" not in serialized


def test_model_budget_proxy_captures_the_exact_canonical_body_before_forwarding(
    tmp_path: Path,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": 1,
            }
        ),
        encoding="utf-8",
    )
    payload = {
        "model": "shared-model",
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": "Q12 only"},
                    {
                        "type": "image_url",
                        "image_url": {"url": "data:image/png;base64,captured"},
                    },
                ],
            }
        ],
        "temperature": 0,
        "max_tokens": 32000,
    }
    expected_body = json.dumps(
        payload,
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    capture_path = tmp_path / "request-capture.json"
    diagnostic_path = tmp_path / "model-proxy.jsonl"

    async def fake_forward(
        upstream: dict[str, str],
        request_payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        assert upstream["api_key"] in {"grading-secret", "config-secret"}
        assert request_payload == payload
        assert capture_path.read_bytes() == expected_body
        return (
            200,
            {"content-type": "application/json"},
            b'{"choices":[{"finish_reason":"stop","message":{"content":"ok"}}]}',
        )

    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=diagnostic_path,
        request_capture_path=capture_path,
        forwarder=fake_forward,
    )
    response = TestClient(app).post(
        "/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json=payload,
    )

    assert response.status_code == 200
    assert capture_path.read_bytes() == expected_body
    assert b"grading-secret" not in capture_path.read_bytes()
    assert b"config-secret" not in capture_path.read_bytes()
    event = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert event["payload_bytes"] == len(expected_body)
    assert event["payload_sha256"] == hashlib.sha256(expected_body).hexdigest()


def test_model_budget_proxy_sends_the_captured_bytes_to_the_real_http_seam(
    tmp_path: Path,
) -> None:
    captured: dict[str, bytes | str] = {}

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self) -> None:
            length = int(self.headers.get("Content-Length") or "0")
            captured["path"] = self.path
            captured["authorization"] = self.headers.get("Authorization") or ""
            captured["body"] = self.rfile.read(length)
            response = (
                b'{"choices":[{"finish_reason":"stop","message":{"content":"ok"}}]}'
            )
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(response)))
            self.end_headers()
            self.wfile.write(response)

        def log_message(self, _format: str, *_args: object) -> None:
            return

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        upstream_path = tmp_path / "upstream.json"
        upstream_path.write_text(
            json.dumps(
                {
                    "grading": {
                        "base_url": f"http://127.0.0.1:{server.server_port}/v1",
                        "api_key": "grading-secret",
                        "model": "shared-model",
                    },
                    "config": {
                        "base_url": f"http://127.0.0.1:{server.server_port}/v1",
                        "api_key": "config-secret",
                        "model": "shared-model",
                    },
                    "max_forwarded_requests": 1,
                }
            ),
            encoding="utf-8",
        )
        payload = {
            "model": "shared-model",
            "temperature": 0,
            "messages": [{"role": "user", "content": "Q12 only"}],
        }
        expected_body = json.dumps(
            payload,
            ensure_ascii=True,
            sort_keys=True,
            separators=(",", ":"),
        ).encode()
        capture_path = tmp_path / "request-capture.json"
        app = acceptance.create_model_budget_proxy(
            upstream_path,
            state_path=tmp_path / "budget-state.json",
            diagnostic_log_path=tmp_path / "model-proxy.jsonl",
            request_capture_path=capture_path,
        )

        response = TestClient(app).post(
            "/v1/chat/completions",
            headers={"Authorization": "Bearer acceptance-proxy"},
            json=payload,
        )

        assert response.status_code == 200
        assert captured == {
            "path": "/v1/chat/completions",
            "authorization": "Bearer config-secret",
            "body": expected_body,
        }
        assert capture_path.read_bytes() == expected_body
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


@pytest.mark.parametrize("transport_failure", [False, True])
def test_model_budget_proxy_diagnostic_write_failure_never_changes_paid_result(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    transport_failure: bool,
) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": 1,
            }
        ),
        encoding="utf-8",
    )
    response_body = (
        b'{"choices":[{"finish_reason":"stop","message":{"content":"paid-ok"}}]}'
    )

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        if transport_failure:
            try:
                raise httpx.ReadTimeout("private-timeout")
            except httpx.ReadTimeout as exc:
                raise acceptance.AcceptanceError(
                    "approved model request failed at the provider"
                ) from exc
        return 200, {"content-type": "application/json"}, response_body

    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=tmp_path / "model-proxy.jsonl",
        forwarder=fake_forward,
    )
    monkeypatch.setattr(
        acceptance,
        "_append_model_proxy_diagnostic",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            OSError("private-log-write-detail")
        ),
    )
    response = TestClient(app).post(
        "/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={"model": "shared-model", "messages": []},
    )

    if transport_failure:
        assert response.status_code == 502
        assert "approved model request failed at the provider" in response.text
    else:
        assert response.status_code == 200
        assert response.content == response_body
    stderr = capsys.readouterr().err
    assert "model_proxy_diagnostic_write_failed" in stderr
    assert "OSError" in stderr
    assert "private-log-write-detail" not in stderr


def test_model_budget_proxy_redacts_unknown_finish_reason(tmp_path: Path) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "shared-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "shared-model",
                },
                "max_forwarded_requests": 1,
            }
        ),
        encoding="utf-8",
    )
    private_reason = "private-provider-content"

    async def fake_forward(
        _upstream: dict[str, str],
        _payload: dict[str, object],
    ) -> tuple[int, dict[str, str], bytes]:
        return (
            200,
            {"content-type": "application/json"},
            json.dumps(
                {
                    "choices": [
                        {
                            "finish_reason": private_reason,
                            "message": {"content": "private-response"},
                        }
                    ]
                }
            ).encode(),
        )

    diagnostic_path = tmp_path / "model-proxy.jsonl"
    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
        diagnostic_log_path=diagnostic_path,
        forwarder=fake_forward,
    )
    response = TestClient(app).post(
        "/v1/chat/completions",
        headers={"Authorization": "Bearer acceptance-proxy"},
        json={"model": "shared-model", "messages": []},
    )

    assert response.status_code == 200
    event = json.loads(diagnostic_path.read_text(encoding="utf-8"))
    assert event["finish_reason"] == "other"
    assert private_reason not in diagnostic_path.read_text(encoding="utf-8")
    assert "private-response" not in diagnostic_path.read_text(encoding="utf-8")


def test_model_budget_proxy_rejects_wrong_key_and_unknown_model(tmp_path: Path) -> None:
    upstream_path = tmp_path / "upstream.json"
    upstream_path.write_text(
        json.dumps(
            {
                "grading": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "grading-secret",
                    "model": "grading-model",
                },
                "config": {
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "config-secret",
                    "model": "config-model",
                },
                "max_forwarded_requests": 4,
            }
        ),
        encoding="utf-8",
    )
    app = acceptance.create_model_budget_proxy(
        upstream_path,
        state_path=tmp_path / "budget-state.json",
    )
    client = TestClient(app)

    assert (
        client.post(
            "/v1/chat/completions",
            headers={"Authorization": "Bearer wrong"},
            json={"model": "grading-model", "messages": []},
        ).status_code
        == 401
    )
    assert (
        client.post(
            "/v1/chat/completions",
            headers={"Authorization": "Bearer acceptance-proxy"},
            json={"model": "unapproved-model", "messages": []},
        ).status_code
        == 403
    )


def test_configure_staged_runtime_uses_only_workspace_relative_paths(
    tmp_path: Path,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(
        json.dumps(
            [
                {
                    "name": "active",
                    "provider": "custom-openai-compatible",
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "secret",
                    "grading_model": "shared-model",
                    "config_model": "shared-model",
                }
            ]
        ),
        encoding="utf-8",
    )
    acceptance.prepare_runtime_profile(
        workspace,
        source_profile_path=profile_path,
        proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
        max_forwarded_requests=4,
    )

    acceptance.configure_staged_runtime(workspace)

    config_text = (workspace / "config" / "app_config.yaml").read_text(
        encoding="utf-8"
    )
    assert "DATA_DIR: acceptance_data" in config_text
    assert "LOGS_DIR: acceptance_logs" in config_text
    assert "user_data" not in config_text
    assert (workspace / "acceptance_data").is_dir()
    assert (workspace / "acceptance_logs").is_dir()
    assert (workspace / "acceptance_ops").is_dir()
    assert acceptance.load_metadata(workspace)["state"] == "runtime_ready"


def test_build_and_copy_frontend_requires_exact_clean_source(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    (repo / "frontend").mkdir()
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    (workspace / "acceptance_data").mkdir()
    (workspace / "acceptance_logs").mkdir()
    calls: list[list[str]] = []
    call_options: list[dict[str, object]] = []

    def fake_run(
        command: list[str],
        **kwargs: object,
    ) -> subprocess.CompletedProcess[str]:
        calls.append(command)
        call_options.append(dict(kwargs))
        if command[:2] == ["git", "rev-parse"]:
            return subprocess.CompletedProcess(command, 0, source_sha, "")
        if command[:2] == ["git", "status"]:
            return subprocess.CompletedProcess(command, 0, "", "")
        if command[-2:] == ["run", "build"]:
            dist = repo / "frontend" / "dist"
            (dist / "assets").mkdir(parents=True)
            (dist / "index.html").write_text("<main>P2-20</main>", encoding="utf-8")
            (dist / "assets" / "app.js").write_text("ok", encoding="utf-8")
            return subprocess.CompletedProcess(command, 0, "built", "")
        raise AssertionError(command)

    monkeypatch.setattr(acceptance.subprocess, "run", fake_run)
    monkeypatch.setattr(acceptance.shutil, "which", lambda _name: "npm")

    acceptance.build_and_copy_frontend(
        workspace,
        repo_root=repo,
        expected_source_sha=source_sha,
    )

    assert (workspace / "frontend" / "dist" / "index.html").is_file()
    assert any(command[-2:] == ["run", "build"] for command in calls)
    build_index = next(
        index for index, command in enumerate(calls) if command[-2:] == ["run", "build"]
    )
    assert call_options[build_index]["encoding"] == "utf-8"
    assert call_options[build_index]["errors"] == "replace"


def test_create_acceptance_app_rejects_missing_frontend_dist(tmp_path: Path) -> None:
    repo, source_sha = _committed_source_repo(tmp_path)
    workspace = tmp_path / "prepared"
    acceptance.prepare_code_workspace(source_sha, workspace, repo_root=repo)
    source_data = _authorized_source_data(tmp_path)
    acceptance.prepare_authorized_inputs(
        workspace,
        source_data_root=source_data,
        session_id=1,
        expected_session_name="0609",
        paper_limit=3,
    )
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(
        json.dumps(
            [
                {
                    "name": "active",
                    "provider": "custom-openai-compatible",
                    "base_url": "https://provider.example.invalid/v1",
                    "api_key": "secret",
                    "grading_model": "shared-model",
                    "config_model": "shared-model",
                }
            ]
        ),
        encoding="utf-8",
    )
    acceptance.prepare_runtime_profile(
        workspace,
        source_profile_path=profile_path,
        proxy_base_url="http://127.0.0.1:8120/acceptance-llm/v1",
        max_forwarded_requests=4,
    )
    acceptance.configure_staged_runtime(workspace)

    with pytest.raises(acceptance.AcceptanceError, match="frontend dist"):
        acceptance.create_acceptance_app(workspace)


@pytest.mark.parametrize(
    ("budget_args", "expected_budget"),
    [
        ([], 4),
        (["--model-request-budget", "5"], 5),
        (["--model-request-budget", "6"], 6),
        (["--model-request-budget", "9"], 9),
    ],
)
def test_main_prepare_orchestrates_exact_authorized_scope(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    budget_args: list[str],
    expected_budget: int,
) -> None:
    workspace = tmp_path / "workspace"
    source_data = tmp_path / "source-user-data"
    profile = tmp_path / "api_profiles.json"
    source_sha = "a" * 40
    calls: list[tuple[str, object]] = []

    monkeypatch.setattr(
        acceptance,
        "prepare_code_workspace",
        lambda source_ref, target, repo_root: calls.append(
            ("code", (source_ref, Path(target), Path(repo_root)))
        )
        or {
            "source_sha": source_sha,
        },
    )
    monkeypatch.setattr(
        acceptance,
        "prepare_authorized_inputs",
        lambda target, **kwargs: calls.append(("inputs", kwargs)) or {},
    )
    monkeypatch.setattr(
        acceptance,
        "prepare_runtime_profile",
        lambda target, **kwargs: calls.append(("profile", kwargs)) or {},
    )
    monkeypatch.setattr(
        acceptance,
        "configure_staged_runtime",
        lambda target: calls.append(("config", Path(target))),
    )
    monkeypatch.setattr(
        acceptance,
        "build_and_copy_frontend",
        lambda target, **kwargs: calls.append(("build", kwargs)),
    )
    monkeypatch.setattr(
        acceptance,
        "safe_status",
        lambda target: {
            "package": "P2-20",
            "state": "runtime_ready",
            "source_sha": source_sha,
            "authorized_session_id": 1,
            "paper_count": 3,
            "model_request_budget": 4,
            "forwarded_requests": 0,
            "local_ocr_requests": 0,
        },
    )

    result = acceptance.main(
        [
            "prepare",
            "--source-ref",
            source_sha,
            "--workspace",
            str(workspace),
            "--source-data-root",
            str(source_data),
            "--session-id",
            "1",
            "--session-name",
            "0609",
            "--profile-path",
            str(profile),
            "--port",
            "8120",
            *budget_args,
        ]
    )

    assert result == 0
    assert [name for name, _details in calls] == [
        "code",
        "inputs",
        "profile",
        "config",
        "build",
    ]
    assert calls[1][1]["paper_limit"] == 3
    assert calls[2][1]["max_forwarded_requests"] == expected_budget
    assert json.loads(capsys.readouterr().out)["state"] == "runtime_ready"


def test_main_prepare_supports_manual_unbounded_count_migration(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workspace = tmp_path / "workspace"
    source_data = tmp_path / "source-user-data"
    profile = tmp_path / "api_profiles.json"
    source_sha = "a" * 40
    profile_calls: list[dict[str, object]] = []

    monkeypatch.setattr(
        acceptance,
        "prepare_code_workspace",
        lambda *_args, **_kwargs: {"source_sha": source_sha},
    )
    monkeypatch.setattr(
        acceptance,
        "prepare_authorized_inputs",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        acceptance,
        "prepare_runtime_profile",
        lambda _target, **kwargs: profile_calls.append(dict(kwargs)) or {},
    )
    monkeypatch.setattr(
        acceptance,
        "configure_staged_runtime",
        lambda _target: None,
    )
    monkeypatch.setattr(
        acceptance,
        "build_and_copy_frontend",
        lambda *_args, **_kwargs: None,
    )
    monkeypatch.setattr(
        acceptance,
        "safe_status",
        lambda _target: {
            "package": "P2-20",
            "state": "runtime_ready",
            "model_request_budget": None,
            "forwarded_requests": 9,
            "direct_diagnostic_requests": 1,
            "total_external_requests": 10,
        },
    )

    result = acceptance.main(
        [
            "prepare",
            "--source-ref",
            source_sha,
            "--workspace",
            str(workspace),
            "--source-data-root",
            str(source_data),
            "--session-id",
            "1",
            "--session-name",
            "0609",
            "--profile-path",
            str(profile),
            "--port",
            "8120",
            "--unlimited-model-requests",
            "--initial-model-request-count",
            "9",
            "--direct-diagnostic-request-count",
            "1",
        ]
    )

    assert result == 0
    assert profile_calls == [
        {
            "source_profile_path": profile,
            "proxy_base_url": "http://127.0.0.1:8120/acceptance-llm/v1",
            "max_forwarded_requests": None,
            "initial_forwarded_requests": 9,
            "direct_diagnostic_requests": 1,
        }
    ]
    assert json.loads(capsys.readouterr().out)["total_external_requests"] == 10


def test_main_start_returns_after_the_acceptance_service_is_healthy(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workspace = tmp_path / "workspace"
    calls: list[tuple[Path, int]] = []

    monkeypatch.setattr(
        acceptance,
        "start_acceptance_server",
        lambda target, port: calls.append((Path(target), port))
        or {
            "package": "P2-20",
            "server_state": "running",
            "port": port,
            "forwarded_requests": 0,
            "local_ocr_requests": 0,
        },
        raising=False,
    )

    result = acceptance.main(
        [
            "start",
            "--workspace",
            str(workspace),
            "--port",
            "8120",
        ]
    )

    assert result == 0
    assert calls == [(workspace, 8120)]
    assert json.loads(capsys.readouterr().out) == {
        "package": "P2-20",
        "server_state": "running",
        "port": 8120,
        "forwarded_requests": 0,
        "local_ocr_requests": 0,
    }


def test_main_stop_returns_after_the_acceptance_service_is_down(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    workspace = tmp_path / "workspace"
    calls: list[Path] = []

    monkeypatch.setattr(
        acceptance,
        "stop_acceptance_server",
        lambda target: calls.append(Path(target))
        or {
            "package": "P2-20",
            "server_state": "stopped",
            "port": 8120,
            "forwarded_requests": 0,
            "local_ocr_requests": 0,
        },
    )

    result = acceptance.main(
        [
            "stop",
            "--workspace",
            str(workspace),
        ]
    )

    assert result == 0
    assert calls == [workspace]
    assert json.loads(capsys.readouterr().out) == {
        "package": "P2-20",
        "server_state": "stopped",
        "port": 8120,
        "forwarded_requests": 0,
        "local_ocr_requests": 0,
    }


def test_status_explicitly_reports_that_no_acceptance_server_is_running(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "workspace"
    (workspace / "acceptance_config").mkdir(parents=True)
    (workspace / acceptance.METADATA_FILENAME).write_text(
        json.dumps(
            {
                "package": "P2-20",
                "source_sha": "a" * 40,
                "state": "runtime_ready",
                "authorized_session_id": 1,
                "paper_count": 3,
                "source_fingerprint": "b" * 64,
                "input_manifest": "acceptance_inputs/manifest.json",
                "profile_fingerprint": "c" * 64,
                "model_request_budget": 4,
            }
        ),
        encoding="utf-8",
    )

    status = acceptance.safe_status(workspace)

    assert status["server_state"] == "stopped"


def test_start_claim_is_atomic_when_two_starts_race(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    (workspace / "acceptance_config").mkdir(parents=True)
    barrier = threading.Barrier(2)

    def claim(token: str) -> str:
        barrier.wait()
        acceptance._claim_server_start(
            workspace,
            source_sha="a" * 40,
            port=8120,
            control_token=token,
        )
        return token

    results: list[str] = []
    errors: list[BaseException] = []
    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [
            executor.submit(claim, "first-token"),
            executor.submit(claim, "second-token"),
        ]
        for future in futures:
            try:
                results.append(future.result())
            except BaseException as exc:
                errors.append(exc)

    assert len(results) == 1
    assert len(errors) == 1
    assert isinstance(errors[0], acceptance.AcceptanceError)


def test_health_probe_rejects_an_unrelated_http_200_service() -> None:
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"ready")

        def log_message(self, _format: str, *_args: object) -> None:
            return

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        assert not acceptance._server_is_healthy(
            int(server.server_port),
            control_token="expected-token",
            source_sha="a" * 40,
        )
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_acceptance_health_requires_the_exact_control_token() -> None:
    app = FastAPI()
    acceptance.mount_acceptance_control_health(
        app,
        control_token="expected-token",
        source_sha="a" * 40,
    )
    client = TestClient(app)

    response = client.get(
        "/acceptance-control/health",
        headers={"X-Acceptance-Control-Token": "expected-token"},
    )

    assert response.status_code == 200
    assert response.json() == {
        "package": "P2-20",
        "source_sha": "a" * 40,
    }
    assert client.get("/acceptance-control/health").status_code == 404
    assert (
        client.get(
            "/acceptance-control/health",
            headers={"X-Acceptance-Control-Token": "wrong-token"},
        ).status_code
        == 404
    )


def test_server_does_not_create_the_app_when_state_registration_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace = tmp_path / "workspace"
    (workspace / "acceptance_config").mkdir(parents=True)
    (workspace / acceptance.METADATA_FILENAME).write_text(
        json.dumps(
            {
                "package": "P2-20",
                "source_sha": "a" * 40,
                "state": "runtime_ready",
                "authorized_session_id": 1,
                "paper_count": 3,
                "source_fingerprint": "b" * 64,
                "input_manifest": "acceptance_inputs/manifest.json",
                "profile_fingerprint": "c" * 64,
                "model_request_budget": 4,
            }
        ),
        encoding="utf-8",
    )
    acceptance._claim_server_start(
        workspace,
        source_sha="a" * 40,
        port=8120,
        control_token="expected-token",
    )
    app_created = False

    def fail_state_write(
        _workspace: Path,
        _state: dict[str, object],
    ) -> None:
        raise OSError("simulated state write failure")

    def mark_app_created(_workspace: Path, _control_token: str):
        nonlocal app_created
        app_created = True
        raise AssertionError("the app must not be created")

    monkeypatch.setattr(acceptance, "_write_server_state", fail_state_write)
    monkeypatch.setattr(acceptance, "create_acceptance_app", mark_app_created)

    with pytest.raises(OSError, match="simulated state write failure"):
        acceptance.run_acceptance_server(
            workspace,
            8120,
            "expected-token",
        )

    assert not app_created
    assert not acceptance._server_claim_path(workspace).exists()


def _detached_lifecycle_workspace(
    tmp_path: Path,
) -> tuple[Path, int]:
    workspace = tmp_path / "workspace"
    (workspace / "acceptance_config").mkdir(parents=True)
    (workspace / "acceptance_logs").mkdir()
    (workspace / "tools").mkdir()
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
        probe.bind(("127.0.0.1", 0))
        port = int(probe.getsockname()[1])
    (workspace / acceptance.METADATA_FILENAME).write_text(
        json.dumps(
            {
                "package": "P2-20",
                "source_sha": "a" * 40,
                "state": "runtime_ready",
                "authorized_session_id": 1,
                "paper_count": 3,
                "source_fingerprint": "b" * 64,
                "input_manifest": "acceptance_inputs/manifest.json",
                "profile_fingerprint": "c" * 64,
                "model_request_budget": 4,
            }
        ),
        encoding="utf-8",
    )
    (workspace / "acceptance_config" / "api_profiles.json").write_text(
        json.dumps(
            [
                {
                    "base_url": (
                        f"http://127.0.0.1:{port}/acceptance-llm/v1"
                    )
                }
            ]
        ),
        encoding="utf-8",
    )
    (workspace / "tools" / "p2_20_acceptance.py").write_text(
        textwrap.dedent(
            """
            import argparse
            import json
            import os
            from http.server import BaseHTTPRequestHandler, HTTPServer
            from pathlib import Path

            parser = argparse.ArgumentParser()
            parser.add_argument("command")
            parser.add_argument("--workspace", required=True)
            parser.add_argument("--port", type=int, required=True)
            parser.add_argument("--control-token", required=True)
            args = parser.parse_args()
            stop_path = Path(args.workspace) / "acceptance_config" / "server-stop.request"
            unhealthy_path = Path(args.workspace) / "acceptance_config" / "server-unhealthy.request"
            state_path = Path(args.workspace) / "acceptance_config" / "server-state.json"
            state_path.write_text(
                json.dumps(
                    {
                        "package": "P2-20",
                        "source_sha": "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa",
                        "state": "starting",
                        "pid": os.getpid(),
                        "port": args.port,
                        "control_token": args.control_token,
                    }
                ),
                encoding="utf-8",
            )

            class Handler(BaseHTTPRequestHandler):
                def do_GET(self):
                    if (
                        self.path != "/acceptance-control/health"
                        or self.headers.get("X-Acceptance-Control-Token")
                        != args.control_token
                        or unhealthy_path.exists()
                    ):
                        self.send_response(503 if unhealthy_path.exists() else 404)
                        self.end_headers()
                        return
                    payload = json.dumps(
                        {
                            "package": "P2-20",
                            "source_sha": "aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa",
                        }
                    ).encode("utf-8")
                    self.send_response(200)
                    self.send_header("Content-Type", "application/json")
                    self.send_header("Content-Length", str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)

                def log_message(self, _format, *_args):
                    return

            server = HTTPServer(("127.0.0.1", args.port), Handler)
            server.timeout = 0.1
            try:
                while not stop_path.exists():
                    server.handle_request()
            finally:
                server.server_close()
            """
        ),
        encoding="utf-8",
    )
    return workspace, port


def test_detached_service_reports_unhealthy_after_losing_health(
    tmp_path: Path,
) -> None:
    workspace, port = _detached_lifecycle_workspace(tmp_path)

    started_at = time.monotonic()
    started = acceptance.start_acceptance_server(
        workspace,
        port,
        ready_timeout=5,
    )
    elapsed = time.monotonic() - started_at
    try:
        assert elapsed < 5
        assert started["server_state"] == "running"
        assert acceptance.safe_status(workspace)["server_state"] == "running"
        (workspace / "acceptance_config" / "server-unhealthy.request").touch()
        deadline = time.monotonic() + 2
        while time.monotonic() < deadline:
            if acceptance.safe_status(workspace)["server_state"] == "unhealthy":
                break
            time.sleep(0.1)
        assert acceptance.safe_status(workspace)["server_state"] == "unhealthy"
    finally:
        stopped = acceptance.stop_acceptance_server(
            workspace,
            stop_timeout=5,
        )

    assert stopped["server_state"] == "stopped"
    assert acceptance.safe_status(workspace)["server_state"] == "stopped"
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
        assert probe.connect_ex(("127.0.0.1", port)) != 0


def test_start_stops_the_service_when_running_state_cannot_be_saved(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, port = _detached_lifecycle_workspace(tmp_path)
    original_write = acceptance._write_server_state

    def fail_running_state(
        target: Path,
        state: dict[str, object],
    ) -> None:
        if state["state"] == "running":
            raise OSError("simulated running state write failure")
        original_write(target, state)

    monkeypatch.setattr(acceptance, "_write_server_state", fail_running_state)

    with pytest.raises(
        acceptance.AcceptanceError,
        match="running state could not be saved",
    ):
        acceptance.start_acceptance_server(
            workspace,
            port,
            ready_timeout=5,
        )

    assert not acceptance._server_state_path(workspace).exists()
    assert not acceptance._server_claim_path(workspace).exists()
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
        assert probe.connect_ex(("127.0.0.1", port)) != 0


def test_second_start_cannot_succeed_before_running_state_is_saved(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace, port = _detached_lifecycle_workspace(tmp_path)
    running_write_reached = threading.Event()
    allow_running_write = threading.Event()
    original_write = acceptance._write_server_state
    first_result: dict[str, object] | None = None

    def pause_running_state(
        target: Path,
        state: dict[str, object],
    ) -> None:
        if state["state"] == "running":
            running_write_reached.set()
            if not allow_running_write.wait(timeout=5):
                raise OSError("timed out waiting to save running state")
        original_write(target, state)

    monkeypatch.setattr(acceptance, "_write_server_state", pause_running_state)
    executor = ThreadPoolExecutor(max_workers=1)
    future = executor.submit(
        acceptance.start_acceptance_server,
        workspace,
        port,
        ready_timeout=5,
    )
    try:
        assert running_write_reached.wait(timeout=5)
        assert acceptance.safe_status(workspace)["server_state"] == "starting"
        with pytest.raises(
            acceptance.AcceptanceError,
            match="start is still in progress",
        ):
            acceptance.start_acceptance_server(
                workspace,
                port,
                ready_timeout=5,
            )
    finally:
        allow_running_write.set()
        try:
            first_result = future.result(timeout=10)
        finally:
            executor.shutdown(wait=True)
            if acceptance._load_server_state(workspace) is not None:
                acceptance.stop_acceptance_server(workspace, stop_timeout=5)

    assert first_result is not None
    assert first_result["server_state"] == "running"
