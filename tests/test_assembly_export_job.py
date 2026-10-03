from __future__ import annotations

from pathlib import Path

import pytest

from backend.jobs.assembly_export import run_assembly_export_job
from backend.jobs.default_handlers import register_default_job_handlers
from backend.jobs.manager import JobManager
from backend.jobs.store import JobStore
from question_bank.models.question import QuestionCreate
from question_bank.services.assembly_workspace_service import AssemblyWorkspaceService
from tests.question_bank_support import QuestionBankTestStore


def test_import_keeps_source_and_only_exports_hide_score_with_original_space(tmp_path):
    import json
    from copy import deepcopy
    from zipfile import ZipFile
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from PIL import Image
    from question_bank.database.schema import connect
    from question_bank.document_pipeline.legacy_exports import published_math_metadata
    from question_bank.exporters.paper_docx_exporter import export_question_paper_docx
    from question_bank.exporters.paper_pdf_exporter import _blocks, _render_source
    from question_bank.importers.batch_importer import import_scanned_papers, ScannedPaper
    from question_bank.personalized_papers.latex_render import render_training_tex
    from question_bank.personalized_papers.rendering import inspect_docx, render_review_docx
    from question_bank.services.rich_content_service import load_question_rich_content
    from question_bank.training_criteria.adapters import QuestionAnalysisInputLoader

    root = tmp_path / 'TEST-score-display'
    bank = root / 'databases/question_bank.db'
    image = tmp_path / 'TEST-figure.png'
    Image.new('RGB', (40, 20), 'navy').save(image)
    source = tmp_path / 'TEST-source-score.docx'
    document = Document()
    stem = document.add_paragraph()
    stem.add_run('1. （').bold = True
    stem.add_run('8').italic = True
    stem.add_run(' 分）小明跑了5分钟，已知 ')
    stem._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>x+1</m:t></m:r></m:oMath>'))
    stem.add_run(' 的值为____。')
    stem.add_run().add_picture(str(image))
    document.add_paragraph('参考答案')
    document.add_paragraph('1. （8分）答案为2，游戏答对可得3分。')
    document.save(source)
    source_bytes, image_bytes = source.read_bytes(), image.read_bytes()
    result = import_scanned_papers([ScannedPaper(source_file=str(source), file_type='docx')], bank,
                                   data_root=root, archive_sources=False)
    assert result.failed_files == 0 and result.question_count == 1
    with connect(bank) as conn:
        row = dict(conn.execute('SELECT * FROM questions').fetchone())
    assert row['question_text'].startswith('（8 分）') and row['question_type'] == '解答题'
    assert row['answer_text'].startswith('（8分）')
    rich_path = root / 'question_bank/rich_content' / f"question_{row['id']}.json"
    rich = load_question_rich_content(row['id'], root=rich_path.parent)
    assert '8 分' in rich['question_blocks'][0]['text']
    assert 'answer_space_lines' not in rich['question_blocks'][0]
    before = bank.read_bytes(), rich_path.read_bytes(), deepcopy(rich)
    loader = QuestionAnalysisInputLoader(db_path=bank, data_root=root)
    (analysis_input,) = loader.load([row['id']])

    word = export_question_paper_docx(bank, [row['id']], root / 'word', title='TEST去分值', include_answer=True)
    exported = Document(word)
    word_text = '\n'.join(p.text for p in exported.paragraphs)
    assert '8 分' not in word_text.split('答案', 1)[0]
    assert '（8分）答案为2' in word_text and '5分钟' in word_text
    assert len(exported.tables[0].rows) == 8
    with ZipFile(word) as archive:
        xml = archive.read('word/document.xml').decode('utf-8')
        assert '<m:t>x+1</m:t>' in xml and '<a:blip ' in xml
        assert '<w:b' in xml and '<w:i' in xml

    row['image_paths'] = json.loads(row['image_paths'])
    metadata = published_math_metadata(row['id'], data_root=root, question_text=row['question_text'],
                                       answer_text=row['answer_text'])
    blocks = _blocks(row, metadata, 'question_blocks', root)
    tex = _render_source([('', [(1, row)])], {row['id']: (blocks, [])}, data_root=root,
                         title='TEST分值显示', header_text='', include_answer=False, include_answer_space=True,
                         include_student_fields=False, choices={}, footer='', probes=False)
    assert '8 分' not in tex and r'\LayoutAnswerSpace{1}{8}' in tex

    snapshot = {
        'paper_instance_id': 'a' * 64, 'series_version': 1,
        'student': {'student_id': 'TEST-1', 'student_name': '测试学生', 'class_id': '测试班'},
        'items': [{'task_item_code': 'TEST-TASK-1', 'question_id': row['id'], 'question_snapshot': {
            'question_id': row['id'], 'tagging_context': analysis_input.tagging_context.to_dict(),
            'rich_question_blocks': rich['question_blocks'], 'images': []}}],
    }
    snapshot_before = deepcopy(snapshot)
    training_word = root / 'TEST-training.docx'
    render_review_docx(snapshot, data_root=root, output_path=training_word)
    trained = Document(training_word)
    assert len(trained.tables[0].rows) == 8
    assert '8 分' not in '\n'.join(p.text for p in trained.paragraphs)
    inspect_docx(training_word, paper_instance_id=snapshot['paper_instance_id'], task_item_codes=['TEST-TASK-1'],
                 question_snapshots=[snapshot['items'][0]['question_snapshot']])
    training_tex = render_training_tex(snapshot, data_root=root)
    assert '8 分' not in training_tex and ('36.0mm' in training_tex or '72mm' in training_tex)
    assert snapshot == snapshot_before
    assert (bank.read_bytes(), rich_path.read_bytes(), rich) == before
    assert loader.load([row['id']]) == (analysis_input,)
    assert (source.read_bytes(), image.read_bytes()) == (source_bytes, image_bytes)


@pytest.mark.parametrize("export_format", ["markdown", "pdf"])
def test_assembly_export_job_publishes_records_and_clears_same_draft(
    tmp_path: Path,
    export_format: str,
) -> None:
    data_root = tmp_path / "data"
    db_path = data_root / "databases" / "question_bank.db"
    question_service = QuestionBankTestStore(db_path)
    question_id = question_service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="选择题",
            question_text="（5分）匿名题目",
            answer_text="A",
        )
    )
    workspace = AssemblyWorkspaceService(data_root)
    empty = workspace.load_draft()
    draft = workspace.save_draft(
        expected_revision=empty.revision,
        draft={
            "basket_ids": [question_id],
            "order_ids": [question_id],
            "sections": [],
            "title": "匿名练习",
            "header_text": "",
            "include_answer": True,
            "layout_mode": "sequential",
            "preview_mode": "student",
        },
    )

    def pdf_exporter(_db, _ids, output_dir, **kwargs):
        import fitz

        output = output_dir / "test-paper.pdf"
        with fitz.open() as pdf:
            pdf.new_page().insert_text((30, 30), "TEST PDF")
            pdf.save(output)
        return output

    def runner(**kwargs):
        return run_assembly_export_job(**kwargs, pdf_exporter=pdf_exporter)

    manager = JobManager(JobStore(tmp_path / "jobs.db"), max_workers=1)
    try:
        register_default_job_handlers(
            manager,
            db_path=tmp_path / "grading.db",
            reports_dir=tmp_path / "reports",
            data_root=data_root,
            question_bank_db_path=db_path,
            assembly_export_runner=runner,
        )
        job = manager.submit(
            "assembly_export",
            {
                "draft_revision": draft.revision,
                "draft": draft.to_payload(),
                "format": export_format,
            },
        )
        manager.wait(job.id, timeout=5)

        loaded = manager.get(job.id)
        assert loaded is not None
        assert loaded.status == "succeeded"
        output = Path(loaded.result["file_path"])
        assert output.is_file()
        assert output.parent == workspace.exports_root
        assert loaded.result["filename"] == output.name
        assert loaded.result["record_id"]
        assert loaded.result["question_count"] == 1
        assert loaded.result["draft_cleared"] is True
        assert workspace.load_draft().basket_ids == ()
        assert workspace.list_records()[0].id == loaded.result["record_id"]
        assert workspace.list_records()[0].export_format == export_format
        resolved = workspace.resolve_record_file(loaded.result["record_id"])
        assert resolved.media_type == (
            "application/pdf" if export_format == "pdf" else "text/markdown"
        )
        from fastapi.testclient import TestClient

        from backend.api.app import create_app
        from backend.api.dependencies import (
            get_assembly_workspace_service,
            get_job_file_service,
            get_job_manager,
            get_question_bank_db_path,
            get_request_diagnosis_profile_service,
        )
        from backend.files.service import JobFileService

        app = create_app()
        app.dependency_overrides[get_job_manager] = lambda: manager
        app.dependency_overrides[get_job_file_service] = lambda: JobFileService(
            tmp_path / "reports", assembly_outputs_dir=workspace.exports_root
        )
        app.dependency_overrides[get_assembly_workspace_service] = lambda: workspace
        app.dependency_overrides[get_question_bank_db_path] = lambda: db_path
        app.dependency_overrides[get_request_diagnosis_profile_service] = lambda: None
        client = TestClient(app)
        public_job = client.get(f"/api/jobs/{job.id}").json()
        assert "file_path" not in public_job["result"]
        records = client.get("/api/question-assembly/records").json()
        assert records["items"][0]["export_format"] == export_format
        for url in (
            public_job["result"]["download_url"],
            records["items"][0]["download_url"],
        ):
            download = client.get(url)
            assert download.status_code == 200
            assert download.headers["content-type"].startswith(resolved.media_type)
            assert download.content == output.read_bytes()
        # Submit through the actual API schema; the frozen draft stays identical.
        if export_format == "pdf":
            refreshed = workspace.load_draft()
            saved = workspace.save_draft(
                expected_revision=refreshed.revision, draft=draft.to_payload()
            )
            submitted = client.post(
                "/api/question-assembly/export",
                json={"draft_revision": saved.revision, "format": "pdf"},
            )
            assert submitted.status_code == 202
            manager.wait(submitted.json()["id"], timeout=5)
            assert manager.get(submitted.json()["id"]).status == "succeeded"
            assert output.is_file()

        assert not list(workspace.exports_root.glob(".job-*"))
    finally:
        manager.shutdown()


def test_pdf_export_compiles_rich_tables_large_figures_and_plain_math(tmp_path: Path):
    import json
    from zipfile import ZipFile

    import fitz
    from docx import Document
    from docx.shared import Inches
    from PIL import Image, ImageDraw

    from question_bank.exporters.paper_docx_exporter import export_question_paper_docx
    from question_bank.exporters.paper_pdf_exporter import export_question_paper_pdf
    from question_bank.personalized_papers.latex_render import TectonicCompiler
    from question_bank.services.assembly_basket_state import SectionSpec

    if not TectonicCompiler().available:
        pytest.skip("Local LaTeX engine is unavailable")
    root = tmp_path / "TEST-pdf"
    db = root / "databases" / "question_bank.db"
    bank = QuestionBankTestStore(db)
    choice = bank.add_question(
        QuestionCreate(
            question_number="7",
            question_type="选择题",
            question_text="（3分）测试选择题 求 $x^2+1$。\nA. 1 B. 2 C. 3 D. 4",
            answer_text="测试选择解析 B",
        )
    )
    solution = bank.add_question(
        QuestionCreate(
            question_number="28",
            question_type="解答题",
            question_text="（8分）测试大图与跨页表格",
            answer_text="（8分）测试解答解析 $x=2$。",
            image_paths=[str(root / "question_bank" / "assets" / "TEST-copy.png")],
        )
    )
    picture = root / "question_bank" / "assets" / "TEST-large.png"
    picture.parent.mkdir(parents=True)
    image = Image.new("RGB", (1400, 700), "white")
    pen = ImageDraw.Draw(image)
    pen.rectangle((30, 30, 1370, 670), outline="black", width=4)
    pen.text((80, 60), "READABLE A B C", fill="black")
    image.save(picture)
    image.save(picture.with_name("TEST-copy.png"), compress_level=0)
    document = Document()
    document.add_paragraph("28. （8分）测试大图与跨页表格，图中要素应保留。")
    document.add_paragraph().add_run().add_picture(str(picture), width=Inches(5.5))
    table = document.add_table(rows=42, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "测试合并单元格"
    for index, row in enumerate(table.rows):
        if index > 2:
            row.cells[0].text = f"项目 {index}"
        row.cells[1].text = f"测试表格第 {index} 行，内容必须完整保留。"
    wide = document.add_table(rows=2, cols=12)
    for column in range(12):
        wide.cell(0, column).text = str(column + 1)
        wide.cell(1, column).text = "值"
    wide.cell(1, 0).paragraphs[0].add_run().add_break()
    wide.cell(1, 0).paragraphs[0].add_run("换行")
    document.add_paragraph("测试末尾小问：依据表格说明理由。")
    rels = {
        rel_id: picture.relative_to(root).as_posix()
        for rel_id, rel in document.part.rels.items()
        if rel.reltype.endswith("/image")
    }
    blocks = [
        {"xml": node.xml, "image_relationships": rels}
        for node in document.element.body
        if node.tag.rsplit("}", 1)[-1] in {"p", "tbl"}
    ]
    rich = root / "question_bank" / "rich_content" / f"question_{solution}.json"
    rich.parent.mkdir(parents=True)
    rich.write_text(
        json.dumps(
            {"version": 3, "question_id": solution, "question_blocks": blocks},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    picture_choice = bank.add_question(
        QuestionCreate(
            question_number="8",
            question_type="选择题",
            question_text="测试图片选项，保留原图尺寸。",
            answer_text="测试图片选项解析 A",
            image_paths=[str(picture.with_name("TEST-answer-only.png"))],
        )
    )
    options = Document()
    options.add_paragraph("8. 测试图片选项，保留原图尺寸。")
    assets = {}
    for index, letter in enumerate("ABCD"):
        option_path = picture.with_name(f"TEST-option-{letter}.png")
        option_image = Image.new("RGB", (500, 500), "white")
        ImageDraw.Draw(option_image).rectangle(
            (20 + index * 10, 20, 480, 480), outline="black", width=4
        )
        option_image.save(option_path)
        paragraph = options.add_paragraph(f"{letter}. ")
        # Cover both labels beside an image and labels in a preceding paragraph.
        if index % 2:
            paragraph = options.add_paragraph()
        run = paragraph.add_run()
        shape = run.add_picture(str(option_path), width=Inches(30 / 25.4))
        assets[shape._inline.graphic.graphicData.pic.blipFill.blip.embed] = (
            option_path.relative_to(root).as_posix()
        )
    answer_picture = picture.with_name("TEST-answer-only.png")
    Image.new("RGB", (400, 300), "lightgray").save(answer_picture)
    answer_document = Document()
    answer_document.add_paragraph("测试图片选项解析 A，仅解析可显示下图。")
    answer_shape = (
        answer_document.add_paragraph()
        .add_run()
        .add_picture(str(answer_picture), width=Inches(1.5))
    )
    answer_assets = {
        answer_shape._inline.graphic.graphicData.pic.blipFill.blip.embed: answer_picture.relative_to(
            root
        ).as_posix()
    }
    (rich.parent / f"question_{picture_choice}.json").write_text(
        json.dumps(
            {
                "version": 3,
                "question_id": picture_choice,
                "question_blocks": [
                    {"xml": node.xml, "image_relationships": assets}
                    for node in options.element.body
                    if node.tag.rsplit("}", 1)[-1] == "p"
                ],
                "answer_blocks": [
                    {"xml": node.xml, "image_relationships": answer_assets}
                    for node in answer_document.element.body
                    if node.tag.rsplit("}", 1)[-1] == "p"
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    ids = [choice, picture_choice, solution]
    output = export_question_paper_pdf(
        db,
        ids,
        root / "exports",
        title="TEST 打印验证",
        include_answer=True,
        sections=[
            SectionSpec(title="基础", question_ids=[choice, picture_choice]),
            SectionSpec(title="应用", question_ids=[solution]),
        ],
    )
    with fitz.open(output) as pdf:
        text = "".join(page.get_text() for page in pdf)
        assert "测试选择题" in text and "测试末尾小问" in text
        student_text = "".join(text.split("答案解析", 1)[0].split())
        assert "3分" not in student_text and "8分" not in student_text
        assert "8分" in "".join(text.split("答案解析", 1)[1].split())
        assert "测试表格第41行" in "".join(text.split()) and "测试合并单元格" in text
        assert (
            text.index("测试末尾小问")
            < text.index("答案解析")
            < text.index("测试解答解析")
        )
        assert "表格续页" in text and "续题" in text and "换行" in text
        images = [info for page in pdf for info in page.get_image_info()]
        assert len(images) == 6
        answer_page = next(
            index for index, page in enumerate(pdf) if "答案解析" in page.get_text()
        )
        assert (
            sum(
                info["width"] == 400
                for page in list(pdf)[:answer_page]
                for info in page.get_image_info()
            )
            == 0
        )
        assert (
            sum(
                info["width"] == 400
                for page in list(pdf)[answer_page:]
                for info in page.get_image_info()
            )
            == 1
        )
        large_image = next(info for info in images if info["width"] == 1400)
        assert large_image["height"] == 700
        width = (large_image["bbox"][2] - large_image["bbox"][0]) * 25.4 / 72
        assert width == pytest.approx(139.7, abs=0.1)
        option_images = [info for info in images if info["width"] == 500]
        assert len(option_images) == 4
        assert (
            max(info["bbox"][1] for info in option_images)
            - min(info["bbox"][1] for info in option_images)
            < 0.1
        )
        for info in option_images:
            assert (info["bbox"][2] - info["bbox"][0]) * 25.4 / 72 == pytest.approx(
                30, abs=0.1
            )
        for page in pdf:
            assert page.rect.width == pytest.approx(595.28, abs=0.1)
            for word in page.get_text("words"):
                assert 0 <= word[0] < word[2] <= page.rect.width
                assert 0 <= word[1] < word[3] <= page.rect.height
    word_output = export_question_paper_docx(
        db, ids, root / "word", title="TEST 同源图片核对", include_answer=False
    )
    with ZipFile(word_output) as archive:
        assert archive.read("word/document.xml").count(b"<a:blip ") == 5
    assert not list((root / "exports").glob(".latex-*"))


def test_pdf_export_failure_hides_compiler_details_and_keeps_no_partial_file(
    tmp_path: Path,
):
    from question_bank.exporters.paper_pdf_exporter import export_question_paper_pdf
    from question_bank.personalized_papers.latex_render import LatexRenderError

    root = tmp_path / "TEST-pdf-failure"
    db = root / "databases" / "question_bank.db"
    question_id = QuestionBankTestStore(db).add_question(
        QuestionCreate(
            question_number="1", question_type="选择题", question_text="测试失败题"
        )
    )

    class BrokenCompiler:
        available = True

        def compile(self, source, destination, **kwargs):
            destination.write_bytes(b"partial")
            raise RuntimeError("private question text and asset path")

    with pytest.raises(LatexRenderError) as failure:
        export_question_paper_pdf(
            db, [question_id], root / "exports", title="TEST", compiler=BrokenCompiler()
        )
    assert "private" not in str(failure.value)
    assert "Word" in str(failure.value)
    assert list((root / "exports").iterdir()) == []


def test_word_export_compact_book_preserves_math_images_and_final_answers(
    tmp_path: Path,
):
    from zipfile import ZipFile

    from docx import Document
    from PIL import Image, ImageDraw

    from question_bank.exporters.paper_docx_exporter import export_question_paper_docx
    from question_bank.services.assembly_basket_state import SectionSpec

    db_path = tmp_path / "data" / "databases" / "question_bank.db"
    bank = QuestionBankTestStore(db_path)
    image = tmp_path / "data" / "question_bank" / "assets" / "test_figure.png"
    image.parent.mkdir(parents=True, exist_ok=True)
    figure = Image.new("RGB", (180, 100), "white")
    drawing = ImageDraw.Draw(figure)
    drawing.line([(25, 80), (90, 15), (155, 80), (25, 80)], fill="black", width=2)
    for point, label in [((15, 80), "A"), ((88, 2), "B"), ((158, 80), "C")]:
        drawing.text(point, label, fill="black")
    figure.save(image)
    first = bank.add_question(
        QuestionCreate(
            question_number="1",
            question_type="选择题",
            question_text="测试选择题 A. 1 B. 2 C. 3 D. 4",
            answer_text="测试选择答案 A",
        )
    )
    second = bank.add_question(
        QuestionCreate(
            question_number="10",
            question_type="解答题",
            question_text="测试几何原题 求 $x^2+1$，图见下方。",
            answer_text="测试几何解析 代入 $x=2$ 得到 $5$。",
            image_paths=[str(image)],
        )
    )
    sections = [
        SectionSpec(title="早期考试", question_ids=[first]),
        SectionSpec(title="后期考试", question_ids=[second]),
    ]
    compact = export_question_paper_docx(
        db_path,
        [first, second],
        tmp_path / "compact",
        title="测试学生 错题本",
        include_answer=True,
        sections=sections,
        include_answer_space=False,
        include_student_fields=False,
        page_header_text="八年级上学期 · 早期考试、后期考试",
    )
    default = export_question_paper_docx(
        db_path,
        [first, second],
        tmp_path / "default",
        title="默认练习",
        include_answer=True,
        sections=sections,
    )
    doc = Document(compact)
    assert doc.sections[0].page_width.mm == pytest.approx(210, abs=0.1)
    assert doc.sections[0].page_height.mm == pytest.approx(297, abs=0.1)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.index("测试几何原题") < text.index("答案") < text.index("测试几何解析")
    assert "姓名：" not in text
    assert "姓名：" in "\n".join(p.text for p in Document(default).paragraphs)
    assert len(doc.tables) < len(Document(default).tables)
    assert "八年级上学期" in doc.sections[0].header.paragraphs[0].text
    with ZipFile(compact) as archive:
        assert b"<m:oMath" in archive.read("word/document.xml")
        assert any(name.startswith("word/media/") for name in archive.namelist())


@pytest.mark.parametrize("export_format", ["markdown", "pdf"])
def test_assembly_export_cancel_after_generation_does_not_publish_or_clear_draft(
    tmp_path: Path,
    export_format: str,
) -> None:
    data_root = tmp_path / "data"
    db_path = data_root / "databases" / "question_bank.db"
    question_service = QuestionBankTestStore(db_path)
    question_id = question_service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="选择题",
            question_text="（5分）匿名题目",
            answer_text="A",
        )
    )
    workspace = AssemblyWorkspaceService(data_root)
    empty = workspace.load_draft()
    draft = workspace.save_draft(
        expected_revision=empty.revision,
        draft={
            "basket_ids": [question_id],
            "order_ids": [question_id],
            "sections": [],
            "title": "取消验证",
            "header_text": "",
            "include_answer": True,
            "layout_mode": "sequential",
            "preview_mode": "student",
        },
    )

    def runner(**kwargs: object) -> dict[str, object]:
        context = kwargs["context"]

        def cancelling_markdown_exporter(
            _db_path: Path,
            _question_ids: list[int],
            output_dir: Path,
            **_export_kwargs: object,
        ) -> str:
            output = output_dir / (
                "cancelled.pdf" if export_format == "pdf" else "cancelled.md"
            )
            output.write_text("# cancelled", encoding="utf-8")
            context.store.request_cancel(context.job_id)
            return str(output)

        return run_assembly_export_job(
            **kwargs,
            markdown_exporter=cancelling_markdown_exporter,
            pdf_exporter=cancelling_markdown_exporter,
        )

    manager = JobManager(JobStore(tmp_path / "jobs.db"), max_workers=1)
    try:
        register_default_job_handlers(
            manager,
            db_path=tmp_path / "grading.db",
            reports_dir=tmp_path / "reports",
            data_root=data_root,
            question_bank_db_path=db_path,
            assembly_export_runner=runner,
        )
        job = manager.submit(
            "assembly_export",
            {
                "draft_revision": draft.revision,
                "draft": draft.to_payload(),
                "format": export_format,
            },
        )
        manager.wait(job.id, timeout=5)

        loaded = manager.get(job.id)
        assert loaded is not None
        assert loaded.status == "cancelled"
        assert not list(workspace.exports_root.glob("*_job-*"))
        assert workspace.list_records() == []
        assert workspace.load_draft().basket_ids == (question_id,)
        assert not list(workspace.exports_root.glob(".job-*"))
    finally:
        manager.shutdown()
