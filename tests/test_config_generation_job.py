from __future__ import annotations

import asyncio
import io
import json
from pathlib import Path
from typing import Any

import fitz
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.config_workspace.sources import (
    ConfigSourceRecord,
    ConfigSourceService,
)
from backend.jobs.config_generation import (
    _run_config_generation_job_impl,
    load_config_generation_input,
    preserve_interrupted_config_generation_checkpoints,
    run_config_generation_job,
    stage_config_source_generation_input,
)
from backend.jobs.manager import JobContext
from backend.jobs.store import JobStore
from db_manager import DBManager
from backend.config_workspace.publish import save_generated_config
from backend.repositories.grading_database import open_grading_repositories


def _valid_config_payload() -> dict[str, object]:
    scores = [17, 17, 17, 17, 17, 15]
    rubric_questions = []
    answer_questions = []
    for index, score in enumerate(scores, start=1):
        question_id = f"Q{index}"
        part_id = f"{question_id}-P1"
        rubric_questions.append(
            {
                "question_id": question_id,
                "question_type": "comprehensive",
                "max_score": score,
                "knowledge_id": f"K{index}",
                "parts": [
                    {
                        "part_id": part_id,
                        "part_score": score,
                        "response_mode": "short_answer_points",
                        "steps": [
                            {
                                "step_id": f"{part_id}-S1",
                                "step_score": score,
                                "core_goal": "answer",
                                "required_elements": [str(index)],
                                "allow_alternative_methods": True,
                            }
                        ],
                    }
                ],
            }
        )
        answer_questions.append(
            {
                "question_id": question_id,
                "canonical_answer": str(index),
                "accepted_forms": [str(index)],
                "method_variants": [],
                "parts": [
                    {
                        "part_id": part_id,
                        "answer": str(index),
                        "analysis": "",
                        "step_milestones": [str(index)],
                    }
                ],
            }
        )
    return {
        "rubric": {
            "exam_title": "Atomic Exam",
            "total_score": 100,
            "questions": rubric_questions,
        },
        "answer_key": {"questions": answer_questions},
        "meta": {"warnings": []},
    }


async def _chunks(content: bytes):
    yield content


def _docx_bytes(text: str = "1. Prove x equals x.\nAnswer: proven") -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Synthetic exam\n1. Compute 1 + 1.\nAnswer\n1. 2",
        fontsize=12,
    )
    content = document.tobytes()
    document.close()
    return content


def _controlled_source(
    tmp_path: Path,
    session_id: int,
    *,
    suffix: str = ".docx",
    text: str = "1. Prove x equals x.\nAnswer: proven",
) -> tuple[ConfigSourceService, ConfigSourceRecord]:
    service = ConfigSourceService(tmp_path / "uploaded")
    content = _pdf_bytes() if suffix == ".pdf" else _docx_bytes(text)
    record = asyncio.run(
        service.stage_and_parse(
            session_id=session_id,
            filename=f"七年级下册测试卷{suffix}",
            chunks=_chunks(content),
        )
    )
    return service, record


def _stage_controlled_input(
    tmp_path: Path,
    source_service: ConfigSourceService,
    source: ConfigSourceRecord,
    expected_paths: tuple[str, str],
    *,
    generation_mode: str,
    sync_to_question_bank: bool = False,
) -> str:
    return stage_config_source_generation_input(
        tmp_path / "uploaded",
        session_id=source.session_id,
        expected_rubric_path=expected_paths[0],
        expected_answer_key_path=expected_paths[1],
        generation_mode=generation_mode,
        source_id=source.source_id,
        source_revision=source.source_revision,
        sync_to_question_bank=sync_to_question_bank,
        curriculum_volume_id=("bnu24-math-g7-upper" if sync_to_question_bank else None),
        decisions=[
            {
                "question_id": source.questions[0].question_id,
                "question_type": "proof",
                "excluded": False,
            }
        ],
    )


def _job_context(
    job_db_path: Path,
    payload: dict[str, object],
) -> tuple[JobContext, JobStore]:
    store = JobStore(job_db_path)
    job = store.create_job("config_generation", payload)
    assert store.mark_running(job.id)
    return (
        JobContext(
            job_id=job.id,
            job_type=job.job_type,
            payload=job.payload,
            store=store,
        ),
        store,
    )


def _db_with_session(tmp_path: Path) -> tuple[DBManager, int, tuple[str, str]]:
    databases_dir = tmp_path / "databases"
    databases_dir.mkdir(parents=True, exist_ok=True)
    db = open_grading_repositories(databases_dir / "grading.db")
    db.initialize()
    initial_dir = tmp_path / "initial"
    rubric_path, answer_path = save_generated_config(
        initial_dir,
        _valid_config_payload(),
        "initial",
    )
    session_id = db.sessions.create_grading_session(
        "Config Job Exam",
        str(rubric_path),
        str(answer_path),
    )
    return db, session_id, (str(rubric_path), str(answer_path))


def _deferred_taxonomy_contract() -> dict[str, Any]:
    return {
        "taxonomy_revision": 2,
        "allowed_dimensions": ["knowledge", "ability", "curriculum"],
        "candidates": {
            "knowledge": [
                {
                    "id": "kp_alg_linear_equation",
                    "name": "一元一次方程",
                    "aliases": ["一次方程"],
                }
            ],
            "ability": [
                {
                    "id": "ability-calculation",
                    "name": "运算能力",
                    "aliases": [],
                }
            ],
        },
        "curriculum_volume": {
            "id": "bnu24-math-g7-upper",
            "sections": [
                {
                    "id": "synthetic-linear-equation-section",
                    "chapter_name": "一元一次方程",
                }
            ],
        },
    }


def _deferred_combined_item(question_id: int) -> dict[str, Any]:
    return {
        "question_id": question_id,
        "tag_analysis": {
            "knowledge_points": ["一元一次方程"],
            "method_tags": [],
            "ability_tags": ["运算能力"],
            "math_model_tags": [],
            "special_type_tags": [],
            "difficulty": 3,
            "error_prone_points": ["运算化简错误"],
            "textbook_chapters": [],
            "curriculum_sections": ["synthetic-linear-equation-section"],
            "suitable_student_level": "",
            "canonical_knowledge_id": "kp_alg_linear_equation",
            "taxonomy_revision": 2,
            "proposed_tags": [],
            "reason": "合成分析。",
            "confidence": 0.95,
        },
        "solution_evidence": {
            "schema_version": "question-solution-evidence-v2",
            "question_id": question_id,
            "parts": [
                {
                    "part_id": f"part-{question_id}",
                    "label": f"第{question_id}题",
                    "response_mode": "process_required",
                    "canonical_answer": "x=1",
                    "accepted_forms": ["x=1"],
                    "full_answer": "移项并化简得 x=1。",
                    "proof_obligations": [],
                    "visual_requirements": [],
                    "deduction_policy": [
                        "未写出移项化简过程或 x=1 结论时，对应证据点未达成"
                    ],
                    "allow_alternative_methods": True,
                    "evidence_points": [
                        {
                            "evidence_point_id": f"step-{question_id}",
                            "step_index": 1,
                            "target": "移项并化简得到 x=1",
                            "justification": "依据等式性质移项并化简",
                            "answer_anchor": "移项并化简",
                            "observable_evidence": "写出正确的移项和化简过程",
                            "depends_on": [],
                            "fine_term_links": [
                                {
                                    "fine_term_id": "kp_alg_linear_equation",
                                    "fine_term_name": "一元一次方程",
                                    "role": "direct",
                                }
                            ],
                            "equivalent_rules": [],
                            "counterexamples": ["移项后未变号"],
                        },
                        {
                            "evidence_point_id": f"result-{question_id}",
                            "step_index": 2,
                            "target": "解得 x=1",
                            "justification": "由前一步的等价方程求解未知数",
                            "answer_anchor": "x=1",
                            "observable_evidence": "写出 x=1 并作为最终结论",
                            "depends_on": [f"step-{question_id}"],
                            "fine_term_links": [
                                {
                                    "fine_term_id": "kp_alg_linear_equation",
                                    "fine_term_name": "一元一次方程",
                                    "role": "direct",
                                }
                            ],
                            "equivalent_rules": ["1=x"],
                            "counterexamples": ["只写中间式未给出解"],
                        },
                    ],
                }
            ],
            "auxiliary_rules": [],
            "rationale": "按可观察步骤拆分。",
            "confidence": 0.95,
        },
    }


class _DeferredProtocolResponse:
    def __init__(self, question_ids: list[int]) -> None:
        self.output_text = json.dumps(
            {
                "results": [
                    _deferred_combined_item(question_id) for question_id in question_ids
                ]
            },
            ensure_ascii=False,
        )
        self.usage = {
            "input_tokens": 100,
            "output_tokens": 50,
            "total_tokens": 150,
        }


class _DeferredProtocol:
    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []

    def responses(self, **kwargs: Any) -> _DeferredProtocolResponse:
        self.calls.append(kwargs)
        prompt = kwargs["kwargs"]["input"][1]["content"][0]["text"]
        questions = json.loads(prompt)["questions"]
        return _DeferredProtocolResponse(
            [int(item["question_id"]) for item in questions]
        )


class _InterruptingDeferredProtocol(_DeferredProtocol):
    def responses(self, **kwargs: Any) -> _DeferredProtocolResponse:
        self.calls.append(kwargs)
        raise KeyboardInterrupt()


class _DeferredTaggingService:
    model = "synthetic-combined-v3"

    def __init__(self, protocol: _DeferredProtocol) -> None:
        self.protocol = protocol

    def _protocol_adapter(self) -> _DeferredProtocol:
        return self.protocol

    def taxonomy_contracts(
        self,
        contexts: dict[int, object],
    ) -> dict[int, dict[str, Any]]:
        return {question_id: _deferred_taxonomy_contract() for question_id in contexts}


def test_evidence_analysis_checkpoint_is_reused_by_score_retry_without_model_replay(
    tmp_path: Path,
) -> None:
    db, session_id, old_paths = _db_with_session(tmp_path)
    _source_service, source = _controlled_source(
        tmp_path,
        session_id,
        text="1. 解方程 x+1=2。\n答案：x=1",
    )
    input_id = _stage_controlled_input(
        tmp_path,
        _source_service,
        source,
        old_paths,
        generation_mode="batched",
        sync_to_question_bank=True,
    )
    first_context, store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "generate",
            "generation_mode": "batched",
            "input_id": input_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
        },
    )
    protocol = _DeferredProtocol()
    tagging_service = _DeferredTaggingService(protocol)
    tagging_factory_calls = 0

    def tagging_factory() -> _DeferredTaggingService:
        nonlocal tagging_factory_calls
        tagging_factory_calls += 1
        return tagging_service

    intake_calls: list[dict[str, Any]] = []

    def intake_runner(**kwargs: Any) -> dict[str, object]:
        intake_calls.append(kwargs)
        return {
            "outcome": "complete",
            "imported_count": 1,
            "tagged_count": 1,
            "criteria_count": 1,
            "failed_count": 0,
        }

    first = run_config_generation_job(
        context=first_context,
        db=db,
        upload_config_dir=tmp_path / "uploaded",
        data_root=tmp_path / "data",
        question_bank_db_path=tmp_path / "data" / "databases" / "question_bank.db",
        tagging_ai_service_factory=tagging_factory,
        taxonomy_governance=object(),
        question_bank_intake_runner=intake_runner,
    )

    assert first["outcome"] == "partial"
    assert first["score_allocation_pending"] is True
    assert first["question_bank_sync_state"] == "ready"
    assert first["question_bank_imported_count"] == 1
    assert first["question_bank_tagged_count"] == 1
    assert first["question_bank_evidence_count"] == 1
    assert len(intake_calls) == 1
    assert intake_calls[0]["artifact"].bundle.status == "succeeded"
    assert len(protocol.calls) == 1
    assert tagging_factory_calls == 1
    store.finish(first_context.job_id, "succeeded", result=first)
    completed_first = store.get_job(first_context.job_id)
    assert completed_first is not None
    assert completed_first.status == "succeeded"
    artifact_files = list(
        (tmp_path / "uploaded").glob("deferred_question_analysis_*.json")
    )
    assert len(artifact_files) == 1

    retry_context, _retry_store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "retry",
            "generation_mode": "batched",
            "source_job_id": first_context.job_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
        },
    )
    retried = run_config_generation_job(
        context=retry_context,
        db=db,
        upload_config_dir=tmp_path / "uploaded",
        data_root=tmp_path / "data",
        question_bank_db_path=tmp_path / "data" / "databases" / "question_bank.db",
        tagging_ai_service_factory=tagging_factory,
        taxonomy_governance=object(),
        question_bank_intake_runner=intake_runner,
    )

    assert retried["outcome"] == "partial"
    assert len(protocol.calls) == 1
    assert tagging_factory_calls == 1
    assert len(intake_calls) == 2


def test_interrupted_evidence_request_is_reported_uncertain_without_model_replay(
    tmp_path: Path,
) -> None:
    db, session_id, old_paths = _db_with_session(tmp_path)
    source_service, source = _controlled_source(
        tmp_path,
        session_id,
        text="1. 解方程 x+1=2。\n答案：x=1",
    )
    input_id = _stage_controlled_input(
        tmp_path,
        source_service,
        source,
        old_paths,
        generation_mode="batched",
        sync_to_question_bank=True,
    )
    first_context, store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "generate",
            "generation_mode": "batched",
            "input_id": input_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
        },
    )
    interrupted_protocol = _InterruptingDeferredProtocol()

    with pytest.raises(KeyboardInterrupt):
        _run_config_generation_job_impl(
            context=first_context,
            db=db,
            upload_config_dir=tmp_path / "uploaded",
            data_root=tmp_path / "data",
            tagging_ai_service_factory=lambda: _DeferredTaggingService(
                interrupted_protocol
            ),
            taxonomy_governance=object(),
        )

    # Calling the implementation directly models a hard process exit: the normal
    # wrapper cannot run cleanup, so both the durable artifact and input survive.
    staged = load_config_generation_input(tmp_path / "uploaded", input_id)
    assert staged["analysis_artifact_id"]
    assert len(interrupted_protocol.calls) == 1
    assert preserve_interrupted_config_generation_checkpoints(
        tmp_path / "uploaded",
        store,
    ) == {input_id}
    interrupted_job = store.get_job(first_context.job_id)
    assert interrupted_job is not None
    assert interrupted_job.result["outcome"] == "partial"
    assert interrupted_job.result["uncertain_question_ids"] == ["Q1"]
    assert interrupted_job.result["needs_teacher_resolution"] is True
    assert interrupted_job.result["retryable"] is False

    replay_protocol = _DeferredProtocol()
    partial_intake_calls: list[dict[str, Any]] = []

    def partial_intake_runner(**kwargs: Any) -> dict[str, object]:
        partial_intake_calls.append(kwargs)
        return {
            "outcome": "partial",
            "imported_count": 1,
            "tagged_count": 0,
            "evidence_count": 0,
            "failed_count": 1,
        }

    resume_context, resume_store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "generate",
            "generation_mode": "batched",
            "input_id": input_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
        },
    )
    resumed = run_config_generation_job(
        context=resume_context,
        db=db,
        upload_config_dir=tmp_path / "uploaded",
        data_root=tmp_path / "data",
        question_bank_db_path=tmp_path / "data" / "databases" / "question_bank.db",
        tagging_ai_service_factory=lambda: _DeferredTaggingService(replay_protocol),
        taxonomy_governance=object(),
        question_bank_intake_runner=partial_intake_runner,
    )
    assert resumed["outcome"] == "partial"
    assert resumed["uncertain_question_ids"] == ["Q1"]
    assert resumed["retryable"] is False
    assert resumed["question_bank_sync_state"] == "partial"
    assert resumed["question_bank_imported_count"] == 1
    assert resumed["question_bank_tagged_count"] == 0
    assert len(partial_intake_calls) == 1
    assert partial_intake_calls[0]["artifact"].bundle.status == "needs_resolution"
    assert replay_protocol.calls == []

    resume_store.finish(resume_context.job_id, "succeeded", result=resumed)
    confirmed_protocol = _DeferredProtocol()

    def confirmed_intake_runner(**_kwargs: Any) -> dict[str, object]:
        return {
            "outcome": "complete",
            "imported_count": 1,
            "tagged_count": 1,
            "criteria_count": 1,
            "failed_count": 0,
        }

    confirmed_context, _confirmed_store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "retry",
            "generation_mode": "batched",
            "source_job_id": resume_context.job_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
            "retry_question_ids": ["Q1"],
            "confirm_uncertain_retry": True,
        },
    )
    confirmed = run_config_generation_job(
        context=confirmed_context,
        db=db,
        upload_config_dir=tmp_path / "uploaded",
        data_root=tmp_path / "data",
        question_bank_db_path=tmp_path / "data" / "databases" / "question_bank.db",
        tagging_ai_service_factory=lambda: _DeferredTaggingService(confirmed_protocol),
        taxonomy_governance=object(),
        question_bank_intake_runner=confirmed_intake_runner,
    )

    assert len(confirmed_protocol.calls) == 1
    assert confirmed["generated_questions"] == 1
    assert confirmed.get("uncertain_question_ids") is None
    assert confirmed["score_allocation_pending"] is True


def _png_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (24, 18), "navy").save(output, format="PNG")
    return output.getvalue()


def _docx_with_question_image() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("1. 解方程 x+1=2。")
    paragraph.add_run().add_picture(io.BytesIO(_png_bytes()), width=Inches(0.25))
    document.add_paragraph("答案：x=1")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _image_source(
    tmp_path: Path,
    session_id: int,
) -> tuple[ConfigSourceService, ConfigSourceRecord, dict[str, Any]]:
    source_service = ConfigSourceService(tmp_path / "uploaded")
    record = asyncio.run(
        source_service.stage_and_parse(
            session_id=session_id,
            filename="七年级下册测试卷.docx",
            chunks=_chunks(_docx_with_question_image()),
        )
    )
    automatic = next(
        item
        for item in record.public_snapshot()["assets"]
        if item["assignment_state"] == "automatic"
    )
    return source_service, record, automatic


def _stage_image_intake_input(
    tmp_path: Path,
    source: ConfigSourceRecord,
    automatic: dict[str, Any],
    expected_paths: tuple[str, str],
) -> str:
    return stage_config_source_generation_input(
        tmp_path / "uploaded",
        session_id=source.session_id,
        expected_rubric_path=expected_paths[0],
        expected_answer_key_path=expected_paths[1],
        generation_mode="batched",
        source_id=source.source_id,
        source_revision=source.source_revision,
        sync_to_question_bank=True,
        curriculum_volume_id="bnu24-math-g7-upper",
        decisions=[
            {
                "question_id": source.questions[0].question_id,
                "question_type": "proof",
                "excluded": False,
            }
        ],
        asset_decisions=[
            {
                "candidate_id": automatic["asset_id"],
                "action": "bind",
                "question_id": source.questions[0].question_id,
                "asset_kind": "answer",
            }
        ],
    )


def test_deferred_intake_blocks_when_asset_decisions_go_stale(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db, session_id, old_paths = _db_with_session(tmp_path)
    _source_service, source, automatic = _image_source(tmp_path, session_id)
    input_id = _stage_image_intake_input(tmp_path, source, automatic, old_paths)
    context, _store = _job_context(
        db.db_path,
        {
            "session_id": session_id,
            "mode": "generate",
            "generation_mode": "batched",
            "input_id": input_id,
            "source_id": source.source_id,
            "source_revision": source.source_revision,
            "sync_to_question_bank": True,
        },
    )
    def stale_resolve(*_args: Any, **_kwargs: Any) -> tuple[dict[str, Any], ...]:
        raise ValueError("invalid ambiguous asset decision")

    monkeypatch.setattr(
        ConfigSourceService,
        "resolve_asset_decision_overrides",
        stale_resolve,
    )

    def intake_runner(**_kwargs: Any) -> dict[str, object]:
        pytest.fail("intake must not start when asset decisions are stale")

    result = run_config_generation_job(
        context=context,
        db=db,
        upload_config_dir=tmp_path / "uploaded",
        data_root=tmp_path / "data",
        question_bank_db_path=tmp_path / "data" / "databases" / "question_bank.db",
        tagging_ai_service_factory=lambda: _DeferredTaggingService(_DeferredProtocol()),
        taxonomy_governance=object(),
        question_bank_intake_runner=intake_runner,
    )

    assert result["question_bank_sync_state"] == "failed"
    assert "图片归属决定已失效" in str(result["question_bank_sync_error"])
    session = db.sessions.get_grading_session(session_id)
    assert session is not None
    assert session["question_bank_sync_state"] == "failed"
    assert "asset_decisions_stale" in str(
        session.get("question_bank_sync_details_json")
        or session.get("question_bank_sync_details")
        or ""
    )
