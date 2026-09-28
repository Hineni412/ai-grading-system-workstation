from __future__ import annotations

import asyncio
import io
import json
import zipfile
from collections.abc import AsyncIterator, Callable
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.config_workspace.sources import (
    AmbiguousAssetDecision,
    ConfigSourceChangedError,
    ConfigSourceInvalidError,
    ConfigSourceNotFoundError,
    ConfigSourceService,
    ConfigSourceTooLargeError,
    ConfigSourceTypeUnsupportedError,
    QuestionDecision,
)


async def chunks(*parts: bytes) -> AsyncIterator[bytes]:
    for part in parts:
        yield part


def service(tmp_path: Path, **kwargs: object) -> ConfigSourceService:
    return ConfigSourceService(tmp_path, **kwargs)


def _png_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (24, 18), "navy").save(output, format="PNG")
    return output.getvalue()


def _docx_bytes(
    *, question: str = "1. Solve x squared.", with_image: bool = False
) -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run(question)
    rich = paragraph.add_run(" rich")
    rich.bold = True
    exponent = paragraph.add_run("2")
    exponent.font.superscript = True
    if with_image:
        paragraph.add_run().add_picture(io.BytesIO(_png_bytes()), width=Inches(0.25))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_ambiguous_floating_image() -> bytes:
    document = Document()
    document.add_paragraph("1. 第一题")
    document.add_paragraph().add_run().add_picture(
        io.BytesIO(_png_bytes()), width=Inches(0.25)
    )
    document.add_paragraph("2. 第二题")
    document.add_paragraph().add_run().add_picture(
        io.BytesIO(_png_bytes()), width=Inches(0.25)
    )
    document.add_paragraph("3. 第三题")
    original = io.BytesIO()
    document.save(original)
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original.getvalue()), "r") as source:
        with zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "word/document.xml":
                    content = content.replace(b"<wp:inline", b"<wp:anchor", 1)
                    content = content.replace(b"</wp:inline>", b"</wp:anchor>", 1)
                target.writestr(info, content)
    return rewritten.getvalue()


def _pdf_bytes(*, question: bool = True, answer: bool = True, pages: int = 1) -> bytes:
    document = fitz.open()
    for index in range(pages):
        page = document.new_page()
        if index == 0:
            lines = ["Synthetic exam"]
            if question:
                lines.extend(["1. Compute 1 + 1.", "A. 1   B. 2   C. 3   D. 4"])
            if answer:
                lines.extend(["Answer", "1. B"])
            page.insert_text((72, 72), "\n".join(lines), fontsize=12)
    payload = document.tobytes()
    document.close()
    return payload


def valid_source_bytes(filename: str) -> bytes:
    return _pdf_bytes() if filename.casefold().endswith(".pdf") else _docx_bytes()


def test_word_preview_reuses_split_xml_after_reopening(tmp_path: Path, monkeypatch) -> None:
    from backend.config_workspace import formula_preview

    source_service = service(tmp_path)
    content = _docx_bytes(question="1. 计算 x²+1。", with_image=True)
    record = asyncio.run(source_service.stage_and_parse(
        session_id=7, filename="formula.docx", chunks=chunks(content),
    ))
    saved_xml = record.private_blocks[0]["_word_paragraphs"]
    expected = formula_preview.load_paragraph_xml_index(content)
    assert formula_preview._index_paragraphs(saved_xml) == expected
    formula_preview._INDEX_CACHE.clear()

    def unexpected_parse(*args, **kwargs):
        pytest.fail("Preview must reuse the persisted split XML")

    monkeypatch.setattr(formula_preview, "load_paragraph_xml_index", unexpected_parse)
    snapshot = record.public_snapshot()
    assert "_word_paragraphs" not in json.dumps(snapshot, ensure_ascii=False)
    prepared = source_service.apply_teacher_decisions(record, [])
    assert all("_word_paragraphs" not in block for block in prepared.confirmed_blocks)
    # Rehydrate from the persisted manifest rather than an in-memory extraction.
    manifest = json.loads(record.manifest_path.read_text(encoding="utf-8"))
    formula_preview._INDEX_CACHE.clear()
    reopened = formula_preview.paragraph_xml_index(
        7, record.source_id, record.source_revision, content,
        paragraphs=manifest["private_blocks"][0]["_word_paragraphs"],
    )
    assert reopened == expected


def test_optional_preview_xml_does_not_reject_large_manifests(tmp_path: Path) -> None:
    source_service = service(tmp_path)
    content = _docx_bytes()
    original = asyncio.run(source_service.stage_and_parse(
        session_id=7, filename="formula.docx", chunks=chunks(content),
    ))
    manifest = json.loads(original.manifest_path.read_text(encoding="utf-8"))
    for block in manifest["private_blocks"]:
        block.pop("_word_paragraphs", None)
    source_service.max_manifest_bytes = len(json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")) + 64
    fallback = asyncio.run(source_service.stage_and_parse(
        session_id=8, filename="formula.docx", chunks=chunks(content),
    ))
    assert all("_word_paragraphs" not in block for block in fallback.private_blocks)
    assert fallback.public_snapshot()["questions"]


def test_ambiguous_floating_image_waits_for_one_manual_adjacent_binding(
    tmp_path: Path,
) -> None:
    source_service = service(tmp_path)
    record = asyncio.run(
        source_service.stage_and_parse(
            session_id=7,
            filename="ambiguous.docx",
            chunks=chunks(_docx_with_ambiguous_floating_image()),
        )
    )

    candidate = record.public_snapshot()["ambiguous_assets"][0]
    assert candidate["previous_question_id"] == "Q1"
    assert candidate["next_question_id"] == "Q2"
    assert "filename" not in candidate
    content, media_type = source_service.read_ambiguous_asset(
        session_id=7,
        source_id=record.source_id,
        candidate_id=candidate["candidate_id"],
    )
    assert content == _png_bytes()
    assert media_type == "image/png"

    unresolved = source_service.apply_teacher_decisions(record, [])
    assert [item["question_id"] for item in unresolved.confirmed_blocks] == ["Q3"]

    bound = source_service.apply_teacher_decisions(
        record,
        [],
        [
            AmbiguousAssetDecision(
                candidate_id=candidate["candidate_id"],
                action="bind",
                question_id="Q2",
                asset_kind="question",
            )
        ],
    )
    assert [item["question_id"] for item in bound.confirmed_blocks] == [
        "Q1",
        "Q2",
        "Q3",
    ]
    assert isinstance(bound.question_images["Q2"]["question"], list)
    assert len(bound.question_images["Q2"]["question"]) == 2
    ignored = source_service.apply_teacher_decisions(
        record,
        [],
        [
            AmbiguousAssetDecision(
                candidate_id=candidate["candidate_id"],
                action="ignore",
            )
        ],
    )
    assert [item["question_id"] for item in ignored.confirmed_blocks] == [
        "Q1",
        "Q2",
        "Q3",
    ]
    moved_beyond_adjacent = source_service.apply_teacher_decisions(
        record,
        [],
        [
            AmbiguousAssetDecision(
                candidate_id=candidate["candidate_id"],
                action="bind",
                question_id="Q3",
                asset_kind="question",
            )
        ],
    )
    assert moved_beyond_adjacent.question_images["Q3"]["question"] is not None
    with pytest.raises(ValueError, match="target is invalid"):
        source_service.apply_teacher_decisions(
            record,
            [],
            [
                AmbiguousAssetDecision(
                    candidate_id=candidate["candidate_id"],
                    action="bind",
                    question_id="missing",
                    asset_kind="question",
                )
            ],
        )


def test_failed_replacement_keeps_previous_active_pointer(tmp_path: Path) -> None:
    current = asyncio.run(
        service(tmp_path).stage_and_parse(
            session_id=7,
            filename="current.pdf",
            chunks=chunks(_pdf_bytes()),
        )
    )

    with pytest.raises(ConfigSourceInvalidError):
        asyncio.run(
            service(tmp_path).stage_and_parse(
                session_id=7,
                filename="bad.pdf",
                chunks=chunks(b"%PDF-not-valid"),
            )
        )

    assert (
        service(tmp_path).load(session_id=7, source_id=current.source_id).source_id
        == current.source_id
    )


def test_parser_failure_after_asset_write_removes_all_new_owned_files(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import rubric_auto_cropper
    import backend.config_workspace.sources as sources_module

    monkeypatch.setattr(
        sources_module,
        "extract_pdf_text",
        lambda _payload: "text",
    )
    monkeypatch.setattr(
        sources_module,
        "parse_plain_question_blocks",
        lambda _text: [
            {"question_id": "Q1", "question_type": "choice", "text": "one"},
            {"question_id": "Q2", "question_type": "choice", "text": "two"},
        ],
    )
    monkeypatch.setattr(
        rubric_auto_cropper,
        "extract_pdf_question_images",
        lambda _payload, _blocks: {
            "Q1": {"question": _png_bytes(), "answer": None},
            "Q2": {"question": b"not-an-image", "answer": None},
        },
    )
    monkeypatch.setattr(
        rubric_auto_cropper,
        "extract_pdf_images",
        lambda _payload: [_png_bytes()],
    )
    original_write = sources_module._write_bytes_atomic
    sentinel_path: Path | None = None

    def write_with_unowned_sentinel(
        path: Path, content: bytes, *args: object, **kwargs: object
    ) -> None:
        nonlocal sentinel_path
        original_write(path, content, *args, **kwargs)
        if sentinel_path is None:
            sentinel_path = path.parent / "not-owned.txt"
            sentinel_path.write_text("keep", encoding="utf-8")

    monkeypatch.setattr(
        sources_module, "_write_bytes_atomic", write_with_unowned_sentinel
    )

    with pytest.raises(ConfigSourceInvalidError):
        asyncio.run(
            service(tmp_path).stage_and_parse(
                session_id=7,
                filename="partial.pdf",
                chunks=chunks(_pdf_bytes()),
            )
        )

    source_root = tmp_path / "config_sources" / "session-7"
    assert sentinel_path is not None
    assert sentinel_path.read_text(encoding="utf-8") == "keep"
    assert [
        path
        for path in source_root.rglob("*")
        if path.is_file() and path != sentinel_path
    ] == []


def test_teacher_question_type_decision_is_preserved_as_a_hard_fact(
    tmp_path: Path,
) -> None:
    record = asyncio.run(
        service(tmp_path).stage_and_parse(
            session_id=7,
            filename="paper.pdf",
            chunks=chunks(_pdf_bytes()),
        )
    )

    prepared = service(tmp_path).apply_teacher_decisions(
        record,
        [QuestionDecision(question_id="Q1", question_type="proof", excluded=False)],
    )

    assert prepared.confirmed_blocks[0]["question_type"] == "proof"
    assert prepared.confirmed_blocks[0]["question_type_confirmed"] is True
    assert prepared.document_text == record.private_document_text
    assert prepared.question_images["Q1"]["question"]
    with pytest.raises(ValueError, match="unknown question"):
        service(tmp_path).apply_teacher_decisions(
            record,
            [
                QuestionDecision(
                    question_id="Q404", question_type="choice", excluded=True
                )
            ],
        )
