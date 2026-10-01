from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from backend.repositories.grading_database import open_grading_repositories


def _client_with_db(tmp_path):
    from backend.api.app import create_app
    from backend.api.dependencies import get_grading_db
    

    db = open_grading_repositories(tmp_path / "grading.db")
    db.initialize()
    app = create_app()
    app.dependency_overrides[get_grading_db] = lambda: db
    return TestClient(app), db


def test_student_workspace_and_import_commit_share_revision_contract(tmp_path) -> None:
    from backend.repositories.students import StudentRecord

    client, db = _client_with_db(tmp_path)
    db.students.upsert_students([StudentRecord("S001", "旧姓名", "一班")])

    workspace = client.get(
        "/api/students/workspace?search=旧姓名&class_name=一班&page=1&page_size=50"
    )

    assert workspace.status_code == 200
    workspace_payload = workspace.json()
    assert workspace_payload["total"] == 1
    assert workspace_payload["page"] == 1
    assert workspace_payload["class_names"] == ["一班"]
    assert len(workspace_payload["roster_revision"]) == 64

    preview = client.post(
        "/api/students/import/preview?filename=students.csv",
        content=(
            "student_code,name,class_name\nS001,新姓名,二班\nS002,匿名学生乙,二班\n"
        ).encode(),
    ).json()
    selected = [
        {
            "student_code": row["student_code"],
            "name": row["name"],
            "class_name": row["class_name"],
        }
        for row in preview["rows"]
        if row["selectable"]
    ]
    committed = client.post(
        "/api/students/import/commit",
        json={
            "expected_revision": preview["roster_revision"],
            "items": selected,
        },
    )

    assert committed.status_code == 200
    assert committed.json()["inserted"] == 1
    assert committed.json()["updated"] == 1
    repeated = client.post(
        "/api/students/import/commit",
        json={
            "expected_revision": preview["roster_revision"],
            "items": selected,
        },
        headers={"x-request-id": "rid-student-conflict"},
    )
    assert repeated.status_code == 409
    assert repeated.json()["error"]["code"] == "student_roster_conflict"
    assert repeated.json()["error"]["request_id"] == "rid-student-conflict"


def test_student_delete_requires_impact_revision_and_confirmation(tmp_path) -> None:
    from backend.repositories.students import StudentRecord

    client, db = _client_with_db(tmp_path)
    db.students.upsert_students([StudentRecord("S001", "Alice", "Class 1")])
    student_id = db.students.list_students()[0]["id"]

    impact = client.get(f"/api/students/{student_id}/deletion-impact")
    unconfirmed = client.delete(
        f"/api/students/{student_id}",
        params={
            "expected_revision": impact.json()["roster_revision"],
            "confirmed": "false",
        },
        headers={"x-request-id": "rid-delete-unconfirmed"},
    )

    assert impact.status_code == 200
    assert impact.json()["counts"]["deleted_students"] == 1
    assert unconfirmed.status_code == 422
    assert unconfirmed.json()["error"]["code"] == "student_roster_invalid"
    assert db.students.list_students()[0]["student_code"] == "S001"


@pytest.mark.parametrize("run_state", ["running", "pause_requested", "paused"])
def test_student_delete_is_rejected_while_grading_is_active(
    tmp_path,
    run_state: str,
) -> None:
    from backend.repositories.students import StudentRecord
    from grading_run_store import GradingRunStore

    client, db = _client_with_db(tmp_path)
    db.students.upsert_students([StudentRecord("S001", "Alice", "Class 1")])
    student_id = int(db.students.list_students()[0]["id"])
    session_id = db.sessions.create_grading_session(
        "Active grading", "rubric.json", "answer.json"
    )
    run_store = GradingRunStore(db.db_path)
    run = run_store.begin(session_id, "a" * 64, "full_paper")
    run_store.add_item(
        run.id,
        source_label="001",
        student_id=student_id,
        paper_fingerprint="b" * 64,
        config_fingerprint="a" * 64,
        status="grading",
    )
    if run_state == "pause_requested":
        assert run_store.request_pause(session_id) is True
    elif run_state == "paused":
        run_store.finish(run.run_token, "paused")
    impact = client.get(f"/api/students/{student_id}/deletion-impact").json()

    response = client.delete(
        f"/api/students/{student_id}",
        params={
            "expected_revision": impact["roster_revision"],
            "confirmed": "true",
        },
        headers={"x-request-id": "rid-delete-active-grading"},
    )

    assert response.status_code == 409
    assert response.json() == {
        "error": {
            "code": "student_grading_active",
            "message": "Wait for grading to finish before deleting this student",
            "details": {},
            "request_id": "rid-delete-active-grading",
        }
    }
    assert db.students.list_students()[0]["student_code"] == "S001"
    assert run_store.get_run(run.id).state == run_state
    assert run_store.counts(run.id)["grading"] == 1
    assert list(db.backup_dir.glob("grading_before_delete_student_*.db")) == []


def _seed_wrong_question_books(tmp_path):
    import json
    import sqlite3
    from pathlib import Path
    from backend.repositories.students import StudentRecord
    from question_bank.models.question import QuestionCreate
    from question_bank.services.source_question_link_service import SourceQuestionLinkService
    from tests.question_bank_support import QuestionBankTestStore

    db = open_grading_repositories(tmp_path / "databases" / "grading.db")
    db.initialize()
    db.students.upsert_students([
        StudentRecord("TEST001", "测试学生甲", "测试班"),
        StudentRecord("TEST002", "测试全对学生", "测试班"),
        StudentRecord("TEST003", "测试学生丙", "测试班"),
    ])
    students = sorted(db.students.list_students(), key=lambda row: row["student_code"])
    rubric = tmp_path / "config" / "uploaded" / "test_rubric.json"
    rubric.parent.mkdir(parents=True, exist_ok=True)
    rubric.write_text(json.dumps({"questions": [
        {"question_id": f"Q{i}", "max_score": 5} for i in range(1, 7)
    ] + [{"question_id": "Q10", "max_score": 5, "parts": [
        {"part_id": "Q10(P1)", "part_score": 2}, {"part_id": "Q10(P2)", "part_score": 3},
    ]}]}), encoding="utf-8")
    sessions = [db.sessions.create_grading_session(
        name, str(rubric), "answer.json", curriculum_volume_id="bnu24-math-g8-upper",
    ) for name in ("测试早期考试", "测试后期考试")]
    other = db.sessions.create_grading_session("测试其他学期", str(rubric), "answer.json", curriculum_volume_id="bnu24-math-g8-lower")
    qb_path = tmp_path / "databases" / "question_bank.db"
    bank = QuestionBankTestStore(qb_path)
    question_ids = [bank.add_question(QuestionCreate(
        question_number=str(i), question_type="解答题",
        question_text=f"测试原题{i} 计算 $x^2+{i}$", answer_text=f"测试解析{i} 用代入法求值",
    )) for i in range(1, 7)]
    links = SourceQuestionLinkService(qb_path)
    for sid in sessions:
        for source, bank_id in [("Q1", question_ids[0]), ("Q2", question_ids[1]), ("Q10", question_ids[2]), ("Q5", question_ids[3]), ("Q4", question_ids[4])]:
            links.confirm_link(grading_session_id=sid, source_question_id=source, bank_question_id=bank_id, link_method="manual")
        links.suggest_link(grading_session_id=sid, source_question_id="Q6", bank_question_id=question_ids[5], confidence=0.8, link_method="similarity")
    with sqlite3.connect(db.db_path) as conn:
        conn.execute("UPDATE grading_sessions SET created_at = '2026-01-01' WHERE id = ?", (sessions[0],))
        conn.execute("UPDATE grading_sessions SET created_at = '2026-02-01' WHERE id = ?", (sessions[1],))
        for sid in [*sessions, other]:
            for index, student in enumerate(students):
                paper_id = conn.execute("INSERT INTO exam_papers (session_id, student_id, front_image, back_image, match_status, processing_status) VALUES (?, ?, '', '', 'matched', 'graded')", (sid, student["id"])).lastrowid
                result_id = conn.execute("INSERT INTO session_results (session_id, student_id, paper_id, student_score, total_score, needs_human_review, raw_json) VALUES (?, ?, ?, 25, 35, 0, '{}')", (sid, student["id"], paper_id)).lastrowid
                details = [("Q1", 0), ("Q5", 0), ("Q6", 0), ("Q10(P1)", 2), ("Q10(P2)", 1)] if sid == sessions[0] else [("Q1", 1), ("Q2", 2)]
                if index == 1:
                    details = [("Q1", 5), ("Q2", 5)]
                if index == 2:
                    details = [("Q1", 5)] if sid == sessions[0] else [("Q2", 2)]
                for qid, score in details:
                    conn.execute("INSERT INTO session_details (result_id, question_id, score_awarded, knowledge_ids) VALUES (?, ?, ?, '[\"UNKNOWN\"]')", (result_id, qid, score))
                if index == 0 and sid == sessions[0]:
                    conn.execute("INSERT INTO teacher_score_locks (session_id, scan_batch_id, student_id, question_id, score_awarded, max_score, source_target_type, source_target_id) VALUES (?, 'test-batch', ?, 'Q5', 5, 5, 'detail', 1)", (sid, student["id"]))
    return db, qb_path, students, sessions, question_ids


def test_wrong_question_books_preview_export_replay_download_and_read_only(tmp_path):
    import sqlite3
    from io import BytesIO
    from zipfile import ZipFile
    from pathlib import Path
    from docx import Document
    from backend.api.dependencies import get_job_manager, get_question_bank_db_path, get_reports_dir
    from backend.jobs.default_handlers import register_default_job_handlers
    from backend.jobs.manager import JobManager
    from backend.jobs.store import JobStore

    db, qb_path, students, sessions, _ = _seed_wrong_question_books(tmp_path)
    from backend.api.app import create_app
    from backend.api.dependencies import get_grading_db
    app = create_app()
    app.dependency_overrides[get_grading_db] = lambda: db
    client = TestClient(app)
    manager = JobManager(JobStore(db.db_path), max_workers=1)
    reports = tmp_path / "reports"
    register_default_job_handlers(manager, db_path=db.db_path, reports_dir=reports, data_root=tmp_path, question_bank_db_path=qb_path)
    client.app.dependency_overrides[get_job_manager] = lambda: manager
    client.app.dependency_overrides[get_question_bank_db_path] = lambda: qb_path
    client.app.dependency_overrides[get_reports_dir] = lambda: reports
    with sqlite3.connect(qb_path) as conn:
        before_bank = list(conn.iterdump())
    with sqlite3.connect(db.db_path) as conn:
        before_grades = list(conn.execute("SELECT * FROM session_details"))
    try:
        scope = {"curriculum_volume_id": "bnu24-math-g8-upper", "include_class": True}
        response = client.post(f"/api/students/{students[0]['id']}/wrong-question-book/preview", json=scope)
        assert response.status_code == 200, response.text
        preview = response.json()
        assert len(preview["students"]) == 3
        assert preview["question_count"] == 4
        assert preview["session_ids"] == sessions
        evidence = client.get(f"/api/students/{students[0]['id']}/exam-results", params={"curriculum_volume_id": scope["curriculum_volume_id"]}).json()
        assert next(item for exam in evidence["sessions"] for item in exam["items"] if item["question_id"] == "Q10(P2)")["bank_question_id"] is not None
        assert [(item["session_name"], item["question_id"]) for item in preview["missing_items"]] == [("测试早期考试", "Q6")]
        selected = client.post(f"/api/students/{students[0]['id']}/wrong-question-book/preview", json=dict(scope, session_ids=[sessions[1]])).json()
        assert selected["question_count"] == 3
        assert selected["missing_items"] == []
        body = {"curriculum_volume_id": scope["curriculum_volume_id"], "student_ids": [row["id"] for row in students], "session_ids": list(reversed(sessions)), "client_request_token": "a" * 32}
        submitted = client.post("/api/students/wrong-question-books", json=body)
        assert submitted.status_code == 202, submitted.text
        job_id = submitted.json()["id"]
        manager.wait(job_id, timeout=30)
        job = manager.get(job_id)
        assert job.status == "succeeded", job.error
        assert len(job.result["generated_students"]) == 2
        assert job.result["empty_students"] == [{"student_id": students[1]["id"], "student_name": students[1]["name"], "reason": "没有错题"}]
        assert job.result["failed_students"] == []
        path = Path(job.result["file_path"])
        assert path.suffix == ".zip"
        with ZipFile(path) as archive:
            assert set(archive.namelist()) == {"TEST001_测试学生甲_错题本.docx", "TEST003_测试学生丙_错题本.docx"}
            doc = Document(BytesIO(archive.read("TEST001_测试学生甲_错题本.docx")))
            text = "\n".join(p.text for p in doc.paragraphs)
            assert text.index("测试早期考试") < text.index("测试后期考试") < text.index("答案")
            assert text.count("测试原题1") == 1
            assert "测试原题3" in text and "测试原题4" not in text and "测试原题5" not in text and "测试原题6" not in text
            assert text.index("测试解析1") > text.index("答案")
            assert "________________" not in text
            assert "测试早期考试" in doc.sections[0].header.paragraphs[0].text
        assert client.post("/api/students/wrong-question-books", json=body).json()["id"] == job_id
        from concurrent.futures import ThreadPoolExecutor
        with ThreadPoolExecutor(max_workers=4) as pool:
            repeated = list(pool.map(lambda _: manager.submit_idempotent_export("wrong_question_export", job.payload), range(4)))
        assert all(item.id == job_id and not created for item, created in repeated)
        assert client.get(f"/api/students/wrong-question-books/by-request/{'a' * 32}").json()["id"] == job_id
        assert client.post("/api/students/wrong-question-books", json=dict(body, session_ids=[sessions[0]])).status_code == 409
        public = client.get(f"/api/jobs/{job_id}").json()["result"]
        assert "file_path" not in public and public["download_url"]
        assert client.get(public["download_url"]).status_code == 200
        assert not path.exists()
        assert not manager.get(job_id).result.get("file_path")
        assert client.get(public["download_url"]).status_code == 410
        single_body = dict(body, student_ids=[students[0]["id"]], client_request_token="b" * 32)
        single_id = client.post("/api/students/wrong-question-books", json=single_body).json()["id"]
        manager.wait(single_id, timeout=30)
        assert manager.get(single_id).result["filename"] == "TEST001_测试学生甲_错题本.docx"
        empty_body = dict(body, student_ids=[students[1]["id"]], client_request_token="c" * 32)
        empty_id = client.post("/api/students/wrong-question-books", json=empty_body).json()["id"]
        manager.wait(empty_id, timeout=30)
        assert manager.get(empty_id).result["empty_students"][0]["student_id"] == students[1]["id"]
        assert not manager.get(empty_id).result.get("file_path")
        with sqlite3.connect(qb_path) as conn:
            assert list(conn.iterdump()) == before_bank
        with sqlite3.connect(db.db_path) as conn:
            assert list(conn.execute("SELECT * FROM session_details")) == before_grades
    finally:
        manager.shutdown()
        db.close()


def test_wrong_question_books_exclude_unknown_scores_and_isolate_student_failures(tmp_path, monkeypatch):
    from pathlib import Path
    from backend.students.wrong_question_book import build_wrong_question_books
    from backend.students.exam_evidence import student_exam_evidence
    from backend.jobs.wrong_question_export import run_wrong_question_export
    from backend.jobs.manager import JobContext
    from backend.jobs.store import JobStore

    db, qb_path, students, sessions, question_ids = _seed_wrong_question_books(tmp_path)
    original = student_exam_evidence(db, qb_path, [row["id"] for row in students], "bnu24-math-g8-upper")
    unknown = dict(original[0], question_id="Q4", score_awarded=None, bank_question_id=question_ids[4])
    monkeypatch.setattr("backend.students.wrong_question_book.student_exam_evidence", lambda *args: [*original, unknown])
    plan = build_wrong_question_books(db, qb_path, [row["id"] for row in students], "bnu24-math-g8-upper", sessions)
    assert plan["question_count"] == 4
    assert question_ids[4] not in [qid for book in plan["books"] for section in book["sections"] for qid in section["question_ids"]]
    store = JobStore(db.db_path)
    payload = {"student_ids": [row["id"] for row in students], "session_ids": sessions, "curriculum_volume_id": "bnu24-math-g8-upper"}
    job = store.create_job("wrong_question_export", payload)

    def exporter(_path, _ids, out, **kwargs):
        if kwargs["title"].startswith("测试学生甲"):
            raise ValueError("synthetic failure")
        file = Path(out) / "sample.docx"
        file.write_bytes(b"synthetic Word")
        return file

    try:
        result = run_wrong_question_export(context=JobContext(job.id, job.job_type, payload, store), db_path=db.db_path, question_bank_db_path=qb_path, reports_dir=tmp_path / "reports", docx_exporter=exporter)
        assert [row["student_name"] for row in result["failed_students"]] == ["测试学生甲"]
        assert [row["student_name"] for row in result["generated_students"]] == ["测试学生丙"]
        assert Path(result["file_path"]).suffix == ".zip"
    finally:
        db.close()
