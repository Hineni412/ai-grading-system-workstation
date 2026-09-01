from __future__ import annotations

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from question_bank.services import question_read_service
from question_bank.services.preview_html import block_preview_html


def _paragraph_xml(*parts: str) -> str:
    """Build a real Word paragraph (runs + raw OMML) and return its XML."""
    document = Document()
    paragraph = document.add_paragraph()
    for part in parts:
        if part.startswith("<m:"):
            paragraph._p.append(parse_xml(part))  # noqa: SLF001 - test fixture needs exact Word structure.
        else:
            paragraph.add_run(part)
    return paragraph._p.xml  # noqa: SLF001 - fixture mirrors the stored payload.


def _omath(inner: str) -> str:
    return f'<m:oMath {nsdecls("m")}>{inner}</m:oMath>'


def test_paragraph_renders_text_and_formula_span() -> None:
    xml = _paragraph_xml(
        "计算：",
        _omath(
            "<m:f>"
            "<m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>3</m:t></m:r></m:den>"
            "</m:f>"
        ),
    )

    html = block_preview_html(xml, {})

    assert html is not None
    assert html.startswith("计算：")
    assert '<span class="qm" data-latex="\\frac{1}{3}">(1)/(3)</span>' in html


def test_run_styles_and_breaks_are_preserved() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("x")
    sup = paragraph.add_run("2")
    sup.font.superscript = True
    paragraph.add_run().add_break()
    underlined = paragraph.add_run("填空")
    underlined.font.underline = True

    html = block_preview_html(paragraph._p.xml, {})  # noqa: SLF001

    assert html is not None
    assert "x<sup>2</sup>" in html
    assert "<br>" in html
    assert "<u>填空</u>" in html


def test_inline_drawing_maps_relationship_to_asset_url() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("如图")
    run._r.append(  # noqa: SLF001 - fixture mirrors the stored drawing XML.
        parse_xml(
            f'<w:drawing {nsdecls("w")}>'
            '<wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
            '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:graphicData><pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            '<pic:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId7"/></pic:blipFill>'
            "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing>"
        )
    )

    html = block_preview_html(
        paragraph._p.xml,  # noqa: SLF001
        {"rId7": "/api/question-bank/questions/9/assets/2"},
    )

    assert html is not None
    assert '<img src="/api/question-bank/questions/9/assets/2"' in html

    # Unknown relationship ids are skipped, never leaking raw rel ids.
    html_missing = block_preview_html(paragraph._p.xml, {})  # noqa: SLF001
    assert html_missing is not None
    assert "<img" not in html_missing


def test_table_xml_renders_as_html_table() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("方法一")
    table.cell(0, 1).paragraphs[0].add_run("OC=OD")
    table.cell(1, 0).paragraphs[0].add_run("作图步骤")

    html = block_preview_html(table._tbl.xml, {})  # noqa: SLF001

    assert html is not None
    assert html.startswith("<table><tbody>")
    assert html.count("<tr>") == 2
    assert "<td>方法一</td>" in html
    assert "<td>OC=OD</td>" in html


def test_text_is_escaped_against_injection() -> None:
    xml = _paragraph_xml("<script>alert(1)</script>")

    html = block_preview_html(xml, {})

    assert html is not None
    assert "<script>" not in html
    assert "&lt;script&gt;" in html


def test_invalid_or_unsupported_xml_returns_none() -> None:
    assert block_preview_html("", {}) is None
    assert block_preview_html("<w:p", {}) is None
    assert block_preview_html(f"<w:sectPr {nsdecls('w')}/>", {}) is None


def test_consistency_check_rejects_drifted_xml() -> None:
    xml = _paragraph_xml(
        "计算：",
        _omath(
            "<m:f>"
            "<m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>3</m:t></m:r></m:den>"
            "</m:f>"
        ),
    )

    assert block_preview_html(xml, {}, expected_text="计算：(1)/(3)") is not None
    assert block_preview_html(xml, {}, expected_text="计算：(2)/(7)") is None


def _rich_payload(question_blocks: list[dict]) -> dict:
    return {"question_blocks": question_blocks, "answer_blocks": []}


def test_public_rich_content_projects_html() -> None:
    xml = _paragraph_xml(
        "计算：",
        _omath(
            "<m:f>"
            "<m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>3</m:t></m:r></m:den>"
            "</m:f>"
        ),
    )
    payload = _rich_payload(
        [{"text": "计算：(1)/(3)", "xml": xml, "image_relationships": {}}]
    )

    rich_blocks = question_read_service._rich_blocks(payload)
    public = question_read_service._public_rich_content(7, True, rich_blocks, [])

    block = public["question_blocks"][0]
    assert '<span class="qm" data-latex="\\frac{1}{3}">' in block["html"]
    assert "_xml" not in block
    assert "_rels" not in block


def test_public_rich_content_falls_back_without_xml_or_on_drift() -> None:
    xml = _paragraph_xml(
        "计算：",
        _omath(
            "<m:f>"
            "<m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>3</m:t></m:r></m:den>"
            "</m:f>"
        ),
    )
    payload = _rich_payload(
        [
            {"text": "计算：(1)/(3)", "xml": "", "image_relationships": {}},
            {"text": "计算：(2)/(7)", "xml": xml, "image_relationships": {}},
        ]
    )

    rich_blocks = question_read_service._rich_blocks(payload)
    public = question_read_service._public_rich_content(7, True, rich_blocks, [])

    assert public["question_blocks"][0]["html"] == ""
    assert public["question_blocks"][1]["html"] == ""
    assert public["question_blocks"][0]["segments"]
