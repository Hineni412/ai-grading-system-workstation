from __future__ import annotations

import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from question_bank.exporters.paper_docx_exporter import export_question_paper_docx
from question_bank.models.question import QuestionCreate
from tests.question_bank_support import QuestionBankTestStore


def test_grouped_docx_export_renumbers_questions_after_grouping(tmp_path: Path) -> None:
    db_path = tmp_path / "question_bank.db"
    service = QuestionBankTestStore(db_path)
    solution_id = service.add_question(
        QuestionCreate(
            question_number="3",
            question_type="解答题",
            question_text="解答题正文",
            answer_text="解答题答案",
        )
    )
    choice_id = service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="选择题",
            question_text="选择题正文",
            answer_text="选择题答案",
        )
    )
    blank_id = service.add_question(
        QuestionCreate(
            question_number="2",
            question_type="填空题",
            question_text="填空题正文",
            answer_text="填空题答案",
        )
    )

    output_path = export_question_paper_docx(
        db_path,
        [solution_id, choice_id, blank_id],
        tmp_path,
        title="分组题号测试",
        grouped_by_type=True,
    )

    paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]
    numbered = [text for text in paragraphs if re.match(r"^\d+\.\s", text)]

    assert [text.split(".", 1)[0] for text in numbered] == ["1", "2", "3"]


def test_exporter_latex_and_options_layout() -> None:
    from question_bank.exporters.base_exporter import parse_latex_runs, extract_and_format_options

    # 1. Test LaTeX parser
    runs = parse_latex_runs("已知 $x=2$ 且 $y=3$，求 $$x^2+y^2$$ 的值。")
    assert runs[0] == ("text", "已知 ")
    assert runs[1] == ("inline_math", "x=2")
    assert runs[2] == ("text", " 且 ")
    assert runs[3] == ("inline_math", "y=3")
    assert runs[4] == ("text", "，求 ")
    assert runs[5] == ("block_math", "x^2+y^2")
    assert runs[6] == ("text", " 的值。")

    # 2. Test Options extractor
    stem, options = extract_and_format_options("下列各数是有理数的是 ( )\nA. 1/2  B. pi  C. sqrt(2)  D. -1")
    assert stem == "下列各数是有理数的是 ( )"
    assert len(options) == 4
    assert options[0] == ("A", "1/2")
    assert options[1] == ("B", "pi")
    assert options[2] == ("C", "sqrt(2)")
    assert options[3] == ("D", "-1")


def test_ordinary_export_uses_shared_native_word_blocks(tmp_path: Path) -> None:
    db_path = tmp_path / "question_bank.db"
    service = QuestionBankTestStore(db_path)
    question_id = service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="解答题",
            question_text="不应显示的<sup>纯文字标签</sup>",
            answer_text="答案",
        )
    )
    source = Document()
    paragraph = source.add_paragraph()
    underlined = paragraph.add_run("1. 富内容 x")
    underlined.underline = True
    superscript = paragraph.add_run("2")
    superscript.font.superscript = True
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "普通组卷富内容表格"
    rich_root = tmp_path / "question_bank" / "rich_content"
    rich_root.mkdir(parents=True)
    rich_root.joinpath(f"question_{question_id}.json").write_text(
        json.dumps(
            {
                "version": 3,
                "question_id": question_id,
                "content_revision": "a" * 64,
                "question_blocks": [
                    {"text": paragraph.text, "xml": paragraph._p.xml},  # noqa: SLF001
                    {"text": "普通组卷富内容表格", "xml": table._tbl.xml},  # noqa: SLF001
                ],
                "answer_blocks": [],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    output = export_question_paper_docx(
        db_path,
        [question_id],
        tmp_path / "output",
        title="共享富内容",
        ensure_previews=False,
    )

    reopened = Document(output)
    visible = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    visible += "\n".join(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    assert "不应显示" not in visible
    assert "普通组卷富内容表格" in visible
    rich_paragraph = next(
        paragraph for paragraph in reopened.paragraphs if "富内容 x" in paragraph.text
    )
    assert rich_paragraph.text.startswith("1. ")
    assert rich_paragraph.paragraph_format.line_spacing == 1.1
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    assert b'<w:u w:val="single"' in document_xml
    assert b'<w:vertAlign w:val="superscript"' in document_xml


def _write_test_image(root: Path, name: str = "figure.png", size=(400, 240)) -> Path:
    from PIL import Image

    image_dir = root / "question_bank" / "extracted_images"
    image_dir.mkdir(parents=True, exist_ok=True)
    image_path = image_dir / name
    Image.new("RGB", size, "white").save(image_path)
    return image_path


def test_first_image_floats_right_on_long_stem(tmp_path: Path) -> None:
    image_path = _write_test_image(tmp_path)
    db_path = tmp_path / "question_bank.db"
    service = QuestionBankTestStore(db_path)
    stem = (
        "如图，在直角三角形中，两条直角边分别为3和4，"
        "请计算斜边的长度，并写出完整的推导过程与理由。"
    )
    question_id = service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="解答题",
            question_text=f"{stem}[[IMAGE:{image_path}|caption=第3题图]]",
            answer_text="答案",
        )
    )

    output = export_question_paper_docx(
        db_path, [question_id], tmp_path / "output", title="浮动图", ensure_previews=False
    )
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "<wp:anchor" in document_xml
    assert 'wp:wrapSquare wrapText="left"' in document_xml
    assert 'descr="第3题图"' in document_xml


def test_short_stem_image_stays_inline_paragraph(tmp_path: Path) -> None:
    image_path = _write_test_image(tmp_path, "short.png")
    db_path = tmp_path / "question_bank.db"
    service = QuestionBankTestStore(db_path)
    question_id = service.add_question(
        QuestionCreate(
            question_number="1",
            question_type="解答题",
            question_text=f"如图，求 x。[[IMAGE:{image_path}]]",
            answer_text="答案",
        )
    )

    output = export_question_paper_docx(
        db_path, [question_id], tmp_path / "output", title="短题干", ensure_previews=False
    )
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "<wp:anchor" not in document_xml
    assert "<w:drawing" in document_xml


def test_option_extraction_relaxed_boundaries() -> None:
    from question_bank.exporters.base_exporter import extract_and_format_options

    _stem, options = extract_and_format_options(
        "下列结论正确的是（）A.25 B.7 C.25或7 D.25 或16"
    )
    assert [label for label, _ in options] == ["A", "B", "C", "D"]
    assert options[0][1] == "25"
    # 3 个选项（OCR 丢一个）也拆分。
    _stem3, options3 = extract_and_format_options(
        "下列结论正确的是（）A.25 B.7 C.25或7"
    )
    assert [label for label, _ in options3] == ["A", "B", "C"]
    # 选项末尾的教辅内嵌答案被剔除。
    _stem4, options4 = extract_and_format_options(
        "下列结论正确的是（）A.25 B.7 C.9 D.25或16（D）"
    )
    assert options4[-1] == ("D", "25或16")


def test_option_extraction_after_chinese_without_space() -> None:
    from question_bank.exporters.base_exporter import extract_and_format_options

    # 中考卷常见 "…的根是A. x=0 B. x=1 …"：A. 紧贴中文也要能拆。
    stem, options = extract_and_format_options(
        "方程x(x-1)=0的根是A. x=0 B. x=1 "
        "C. $x_1=0$ $x_2=-1$ D. $x_1=0$ $x_2=1$"
    )
    assert stem == "方程x(x-1)=0的根是"
    assert [label for label, _ in options] == ["A", "B", "C", "D"]
    assert options[0][1] == "x=0"
    # 不足 3 个标签不误拆。
    stem, options = extract_and_format_options("点A.在圆上，则B. 不成立")
    assert options == []
    assert stem == "点A.在圆上，则B. 不成立"


def test_visible_length_counts_rendered_chars() -> None:
    from question_bank.exporters.base_exporter import _visible_length

    assert _visible_length("$\\frac{3}{2}$") == 3
    assert _visible_length("$x_{1} = 0$ $x_{2} = - 1$") == 9
    assert _visible_length("不能确定") == 4
