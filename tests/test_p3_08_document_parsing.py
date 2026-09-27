from __future__ import annotations

import io
from pathlib import Path
import zipfile

import fitz
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image


def _save_docx_bytes(document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _number_paragraph(
    paragraph, num_id: int, *, start: int | None = None, label: str = "%1."
):
    paragraph._element.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )
    if start is not None:
        root = paragraph.part.numbering_part.element
        root.append(
            parse_xml(
                f'<w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{num_id}">'
                f'<w:lvl w:ilvl="0"><w:start w:val="{start}"/><w:numFmt w:val="decimal"/><w:lvlText w:val="{label}"/></w:lvl></w:abstractNum>'
            )
        )
        root.append(
            parse_xml(
                f'<w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="{num_id}">'
                f'<w:abstractNumId w:val="{num_id}"/></w:num>'
            )
        )


def _parse_synthetic_docx(document, tmp_path):
    from backend.document_parsing import parse_docx_question_blocks
    from backend.config_workspace.sources import (
        _strip_embedded_question_section_heading,
        _align_question_type_with_visible_blank,
    )

    blocks = parse_docx_question_blocks(
        _save_docx_bytes(document),
        temporary_root=tmp_path,
        asset_root=tmp_path / "assets",
    )
    for block in blocks:
        _strip_embedded_question_section_heading(block)
        _align_question_type_with_visible_blank(block)
    return {block["question_id"]: block for block in blocks}


def _synthetic_picture(paragraph, color="navy"):
    data = io.BytesIO()
    Image.new("RGB", (80, 40), color).save(data, format="PNG")
    return paragraph.add_run().add_picture(io.BytesIO(data.getvalue()), width=Inches(1))


def test_actual_word_subquestion_start_and_style_are_preserved(tmp_path):
    document = Document()
    document.add_paragraph("1. 选择正确答案。 A.甲 B.乙 C.丙 D.丁")
    document.add_paragraph("二、计算题")
    document.add_paragraph("12. 计算以下各式。")
    _number_paragraph(
        document.add_paragraph("√(50) - √(2)；"), 81, start=4, label="(%1)"
    )
    _number_paragraph(
        document.add_paragraph("√(72) - √(8)。"), 82, start=8, label="(%1)"
    )
    blocks = _parse_synthetic_docx(document, tmp_path)
    assert list(blocks) == ["Q1", "Q12"]
    assert blocks["Q1"]["question_type"] == "choice"
    assert (
        "(4)" in blocks["Q12"]["question_text"]
        and "(8)" in blocks["Q12"]["question_text"]
    )
    assert "1. √" not in blocks["Q12"]["question_text"]


def test_picture_in_question_table_retains_table_cells_and_image(tmp_path):
    document = Document()
    document.add_paragraph("1. 请根据表格及配图求阴影面积。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "图形"
    table.cell(0, 1).text = "边长"
    _synthetic_picture(table.cell(1, 0).paragraphs[0])
    table.cell(1, 1).text = "5 cm"
    document.add_paragraph("2. 求 3 + 4 的值。")
    blocks = _parse_synthetic_docx(document, tmp_path)
    assert (
        "table" in blocks["Q1"]["question_html"]
        and "5 cm" in blocks["Q1"]["question_text"]
    )
    assert len(blocks["Q1"]["image_paths"]) == 1
    assert not blocks["Q2"]["image_paths"]
