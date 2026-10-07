from __future__ import annotations

import json
import os
import sqlite3
import subprocess
import tempfile
import warnings
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path

import pytest


warnings.filterwarnings(
    "ignore",
    message="Using `httpx` with `starlette.testclient` is deprecated.*",
)

from fastapi.testclient import TestClient

from question_bank.database.schema import initialize_database
from question_bank.services.question_read_service import QuestionBankReadService
from question_bank.solution_evidence import (
    CoreResolution,
    QuestionSolutionEvidence,
    SolutionEvidenceRepository,
)
from question_bank.training_criteria import (
    QuestionAnalysisInputLoader,
    solution_evidence_source_content_hash,
)


@pytest.fixture
def question_bank_fixture(
    tmp_path: Path,
    question_bank_database,
) -> tuple[QuestionBankReadService, Path, bytes]:
    db_path = tmp_path / "question_bank.db"
    question_bank_database(db_path, taxonomy_revision=3)

    with closing(sqlite3.connect(db_path)) as conn:
        conn.executemany(
            """
            INSERT INTO papers (
                id, title, source_file, content_fingerprint, import_status, created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    1,
                    "Empty",
                    "C:/private/empty.docx",
                    "empty-secret",
                    "success",
                    "2026-01-01 00:00:00",
                ),
                (
                    2,
                    "Newest",
                    "C:/private/newest.docx",
                    "newest-secret",
                    "success",
                    "2026-01-02 00:00:00",
                ),
            ],
        )
        conn.executemany(
            """
            INSERT INTO questions (
                id, paper_id, question_number, question_text, is_deleted
            ) VALUES (?, ?, ?, ?, ?)
            """,
            [
                (1, 2, "1", "Active question", 0),
                (2, 2, "2", "Deleted question", 1),
            ],
        )
        conn.executemany(
            """
            INSERT INTO question_tags (question_id, tag_type, tag_value)
            VALUES (?, ?, ?)
            """,
            [
                (1, "knowledge_point", "一次函数"),
                (2, "ability", "Deleted tag"),
            ],
        )
        conn.commit()

    before = db_path.read_bytes()
    return QuestionBankReadService(db_path), db_path, before


def test_skill_routes_remain_read_only_and_optional_list_fields(question_bank_fixture):
    service, db_path, _ = question_bank_fixture
    client = _question_bank_client(service, question_bank_db_path=db_path)
    plain = client.get('/api/question-bank/questions').json()['items'][0]
    assert 'skills' not in plain and 'skill_hits' not in plain
    enriched = client.get('/api/question-bank/questions?include_skills=true').json()['items'][0]
    assert enriched == {**plain, 'skills': [], 'skill_hits': [], 'evidence_point_count': 0, 'duplicate_members': []}
    assert client.get('/api/question-bank/questions?skill_keys=sk_TEST_missing').json()['total'] == 0
    assert client.get('/api/question-bank/facets?skill_keys=sk_TEST_missing').json()['question_types'] == []
    assert client.get('/api/question-bank/skill-index?curriculum_volume_id=invalid').status_code == 422
    index = client.get('/api/question-bank/skill-index?curriculum_volume_id=bnu24-math-g8-upper')
    assert index.status_code == 200 and index.json()['model_calls'] == 0
    assert client.get('/api/question-bank/papers').json()['items'][0]['skill_unlinked_question_count'] == 1


def test_chapter_exam_profile_route_and_question_ids_filter(question_bank_fixture):
    service, db_path, _ = question_bank_fixture
    client = _question_bank_client(service, question_bank_db_path=db_path)
    assert client.get('/api/question-bank/chapter-exam-profile?curriculum_volume_id=invalid').status_code == 422
    response = client.get('/api/question-bank/chapter-exam-profile?curriculum_volume_id=bnu24-math-g8-upper')
    assert response.status_code == 200
    payload = response.json()
    assert payload['curriculum_volume_id'] == 'bnu24-math-g8-upper'
    assert payload['model_calls'] == 0
    assert {row['stage'] for row in payload['stages']} == {'midterm', 'final'}
    assert isinstance(payload['chapters'], list)
    selected = client.get('/api/question-bank/questions?question_ids=1&tag_status=all').json()
    assert [item['id'] for item in selected['items']] == [1]
    assert client.get('/api/question-bank/questions?question_ids=999&tag_status=all').json()['total'] == 0
    assert client.get('/api/question-bank/questions?' + '&'.join(
        f'question_ids={number}' for number in range(1, 502)
    )).status_code == 422


@pytest.mark.parametrize(("original", "expected"), [
    ("（8分）求未知数的值。", "求未知数的值。"),
    ("( 8 分 ) 求未知数的值。", "求未知数的值。"),
    ("（2.5分）求未知数的值。", "求未知数的值。"),
    ("<b>（</b><i>8</i>分）求未知数的值。", "<b></b><i></i>求未知数的值。"),
    ("（1）求未知数的值。", "（1）求未知数的值。"),
    ("小明跑了5分钟。", "小明跑了5分钟。"),
    ("游戏答对一题得3分。", "游戏答对一题得3分。"),
    ("函数 $f(x)=(1)/(2)$。", "函数 $f(x)=(1)/(2)$。"),
    ("求未知数的值。（3分）", "求未知数的值。（3分）"),
])
def test_stem_display_score_rule_preserves_body(original, expected):
    from question_bank.services.rich_content_service import strip_question_source_score

    assert strip_question_source_score(original) == expected
    assert strip_question_source_score(expected) == expected


def test_public_stems_hide_source_scores_without_changing_storage_or_analysis(question_bank_fixture, monkeypatch):
    from copy import deepcopy
    from docx import Document
    from question_bank.services.rich_content_service import load_question_rich_content, save_question_rich_content
    from tools import maintain_question_bank as maintenance

    _, db_path, _ = question_bank_fixture
    root = db_path.parent
    original = "（8 分）小明跑了5分钟，求未知数的值。"
    with closing(sqlite3.connect(db_path)) as conn:
        conn.execute("UPDATE questions SET question_text=?,answer_text='（8分）答案为2。',question_type='解答题' WHERE id=1", (original,))
        conn.commit()
    document = Document()
    stem = document.add_paragraph()
    stem.add_run("1. （").bold = True
    stem.add_run("8").italic = True
    stem.add_run(" 分）小明跑了5分钟，求未知数的值。").underline = True
    answer = document.add_paragraph("1. （8分）答案为2。")
    path = save_question_rich_content(1, question_blocks=[{"text": stem.text, "xml": stem._p.xml}],
                                      answer_blocks=[{"text": answer.text, "xml": answer._p.xml}],
                                      root=root / "question_bank/rich_content")
    service = QuestionBankReadService(db_path, data_root=root)
    loader = QuestionAnalysisInputLoader(db_path=db_path, data_root=root)
    (raw_input,) = loader.load([1])
    cached_original = deepcopy(load_question_rich_content(1, root=root / "question_bank/rich_content"))
    before = db_path.read_bytes(), path.read_bytes()
    with _question_bank_client(service, question_bank_db_path=db_path, data_root=root) as client:
        listed = client.get('/api/question-bank/questions').json()['items'][0]
        detail = client.get('/api/question-bank/questions/1').json()
    for item in (listed, detail, service.get_questions([1])[0]):
        assert item['question_text'] == "小明跑了5分钟，求未知数的值。"
        block = item['rich_content']['question_blocks'][0]
        assert "8 分" not in block['text'] and "8 分" not in block['html']
        assert "5分钟" in block['html'] and "<u>" in block['html']
        assert "（8分）" in item['rich_content']['answer_blocks'][0]['text']
        assert item['revision'] == service.get_questions_for_export([1])[0]['revision']
    assert service.get_questions_for_export([1])[0]['question_text'] == original
    (after_input,) = loader.load([1])
    assert after_input == raw_input and after_input.tagging_context.question_text == original
    assert cached_original == load_question_rich_content(1, root=root / "question_bank/rich_content")
    report = maintenance.preview_source_scores(db_path, root)
    assert report['changed_stem_questions'] == report['changed_rich_stem_questions'] == 1
    assert report['questions_using_original_answer_space'] == 1
    assert report['unexpected_xml_structure_changes'] == report['blocked_rich_questions'] == 0
    assert report['source_content_unchanged'] and report['analysis_unchanged']
    assert report['answers_changed'] == 0 and report['applied'] is False
    monkeypatch.setattr('sys.argv', ['maintain_question_bank.py', 'source-scores', '--apply', '--data-root', str(root)])
    with pytest.raises(SystemExit) as error:
        maintenance.main()
    assert error.value.code == 2
    assert (db_path.read_bytes(), path.read_bytes()) == before


@pytest.mark.parametrize("changed_content", ["stem", "answer", "media", "rich_body", "media_body"])
def test_solution_evidence_route_returns_latest_point_level_union(
    question_bank_fixture,
    changed_content: str,
    monkeypatch,
) -> None:
    service, db_path, _ = question_bank_fixture
    data_root = db_path.parent / "data"
    from question_bank.services.rich_content_service import save_question_rich_content
    from question_bank.taxonomy.curriculum_catalog import curriculum_volume

    volume = curriculum_volume(volume_id="bnu24-math-g8-upper")
    service = QuestionBankReadService(db_path, data_root=data_root)
    with sqlite3.connect(db_path) as conn:
        conn.execute("UPDATE papers SET grade=?,semester=?,textbook_version=? WHERE id=2",
                     (volume['grade'], volume['semester'], volume['textbook_version']))
        conn.execute("UPDATE questions SET question_type='解答题',difficulty='3',is_deleted=0")
        conn.execute("INSERT INTO questions(id,paper_id,question_number,question_text,question_type,difficulty) "
                     "VALUES(3,2,'3','TEST-missing-analysis','解答题','3')")
        conn.executemany("INSERT INTO question_tags(question_id,tag_type,tag_value) VALUES(?,?,?)", [
            (qid, tag_type, value) for qid in (1, 2, 3)
            for tag_type, value in (("ability", "运算求解"), ("exam_scope", "八年级上册"))
        ])
        conn.executemany("INSERT INTO question_tags(question_id,tag_type,tag_value) VALUES(?,'knowledge_point','一次函数')",
                         [(2,), (3,)])
        conn.executemany("INSERT INTO grading_question_links "
            "(grading_session_id,source_question_id,bank_question_id,link_method,status) "
            "VALUES ('7',?,?, 'TEST', 'confirmed')", [(str(qid), qid) for qid in (1, 2, 3)])
        if changed_content == "media_body":
            image = data_root / 'question_bank/extracted_images/TEST-source.png'
            image.parent.mkdir(parents=True)
            image.write_bytes(b'TEST-original-image-content')
            conn.execute("UPDATE questions SET has_images=1,image_paths=? WHERE id=1",
                         (json.dumps(['question_bank/extracted_images/TEST-source.png']),))
    if changed_content == "rich_body":
        save_question_rich_content(1, question_blocks=[{'text': 'TEST-original-rich-content'}],
                                   root=data_root / 'question_bank/rich_content')
    client = _question_bank_client(
        service,
        question_bank_db_path=db_path,
        data_root=data_root,
    )

    empty = client.get("/api/question-bank/questions/1/solution-evidence")
    assert empty.status_code == 200
    assert empty.json() == {
        "question_id": 1,
        "available": False,
        "evidence_version_id": None,
        "status": None,
        "evidence": None,
        "part_assessments": [],
        "assessment_revision": None,
    }

    class Resolver:
        def resolve(self, fine_term_id: str) -> CoreResolution:
            if fine_term_id == "fine-direct":
                return CoreResolution(
                    status="resolved",
                    stable_keys=("kp_equation",),
                    reason="test mapping",
                )
            return CoreResolution(status="unmapped", reason="test unmapped")

    evidence = QuestionSolutionEvidence.from_model_dict(
        {
            "schema_version": "question-solution-evidence-v1",
            "question_id": 1,
            "parts": [
                {
                    "part_id": "part-1",
                    "label": "（1）",
                    "response_mode": "process_required",
                    "canonical_answer": "x=2",
                    "accepted_forms": ["x = 2"],
                    "full_answer": "移项后求得 x=2。",
                    "proof_obligations": [],
                    "visual_requirements": [],
                    "deduction_policy": ["没有等价变形过程则该点未达成"],
                    "allow_alternative_methods": True,
                    "evidence_points": [
                        {
                            "evidence_point_id": "point-1",
                            "target": "求出方程的解",
                            "observable_evidence": "给出等价变形并写出正确解。",
                            "fine_term_links": [
                                {
                                    "fine_term_id": "fine-direct",
                                    "fine_term_name": "一元一次方程求解",
                                    "role": "direct",
                                },
                                {
                                    "fine_term_id": "fine-support",
                                    "fine_term_name": "规范移项",
                                    "role": "supporting_prerequisite",
                                },
                            ],
                            "equivalent_rules": [],
                            "counterexamples": [],
                        }
                    ],
                }
            ],
            "auxiliary_rules": [],
            "rationale": "按踩分点拆分。",
            "confidence": 0.9,
        },
        question_id=1,
        source_content_hash=solution_evidence_source_content_hash(
            QuestionAnalysisInputLoader(
                db_path=db_path,
                data_root=data_root,
            ).load((1,))[0]
        ),
        resolver=Resolver(),
    )
    version_id = SolutionEvidenceRepository(db_path).save(
        evidence,
        source_kind="combined_model",
        source_reference="analysis:test-route:1",
        created_by="model:synthetic",
    )

    response = client.get("/api/question-bank/questions/1/solution-evidence")
    assert response.status_code == 200
    body = response.json()
    assert body["available"] is True
    assert body["evidence_version_id"] == version_id
    assert (
        body["evidence"]["parts"][0]["evidence_points"][0]["target"] == "求出方程的解"
    )
    union = body["evidence"]["whole_question_classification"]
    assert union["direct_fine_terms"] == [
        {"fine_term_id": "fine-direct", "fine_term_name": "一元一次方程求解"}
    ]
    assert union["supporting_prerequisite_fine_terms"] == [
        {"fine_term_id": "fine-support", "fine_term_name": "规范移项"}
    ]
    assert union["resolved_core_node_ids"] == ["kp_equation"]
    assert union["unmapped_fine_term_ids"] == ["fine-support"]

    second_input = QuestionAnalysisInputLoader(db_path=db_path, data_root=data_root).load((2,))[0]
    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA ignore_check_constraints=ON")
        conn.execute("""INSERT INTO question_solution_evidence_versions
            (evidence_version_id,question_id,source_content_hash,schema_version,content_hash,
             evidence_json,status,source_kind,source_reference,created_by)
            VALUES (?,2,?,'question-solution-evidence-v2',?,'{','proposed','combined_model','TEST','TEST')""",
            ('f' * 64, solution_evidence_source_content_hash(second_input), 'e' * 64))
        assert conn.execute("SELECT COUNT(*) FROM training_criterion_heads").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM training_criterion_versions").fetchone()[0] == 0
    before_read = db_path.read_bytes()
    complete = client.get('/api/question-bank/questions?analysis_status=complete').json()
    assert [item['id'] for item in complete['items']] == [1]
    assert client.get('/api/question-bank/question-refs?analysis_status=complete').json()['total'] == 1
    incomplete = client.get('/api/question-bank/questions?analysis_status=incomplete').json()
    assert {item['id'] for item in incomplete['items']} == {2, 3}
    needs_review = client.get('/api/question-bank/questions?criteria_needs_review=true').json()
    assert [item['id'] for item in needs_review['items']] == [2]
    paper = next(item for item in client.get('/api/question-bank/papers').json()['items'] if item['id'] == 2)
    assert paper['question_count'] == 3
    assert paper['evidence_question_count'] == paper['criteria_question_count'] == paper['complete_analysis_count'] == 1
    assert paper['criteria_needs_review_count'] == 1
    plain = client.get('/api/question-bank/questions').json()
    assert next(item for item in plain['items'] if item['id'] == 1)['criteria_needs_review'] is False
    assert client.get('/api/question-bank/facets?analysis_status=complete').json()['question_types'] == [
        {'value': '解答题', 'count': 1}]
    assert client.get('/api/question-bank/skill-index?curriculum_volume_id=bnu24-math-g8-upper').json()[
        'unlinked']['no_usable_evidence'] == 2
    assert 1 in service._skill_snapshot()['evidence_versions']
    if changed_content in {'rich_body', 'media_body'}:
        other_root = db_path.parent / 'TEST-other-data-root'
        other_service = QuestionBankReadService(db_path, data_root=other_root)
        other_paper = next(item for item in other_service.list_papers() if item['id'] == 2)
        assert other_paper['evidence_question_count'] == other_paper['criteria_question_count'] == 0
        assert next(item for item in service.list_papers() if item['id'] == 2)['complete_analysis_count'] == 1
    # Browser and repair planning agree on actual usable material without writes.
    status = service.session_analysis_status(7)
    assert status['evidence_count'] == status['criteria_count'] == status['complete_count'] == 1
    assert status['incomplete_question_ids'] == [2, 3]
    assert db_path.read_bytes() == before_read

    from question_bank.services import standard_difficulty

    difficulty_input = QuestionAnalysisInputLoader(db_path=db_path, data_root=data_root).load((1,))[0]
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO question_part_difficulty_features (
                question_id, part_id, features_json, formula_difficulty,
                formula_version, source_content_hash, is_active
            ) VALUES (1, 'part-1', '{"evidence":"synthetic estimate"}', 3,
                      'std-difficulty-v1', ?, 1)
            """,
            (
                standard_difficulty.question_content_fingerprint(
                    difficulty_input
                ),
            ),
        )
    supplemented = client.get("/api/question-bank/questions/1/solution-evidence").json()
    assert supplemented['available'], supplemented
    assert supplemented["part_assessments"][0]["difficulty"] == 3
    assert supplemented["part_assessments"][0]["source"] == "formula"
    assert isinstance(supplemented["assessment_revision"], str)
    assert supplemented["assessment_revision"]

    # Warm every read after the last database write, then change only the files.
    client.get('/api/question-bank/questions?analysis_status=complete')
    client.get('/api/question-bank/question-refs?analysis_status=complete')
    client.get('/api/question-bank/questions')
    client.get('/api/question-bank/papers')
    client.get('/api/question-bank/facets?analysis_status=complete')
    client.get('/api/question-bank/skill-index?curriculum_volume_id=bnu24-math-g8-upper')
    unchanged_database = db_path.read_bytes()
    with sqlite3.connect(db_path) as conn:
        if changed_content == "stem":
            conn.execute(
                "UPDATE questions SET question_text = 'Changed stem' WHERE id = 1"
            )
        elif changed_content == "answer":
            conn.execute(
                "UPDATE questions SET answer_text = 'Changed answer' WHERE id = 1"
            )
        elif changed_content == "media":
            image = data_root / "question_bank" / "extracted_images" / "changed.png"
            image.parent.mkdir(parents=True, exist_ok=True)
            image.write_bytes(b"synthetic changed image")
            conn.execute(
                """
                UPDATE questions
                SET has_images = 1, image_paths = ?
                WHERE id = 1
                """,
                (json.dumps(["question_bank/extracted_images/changed.png"]),),
            )
        elif changed_content == "rich_body":
            save_question_rich_content(1, question_blocks=[{'text': 'TEST-changed-rich-content-plus-new-requirement'}],
                                       root=data_root / 'question_bank/rich_content')
        else:
            image.write_bytes(b'TEST-changed-image-content-plus-new-requirement')
        conn.commit()
    if changed_content in {'rich_body', 'media_body'}:
        assert db_path.read_bytes() == unchanged_database

    stale = client.get("/api/question-bank/questions/1/solution-evidence")
    assert stale.status_code == 200
    assert stale.json() == {
        "question_id": 1,
        "available": False,
        "evidence_version_id": version_id,
        "status": "stale",
        "evidence": None,
        "part_assessments": [],
        "assessment_revision": None,
    }
    assert client.get('/api/question-bank/questions?analysis_status=complete').json()['total'] == 0
    needs_review = client.get('/api/question-bank/questions?criteria_needs_review=true').json()
    assert {item['id'] for item in needs_review['items']} == {1, 2}
    paper = next(item for item in client.get('/api/question-bank/papers').json()['items'] if item['id'] == 2)
    assert paper['evidence_question_count'] == paper['criteria_question_count'] == paper['complete_analysis_count'] == 0
    assert paper['criteria_needs_review_count'] == 2
    assert client.get('/api/question-bank/question-refs?analysis_status=complete').json()['total'] == 0
    plain = client.get('/api/question-bank/questions').json()
    assert next(item for item in plain['items'] if item['id'] == 1)['criteria_needs_review'] is True
    assert client.get('/api/question-bank/facets?analysis_status=complete').json()['question_types'] == []
    assert client.get('/api/question-bank/skill-index?curriculum_volume_id=bnu24-math-g8-upper').json()[
        'unlinked']['no_usable_evidence'] == 3
    assert 1 in service._skill_snapshot()['no_usable']
    assert client.get('/api/question-bank/questions?analysis_status=complete').json()['total'] == 0
    assert next(item for item in service.list_papers() if item['id'] == 2)['complete_analysis_count'] == 0
    if changed_content in {'rich_body', 'media_body'}:
        from question_bank.services import question_read_service as read_module

        monkeypatch.setattr(read_module, '_skill_asset_manifest', lambda _root: None)
        if changed_content == 'rich_body':
            save_question_rich_content(1, question_blocks=[{'text': 'TEST-original-rich-content'}],
                                       root=data_root / 'question_bank/rich_content')
        else:
            image.write_bytes(b'TEST-original-image-content')
        assert client.get('/api/question-bank/questions?analysis_status=complete').json()['total'] == 1
        assert client.get('/api/question-bank/question-refs?analysis_status=complete').json()['total'] == 1
        assert next(item for item in service.list_papers() if item['id'] == 2)['complete_analysis_count'] == 1
        if changed_content == 'rich_body':
            save_question_rich_content(1, question_blocks=[{'text': 'TEST-changed-rich-content-again'}],
                                       root=data_root / 'question_bank/rich_content')
        else:
            image.write_bytes(b'TEST-changed-image-content-again')
        assert client.get('/api/question-bank/questions?analysis_status=complete').json()['total'] == 0
        assert client.get('/api/question-bank/question-refs?analysis_status=complete').json()['total'] == 0
        assert next(item for item in service.list_papers() if item['id'] == 2)['complete_analysis_count'] == 0
        assert db_path.read_bytes() == unchanged_database


def _question_bank_client(
    service: QuestionBankReadService,
    *,
    question_bank_db_path: Path | None = None,
    data_root: Path | None = None,
    raise_server_exceptions: bool = True,
) -> TestClient:
    from backend.api.app import create_app
    from backend.api.dependencies import (
        get_data_root,
        get_question_bank_db_path,
        get_question_bank_read_service,
    )

    app = create_app()
    app.dependency_overrides[get_question_bank_read_service] = lambda: service
    if question_bank_db_path is not None:
        app.dependency_overrides[get_question_bank_db_path] = lambda: (
            question_bank_db_path
        )
    if data_root is not None:
        app.dependency_overrides[get_data_root] = lambda: data_root
    return TestClient(app, raise_server_exceptions=raise_server_exceptions)


def _seed_combination_filter_questions(db_path: Path) -> None:
    papers = [
        (1, "Target", "2026", "期中", "八年级"),
        (2, "Wrong year", "2025", "期中", "八年级"),
        (3, "Wrong exam", "2026", "期末", "八年级"),
        (4, "Wrong grade", "2026", "期中", "七年级"),
        (99, "Wrong paper", "2026", "期中", "八年级"),
    ]
    questions = [
        (
            100,
            1,
            "7",
            "选择题",
            "needle target [[IMAGE:C:/private/stem.png]]",
            "answer [[IMAGE:C:/private/answer.png]]",
            "6",
            0,
        ),
        (101, 99, "7", "选择题", "needle wrong paper", "answer", "6", 0),
        (102, 2, "7", "选择题", "needle wrong year", "answer", "6", 0),
        (103, 3, "7", "选择题", "needle wrong exam", "answer", "6", 0),
        (104, 4, "7", "选择题", "needle wrong grade", "answer", "6", 0),
        (105, 1, "7", "解答题", "needle wrong type", "answer", "6", 0),
        (106, 1, "7", "选择题", "needle wrong difficulty", "answer", "9", 0),
        (107, 1, "7", "选择题", "needle incomplete tags", "answer", None, 0),
        (108, 1, "7", "选择题", "needle wrong scope", "answer", "6", 0),
        (109, 1, "7", "选择题", "needle wrong knowledge", "answer", "6", 0),
        (110, 1, "7", "选择题", "unrelated text", "different answer", "6", 0),
        (111, 1, "8", "选择题", "needle wrong number", "answer", "6", 0),
        (112, 1, "7", "选择题", "needle deleted", "answer", "6", 1),
    ]
    with sqlite3.connect(db_path) as conn:
        conn.executemany(
            """
            INSERT INTO papers (
                id, title, source_file, year, exam_type, grade, import_status,
                content_fingerprint, created_at
            ) VALUES (?, ?, 'C:/private/paper.docx', ?, ?, ?, 'success',
                      'paper-secret', '2026-01-01 00:00:00')
            """,
            papers,
        )
        conn.executemany(
            """
            INSERT INTO questions (
                id, paper_id, question_number, question_type, question_text,
                answer_text, source_file, image_paths, difficulty, is_deleted,
                created_at
            ) VALUES (?, ?, ?, ?, ?, ?, 'C:/private/question.docx',
                      '["C:/private/list.png"]', ?, ?, '2026-02-01 00:00:00')
            """,
            questions,
        )
        for question_id, *_ in questions:
            tags = [
                (
                    question_id,
                    "knowledge_point",
                    "二次函数" if question_id == 109 else "一次函数",
                ),
                (question_id, "ability", "运算求解"),
                (
                    question_id,
                    "exam_scope",
                    "九年级下册" if question_id == 108 else "八年级上册",
                ),
            ]
            if question_id != 107:
                tags.append((question_id, "student_level", "中档提升"))
            conn.executemany(
                """
                INSERT INTO question_tags (question_id, tag_type, tag_value)
                VALUES (?, ?, ?)
                """,
                tags,
            )
        conn.execute(
            """
            INSERT INTO question_fingerprints (
                question_id, base_fingerprint, style_features_json
            ) VALUES (100, 'question-secret', '{"private": "C:/private/fingerprint"}')
            """
        )
        conn.commit()


def test_questions_combined_filters_and_public_projection(
    tmp_path: Path,
    question_bank_database,
) -> None:
    db_path = tmp_path / "question_bank.db"
    question_bank_database(db_path, taxonomy_revision=3)
    _seed_combination_filter_questions(db_path)
    with sqlite3.connect(db_path) as conn:
        conn.executemany(
            """
            INSERT INTO question_error_patterns (
                question_id, category, pattern, status, source
            ) VALUES (?, ?, ?, ?, 'TEST')
            """,
            [
                (100, "方法与思路", "TEST-错因", "confirmed"),
                (101, "方法与思路", "TEST-候选错因", "candidate"),
            ],
        )
    client = _question_bank_client(QuestionBankReadService(db_path))

    response = client.get(
        "/api/question-bank/questions",
        params=[
            ("question_number", "7"),
            ("keyword", "needle"),
            ("knowledge_point", "一次函数"),
            ("difficulty_min", "5"),
            ("difficulty_max", "7"),
            ("question_types", "选择题"),
            ("question_types", "填空题"),
            ("paper_ids", "1"),
            ("paper_ids", "2"),
            ("paper_ids", "3"),
            ("paper_ids", "4"),
            ("years", "2026"),
            ("years", "2027"),
            ("exam_types", "期中"),
            ("exam_types", "模拟"),
            ("grades", "八年级"),
            ("grades", "九年级"),
            ("exam_scopes", "八年级上册"),
            ("exam_scopes", "八年级下册"),
            ("tag_status", "tagged"),
            ("sort", "newest"),
        ],
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 1
    assert [item["id"] for item in payload["items"]] == [100]
    item = payload["items"][0]
    assert item["question_text"] == "needle target "
    assert item["answer_text"] == "answer "
    assert item["asset_urls"] == [
        "/api/question-bank/questions/100/assets/0",
        "/api/question-bank/questions/100/assets/1",
        "/api/question-bank/questions/100/assets/2",
    ]
    serialized = response.text
    for forbidden in (
        "source_file",
        "fingerprint",
        "image_paths",
        "C:/private",
        "[[IMAGE:",
        "paper-secret",
        "question-secret",
    ):
        assert forbidden not in serialized

    by_category = client.get(
        "/api/question-bank/questions",
        params={"error_pattern_categories": "方法与思路"},
    ).json()
    assert [item["id"] for item in by_category["items"]] == [100]
    assert client.get(
        "/api/question-bank/questions",
        params={"error_pattern_categories": "审题与条件"},
    ).json()["total"] == 0
    assert client.get("/api/question-bank/facets").json()[
        "error_pattern_categories"
    ] == [{"value": "方法与思路", "count": 1}]


def _json_strings(value: object):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for key, item in value.items():
            yield str(key)
            yield from _json_strings(item)
    elif isinstance(value, list):
        for item in value:
            yield from _json_strings(item)


def _json_keys(value: object):
    if isinstance(value, dict):
        for key, item in value.items():
            yield str(key)
            yield from _json_keys(item)
    elif isinstance(value, list):
        for item in value:
            yield from _json_keys(item)


@dataclass(frozen=True)
class _QuestionBankMediaSeed:
    client: TestClient
    db_path: Path
    data_root: Path
    asset_root: Path
    preview_root: Path
    question_id: int
    ordered_assets: tuple[tuple[bytes, str], ...]
    current_preview_bytes: bytes
    db_before: bytes
    sqlite_sidecars_before: dict[str, bytes | None]


def _sqlite_sidecar_snapshot(db_path: Path) -> dict[str, bytes | None]:
    snapshot: dict[str, bytes | None] = {}
    for suffix in ("-wal", "-shm"):
        sidecar = Path(f"{db_path}{suffix}")
        snapshot[suffix] = sidecar.read_bytes() if sidecar.exists() else None
    return snapshot


@pytest.fixture
def question_bank_media_seed(
    tmp_path: Path,
    question_bank_database,
) -> _QuestionBankMediaSeed:
    db_path = tmp_path / "question_bank.db"
    data_root = tmp_path / "isolated-data-root"
    asset_root = data_root / "question_bank" / "extracted_images"
    preview_root = data_root / "question_bank" / "previews"
    asset_root.mkdir(parents=True)
    preview_root.mkdir(parents=True)
    question_bank_database(db_path)

    db_asset = asset_root / "set" / "shared-fallback.png"
    question_asset = asset_root / "question-marker.jpg"
    answer_asset = asset_root / "answer" / "answer-marker.webp"
    rich_asset = asset_root / "rich" / "rich.bmp"
    current_preview = preview_root / "current" / "shared-preview.jpg"
    old_preview = preview_root / "old-question.jpg"
    old_answer_preview = preview_root / "old-answer.jpg"
    for path, content in (
        (db_asset, b"asset-png-current"),
        (question_asset, b"asset-jpeg-question"),
        (answer_asset, b"asset-webp-answer"),
        (rich_asset, b"asset-bmp-rich"),
        (current_preview, b"preview-jpeg-current"),
        (old_preview, b"preview-jpeg-old"),
        (old_answer_preview, b"preview-jpeg-old-answer"),
        (
            preview_root / "shadow" / "shared-fallback.png",
            b"wrong-preview-search-root",
        ),
        (
            asset_root / "shadow" / "shared-preview.jpg",
            b"wrong-asset-search-root",
        ),
    ):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)

    outside_asset = tmp_path / "outside-private.png"
    outside_asset.write_bytes(b"outside-private")
    traversal_asset = tmp_path / "outside-traversal.png"
    traversal_asset.write_bytes(b"outside-traversal")
    unsafe_asset = asset_root / "unsafe.txt"
    unsafe_asset.write_bytes(b"unsafe-asset-type")
    outside_preview = asset_root / "outside-preview.jpg"
    outside_preview.write_bytes(b"outside-preview-root")
    unsafe_preview = preview_root / "unsafe-preview.txt"
    unsafe_preview.write_bytes(b"unsafe-preview-type")
    for folder in ("ambiguous-a", "ambiguous-b"):
        ambiguous = asset_root / folder / "ambiguous.png"
        ambiguous.parent.mkdir(parents=True)
        ambiguous.write_bytes(folder.encode())

    with closing(sqlite3.connect(db_path)) as conn:
        conn.executemany(
            """
            INSERT INTO papers (id, title, import_status)
            VALUES (?, ?, ?)
            """,
            [
                (1, "Active", "success"),
                (2, "Deleted", "deleted"),
            ],
        )
        conn.executemany(
            """
            INSERT INTO questions (
                id, paper_id, question_number, question_text, answer_text,
                image_paths, is_deleted
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    100,
                    1,
                    "100",
                    (
                        "Stem[[IMAGE:question_bank/extracted_images/"
                        "question-marker.jpg]]"
                    ),
                    "Answer[[IMAGE:C:/offline/answer-marker.webp]]",
                    json.dumps(["C:/offline/shared-fallback.png"]),
                    0,
                ),
                (
                    101,
                    1,
                    "101",
                    "Soft deleted",
                    None,
                    json.dumps(["question_bank/extracted_images/question-marker.jpg"]),
                    1,
                ),
                (
                    102,
                    2,
                    "102",
                    "Deleted parent",
                    None,
                    json.dumps(["question_bank/extracted_images/question-marker.jpg"]),
                    0,
                ),
                (
                    110,
                    1,
                    "110",
                    "Missing asset",
                    None,
                    json.dumps(["question_bank/extracted_images/missing-current.png"]),
                    0,
                ),
                (111, 1, "111", "Missing preview", None, "[]", 0),
                (
                    112,
                    1,
                    "112",
                    "Outside asset",
                    None,
                    json.dumps([str(outside_asset.resolve())]),
                    0,
                ),
                (
                    113,
                    1,
                    "113",
                    "Traversal asset",
                    None,
                    json.dumps(["../outside-traversal.png"]),
                    0,
                ),
                (
                    114,
                    1,
                    "114",
                    "Unsafe asset",
                    None,
                    json.dumps(["question_bank/extracted_images/unsafe.txt"]),
                    0,
                ),
                (
                    115,
                    1,
                    "115",
                    "Ambiguous asset",
                    None,
                    json.dumps(["C:/offline/ambiguous.png"]),
                    0,
                ),
                (116, 1, "116", "Outside preview", None, "[]", 0),
                (117, 1, "117", "Unsafe preview", None, "[]", 0),
            ],
        )
        conn.executemany(
            """
            INSERT INTO question_previews (
                id, question_id, preview_type, image_path, status, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    1,
                    100,
                    "question",
                    "question_bank/previews/old-question.jpg",
                    "ready",
                    "2026-01-01 00:00:00",
                ),
                (
                    2,
                    100,
                    "question",
                    "C:/offline/shared-preview.jpg",
                    "ready",
                    "2026-02-01 00:00:00",
                ),
                (
                    3,
                    100,
                    "answer",
                    "question_bank/previews/old-answer.jpg",
                    "ready",
                    "2026-01-01 00:00:00",
                ),
                (4, 100, "answer", None, "failed", "2026-02-01 00:00:00"),
                (
                    5,
                    111,
                    "question",
                    "question_bank/previews/missing-current.jpg",
                    "ready",
                    "2026-03-01 00:00:00",
                ),
                (
                    6,
                    116,
                    "question",
                    str(outside_preview),
                    "ready",
                    "2026-03-01 00:00:00",
                ),
                (
                    7,
                    117,
                    "question",
                    "question_bank/previews/unsafe-preview.txt",
                    "ready",
                    "2026-03-01 00:00:00",
                ),
                (
                    8,
                    101,
                    "question",
                    "question_bank/previews/old-question.jpg",
                    "ready",
                    "2026-03-01 00:00:00",
                ),
                (
                    9,
                    102,
                    "question",
                    "question_bank/previews/old-question.jpg",
                    "ready",
                    "2026-03-01 00:00:00",
                ),
            ],
        )
        conn.commit()
        conn.execute("PRAGMA journal_mode = DELETE").fetchone()

    rich_root = data_root / "question_bank" / "rich_content"
    rich_root.mkdir(parents=True)
    rich_root.joinpath("question_100.json").write_text(
        json.dumps(
            {
                "version": 3,
                "question_id": 100,
                "question_blocks": [
                    {
                        "text": (
                            "Rich[[IMAGE:question_bank/extracted_images/rich/rich.bmp]]"
                        ),
                        "image_relationships": {
                            "rId1": "question_bank/extracted_images/rich/rich.bmp"
                        },
                    }
                ],
                "answer_blocks": [],
            }
        ),
        encoding="utf-8",
    )

    client = _question_bank_client(
        QuestionBankReadService(db_path, data_root=data_root)
    )
    return _QuestionBankMediaSeed(
        client=client,
        db_path=db_path,
        data_root=data_root,
        asset_root=asset_root,
        preview_root=preview_root,
        question_id=100,
        ordered_assets=(
            (b"asset-png-current", "image/png"),
            (b"asset-jpeg-question", "image/jpeg"),
            (b"asset-webp-answer", "image/webp"),
            (b"asset-bmp-rich", "image/bmp"),
        ),
        current_preview_bytes=b"preview-jpeg-current",
        db_before=db_path.read_bytes(),
        sqlite_sidecars_before=_sqlite_sidecar_snapshot(db_path),
    )
