from __future__ import annotations

import json
import sqlite3
import warnings
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

warnings.filterwarnings(
    "ignore",
    message="Using `httpx` with `starlette.testclient` is deprecated.*",
)

from fastapi.testclient import TestClient


COMPLETE_RAW_JSON = json.dumps(
    {"grading_completeness": {"status": "complete"}},
    ensure_ascii=False,
)

PERSONAL_NARRATIVE = {
    "overall_comment": "整体表现稳定，丢分集中在证明题。",
    "question_analyses": [
        {
            "question_id": "Q2",
            "analysis": "证明过程缺关键步骤，建议写出中间结论。",
            "root_cause_type": "method",
        }
    ],
    "strengths": ["选择题全对，基础扎实"],
    "problems": [{"title": "证明书写不完整", "detail": "第 2 题丢了过程分"}],
    "suggestions": [
        {"title": "重做证明题", "timeframe": "本周", "detail": "练 3 道同类题"}
    ],
}

CLASS_NARRATIVE = {
    "key_findings": [
        {
            "title": "证明题得分率低",
            "detail": "第 2 题全班得分率不足一半",
            "severity": "high",
        }
    ],
    "common_issues": [
        {
            "title": "证明步骤缺失",
            "evidence": "S1、S2 均缺关键步骤",
            "teaching_action": "板书示范 10 分钟",
        }
    ],
    "student_notes": [
        {"alias": "S1", "note": "高但证明扣分", "suggestion": "保持", "flags": []},
        {
            "alias": "S2",
            "note": "客观题失分多",
            "suggestion": "建议当面了解",
            "flags": ["needs_review", "priority_talk"],
        },
    ],
    "grouping_advice": "按分数断层分层布置作业",
}
from backend.repositories.grading_database import open_grading_repositories


class FakeLLMClient:
    def __init__(self, narrative: dict | None = None, error: Exception | None = None):
        self.calls = 0
        self.requests: list[dict] = []
        self._narrative = dict(narrative or PERSONAL_NARRATIVE)
        self._error = error

    def json_from_text(self, prompt, extra_kwargs=None, **_kwargs):
        return self._respond(prompt, [], extra_kwargs, _kwargs)

    def json_from_images_once(self, prompt, image_blobs, extra_kwargs=None, **kwargs):
        return self._respond(prompt, image_blobs, extra_kwargs, kwargs)

    def _respond(self, prompt, images, extra_kwargs, kwargs):
        self.calls += 1
        self.requests.append(
            {"prompt": prompt, "images": images, "options": extra_kwargs, **kwargs}
        )
        if self._error is not None:
            raise self._error
        return dict(self._narrative)


def _seed_analysis_session(db, root: Path) -> int:
    """三名学生：张三/李四正常参考，王五缺考。rubric 两题共 100 分。"""
    config_dir = root / "config" / "uploaded"
    config_dir.mkdir(parents=True, exist_ok=True)
    rubric = {
        "exam_title": "单元测试",
        "subject": "数学",
        "total_score": 100,
        "questions": [
            {
                "question_id": "Q1",
                "question_type": "choice",
                "max_score": 60,
                "stem_summary": "识别轴对称图形",
            },
            {
                "question_id": "Q2",
                "question_type": "proof",
                "max_score": 40,
                "stem_summary": "证明线段数量关系",
            },
        ],
    }
    answer_key = {
        "questions": [
            {"question_id": "Q1", "canonical_answer": "B"},
            {"question_id": "Q2", "canonical_answer": "AB=BD+DH"},
        ]
    }
    rubric_path = config_dir / "rubric.json"
    answer_path = config_dir / "answer_key.json"
    rubric_path.write_text(json.dumps(rubric, ensure_ascii=False), encoding="utf-8")
    answer_path.write_text(json.dumps(answer_key, ensure_ascii=False), encoding="utf-8")
    session_id = db.sessions.create_grading_session(
        "单元测试", str(rubric_path), str(answer_path)
    )

    with sqlite3.connect(db.db_path) as conn:
        students = {}
        for code, name in (("001", "张三"), ("002", "李四"), ("003", "王五")):
            students[name] = int(
                conn.execute(
                    "INSERT INTO students (student_code, name, class_name) VALUES (?, ?, ?)",
                    (code, name, "1 班"),
                ).lastrowid
            )
        for name, total, awarded, review in (
            ("张三", 100, 90, 0),
            ("李四", 100, 50, 1),
        ):
            student_id = students[name]
            paper_id = int(
                conn.execute(
                    """
                    INSERT INTO exam_papers (
                        session_id, front_image, back_image, student_id,
                        match_status, processing_status
                    ) VALUES (?, '', '', ?, 'matched', 'graded')
                    """,
                    (session_id, student_id),
                ).lastrowid
            )
            result_id = int(
                conn.execute(
                    """
                    INSERT INTO session_results (
                        session_id, student_id, paper_id, total_score, student_score,
                        needs_human_review, raw_json
                    ) VALUES (?, ?, ?, 100, ?, ?, ?)
                    """,
                    (
                        session_id,
                        student_id,
                        paper_id,
                        awarded,
                        review,
                        COMPLETE_RAW_JSON,
                    ),
                ).lastrowid
            )
            conn.execute(
                """
                INSERT INTO session_details (
                    result_id, question_id, score_awarded, deduction_reason,
                    knowledge_ids, error_category, error_summary
                ) VALUES (?, 'Q1', ?, ?, '["UNKNOWN"]', ?, ?)
                """,
                (
                    result_id,
                    60 if awarded >= 60 else 30,
                    None if awarded >= 60 else "objective_answer=C",
                    None if awarded >= 60 else "答错",
                    None,
                ),
            )
            if review:
                conn.execute(
                    "UPDATE session_details SET confidence_score=70 WHERE result_id=? AND question_id='Q1'",
                    (result_id,),
                )
            conn.execute(
                """
                INSERT INTO session_details (
                    result_id, question_id, score_awarded, deduction_reason,
                    knowledge_ids, error_category, error_summary
                ) VALUES (?, 'Q2', ?, ?, '["UNKNOWN"]', ?, ?)
                """,
                (
                    result_id,
                    30 if awarded >= 60 else 20,
                    "缺关键步骤",
                    "过程不完整",
                    "缺 BE⊥AC 步骤",
                ),
            )
        for name, status in (
            ("张三", "present"),
            ("李四", "present"),
            ("王五", "absent"),
        ):
            conn.execute(
                """
                INSERT INTO session_attendance (
                    session_id, student_id, attendance_status, source_reason
                ) VALUES (?, ?, ?, ?)
                """,
                (
                    session_id,
                    students[name],
                    status,
                    "测试" if status != "present" else None,
                ),
            )
    return session_id


@pytest.fixture
def analysis_db(tmp_path: Path):
    

    db = open_grading_repositories(tmp_path / "databases" / "grading.db")
    db.initialize()
    session_id = _seed_analysis_session(db, tmp_path)
    return db, session_id, tmp_path


def _make_generator(db, out_dir: Path, reports_dir: Path, llm_client) -> "object":
    from backend.reporting.analysis_report_exporter import AnalysisReportGenerator

    return AnalysisReportGenerator(
        db,
        out_dir,
        llm_client_factory=lambda: llm_client,
        reports_dir=reports_dir,
        data_root=out_dir.parent,
    )


def _seed_report_knowledge(
    analysis_db, question_bank_database, *, type_mode: bool, link_second: bool = True,
    volume_id: str = "bnu24-math-g8-upper", standard_revision: int | None = None,
) -> SimpleNamespace:
    """Synthetic bank links and frozen evidence for the v8/v9 report boundary."""
    from question_bank.database.schema import connect
    from question_bank.knowledge_graph_release import bootstrap_release
    from question_bank.knowledge_graph_release.contracts import KnowledgeGraphRelease
    from question_bank.knowledge_graph_release.loader import load_release_for_taxonomy_revision
    from question_bank.services.source_question_link_service import SourceQuestionLinkService
    from question_bank.solution_evidence.evidence_snapshot import freeze_session_evidence_snapshot

    db, session, root = analysis_db
    path = root / "databases" / "question_bank.db"
    question_bank_database(path)
    type_key = "kp_bnu24_math_g8_upper_1_1_t01"
    section = "kp_bnu24_math_g7_upper_3_3" if volume_id == "bnu24-math-g7-upper" else "kp_bnu24_math_g8_upper_1_1"
    skill_key = "sk_" + section[3:] + "_101"
    topic_key = section + "_1"
    names = {type_key: "题型·合成题型", "kp_bnu24_math_g8_upper_1_1_t02": "题型·合成题型",
             skill_key: "技能·合成技能", topic_key: "合成知识点"}
    payload = load_release_for_taxonomy_revision(standard_revision or (11 if type_mode else 10)).to_dict()
    payload.pop("content_hash", None)
    payload["release_id"] += "-synthetic-report"
    for node in payload["core_nodes"]:
        if node["stable_key"] in names:
            node["display_name"] = names[node["stable_key"]]
    release = KnowledgeGraphRelease.from_mapping(payload)
    bootstrap_release(path, release, actor_ref="test-suite", source_reference="synthetic-report",
                      reason="verify v8 and v9 exam reports")
    rubric_path = root / "config" / "uploaded" / "rubric.json"
    rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
    # A legacy rubric label must not override confirmed current bank identities.
    rubric["questions"][0]["knowledge_points"] = [{"knowledge_id": "UNKNOWN", "name": "旧评分知识点"}]
    for question in rubric["questions"]:
        question["steps"] = [{"step_id": "S1", "step_score": question["max_score"],
                              "evidence_point_ids": ["p1"]}]
    rubric_path.write_text(json.dumps(rubric, ensure_ascii=False), encoding="utf-8")
    with sqlite3.connect(db.db_path) as connection:
        connection.execute("UPDATE grading_sessions SET curriculum_volume_id=? WHERE id=?", (volume_id, session))
    with connect(path) as connection:
        for number, question in enumerate(rubric["questions"], 1):
            version = f"{number:064x}"
            connection.execute(
                "INSERT INTO questions(id,question_number,question_type,question_text,answer_text) VALUES(?,?,?,?,?)",
                (number, str(number), question["question_type"], f"TEST-报告合成题 {number}", "合成答案"),
            )
            tags = [skill_key, topic_key, *([type_key] if type_mode else [])]
            connection.executemany(
                "INSERT INTO question_tags(question_id,tag_type,tag_value,confidence,source) VALUES(?,'knowledge_point',?,1,'taxonomy')",
                [(number, key) for key in tags],
            )
            if type_mode:
                connection.execute(
                    "INSERT INTO question_tags(question_id,tag_type,tag_value,source) VALUES(?,'secondary_type',?,'taxonomy')",
                    (number, "kp_bnu24_math_g8_upper_1_1_t02"),
                )
            evidence = {"parts": [{"part_id": question["question_id"], "response_mode": "process",
                "evidence_points": [{"evidence_point_id": "p1", "target": "合成判定点"}]}]}
            connection.execute(
                """INSERT INTO question_solution_evidence_versions(
                    evidence_version_id,question_id,source_content_hash,schema_version,content_hash,
                    evidence_json,status,source_kind,source_reference,created_by,graph_release_id
                ) VALUES(?,?,?,'question-solution-evidence-v2',?,?,'approved','backfill','synthetic','test',?)""",
                (version, number, "a" * 64, "b" * 64, json.dumps(evidence, ensure_ascii=False), release.release_id),
            )
            connection.executemany(
                """INSERT INTO evidence_point_knowledge_links(
                    evidence_version_id,question_id,part_id,evidence_point_id,graph_release_id,
                    role,term_id,stable_key,resolution_status,weight,source_kind,source_reference
                ) VALUES(?,?,?,'p1',?,'direct',?,?,'resolved',1,'link_job','synthetic')""",
                [(version, number, question["question_id"], release.release_id, key, key)
                 for key in (skill_key, topic_key)],
            )
    links = SourceQuestionLinkService(path)
    for number in (1, 2) if link_second else (1,):
        links.confirm_link(grading_session_id=session, source_question_id=f"Q{number}",
                           bank_question_id=number, link_method="synthetic")
    freeze_session_evidence_snapshot(path, grading_session_id=session,
        upload_config_dir=rubric_path.parent, data_root=root)
    return SimpleNamespace(path=path, type_key=type_key, skill_key=skill_key, topic_key=topic_key)


def test_llm_failure_degrades_to_data_only_report(analysis_db, tmp_path: Path) -> None:
    from backend.llm.llm_client import LLMResponseFormatError

    db, session_id, _root = analysis_db
    client = FakeLLMClient(error=LLMResponseFormatError("bad json"))
    generator = _make_generator(db, tmp_path / "out", tmp_path / "cache", client)

    zip_path = Path(
        generator.export_session(
            session_id, "personal_analysis_html", score_revision="rev-1"
        )
    )

    assert client.calls == 2  # 失败不暗中重发
    with zipfile.ZipFile(zip_path) as archive:
        personal_html = archive.read("001_张三_个人报告.html").decode("utf-8")
    # 无 AI 叙述版降级为老师批语/参考答案，不给家长看内部失败提示与原始批改记录。
    assert "AI 分析生成失败" not in personal_html
    assert "缺关键步骤" not in personal_html
    # 数据段照常渲染。
    assert "本卷答题一览" in personal_html
    assert "全部失分题详解" in personal_html


@pytest.fixture
def analysis_api_client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from backend.api.app import create_app
    from backend.api.dependencies import (
        get_grading_db,
        get_job_manager,
        get_reports_dir,
    )
    from backend.jobs.manager import JobContext, JobManager
    from backend.jobs.store import JobStore
    

    db = open_grading_repositories(tmp_path / "databases" / "grading.db")
    db.initialize()
    session_id = _seed_analysis_session(db, tmp_path)
    reports_dir = tmp_path / "databases" / ".." / "reports"
    reports_dir = (tmp_path / "reports").resolve()
    reports_dir.mkdir(parents=True, exist_ok=True)

    manager = JobManager(JobStore(tmp_path / "jobs.db"), max_workers=1)

    def report_handler(context: JobContext) -> dict[str, object]:
        report_type = str(context.payload.get("report_type") or "score_excel")
        suffix = ".zip" if report_type == "personal_analysis_html" else ".xlsx"
        report_path = reports_dir / f"report{suffix}"
        report_path.write_bytes(b"fake")
        return {
            "session_id": int(context.payload["session_id"]),
            "file_path": str(report_path),
            "filename": report_path.name,
        }

    manager.register("report_export", report_handler)
    app = create_app()
    app.dependency_overrides[get_grading_db] = lambda: db
    app.dependency_overrides[get_job_manager] = lambda: manager
    app.dependency_overrides[get_reports_dir] = lambda: reports_dir
    with TestClient(app) as client:
        try:
            yield client, db, session_id, reports_dir, manager, monkeypatch
        finally:
            manager.shutdown()


def _patch_configured(monkeypatch: pytest.MonkeyPatch, configured: bool) -> None:
    import backend.reporting.analysis_report_exporter as exporter
    import backend.model_profiles.content_generation as content_generation
    from backend.llm.llm_client import LLMSettings

    settings = (
        LLMSettings(
            api_key="k",
            base_url="https://example.com/v1",
            ocr_model="m",
            grading_model="m",
            config_model="content-model",
        )
        if configured
        else None
    )
    monkeypatch.setattr(
        exporter, "resolve_content_generation_settings", lambda: settings
    )
    monkeypatch.setattr(
        exporter,
        "content_generation_public_info",
        lambda: (
            ("测试服务 @ example.com", "content-model") if configured else (None, None)
        ),
    )
    monkeypatch.setattr(
        content_generation, "resolve_content_generation_settings", lambda: settings
    )
    monkeypatch.setattr(
        content_generation,
        "content_generation_public_info",
        lambda: (
            ("测试服务 @ example.com", "content-model") if configured else (None, None)
        ),
    )


# ---------------------------------------------------------------------------
# 报告内容口径：题型中译、批改记录清洗、标准答案收敛、学生作答回填
# ---------------------------------------------------------------------------


def _add_personal_report_scans(db, session_id: int, root: Path) -> list[dict]:
    from PIL import Image

    with sqlite3.connect(db.db_path) as conn:
        conn.row_factory = sqlite3.Row
        papers = [
            dict(row)
            for row in conn.execute(
                "SELECT id, student_id FROM exam_papers WHERE session_id=? ORDER BY id",
                (session_id,),
            )
        ]
        for index, paper in enumerate(papers):
            path = root / f"synthetic-paper-{paper['id']}.png"
            Image.new(
                "RGB", (400, 300), (220, 30, 30) if index == 0 else (30, 30, 220)
            ).save(path)
            conn.execute(
                "UPDATE exam_papers SET front_image=? WHERE id=?",
                (str(path), paper["id"]),
            )
        template_id = conn.execute(
            "INSERT INTO session_templates (session_id, front_template_path, back_template_path) VALUES (?, ?, '')",
            (session_id, str(root / f"synthetic-paper-{papers[0]['id']}.png")),
        ).lastrowid
        for index, qid in enumerate(("Q1", "Q2")):
            conn.execute(
                """INSERT INTO answer_regions (
                    region_uuid, session_id, template_id, page, region_order,
                    x, y, w, h, mapped_question_id
                ) VALUES (?, ?, ?, 'front', ?, 10, ?, 300, 100, ?)""",
                (
                    f"report-{qid}",
                    session_id,
                    template_id,
                    index,
                    10 + index * 110,
                    qid,
                ),
            )
    return papers


def _score_state(db_path: Path) -> list[list[tuple]]:
    with sqlite3.connect(db_path) as conn:
        return [
            conn.execute(sql).fetchall()
            for sql in (
                "SELECT * FROM session_results ORDER BY id",
                "SELECT * FROM session_details ORDER BY id",
                "SELECT * FROM teacher_score_locks ORDER BY id",
            )
        ]


@pytest.mark.parametrize("type_mode,volume_id,standard_revision", [
    (False, "bnu24-math-g8-upper", 10), (True, "bnu24-math-g8-upper", 11),
    (False, "bnu24-math-g7-upper", 11),
], ids=["v8", "v9-typed", "v9-unconverted"])
@pytest.mark.parametrize("link_second", [False, True], ids=["unlinked", "same-target"])
def test_report_target_buckets_keep_final_scores_and_ranks(
    analysis_db, question_bank_database, type_mode, volume_id, standard_revision, link_second,
):
    from openpyxl import load_workbook
    from backend.reporting.report import ReportGenerator
    from backend.session_analysis import assemble_session_analysis, enrich_personal_knowledge
    from backend.reporting.analysis_report_exporter import build_personal_payload, _render_personal_html
    from bs4 import BeautifulSoup

    db, session, root = analysis_db
    bank = _seed_report_knowledge(analysis_db, question_bank_database,
                                 type_mode=type_mode, link_second=link_second,
                                 volume_id=volume_id, standard_revision=standard_revision)
    with sqlite3.connect(db.db_path) as connection:
        student_id, detail_id = connection.execute(
            """SELECT result.student_id,detail.id FROM session_details detail
               JOIN session_results result ON result.id=detail.result_id
               WHERE result.session_id=? AND detail.question_id='Q1' ORDER BY detail.id LIMIT 1""", (session,),
        ).fetchone()
    # Use the application's confirmation path so detail, total, and lock agree.
    db.reviews.confirm_teacher_score_lock(session_id=session, scan_batch_id="TEST-report-types",
        student_id=student_id, question_id="Q1", score_awarded=10, max_score=60,
        deduction_reason="教师确认", source_target_type="session_detail", source_target_id=detail_id,
        expected_revision=0)
    before = _score_state(db.db_path)
    data = assemble_session_analysis(db, session, data_root=root)
    enrich_personal_knowledge(db, data, root)
    assert [(student.student_code, student.student_score, student.rank) for student in data.students] == [
        ("002", 50, 1), ("001", 40, 2),
    ]
    locked_student = next(student for student in data.students if student.student_id == student_id)
    locked = next(record for record in locked_student.records if record.question_id == "Q1")
    assert (locked.score, locked.max_score, locked.teacher_confirmed) == (10, 60, True)
    payload = build_personal_payload(data, locked_student)
    assert (payload["student"]["total_score"], payload["student"]["class_stats"]["rank"]) == (40, 2)
    question = next(question for question in payload["questions"] if question["question_id"] == "Q1")
    assert (question["score"], question["max_score"], question["grading_record"]["teacher_confirmed"]) == (10, 60, True)
    report = BeautifulSoup(_render_personal_html(data, locked_student, None, {}), "html.parser")
    assert report.select_one(".score-line .big").get_text() == "40"
    assert report.select_one(".stat3 .cell .v").get_text() == "2/2"
    assert data.knowledge_structure["target_kind"] == ("type" if type_mode else "skill")
    keys = {entry["stable_key"] for entry in data.knowledge_backfill["Q1"]}
    assert keys == ({bank.type_key} if type_mode else {bank.topic_key, bank.skill_key})
    if volume_id == "bnu24-math-g7-upper":
        from integration.diagnosis_profile_service import DiagnosisProfileService
        service = DiagnosisProfileService(db.db_path, bank.path, data_root=root)
        profile = service.build_profiles(scope={"mode": "all"},
            exam_scope={"mode": "semester", "curriculum_volume_id": volume_id})
        assert profile["target_kind"] == "skill"
        assert any(point["knowledge_key"] == bank.skill_key and point["evidence_count"] > 0
                   for student in profile["students"] for point in student["weak_points"])
        current = service.build_profiles(scope={"mode": "all"},
            exam_scope={"mode": "current", "session_ids": [session]})
        assert current["target_kind"] == "skill"
        exams = service.assembly_exam_questions(class_ids=["1 班"], volume_id=volume_id)
        assert exams["target_kind"] == "skill"
        assert any(bank.skill_key in item["skill_keys"]
                   for exam in exams["exams"] for item in exam["questions"])
        empty = service.assembly_exam_questions(class_ids=["TEST-empty-class"], volume_id=volume_id)
        assert empty == {"student_count": 0, "exams": [], "target_kind": "skill"}

    generator = ReportGenerator(db, root / "out")
    workbook = load_workbook(generator.export_session(session))
    try:
        label = "题型" if type_mode else "知识点"
        assert f"{label}分析" in workbook.sheetnames
        assert f"{'知识点' if type_mode else '题型'}分析" not in workbook.sheetnames

        def table(sheet):
            rows = list(sheet.iter_rows(values_only=True))
            header_index = next(index for index, row in enumerate(rows) if row[:2] == ("班级", label))
            header = rows[header_index]
            return [dict(zip(header, row, strict=True)) for row in rows[header_index + 1:]], rows

        buckets, raw_rows = table(workbook[f"{label}分析"])
        unlinked_label = f"未命名{label}" if type_mode else "旧评分知识点"
        named = [row for row in buckets if row[label] != unlinked_label]
        assert {row[label] for row in named} == ({"合成题型"} if type_mode else {"合成知识点", "技能·合成技能"})
        for row in named:
            assert (row["累计得分"], row["累计满分"], row["失分人数"]) == (
                (90, 200, 2) if link_second else (40, 120, 2)
            )
            assert row["涉及题目"] == ("Q1、Q2" if link_second else "Q1")
            assert row["得分率"] == (0.45 if link_second else 0.3333)
        unknown = [row for row in buckets if row[label] == unlinked_label]
        if link_second:
            assert unknown == []
        else:
            assert len(unknown) == 1
            assert (unknown[0]["涉及题目"], unknown[0]["累计得分"], unknown[0]["累计满分"]) == ("Q2", 50, 80)
            assert unknown[0]["得分率"] == 0.625
            if type_mode:
                assert f"单元测试 · {label}分析" in raw_rows[0]
                assert f"未命名{label}" in str(raw_rows[1])
            else:
                # v8 keeps the historical rubric label when the bank has no link.
                assert raw_rows[0][:2] == ("班级", "知识点")

        detail_rows = list(workbook["成绩与小题明细"].iter_rows(values_only=True))
        details = [dict(zip(detail_rows[0], row, strict=True)) for row in detail_rows[1:]]
        assert [(row["学号"], row["总分"], row["班级排名"]) for row in details] == [
            ("002", 50, 1), ("001", 40, 2),
        ]
    finally:
        workbook.close()
    assert _score_state(db.db_path) == before


def test_report_keeps_same_named_types_separate_by_current_identity(
    analysis_db, question_bank_database,
):
    from openpyxl import load_workbook
    from backend.reporting.analysis_report_exporter import _class_knowledge_view
    from backend.reporting.report import ReportGenerator
    from backend.session_analysis import assemble_session_analysis, enrich_personal_knowledge

    db, session, root = analysis_db
    bank = _seed_report_knowledge(analysis_db, question_bank_database, type_mode=True)
    second_type = "kp_bnu24_math_g8_upper_1_1_t02"
    with sqlite3.connect(bank.path) as connection:
        connection.execute(
            "UPDATE question_tags SET tag_value=? WHERE question_id=2 AND tag_type='knowledge_point' AND tag_value=?",
            (second_type, bank.type_key),
        )
    before = _score_state(db.db_path)
    data = assemble_session_analysis(db, session, data_root=root)
    enrich_personal_knowledge(db, data, root)
    nodes = _class_knowledge_view(data)["nodes"]
    assert {node["key"] for node in nodes} == {bank.type_key, second_type}
    assert [node["label"] for node in nodes] == ["合成题型", "合成题型"]
    assert sorted((node["score"], node["full"]) for node in nodes) == [(50, 80), (90, 120)]
    workbook = load_workbook(ReportGenerator(db, root / "out").export_session(session))
    try:
        rows = list(workbook["题型分析"].iter_rows(values_only=True))
        assert rows[0][:2] == ("班级", "题型")
        assert [(row[1], row[2], row[3], row[4], row[5]) for row in rows[1:]] == [
            ("合成题型", "Q2", 50, 80, 0.625), ("合成题型", "Q1", 90, 120, 0.75),
        ]
    finally:
        workbook.close()
    assert [(student.student_score, student.rank) for student in data.students] == [(90, 1), (50, 2)]
    assert _score_state(db.db_path) == before


@pytest.mark.parametrize("manual_only", [True, False])
def test_visual_reports_include_manual_students_and_keep_final_scores(
    analysis_db,
    tmp_path: Path,
    manual_only: bool,
) -> None:
    import io
    from PIL import Image
    from backend.session_analysis import assemble_session_analysis

    db, session_id, root = analysis_db
    papers = _add_personal_report_scans(db, session_id, root)
    selected = papers if manual_only else papers[1:]
    with sqlite3.connect(db.db_path) as conn:
        for paper in selected:
            student_id = paper["student_id"]
            conn.execute(
                "DELETE FROM session_details WHERE result_id IN (SELECT id FROM session_results WHERE student_id=?)",
                (student_id,),
            )
            conn.execute(
                "DELETE FROM session_results WHERE student_id=?", (student_id,)
            )
            for qid, score, maximum in (("Q1", 40, 60), ("Q2", 15, 40)):
                conn.execute(
                    """INSERT INTO teacher_score_locks (
                        session_id, scan_batch_id, student_id, question_id,
                        score_awarded, max_score, deduction_reason,
                        source_target_type, source_target_id
                    ) VALUES (?, 'synthetic-batch', ?, ?, ?, ?, '教师确认的过程分', 'exam_paper', ?)""",
                    (session_id, student_id, qid, score, maximum, paper["id"]),
                )
    before = _score_state(db.db_path)
    data = assemble_session_analysis(db, session_id, data_root=root)
    assert len(data.students) == 2
    assert sum(student.result_id < 0 for student in data.students) == len(selected)
    client = FakeLLMClient()
    generator = _make_generator(db, tmp_path / "out", tmp_path / "cache", client)
    output = generator.export_session(
        session_id, "personal_analysis_html", score_revision="manual-v1"
    )
    assert _score_state(db.db_path) == before
    assert client.calls == 2
    with zipfile.ZipFile(output) as archive:
        for student, request in zip(data.students, client.requests, strict=True):
            payload = json.loads(request["prompt"].split("输入 JSON：\n", 1)[1])
            assert payload["student"]["total_score"] == student.student_score
            assert request["use_config_client"] is True
            assert len(request["images"]) == len(payload["image_map"]) >= 2
            assert [item["image_number"] for item in payload["image_map"]] == list(
                range(1, len(request["images"]) + 1)
            )
            assert any(
                item["kind"] == "student_overview" for item in payload["image_map"]
            )
            # Different solid colours ensure the report receives its own student's image.
            with Image.open(io.BytesIO(request["images"][0])) as image:
                red, _, blue = image.convert("RGB").getpixel((0, 0))
            assert (red > blue) == (student.student_id == papers[0]["student_id"])
            if student.result_id < 0:
                assert student.student_score == 55
                assert all(
                    q["grading_record"]["teacher_confirmed"]
                    for q in payload["questions"]
                    if q["lost"]
                )
            text = archive.read(
                f"{student.student_code}_{student.student_name}_个人报告.html"
            ).decode("utf-8")
            assert "原卷截图" in text
            assert "未取得该生可读取的答卷图片" not in text
    # Re-entering export reuses the existing narrative; it does not regrade or call twice.
    generator.export_session(
        session_id, "personal_analysis_html", score_revision="manual-v1"
    )
    assert client.calls == 2
    assert _score_state(db.db_path) == before


# ---------------------------------------------------------------------------
# 个人学情报告留存：删除接口与重复提交复用
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("legacy_version", ["personal_analysis_html_v11_parent_brief", "personal_analysis_html_v10_actions", "personal_analysis_html_v9_problem_refs", "personal_analysis_html_v8_parts"])
def test_old_prompt_personal_reports_regenerate_and_store(analysis_db, tmp_path, legacy_version):
    """输入一致的旧提示词结果（old_prompt）不命中生成缓存：导出重新生成并
    把新叙述写回 PersonalReportStore。"""
    from backend.personal_reports import student_report_digests
    from backend.report_results import PersonalReportStore, prompt_version
    from backend.session_analysis import (
        assemble_session_analysis,
        enrich_personal_questions,
    )

    db, session_id, root = analysis_db
    reports_dir = tmp_path / "reports"
    data = assemble_session_analysis(db, session_id, data_root=root)
    enrich_personal_questions(db, data, root)
    digests = student_report_digests(
        db, session_id, data, reports_dir=reports_dir
    )
    store = PersonalReportStore(reports_dir)
    for student in data.students:
        store.save(
            session_id, student.student_id,
            narrative=PERSONAL_NARRATIVE,
            input_digest=digests[student.student_id],
            prompt_version=legacy_version,
        )
    client = FakeLLMClient()
    generator = _make_generator(db, tmp_path / "out", reports_dir, client)
    output = generator.export_session(session_id, "personal_analysis_html", html_only=True)
    assert client.calls == len(data.students)

    # 旧提示词条目被新结果覆盖为当前版本。
    for student in data.students:
        entry = store.load(session_id, student.student_id)
        assert entry["prompt_version"] == prompt_version("personal_report")
        assert entry["input_digest"] == digests[student.student_id]
    assert output.is_dir()
    assert len(list(output.glob("*.html"))) == len(data.students)
    assert all(path.suffix == ".html" for path in output.iterdir())
    assert not list((tmp_path / "out").rglob("*.zip"))
    assert not list((tmp_path / "out").rglob("*.pdf"))
    for request in client.requests:
        assert "knowledge_focus" in request["prompt"]
    report = (output / "001_张三_个人报告.html").read_text(encoding="utf-8")
    assert "先做：" in report and "检查：" in report
    assert '<span class="when">本周</span>' in report
    generator.export_session(session_id, "personal_analysis_html", html_only=True)
    assert client.calls == len(data.students)

def test_report_status_and_digests_share_the_same_inputs(analysis_db, monkeypatch):
    from backend.personal_reports import personal_report_states, student_report_digests
    from backend.session_analysis import assemble_session_analysis
    db, sid, root = analysis_db
    data = assemble_session_analysis(db, sid, data_root=root, page_only=True)
    states = personal_report_states(db, sid, root / 'reports', data=data)
    assert {s["student_id"] for s in states["students"]
            if s["status"] != "unavailable"} == {s.student_id for s in data.students}
    # 同一输入重复计算指纹一致且确定性输出。
    reports_dir = root / 'reports'
    digests = student_report_digests(db, sid, data, reports_dir=reports_dir)
    assert student_report_digests(db, sid, data, reports_dir=reports_dir) == digests
    assert set(digests) == {s.student_id for s in data.students}


def test_personal_context_cold_exam_does_not_block_cached_exam_and_shares_same_key(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from types import SimpleNamespace
    from collections import Counter
    import backend.personal_reports as personal
    import backend.session_analysis as analysis
    import backend.reporting.analysis_report_exporter as exporter
    import backend.class_analysis as causes
    import backend.files.session_originals as originals
    entered, release = Event(), Event()
    calls = Counter()
    repos = SimpleNamespace(db_path=tmp_path/'grading.db',
                            sessions=SimpleNamespace(list_grading_sessions=lambda: []))
    monkeypatch.setattr(personal, '_read_generation', lambda *_: ('TEST',))
    monkeypatch.setattr(originals, 'originals_state', lambda *_: 'complete')
    def assemble(_repos, sid, **_):
        calls[sid] += 1
        return SimpleNamespace(session_id=sid)
    monkeypatch.setattr(personal, 'assemble_session_analysis', assemble)
    monkeypatch.setattr(personal, 'split_session_analysis_by_class', lambda _: {})
    monkeypatch.setattr(analysis, 'enrich_personal_questions', lambda *_: None)
    def knowledge(_repos, data, _root):
        if data.session_id == 1:
            entered.set()
            assert release.wait(3)
    monkeypatch.setattr(analysis, 'enrich_personal_knowledge', knowledge)
    monkeypatch.setattr(exporter, '_load_student_histories', lambda *_: {})
    monkeypatch.setattr(exporter, 'load_session_regions', lambda *_, **__: {})
    monkeypatch.setattr(exporter, '_load_personal_error_histories', lambda *_: {})
    monkeypatch.setattr(causes, 'build_cause_inputs', lambda _: {})
    personal._contexts.clear()
    reports = tmp_path/'reports'
    cached = personal.personal_render_context(repos, 2, reports)
    try:
        with ThreadPoolExecutor(max_workers=3) as pool:
            first = pool.submit(personal.personal_render_context, repos, 1, reports)
            assert entered.wait(2)
            second = pool.submit(personal.personal_render_context, repos, 1, reports)
            hot = pool.submit(personal.personal_render_context, repos, 2, reports)
            try:
                assert hot.result(timeout=1) is cached
            finally:
                release.set()
            assert first.result(timeout=2) is second.result(timeout=2)
    finally:
        release.set()
        personal._contexts.clear()
    assert calls == {1: 1, 2: 1}


def test_personal_report_revision_keeps_peers_current_and_reads_stale_text(analysis_db, tmp_path):
    from backend.reporting.analysis_report_exporter import AnalysisReportGenerator
    from backend.personal_reports import personal_report_states, render_personal_report
    db, sid, root = analysis_db
    with sqlite3.connect(db.db_path) as conn:
        manual_id = conn.execute("INSERT INTO students(student_code,name,class_name) VALUES ('004','合成纯人工','1 班')").lastrowid
        conn.execute("INSERT INTO session_attendance(session_id,student_id,attendance_status) VALUES (?,?,'present')", (sid, manual_id))
        for qid, score, maximum in [('Q1', 50, 60), ('Q2', 30, 40)]:
            conn.execute("""INSERT INTO teacher_score_locks(session_id,scan_batch_id,student_id,question_id,score_awarded,max_score,deduction_reason,source_target_type,source_target_id)
                VALUES (?,0,?,?,?,?,?,'manual',1)""", (sid, manual_id, qid, score, maximum, '合成批语'))
    reports = root / "reports"
    client = FakeLLMClient()
    generator = AnalysisReportGenerator(db, root / "test-generated", data_root=root, reports_dir=reports,
        llm_client_factory=lambda: client)
    generator.export_session(sid, "personal_analysis_html", html_only=True)
    states = personal_report_states(db, sid, reports)["students"]
    current = [s["student_id"] for s in states if s["status"] == "current"]
    assert len(current) == 3
    assert any(s["status"] == "unavailable" for s in states)
    files_before = {p.relative_to(reports): p.read_bytes() for p in reports.rglob("*") if p.is_file()}
    with sqlite3.connect(db.db_path) as conn:
        conn.execute("UPDATE session_details SET score_awarded=score_awarded-1 WHERE result_id IN (SELECT id FROM session_results WHERE student_id=?) AND question_id='Q2'", (current[0],))
    states = {s["student_id"]: s["status"] for s in personal_report_states(db, sid, reports)["students"]}
    assert states[current[0]] == "stale"
    assert all(states[student_id] == "current" for student_id in current[1:])
    html = render_personal_report(db, sid, current[0], reports, review_links=True)
    assert "选择题全对，基础扎实" in html
    assert 'class="review-link"' in html
    assert files_before == {p.relative_to(reports): p.read_bytes() for p in reports.rglob("*") if p.is_file()}
    assert client.calls == 3


def test_personal_cache_only_reads_store_without_client(analysis_db, tmp_path):
    """cache_only 导出只读 PersonalReportStore：不建模型客户端、不写存储。"""
    from backend.personal_reports import (
        personal_report_states,
        student_report_digests,
    )
    from backend.report_results import PersonalReportStore, prompt_version
    from backend.reporting.analysis_report_exporter import AnalysisReportGenerator
    from backend.session_analysis import (
        assemble_session_analysis,
        enrich_personal_questions,
    )

    db, sid, root = analysis_db
    reports_dir = root / "reports"
    data = assemble_session_analysis(db, sid)
    enrich_personal_questions(db, data, root)
    student = data.students[0]
    store = PersonalReportStore(reports_dir)
    store.save(
        sid, student.student_id, narrative=PERSONAL_NARRATIVE,
        input_digest=student_report_digests(
            db, sid, data, reports_dir=reports_dir
        )[student.student_id],
        prompt_version=prompt_version("personal_report"),
    )

    def forbidden():
        raise AssertionError("只读导出不能初始化模型")

    generator = AnalysisReportGenerator(db, root / "test-bundle", data_root=root,
        reports_dir=reports_dir, llm_client_factory=forbidden)
    generator.export_session(sid, "personal_analysis_html", narrative_mode="cache_only", html_only=True)
    assert generator.last_personal_summary == dict(generated=1, failed=0, skipped=2)
    assert {s["reason"] for s in generator.last_personal_missing} >= {"未生成", "缺考"}
    assert store.load(sid, student.student_id)["narrative"] == PERSONAL_NARRATIVE
    assert not (reports_dir / ".analysis_review_notes").exists()
    assert personal_report_states(db, sid, reports_dir)["students"][0]["status"] == "current"


def test_personal_online_and_offline_images_share_content_and_released_notice(analysis_db, monkeypatch):
    import re
    from backend.personal_reports import render_personal_report, RELEASED_SHOT_NOTE
    db, sid, root = analysis_db
    papers = _add_personal_report_scans(db, sid, root)
    student_id = papers[0]["student_id"]
    from unittest.mock import Mock
    from backend.session_analysis import enrich_personal_knowledge
    prepare = Mock(wraps=enrich_personal_knowledge)
    monkeypatch.setattr('backend.session_analysis.enrich_personal_knowledge', prepare)
    online = render_personal_report(db, sid, student_id, root / "reports", narrative_mode="none")
    offline = render_personal_report(db, sid, student_id, root / "reports", narrative_mode="none", online=False)
    assert f'/personal-reports/{student_id}/shots/Q2' in online
    assert 'data:image/jpeg' not in online
    assert 'data:image/jpeg' in offline
    assert 'class="review-link"' not in online
    from backend.reporting.analysis_report_exporter import _PERSONAL_KEYBOARD_JS
    assert 'personal-report:key' in online and 'personal-report:key' not in offline
    # 在线数据版带「AI 分析部分尚未整理」占位提示；离线导出不带。
    assert "AI 分析部分尚未整理" in online and "AI 分析部分尚未整理" not in offline
    strip_banner = lambda html: re.sub(
        r'<div class="card"><div class="review-banner">AI 分析部分尚未整理.*?</div></div>'
        r'|<div class="note">以下按失分排序；AI 跟进建议整理后显示。</div>',
        '', html, flags=re.S,
    )
    normalize = lambda html: strip_banner(re.sub(r'src="(?:/api/[^\"]+|data:image/jpeg[^\"]+)"', 'src="SHOT"', html.replace(_PERSONAL_KEYBOARD_JS, '')))
    assert normalize(online) == normalize(offline)
    assert prepare.call_count == 1  # 同场连续读取只准备一次，在线与离线结果一致。
    from backend.personal_reports import personal_render_context
    from dataclasses import asdict
    before = asdict(personal_render_context(db, sid, root / 'reports')['data'])
    render_personal_report(db, sid, papers[1]['student_id'], root / 'reports', narrative_mode='none')
    assert asdict(personal_render_context(db, sid, root / 'reports')['data']) == before
    monkeypatch.setattr("backend.files.session_originals.originals_state", lambda *_a: "cleared")
    for mode in (True, False):
        html = render_personal_report(db, sid, student_id, root / "reports", narrative_mode="none", online=mode)
        assert RELEASED_SHOT_NOTE in html
        assert '/shots/' not in html and 'data:image/jpeg' not in html


def test_personal_report_teacher_lock_change_invalidates_input(analysis_db):
    """教师最终分锁进入该生输入指纹：加锁使该生过期，其他学生不受影响。"""
    from backend.personal_reports import personal_report_states, student_report_digests
    from backend.session_analysis import (
        assemble_session_analysis,
        enrich_personal_questions,
    )

    db, sid, root = analysis_db
    reports_dir = root / "reports"
    data = assemble_session_analysis(db, sid, data_root=root)
    enrich_personal_questions(db, data, root)
    before = student_report_digests(db, sid, data, reports_dir=reports_dir)
    from backend.report_results import PersonalReportStore, prompt_version
    for saved_student_id, digest in before.items():
        PersonalReportStore(reports_dir).save(sid, saved_student_id, narrative={"summary": "TEST-lock"},
                                             input_digest=digest, prompt_version=prompt_version("personal_report"))
    initial_states = personal_report_states(db, sid, reports_dir)
    assert all(row["status"] == "current" for row in initial_states["students"] if row["student_id"] in before)
    student = db.results.get_session_results(sid)[0]
    student_id = student["student_id"]
    with sqlite3.connect(db.db_path) as conn:
        conn.execute("""INSERT INTO teacher_score_locks(session_id,scan_batch_id,student_id,question_id,score_awarded,max_score,deduction_reason,source_target_type,source_target_id)
            VALUES (?,'test-personal',?,'Q2',29,40,'教师补充批语','exam_paper',?)""", (sid, student_id, db.results.get_result_context(student["result_id"])["paper_id"]))
    after = student_report_digests(db, sid, data, reports_dir=reports_dir)
    assert after[student_id] != before[student_id]
    assert all(after[other] == before[other] for other in before if other != student_id)
    refreshed = personal_report_states(db, sid, reports_dir)
    assert next(row for row in refreshed["students"] if row["student_id"] == student_id)["status"] == "stale"
    assert all(row["status"] == "current" for row in refreshed["students"]
               if row["student_id"] in before and row["student_id"] != student_id)


@pytest.mark.parametrize("limit", [8, 1])
def test_personal_report_executor_follows_configured_concurrency(
    analysis_db, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, limit: int,
) -> None:
    """个人报告线程池并发数跟随模型配置，不再被固定上限 3 截断。"""
    import backend.reporting.analysis_report_exporter as exporter_module

    db, session_id, _root = analysis_db
    client = FakeLLMClient()
    client.config_gateway = SimpleNamespace(
        execution_snapshot=SimpleNamespace(max_in_flight=limit)
    )
    captured: list[int] = []
    real_executor = exporter_module.ThreadPoolExecutor

    class Recorder:
        def __init__(self, *args, **kwargs):
            if kwargs.get("thread_name_prefix") == "personal-report":
                captured.append(kwargs["max_workers"])
            self._pool = real_executor(*args, **kwargs)

        def __enter__(self):
            return self._pool.__enter__()

        def __exit__(self, *exc):
            return self._pool.__exit__(*exc)

        def submit(self, *args, **kwargs):
            return self._pool.submit(*args, **kwargs)

    monkeypatch.setattr(exporter_module, "ThreadPoolExecutor", Recorder)
    generator = _make_generator(db, tmp_path / "out", tmp_path / "cache", client)
    generator.export_session(
        session_id, "personal_analysis_html", score_revision="rev-1"
    )
    assert captured == [limit]


def test_personal_store_writes_on_caller_thread_and_tolerates_corruption(analysis_db, monkeypatch):
    """生成的叙述由主线程写入 PersonalReportStore；损坏文件按缺失处理。"""
    import threading
    import backend.personal_reports as personal
    from backend.report_results import PersonalReportStore
    from backend.reporting.analysis_report_exporter import AnalysisReportGenerator

    db, sid, root = analysis_db
    owner = threading.get_ident()
    original = PersonalReportStore.save
    writes = []
    def record(self, *args, **kwargs):
        writes.append(threading.get_ident())
        return original(self, *args, **kwargs)
    monkeypatch.setattr(PersonalReportStore, "save", record)
    generator = AnalysisReportGenerator(db, root / "test-index", data_root=root,
        reports_dir=root / "reports", llm_client_factory=lambda: FakeLLMClient())
    generator.export_session(sid, "personal_analysis_html", html_only=True)
    assert writes and set(writes) == {owner}
    session_dir = root / "reports" / ".personal_reports" / f"session_{sid}"
    files = sorted(session_dir.glob("*.json"))
    assert files
    corrupted_id = int(files[0].stem)
    files[0].write_text("[invalid", encoding="utf-8")
    states = personal.personal_report_states(db, sid, root / "reports")["students"]
    by_id = {s["student_id"]: s["status"] for s in states}
    assert by_id[corrupted_id] == "missing"
    assert sum(1 for s in states if s["status"] == "current") == len(files) - 1
