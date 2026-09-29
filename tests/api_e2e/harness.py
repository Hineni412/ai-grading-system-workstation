from __future__ import annotations

import json
import re
import threading
import time
from collections.abc import Iterator
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from shutil import copy2
from types import SimpleNamespace
from typing import Any

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from backend.domain_models import GradingResult, QuestionGradingDetail
from backend.api.dependencies import (
    get_annotated_dir,
    get_backups_dir,
    get_data_root,
    get_exams_dir,
    get_grading_db,
    get_job_file_service,
    get_job_manager,
    get_ops_write_service,
    get_outputs_dir,
    get_reports_dir,
    get_templates_dir,
    get_upload_config_dir,
)
from backend.files.service import JobFileService
from backend.jobs.config_generation import run_config_generation_job
from backend.jobs.default_handlers import register_default_job_handlers
from backend.jobs.grading_run import run_grading_job
from backend.jobs.manager import JobManager
from backend.jobs.scan_analysis import run_scan_analysis
from backend.jobs.store import JobStore
from report import ReportGenerator
from scanner import ExamPaperGroup, ScanAnalysis
from backend.repositories.grading_database import open_grading_repositories


@dataclass
class E2EControls:
    scan_failures_remaining: int = 0
    grading_runs: int = 0
    scan_runner_calls: int = 0
    scanner_analyze_calls: int = 0
    fake_llm_calls: list[str] = field(default_factory=list)
    grading_qbank_paths: list[Path | None] = field(default_factory=list)
    uploaded_scan_paths: dict[str, Path] = field(default_factory=dict)
    _lock: threading.Lock = field(default_factory=threading.Lock, repr=False)

    def record_fake_llm_call(self, call: str) -> None:
        with self._lock:
            self.fake_llm_calls.append(call)

    def record_scan_runner_call(self) -> None:
        with self._lock:
            self.scan_runner_calls += 1

    def begin_scanner_analysis(self) -> bool:
        with self._lock:
            self.scanner_analyze_calls += 1
            if self.scan_failures_remaining <= 0:
                return False
            self.scan_failures_remaining -= 1
            return True

    def record_grading_qbank_path(self, path: Path | None) -> None:
        with self._lock:
            self.grading_qbank_paths.append(path)


@dataclass(frozen=True)
class E2EPaths:
    data_root: Path
    db_path: Path
    qb_db_path: Path
    reports_dir: Path
    exams_dir: Path
    templates_dir: Path
    annotated_dir: Path
    outputs_dir: Path
    backups_dir: Path
    upload_config_dir: Path
    bootstrap_rubric: Path
    bootstrap_answer: Path
    template_front: Path
    template_back: Path


class FakeLLM:
    _BATCH_QUESTION_IDS = re.compile(
        r"^BATCH_QUESTION_IDS_JSON=(\[.*\])$",
        re.MULTILINE,
    )

    def __init__(self, controls: E2EControls) -> None:
        self.controls = controls
        self.settings = SimpleNamespace(
            config_model=None,
            ocr_model=None,
            grading_model=None,
        )

    def json_from_text(
        self,
        prompt: str,
        model: str | None = None,
        **_kwargs: Any,
    ) -> dict[str, Any]:
        del model
        if "SCORE_QUESTION_IDS_JSON=" in prompt:
            self.controls.record_fake_llm_call("score_allocation")
            return _fake_score_allocation()
        match = self._BATCH_QUESTION_IDS.search(prompt)
        if match is None:
            raise AssertionError("unexpected synthetic config-generation prompt")
        question_ids = json.loads(match.group(1))
        if (
            not isinstance(question_ids, list)
            or not question_ids
            or any(
                not isinstance(question_id, str)
                or re.fullmatch(r"Q\d+", question_id) is None
                for question_id in question_ids
            )
        ):
            raise AssertionError("unexpected synthetic config-generation question ids")
        for question_id in question_ids:
            self.controls.record_fake_llm_call(f"question:{question_id}")
        payloads = [_fake_single_question_payload(question_id) for question_id in question_ids]
        return {
            "rubric": {
                "questions": [
                    question
                    for payload in payloads
                    for question in payload["rubric"]["questions"]
                ]
            },
            "answer_key": {
                "questions": [
                    question
                    for payload in payloads
                    for question in payload["answer_key"]["questions"]
                ]
            },
        }

    def json_from_text_once(
        self,
        prompt: str,
        model: str | None = None,
        **kwargs: Any,
    ) -> dict[str, Any]:
        return self.json_from_text(prompt, model=model, **kwargs)

    def json_from_images(
        self,
        prompt: str,
        _images: list[bytes],
        model: str | None = None,
        **kwargs: Any,
    ) -> dict[str, Any]:
        return self.json_from_text(prompt, model=model, **kwargs)


def _fake_single_question_payload(question_id: str) -> dict[str, Any]:
    part_id = question_id
    step_id = "S1"
    return {
        "rubric": {
            "questions": [
                {
                    "question_id": question_id,
                    "question_type": "comprehensive",
                    "max_score": 1,
                    "parts": [
                        {
                            "part_id": part_id,
                            "part_score": 1,
                            "response_mode": "short_answer_points",
                            "visual_requirements": [],
                            "steps": [
                                {
                                    "step_id": step_id,
                                    "step_score": 1,
                                    "core_goal": "synthetic answer",
                                    "required_elements": ["synthetic evidence"],
                                    "allow_alternative_methods": True,
                                }
                            ],
                        }
                    ],
                }
            ]
        },
        "answer_key": {
            "questions": [
                {
                    "question_id": question_id,
                    "canonical_answer": f"synthetic answer {question_id}",
                    "accepted_forms": [f"synthetic answer {question_id}"],
                    "method_variants": [],
                    "parts": [
                        {
                            "part_id": part_id,
                            "answer": f"synthetic answer {question_id}",
                            "accepted_forms": [],
                            "analysis": "synthetic analysis",
                            "step_milestones": ["synthetic evidence"],
                        }
                    ],
                }
            ]
        },
    }


def _fake_score_allocation() -> dict[str, Any]:
    question_scores = []
    for index, score in enumerate([17, 17, 17, 17, 17, 15], start=1):
        question_id = f"Q{index}"
        part_id = question_id
        question_scores.append(
            {
                "question_id": question_id,
                "max_score": score,
                "parts": [
                    {
                        "part_id": part_id,
                        "part_score": score,
                        "steps": [
                            {
                                "step_id": "S1",
                                "step_score": score,
                            }
                        ],
                    }
                ],
            }
        )
    return {"question_scores": question_scores}


class SyntheticScanner:
    def __init__(
        self,
        paths: E2EPaths,
        controls: E2EControls,
        scanner_kwargs: dict[str, Any],
    ) -> None:
        self.paths = paths
        self.controls = controls
        self.scanner_kwargs = dict(scanner_kwargs)

    def analyze(self, students: list[dict[str, Any]]) -> ScanAnalysis:
        if self.controls.begin_scanner_analysis():
            raise RuntimeError("synthetic scan failure")
        students_by_code = {
            str(student.get("student_code") or ""): student
            for student in students
        }
        groups: list[ExamPaperGroup] = []
        for student_code in ("SYN-001", "SYN-002"):
            student = students_by_code[student_code]
            groups.append(
                ExamPaperGroup(
                    front_image=self.controls.uploaded_scan_paths.get(
                        f"{student_code}_front.png", self.paths.exams_dir / f"{student_code}_front.png",
                    ),
                    back_image=self.controls.uploaded_scan_paths.get(
                        f"{student_code}_back.png", self.paths.exams_dir / f"{student_code}_back.png",
                    ),
                    student_name=str(student["name"]),
                    student_id=int(student["id"]),
                    detected_name=str(student["name"]),
                    source_label=f"{student_code} synthetic scan",
                    match_method="exact",
                    match_score=1.0,
                )
            )
        return ScanAnalysis(groups=groups, total_pages=4)


class SyntheticGradingService:
    def __init__(
        self,
        db_manager: Any,
        llm_client: Any,
        *,
        controls: E2EControls,
        question_bank_db_path: Path | None = None,
    ) -> None:
        self.db = db_manager
        self.llm_client = llm_client
        self.controls = controls
        self.question_bank_db_path = question_bank_db_path
        self.controls.record_grading_qbank_path(question_bank_db_path)

    def run_session_grading(
        self,
        *,
        session_id: int,
        scan_analysis: dict[str, Any] | None,
        failed_only: bool,
        **_kwargs: Any,
    ) -> Iterator[dict[str, Any]]:
        self.controls.grading_runs += 1
        self.db.sessions.try_start_session_run(session_id)
        if failed_only:
            failed_papers = self.db.papers.list_failed_papers_detailed(session_id)
            total = len(failed_papers)
            for current, paper in enumerate(failed_papers, start=1):
                student_id = int(paper["student_id"])
                student = self._student_by_id(student_id)
                result_id = self.db.results.save_session_result(
                    session_id,
                    student_id,
                    int(paper["paper_id"]),
                    _synthetic_grading_result(
                        student_name=str(student["name"]),
                        scores=[10, 12, 12, 12, 11, 13],
                        needs_human_review=False,
                    ),
                )
                self.db.papers.update_exam_paper_status(int(paper["paper_id"]), "graded")
                yield {
                    "event": "graded",
                    "student_name": str(student["name"]),
                    "result_id": result_id,
                    "score": 70.0,
                    "total_score": 100.0,
                    "current": current,
                    "total": total,
                }
        else:
            groups = list((scan_analysis or {}).get("groups") or [])
            total = len(groups)
            for current, group in enumerate(groups, start=1):
                student_id = int(group["student_id"])
                student = self._student_by_id(student_id)
                student_code = str(student["student_code"])
                paper_id = self.db.papers.create_exam_paper(
                    session_id=session_id,
                    front_image=str(group["front_image"]),
                    back_image=str(group["back_image"]),
                    ocr_name=str(group.get("detected_name") or group["student_name"]),
                    student_id=student_id,
                    match_status="matched",
                    processing_status="grading",
                )
                if student_code == "SYN-001":
                    result_id = self.db.results.save_session_result(
                        session_id,
                        student_id,
                        paper_id,
                        _synthetic_grading_result(
                            student_name=str(student["name"]),
                            scores=[12, 17, 17, 17, 17, 5],
                            needs_human_review=True,
                        ),
                    )
                    self.db.papers.update_exam_paper_status(paper_id, "graded")
                    yield {
                        "event": "graded",
                        "student_name": str(student["name"]),
                        "result_id": result_id,
                        "score": 85.0,
                        "total_score": 100.0,
                        "current": current,
                        "total": total,
                    }
                else:
                    self.db.papers.update_exam_paper_status(
                        paper_id,
                        "failed",
                        "synthetic grading failure",
                    )
                    yield {
                        "event": "grading_failed",
                        "student_name": str(student["name"]),
                        "current": current,
                        "total": total,
                    }

        self.db.sessions.finish_session_run(session_id, "completed")
        yield {
            "event": "session_completed",
            "progress": self.db.papers.get_session_progress(session_id),
        }

    def _student_by_id(self, student_id: int) -> dict[str, Any]:
        return next(
            student
            for student in self.db.students.list_students()
            if int(student["id"]) == int(student_id)
        )


def _synthetic_grading_result(
    *,
    student_name: str,
    scores: list[float],
    needs_human_review: bool,
) -> GradingResult:
    details = [
        QuestionGradingDetail(
            question_id=f"Q{index}",
            score_awarded=float(score),
            deduction_reason=(
                "synthetic review required"
                if needs_human_review and index == 1
                else "synthetic grading"
            ),
            knowledge_id=f"SYN-K{index}",
            confidence_score=50.0 if needs_human_review and index == 1 else 100.0,
            knowledge_ids=[f"SYN-K{index}"],
        )
        for index, score in enumerate(scores, start=1)
    ]
    return GradingResult(
        student_name=student_name,
        total_score=100.0,
        student_score=float(sum(scores)),
        needs_human_review=needs_human_review,
        grading_details=details,
        raw_json={
            "source": "synthetic_api_e2e",
            "grading_completeness": {"status": "complete"},
        },
    )


def make_scan_runner(paths: E2EPaths, controls: E2EControls):
    def scan_runner(**kwargs: Any) -> dict[str, object]:
        controls.record_scan_runner_call()
        return run_scan_analysis(
            **kwargs,
            scanner_factory=lambda **scanner_kwargs: SyntheticScanner(
                paths,
                controls,
                scanner_kwargs,
            ),
        )

    return scan_runner


def make_grading_runner(controls: E2EControls):
    def grading_runner(**kwargs: Any) -> dict[str, object]:
        return run_grading_job(
            **kwargs,
            service_factory=lambda db_manager, llm_client, question_bank_db_path=None: (
                SyntheticGradingService(
                    db_manager,
                    llm_client,
                    controls=controls,
                    question_bank_db_path=question_bank_db_path,
                )
            ),
        )

    return grading_runner


def _iter_strings(value: Any) -> Iterator[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for key, item in value.items():
            yield from _iter_strings(key)
            yield from _iter_strings(item)
    elif isinstance(value, (list, tuple)):
        for item in value:
            yield from _iter_strings(item)


def _normalized_public_text(value: str) -> str:
    return value.replace("\\", "/").casefold()


class ApiE2EHarness:
    TERMINAL = {"succeeded", "failed", "cancelled"}

    def __init__(
        self,
        client: TestClient,
        db: Any,
        manager: Any,
        paths: E2EPaths,
        controls: E2EControls,
    ) -> None:
        self.client = client
        self.db = db
        self.manager = manager
        self.paths = paths
        self.controls = controls

    def poll_job(
        self,
        job_id: int,
        expected_status: str,
        timeout: float = 5.0,
    ) -> dict[str, Any]:
        deadline = time.monotonic() + timeout
        last: dict[str, Any] = {}
        while time.monotonic() < deadline:
            response = self.client.get(f"/api/jobs/{job_id}")
            assert response.status_code == 200
            last = response.json()
            normalized_data_root = _normalized_public_text(
                str(self.paths.data_root)
            ).rstrip("/")
            for value in _iter_strings(last):
                assert normalized_data_root not in _normalized_public_text(value), (
                    "job response exposed the temporary data root"
                )
                assert "api_key" not in value.casefold(), (
                    "job response exposed an API key field or value"
                )
            if last["status"] in self.TERMINAL:
                assert last["status"] == expected_status
                return last
            time.sleep(0.01)
        raise AssertionError(f"job {job_id} did not reach a terminal state: {last}")

    def create_configured_session(self) -> int:
        created = self.client.post(
            "/api/sessions",
            json={
                "name": "Synthetic E2E Exam",
                "rubric_path": str(self.paths.bootstrap_rubric),
                "answer_key_path": str(self.paths.bootstrap_answer),
            },
        )
        assert created.status_code == 201
        session_id = int(created.json()["id"])

        submitted = self.client.post(
            f"/api/sessions/{session_id}/config/generate",
            json=self.config_request(),
        )
        assert submitted.status_code == 202
        config_job = self.poll_job(submitted.json()["id"], "succeeded")
        assert config_job["result"]["outcome"] == "complete"

        committed = self.bind_and_commit_template(session_id)
        snapshots = list(
            (
                self.paths.templates_dir / f"session_{session_id}"
            ).glob("regions_confirmed_*.json")
        )
        assert len(snapshots) == 1
        snapshot_path = snapshots[0]
        assert snapshot_path.is_file()
        snapshot_path.resolve().relative_to(
            (self.paths.templates_dir / f"session_{session_id}").resolve()
        )
        assert committed == {
            "committed": True,
            "snapshot_pending": False,
            "issues": [],
            "region_count": 1,
            "error": None,
        }

        upload_dir = (
            self.paths.exams_dir
            / f"session_{session_id}"
            / "uploaded_scans"
        )
        upload_dir.mkdir(parents=True, exist_ok=True)
        for source in sorted(self.paths.exams_dir.glob("SYN-*.png")):
            copy2(source, upload_dir / source.name)
        return session_id

    def scan(self, session_id: int) -> dict[str, Any]:
        submitted = self.client.post(
            f"/api/sessions/{session_id}/scan/analyze",
            json={"enhance_images": False},
        )
        assert submitted.status_code == 202
        return self.poll_job(submitted.json()["id"], "succeeded")

    def grade(
        self,
        session_id: int,
        *,
        failed_only: bool = False,
    ) -> dict[str, Any]:
        submitted = self.client.post(
            f"/api/sessions/{session_id}/grading/run",
            json={
                "grading_mode": "ai",
                "failed_only": failed_only,
                "enhance_images": False,
            },
        )
        assert submitted.status_code == 202
        return self.poll_job(submitted.json()["id"], "succeeded")

    @staticmethod
    def xlsx_score(content: bytes, student_code: str) -> float:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
        try:
            sheet = workbook["成绩与小题明细"]
            rows = sheet.iter_rows(values_only=True)
            headers = {
                str(value): index
                for index, value in enumerate(next(rows))
                if value is not None
            }
            code_index = headers["学号"]
            score_index = headers["总分"]
            for row in rows:
                if str(row[code_index]) == str(student_code):
                    return float(row[score_index])
            raise AssertionError(
                f"student {student_code!r} was not found in the downloaded report"
            )
        finally:
            workbook.close()

    @staticmethod
    def config_request() -> dict[str, Any]:
        return {
            "confirmed_blocks": [
                {
                    "question_id": f"Q{index}",
                    "question_type": "comprehensive",
                    "text": f"synthetic prompt {index}",
                    "answer_text": f"synthetic answer {index}",
                    "canonical_answer": f"synthetic answer {index}",
                }
                for index in range(1, 7)
            ],
            "document_text": "synthetic exam text",
            "question_images": {},
        }

    def bind_and_commit_template(self, session_id: int) -> dict[str, Any]:
        bound = self.client.put(
            f"/api/sessions/{session_id}/template",
            json={
                "front_template_path": str(self.paths.template_front),
                "back_template_path": str(self.paths.template_back),
            },
        )
        assert bound.status_code == 200
        region = {
            "region_uuid": "q1-region",
            "page": "front",
            "region_order": 1,
            "x": 10,
            "y": 10,
            "w": 80,
            "h": 60,
            "mapped_question_id": "Q1",
            "mapping_status": "manual",
            "is_confirmed": True,
        }
        draft = self.client.put(
            f"/api/sessions/{session_id}/regions/draft",
            json={"revision": 1, "regions": [region]},
        )
        assert draft.status_code == 200
        committed = self.client.post(
            f"/api/sessions/{session_id}/regions/commit",
            json={
                "regions": [region],
                "image_sizes": {"front": [120, 160], "back": [120, 160]},
                "template_matches": True,
                "expected_template_fingerprint": draft.json()[
                    "template_fingerprint"
                ],
            },
        )
        assert committed.status_code == 200
        return committed.json()

    def paper_statuses(self, session_id: int) -> list[str]:
        with self.db._connect() as connection:
            rows = connection.execute(
                "SELECT processing_status FROM exam_papers "
                "WHERE session_id = ? ORDER BY processing_status, id",
                (session_id,),
            ).fetchall()
        return [str(row["processing_status"]) for row in rows]

    def result_scores(self, session_id: int) -> dict[str, float]:
        with self.db._connect() as connection:
            rows = connection.execute(
                "SELECT s.student_code, sr.student_score "
                "FROM session_results sr "
                "JOIN students s ON s.id = sr.student_id "
                "WHERE sr.session_id = ? ORDER BY s.student_code",
                (session_id,),
            ).fetchall()
        return {
            str(row["student_code"]): float(row["student_score"])
            for row in rows
        }

    def result_ids(self, session_id: int) -> dict[str, int]:
        with self.db._connect() as connection:
            rows = connection.execute(
                "SELECT s.student_code, sr.id "
                "FROM session_results sr "
                "JOIN students s ON s.id = sr.student_id "
                "WHERE sr.session_id = ? ORDER BY s.student_code",
                (session_id,),
            ).fetchall()
        return {str(row["student_code"]): int(row["id"]) for row in rows}

    def result_review_flags(self, session_id: int) -> dict[str, bool]:
        with self.db._connect() as connection:
            rows = connection.execute(
                "SELECT s.student_code, sr.needs_human_review "
                "FROM session_results sr "
                "JOIN students s ON s.id = sr.student_id "
                "WHERE sr.session_id = ? ORDER BY s.student_code",
                (session_id,),
            ).fetchall()
        return {
            str(row["student_code"]): bool(row["needs_human_review"])
            for row in rows
        }

    def result_details(
        self,
        session_id: int,
    ) -> dict[str, list[tuple[str, float]]]:
        with self.db._connect() as connection:
            rows = connection.execute(
                "SELECT s.student_code, sd.question_id, sd.score_awarded "
                "FROM session_details sd "
                "JOIN session_results sr ON sr.id = sd.result_id "
                "JOIN students s ON s.id = sr.student_id "
                "WHERE sr.session_id = ? "
                "ORDER BY s.student_code, sd.id",
                (session_id,),
            ).fetchall()
        details: dict[str, list[tuple[str, float]]] = {}
        for row in rows:
            details.setdefault(str(row["student_code"]), []).append(
                (str(row["question_id"]), float(row["score_awarded"]))
            )
        return details

def build_paths(tmp_path: Path) -> E2EPaths:
    data_root = tmp_path / "synthetic_data"
    databases_dir = data_root / "databases"
    reports_dir = data_root / "reports"
    exams_dir = data_root / "exams"
    templates_dir = data_root / "templates"
    annotated_dir = data_root / "annotated"
    outputs_dir = data_root / "outputs"
    backups_dir = data_root / "backups"
    upload_config_dir = data_root / "config" / "uploaded"
    for directory in (
        databases_dir,
        reports_dir,
        exams_dir,
        templates_dir,
        annotated_dir,
        outputs_dir,
        backups_dir,
        upload_config_dir,
    ):
        directory.mkdir(parents=True, exist_ok=True)

    bootstrap_rubric = upload_config_dir / "bootstrap_rubric.json"
    bootstrap_answer = upload_config_dir / "bootstrap_answer.json"
    for path in (bootstrap_rubric, bootstrap_answer):
        path.write_text(json.dumps({}), encoding="utf-8")

    template_front = templates_dir / "template_front.png"
    template_back = templates_dir / "template_back.png"
    image_assets = (
        (template_front, "white"),
        (template_back, "white"),
        (exams_dir / "SYN-001_front.png", "lightgray"),
        (exams_dir / "SYN-001_back.png", "lightgray"),
        (exams_dir / "SYN-002_front.png", "gainsboro"),
        (exams_dir / "SYN-002_back.png", "gainsboro"),
    )
    for path, color in image_assets:
        Image.new("RGB", (120, 160), color=color).save(path, format="PNG")

    return E2EPaths(
        data_root=data_root,
        db_path=databases_dir / "grading_system.db",
        qb_db_path=databases_dir / "question_bank.db",
        reports_dir=reports_dir,
        exams_dir=exams_dir,
        templates_dir=templates_dir,
        annotated_dir=annotated_dir,
        outputs_dir=outputs_dir,
        backups_dir=backups_dir,
        upload_config_dir=upload_config_dir,
        bootstrap_rubric=bootstrap_rubric,
        bootstrap_answer=bootstrap_answer,
        template_front=template_front,
        template_back=template_back,
    )


def build_job_manager(paths: E2EPaths, *, controls: E2EControls) -> JobManager:
    manager = JobManager(JobStore(paths.db_path), max_workers=1)
    try:
        register_default_job_handlers(
            manager,
            db_path=paths.db_path,
            reports_dir=paths.reports_dir,
            exams_dir=paths.exams_dir,
            templates_dir=paths.templates_dir,
            data_root=paths.data_root,
            question_bank_db_path=paths.qb_db_path,
            upload_config_dir=paths.upload_config_dir,
            report_generator_factory=ReportGenerator,
            scan_runner=make_scan_runner(paths, controls),
            grading_runner=make_grading_runner(controls),
            config_generation_runner=run_config_generation_job,
            llm_client_factory=lambda: FakeLLM(controls),
        )
    except Exception:
        manager.shutdown()
        raise
    return manager


def install_dependency_overrides(
    app: FastAPI,
    *,
    db: Any,
    manager: JobManager,
    paths: E2EPaths,
) -> None:
    job_files = JobFileService(
        paths.reports_dir,
        backups_dir=paths.backups_dir,
        ops_outputs_dir=paths.outputs_dir / "ops",
    )
    inert_ops_service = SimpleNamespace()
    app.dependency_overrides.update(
        {
            get_grading_db: lambda: db,
            get_job_manager: lambda: manager,
            get_upload_config_dir: lambda: paths.upload_config_dir,
            get_templates_dir: lambda: paths.templates_dir,
            get_data_root: lambda: paths.data_root,
            get_exams_dir: lambda: paths.exams_dir,
            get_reports_dir: lambda: paths.reports_dir,
            get_annotated_dir: lambda: paths.annotated_dir,
            get_outputs_dir: lambda: paths.outputs_dir,
            get_backups_dir: lambda: paths.backups_dir,
            get_job_file_service: lambda: job_files,
            get_ops_write_service: lambda: inert_ops_service,
        }
    )


def serve_browser_review() -> None:
    """Serve the existing synthetic API harness for the focused browser check."""
    import argparse
    import hashlib
    import os

    import path_manager
    import uvicorn
    from backend.api.app import create_app
    from db_manager import StudentRecord
    from backend.scan_grading.workspace import ScanGradingWorkspace
    from question_bank.database.schema import initialize_database

    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, required=True)
    args = parser.parse_args()
    # The suite supplies a fresh sandbox and removes real model configuration.
    sandbox = Path(os.environ["AI_GRADING_WORKTREE_DATA_DIR"])
    paths = build_paths(sandbox)
    application_paths = path_manager.get_path_manager()
    application_paths._data_root = paths.data_root
    application_paths._logs_root = sandbox / "logs"
    db = open_grading_repositories(paths.db_path)
    db.initialize()
    initialize_database(paths.qb_db_path)
    db.students.upsert_students([
        StudentRecord("SYN-001", "Synthetic Student A", "Synthetic Class"),
        StudentRecord("SYN-002", "Synthetic Student B", "Synthetic Class"),
    ])
    controls = E2EControls()
    manager = build_job_manager(paths, controls=controls)
    try:
        app = create_app(path_manager=application_paths)
        install_dependency_overrides(app, db=db, manager=manager, paths=paths)
        with TestClient(app) as client:
            harness = ApiE2EHarness(client, db, manager, paths, controls)
            session_id = harness.create_configured_session()
            # Use the current frozen upload batch, not the legacy no-batch API path.
            workspace = ScanGradingWorkspace(
                exams_root=paths.exams_dir, templates_root=paths.templates_dir,
                grading_db_path=paths.db_path,
            )
            uploaded: dict[str, str] = {}
            for index, source in enumerate(sorted(paths.exams_dir.glob("SYN-*.png"))):
                with Image.open(source) as scan:
                    scan.putpixel((0, 0), (index, 0, 0))
                    scan.save(source)
                content = source.read_bytes()
                digest = hashlib.sha256(content).hexdigest()
                response = client.post(
                    f"/api/sessions/{session_id}/scan-uploads", content=content,
                    headers={"content-type": "image/png", "x-upload-filename": source.name,
                             "x-content-sha256": digest},
                )
                assert response.status_code == 201, response.text
                uploaded[source.name] = digest
            frozen = client.post(
                f"/api/sessions/{session_id}/scan-uploads/freeze",
                json={"expected_revision": len(uploaded)},
            )
            assert frozen.status_code == 200, frozen.text
            scan_dir = workspace.frozen_scan_dir(session_id)
            controls.uploaded_scan_paths = {
                name: scan_dir / f"{digest}.png" for name, digest in uploaded.items()
            }
            harness.scan(session_id)
            preflight = client.get(f"/api/sessions/{session_id}/scan/preflight")
            assert preflight.status_code == 200, preflight.text
            # Seed known AI results; this acceptance starts at teacher review.
            # All saves under test subsequently go through the real browser/API.
            db.sessions.try_start_session_run(session_id)
            for student, scores in zip(db.students.list_students(), (
                [12, 17, 17, 16, 16, 7], [10, 12, 12, 12, 11, 13],
            ), strict=True):
                code = student["student_code"]
                paper_id = db.papers.create_exam_paper(
                    session_id=session_id,
                    front_image=str(controls.uploaded_scan_paths[f"{code}_front.png"]),
                    back_image=str(controls.uploaded_scan_paths[f"{code}_back.png"]),
                    ocr_name=student["name"], student_id=student["id"],
                    match_status="matched", processing_status="graded",
                )
                db.results.save_session_result(
                    session_id, student["id"], paper_id,
                    _synthetic_grading_result(
                        student_name=student["name"], scores=scores,
                        needs_human_review=code == "SYN-001",
                    ),
                    scan_batch_id=frozen.json()["batch_id"],
                )
            db.sessions.finish_session_run(session_id, "completed")
        uvicorn.run(app, host="127.0.0.1", port=args.port)
    finally:
        manager.shutdown()


if __name__ == "__main__":
    serve_browser_review()
