from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from question_bank.personalized_papers.latex_render import (
    LatexRenderError,
    render_training_tex,
)


@pytest.mark.parametrize("inside_word_run", [False, True])
@pytest.mark.parametrize(("math_xml", "expected"), [
    ('<m:rad><m:deg/><m:e><m:r><m:t>18</m:t></m:r></m:e></m:rad>', r'\sqrt{18}'),
    ('<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:rad><m:deg><m:r><m:t>3</m:t></m:r></m:deg><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad></m:den></m:f>', r'\frac{1}{\sqrt[3]{x}}'),
    ('<m:sSubSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSubSup>', r'{x}_{1}^{2}'),
])
def test_word_math_keeps_structure(tmp_path: Path, math_xml: str, expected: str, inside_word_run: bool) -> None:
    content = f'<m:oMath>{math_xml}</m:oMath>'
    if inside_word_run:
        content = f'<w:r><w:t>公式：</w:t>{content}<w:t>。</w:t></w:r>'
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
           f'{content}</w:p>')
    tex = render_training_tex(_snapshot([{"text": "合成公式", "xml": xml}]), data_root=tmp_path)
    assert f'\\({expected}\\)' in tex


def test_unknown_word_math_uses_existing_export_fallback(tmp_path: Path) -> None:
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
           '<m:oMath><m:unknownStructure><m:r><m:t>x</m:t></m:r></m:unknownStructure></m:oMath></w:p>')
    with pytest.raises(LatexRenderError, match="math"):
        render_training_tex(_snapshot([{"text": "合成公式", "xml": xml}]), data_root=tmp_path)


def test_unknown_math_property_does_not_silently_change_the_equation(tmp_path: Path) -> None:
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath>'
           '<m:rad><m:radPr><m:unknownEffect/></m:radPr><m:deg/><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>'
           '</m:oMath></w:p>')
    with pytest.raises(LatexRenderError, match="math property"):
        render_training_tex(_snapshot([{"text": "合成公式", "xml": xml}]), data_root=tmp_path)


@pytest.mark.parametrize("alphabet", ["double-struck", "script", "fraktur"])
def test_unsupported_math_alphabet_keeps_existing_fallback(tmp_path: Path, alphabet: str) -> None:
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath>'
           f'<m:r><m:rPr><m:scr m:val="{alphabet}"/></m:rPr><m:t>R</m:t></m:r>'
           '</m:oMath></w:p>')
    with pytest.raises(LatexRenderError, match="math alphabet"):
        render_training_tex(_snapshot([{"text": "合成符号", "xml": xml}]), data_root=tmp_path)


def _asset(tmp_path: Path, name: str = "pic.png") -> Path:
    target = tmp_path / "question_bank" / "personalized_papers" / "inst" / "assets" / name
    target.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (80, 40), "white").save(target, format="PNG")
    return target


def _picture_paragraph_xml(asset: Path, tmp_path: Path) -> tuple[str, dict[str, str]]:
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run().add_picture(str(asset), width=Inches(1.25))
    blip = paragraph._p.xpath(".//a:blip")[0]  # noqa: SLF001
    relationship_id = str(blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    ))
    relative = asset.relative_to(tmp_path).as_posix()
    return paragraph._p.xml, {relationship_id: relative}  # noqa: SLF001


def _paragraph_xml(runs: list[tuple[str, dict[str, bool]]]) -> str:
    source = Document()
    paragraph = source.add_paragraph()
    for text, style in runs:
        run = paragraph.add_run(text)
        run.underline = style.get("underline", False)
        run.font.superscript = style.get("superscript", False)
        run.font.subscript = style.get("subscript", False)
    return paragraph._p.xml  # noqa: SLF001


def _snapshot(blocks: list[dict[str, object]], **overrides) -> dict[str, object]:
    question = {
        "question_id": 1,
        "tagging_context": {
            "question_number": "1",
            "question_type": overrides.get("question_type", "解答题"),
            "question_text": overrides.get("question_text", "（12 分）合成题干"),
        },
        "rich_question_blocks": blocks,
    }
    return {
        "paper_instance_id": "c" * 64,
        "student": {
            "student_id": "1",
            "student_name": "合成学生",
            "class_id": "七年级一班",
        },
        "items": [{
            "task_item_code": "TASK-1",
            "question_id": 1,
            "question_snapshot": question,
        }],
    }


def test_render_tex_escapes_specials_and_styles(tmp_path: Path) -> None:
    tex = render_training_tex(
        _snapshot([
            {
                "text": "已知 x_1 & y^2 {100%}",
                "xml": _paragraph_xml([
                    ("已知 x", {}),
                    ("1", {"subscript": True}),
                    (" & y", {}),
                    ("2", {"superscript": True}),
                    (" {100%}", {}),
                    ("重点", {"underline": True}),
                ]),
            },
        ]),
        data_root=tmp_path,
    )
    assert "x\\_1" not in tex  # 下标来自 XML vertAlign，不是文本
    assert "\\&" in tex and "\\%" in tex
    assert "\\textbf{1.}" in tex
    assert "\\textsubscript{1}" in tex
    assert "\\textsuperscript{2}" in tex
    assert "\\underline{重点}" in tex
    assert "姓名：合成学生" in tex and "班级：七年级一班" in tex


def test_render_tex_table_becomes_tabular(tmp_path: Path) -> None:
    source = Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "售价"
    table.cell(0, 1).text = "29"
    table.cell(1, 0).text = "销量"
    table.cell(1, 1).text = "220"
    tex = render_training_tex(
        _snapshot([{"text": "表格", "xml": table._tbl.xml}]),  # noqa: SLF001
        data_root=tmp_path,
    )
    assert "\\begin{tabular}" in tex
    assert "售价 & 29" in tex
    assert "\\hline" in tex


def test_render_tex_images_single_and_row(tmp_path: Path) -> None:
    first = _asset(tmp_path, "a.png")
    second = _asset(tmp_path, "b.png")
    xml1, rels1 = _picture_paragraph_xml(first, tmp_path)
    xml2, rels2 = _picture_paragraph_xml(second, tmp_path)
    tex = render_training_tex(
        _snapshot(
            [
                {"text": "", "xml": xml1, "image_relationships": rels1},
                {"text": "", "xml": xml2, "image_relationships": rels2},
            ],
            question_type="选择题",
            question_text="（3 分）选择题",
        ),
        data_root=tmp_path,
    )
    assert tex.count("\\begin{minipage}") == 2
    # 图片统一压缩到印刷缓存，tex 引用缓存副本。
    assert tex.count("image_cache/") == 2

    single = render_training_tex(
        _snapshot(
            [{"text": "", "xml": xml1, "image_relationships": rels1}],
            question_type="选择题",
            question_text="（3 分）选择题",
        ),
        data_root=tmp_path,
    )
    assert "\\begin{center}" in single
    assert "\\begin{minipage}" not in single


def test_render_tex_option_inline_image_is_height_capped(tmp_path: Path) -> None:
    asset = _asset(tmp_path)
    source = Document()
    paragraph = source.add_paragraph("A. ")
    paragraph.add_run().add_picture(str(asset), width=Inches(2.5))
    paragraph.add_run("选项说明")
    blip = paragraph._p.xpath(".//a:blip")[0]  # noqa: SLF001
    relationship_id = str(blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    ))
    tex = render_training_tex(
        _snapshot([{
            "text": "A. 选项说明",
            "xml": paragraph._p.xml,  # noqa: SLF001
            "image_relationships": {relationship_id: asset.relative_to(tmp_path).as_posix()},
        }]),
        data_root=tmp_path,
    )
    assert "height=22.0mm" in tex
    assert "选项说明" in tex


def test_render_tex_answer_space_follows_score(tmp_path: Path) -> None:
    tex = render_training_tex(
        _snapshot(
            [{"text": "题干", "xml": _paragraph_xml([("题干", {})])}],
            question_type="填空题",
            question_text="（12 分）光明乳鸽",
        ),
        data_root=tmp_path,
    )
    assert "\\vspace*{90mm}" in tex

    no_space = render_training_tex(
        _snapshot(
            [{"text": "题干", "xml": _paragraph_xml([("题干", {})])}],
            question_type="选择题",
            question_text="（3 分）选择题",
        ),
        data_root=tmp_path,
    )
    assert "\\vspace*" not in no_space


def test_render_tex_fails_closed_on_missing_image_or_text_block(
    tmp_path: Path,
) -> None:
    xml, rels = _picture_paragraph_xml(_asset(tmp_path), tmp_path)
    broken = dict(rels)
    broken[next(iter(broken))] = "question_bank/personalized_papers/inst/assets/missing.png"
    with pytest.raises(LatexRenderError):
        render_training_tex(
            _snapshot([{"text": "", "xml": xml, "image_relationships": broken}]),
            data_root=tmp_path,
        )
    with pytest.raises(LatexRenderError):
        render_training_tex(
            _snapshot([{"text": "只有文本没有XML", "xml": ""}]),
            data_root=tmp_path,
        )
