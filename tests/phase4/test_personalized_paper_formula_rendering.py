from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Mm
from PIL import Image

from question_bank.personalized_papers.rendering import (
    PaperRenderError,
    inspect_docx,
    render_review_docx,
)


def test_review_docx_uses_compact_student_bar_and_borderless_three_line_answer_space(
    tmp_path: Path,
) -> None:
    output = tmp_path / "compact-review.docx"
    render_review_docx(
        {
            "paper_instance_id": "f" * 64,
            "series_version": 1,
            "student": {
                "student_id": "1",
                "student_name": "贾浩然",
                "class_id": "七年级一班",
            },
            "estimated_minutes": 20,
            "budget": {
                "estimated_total_tokens": 100,
                "context_window_tokens": 32768,
            },
            "items": [
                {
                    "task_item_code": "TASK-CHOICE-1",
                    "question_id": 1,
                    "question_snapshot": {
                        "question_id": 1,
                        "tagging_context": {
                            "question_number": "3",
                            "question_type": "选择题",
                            "question_text": "选择题正文（　　）",
                        },
                        "images": [],
                    },
                    "recommendation_snapshot": {
                        "question_number": "3",
                        "source_paper": "2025-2026学年广东省深圳市福田区七年级期末数学试卷",
                    },
                },
                {
                    "task_item_code": "TASK-SOLUTION-2",
                    "question_id": 2,
                    "question_snapshot": {
                        "question_id": 2,
                        "tagging_context": {
                            "question_number": "18",
                            "question_type": "解答题",
                            "question_text": "解答题正文。",
                        },
                        "images": [],
                    },
                    "recommendation_snapshot": {
                        "question_number": "18",
                        "source_paper": "2025-2026学年广东省深圳市福田区七年级期末数学试卷",
                    },
                },
            ],
        },
        data_root=tmp_path,
        output_path=output,
    )

    document = Document(output)
    visible_paragraphs = [paragraph.text for paragraph in document.paragraphs]
    visible_text = "\n".join(visible_paragraphs)
    assert "个性化训练卷（审核稿）" not in visible_text
    assert "检查题目、分页与留白" not in visible_text
    assert "判定预算" not in visible_text
    assert "姓名：贾浩然" in visible_text
    assert "1. 选择题正文" in visible_text
    assert "2. 解答题正文" in visible_text
    assert "深圳期末" not in visible_text
    assert "原题第" not in visible_text
    assert len(document.tables) == 1
    assert document.sections[0].bottom_margin >= Mm(27)
    body_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "选择题正文" in paragraph.text
    )
    assert body_paragraph.paragraph_format.line_spacing == 1.1
    solution_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "解答题正文" in paragraph.text
    )
    assert visible_paragraphs.index(solution_paragraph.text) == (
        visible_paragraphs.index(body_paragraph.text) + 1
    )

    answer_table = document.tables[0]
    answer_xml = answer_table._tbl.xml  # noqa: SLF001
    assert 'w:val="single"' not in answer_xml
    assert 'w:val="nil"' in answer_xml
    height = answer_table.rows[0]._tr.xpath("./w:trPr/w:trHeight")[0]  # noqa: SLF001
    assert height.get(qn("w:hRule")) == "atLeast"
    assert int(height.get(qn("w:val"))) >= 1530

    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    assert b'w:hRule="atLeast"' in document_xml



def test_review_docx_uses_editable_omml_for_supported_formula(tmp_path: Path) -> None:
    output = tmp_path / "review.docx"
    fallbacks = render_review_docx(
        {
            "paper_instance_id": "a" * 64,
            "series_version": 1,
            "student": {
                "student_id": "1",
                "student_code": "001",
                "student_name": "张三",
                "class_id": "1班",
            },
            "estimated_minutes": 10,
            "budget": {
                "estimated_total_tokens": 100,
                "context_window_tokens": 32768,
            },
            "items": [{
                "task_item_code": "TASK-1",
                "question_id": 1,
                "question_snapshot": {
                    "question_id": 1,
                    "tagging_context": {
                        "question_number": "1",
                        "question_type": "解答题",
                        "question_text": "计算 $x^2+1$ 的值。",
                    },
                    "images": [],
                },
                "recommendation_snapshot": {"question_number": "1"},
            }],
        },
        data_root=tmp_path,
        output_path=output,
    )

    assert fallbacks == ()
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    assert b"<m:oMath" in document_xml


def test_choice_question_places_small_standalone_picture_beside_stem(
    tmp_path: Path,
) -> None:
    asset = tmp_path / "assets" / "choice.png"
    asset.parent.mkdir(parents=True)
    Image.new("RGB", (180, 150), "white").save(asset)
    source = Document()
    stem = source.add_paragraph("1. 选择题图文并排正文")
    picture = source.add_paragraph()
    picture.add_run().add_picture(str(asset), width=Inches(1.2))
    relationship_id = str(
        picture._p.xpath(".//a:blip")[0].get(qn("r:embed"))  # noqa: SLF001
    )
    output = tmp_path / "choice-image.docx"

    render_review_docx(
        {
            "paper_instance_id": "9" * 64,
            "series_version": 1,
            "student": {"student_id": "1", "student_name": "贾浩然"},
            "items": [{
                "task_item_code": "TASK-CHOICE-IMAGE",
                "question_id": 1,
                "question_snapshot": {
                    "question_id": 1,
                    "tagging_context": {
                        "question_type": "选择题",
                        "question_text": "纯文本后备题干",
                    },
                    "rich_question_blocks": [
                        {"text": stem.text, "xml": stem._p.xml},  # noqa: SLF001
                        {
                            "text": "",
                            "xml": picture._p.xml,  # noqa: SLF001
                            "image_relationships": {
                                relationship_id: "assets/choice.png",
                            },
                        },
                    ],
                    "images": [{
                        "role": "question",
                        "mime_type": "image/png",
                        "sha256": "a" * 64,
                        "asset_path": "assets/choice.png",
                    }],
                },
                "recommendation_snapshot": {
                    "source_paper": "2025-2026学年广东省深圳市期末数学试卷",
                },
            }],
        },
        data_root=tmp_path,
        output_path=output,
    )

    document = Document(output)
    assert len(document.tables) == 1
    layout = document.tables[0]
    assert len(layout.columns) == 2
    assert layout.cell(0, 0).text.startswith("1. 选择题图文并排正文")
    assert any(
        paragraph._p.xpath(".//a:blip")  # noqa: SLF001
        for paragraph in layout.cell(0, 1).paragraphs
    )
    assert "深圳期末" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )


def test_unsupported_formula_without_source_image_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(PaperRenderError, match="source image fallback"):
        render_review_docx(
            {
                "paper_instance_id": "b" * 64,
                "series_version": 1,
                "student": {"student_id": "1", "student_name": "张三"},
                "budget": {"estimated_total_tokens": 100, "context_window_tokens": 32768},
                "items": [{
                    "task_item_code": "TASK-1",
                    "question_id": 1,
                    "question_snapshot": {
                        "question_id": 1,
                        "tagging_context": {
                            "question_number": "1",
                            "question_type": "解答题",
                            "question_text": "计算 $\\input{unsafe}$。",
                        },
                        "images": [],
                    },
                    "recommendation_snapshot": {"question_number": "1"},
                }],
            },
            data_root=tmp_path,
            output_path=tmp_path / "unsafe.docx",
        )


def test_review_docx_prefers_frozen_word_blocks_and_keeps_original_image_size(
    tmp_path: Path,
) -> None:
    asset = tmp_path / "question_bank" / "personalized_papers" / "paper" / "assets" / "figure.png"
    asset.parent.mkdir(parents=True)
    image = Image.new("RGB", (300, 120), "white")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    asset.write_bytes(buffer.getvalue())

    source = Document()
    stem = source.add_paragraph()
    underlined = stem.add_run("1. 已知 x")
    underlined.underline = True
    superscript = stem.add_run("2")
    superscript.font.superscript = True
    stem.add_run("，完成下表。")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "富文本表格内容"
    picture = source.add_paragraph()
    picture.add_run().add_picture(str(asset), width=Inches(1.25))
    blip = picture._p.xpath(".//a:blip")[0]  # noqa: SLF001
    relationship_id = str(blip.get(qn("r:embed")))
    relative_asset = asset.relative_to(tmp_path).as_posix()

    output = tmp_path / "rich-review.docx"
    render_review_docx(
        {
            "paper_instance_id": "c" * 64,
            "series_version": 1,
            "student": {
                "student_id": "1",
                "student_name": "贾浩然",
                "class_id": "1班",
            },
            "estimated_minutes": 10,
            "budget": {
                "estimated_total_tokens": 100,
                "context_window_tokens": 32768,
            },
            "items": [{
                "task_item_code": "TASK-RICH-1",
                "question_id": 1,
                "question_snapshot": {
                    "question_id": 1,
                    "tagging_context": {
                        "question_number": "1",
                        "question_type": "解答题",
                        "question_text": "不应显示的<sup>原始标签</sup>",
                    },
                    "rich_question_blocks": [
                        {"text": stem.text, "xml": stem._p.xml},  # noqa: SLF001
                        {"text": "富文本表格内容", "xml": table._tbl.xml},  # noqa: SLF001
                        {
                            "text": "",
                            "xml": picture._p.xml,  # noqa: SLF001
                            "image_relationships": {
                                relationship_id: relative_asset,
                            },
                        },
                    ],
                    "images": [{
                        "role": "question",
                        "mime_type": "image/png",
                        "sha256": "d" * 64,
                        "asset_path": relative_asset,
                    }],
                },
                "recommendation_snapshot": {"question_number": "1"},
            }],
        },
        data_root=tmp_path,
        output_path=output,
    )

    reopened = Document(output)
    visible_text = "\n".join(
        [*(paragraph.text for paragraph in reopened.paragraphs)]
        + [
            paragraph.text
            for output_table in reopened.tables
            for row in output_table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
    )
    assert "不应显示" not in visible_text
    assert "<sup>" not in visible_text
    assert "富文本表格内容" in visible_text
    assert len(reopened.tables) == 2
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert b'<w:u w:val="single"' in document_xml
    assert b'<w:vertAlign w:val="superscript"' in document_xml
    assert b'wp:extent cx="1143000"' in document_xml
    assert len(media) == 1
    rich_stem = next(
        paragraph for paragraph in reopened.paragraphs if "已知 x" in paragraph.text
    )
    assert rich_stem.text.startswith("1. ")
    assert rich_stem.paragraph_format.line_spacing == 1.1
    assert rich_stem.paragraph_format.keep_with_next is True
    assert not any(
        paragraph._p.xpath(".//a:blip")  # noqa: SLF001
        for paragraph in reopened.paragraphs
    )
    answer_space = reopened.tables[-1]
    assert len(answer_space.columns) == 2
    rich_picture = next(
        paragraph
        for paragraph in answer_space.cell(0, 1).paragraphs
        if paragraph._p.xpath(".//a:blip")  # noqa: SLF001
    )
    assert rich_picture.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert 'w:val="single"' not in answer_space._tbl.xml  # noqa: SLF001
    inspect_docx(
        output,
        paper_instance_id="c" * 64,
        task_item_codes=("TASK-RICH-1",),
        question_snapshots=({
            "tagging_context": {
                "question_text": "不应显示的<sup>原始标签</sup>",
            },
            "rich_question_blocks": [
                {"text": stem.text, "xml": stem._p.xml},  # noqa: SLF001
                {"text": "富文本表格内容", "xml": table._tbl.xml},  # noqa: SLF001
                {
                    "text": "",
                    "xml": picture._p.xml,  # noqa: SLF001
                    "image_relationships": {relationship_id: relative_asset},
                },
            ],
        },),
    )
