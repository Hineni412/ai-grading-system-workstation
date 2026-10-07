from __future__ import annotations

import hashlib
import json
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from question_bank.document_pipeline import OcrLine, QuestionDocumentPipeline
from question_bank.importers import batch_importer
from question_bank.importers.batch_importer import (
    PaperMetadata,
    ScannedPaper,
    parse_paper_text,
)
from question_bank.database.schema import connect
from question_bank.importers.types import ExtractedDocument
from question_bank.services.rich_content_service import (
    RichContentReadError,
    load_question_rich_content,
)


def test_word_picture_alone_does_not_require_teacher_review(tmp_path):
    from io import BytesIO
    from base64 import b64decode
    from question_bank.importers.docx_importer import import_docx
    document = Document()
    document.add_paragraph("1. 看图填空。")
    document.add_picture(BytesIO(b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6rCEAAAAASUVORK5CYII=")))
    file = tmp_path / "synthetic-image.docx"
    document.save(file)
    extracted = import_docx(file, asset_root=tmp_path / "assets")
    assert extracted.has_images and extracted.image_paths
    assert not extracted.needs_image_review


@pytest.mark.parametrize("version", [1, 2, 3, "2", "legacy", None, {"label": "legacy"}])
def test_rich_content_reads_supported_blocks_without_version_rewrites(tmp_path, version):
    document = Document()
    paragraph = document.add_paragraph("TEST-计算 x")
    paragraph.add_run("2").font.superscript = True
    block = {
        "text": "TEST-计算 x2。[[IMAGE:TEST-figure.png]]",
        "xml": paragraph._p.xml,
        "image_relationships": {"rId9": "TEST-figure.png"},
    }
    if version == 2:
        block["xml"] = '<?xml version="1.0" encoding="UTF-8"?>' + block["xml"]
    payload = {
        "version": version,
        "question_id": 1,
        "question_blocks": [{"text": "一、选择题"}, block],
        "answer_blocks": [{"text": "TEST-参考答案", "xml": "", "image_relationships": {}}],
    }
    sidecar = tmp_path / "question_1.json"
    sidecar.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8-sig" if version == 1 else "utf-8")
    original = sidecar.read_bytes()
    result = load_question_rich_content(1, root=tmp_path, strict=True)
    assert result is not None
    assert result["question_blocks"] == [block]
    assert result["answer_blocks"] == payload["answer_blocks"]
    assert result["version"] == version
    assert sidecar.read_bytes() == original


@pytest.mark.parametrize("body,code", [
    ("{TEST", "rich_content_invalid_json"),
    (json.dumps({"version": 3, "paragraphs": ["TEST-unknown"]}), "rich_content_unknown_structure"),
    (json.dumps({"question_blocks": [{"text": 123}]}), "rich_content_invalid_blocks"),
    (json.dumps({"question_blocks": [{"text": "TEST", "xml": "<w:p>"}]}), "rich_content_invalid_xml"),
    (json.dumps({"question_blocks": [{"text": "TEST", "xml": "<future>TEST</future>"}]}), "rich_content_unknown_structure"),
    (json.dumps({"question_blocks": [{"text": "TEST", "image_relationships": []}]}), "rich_content_invalid_blocks"),
    (json.dumps({"question_id": 2, "question_blocks": [{"text": "TEST"}]}), "rich_content_question_mismatch"),
])
def test_rich_content_distinguishes_damaged_sidecar_from_absent_file(tmp_path, body, code):
    assert load_question_rich_content(1, root=tmp_path, strict=True) is None
    sidecar = tmp_path / "question_1.json"
    sidecar.write_text(body, encoding="utf-8")
    assert load_question_rich_content(1, root=tmp_path) is None
    with pytest.raises(RichContentReadError) as error:
        load_question_rich_content(1, root=tmp_path, strict=True)
    assert error.value.code == code


def test_rich_content_reports_unreadable_file_without_exposing_source(tmp_path, monkeypatch):
    from question_bank.services import rich_content_service

    def unreadable(path):
        raise PermissionError('TEST-controlled-file-error')

    monkeypatch.setattr(rich_content_service, 'file_is_file', unreadable)
    assert load_question_rich_content(1, root=tmp_path) is None
    with pytest.raises(RichContentReadError) as error:
        load_question_rich_content(1, root=tmp_path, strict=True)
    assert str(error.value) == 'rich_content_unreadable'


def _extracted_with_floating_image(tmp_path: Path) -> ExtractedDocument:
    floating = tmp_path / "floating.png"
    floating.write_bytes(b"floating-image-bytes")
    inline = tmp_path / "inline.png"
    inline.write_bytes(b"inline-image-bytes")
    paragraphs = [
        {"text": "1. 第一题题干内容", "xml": "", "image_relationships": {}},
        {
            "text": f"[[IMAGE:{floating}]]",
            "xml": "<wp:anchor/>",
            "image_relationships": {},
        },
        {"text": "2. 第二题题干内容", "xml": "", "image_relationships": {}},
        {
            "text": f"第二题已有插图\n[[IMAGE:{inline}]]",
            "xml": "",
            "image_relationships": {},
        },
        {"text": "3. 第三题题干内容", "xml": "", "image_relationships": {}},
    ]
    return ExtractedDocument(
        source_file="tmp_asset_override_paper.docx",
        page_range="document",
        text="\n".join(str(item["text"]) for item in paragraphs),
        has_images=True,
        needs_image_review=True,
        image_paths=[str(floating), str(inline)],
        rich_paragraphs=paragraphs,
    )


def test_apply_asset_overrides_bind_moves_image_to_next_question(
    tmp_path: Path,
) -> None:
    extracted = _extracted_with_floating_image(tmp_path)
    floating, inline = extracted.image_paths
    digest = hashlib.sha256(Path(floating).read_bytes()).hexdigest()

    result = batch_importer.apply_asset_overrides(
        extracted,
        [
            {
                "sha256": digest,
                "action": "bind",
                "question_number": "2",
                "asset_kind": "question",
            }
        ],
    )

    assert [str(item["text"]) for item in result.rich_paragraphs] == [
        "1. 第一题题干内容",
        "2. 第二题题干内容",
        f"第二题已有插图\n[[IMAGE:{inline}]]",
        f"[[IMAGE:{floating}]]",
        "3. 第三题题干内容",
    ]
    parsed = parse_paper_text(
        result.text,
        source_file="tmp_asset_override_paper.docx",
        page_range="document",
    )
    by_number = {item.question_number: item for item in parsed.questions}
    assert floating not in by_number["1"].image_paths
    assert by_number["2"].image_paths == [inline, floating]
    assert floating not in by_number["3"].image_paths


def _paragraph_record_for_test(text: str, num_id: int, ilvl: int):
    """构造带 Word 自动编号（numPr）的段落。"""
    document = Document()
    paragraph = document.add_paragraph(text)
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - build the exact numbering structure.
    p_pr.append(
        parse_xml(
            f"<w:numPr {nsdecls('w')}>"
            f'<w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/>'
            f"</w:numPr>"
        )
    )
    return document, paragraph


def test_docx_importer_restarts_auto_number_per_word_list(tmp_path: Path) -> None:
    from question_bank.importers.docx_importer import _paragraph_record

    document, first = _paragraph_record_for_test("第一题题干", num_id=5, ilvl=0)
    _, second = _paragraph_record_for_test("第二题题干", num_id=5, ilvl=0)
    _, answer_first = _paragraph_record_for_test("第一题答案", num_id=6, ilvl=0)

    state: dict = {"counters": {}}
    record_first = _paragraph_record(first, document, tmp_path, {}, state)
    record_second = _paragraph_record(second, document, tmp_path, {}, state)
    record_answer = _paragraph_record(answer_first, document, tmp_path, {}, state)

    assert record_first["text"].startswith("1. ")
    assert record_second["text"].startswith("2. ")
    # 答案区另起编号列表时重新从 1 开始，而不是接着题目继续数
    assert record_answer["text"].startswith("1. ")
    assert record_first["numbering_level"] == 0


def _scanned_pdf_from_lines(lines: list[str]) -> bytes:
    """把文字渲染成图片再嵌入新 PDF，保证页面没有可抽取的文字层。"""
    text_document = fitz.open()
    page = text_document.new_page(width=600, height=800)
    for index, line in enumerate(lines):
        page.insert_text((40, 60 + index * 28), line, fontsize=12, fontname="china-s")
    pixmap = page.get_pixmap(alpha=False)
    png = pixmap.tobytes("png")
    text_document.close()
    document = fitz.open()
    scanned_page = document.new_page(width=pixmap.width, height=pixmap.height)
    scanned_page.insert_image(scanned_page.rect, stream=png)
    content = document.tobytes()
    document.close()
    return content


class _FakeOcrAdapter:
    """按页返回预设文本行的 OCR 替身，记录调用页码。"""

    def __init__(self, lines_by_page: dict[int, tuple[OcrLine, ...]]) -> None:
        self.lines_by_page = lines_by_page
        self.calls: list[int] = []

    def recognize(self, png_bytes: bytes, *, page_number: int) -> tuple[OcrLine, ...]:
        assert png_bytes.startswith(b"\x89PNG")
        self.calls.append(page_number)
        return self.lines_by_page.get(page_number, ())


def _ocr_line(text: str, y0: float, y1: float) -> OcrLine:
    return OcrLine(
        text=text,
        polygon=((0.05, y0), (0.95, y0), (0.95, y1), (0.05, y1)),
        confidence=0.9,
        engine_version="fake-ocr/1",
    )


_PAPER_LINES = [
    "1. 题干甲内容足够长",
    "2. 题干乙内容足够长",
    "参考答案",
    "1. 答案甲",
    "2. 答案乙",
]


def test_scanned_pdf_import_runs_local_ocr_and_marks_questions_for_review(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scan.pdf"
    source.write_bytes(_scanned_pdf_from_lines(_PAPER_LINES))
    workspace = tmp_path / "ws"
    ocr = _FakeOcrAdapter(
        {
            1: tuple(
                _ocr_line(text, 0.08 + index * 0.1, 0.16 + index * 0.1)
                for index, text in enumerate(_PAPER_LINES)
            )
        }
    )
    pipeline = QuestionDocumentPipeline(workspace, ocr_adapter=ocr)
    db_path = tmp_path / "data" / "databases" / "question_bank.db"

    result = batch_importer.import_scanned_papers(
        [ScannedPaper(source_file=str(source), file_type="pdf")],
        db_path,
        data_root=tmp_path / "data",
        document_pipeline=pipeline,
        # 固定跳过完整解析，避免本机模型文件齐备时跑真实 MinerU。
        full_parser=lambda path: None,
    )

    assert ocr.calls == [1]
    (file_result,) = result.files
    assert file_result.status == "needs_review"
    assert file_result.question_count == 2
    with connect(db_path) as conn:
        rows = conn.execute(
            "SELECT question_number, question_text, answer_text, needs_review "
            "FROM questions ORDER BY question_number"
        ).fetchall()
        (paper_status,) = conn.execute("SELECT import_status FROM papers").fetchone()
    assert paper_status == "needs_review"
    assert [row["question_number"] for row in rows] == ["1", "2"]
    assert all(int(row["needs_review"]) == 1 for row in rows)
    assert "题干甲" in rows[0]["question_text"]
    assert "答案甲" in rows[0]["answer_text"]
    assert "答案甲" not in rows[0]["question_text"]
    assert "答案乙" in rows[1]["answer_text"]
    # 文档管线快照落在注入的 workspace_root 下。
    assert list(workspace.rglob("snapshot.json"))
