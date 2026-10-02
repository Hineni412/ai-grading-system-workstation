from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from question_bank.personalized_papers.latex_render import (
    LatexRenderError,
    _math_latex,
    render_training_tex,
)


def test_math_delimiter_separators_and_upper_limits_preserve_structure():
    namespace = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
    pair = ET.fromstring(
        f'<m:d {namespace}><m:dPr><m:sepChr m:val=","/></m:dPr>'
        "<m:e><m:r><m:t>x</m:t></m:r></m:e>"
        "<m:e><m:r><m:t>y</m:t></m:r></m:e></m:d>"
    )
    assert _math_latex(pair) == r"\left(x,y\right)"
    limit = ET.fromstring(
        f"<m:limUpp {namespace}><m:limUppPr/>"
        "<m:e><m:r><m:t>AB</m:t></m:r></m:e>"
        "<m:lim><m:r><m:t>→</m:t></m:r></m:lim></m:limUpp>"
    )
    # Use supported observed text rather than accepting unknown symbols.
    limit.find(
        ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}lim/"
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}r/"
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"
    ).text = "x"
    assert _math_latex(limit) == r"\overset{x}{AB}"


def test_unknown_math_property_does_not_silently_change_the_equation(
    tmp_path: Path,
) -> None:
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath>'
        "<m:rad><m:radPr><m:unknownEffect/></m:radPr><m:deg/><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>"
        "</m:oMath></w:p>"
    )
    with pytest.raises(LatexRenderError, match="math property"):
        render_training_tex(
            _snapshot([{"text": "合成公式", "xml": xml}]), data_root=tmp_path
        )


def _asset(tmp_path: Path, name: str = "pic.png") -> Path:
    target = (
        tmp_path / "question_bank" / "personalized_papers" / "inst" / "assets" / name
    )
    target.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (80, 40), "white").save(target, format="PNG")
    return target


def _picture_paragraph_xml(asset: Path, tmp_path: Path) -> tuple[str, dict[str, str]]:
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run().add_picture(str(asset), width=Inches(1.25))
    blip = paragraph._p.xpath(".//a:blip")[0]
    relationship_id = str(
        blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
    )
    relative = asset.relative_to(tmp_path).as_posix()
    return paragraph._p.xml, {relationship_id: relative}


def _snapshot(blocks: list[dict[str, object]], **overrides) -> dict[str, object]:
    question = {
        "question_id": 1,
        "tagging_context": {
            "question_number": "1",
            "question_type": overrides.get("question_type", "解答题"),
            "question_text": overrides.get("question_text", "（12 分）合成题干"),
        },
        "rich_question_blocks": blocks,
    }
    return {
        "paper_instance_id": "c" * 64,
        "student": {
            "student_id": "1",
            "student_name": "合成学生",
            "class_id": "七年级一班",
        },
        "items": [
            {
                "task_item_code": "TASK-1",
                "question_id": 1,
                "question_snapshot": question,
            }
        ],
    }


def test_render_tex_fails_closed_on_missing_image_or_text_block(
    tmp_path: Path,
) -> None:
    xml, rels = _picture_paragraph_xml(_asset(tmp_path), tmp_path)
    broken = dict(rels)
    broken[next(iter(broken))] = (
        "question_bank/personalized_papers/inst/assets/missing.png"
    )
    with pytest.raises(LatexRenderError):
        render_training_tex(
            _snapshot([{"text": "", "xml": xml, "image_relationships": broken}]),
            data_root=tmp_path,
        )
    with pytest.raises(LatexRenderError):
        render_training_tex(
            _snapshot([{"text": "只有文本没有XML", "xml": ""}]),
            data_root=tmp_path,
        )
