from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from question_bank.personalized_papers.rendering import (
    inspect_docx,
    render_review_docx,
)


def test_review_docx_uses_editable_omml_for_supported_formula(tmp_path: Path) -> None:
    output = tmp_path / "review.docx"
    fallbacks = render_review_docx(
        {
            "paper_instance_id": "a" * 64,
            "series_version": 1,
            "student": {
                "student_id": "1",
                "student_code": "001",
                "student_name": "张三",
                "class_id": "1班",
            },
            "estimated_minutes": 10,
            "budget": {
                "estimated_total_tokens": 100,
                "context_window_tokens": 32768,
            },
            "items": [
                {
                    "task_item_code": "TASK-1",
                    "question_id": 1,
                    "question_snapshot": {
                        "question_id": 1,
                        "tagging_context": {
                            "question_number": "1",
                            "question_type": "解答题",
                            "question_text": "计算 $x^2+1$ 的值。",
                        },
                        "images": [],
                    },
                    "recommendation_snapshot": {"question_number": "1"},
                }
            ],
        },
        data_root=tmp_path,
        output_path=output,
    )

    assert fallbacks == ()
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    assert b"<m:oMath" in document_xml


def test_review_docx_prefers_frozen_word_blocks_and_keeps_original_image_size(
    tmp_path: Path,
) -> None:
    asset = (
        tmp_path
        / "question_bank"
        / "personalized_papers"
        / "paper"
        / "assets"
        / "figure.png"
    )
    asset.parent.mkdir(parents=True)
    image = Image.new("RGB", (300, 120), "white")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    asset.write_bytes(buffer.getvalue())

    source = Document()
    stem = source.add_paragraph()
    underlined = stem.add_run("1. 已知 x")
    underlined.underline = True
    superscript = stem.add_run("2")
    superscript.font.superscript = True
    stem.add_run("，完成下表。")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "富文本表格内容"
    picture = source.add_paragraph()
    picture.add_run().add_picture(str(asset), width=Inches(1.25))
    blip = picture._p.xpath(".//a:blip")[0]  # noqa: SLF001
    relationship_id = str(blip.get(qn("r:embed")))
    relative_asset = asset.relative_to(tmp_path).as_posix()

    output = tmp_path / "rich-review.docx"
    render_review_docx(
        {
            "paper_instance_id": "c" * 64,
            "series_version": 1,
            "student": {
                "student_id": "1",
                "student_name": "贾浩然",
                "class_id": "1班",
            },
            "estimated_minutes": 10,
            "budget": {
                "estimated_total_tokens": 100,
                "context_window_tokens": 32768,
            },
            "items": [
                {
                    "task_item_code": "TASK-RICH-1",
                    "question_id": 1,
                    "question_snapshot": {
                        "question_id": 1,
                        "tagging_context": {
                            "question_number": "1",
                            "question_type": "解答题",
                            "question_text": "不应显示的<sup>原始标签</sup>",
                        },
                        "rich_question_blocks": [
                            {"text": stem.text, "xml": stem._p.xml},  # noqa: SLF001
                            {"text": "富文本表格内容", "xml": table._tbl.xml},  # noqa: SLF001
                            {
                                "text": "",
                                "xml": picture._p.xml,  # noqa: SLF001
                                "image_relationships": {
                                    relationship_id: relative_asset,
                                },
                            },
                        ],
                        "images": [
                            {
                                "role": "question",
                                "mime_type": "image/png",
                                "sha256": "d" * 64,
                                "asset_path": relative_asset,
                            }
                        ],
                    },
                    "recommendation_snapshot": {"question_number": "1"},
                }
            ],
        },
        data_root=tmp_path,
        output_path=output,
    )

    reopened = Document(output)
    visible_text = "\n".join(
        [*(paragraph.text for paragraph in reopened.paragraphs)]
        + [
            paragraph.text
            for output_table in reopened.tables
            for row in output_table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
    )
    assert "不应显示" not in visible_text
    assert "<sup>" not in visible_text
    assert "富文本表格内容" in visible_text
    assert len(reopened.tables) == 2
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert b'<w:u w:val="single"' in document_xml
    assert b'<w:vertAlign w:val="superscript"' in document_xml
    assert b'wp:extent cx="1143000"' in document_xml
    assert len(media) == 1
    rich_stem = next(
        paragraph for paragraph in reopened.paragraphs if "已知 x" in paragraph.text
    )
    assert rich_stem.text.startswith("1. ")
    assert rich_stem.paragraph_format.line_spacing == 1.1
    assert rich_stem.paragraph_format.keep_with_next is True
    # 题末带图：图片排右格，作答区排左列（高度约为普通留白的一半）。
    answer_space = reopened.tables[-1]
    assert len(answer_space.columns) == 2
    rich_picture = next(
        paragraph
        for paragraph in answer_space.cell(0, 1).paragraphs
        if paragraph._p.xpath(".//a:blip")  # noqa: SLF001
    )
    assert rich_picture is not None
    assert 'w:val="single"' not in answer_space._tbl.xml  # noqa: SLF001
    inspect_docx(
        output,
        paper_instance_id="c" * 64,
        task_item_codes=("TASK-RICH-1",),
        question_snapshots=(
            {
                "tagging_context": {
                    "question_text": "不应显示的<sup>原始标签</sup>",
                },
                "rich_question_blocks": [
                    {"text": stem.text, "xml": stem._p.xml},  # noqa: SLF001
                    {"text": "富文本表格内容", "xml": table._tbl.xml},  # noqa: SLF001
                    {
                        "text": "",
                        "xml": picture._p.xml,  # noqa: SLF001
                        "image_relationships": {relationship_id: relative_asset},
                    },
                ],
            },
        ),
    )
