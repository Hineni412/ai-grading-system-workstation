from __future__ import annotations

import asyncio
import io
import json
import warnings
from pathlib import Path
from urllib.parse import quote

import fitz

warnings.filterwarnings(
    "ignore",
    message="Using `httpx` with `starlette.testclient` is deprecated.*",
)

from fastapi.testclient import TestClient
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.api.app import create_app
from backend.api.dependencies import (
    get_config_source_service,
    get_grading_db,
    get_job_manager,
    get_upload_config_dir,
)
from backend.config_workspace.sources import ConfigSourceService
from backend.jobs.manager import JobManager
from backend.jobs.store import JobStore
from backend.api.routers.config import upload_config_source
from db_manager import DBManager


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Synthetic private exam text\n1. Compute 1 + 1.\nA. 1  B. 2\nAnswer\n1. B",
        fontsize=12,
    )
    payload = document.tobytes()
    document.close()
    return payload


def _docx_with_inline_main_question_marker_bytes() -> bytes:
    document = Document()
    document.add_paragraph("10. Tenth question.")
    document.add_paragraph(
        "Continuation of question ten. 11. Eleventh question."
    )
    document.add_paragraph("12. Twelfth question.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_inline_question_after_breaks_bytes() -> bytes:
    document = Document()
    document.add_paragraph("10. Tenth question.")
    paragraph = document.add_paragraph("Continuation of question ten.")
    breaks = paragraph.add_run()
    for _index in range(4):
        breaks.add_break()
    paragraph.add_run("11. Eleventh question.")
    document.add_paragraph("12. Twelfth question.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_inline_question_media_bytes() -> bytes:
    image = io.BytesIO()
    Image.new("RGB", (24, 18), "navy").save(image, format="PNG")
    document = Document()
    document.add_paragraph("10. Tenth question.")
    paragraph = document.add_paragraph(
        "Continuation of question ten. 11. Eleventh question x"
    )
    exponent = paragraph.add_run("2")
    exponent.font.superscript = True
    paragraph.add_run().add_picture(io.BytesIO(image.getvalue()), width=Inches(0.25))
    document.add_paragraph("12. Twelfth question.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_image_before_inline_question_marker_bytes() -> bytes:
    image = io.BytesIO()
    Image.new("RGB", (24, 18), "maroon").save(image, format="PNG")
    document = Document()
    document.add_paragraph("10. Tenth question.")
    paragraph = document.add_paragraph("Continuation of question ten.")
    paragraph.add_run().add_picture(io.BytesIO(image.getvalue()), width=Inches(0.25))
    paragraph.add_run(" 11. Eleventh question.")
    document.add_paragraph("12. Twelfth question.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _client(tmp_path: Path) -> tuple[TestClient, DBManager, Path]:
    db = DBManager(tmp_path / "grading.db")
    db.initialize()
    upload_root = tmp_path / "uploaded"
    app = create_app()
    manager = JobManager(JobStore(db.db_path), max_workers=1)
    app.dependency_overrides[get_grading_db] = lambda: db
    app.dependency_overrides[get_job_manager] = lambda: manager
    app.dependency_overrides[get_upload_config_dir] = lambda: upload_root
    return TestClient(app), db, upload_root


def _session(db: DBManager, tmp_path: Path) -> int:
    rubric = tmp_path / "rubric.json"
    answer = tmp_path / "answer.json"
    rubric.write_text("{}", encoding="utf-8")
    answer.write_text("{}", encoding="utf-8")
    return db.create_grading_session("Config Source", str(rubric), str(answer))


def _upload(
    client: TestClient,
    session_id: int,
    *,
    filename: str = "数学卷.pdf",
    request_token: str | None = None,
):
    headers = {
        "content-type": "application/octet-stream",
        "x-upload-filename": quote(filename),
    }
    if request_token is not None:
        headers["x-client-request-token"] = request_token
    return client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes(),
        headers=headers,
    )


def test_abandon_unseen_upload_token_blocks_a_late_request(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "9" * 32

    abandoned = client.post(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}/abandon"
    )
    replay = client.post(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}/abandon"
    )
    late_upload = _upload(client, session_id, request_token=token)
    lookup = client.get(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}"
    )

    assert abandoned.status_code == 200
    assert abandoned.json() == {"status": "abandoned"}
    assert replay.status_code == 200
    assert late_upload.status_code == 409
    assert lookup.status_code == 404


def test_abandon_upload_token_rejects_an_already_received_request(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "8" * 32

    uploaded = _upload(client, session_id, request_token=token)
    abandoned = client.post(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}/abandon"
    )
    lookup = client.get(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}"
    )

    assert uploaded.status_code == 201
    assert abandoned.status_code == 409
    assert lookup.status_code == 200
    assert lookup.json()["source"] == uploaded.json()


def test_upload_and_restart_get_return_only_bounded_public_projection(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)

    uploaded = _upload(client, session_id, filename="../数学卷.pdf")

    assert uploaded.status_code == 201
    body = uploaded.json()
    assert body["safe_filename"] == "数学卷.pdf"
    assert body["parse_state"] == "ready"
    assert len(body["sha256_prefix"]) == 12
    assert len(body["source_id"]) == 32
    assert len(body["source_revision"]) == 64
    encoded = json.dumps(body, ensure_ascii=False)
    assert "Synthetic private exam text" not in encoded
    assert str(tmp_path) not in encoded
    assert "base64" not in encoded.casefold()

    reloaded = client.get(
        f"/api/sessions/{session_id}/config/sources/{body['source_id']}"
    )
    assert reloaded.status_code == 200
    assert reloaded.json() == body


def test_docx_upload_splits_consecutive_main_question_marker_mid_paragraph(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)

    uploaded = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_docx_with_inline_main_question_marker_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("inline-question.docx"),
        },
    )

    assert uploaded.status_code == 201
    questions = uploaded.json()["questions"]
    assert [question["question_id"] for question in questions] == [
        "Q10",
        "Q11",
        "Q12",
    ]
    previews = {
        question["question_id"]: question["question_preview"]
        for question in questions
    }
    assert "Eleventh question" not in previews["Q10"]
    assert "Eleventh question" in previews["Q11"]


def test_docx_upload_splits_inline_question_after_same_paragraph_breaks(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)

    uploaded = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_docx_with_inline_question_after_breaks_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("inline-question-after-breaks.docx"),
        },
    )

    assert uploaded.status_code == 201
    questions = uploaded.json()["questions"]
    assert [question["question_id"] for question in questions] == [
        "Q10",
        "Q11",
        "Q12",
    ]
    previews = {
        question["question_id"]: question["question_preview"]
        for question in questions
    }
    assert "Eleventh question" not in previews["Q10"]
    assert "Eleventh question" in previews["Q11"]


def test_docx_inline_question_keeps_following_formula_and_image_with_new_block(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)

    uploaded = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_docx_with_inline_question_media_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("inline-question-media.docx"),
        },
    )

    assert uploaded.status_code == 201
    questions = {
        question["question_id"]: question
        for question in uploaded.json()["questions"]
    }
    assert questions["Q10"]["has_question_asset"] is False
    assert questions["Q11"]["has_question_asset"] is True
    assert "Eleventh question x" in questions["Q11"]["question_preview"]
    assert questions["Q11"]["question_preview"].endswith("2")


def test_docx_inline_question_keeps_preceding_image_with_current_block(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)

    uploaded = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_docx_with_image_before_inline_question_marker_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("inline-question-preceding-image.docx"),
        },
    )

    assert uploaded.status_code == 201
    questions = {
        question["question_id"]: question
        for question in uploaded.json()["questions"]
    }
    assert questions["Q10"]["has_question_asset"] is True
    assert questions["Q11"]["has_question_asset"] is False


def test_active_source_endpoint_returns_replacement_without_old_source_id(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    first = _upload(client, session_id, filename="first.pdf").json()
    second = _upload(client, session_id, filename="second.pdf").json()

    active = client.get(f"/api/sessions/{session_id}/config/sources/active")

    assert active.status_code == 200
    assert active.json() == second
    assert active.json()["source_id"] != first["source_id"]


def test_active_config_job_rejects_source_replacement(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    first = _upload(client, session_id, filename="first.pdf").json()
    manager = client.app.dependency_overrides[get_job_manager]()
    job = manager.store.create_claimed_config_job(
        {"session_id": session_id, "mode": "generate"}
    )

    response = _upload(client, session_id, filename="replacement.pdf")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "config_generation_in_progress"
    active = client.get(f"/api/sessions/{session_id}/config/sources/active")
    assert active.status_code == 200
    assert active.json() == first
    manager.store.finish(job.id, "failed", error="test cleanup")


def test_upload_submission_token_is_queryable_and_repeated_post_is_rejected(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "a" * 32
    headers = {
        "content-type": "application/octet-stream",
        "x-upload-filename": quote("token-paper.pdf"),
        "x-client-request-token": token,
    }

    first = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes(),
        headers=headers,
    )
    replay = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes(),
        headers=headers,
    )
    queried = client.get(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}"
    )

    assert first.status_code == 201
    assert replay.status_code == 409
    assert replay.json()["error"]["code"] == "config_source_upload_already_submitted"
    assert first.json()["source_id"] == token
    assert queried.status_code == 200
    assert queried.json() == {"status": "succeeded", "source": first.json()}


def test_upload_submission_token_never_accepts_different_same_length_content(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "9" * 32
    content = _pdf_bytes()
    headers = {
        "content-type": "application/octet-stream",
        "x-upload-filename": quote("same-metadata.pdf"),
        "x-client-request-token": token,
    }

    first = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=content,
        headers=headers,
    )
    changed = bytearray(content)
    changed[-1] ^= 1
    replay = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=bytes(changed),
        headers=headers,
    )

    assert first.status_code == 201
    assert replay.status_code == 409
    assert replay.json()["error"]["code"] == "config_source_upload_already_submitted"


def test_upload_rejects_a_deleted_session_without_creating_source_files(
    tmp_path: Path,
) -> None:
    client, db, upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    db.soft_delete_grading_session(session_id)

    response = _upload(client, session_id)

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "session_not_found"
    assert not (upload_root / "config_sources" / f"session-{session_id}").exists()


def test_upload_submission_token_rejects_different_request_metadata(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "b" * 32
    first = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("first.pdf"),
            "x-client-request-token": token,
        },
    )

    conflict = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes() + b"different-size",
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("second.pdf"),
            "x-client-request-token": token,
        },
    )

    assert first.status_code == 201
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "config_source_submission_conflict"


def test_upload_submission_reports_replaced_instead_of_returning_stale_source(
    tmp_path: Path,
) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "c" * 32
    first = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=_pdf_bytes(),
        headers={
            "content-type": "application/octet-stream",
            "x-upload-filename": quote("first.pdf"),
            "x-client-request-token": token,
        },
    )
    assert first.status_code == 201
    assert _upload(client, session_id, filename="replacement.pdf").status_code == 201

    queried = client.get(
        f"/api/sessions/{session_id}/config/sources/submissions/{token}"
    )

    assert queried.status_code == 200
    assert queried.json() == {"status": "replaced", "source": None}


def test_cancelled_upload_marks_exact_submission_failed(tmp_path: Path) -> None:
    client, db, upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "d" * 32

    class CancelledService(ConfigSourceService):
        async def stage_and_parse(self, **_kwargs):
            raise asyncio.CancelledError()

    class FakeRequest:
        headers = {
            "content-length": "4",
            "x-upload-filename": quote("cancelled.pdf"),
            "x-client-request-token": token,
        }

        async def stream(self):
            yield b"%PDF"

    service = CancelledService(upload_root)
    manager = client.app.dependency_overrides[get_job_manager]()
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(upload_config_source(session_id, FakeRequest(), db, service, manager))

    assert service.submission_public(
        session_id=session_id,
        request_token=token,
    ) == {"status": "failed", "source": None}


def test_submission_recovers_succeeded_from_active_source_before_success_marker(
    tmp_path: Path,
) -> None:
    _client_value, db, upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    token = "e" * 32
    payload = _pdf_bytes()
    service = ConfigSourceService(upload_root)
    assert service.begin_submission(
        session_id=session_id,
        request_token=token,
        filename="crash-window.pdf",
        content_length=len(payload),
    ) == "started"

    async def chunks():
        yield payload

    record = asyncio.run(service.stage_and_parse(
        session_id=session_id,
        filename="crash-window.pdf",
        chunks=chunks(),
        source_id=token,
    ))

    assert service.submission_public(
        session_id=session_id,
        request_token=token,
    ) == {"status": "succeeded", "source": record.public_snapshot()}


def test_public_get_does_not_read_or_encode_private_images(
    tmp_path: Path,
    monkeypatch,
) -> None:
    import backend.config_workspace.sources as sources_module

    client, db, upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    source_service = ConfigSourceService(upload_root)
    client.app.dependency_overrides[get_config_source_service] = lambda: source_service
    uploaded = _upload(client, session_id).json()
    original_read_bytes = source_service._files.read_bytes

    def reject_private_image_read(path: Path, **kwargs: object) -> bytes:
        if path.name.startswith(("asset-", "whole-page-")):
            raise AssertionError("public metadata load read a private image")
        return original_read_bytes(path, **kwargs)

    def reject_encoding(_content: bytes) -> bytes:
        raise AssertionError("public metadata load encoded a private image")

    monkeypatch.setattr(source_service._files, "read_bytes", reject_private_image_read)
    monkeypatch.setattr(sources_module.base64, "b64encode", reject_encoding)

    response = client.get(
        f"/api/sessions/{session_id}/config/sources/{uploaded['source_id']}"
    )

    assert response.status_code == 200
    assert response.json() == uploaded


def test_upload_rejects_missing_session_without_creating_source(tmp_path: Path) -> None:
    client, _db, upload_root = _client(tmp_path)

    response = _upload(client, 404)

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "session_not_found"
    assert list(upload_root.rglob("manifest.json")) == []


def test_upload_maps_size_type_and_parser_errors_without_echoing_input(tmp_path: Path) -> None:
    client, db, upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    client.app.dependency_overrides[get_config_source_service] = lambda: ConfigSourceService(
        upload_root,
        max_upload_bytes=4,
    )
    too_large = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=b"12345",
        headers={"x-upload-filename": quote("C:/private/secret.pdf")},
    )
    assert too_large.status_code == 413
    assert too_large.json()["error"]["code"] == "config_source_too_large"
    assert "secret" not in too_large.text

    client.app.dependency_overrides[get_config_source_service] = (
        lambda: ConfigSourceService(upload_root)
    )
    unsupported = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=b"private",
        headers={"x-upload-filename": quote("C:/private/secret.txt")},
    )
    assert unsupported.status_code == 415
    assert unsupported.json()["error"]["code"] == "config_source_type_unsupported"
    assert "secret" not in unsupported.text

    invalid = client.post(
        f"/api/sessions/{session_id}/config/sources",
        content=b"%PDF-private-parser-detail",
        headers={"x-upload-filename": quote("C:/private/secret.pdf")},
    )
    assert invalid.status_code == 422
    assert invalid.json()["error"]["code"] == "config_source_invalid"
    assert "secret" not in invalid.text
    assert "parser" not in invalid.text


def test_controlled_question_and_answer_assets_are_private_no_store(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    uploaded = _upload(client, session_id).json()
    question = uploaded["questions"][0]

    for kind in ("question", "answer"):
        response = client.get(
            f"/api/sessions/{session_id}/config/sources/{uploaded['source_id']}"
            f"/questions/{question['question_id']}/assets/{kind}"
        )
        assert response.status_code == 200
        assert response.headers["cache-control"] == "private, no-store"
        assert response.headers["content-type"].startswith("image/")
        assert response.content

    missing = client.get(
        f"/api/sessions/{session_id}/config/sources/{uploaded['source_id']}"
        f"/questions/{question['question_id']}/assets/private-path"
    )
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "config_asset_not_found"


def test_replaced_source_get_fails_closed_but_new_source_remains_available(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    session_id = _session(db, tmp_path)
    first = _upload(client, session_id, filename="first.pdf").json()
    second = _upload(client, session_id, filename="second.pdf").json()

    stale = client.get(
        f"/api/sessions/{session_id}/config/sources/{first['source_id']}"
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "config_source_changed"
    current = client.get(
        f"/api/sessions/{session_id}/config/sources/{second['source_id']}"
    )
    assert current.status_code == 200


def test_source_ids_cannot_select_another_session_or_path(tmp_path: Path) -> None:
    client, db, _upload_root = _client(tmp_path)
    first_session = _session(db, tmp_path)
    second_session = db.create_grading_session(
        "Other",
        str(tmp_path / "rubric.json"),
        str(tmp_path / "answer.json"),
    )
    source = _upload(client, first_session).json()

    mismatch = client.get(
        f"/api/sessions/{second_session}/config/sources/{source['source_id']}"
    )
    assert mismatch.status_code == 404
    assert mismatch.json()["error"]["code"] == "config_source_not_found"
    invalid_id = client.get(
        f"/api/sessions/{first_session}/config/sources/{'A' * 32}"
    )
    assert invalid_id.status_code == 404
    assert invalid_id.json()["error"]["code"] == "config_source_not_found"
